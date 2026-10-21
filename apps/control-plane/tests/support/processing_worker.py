# Copyright Thales 2026
# SPDX-License-Identifier: Apache-2.0
"""Existing KF workers on disposable stores; only the model provider is simulated."""

import asyncio
import json
import os
import sys
import threading
from contextlib import AsyncExitStack
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit


class ModelHandler(BaseHTTPRequestHandler):
    failure_file: Path

    def log_message(self, *_):
        pass

    def do_POST(self):
        request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        if self.failure_file.exists():
            self.send_response(503)
            self.send_header("Content-Length", "0")
            self.end_headers()
            return
        if self.path.endswith("/embeddings"):
            values = request["input"]
            if isinstance(values, str):
                values = [values]
            body = {
                "object": "list",
                "data": [
                    {
                        "object": "embedding",
                        "index": i,
                        "embedding": [0.1, 0.2, 0.3, 0.4],
                    }
                    for i in range(len(values))
                ],
                "model": "fixture",
                "usage": {"prompt_tokens": len(values), "total_tokens": len(values)},
            }
        else:
            body = {
                "id": "fixture",
                "object": "chat.completion",
                "created": 1,
                "model": "fixture",
                "choices": [
                    {
                        "index": 0,
                        "message": {
                            "role": "assistant",
                            "content": "A short document about manager ingestion.",
                        },
                        "finish_reason": "stop",
                    }
                ],
                "usage": {
                    "prompt_tokens": 1,
                    "completion_tokens": 1,
                    "total_tokens": 2,
                },
            }
        encoded = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)


def create_inputs(folder):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image

    image = folder / "diagram.png"
    Image.new("RGB", (100, 60), (40, 100, 180)).save(image)
    pdf = PdfWriter()
    page = pdf.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): pdf._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 50 750 Td (Manager PDF ingestion produces inspectable chunks.) Tj ET"
    )
    page[NameObject("/Contents")] = pdf._add_object(stream)
    pdf.write(folder / "report.pdf")
    document = Document()
    document.add_paragraph("Manager Office ingestion preserves the existing pipeline.")
    document.add_picture(str(image))
    document.save(folder / "report.docx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Manager presentation ingestion"
    slide.placeholders[1].text = "The existing worker extracts this content."
    slide.shapes.add_picture(str(image), Inches(1), Inches(4))
    deck.save(folder / "report.pptx")
    workbook = Workbook()
    workbook.active.append(["name", "value"])
    workbook.active.append(["one", 1])
    workbook.active.append(["two", 2])
    workbook.save(folder / "dataset.xlsx")
    import subprocess

    for original, target in [
        ("report.docx", "doc:MS Word 97"),
        ("report.docx", "odt:writer8"),
        ("report.pptx", "ppt:MS PowerPoint 97"),
    ]:
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation={(folder / 'office-profile').as_uri()}",
                "--headless",
                "--convert-to",
                target,
                "--outdir",
                str(folder),
                str(folder / original),
            ],
            check=True,
            capture_output=True,
            timeout=30,
        )
        assert (folder / ("report." + target.split(":")[0])).exists()


async def main():
    fixture = json.loads(Path(sys.argv[1]).read_text())
    folder = Path(sys.argv[1]).parent
    if not (folder / "report.pdf").exists():
        create_inputs(folder)
    import shutil
    import shlex

    real_soffice = shutil.which("soffice")
    assert real_soffice
    wrapper = folder / "bin"
    wrapper.mkdir(exist_ok=True)
    executable = wrapper / "soffice"
    executable.write_text(
        "#!/bin/sh\nif [ -f "
        + shlex.quote(str(folder / "fail-render"))
        + " ]; then exit 1; fi\nexec "
        + shlex.quote(real_soffice)
        + ' "$@"\n'
    )
    executable.chmod(0o700)
    os.environ["PATH"] = str(wrapper) + os.pathsep + os.environ["PATH"]
    ModelHandler.failure_file = folder / "fail-model"
    os.environ.update(
        OPENSEARCH_PASSWORD="fixture",
        OPENAI_API_KEY="fixture",
        MANAGER_WORKER_CLIENT_SECRET="fixture",
        MANAGER_DERIVED_ACCESS_KEY="fixture-access",
        MANAGER_DERIVED_SECRET_KEY="fixture-secret",
    )
    from knowledge_flow_backend.common.config_loader import load_configuration
    from knowledge_flow_backend.common.structures import (
        Configuration,
        IngestionWorkerRole,
    )
    from knowledge_flow_backend.application_context import ApplicationContext
    from knowledge_flow_backend.features.scheduler.worker import _build_worker
    from fred_core.kpi.noop_kpi_writer import NoOpKPIWriter
    from temporalio.client import Client

    model = ThreadingHTTPServer(("127.0.0.1", 0), ModelHandler)
    threading.Thread(target=model.serve_forever, daemon=True).start()
    config = load_configuration().model_dump(mode="json")
    database = urlsplit(fixture["postgres"])
    config["storage"]["postgres"] = {
        "host": database.hostname,
        "port": database.port,
        "database": database.path.lstrip("/"),
        "username": database.username,
        "password": database.password,
    }
    config["storage"]["vector_store"] = {
        "type": "pgvector",
        "collection_name": "legacy_fixture",
    }
    config["manager_processing"] = {
        "enabled": True,
        "url": fixture["manager"],
        "issuer": fixture["issuer"],
        "audience": "app",
        "client_id": "manager-worker",
        "endpoint": fixture["object_endpoint"],
        "bucket": "originals",
        "allow_insecure": True,
        "vector_store": {"type": "pgvector", "collection_name": "manager_fixture"},
    }
    for name in ("embedding_model", "chat_model", "vision_model"):
        config[name] = {
            "provider": "openai",
            "name": "fixture",
            "settings": {
                "base_url": f"http://127.0.0.1:{model.server_port}/v1",
                "max_retries": 0,
                "check_embedding_ctx_length": False,
            }
            if name == "embedding_model"
            else {
                "base_url": f"http://127.0.0.1:{model.server_port}/v1",
                "max_retries": 0,
            },
        }
    config["scheduler"]["temporal"]["host"] = fixture["temporal"]
    config["scheduler"]["temporal"]["task_queue"] = "ingestion"
    config["observability"]["kpi"] = {"enabled": False}
    config["processing"]["profiles"] = json.loads(
        Path(fixture["contract"]).read_text()
    )["processor_profiles"]
    config_file = folder / "worker.json"
    config_file.write_text(json.dumps(config))
    os.environ["CONFIG_FILE"] = str(config_file)
    context = ApplicationContext(Configuration.model_validate(config))
    context._kpi_writer = NoOpKPIWriter()

    from langchain_core.documents import Document

    legacy_uid = "00000000-0000-0000-0000-000000000001"
    legacy = context.get_create_vector_store(context.get_embedder())
    if not legacy.list_document_uids():
        legacy.add_documents(
            [
                Document(
                    page_content="Legacy corpus stays isolated.",
                    metadata={
                        "document_uid": legacy_uid,
                        "document_name": "legacy.txt",
                    },
                )
            ],
            ids=[legacy_uid],
        )
    assert legacy.list_document_uids() == [legacy_uid]

    # Fail if the adapter reaches a legacy resource store.
    def guarded(original):
        def invoke(self, *args, **kwargs):
            if self is context:
                raise AssertionError(
                    f"Manager ingestion reached legacy {original.__name__}"
                )
            return original(self, *args, **kwargs)

        return invoke

    for method in ("get_metadata_store", "get_content_store", "get_tag_store"):
        setattr(
            ApplicationContext, method, guarded(getattr(ApplicationContext, method))
        )
    client = None
    for _ in range(60):
        try:
            client = await Client.connect(fixture["temporal"])
            break
        except Exception:
            await asyncio.sleep(0.5)
    assert client is not None
    (folder / "ready").write_text("ready")
    while not (folder / "release-worker").exists():
        await asyncio.sleep(0.1)
    async with AsyncExitStack() as stack:
        for role in IngestionWorkerRole:
            await stack.enter_async_context(
                _build_worker(
                    client=client,
                    config=context.configuration.scheduler.temporal,
                    role=role,
                    workflow_task_concurrency=4,
                    activity_concurrency=2,
                )
            )
        while not (folder / "stop").exists():
            await asyncio.sleep(0.2)
    manager_vectors = context.get_manager_vector_store()
    manager_ids = set(manager_vectors.list_document_uids())
    assert manager_ids and legacy_uid not in manager_ids
    assert legacy.list_document_uids() == [legacy_uid]
    legacy.delete_vectors_for_document(document_uid=next(iter(manager_ids)))
    assert set(manager_vectors.list_document_uids()) == manager_ids
    model.shutdown()


if __name__ == "__main__":
    asyncio.run(main())
