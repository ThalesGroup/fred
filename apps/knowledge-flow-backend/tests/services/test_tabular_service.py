# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

from __future__ import annotations

import resource
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import duckdb
import pytest
from fred_core import KeycloakUser
from fred_core.documents.document_structures import (
    DocumentMetadata,
    FileInfo,
    FileType,
    Identity,
    ProcessingStage,
    ProcessingStatus,
    SourceInfo,
    SourceType,
    Tagging,
)

from knowledge_flow_backend.application_context import ApplicationContext
from knowledge_flow_backend.core.processors.output.tabular_processor.tabular_processor import TabularProcessor
from knowledge_flow_backend.core.stores.content.filesystem_content_store import FileSystemContentStore
from knowledge_flow_backend.features.tabular.artifacts import (
    TABULAR_EXTENSION_KEY,
    TABULAR_MULTI_EXTENSION_KEY,
    document_artifact_prefix,
    duckdb_dtype_to_literal,
    max_categories,
    read_tabular_artifact,
)
from knowledge_flow_backend.features.tabular.execution import open_duckdb_connection
from knowledge_flow_backend.features.tabular.service import TabularDatasetReadError
from knowledge_flow_backend.features.tabular.structures import TabularQueryRequest
from tests.services.tabular_engine import TabularEngineHarness


def _user(uid: str = "u-1") -> KeycloakUser:
    return KeycloakUser(
        uid=uid,
        username=uid,
        email=f"{uid}@example.com",
        roles=["admin"],
    )


def _metadata(
    *,
    document_uid: str,
    file_name: str,
    tag_ids: list[str] | None = None,
    tag_names: list[str] | None = None,
) -> DocumentMetadata:
    return DocumentMetadata(
        identity=Identity(document_name=file_name, document_uid=document_uid, title=file_name),
        source=SourceInfo(source_type=SourceType.PUSH, source_tag="uploads"),
        file=FileInfo(file_type=FileType.CSV, mime_type="text/csv"),
        tags=Tagging(tag_ids=_library_of(document_uid) if tag_ids is None else tag_ids, tag_names=tag_names or []),
    )


def _library_of(document_uid: str) -> list[str]:
    """A folder per test document: ReBAC grants documents through folders only."""
    return [f"lib-{document_uid}"]


async def _ingest_csv(
    *,
    tmp_path: Path,
    metadata_store,
    document_uid: str,
    file_name: str,
    content: str,
    tag_ids: list[str] | None = None,
    tag_names: list[str] | None = None,
) -> DocumentMetadata:
    csv_path = tmp_path / file_name
    csv_path.write_text(content, encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(
        document_uid=document_uid,
        file_name=file_name,
        tag_ids=tag_ids,
        tag_names=tag_names,
    )
    processed_metadata = processor.process(str(csv_path), metadata)
    await ApplicationContext.get_instance().get_metadata_store().save_metadata(processed_metadata)
    return processed_metadata


async def _ingest_excel_workbook(*, tmp_path: Path, document_uid: str, extra_tables: int = 0) -> tuple[DocumentMetadata, Path]:
    """Run spreadsheet input + output for Ventes, Cibles and optional extra sheets.

    Returns the saved metadata and the local output directory (output.md +
    tables.json), which callers can upload as document preview content.
    """
    from openpyxl import Workbook

    from knowledge_flow_backend.core.processors.input.excel_processor.excel_processor import ExcelProcessor
    from knowledge_flow_backend.core.processors.output.excel_processor.excel_table_registration_processor import ExcelTableRegistrationProcessor

    workbook_path = tmp_path / f"{document_uid}.xlsx"
    workbook = Workbook()
    ventes = workbook.active
    assert ventes is not None
    ventes.title = "Ventes"
    for row in [("city", "amount"), ("Paris", 10), ("Lyon", 20)]:
        ventes.append(row)
    cibles = workbook.create_sheet("Cibles")
    for row in [("city", "target"), ("Paris", 15), ("Lyon", 25)]:
        cibles.append(row)
    for index in range(extra_tables):
        extra = workbook.create_sheet(f"Extra {index + 1}")
        for row in [("city", "amount"), ("Paris", index), ("Lyon", index + 10)]:
            extra.append(row)
    workbook.save(str(workbook_path))

    metadata = DocumentMetadata(
        identity=Identity(document_name=f"{document_uid}.xlsx", document_uid=document_uid, title=document_uid),
        source=SourceInfo(source_type=SourceType.PUSH, source_tag="uploads"),
        file=FileInfo(file_type=FileType.XLSX),
        tags=Tagging(tag_ids=_library_of(document_uid), tag_names=[]),
    )

    output_dir = tmp_path / f"{document_uid}-output"
    processor = ExcelProcessor()
    # This helper builds a deterministic workbook with openpyxl. Formula
    # recalculation through LibreOffice is covered by the integration suite.
    processor.recalc = False
    processor.convert_file_to_markdown(workbook_path, output_dir, document_uid)
    metadata = ExcelTableRegistrationProcessor().process(str(output_dir / "output.md"), metadata)
    await ApplicationContext.get_instance().get_metadata_store().save_metadata(metadata)
    return metadata, output_dir


@pytest.mark.asyncio
async def test_tabular_processor_stores_distinct_artifacts_for_each_revision(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-1",
        file_name="sales.csv",
        content="city,amount,created_at\nParis,10,2024-01-01\nLyon,20,2024-01-02\n",
    )
    artifact = read_tabular_artifact(metadata)

    assert artifact is not None
    assert artifact.dataset_uid == "doc-1"
    assert artifact.row_count == 2
    assert [column.name for column in artifact.columns] == ["city", "amount", "created_at"]
    assert metadata.extensions is not None
    assert TABULAR_EXTENSION_KEY in metadata.extensions

    tabular_config = ApplicationContext.get_instance().get_config().storage.tabular_store
    object_prefix = document_artifact_prefix(
        artifacts_prefix=tabular_config.artifacts_prefix,
        document_uid="doc-1",
    )
    stored_objects = content_store.list_objects(object_prefix)
    assert len(stored_objects) == 1
    assert stored_objects[0].key == artifact.object_key

    updated_csv = tmp_path / "sales.csv"
    updated_csv.write_text("city,amount,created_at\nParis,30,2024-02-01\n", encoding="utf-8")

    processor = TabularProcessor()
    updated_metadata = processor.process(str(updated_csv), metadata)
    updated_artifact = read_tabular_artifact(updated_metadata)
    assert updated_artifact is not None
    assert updated_artifact.object_key != artifact.object_key

    stored_objects = content_store.list_objects(object_prefix)
    assert {stored_object.key for stored_object in stored_objects} == {
        artifact.object_key,
        updated_artifact.object_key,
    }

    await ApplicationContext.get_instance().get_metadata_store().save_metadata(updated_metadata)

    # Metadata selects the new revision. Artifact retirement belongs to the
    # manager processing lifecycle, not this processor's metadata fixture.
    saved = await ApplicationContext.get_instance().get_metadata_store().get_metadata_by_uid("doc-1")
    assert saved is not None
    assert read_tabular_artifact(saved) == updated_artifact
    assert updated_artifact.row_count == 1
    assert updated_artifact.source_revision != artifact.source_revision


@pytest.mark.asyncio
async def test_tabular_processor_converts_csv_without_pandas_read_csv(tmp_path, monkeypatch):
    import pandas as pd

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    monkeypatch.setattr(pd, "read_csv", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("pandas path should not be used")))

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-no-pandas", file_name="sales.csv")

    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    assert artifact.row_count == 2
    assert [column.name for column in artifact.columns] == ["city", "amount"]
    assert [column.dtype for column in artifact.columns] == ["string", "integer"]
    assert artifact.columns[0].is_categorical is True
    assert artifact.columns[0].sample_values is None
    assert artifact.columns[1].min_value is None
    assert artifact.columns[1].max_value is None
    assert processed_metadata.processing.stages[ProcessingStage.PREVIEW_READY] == ProcessingStatus.DONE
    assert processed_metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE


@pytest.mark.asyncio
async def test_tabular_processor_keeps_mixed_numeric_and_text_column_as_string(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    csv_path = tmp_path / "mixed-values.csv"
    with csv_path.open("w", encoding="utf-8") as file_handle:
        file_handle.write("id,label,value\n")
        for index in range(30_000):
            file_handle.write(f"{index},row,{index / 10}\n")
        file_handle.write("30000,tail,xxxxxxxxxxxxxxx\n")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-mixed-types", file_name="mixed-values.csv")

    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    assert artifact.row_count == 30_001
    assert [column.name for column in artifact.columns] == ["id", "label", "value"]
    assert [column.dtype for column in artifact.columns] == ["integer", "string", "string"]


@pytest.mark.asyncio
async def test_tabular_processor_types_float_without_storing_bounds(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    csv_path = tmp_path / "scores.csv"
    csv_path.write_text("name,score\nfirst,-2.25\nsecond,1.5\nthird,\n", encoding="utf-8")
    metadata = _metadata(document_uid="doc-scores", file_name="scores.csv")
    artifact = read_tabular_artifact(TabularProcessor().process(str(csv_path), metadata))

    assert artifact is not None
    score = next(column for column in artifact.columns if column.name == "score")
    assert score.dtype == "float"
    assert score.min_value is None
    assert score.max_value is None


@pytest.mark.parametrize(
    ("row_count", "expected"),
    [(0, 2), (10, 2), (11, 2), (20, 3), (30, 4), (100, 10), (100_000, 256)],
)
def test_max_categories_uses_sublinear_growth_and_cap(row_count, expected):
    assert max_categories(row_count) == expected


def test_decimal_parquet_schema_is_reported_as_float():
    assert duckdb_dtype_to_literal("DECIMAL(10,2)") == "float"


@pytest.mark.asyncio
async def test_tabular_processor_classifies_strings_without_storing_category_values(tmp_path):
    """
    Ingestion classifies a categorical string column; the description tool
    reads its exact distinct values from Parquet later.
    """
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    rows = "\n".join(f"{i},{'CRITICAL' if i % 2 == 0 else 'LOW'}" for i in range(10))
    csv_path = tmp_path / "scan.csv"
    csv_path.write_text(f"id,severity\n{rows}\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-severity", file_name="scan.csv")
    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    severity_column = next(column for column in artifact.columns if column.name == "severity")
    assert severity_column.is_categorical is True
    assert severity_column.has_two_values is True
    assert severity_column.sample_values is None

    # Non-string columns are never sampled, regardless of cardinality.
    id_column = next(column for column in artifact.columns if column.name == "id")
    assert id_column.is_categorical is None
    assert id_column.has_two_values is None
    assert id_column.sample_values is None


@pytest.mark.asyncio
async def test_tabular_processor_excludes_high_cardinality_string_columns(tmp_path):
    """
    A string column above the category limit gets no sample_values
    — there is nothing useful to ground an agent's query with once every row
    is (close to) unique, and listing them all would bloat the schema payload.
    """
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    rows = "\n".join(f"{i},unique-label-{i}" for i in range(30))
    csv_path = tmp_path / "wide.csv"
    csv_path.write_text(f"row,label\n{rows}\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-high-cardinality", file_name="wide.csv")
    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    label_column = next(column for column in artifact.columns if column.name == "label")
    assert label_column.is_categorical is False
    assert label_column.has_two_values is False
    assert label_column.sample_values is None


@pytest.mark.asyncio
async def test_tabular_processor_counts_null_rows_for_category_limit(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    values = ["RED", "GREEN", "BLUE"] * 3 + ["RED"] + [""] * 10
    csv_path = tmp_path / "nullable-status.csv"
    rows = "\n".join(f"{index},{value}," for index, value in enumerate(values))
    csv_path.write_text(f"id,status,blank\n{rows}\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-nullable-status", file_name="nullable-status.csv")
    artifact = read_tabular_artifact(processor.process(str(csv_path), metadata))

    assert artifact is not None
    assert artifact.row_count == 20
    status_column = next(column for column in artifact.columns if column.name == "status")
    assert status_column.is_categorical is True
    assert status_column.has_two_values is False
    assert status_column.sample_values is None
    blank_column = next(column for column in artifact.columns if column.name == "blank")
    assert blank_column.dtype == "string"
    assert blank_column.is_categorical is False
    assert blank_column.has_two_values is False
    assert blank_column.sample_values is None


@pytest.mark.asyncio
async def test_tabular_processor_preserves_duckdb_boolean_columns(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    csv_path = tmp_path / "flags.csv"
    csv_path.write_text("id,active,available,french\n1,true,yes,oui\n2,false,no,non\n3,,,\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-flags", file_name="flags.csv")
    artifact = read_tabular_artifact(processor.process(str(csv_path), metadata))

    assert artifact is not None
    assert artifact.row_count == 3
    columns = {column.name: column for column in artifact.columns}
    assert columns["active"].dtype == "boolean"
    assert columns["available"].dtype == "boolean"
    assert columns["active"].is_categorical is None
    assert columns["available"].is_categorical is None
    assert columns["active"].has_two_values is None
    assert columns["available"].has_two_values is None
    assert columns["french"].dtype == "string"
    assert columns["french"].has_two_values is True
    assert columns["french"].sample_values is None


@pytest.mark.asyncio
async def test_tabular_processor_excludes_one_value_above_category_limit(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    rows = "\n".join(f"{index},status-{index % 4}" for index in range(20))
    csv_path = tmp_path / "four-statuses.csv"
    csv_path.write_text(f"id,status\n{rows}\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-four-statuses", file_name="four-statuses.csv")
    artifact = read_tabular_artifact(processor.process(str(csv_path), metadata))

    assert artifact is not None
    assert artifact.row_count == 20
    status_column = next(column for column in artifact.columns if column.name == "status")
    assert status_column.is_categorical is False
    assert status_column.has_two_values is False
    assert status_column.sample_values is None


@pytest.mark.asyncio
async def test_tabular_processor_cleans_temporary_parquet_after_upload(tmp_path, monkeypatch):
    """
    Ensure the DuckDB-generated Parquet temp file is deleted after content-store upload.

    Why this exists:
    - Tabular ingestion briefly materializes a Parquet artifact on local disk
      before uploading it to shared content storage.
    - The temporary file should not remain on disk once processing completes.

    How to use:
    - Patch the generated Parquet path to a known location, run
      `TabularProcessor.process(...)`, and assert the file existed during
      upload but not afterwards.
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    tracked_parquet_path = tmp_path / "tracked-temp.parquet"
    observed: dict[str, object] = {"seen_during_upload": False}

    class _TrackedNamedTemporaryFile:
        def __enter__(self):
            self._handle = tracked_parquet_path.open("w+b")
            return self._handle

        def __exit__(self, exc_type, exc, tb):
            self._handle.close()

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-temp-cleanup", file_name="sales.csv")
    original_put_file = processor.content_store.put_file

    def _tracked_put_file(key: str, file_path: Path, *, content_type: str):
        observed["seen_during_upload"] = tracked_parquet_path.exists()
        assert file_path == tracked_parquet_path
        return original_put_file(key, file_path, content_type=content_type)

    monkeypatch.setattr(
        "knowledge_flow_backend.core.processors.output.tabular_processor.tabular_processor.tempfile.NamedTemporaryFile",
        lambda *args, **kwargs: _TrackedNamedTemporaryFile(),
    )
    monkeypatch.setattr(processor.content_store, "put_file", _tracked_put_file)

    processed_metadata = processor.process(str(csv_path), metadata)

    assert processed_metadata.file.row_count == 2
    assert observed["seen_during_upload"] is True
    assert not tracked_parquet_path.exists()


@pytest.mark.asyncio
async def test_tabular_processor_logs_parquet_artifact_write(tmp_path, caplog):
    """
    Ensure a successful Parquet upload leaves a durable, gathersable log line.

    Why this exists:
    - Before this fix, `_persist_parquet_artifact` uploaded the artifact to
      content storage without logging anything on success, so operators had
      no positive confirmation in stdout/OpenSearch that ingestion actually
      produced a queryable dataset.
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-log-check", file_name="sales.csv")

    with caplog.at_level("INFO"):
        processed_metadata = processor.process(str(csv_path), metadata)

    artifact = read_tabular_artifact(processed_metadata)
    assert artifact is not None

    tabular_log_records = [record.message for record in caplog.records if "[TABULAR]" in record.message]
    assert any("document_uid=doc-log-check" in message and f"object_key={artifact.object_key}" in message and "rows=2" in message for message in tabular_log_records)


@pytest.mark.integration
def test_tabular_processor_limits_python_rss_growth_for_large_csv(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    row_count = 400_000
    csv_path = tmp_path / "large.csv"
    with csv_path.open("w", encoding="utf-8") as file_handle:
        file_handle.write("city,amount,created_at\n")
        for index in range(row_count):
            file_handle.write(f"city_{index % 100},{index},2024-01-{(index % 28) + 1:02d}\n")

    assert csv_path.stat().st_size > 9 * 1024 * 1024

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-large", file_name="large.csv")
    rss_before_kb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss

    processed_metadata = processor.process(str(csv_path), metadata)

    rss_after_kb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    rss_growth_bytes = max(0, rss_after_kb - rss_before_kb) * 1024

    assert processed_metadata.file.row_count == row_count
    assert rss_growth_bytes < 128 * 1024 * 1024


@pytest.mark.asyncio
async def test_tabular_service_lists_context_and_queries_datasets(tmp_path):
    app_context = ApplicationContext.get_instance()
    content_store = app_context.get_content_store()
    metadata_store = app_context.get_metadata_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-targets",
        file_name="targets.csv",
        content="city,target\nParis,15\nLyon,25\n",
    )

    service = TabularEngineHarness()
    datasets = await service.list_datasets(_user())
    assert {dataset.document_uid for dataset in datasets} == {"doc-sales", "doc-targets"}

    dataset_by_uid = {dataset.document_uid: dataset for dataset in datasets}
    sales_alias = dataset_by_uid["doc-sales"].query_alias
    targets_alias = dataset_by_uid["doc-targets"].query_alias

    sales_frame = await service.read_dataset_frame(_user(), "doc-sales")
    assert sales_frame.to_dict(orient="records") == [
        {"city": "Paris", "amount": 10},
        {"city": "Lyon", "amount": 20},
    ]

    query_response = await service.query_read(
        _user(),
        request=TabularQueryRequest(
            sql=(f"SELECT s.city, s.amount, t.target FROM {sales_alias} AS s JOIN {targets_alias} AS t ON s.city = t.city ORDER BY s.amount DESC"),
        ),
    )
    assert query_response.rows == [
        {"city": "Lyon", "amount": 20, "target": 25},
        {"city": "Paris", "amount": 10, "target": 15},
    ]

    count_query = await service.query_read(_user(), request=TabularQueryRequest(sql=f"SELECT COUNT(*) AS total_rows FROM {sales_alias}"))
    assert count_query.rows == [{"total_rows": 2}]

    with pytest.raises(ValueError, match="Only SELECT or WITH statements are allowed"):
        await service.query_read(
            _user(),
            request=TabularQueryRequest(sql=f"DROP TABLE {sales_alias}"),
        )


@pytest.mark.asyncio
async def test_tabular_service_rejects_duckdb_table_functions_outside_authorized_datasets(tmp_path):
    """
    Verify read-only tabular queries cannot bypass dataset scoping with DuckDB functions.

    Why this exists:
    - Bare table functions such as `read_parquet(...)` are not mounted Fred
      datasets and must be blocked by the dataset allowlist before execution.

    How to use:
    - Ingest one valid dataset, then assert that a query using
      `read_parquet(...)` is rejected.
    """

    app_context = ApplicationContext.get_instance()
    content_store = app_context.get_content_store()
    metadata_store = app_context.get_metadata_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )

    service = TabularEngineHarness()

    with pytest.raises(ValueError, match=r"unauthorized datasets: read_parquet\(\)"):
        await service.query_read(
            _user(),
            request=TabularQueryRequest(
                sql="SELECT * FROM read_parquet('/tmp/forbidden.parquet')",
            ),
        )


@pytest.mark.asyncio
async def test_tabular_service_expands_excel_workbook_into_multiple_datasets(tmp_path):
    """
    Vérifie la chaîne spreadsheet complète : input Excel (upload Parquet en clés
    canoniques + sidecar), output (promotion en `tabular_multi_v1`), puis
    éclatement 1 document → N datasets interrogeables en SQL avec les alias
    déterministes publiés dans le catalogue.
    """
    from knowledge_flow_backend.features.tabular.artifacts import build_table_query_alias, read_tabular_multi_artifact

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    document_uid = "doc-excel"
    # Input stage (output.md + tables.json + Parquet en clés canoniques) puis
    # output stage (promotion du sidecar dans metadata.extensions)
    metadata, output_dir = await _ingest_excel_workbook(tmp_path=tmp_path, document_uid=document_uid)
    assert (output_dir / "output.md").is_file()
    multi = read_tabular_multi_artifact(metadata)
    assert multi is not None and len(multi.tables) == 2
    assert metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE

    # La sauvegarde (pruning des révisions) ne doit pas supprimer les tables courantes
    prefix = document_artifact_prefix(
        artifacts_prefix=ApplicationContext.get_instance().get_config().storage.tabular_store.artifacts_prefix,
        document_uid=document_uid,
    )
    assert {stored.key for stored in content_store.list_objects(prefix)} == {table.object_key for table in multi.tables}

    # Le service éclate le document en un dataset par table, alias stockés
    service = TabularEngineHarness()
    datasets = await service.list_datasets(_user())
    ventes_alias = build_table_query_alias(document_uid, "Ventes", 1)
    cibles_alias = build_table_query_alias(document_uid, "Cibles", 1)
    assert {dataset.query_alias for dataset in datasets} == {ventes_alias, cibles_alias}
    assert {dataset.document_uid for dataset in datasets} == {document_uid}

    # Les deux tables se montent et se joignent en SQL
    joined = await service.query_read(
        _user(),
        request=TabularQueryRequest(
            sql=f"SELECT v.city, v.amount, c.target FROM {ventes_alias} AS v JOIN {cibles_alias} AS c ON v.city = c.city ORDER BY v.city",
        ),
    )
    assert [row["city"] for row in joined.rows] == ["Lyon", "Paris"]

    # Scoper la requête sur l'uid du document monte bien TOUTES ses tables
    scoped = await service.query_read(
        _user(),
        request=TabularQueryRequest(
            sql=f"SELECT COUNT(*) AS n FROM {cibles_alias}",
            dataset_uids=[document_uid],
        ),
    )
    assert scoped.rows == [{"n": 2}]


@pytest.mark.asyncio
async def test_tabular_service_lists_documents_mixing_csv_and_spreadsheet(tmp_path, metadata_store):
    """
    Vérifie la surface document-centrique : un document par entrée (CSV comme
    classeur), kind correct, tables imbriquées avec alias, et liste allégée
    (pas de colonnes — elles relèvent de l'endpoint schemas).
    """
    from knowledge_flow_backend.features.tabular.artifacts import build_table_query_alias

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    await _ingest_excel_workbook(tmp_path=tmp_path, document_uid="doc-excel")

    service = TabularEngineHarness()
    documents = await service.list_documents(_user())
    documents_by_uid = {document.document_uid: document for document in documents}
    assert set(documents_by_uid) == {"doc-sales", "doc-excel"}

    csv_document = documents_by_uid["doc-sales"]
    assert csv_document.kind == "csv"
    assert len(csv_document.tables) == 1
    assert csv_document.tables[0].sheet is None
    assert csv_document.tables[0].row_count == 2

    excel_document = documents_by_uid["doc-excel"]
    assert excel_document.kind == "spreadsheet"
    assert {table.sheet for table in excel_document.tables} == {"Ventes", "Cibles"}
    assert {table.query_alias for table in excel_document.tables} == {
        build_table_query_alias("doc-excel", "Ventes", 1),
        build_table_query_alias("doc-excel", "Cibles", 1),
    }

    # La liste reste légère : le résumé de table n'expose pas les colonnes.
    assert "columns" not in csv_document.tables[0].model_dump()


@pytest.mark.asyncio
async def test_tabular_service_describes_documents_in_batch_with_all_tables(tmp_path, metadata_store):
    """
    Verify one mixed batch includes CSV types and the complete Excel catalog
    before every typed table, while denied and missing uids fail closed.
    """
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    metadata, output_dir = await _ingest_excel_workbook(tmp_path=tmp_path, document_uid="doc-excel")
    content_store.save_output("doc-excel", output_dir)
    metadata.mark_stage_done(ProcessingStage.PREVIEW_READY)
    await ApplicationContext.get_instance().get_metadata_store().save_metadata(metadata)

    service = TabularEngineHarness()
    schemas = await service.describe_documents(_user(), ["doc-sales", "doc-excel"])
    assert [schema.document_uid for schema in schemas] == ["doc-sales", "doc-excel"]

    csv_schema, excel_schema = schemas
    assert csv_schema.kind == "csv"
    assert csv_schema.markdown is None
    assert [column.name for column in csv_schema.tables[0].columns] == ["city", "amount"]
    assert csv_schema.tables[0].columns[0].is_categorical is True
    assert csv_schema.tables[0].columns[0].sample_values == ["Lyon", "Paris"]
    assert csv_schema.tables[0].columns[1].min_value == 10
    assert csv_schema.tables[0].columns[1].max_value == 20

    assert excel_schema.kind == "spreadsheet"
    assert excel_schema.markdown is not None
    assert "# Extraction summary" in excel_schema.markdown
    assert excel_schema.model_dump_json().index('"markdown"') < excel_schema.model_dump_json().index('"tables"')
    assert len(excel_schema.tables) == 2
    columns_by_sheet = {table.sheet: [column.name for column in table.columns] for table in excel_schema.tables}
    assert columns_by_sheet == {"Ventes": ["city", "amount"], "Cibles": ["city", "target"]}
    ventes_columns = next(table.columns for table in excel_schema.tables if table.sheet == "Ventes")
    assert ventes_columns[0].is_categorical is True
    assert ventes_columns[0].sample_values == ["Lyon", "Paris"]
    assert ventes_columns[1].min_value == 10
    assert ventes_columns[1].max_value == 20
    assert all(table.query_alias in excel_schema.markdown for table in excel_schema.tables)

    with pytest.raises(ValueError, match="At least one document uid"):
        await service.describe_documents(_user(), [])

    with pytest.raises(PermissionError, match="doc-unknown"):
        await service.describe_documents(_user(), ["doc-sales", "doc-unknown"])

    with pytest.raises(PermissionError, match="doc-hidden"):
        await service.describe_documents(_user(), ["doc-sales", "doc-hidden"])


@pytest.mark.asyncio
async def test_describe_documents_reads_values_instead_of_legacy_metadata(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    csv_metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount,rate\nParis,10,-1.5\nLyon,20,2.25\n",
    )
    for column in csv_metadata.extensions[TABULAR_EXTENSION_KEY]["columns"]:
        if column["name"] == "city":
            column["sample_values"] = ["stale city"]
        elif column["name"] == "amount":
            column["min_value"] = 999
            column["max_value"] = 999
        else:
            column.pop("min_value", None)
            column.pop("max_value", None)
    await ApplicationContext.get_instance().get_metadata_store().save_metadata(csv_metadata)

    excel_metadata, output_dir = await _ingest_excel_workbook(tmp_path=tmp_path, document_uid="doc-excel")
    content_store.save_output("doc-excel", output_dir)
    excel_metadata.mark_stage_done(ProcessingStage.PREVIEW_READY)
    for table in excel_metadata.extensions[TABULAR_MULTI_EXTENSION_KEY]["tables"]:
        for column in table["columns"]:
            if column["name"] == "city":
                column["sample_values"] = ["stale city"]
            else:
                column.pop("min_value", None)
                column.pop("max_value", None)
    await ApplicationContext.get_instance().get_metadata_store().save_metadata(excel_metadata)

    service = TabularEngineHarness()
    descriptions = await service.describe_documents(_user(), ["doc-sales", "doc-excel"])
    csv_columns = {column.name: column for column in descriptions[0].tables[0].columns}
    assert csv_columns["city"].sample_values == ["Lyon", "Paris"]
    assert (csv_columns["amount"].min_value, csv_columns["amount"].max_value) == (10, 20)
    assert (csv_columns["rate"].min_value, csv_columns["rate"].max_value) == (-1.5, 2.25)
    excel_bounds = {table.sheet: (table.columns[1].min_value, table.columns[1].max_value) for table in descriptions[1].tables}
    assert excel_bounds == {"Ventes": (10, 20), "Cibles": (15, 25)}
    assert all(table.columns[0].sample_values == ["Lyon", "Paris"] for table in descriptions[1].tables)
    assert descriptions[1].markdown is not None


@pytest.mark.asyncio
async def test_describe_documents_reads_current_values_on_every_call(tmp_path, metadata_store):
    import pandas as pd

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()
    metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-current",
        file_name="current.csv",
        content="city,score\nParis,-1.5\nLyon,2.25\n",
    )
    artifact = read_tabular_artifact(metadata)
    assert artifact is not None

    service = TabularEngineHarness()
    first = (await service.describe_documents(_user(), ["doc-current"]))[0].tables[0].columns
    assert first[0].sample_values == ["Lyon", "Paris"]
    assert (first[1].min_value, first[1].max_value) == (-1.5, 2.25)

    replacement = tmp_path / "replacement.parquet"
    pd.DataFrame({"city": ["Milan", "Rome", "Turin"], "score": [-5.0, 30.0, 15.0]}).to_parquet(replacement, index=False)
    content_store.put_file(artifact.object_key, replacement, content_type="application/vnd.apache.parquet")
    second = (await service.describe_documents(_user(), ["doc-current"]))[0].tables[0].columns
    assert second[0].sample_values == ["Milan", "Rome", "Turin"]
    assert (second[1].min_value, second[1].max_value) == (-5.0, 30.0)

    pd.DataFrame({"city": [None, None], "score": [float("inf"), float("nan")]}).to_parquet(replacement, index=False)
    content_store.put_file(artifact.object_key, replacement, content_type="application/vnd.apache.parquet")
    third = (await service.describe_documents(_user(), ["doc-current"]))[0].tables[0].columns
    assert third[0].sample_values is None
    assert (third[1].min_value, third[1].max_value) == (None, None)

    stored = read_tabular_artifact(await metadata_store.get_metadata_by_uid("doc-current"))
    assert stored is not None
    assert stored.columns[0].sample_values is None
    assert (stored.columns[1].min_value, stored.columns[1].max_value) == (None, None)


@pytest.mark.asyncio
async def test_describe_documents_reads_all_twenty_workbook_tables(tmp_path, metadata_store, monkeypatch):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()
    metadata, output_dir = await _ingest_excel_workbook(tmp_path=tmp_path, document_uid="doc-twenty", extra_tables=18)
    content_store.save_output("doc-twenty", output_dir)
    metadata.mark_stage_done(ProcessingStage.PREVIEW_READY)
    await ApplicationContext.get_instance().get_metadata_store().save_metadata(metadata)

    tables = (await TabularEngineHarness().describe_documents(_user(), ["doc-twenty"]))[0].tables
    assert len(tables) == 20
    assert all(table.columns[0].sample_values == ["Lyon", "Paris"] for table in tables)
    assert [(table.columns[1].min_value, table.columns[1].max_value) for table in tables[:2]] == [(10, 20), (15, 25)]
    assert [(table.columns[1].min_value, table.columns[1].max_value) for table in tables[2:]] == [(index, index + 10) for index in range(18)]

    limited_service = TabularEngineHarness()
    limited_service.tabular_config = limited_service.tabular_config.model_copy(update={"query": limited_service.tabular_config.query.model_copy(update={"max_selected_datasets": 19})})
    monkeypatch.setattr(limited_service, "_resolve_dataset_location", lambda *_: pytest.fail("No table scan above the configured limit"))
    with pytest.raises(ValueError, match="above the limit of 19"):
        await limited_service.describe_documents(_user(), ["doc-twenty"])


@pytest.mark.asyncio
async def test_describe_documents_fails_when_a_table_artifact_is_missing(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()
    metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-missing-artifact",
        file_name="missing.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    artifact = read_tabular_artifact(metadata)
    assert artifact is not None
    assert isinstance(content_store, FileSystemContentStore)
    (content_store.object_root / artifact.object_key).unlink()

    with pytest.raises(FileNotFoundError, match="Tabular artifact"):
        await TabularEngineHarness().describe_documents(_user(), ["doc-missing-artifact"])


@pytest.mark.asyncio
async def test_describe_documents_checks_artifacts_without_value_columns(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()
    metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-plain-strings",
        file_name="plain-strings.csv",
        content="name\nAlice\nBob\nCarla\n",
    )
    artifact = read_tabular_artifact(metadata)
    assert artifact is not None
    assert artifact.columns[0].is_categorical is False
    assert isinstance(content_store, FileSystemContentStore)
    (content_store.object_root / artifact.object_key).unlink()

    with pytest.raises(FileNotFoundError, match="Tabular artifact"):
        await TabularEngineHarness().describe_documents(_user(), ["doc-plain-strings"])


@pytest.mark.asyncio
async def test_describe_documents_redacts_signed_url_from_parquet_errors(tmp_path, metadata_store, monkeypatch):
    import duckdb

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()
    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-signed",
        file_name="signed.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )

    signed_url = "https://signed.example.invalid/data.parquet?X-Amz-Signature=secret"

    class FailingConnection:
        def from_parquet(self, _location):
            raise duckdb.IOException(f"Could not read {signed_url}")

    service = TabularEngineHarness()
    monkeypatch.setattr(service, "_resolve_dataset_location", lambda *_: signed_url)
    monkeypatch.setattr(service, "_ensure_httpfs_ready", lambda *_: None)
    monkeypatch.setattr("knowledge_flow_backend.features.tabular.service.open_duckdb_connection", lambda *_args, **_kwargs: FailingConnection())
    monkeypatch.setattr("knowledge_flow_backend.features.tabular.service.close_duckdb_connection", lambda *_args: None)

    with pytest.raises(TabularDatasetReadError) as exc_info:
        await service.describe_documents(_user(), ["doc-signed"])
    assert "secret" not in str(exc_info.value)
    assert "<redacted-signed-url>" in str(exc_info.value)


def test_tabular_service_httpfs_install_is_attempted_after_load_failure():
    """
    Verify the tabular runtime tries `INSTALL httpfs` after a failed load.

    Why this exists:
    - Connected environments may start without the extension preloaded even
      though remote tabular access is still expected to work.

    How to use:
    - Call `_ensure_httpfs_ready(...)` with a fake DuckDB connection whose
      first `LOAD httpfs` fails and whose subsequent install+load succeeds.

    Example:
    - `service._ensure_httpfs_ready(fake_connection)`
    """

    class _ConnectionProbe:
        def __init__(self) -> None:
            self.commands: list[str] = []
            self._first_load = True

        def execute(self, sql: str) -> None:
            self.commands.append(sql)
            if sql == "LOAD httpfs" and self._first_load:
                self._first_load = False
                raise RuntimeError("missing extension")

    service = TabularEngineHarness()
    connection = _ConnectionProbe()

    service._ensure_httpfs_ready(connection)  # type: ignore[arg-type]

    assert connection.commands == ["LOAD httpfs", "INSTALL httpfs", "LOAD httpfs"]


class _FakeVectorStore:
    """Records `add_documents` calls without touching a real vector backend."""

    def __init__(self, *, raise_on_add: bool = False):
        self.raise_on_add = raise_on_add
        self.added_documents: list = []

    def add_documents(self, documents):
        if self.raise_on_add:
            raise RuntimeError("simulated vector store failure")
        self.added_documents.extend(documents)
        return [doc.metadata.get("chunk_uid") for doc in documents]


def test_tabular_processor_pointer_chunks_disabled_by_default(tmp_path):
    """
    Verify the disabled-by-default path is untouched (RAG-DATASET-DISCOVERY-RFC.md §8:
    measured activation — no embedder/vector-store dependency paid unless enabled).
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    assert processor.tabular_config.pointer_chunks_enabled is False
    assert processor.vector_store is None
    assert processor.embedder is None

    metadata = _metadata(document_uid="doc-no-pointer", file_name="sales.csv")
    processed_metadata = processor.process(str(csv_path), metadata)

    artifact = read_tabular_artifact(processed_metadata)
    assert artifact is not None
    assert processed_metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE


def test_tabular_processor_emits_dataset_pointer_chunk_when_enabled(tmp_path):
    """
    Verify the pointer-chunk content, id, and kind match RAG-DATASET-DISCOVERY-RFC.md
    §2.1/§2.2/§2.4: one deterministic chunk, injection-resistant fixed template, no
    sample values in this increment.
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    processor.tabular_config.pointer_chunks_enabled = True
    fake_vector_store = _FakeVectorStore()
    processor.vector_store = fake_vector_store

    metadata = _metadata(document_uid="doc-pointer", file_name="sales.csv")
    processed_metadata = processor.process(str(csv_path), metadata)

    assert len(fake_vector_store.added_documents) == 1
    pointer_document = fake_vector_store.added_documents[0]

    assert pointer_document.metadata["chunk_uid"] == "doc-pointer::pointer"
    assert pointer_document.metadata["chunk_kind"] == "dataset_pointer"
    assert pointer_document.metadata["document_uid"] == "doc-pointer"
    # VectorSearchService unconditionally filters on metadata.retrievable=true
    # (vector_search_service.py:335,439,518) — without this the pointer chunk
    # is written but invisible to every search call. Caught live (2026-07-19)
    # by inspecting the actual OpenSearch document after a real ingestion.
    assert pointer_document.metadata["retrievable"] is True
    # Every deletion/consistency path in metadata/service.py gates its vector
    # cleanup on ProcessingStage.VECTORIZED being marked — without this, a
    # deleted tabular dataset leaves its pointer chunk permanently orphaned.
    # Caught live (2026-07-19): tabular artifacts/content were deleted but the
    # vector chunk was left behind because this guard never saw VECTORIZED.
    assert processed_metadata.processing.stages[ProcessingStage.VECTORIZED] == ProcessingStatus.DONE

    text = pointer_document.page_content
    assert "[DATASET POINTER" in text
    assert "[END DATASET POINTER]" in text
    assert "city" in text and "amount" in text
    assert "query_tabular_documents" in text
    assert "dataset_uid=doc-pointer" in text
    # Increment 1 explicitly excludes sample values (§2.4) — guard against
    # accidentally reintroducing real cell values into the pointer text.
    assert "Paris" not in text
    assert "Lyon" not in text
    assert "Sample values" not in text


def test_tabular_processor_skips_pointer_chunk_when_disabled_per_call(tmp_path):
    """
    Verify `emit_pointer_chunk=False` overrides an enabled deployment config.

    Why this exists:
    - `flat_metadata_from` (used by `_emit_pointer_chunk`) carries no
      `scope`/`user_id` fields, so a pointer chunk for a session-scoped chat
      attachment would be visible to every user's search regardless of the
      deployment's `pointer_chunks_enabled` setting (DESIGN.md,
      "Session-Scoped Attachment Datasets"). The per-call override is what
      the attachment ingestion path relies on to prevent that leak.
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    processor.tabular_config.pointer_chunks_enabled = True
    fake_vector_store = _FakeVectorStore()
    processor.vector_store = fake_vector_store

    metadata = _metadata(document_uid="doc-attachment", file_name="sales.csv")
    processed_metadata = processor.process(str(csv_path), metadata, emit_pointer_chunk=False)

    assert fake_vector_store.added_documents == []
    artifact = read_tabular_artifact(processed_metadata)
    assert artifact is not None
    assert processed_metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE
    assert ProcessingStage.VECTORIZED not in processed_metadata.processing.stages


def test_tabular_processor_pointer_chunk_failure_is_non_fatal(tmp_path, caplog):
    """
    Verify a pointer-chunk write failure never breaks Parquet ingestion — the
    processor's primary contract (RAG-DATASET-DISCOVERY-RFC.md §2.1: best-effort).
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    processor.tabular_config.pointer_chunks_enabled = True
    processor.vector_store = _FakeVectorStore(raise_on_add=True)

    metadata = _metadata(document_uid="doc-pointer-fail", file_name="sales.csv")
    with caplog.at_level("WARNING"):
        processed_metadata = processor.process(str(csv_path), metadata)

    artifact = read_tabular_artifact(processed_metadata)
    assert artifact is not None
    assert processed_metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE
    assert any("failed to write dataset pointer chunk" in record.message for record in caplog.records)


def test_tabular_store_config_pointer_chunks_disabled_by_default():
    """Guard the measured-activation default (RAG-DATASET-DISCOVERY-RFC.md §8)."""
    from knowledge_flow_backend.common.structures import TabularStoreConfig

    assert TabularStoreConfig().pointer_chunks_enabled is False


def test_allowed_chunk_keys_includes_chunk_kind():
    """chunk_kind must survive `sanitize_chunk_metadata`'s whitelist or pointer chunks silently lose it."""
    from knowledge_flow_backend.core.processors.output.vectorization_processor.vectorization_utils import sanitize_chunk_metadata

    clean, dropped = sanitize_chunk_metadata({"chunk_uid": "x::pointer", "chunk_kind": "dataset_pointer"})
    assert clean["chunk_kind"] == "dataset_pointer"
    assert "chunk_kind" not in dropped


@pytest.mark.asyncio
async def test_query_mounts_only_the_datasets_the_sql_references(tmp_path):
    """
    Verify an unscoped query pays only for the datasets it actually reads.

    Why this exists:
    - Leaving `dataset_uids` empty selects every readable dataset, and every
      selected dataset used to be mounted regardless of the SQL. On a remote
      content store that is one Parquet metadata fetch per dataset before any
      query runs, so an unscoped query on a large tenant was the cheapest way to
      overload the pod (issue #2182 finding D).
    """
    app_context = ApplicationContext.get_instance()
    content_store = app_context.get_content_store()
    metadata_store = app_context.get_metadata_store()
    content_store.clear()

    for index in range(4):
        await _ingest_csv(
            tmp_path=tmp_path,
            metadata_store=metadata_store,
            document_uid=f"doc-{index}",
            file_name=f"table-{index}.csv",
            content=f"city,amount\nParis,{index}\n",
        )

    service = TabularEngineHarness()
    datasets = await service.list_datasets(_user())
    assert len(datasets) == 4
    target_alias = {dataset.document_uid: dataset.query_alias for dataset in datasets}["doc-2"]

    mounted: list[list[str]] = []
    original_mount = service._mount_datasets

    def _record(*, connection, datasets, handle):
        mounted.append([dataset.query_alias for dataset in datasets])
        return original_mount(connection=connection, datasets=datasets, handle=handle)

    service._mount_datasets = _record

    # No dataset_uids at all: every dataset is authorized and selected, but only
    # the one named in the SQL may be mounted.
    response = await service.query_read(_user(), request=TabularQueryRequest(sql=f"SELECT amount FROM {target_alias}"))

    assert response.rows == [{"amount": 2}]
    assert mounted == [[target_alias]]
    assert response.query_aliases == [target_alias]
    assert response.dataset_uids == ["doc-2"]


@pytest.mark.asyncio
async def test_query_referencing_no_dataset_is_rejected(tmp_path):
    """
    Verify a query that reads no dataset is refused rather than executed.

    Why this exists:
    - A relation-free query is pointless on a tabular API, and it is also the
      cheapest shape for a pure scalar-allocation attack, which DuckDB's
      `memory_limit` does not bound (issue #2182). Rejecting it closes that
      shape; it does not close the same expression with a FROM clause attached.
    """
    app_context = ApplicationContext.get_instance()
    app_context.get_content_store().clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=app_context.get_metadata_store(),
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\n",
    )

    service = TabularEngineHarness()
    with pytest.raises(ValueError, match="references no authorized dataset"):
        await service.query_read(_user(), request=TabularQueryRequest(sql="SELECT 1"))


@pytest.mark.asyncio
async def test_query_referencing_more_datasets_than_the_cap_is_rejected(tmp_path):
    """
    Verify the dataset cap is a hard error, never a silent truncation.

    Why this exists:
    - Silently dropping an alias would resurface as a confusing "unauthorized
      dataset" error from `validate_read_query` instead of a clear message the
      caller can act on.
    """
    app_context = ApplicationContext.get_instance()
    content_store = app_context.get_content_store()
    metadata_store = app_context.get_metadata_store()
    content_store.clear()

    for index in range(3):
        await _ingest_csv(
            tmp_path=tmp_path,
            metadata_store=metadata_store,
            document_uid=f"doc-wide-{index}",
            file_name=f"wide-{index}.csv",
            content="city\nParis\n",
        )

    service = TabularEngineHarness()
    service.tabular_config = service.tabular_config.model_copy(update={"query": service.tabular_config.query.model_copy(update={"max_selected_datasets": 2})})
    aliases = [dataset.query_alias for dataset in await service.list_datasets(_user())]
    union_sql = " UNION ALL ".join(f"SELECT city FROM {alias}" for alias in aliases)

    with pytest.raises(ValueError, match="above the limit of 2"):
        await service.query_read(_user(), request=TabularQueryRequest(sql=union_sql))


@pytest.mark.asyncio
async def test_duckdb_resource_limits_are_applied_on_the_query_path(tmp_path, monkeypatch):
    """Every online connection must carry explicit threads/memory/spill settings."""
    observed_settings: list[tuple[int, str]] = []

    def record_connection(handle, *, config):
        connection = open_duckdb_connection(handle, config=config)
        observed_settings.append(connection.execute("SELECT current_setting('threads'), current_setting('temp_directory')").fetchone())
        return connection

    monkeypatch.setattr("knowledge_flow_backend.features.tabular.service.open_duckdb_connection", record_connection)
    app_context = ApplicationContext.get_instance()
    app_context.get_content_store().clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=app_context.get_metadata_store(),
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\n",
    )

    service = TabularEngineHarness()
    alias = (await service.list_datasets(_user()))[0].query_alias
    response = await service.query_read(
        _user(),
        request=TabularQueryRequest(sql=f"SELECT amount FROM {alias} LIMIT 1"),
    )

    assert response.rows == [{"amount": 10}]
    assert observed_settings == [(service.tabular_config.query.duckdb_threads, "")]


@pytest.mark.asyncio
async def test_query_cannot_read_outside_its_local_dataset(tmp_path, monkeypatch):
    app_context = ApplicationContext.get_instance()
    app_context.get_content_store().clear()
    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=app_context.get_metadata_store(),
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\n",
    )

    service = TabularEngineHarness()
    original_mount = service._mount_datasets

    def inspect_mount(*, connection, datasets, handle):
        original_mount(connection=connection, datasets=datasets, handle=handle)
        assert connection.execute("SELECT current_setting('enable_external_access')").fetchone() == (False,)
        with pytest.raises(duckdb.PermissionException):
            connection.execute("SELECT * FROM read_parquet('/etc/passwd')")

    monkeypatch.setattr(service, "_mount_datasets", inspect_mount)
    alias = (await service.list_datasets(_user()))[0].query_alias
    analytical_query = f"SELECT regexp_replace(lower(city), 'a', '_') AS city, quantile_cont(amount, 0.5) AS median_amount FROM {alias} GROUP BY city"
    response = await service.query_read(_user(), request=TabularQueryRequest(sql=analytical_query))
    assert response.rows == [{"city": "p_ris", "median_amount": 10.0}]

    with pytest.raises(ValueError, match="restricted SQL function"):
        await service.query_read(_user(), request=TabularQueryRequest(sql=f"SELECT current_setting('allowed_paths') FROM {alias}"))
    with pytest.raises(ValueError, match=r"unauthorized datasets: read_parquet\(\)"):
        await service.query_read(_user(), request=TabularQueryRequest(sql="SELECT * FROM read_parquet('/etc/passwd')"))
    with pytest.raises(ValueError, match=r"unauthorized datasets: duckdb_settings\(\)"):
        await service.query_read(_user(), request=TabularQueryRequest(sql=f"SELECT (SELECT value FROM duckdb_settings() LIMIT 1) FROM {alias}"))

    response = await service.query_read(_user(), request=TabularQueryRequest(sql=analytical_query))
    assert response.rows == [{"city": "p_ris", "median_amount": 10.0}]


@pytest.mark.integration
@pytest.mark.asyncio
async def test_query_reads_only_its_signed_parquet_url(tmp_path, monkeypatch):
    app_context = ApplicationContext.get_instance()
    content_store = app_context.get_content_store()
    content_store.clear()
    metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=app_context.get_metadata_store(),
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\n",
    )
    artifact = read_tabular_artifact(metadata)
    assert artifact is not None
    parquet_bytes = (content_store.object_root / artifact.object_key).read_bytes()

    class ParquetHandler(BaseHTTPRequestHandler):
        def log_message(self, *_args):
            pass

        def do_HEAD(self):
            self._respond(include_body=False)

        def do_GET(self):
            self._respond(include_body=True)

        def _respond(self, *, include_body: bool):
            range_header = self.headers.get("Range")
            start, end = 0, len(parquet_bytes) - 1
            if range_header:
                first, last = range_header.removeprefix("bytes=").split("-")
                start = int(first) if first else 0
                end = min(int(last), end) if last else end
            self.send_response(206 if range_header else 200)
            self.send_header("Content-Length", str(end - start + 1))
            self.send_header("Accept-Ranges", "bytes")
            if range_header:
                self.send_header("Content-Range", f"bytes {start}-{end}/{len(parquet_bytes)}")
            self.end_headers()
            if include_body:
                self.wfile.write(parquet_bytes[start : end + 1])

    server = ThreadingHTTPServer(("127.0.0.1", 0), ParquetHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    signed_url = f"http://127.0.0.1:{server.server_port}/data.parquet?signature=test"
    try:
        service = TabularEngineHarness()
        monkeypatch.setattr(service, "_resolve_dataset_location", lambda _key: signed_url)
        original_mount = service._mount_datasets

        def inspect_mount(*, connection, datasets, handle):
            original_mount(connection=connection, datasets=datasets, handle=handle)
            assert connection.execute("SELECT current_setting('allowed_paths')").fetchone()[0] == [signed_url]
            with pytest.raises(duckdb.PermissionException):
                connection.execute("SELECT * FROM read_parquet('/etc/passwd')")

        monkeypatch.setattr(service, "_mount_datasets", inspect_mount)
        alias = (await service.list_datasets(_user()))[0].query_alias
        response = await service.query_read(_user(), request=TabularQueryRequest(sql=f"SELECT amount FROM {alias}"))
        assert response.rows == [{"amount": 10}]
        for inspection_query in (
            f"SELECT current_setting('allowed_paths') FROM {alias}",
            f"SELECT pg_get_viewdef(1) FROM {alias}",
        ):
            with pytest.raises(ValueError, match="restricted SQL function") as error:
                await service.query_read(_user(), request=TabularQueryRequest(sql=inspection_query))
            assert signed_url not in str(error.value)
    finally:
        server.shutdown()
        server.server_close()
