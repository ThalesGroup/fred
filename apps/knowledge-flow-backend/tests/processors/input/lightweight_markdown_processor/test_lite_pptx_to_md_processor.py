# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Slide headings, speaker notes and native chart data in the fast PPTX extractor, decks generated per test."""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from knowledge_flow_backend.core.processors.input.fast_text_processor.fast_lite_pptx_processor import FastLitePptxProcessor
from knowledge_flow_backend.core.processors.input.lightweight_markdown_processor.lite_pptx_to_md_processor import LitePptxToMdExtractor

NOTES = {
    3: "Information non publiée : la marge inclut un produit exceptionnel de 340 000 € lié à la cession de l'entrepôt de Saint-Priest.",
    10: "Négociation : Méridia a obtenu une remise de 4 % en échange d'une clause d'exclusivité jusqu'en 2029.",
}


def _add_chart(slide, chart_type, categories: list[str], series: list[tuple[str, list[float]]], title: str | None = None):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(chart_type, Inches(1), Inches(1.5), Inches(6), Inches(4), data).chart
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
    return chart


def review_deck(path: Path) -> Path:
    deck = Presentation()
    for number in range(1, 11):
        slide = deck.slides.add_slide(deck.slide_layouts[5 if number == 5 else 1])
        slide.shapes.title.text = f"Titre {number}"
        if number != 5:
            slide.placeholders[1].text = f"Contenu de la diapositive {number}"
        if number in NOTES:
            slide.notes_slide.notes_text_frame.text = NOTES[number]
    _add_chart(
        deck.slides[4],
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        ["Juillet", "Août", "Septembre"],
        [("Gamme Atlas", [3450, 2980, 4120]), ("Gamme Boréal", [1870, 1640, 2210])],
        title="Record pour la gamme Atlas",
    )
    deck.save(str(path))
    return path


def text_of(path: Path) -> str:
    return FastLitePptxProcessor().extract(path).text


def test_slides_notes_and_chart_data_are_extracted(tmp_path):
    md = text_of(review_deck(tmp_path / "revue.pptx"))

    assert "## Slide 8: Titre 8" in md
    assert "Notes:\n" + NOTES[3] in md and "340 000 €" in md and "Saint-Priest" in md
    assert "Méridia" in md and "4 %" in md and "exclusivité jusqu'en 2029" in md
    assert "Chart (column clustered): Record pour la gamme Atlas" in md
    assert "| Category | Gamme Atlas | Gamme Boréal |" in md
    assert "| Septembre | 4120 | 2210 |" in md
    # Notes stay within their slide.
    assert md.index("Saint-Priest") < md.index("## Slide 4")


def test_deck_without_notes_or_charts_is_unchanged(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Planning"
    body = slide.placeholders[1].text_frame
    body.text = "Phase 1"
    sub = body.add_paragraph()
    sub.text, sub.level = "Lancement le 3 mars", 1
    table = deck.slides.add_slide(deck.slide_layouts[5])
    table.shapes.title.text = "Budget"
    grid = table.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (r, c), value in {(0, 0): "Poste", (0, 1): "k€", (1, 0): "Robots", (1, 1): "2 300"}.items():
        grid.cell(r, c).text = value
    deck.save(str(tmp_path / "plain.pptx"))

    assert text_of(tmp_path / "plain.pptx") == ("## Slide 1: Planning\n\nPlanning\nPhase 1\n  - Lancement le 3 mars\n\n## Slide 2: Budget\n\nBudget\n| Poste | k€ |\n| --- | --- |\n| Robots | 2 300 |")


def test_pie_keeps_one_series_and_large_charts_are_truncated(tmp_path):
    deck = Presentation()
    pie = deck.slides.add_slide(deck.slide_layouts[6])
    _add_chart(pie, XL_CHART_TYPE.PIE, ["Nord", "Sud"], [("Part", [60, 40]), ("Autre", [1, 2])])
    big = deck.slides.add_slide(deck.slide_layouts[6])
    _add_chart(big, XL_CHART_TYPE.LINE, [f"C{i}" for i in range(60)], [(f"S{j}", [float(i) + 0.5 for i in range(60)]) for j in range(12)])
    deck.save(str(tmp_path / "charts.pptx"))

    md = text_of(tmp_path / "charts.pptx")

    assert "Chart (pie)\n\n| Category | Part |\n| --- | --- |\n| Nord | 60 |\n| Sud | 40 |" in md
    assert "| C49 | 49.5 |" in md and "C50" not in md
    assert "| S9 |" in md and "S10" not in md
    assert "… (chart data truncated)" in md


def test_a_chart_that_fails_to_parse_is_skipped(tmp_path, monkeypatch):
    def broken(self, chart):
        raise ValueError("bad chart xml")

    monkeypatch.setattr(LitePptxToMdExtractor, "_chart_to_markdown", broken)
    md = text_of(review_deck(tmp_path / "revue.pptx"))

    assert "Chart (" not in md
    assert "## Slide 5: Titre 5" in md and "Saint-Priest" in md
