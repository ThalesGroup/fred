# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Text of document attachments appended to the mail markdown, documents generated per test."""

import importlib.util
import io
import time
from email.message import EmailMessage
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from knowledge_flow_backend.core.processors.input.email_processor import email_message_parser
from knowledge_flow_backend.core.processors.input.email_processor.email_markdown_processor import EmailMarkdownProcessor
from knowledge_flow_backend.core.processors.input.email_processor.email_message_parser import ATTACHMENT_NOT_EXTRACTABLE
from knowledge_flow_backend.core.processors.input.fast_text_processor.base_fast_text_processor import FastTextOptions
from knowledge_flow_backend.core.processors.input.fast_text_processor.fast_email_processor import FastEmailProcessor

_BUILDER_PATH = Path(__file__).parent / "fixtures" / "build_msg_fixtures.py"
_spec = importlib.util.spec_from_file_location("build_msg_fixtures", _BUILDER_PATH)
assert _spec and _spec.loader
msg_builder = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(msg_builder)


def pdf_bytes(text: str, *, password: str | None = None) -> bytes:
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer)
    page.drawString(72, 720, text)
    page.save()
    if password is None:
        return buffer.getvalue()
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(buffer.getvalue())))
    writer.encrypt(user_password=password, owner_password=password)
    encrypted = io.BytesIO()
    writer.write(encrypted)
    return encrypted.getvalue()


def docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_heading("Conditions", level=1)
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(text: str) -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Planning"
    slide.placeholders[1].text = text
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


_MIME = {".pdf": ("application", "pdf"), ".docx": ("application", "vnd.openxmlformats-officedocument.wordprocessingml.document"), ".txt": ("text", "plain")}


def mail(subject: str, body: str, files: list[tuple[str, bytes]] = (), attached: list[EmailMessage] = ()) -> EmailMessage:
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = "a@example.com", "b@example.com", subject
    msg.set_content(body)
    for name, data in files:
        maintype, subtype = _MIME.get(Path(name).suffix, ("application", "octet-stream"))
        if maintype == "text":
            msg.add_attachment(data.decode("utf-8"), subtype=subtype, filename=name)
        else:
            msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    for inner in attached:
        msg.add_attachment(inner)
    return msg


def library_markdown(tmp_path: Path, msg: EmailMessage) -> tuple[Path, str]:
    file_path = tmp_path / "mail.eml"
    file_path.write_bytes(msg.as_bytes())
    result = EmailMarkdownProcessor().convert_file_to_markdown(file_path, tmp_path / "out", None)
    return file_path, Path(result["md_file"]).read_text(encoding="utf-8")


def test_pdf_docx_and_pptx_texts_follow_the_attachment_list(tmp_path):
    msg = mail(
        "Offre",
        "Voir les pièces jointes.",
        [
            ("offre.pdf", pdf_bytes("Budget total 482 000 euros")),
            ("conditions.docx", docx_bytes("Paiement a 45 jours")),
            ("planning.pptx", pptx_bytes("Lancement le 3 mars")),
            ("photo.png", b"\x89PNG"),
            ("code.zip", b"PK\x03\x04"),
        ],
        [mail("Inner", "inner body")],
    )
    file_path, markdown = library_markdown(tmp_path, msg)

    listing, attached = markdown.index("## Attachments"), markdown.index("## Attached message 1")
    for heading, fact in [("### Attachment: offre.pdf (", "482 000"), ("### Attachment: conditions.docx (", "45 jours"), ("### Attachment: planning.pptx (", "3 mars")]:
        assert listing < markdown.index(heading) < markdown.index(fact) < attached
    # Attachment headings are pushed below the attachment heading level.
    assert "#### Conditions" in markdown
    # Slide headings nest under their attachment; PDF page markers stay off.
    assert markdown.index("### Attachment: planning.pptx (") < markdown.index("##### Slide 1: Planning") < markdown.index("3 mars")
    assert "Page 1" not in markdown
    # Images and archives stay names-only.
    assert "- photo.png (" in markdown and "- code.zip (" in markdown
    assert "Attachment: photo.png" not in markdown and "Attachment: code.zip" not in markdown
    # Metadata stays the parent's.
    assert EmailMarkdownProcessor().process_metadata(file_path, [], "uploads").identity.title == "Offre"
    # The chat attachment path renders the same text.
    assert FastEmailProcessor().extract(file_path).text.strip() == markdown.strip()


def test_same_bytes_are_extracted_once_across_the_file(tmp_path):
    proposal = pdf_bytes("Clause de revision annuelle")
    msg = mail("Top", "parent", [("Proposition.pdf", proposal), ("copie.pdf", proposal)], [mail("Fwd", "inner", [("Proposition.pdf", proposal)])])
    _, markdown = library_markdown(tmp_path, msg)

    assert markdown.count("Clause de revision annuelle") == 1
    assert markdown.count("_Same file as Proposition.pdf above._") == 2
    assert "### Attachment: copie.pdf (" in markdown
    assert "#### Attachment: Proposition.pdf (" in markdown.split("## Attached message 1")[1]


def test_unreadable_attachments_get_one_line_and_never_fail_the_mail(tmp_path):
    msg = mail(
        "Top",
        "parent stays readable",
        [
            ("secret.pdf", pdf_bytes("hidden text", password="pw")),  # pragma: allowlist secret
            ("broken.docx", b"PK not a docx"),
            ("protected.pptx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 EncryptedPackage"),
            ("ok.txt", b"plain words"),
        ],
    )
    _, markdown = library_markdown(tmp_path, msg)

    assert "parent stays readable" in markdown and "plain words" in markdown
    assert markdown.count(ATTACHMENT_NOT_EXTRACTABLE) == 3
    assert "hidden text" not in markdown and "EncryptedPackage" not in markdown


def test_oversized_attachments_stay_names_only(tmp_path, monkeypatch):
    monkeypatch.setattr(email_message_parser, "MAX_ATTACHMENT_EXTRACT_BYTES", 100)
    _, markdown = library_markdown(tmp_path, mail("Top", "parent", [("big.txt", b"x" * 101), ("small.txt", b"small text")]))

    assert "- big.txt (" in markdown and "Attachment: big.txt" not in markdown
    assert "small text" in markdown


def test_at_most_ten_distinct_attachments_are_extracted(tmp_path):
    files = [(f"note{i}.txt", f"note number {i}".encode()) for i in range(1, 13)]
    _, markdown = library_markdown(tmp_path, mail("Top", "parent", files))

    assert markdown.count("### Attachment: ") == 10
    assert "note number 10" in markdown and "note number 11" not in markdown
    assert "- note12.txt (" in markdown


def test_each_attachment_is_capped(tmp_path, monkeypatch):
    monkeypatch.setattr(email_message_parser, "MAX_ATTACHMENT_CHARS", 50)
    _, markdown = library_markdown(tmp_path, mail("Top", "parent", [("long.txt", b"word " * 100)]))

    section = markdown.split("### Attachment: long.txt")[1]
    assert section.count("word") < 15 and "…" in section


def test_body_budget_covers_attachment_texts_in_rendering_order(tmp_path, monkeypatch):
    monkeypatch.setattr("knowledge_flow_backend.core.processors.input.email_processor.email_markdown_processor.LIBRARY_MAX_BODY_CHARS", 100)
    msg = mail("Top", "short parent body", [("first.txt", b"word " * 100), ("second.txt", b"never extracted")], [mail("Inner", "inner body never shown")])
    _, markdown = library_markdown(tmp_path, msg)

    assert "short parent body" in markdown
    assert "### Attachment: first.txt (" in markdown and markdown.count("word") < 30
    assert "Attachment: second.txt" not in markdown and "never extracted" not in markdown
    assert "inner body never shown" not in markdown
    assert markdown.count("_[Body truncated: the message exceeds 100 characters.]_") == 2


def test_a_timeout_gives_the_failure_line_and_stops_extraction(tmp_path, monkeypatch):
    monkeypatch.setattr(email_message_parser, "ATTACHMENT_TIMEOUT_SECONDS", 0.2)
    monkeypatch.setattr(email_message_parser, "_extract_text", lambda name, path, outcome: time.sleep(1))
    _, markdown = library_markdown(tmp_path, mail("Top", "parent", [("slow.txt", b"slow"), ("next.txt", b"next")]))

    assert "### Attachment: slow.txt (" in markdown and ATTACHMENT_NOT_EXTRACTABLE in markdown
    assert "Attachment: next.txt" not in markdown


def test_msg_attachments_and_embedded_messages_are_extracted(tmp_path):
    proposal = pdf_bytes("Remise de 12 pour cent")
    inner = msg_builder.message(
        [(0x001A, msg_builder.PT_STRING, "IPM.Note"), (0x0037, msg_builder.PT_STRING, "Inner"), (0x1000, msg_builder.PT_STRING, "inner body")], attachments=[("offre-v2.pdf", proposal)], embedded=True
    )
    file_path = tmp_path / "mail.msg"
    msg_builder.write_cfb(
        file_path,
        msg_builder.message(
            [(0x001A, msg_builder.PT_STRING, "IPM.Note"), (0x0037, msg_builder.PT_STRING, "Top"), (0x1000, msg_builder.PT_STRING, "parent body")],
            attachments=[("offre.pdf", proposal), ("Inner", inner)],
        ),
    )
    markdown = Path(EmailMarkdownProcessor().convert_file_to_markdown(file_path, tmp_path / "out", None)["md_file"]).read_text(encoding="utf-8")

    assert markdown.count("Remise de 12 pour cent") == 1
    assert "_Same file as offre.pdf above._" in markdown.split("## Attached message 1")[1]


def test_chat_path_keeps_its_cap_over_attachment_texts(tmp_path):
    file_path = tmp_path / "mail.eml"
    file_path.write_bytes(mail("Top", "parent", [("a.txt", b"alpha " * 200), ("b.txt", b"bravo text")]).as_bytes())

    capped = FastEmailProcessor().extract(file_path, FastTextOptions(max_chars=400))
    assert capped.truncated is True and len(capped.text) <= 402
    assert "alpha" in capped.text and "bravo text" not in capped.text
