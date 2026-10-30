# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Analytics Excel export: workbook content, texts shared with the page, route and auth."""

from __future__ import annotations

import dataclasses
import io
import json
from datetime import date, datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock

import pytest
from control_plane_backend.kpi import api, export
from control_plane_backend.kpi import scope as kpi_scope
from control_plane_backend.kpi.export_texts import EXPORT_ONLY, SHARED, Lang
from control_plane_backend.kpi.export_workbook import (
    XLSX_MEDIA_TYPE,
    ExportData,
    build_workbook,
    export_filename,
)
from control_plane_backend.kpi.presets.analytics_activity_over_time import (
    AnalyticsActivityResponse,
)
from control_plane_backend.kpi.presets.analytics_adoption import (
    AnalyticsAdoptionResponse,
    AnalyticsCohortRow,
)
from control_plane_backend.kpi.presets.analytics_agents import (
    AnalyticsAgentRow,
    AnalyticsAgentsResponse,
    AnalyticsCapabilityRow,
)
from control_plane_backend.kpi.presets.analytics_common import BucketPoint, PeriodValue
from control_plane_backend.kpi.presets.analytics_quality_cost import (
    AnalyticsAgentCostRow,
    AnalyticsModelRow,
    AnalyticsQualityCostResponse,
)
from control_plane_backend.kpi.presets.analytics_teams import (
    AnalyticsTeamsResponse,
    AnalyticsTeamUsageRow,
)
from control_plane_backend.kpi.presets.analytics_teams_catalog import AnalyticsTeamRow
from control_plane_backend.kpi.presets.analytics_usage import AnalyticsUsageResponse
from control_plane_backend.kpi.presets.common import CREATION_ASSISTANT_LABEL
from control_plane_backend.kpi.user_filter import (
    AnalyticsUserFilter,
    ResolvedUserFilter,
    UserFilterMatch,
)
from fastapi import HTTPException, Request
from fastapi.routing import APIRoute
from fred_core import KeycloakUser
from fred_core.common import TeamId
from fred_core.security.models import AuthorizationError, Resource
from openpyxl import load_workbook

SINCE = datetime(2026, 9, 25, 22, tzinfo=timezone.utc)  # 26/09 in Paris
UNTIL = datetime(2026, 10, 25, 22, tzinfo=timezone.utc)  # end of 25/10 in Paris
PREVIOUS = datetime(2026, 8, 26, 22, tzinfo=timezone.utc)
USER = KeycloakUser(uid="user-secret-uid", username="secret.name", roles=[])
LOCALES = Path(__file__).resolve().parents[2] / "frontend" / "src" / "locales"


def pv(current: float | None, previous: float | None = None) -> PeriodValue:
    return PeriodValue(current=current, previous=previous)


def _responses() -> dict[str, Any]:
    return dict(
        adoption=AnalyticsAdoptionResponse(
            people_with_access=2950,
            visitors=pv(1560, 1410),
            active_users=pv(1120, 980),
            visit_rate=pv(0.53, 0.49),
            visitor_to_active_rate=pv(0.72, 0.70),
            adoption_rate=pv(0.38, 0.34),
            new_users=pv(186, 154),
            first_message_within_7_days_rate=pv(0.6),
            median_days_to_first_message=pv(1),
            next_month_return=pv(0.64, 0.61),
            cohorts=[
                AnalyticsCohortRow(month="2026-07", size=212, returns=[0.58, 0.47, 0.42]),
                AnalyticsCohortRow(month="2026-10", size=186, returns=[None, None, None]),
            ],
            since=SINCE, until=UNTIL, previous_since=PREVIOUS,
        ),
        usage=AnalyticsUsageResponse(
            conversations=pv(6240, 5480),
            messages=pv(18400, 15900),
            active_users=pv(1120, 980),
            active_users_per_workday=pv(268, 241),
            daily_usage_share=pv(0.2),
            days_per_week_median=pv(2.4, 2.2),
            conversations_per_user_median=pv(3),
            messages_per_conversation_median=pv(2.9, 2.9),
            single_exchange_share=pv(0.3),
            top_decile_conversation_share=pv(0.4),
            conversation_distribution=[],
            since=SINCE, until=UNTIL, previous_since=PREVIOUS,
        ),
        teams=AnalyticsTeamsResponse(
            active_teams=pv(2, 1),
            rows=[
                AnalyticsTeamUsageRow(
                    team_id="t-audit", name="Audit interne", active_members=2,
                    messages=140, messages_previous=100,
                ),
                AnalyticsTeamUsageRow(
                    team_id="t-big", name="Ingénierie", active_members=40,
                    messages=900, messages_previous=0,
                ),
            ],
            since=SINCE, until=UNTIL, previous_since=PREVIOUS,
        ),
        agents=AnalyticsAgentsResponse(
            agents_used=pv(8),
            agents_created=pv(2),
            agents_created_published=pv(1),
            documents_added=pv(10),
            teams_feeding_corpus=pv(3),
            top_agents=[
                AnalyticsAgentRow(
                    agent_instance_id="agent-instance-secret", label="Assistant RH",
                    users=412, conversations=900, messages=2700,
                )
            ],
            capabilities=[
                AnalyticsCapabilityRow(
                    capability_id="web_research",
                    agents_equipped=28,
                    agents_used=7,
                    calls=40,
                    users=12,
                    failed_calls=2,
                ),
                AnalyticsCapabilityRow(capability_id="unknown_cap", agents_equipped=1),
            ],
            since=SINCE, until=UNTIL, previous_since=PREVIOUS,
        ),
        quality=AnalyticsQualityCostResponse(
            failed_turn_rate=pv(0.018, 0.024),
            cost_eur=pv(4210, 3870),
            cost_eur_per_active_user=pv(3.76),
            input_tokens=pv(412_000_000, 361_000_000),
            output_tokens=pv(38_500_000, 33_900_000),
            cached_input_tokens=pv(0),
            co2e_grams=pv(96_000, 88_000),
            kwh=pv(200),
            models=[
                AnalyticsModelRow(
                    model="model-a", input_tokens=228, output_tokens=21,
                    cost_eur=2890, co2e_grams=61_000,
                )
            ],
            agents=[
                AnalyticsAgentCostRow(
                    label=CREATION_ASSISTANT_LABEL, tokens=500, cost_eur=12.5,
                    users=5, cost_eur_per_user=2.5,
                )
            ],
            since=SINCE, until=UNTIL, previous_since=PREVIOUS,
        ),
        activity=AnalyticsActivityResponse(
            bucket="day",
            series=["visitors"],
            points=[
                BucketPoint(
                    start="2026-09-26",
                    values={"visitors": 30, "active_users": 20, "conversations": 9,
                            "messages": 27, "cost_eur": 6.25},
                )
            ],
            since=SINCE,
            until=UNTIL,
        ),
    )  # fmt: skip


def _catalog() -> list[AnalyticsTeamRow]:
    return [
        AnalyticsTeamRow(team_id="t-audit", name="Audit interne", member_count=2),
        AnalyticsTeamRow(team_id="t-big", name="Ingénierie", member_count=131),
        AnalyticsTeamRow(team_id="t-idle", name="Achats", member_count=35),
    ]


def _data(team_ids: tuple[str, ...] = ()) -> ExportData:
    return ExportData(
        since=SINCE,
        until=UNTIL,
        previous_since=PREVIOUS,
        tz="Europe/Paris",
        team_ids=team_ids,
        exported_at=datetime(2026, 10, 26, 9, 42, tzinfo=timezone.utc),
        catalog=[r for r in _catalog() if not team_ids or r.team_id in team_ids],
        **_responses(),
    )


def _book(lang: Lang = "fr", team_ids: tuple[str, ...] = ()) -> Any:
    return load_workbook(io.BytesIO(build_workbook(_data(team_ids), lang)))  # type: ignore[arg-type]


def _filtered(lang: Lang) -> ExportData:
    user_filter = AnalyticsUserFilter.model_validate(
        {
            "combination": "any",
            "conditions": [
                {"claim": ["organisation"], "operator": "equals", "value": "=A"},
                {
                    "claim": ["profile", "unit"],
                    "operator": "contains",
                    "value": "R&D",
                    "case_sensitive": True,
                },
            ],
        }
    )
    match = UserFilterMatch(user_ids=("u1", "u2"), known_users=1200, no_data=1100)
    return dataclasses.replace(
        _data(), user_filter=ResolvedUserFilter(filter=user_filter, match=match)
    )


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_about_sheet_states_the_user_filter_and_its_coverage(lang: Lang) -> None:
    data = _filtered(lang)
    ws = load_workbook(io.BytesIO(build_workbook(data, lang)))[
        "À propos" if lang == "fr" else "About"
    ]
    values = {ws.cell(r, 1).value: ws.cell(r, 2) for r in range(3, 10)}
    if lang == "fr":
        text = values["Filtre utilisateurs"]
        assert text.value == (
            'Au moins une (OU) : organisation Est égal à "=A"; '
            'profile > unit Contient "R&D" (Sensible à la casse)'
        )
        assert values["Utilisateurs filtrés"].value == (
            "2 correspondent, 1\u202f100 sans donnée, sur 1\u202f200 utilisateurs connus"
        )
    else:
        text = values["User filter"]
        assert str(text.value).startswith('At least one (OR): organisation Equals "=A"')
        assert values["Filtered users"].value == (
            "2 matching, 1,100 without data, out of 1,200 known users"
        )
    # A value from a user is text, never a formula.
    assert text.data_type == "s"
    assert export_filename(data, lang) == (
        "analytics_plateforme-filtre_2026-09-26_2026-10-25.xlsx"
        if lang == "fr"
        else "analytics_platform-filtered_2026-09-26_2026-10-25.xlsx"
    )


def _header(ws: Any, row: int = 1) -> list[Any]:
    return [c.value for c in ws[row] if c.value is not None]


def _fr(key: str) -> str:
    return SHARED["fr"][f"rework.analytics.{key}"]


def test_sheets_tables_and_frozen_headers() -> None:
    wb = _book()
    assert wb.sheetnames == [
        "À propos", "Synthèse", "Évolution", "Adoption",
        "Équipes", "Agents", "Capacités", "Modèles et coût",
    ]  # fmt: skip
    tables = {name for ws in wb.worksheets for name in ws.tables}
    assert tables == {
        "Synthese", "Evolution", "Entonnoir", "Retour", "Equipes",
        "Agents", "Capacites", "Modeles", "CoutAgents",
    }  # fmt: skip
    assert wb["Synthèse"].freeze_panes == "A2"
    assert wb["À propos"].sheet_view.showGridLines is False
    assert wb["Synthèse"].sheet_view.showGridLines is True


def test_about_sheet_states_scope_period_and_comparison() -> None:
    ws = _book(team_ids=("t-audit", "t-big"))["À propos"]
    values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(3, 8)}
    assert values["Périmètre"] == "2 équipes (2 950 membres)"
    assert values["Période"] == "26/09/2026 – 25/10/2026"
    assert values["Période de comparaison"] == "27/08/2026 – 25/09/2026"
    assert values["Exporté le"] == datetime(2026, 10, 26, 10, 42)  # winter time
    assert _book()["À propos"].cell(3, 2).value == _fr("scope.platform")


def test_about_sheet_names_a_single_team_without_the_platform_team_count() -> None:
    ws = _book(team_ids=("t-audit",))["À propos"]
    assert ws.cell(3, 2).value.startswith("Audit interne (")
    texts = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "57" not in texts and " sur " not in texts


def test_summary_rows_formats_and_comments_reuse_page_texts() -> None:
    ws = _book()["Synthèse"]
    assert _header(ws) == [
        "Thème", "Indicateur", "26/09/2026 – 25/10/2026", "Période précédente",
        "Évolution", "Unité", "Commentaire",
    ]  # fmt: skip
    rows = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    adoption = rows[_fr("funnel.adoption.label")]
    assert (adoption[2].value, adoption[2].number_format) == (0.38, "0%")
    assert adoption[4].value == pytest.approx(4.0)
    assert "pts" in adoption[4].number_format
    assert adoption[6].value == (
        f"{_fr('funnel.adoption.def')} {_fr('howToRead')} {_fr('funnel.adoption.how')}"
    )
    cost = rows[_fr("tiles.cost.label")]
    assert (cost[2].value, cost[2].number_format, cost[5].value) == (
        4210,
        '#,##0 "€"',
        "€",
    )
    assert cost[4].value == pytest.approx(340 / 3870)
    footprint = rows[_fr("tiles.footprint.label")]
    assert footprint[2].value == pytest.approx(96)
    # A definition without a "how to read" stays alone.
    assert rows[_fr("tiles.conversations.label")][6].value == _fr(
        "tiles.conversations.def"
    )


def test_over_time_adoption_teams_and_models_sheets() -> None:
    wb = _book()
    ws = wb["Évolution"]
    assert _header(ws) == [
        "Date", "Visiteurs", "Utilisateurs actifs", "Conversations",
        "Messages envoyés", "Coût estimé",
    ]  # fmt: skip
    assert ws["A2"].value == datetime(2026, 9, 26)
    assert ws["A2"].number_format == "dd/mm/yyyy"
    assert ws["F2"].value == 6.25
    assert ws["B1"].comment.text == _fr("funnel.steps.visitors.def")

    ws = wb["Adoption"]
    assert [ws.cell(r, 2).value for r in (3, 4, 5)] == [2950, 1560, 1120]
    assert ws["C4"].value == pytest.approx(1560 / 2950)
    assert ws["A2"].comment.text.startswith(_fr("funnel.def"))
    assert ws["A11"].value == datetime(2026, 7, 1)
    assert ws["A11"].number_format == "mmmm yyyy"
    assert ws["E11"].value == 0.42 and ws["C12"].value is None

    ws = wb["Équipes"]
    names = [ws.cell(r, 1).value for r in range(2, 5)]
    # Every listed team, active or not, two-member teams included.
    assert names == ["Ingénierie", "Audit interne", "Achats"]
    assert [ws.cell(3, c).value for c in range(2, 7)] == [2, 2, 1.0, 140, 0.4]
    assert ws.cell(2, 6).value is None  # no previous messages: no change
    assert ws["C1"].comment is not None

    assert wb["Capacités"]["A2"].value == SHARED["fr"]["capability.web_research.name"]
    assert wb["Capacités"]["A3"].value == "unknown_cap"
    assert [wb["Capacités"].cell(2, c).value for c in range(2, 8)] == [
        28, 7, 0.25, 40, 12, 0.05,
    ]  # fmt: skip
    assert wb["Capacités"]["D2"].number_format == "0%"
    assert wb["Capacités"]["D3"].value == 0.0 and wb["Capacités"]["G3"].value is None
    ws = wb["Modèles et coût"]
    assert ws["E3"].value == pytest.approx(61)
    assert ws["A7"].value == _fr("tokenUsage.byAgent.creationAssistant")


def test_english_workbook_and_filenames() -> None:
    wb = _book("en")
    assert wb.sheetnames[:3] == ["About", "Summary", "Over time"]
    assert wb["Over time"]["A2"].number_format == "yyyy-mm-dd"
    assert export_filename(_data(), "fr") == (
        "analytics_plateforme_2026-09-26_2026-10-25.xlsx"
    )
    assert export_filename(_data(("a", "b", "c")), "en") == (
        "analytics_3-teams_2026-09-26_2026-10-25.xlsx"
    )
    assert export_filename(_data(("a", "b")), "fr").startswith("analytics_2-equipes_")


def test_workbook_never_contains_user_or_agent_ids() -> None:
    langs: tuple[Lang, ...] = ("fr", "en")
    for lang in langs:
        wb = _book(lang)
        cells = [
            str(c.value)
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for c in row
            if c.value is not None
        ]
        blob = "\n".join(cells)
        for secret in ("user-secret-uid", "secret.name", "agent-instance-secret"):
            assert secret not in blob


def _frontend(lang: Lang) -> dict[str, Any]:
    return json.loads((LOCALES / lang / "translation.json").read_text("utf-8"))


@pytest.mark.skipif(not LOCALES.exists(), reason="frontend sources not checked out")
@pytest.mark.parametrize("lang", ["fr", "en"])
def test_shared_texts_match_the_page_translations(lang: Lang) -> None:
    translations = _frontend(lang)
    for key, value in SHARED[lang].items():
        node: Any = translations
        for part in key.split("."):
            node = node[part]
        assert node == value, f"{lang}:{key} differs from the page translation"


def test_both_languages_have_every_text() -> None:
    assert SHARED["fr"].keys() == SHARED["en"].keys()
    assert EXPORT_ONLY["fr"].keys() == EXPORT_ONLY["en"].keys()


# --- route -------------------------------------------------------------------


def _route() -> APIRoute:
    return next(
        r
        for r in api.build_kpi_router().routes
        if isinstance(r, APIRoute) and r.path == "/kpi/analytics/export"
    )


def _call(**overrides: Any) -> dict[str, Any]:
    return {
        **dict(
            request=Request({"type": "http"}),
            since=SINCE,
            until=UNTIL,
            team_ids=[],
            bucket="auto",
            tz="Europe/Paris",
            lang="fr",
            user=USER,
            store=SimpleNamespace(),
        ),
        **overrides,
    }


@pytest.fixture
def presets(monkeypatch: pytest.MonkeyPatch) -> dict[str, AsyncMock]:
    responses = _responses()
    mocks = {
        "query_analytics_adoption": AsyncMock(return_value=responses["adoption"]),
        "query_analytics_usage": AsyncMock(return_value=responses["usage"]),
        "query_analytics_teams": AsyncMock(return_value=responses["teams"]),
        "query_analytics_agents": AsyncMock(return_value=responses["agents"]),
        "query_analytics_quality_cost": AsyncMock(return_value=responses["quality"]),
        "query_analytics_activity_over_time": AsyncMock(
            return_value=responses["activity"]
        ),
        "team_catalog": AsyncMock(return_value=_catalog()),
    }
    for name, mock in mocks.items():
        monkeypatch.setattr(export, name, mock)
    return mocks


@pytest.mark.asyncio
async def test_route_returns_the_workbook_with_its_filename(
    presets: dict[str, AsyncMock], monkeypatch: pytest.MonkeyPatch
) -> None:
    scope = AsyncMock(return_value=kpi_scope.KpiScope(team_ids=(TeamId("t-big"),)))
    monkeypatch.setattr(export, "resolve_kpi_scope", scope)

    response = await _route().endpoint(**_call(team_ids=[TeamId("t-big")], lang="en"))

    assert response.media_type == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="analytics_1-teams_2026-09-26_2026-10-25.xlsx"'
    )
    assert load_workbook(io.BytesIO(response.body)).sheetnames[0] == "About"
    awaited = presets["query_analytics_activity_over_time"].await_args
    assert awaited is not None
    kwargs = awaited.kwargs
    assert kwargs["team_ids"] == (TeamId("t-big"),)
    assert (kwargs["bucket"], kwargs["tz"]) == ("auto", "Europe/Paris")
    presets["team_catalog"].assert_awaited_once()


@pytest.mark.asyncio
async def test_route_passes_the_user_filter_to_scope_and_sections(
    presets: dict[str, AsyncMock], monkeypatch: pytest.MonkeyPatch
) -> None:
    resolved = _filtered("en").user_filter
    scope = AsyncMock(return_value=kpi_scope.KpiScope(user_filter=resolved))
    monkeypatch.setattr(export, "resolve_kpi_scope", scope)
    raw = '{"conditions":[{"claim":["organisation"],"operator":"equals","value":"A"}]}'

    response = await _route().endpoint(**_call(user_filter=raw, lang="en"))

    passed = scope.await_args
    assert passed is not None
    assert passed.kwargs["user_filter"].conditions[0].value == "A"
    for name, mock in presets.items():
        if name != "team_catalog":
            assert mock.await_args is not None
            assert mock.await_args.kwargs["user_ids"] == ("u1", "u2")
    assert "platform-filtered" in response.headers["content-disposition"]


@pytest.mark.asyncio
async def test_route_rejects_unknown_time_zones(presets: dict[str, AsyncMock]) -> None:
    with pytest.raises(HTTPException) as exc:
        await _route().endpoint(**_call(tz="Mars/Olympus"))
    assert exc.value.status_code == 400


class _DenyAll:
    async def has_user_permission(self, *args: Any, **kwargs: Any) -> bool:
        return False

    async def check_user_permission_or_raise(self, *args: Any, **kwargs: Any) -> None:
        raise AuthorizationError("u", "can_observe_platform", Resource.ORGANIZATION)


@pytest.mark.asyncio
async def test_non_observer_is_forbidden_before_any_read(
    presets: dict[str, AsyncMock], monkeypatch: pytest.MonkeyPatch
) -> None:
    container = SimpleNamespace(get_rebac_engine=lambda: _DenyAll())
    monkeypatch.setattr(kpi_scope, "get_application_container", lambda r: container)

    with pytest.raises(AuthorizationError):
        await _route().endpoint(**_call())
    for mock in presets.values():
        mock.assert_not_awaited()


def test_cohort_months_are_real_dates() -> None:
    ws = _book()["Adoption"]
    assert isinstance(ws["A12"].value, datetime)
    assert ws["A12"].value.date() == date(2026, 10, 1)


def test_user_named_cells_stay_text_and_never_break_the_file() -> None:
    data = _data()
    data.agents.top_agents[0].label = '=HYPERLINK("http://evil","x")'
    data.catalog[0].name = "Audit\x01 interne"
    wb = load_workbook(io.BytesIO(build_workbook(data, "fr")))  # type: ignore[arg-type]
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert not [c for c in cells if c.data_type == "f"]
    values = [c.value for c in cells]
    assert '=HYPERLINK("http://evil","x")' in values
    assert "Audit interne" in values
