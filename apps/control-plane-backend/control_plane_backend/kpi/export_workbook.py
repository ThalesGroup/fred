# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The analytics Excel workbook, rendered from the section presets' responses.

Only aggregates and names of teams, agents, capabilities and models are
written: no user name, email or id ever reaches a sheet.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from control_plane_backend.kpi.counted_days import ALL_DAYS, CountedDays
from control_plane_backend.kpi.export_texts import SHARED, Lang, text
from control_plane_backend.kpi.presets.analytics_activity_over_time import (
    AnalyticsActivityResponse,
)
from control_plane_backend.kpi.presets.analytics_adoption import (
    AnalyticsAdoptionResponse,
)
from control_plane_backend.kpi.presets.analytics_agents import AnalyticsAgentsResponse
from control_plane_backend.kpi.presets.analytics_common import PeriodValue
from control_plane_backend.kpi.presets.analytics_quality_cost import (
    AnalyticsQualityCostResponse,
)
from control_plane_backend.kpi.presets.analytics_teams import AnalyticsTeamsResponse
from control_plane_backend.kpi.presets.analytics_teams_catalog import AnalyticsTeamRow
from control_plane_backend.kpi.presets.analytics_usage import AnalyticsUsageResponse
from control_plane_backend.kpi.presets.common import CREATION_ASSISTANT_LABEL
from control_plane_backend.kpi.user_filter import ResolvedUserFilter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

PCT, INT, DEC = "0%", "#,##0", "#,##0.0"
EUR, EUR_CENTS = '#,##0 "€"', '#,##0.00 "€"'
POINTS = '+0.0" pts";-0.0" pts";0" pts"'
DELTA = "+0%;-0%;0%"
DATE_FORMAT: dict[Lang, str] = {"fr": "dd/mm/yyyy", "en": "yyyy-mm-dd"}
STRFTIME: dict[Lang, str] = {"fr": "%d/%m/%Y", "en": "%Y-%m-%d"}
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


@dataclass(frozen=True)
class ExportData:
    since: datetime
    until: datetime
    previous_since: datetime
    tz: str
    team_ids: tuple[str, ...]
    exported_at: datetime
    adoption: AnalyticsAdoptionResponse
    usage: AnalyticsUsageResponse
    teams: AnalyticsTeamsResponse
    agents: AnalyticsAgentsResponse
    quality: AnalyticsQualityCostResponse
    activity: AnalyticsActivityResponse
    # Selected teams (every team for the platform).
    catalog: list[AnalyticsTeamRow]
    user_filter: ResolvedUserFilter | None = None
    days: CountedDays = ALL_DAYS

    def local_day(self, moment: datetime, *, end: bool = False) -> date:
        # The period end is an instant: its last covered day is the one just before it.
        shifted = moment - timedelta(microseconds=1) if end else moment
        return shifted.astimezone(ZoneInfo(self.tz)).date()


def export_filename(data: ExportData, lang: Lang) -> str:
    scope = (
        text(lang, "filename.teams").format(count=len(data.team_ids))
        if data.team_ids
        else text(lang, "filename.platform")
    )
    if data.user_filter is not None:
        scope = f"{scope}-{text(lang, 'filename.filtered')}"
    start = data.local_day(data.since)
    end = data.local_day(data.until, end=True)
    return f"analytics_{scope}_{start.isoformat()}_{end.isoformat()}.xlsx"


def build_workbook(data: ExportData, lang: Lang) -> bytes:
    """Render every sheet and return the `.xlsx` bytes (CPU-bound: run off the loop)."""
    book = _Book(data, lang)
    book.about()
    book.summary()
    book.over_time()
    book.adoption()
    book.teams()
    book.agents()
    book.capabilities()
    book.models()
    for sheet in book.wb.worksheets:
        sheet.sheet_view.showGridLines = sheet.title != text(lang, "sheet.about")
    buffer = io.BytesIO()
    book.wb.save(buffer)
    return buffer.getvalue()


def _write(ws: Worksheet, row: int, column: int, value: Any) -> Any:
    """Names come from users: never let one become a formula or break the file."""
    if not isinstance(value, str):
        return ws.cell(row, column, value)
    cell = ws.cell(row, column, ILLEGAL_CHARACTERS_RE.sub("", value))
    cell.data_type = "s"
    return cell


def _change(current: float | None, previous: float | None, percent: bool) -> Any:
    if current is None or previous is None:
        return None
    if percent:
        return (current - previous) * 100
    return (current - previous) / previous if previous else None


def _scaled(value: PeriodValue, factor: float) -> PeriodValue:
    def one(v: float | None) -> float | None:
        return v * factor if v is not None else None

    return PeriodValue(current=one(value.current), previous=one(value.previous))


class _Book:
    def __init__(self, data: ExportData, lang: Lang) -> None:
        self.data = data
        self.lang: Lang = lang
        self.wb = Workbook()
        self._first = True

    def t(self, key: str) -> str:
        return text(self.lang, key)

    def explain(self, prefix: str) -> str | None:
        """`<prefix>.def`, then "How to read it:" and `<prefix>.how` when there is one."""
        shared = SHARED[self.lang]
        definition = shared.get(f"rework.analytics.{prefix}.def")
        if definition is None:
            return None
        how = shared.get(f"rework.analytics.{prefix}.how")
        return f"{definition} {self.t('howToRead')} {how}" if how else definition

    def sheet(self, key: str) -> Worksheet:
        if self._first:
            self._first = False
            ws = self.wb.active
            assert ws is not None
            ws.title = self.t(key)
            return ws
        return self.wb.create_sheet(self.t(key))

    def day(self, moment: datetime, *, end: bool = False) -> str:
        return self.data.local_day(moment, end=end).strftime(STRFTIME[self.lang])

    def period(self, start: datetime, end: datetime) -> str:
        return f"{self.day(start)} – {self.day(end, end=True)}"

    def table(
        self,
        ws: Worksheet,
        name: str,
        headers: Sequence[tuple[str, str | None]],
        rows: Sequence[Sequence[Any]],
        formats: Sequence[str | None],
        *,
        start_row: int = 1,
        widths: dict[int, int] | None = None,
        wrap_last: bool = False,
    ) -> None:
        """An Excel table with filters; each header is (label, note)."""
        for c, (label, note) in enumerate(headers, 1):
            cell = ws.cell(start_row, c, label)
            if note:
                cell.comment = Comment(note, "Analytics", width=400, height=160)
        # A table needs at least one data row to open without a repair prompt.
        body = rows or [[None] * len(headers)]
        for r, row in enumerate(body, start_row + 1):
            for c, value in enumerate(row, 1):
                cell = _write(ws, r, c, value)
                number_format = formats[c - 1]
                if number_format:
                    cell.number_format = number_format
                cell.alignment = WRAP if wrap_last and c == len(row) else TOP
        ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(body)}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        ws.add_table(table)
        ws.freeze_panes = f"A{start_row + 1}"
        for i, (label, _) in enumerate(headers, 1):
            width = (widths or {}).get(i) or max(12, len(label) + 4)
            current = ws.column_dimensions[get_column_letter(i)].width or 0
            ws.column_dimensions[get_column_letter(i)].width = max(width, current)

    def heading(self, ws: Worksheet, cell: str, label: str) -> None:
        ws[cell] = label
        ws[cell].font = Font(bold=True, size=12)

    # --- sheets ----------------------------------------------------------------

    def about(self) -> None:
        data = self.data
        ws = self.sheet("sheet.about")
        ws["A1"] = f"{self.t('title')} — export"
        ws["A1"].font = Font(bold=True, size=14)
        if data.team_ids:
            # A team admin exports their own team: no platform-wide team count.
            names = [row.name for row in data.catalog]
            scope = (
                names[0]
                if len(data.team_ids) == 1 and len(names) == 1
                else self.t("about.teams").format(count=len(data.team_ids))
            )
            people = data.adoption.people_with_access
            if people is not None:
                members = self.t("about.members").format(members=self.count(people))
                scope = f"{scope} ({members})"
        else:
            scope = self.t("scope.platform")
        exported = data.exported_at.astimezone(ZoneInfo(data.tz)).replace(tzinfo=None)
        info: list[tuple[str, Any]] = [
            (self.t("about.scope"), scope),
            (self.t("about.period"), self.period(data.since, data.until)),
            (
                self.t("about.comparison"),
                self.period(data.previous_since, data.since),
            ),
            (self.t("about.days"), self.days_text(data.days)),
            (self.t("about.exportedAt"), exported),
            (self.t("about.privacy"), self.t("about.privacyText")),
        ]
        if data.user_filter is not None:
            match = data.user_filter.match
            info[1:1] = [
                (self.t("about.userFilter"), self.filter_text(data.user_filter)),
                (
                    self.t("about.userFilterCoverage"),
                    self.t("about.userFilterCounts").format(
                        matched=self.count(len(match.user_ids)),
                        no_data=self.count(match.no_data),
                        known=self.count(match.known_users),
                    ),
                ),
            ]
        for i, (label, value) in enumerate(info, 3):
            ws.cell(i, 1, label).font = Font(bold=True)
            # Team names and filter values come from users.
            cell = _write(ws, i, 2, value)
            if isinstance(value, datetime):
                cell.number_format = f"{DATE_FORMAT[self.lang]} hh:mm"
                cell.alignment = Alignment(horizontal="left")
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 60

    def days_text(self, days: CountedDays) -> str:
        key = {
            (False, False): "all",
            (True, False): "noWeekends",
            (False, True): "noHolidays",
            (True, True): "workdays",
        }[(days.exclude_weekends, days.exclude_holidays)]
        return self.t(f"about.days.{key}")

    def count(self, value: int) -> str:
        return f"{value:,}".replace(",", "\u202f" if self.lang == "fr" else ",")

    def filter_text(self, resolved: ResolvedUserFilter) -> str:
        """The conditions as the page shows them, e.g. `organisation Equals "A"`."""
        rule = "rework.platformAccess.rule"
        conditions = [
            " ".join(
                [
                    " > ".join(c.claim),
                    self.t(f"{rule}.operatorLabel.{c.operator}"),
                    f'"{c.value}"',
                    *(
                        [f"({self.t(f'{rule}.caseSensitive')})"]
                        if c.case_sensitive
                        else []
                    ),
                ]
            )
            for c in resolved.filter.conditions
        ]
        if len(conditions) == 1:
            return conditions[0]
        combination = self.t(f"{rule}.{resolved.filter.combination}")
        separator = " : " if self.lang == "fr" else ": "
        return f"{combination}{separator}" + "; ".join(conditions)

    def summary(self) -> None:
        data = self.data
        adoption, usage, quality = data.adoption, data.usage, data.quality
        adoption_theme, usage_theme = self.t("tabs.adoption"), self.t("tabs.usage")
        quality_theme, cost_theme = self.t("theme.quality"), self.t("theme.cost")
        access = PeriodValue(
            current=adoption.people_with_access, previous=adoption.people_with_access
        )
        # (theme, label key, explanation prefix, value, unit key, number format)
        indicators: list[tuple[str, str, str, PeriodValue, str, str]] = [
            (adoption_theme, "funnel.steps.peopleWithAccess.label", "funnel.steps.peopleWithAccess", access, "unit.people", INT),
            (adoption_theme, "funnel.steps.visitors.label", "funnel.steps.visitors", adoption.visitors, "unit.people", INT),
            (adoption_theme, "tiles.activeUsers.label", "tiles.activeUsers", adoption.active_users, "unit.people", INT),
            (adoption_theme, "funnel.adoption.label", "funnel.adoption", adoption.adoption_rate, "%", PCT),
            (adoption_theme, "funnel.rates.visitRate.name", "funnel.rates.visitRate", adoption.visit_rate, "%", PCT),
            (adoption_theme, "funnel.rates.visitorToActive.name", "funnel.rates.visitorToActive", adoption.visitor_to_active_rate, "%", PCT),
            (adoption_theme, "tiles.newUsers.label", "tiles.newUsers", adoption.new_users, "unit.people", INT),
            (adoption_theme, "tiles.nextMonthReturn.label", "tiles.nextMonthReturn", adoption.next_month_return, "%", PCT),
            (adoption_theme, "tiles.activeTeams.label", "tiles.activeTeams", data.teams.active_teams, "unit.teams", INT),
            (usage_theme, "tiles.conversations.label", "tiles.conversations", usage.conversations, "unit.conversations", INT),
            (usage_theme, "tiles.messages.label", "tiles.messages", usage.messages, "unit.messages", INT),
            (usage_theme, "tiles.activeUsersPerWorkday.label", "tiles.activeUsersPerWorkday", usage.active_users_per_workday, "unit.people", DEC),
            (usage_theme, "tiles.daysPerWeek.label", "tiles.daysPerWeek", usage.days_per_week_median, "unit.days", DEC),
            (usage_theme, "tiles.messagesPerConversation.label", "tiles.messagesPerConversation", usage.messages_per_conversation_median, "unit.messages", DEC),
            (quality_theme, "tiles.failedAnswers.label", "tiles.failedAnswers", quality.failed_turn_rate, "%", PCT),
            (cost_theme, "tiles.cost.label", "tiles.cost", quality.cost_eur, "€", EUR),
            (cost_theme, "tokensSent", "tiles.tokens", quality.input_tokens, "unit.tokens", INT),
            (cost_theme, "tokensGenerated", "tiles.tokens", quality.output_tokens, "unit.tokens", INT),
            (cost_theme, "tiles.footprint.label", "tiles.footprint", _scaled(quality.co2e_grams, 0.001), "unit.kgCo2e", DEC),
        ]  # fmt: skip
        ws = self.sheet("sheet.summary")
        rows = []
        for theme, label, prefix, value, unit, _ in indicators:
            percent = unit == "%"
            rows.append(
                (
                    theme,
                    self.t(label),
                    value.current,
                    value.previous,
                    _change(value.current, value.previous, percent),
                    unit if unit in ("%", "€") else self.t(unit),
                    self.explain(prefix),
                )
            )
        headers = [
            (self.t("summary.theme"), None),
            (self.t("summary.indicator"), None),
            (self.period(data.since, data.until), None),
            (self.t("summary.previous"), None),
            (self.t("summary.change"), None),
            (self.t("summary.unit"), None),
            (self.t("summary.comment"), None),
        ]
        self.table(
            ws,
            "Synthese",
            headers,
            rows,
            [None] * 7,
            widths={1: 14, 2: 44, 3: 26, 4: 20, 5: 12, 6: 14, 7: 80},
            wrap_last=True,
        )
        for r, (*_, unit, number_format) in enumerate(indicators, 2):
            ws.cell(r, 3).number_format = number_format
            ws.cell(r, 4).number_format = number_format
            ws.cell(r, 5).number_format = POINTS if unit == "%" else DELTA

    def over_time(self) -> None:
        activity = self.data.activity
        ws = self.sheet("sheet.overTime")
        hourly = activity.bucket == "hour"
        date_format = DATE_FORMAT[self.lang] + (" hh:mm" if hourly else "")
        series = [
            ("visitors", "charts.activityOverTime.visitors", "funnel.steps.visitors"),
            ("active_users", "charts.activityOverTime.activeUsers", "tiles.activeUsers"),
            ("conversations", "tiles.conversations.label", "tiles.conversations"),
            ("messages", "tiles.messages.label", "tiles.messages"),
            ("cost_eur", "tiles.cost.label", "tiles.cost"),
        ]  # fmt: skip
        rows = [
            (
                datetime.fromisoformat(point.start),
                *(point.values.get(key) for key, _, _ in series),
            )
            for point in activity.points
        ]
        headers = [(self.t("date"), None)] + [
            (self.t(label), self.explain(prefix)) for _, label, prefix in series
        ]
        self.table(
            ws,
            "Evolution",
            headers,
            rows,
            [date_format, INT, INT, INT, INT, EUR_CENTS],
            widths={1: 18 if hourly else 14},
        )

    def adoption(self) -> None:
        adoption = self.data.adoption
        ws = self.sheet("sheet.adoption")
        self.heading(ws, "A1", self.t("funnel.title"))
        access = adoption.people_with_access
        visitors, active = adoption.visitors.current, adoption.active_users.current

        def share(value: float | None, base: float | None) -> float | None:
            return value / base if value is not None and base else None

        funnel = [
            ("funnel.steps.peopleWithAccess", access, None, 1.0 if access else None),
            ("funnel.steps.visitors", visitors, share(visitors, access), share(visitors, access)),
            ("funnel.steps.activeUsers", active, share(active, visitors), share(active, access)),
        ]  # fmt: skip
        rows = [
            (self.t(f"{key}.label"), people, step, of_access, self.explain(key))
            for key, people, step, of_access in funnel
        ]
        headers = [
            (self.t("funnel.step"), self.explain("funnel")),
            (self.t("funnel.people"), None),
            (self.t("funnel.stepRate"), None),
            (self.t("funnel.accessShare"), None),
            (self.t("summary.comment"), None),
        ]
        self.table(
            ws,
            "Entonnoir",
            headers,
            rows,
            [None, INT, PCT, PCT, None],
            start_row=2,
            widths={1: 26, 3: 34, 4: 26, 5: 80},
            wrap_last=True,
        )

        self.heading(ws, "A8", self.t("tables.cohorts.title"))
        months = max((len(row.returns) for row in adoption.cohorts), default=3)
        note = self.explain("tables.cohorts")
        headers = [
            (self.t("tables.cohorts.month"), note),
            (self.t("tables.cohorts.size"), None),
        ] + [
            (self.t("tables.cohorts.monthPlus").replace("{{n}}", str(n)), note)
            for n in range(1, months + 1)
        ]
        rows = []
        for row in adoption.cohorts:
            year, month = (int(part) for part in row.month.split("-"))
            returns = list(row.returns) + [None] * (months - len(row.returns))
            rows.append((date(year, month, 1), row.size, *returns))
        self.table(
            ws,
            "Retour",
            headers,
            rows,
            ["mmmm yyyy", INT, *([PCT] * months)],
            start_row=10,
            widths={1: 18},
        )
        ws.freeze_panes = None

    def teams(self) -> None:
        data = self.data
        usage = {row.team_id: row for row in data.teams.rows}
        known = {row.team_id: row for row in data.catalog}
        ids = list(dict.fromkeys([*known, *usage]))
        rows = []
        for team_id in ids:
            row = usage.get(team_id)
            members = known[team_id].member_count if team_id in known else None
            active = row.active_members if row else 0
            messages = row.messages if row else 0
            previous = row.messages_previous if row else 0
            name = row.name if row else known[team_id].name
            rows.append(
                (
                    name,
                    members,
                    active,
                    min(1.0, active / members) if members else None,
                    messages,
                    (messages - previous) / previous if previous else None,
                )
            )
        rows.sort(key=lambda r: (-r[4], r[0].casefold()))
        note = self.explain("tables.teams")
        headers = [
            (self.t("tables.teams.team"), None),
            (self.t("tables.teams.members"), None),
            (self.t("tables.teams.active"), note),
            (self.t("tables.teams.share"), note),
            (self.t("tables.teams.messages"), self.explain("tiles.messages")),
            (self.t("tables.teams.change"), None),
        ]
        ws = self.sheet("sheet.teams")
        self.table(
            ws,
            "Equipes",
            headers,
            rows,
            [None, INT, INT, PCT, INT, DELTA],
            widths={1: 30},
        )

    def agents(self) -> None:
        rows = [
            (row.label, row.conversations, row.users, row.messages)
            for row in self.data.agents.top_agents
        ]
        headers = [
            (self.t("tables.agentCost.agent"), None),
            (self.t("tiles.conversations.label"), self.explain("tiles.conversations")),
            (self.t("agents.users"), None),
            (self.t("tiles.messages.label"), self.explain("tiles.messages")),
        ]
        ws = self.sheet("sheet.agents")
        self.table(ws, "Agents", headers, rows, [None, INT, INT, INT], widths={1: 34})

    def capabilities(self) -> None:
        names = SHARED[self.lang]
        rows = [
            (
                names.get(f"capability.{row.capability_id}.name", row.capability_id),
                row.agents_equipped,
                row.agents_used,
                row.agents_used / row.agents_equipped if row.agents_equipped else None,
                row.calls,
                row.users,
                row.failed_calls / row.calls if row.calls else None,
            )
            for row in self.data.agents.capabilities
        ]
        note = self.explain("tables.capabilities")
        headers = [
            (self.t("tables.capabilities.capability"), None),
            (self.t("tables.capabilities.equipped"), note),
            (self.t("tables.capabilities.used"), note),
            (self.t("tables.capabilities.usageRate"), note),
            (self.t("tables.capabilities.calls"), None),
            (self.t("tables.capabilities.users"), None),
            (self.t("tables.capabilities.failed"), None),
        ]
        ws = self.sheet("sheet.capabilities")
        self.table(
            ws,
            "Capacites",
            headers,
            rows,
            [None, INT, INT, PCT, INT, INT, PCT],
            widths={1: 40},
        )

    def models(self) -> None:
        quality = self.data.quality
        ws = self.sheet("sheet.models")
        self.heading(ws, "A1", self.t("tables.models.title"))
        tokens_note = self.explain("tiles.tokens")
        headers = [
            (self.t("tables.models.model"), None),
            (self.t("tokensSent"), tokens_note),
            (self.t("tokensGenerated"), tokens_note),
            (self.t("tables.models.cost"), self.explain("tiles.cost")),
            (self.t("tiles.footprint.label"), self.explain("tiles.footprint")),
        ]
        rows = [
            (m.model, m.input_tokens, m.output_tokens, m.cost_eur, m.co2e_grams / 1000)
            for m in quality.models
        ]
        self.table(
            ws,
            "Modeles",
            headers,
            rows,
            [None, INT, INT, EUR, DEC],
            start_row=2,
            widths={1: 34, 2: 18, 3: 18, 4: 16, 5: 28},
        )

        start = len(rows or [None]) + 5
        self.heading(ws, f"A{start - 1}", self.t("tables.agentCost.title"))
        assistant = self.t("tokenUsage.byAgent.creationAssistant")
        headers = [
            (self.t("tables.agentCost.agent"), self.explain("tables.agentCost")),
            (self.t("tables.agentCost.tokens"), tokens_note),
            (self.t("tables.agentCost.cost"), self.explain("tiles.cost")),
            (self.t("agentCost.users"), None),
            (self.t("tables.agentCost.perUser"), None),
        ]
        rows = [
            (
                assistant if a.label == CREATION_ASSISTANT_LABEL else a.label,
                a.tokens,
                a.cost_eur,
                a.users,
                a.cost_eur_per_user,
            )
            for a in quality.agents
        ]
        self.table(
            ws,
            "CoutAgents",
            headers,
            rows,
            [None, INT, EUR, INT, EUR_CENTS],
            start_row=start,
        )
        ws.freeze_panes = None
