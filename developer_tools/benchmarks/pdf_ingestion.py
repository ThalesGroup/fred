"""Generate a synthetic, redistributable PDF/PPTX corpus and measure real processors.

Run with Knowledge Flow's Python: pdf_ingestion.py generate /tmp/pdf-corpus,
then pdf_ingestion.py measure /tmp/pdf-corpus --engine lite --output /tmp/result.json.
Select the checkout under test with PYTHONPATH. All generated content is Apache-2.0.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import signal
import statistics
import subprocess
import sys
import tempfile
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path


def generate(folder: Path) -> None:
    """Write deterministic fixtures and page-level reference text, independent of extractors."""
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    folder.mkdir(parents=True, exist_ok=True)
    font = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    pdfmetrics.registerFont(TTFont("Corpus", str(font)))
    records = []
    for name, pages in [
        ("text", 1),
        ("text-long", 5),
        ("multipage", 4),
        ("hundred-pages", 100),
        ("table", 1),
        ("columns", 2),
        ("unicode", 1),
        ("malformed-text", 1),
        ("scan", 1),
        ("scan-multi", 3),
    ]:
        path = folder / f"{name}.pdf"
        pdf = canvas.Canvas(
            str(path), pagesize=(600, 800), invariant=True, pageCompression=0
        )
        pdf.setTitle(name)
        expected = []
        for number in range(1, pages + 1):
            lines = [
                f"Page {number} reference ANCHOR{number}",
                "Fred document extraction validation.",
                "Airport operations report. Gate A12. Flight 456.",
            ]
            if name == "unicode":
                lines += ["Aéroport été Noël - Ελληνικά - Привет - €"]
            if name == "malformed-text":
                lines += ["Broken glyph X end."]
            if name in ("table", "columns"):
                lines = [
                    lines[0],
                    "Gate Flight Status",
                    "A12 456 Ready",
                    "B34 789 Waiting",
                ]
            expected.append("\n".join(lines))
            if name.startswith("scan"):
                img = Image.new("RGB", (1200, 1600), "white")
                draw = ImageDraw.Draw(img)
                draw.multiline_text(
                    (80, 100),
                    "\n".join(lines),
                    font=ImageFont.truetype(str(font), 30),
                    fill="black",
                    spacing=20,
                )
                pdf.drawImage(ImageReader(img), 0, 0, width=600, height=800)
            else:
                pdf.setFont("Corpus", 14)
                for row, line in enumerate(lines):
                    if name in ("table", "columns") and row > 0:
                        for col, cell in enumerate(line.split()):
                            pdf.drawString(40 + col * 170, 740 - row * 28, cell)
                        pdf.line(35, 730 - row * 28, 550, 730 - row * 28)
                    else:
                        pdf.drawString(40, 740 - row * 28, line)
            pdf.showPage()
        pdf.save()
        if name == "malformed-text":
            # Keep byte offsets intact while introducing an unpaired UTF-16 CMap value.
            data = path.read_bytes().replace(b"<58> <0058>", b"<58> <D800>")
            path.write_bytes(data)
            expected = [text.replace("glyph X", "glyph ") for text in expected]
        records.append(
            {
                "file": path.name,
                "kind": "ocr" if name.startswith("scan") else "text",
                "pages": expected,
            }
        )
    for mixed in (False, True):
        deck = Presentation()
        for number in range(1, 4):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
            box.text_frame.text = f"Slide {number} ANCHOR{number}\nAirport operations"
            if mixed:
                slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(2), Inches(1)
                )
                picture = Image.new("RGB", (100, 100), "blue")
                from io import BytesIO

                buffer = BytesIO()
                picture.save(buffer, format="PNG")
                buffer.seek(0)
                slide.shapes.add_picture(
                    buffer, Inches(4), Inches(4), Inches(1), Inches(1)
                )
        path = folder / ("slides-mixed.pptx" if mixed else "slides-text.pptx")
        deck.save(path)
        records.append(
            {
                "file": path.name,
                "kind": "pptx",
                "pages": [
                    f"Slide {n} ANCHOR{n}\nAirport operations" for n in range(1, 4)
                ],
            }
        )
    for record in records:
        record["sha256"] = hashlib.sha256(
            (folder / record["file"]).read_bytes()
        ).hexdigest()
    (folder / "manifest.json").write_text(
        json.dumps(records, indent=2), encoding="utf-8"
    )


def extract(path: Path, engine: str, output: Path) -> None:
    """Exercise the installed Fred implementation and retain text or rendered slide sizes."""
    with tempfile.TemporaryDirectory() as work:
        if engine == "render":
            from knowledge_flow_backend.core.processors.input.pptx_markdown_processor.utils.pptx_slide_renderer import (
                render_pdf_pages_to_png,
            )
            from PIL import Image

            images = render_pdf_pages_to_png(path, [1, 2, 3, 999], Path(work))
            sizes = {}
            for number, image_path in images.items():
                with Image.open(image_path) as image:
                    sizes[str(number)] = list(image.size)
            result = {"images": sizes}
        else:
            if engine == "lite":
                from knowledge_flow_backend.core.processors.input.lightweight_markdown_processor.lite_pdf_to_md_processor import (
                    LitePdfToMdProcessor,
                )

                processor = LitePdfToMdProcessor()
            else:
                from knowledge_flow_backend.core.processors.input.lightweight_markdown_processor.lite2_pdf_to_md_processor import (
                    LitePdfToMdExtractor,
                )

                processor = LitePdfToMdExtractor()
            extracted = processor.extract(path)
            result = {
                "text": extracted.markdown,
                "pages": [
                    {"number": p.page_no, "text": p.markdown} for p in extracted.pages
                ],
                "page_count": extracted.page_count,
            }
        output.write_text(json.dumps(result, ensure_ascii=True), encoding="utf-8")


def measure(folder: Path, engine: str, output: Path, jobs: int) -> None:
    """Measure cold/warm samples in fresh processes, including descendant RSS and startup."""
    import psutil

    records = json.loads((folder / "manifest.json").read_text())
    records = [
        r for r in records if r["kind"] == ("pptx" if engine == "render" else "text")
    ]

    if records and len(records) < jobs:
        records *= (jobs + len(records) - 1) // len(records)

    def run(record: dict) -> dict:
        path = folder / record["file"]
        if engine == "render":
            path = path.with_suffix(".pdf")
            if not path.exists():
                raise FileNotFoundError(
                    f"Convert the PPTX with LibreOffice before measuring: {path}"
                )
        with tempfile.TemporaryDirectory() as work:
            result_path = Path(work) / "output.json"
            start = time.perf_counter()
            process = subprocess.Popen(
                [
                    sys.executable,
                    __file__,
                    "extract",
                    str(path),
                    "--engine",
                    engine,
                    "--output",
                    str(result_path),
                ],
                stdout=subprocess.DEVNULL,
                start_new_session=True,
            )
            peak = 0
            monitored = psutil.Process(process.pid)
            while process.poll() is None:
                rss = 0
                try:
                    for child in [monitored, *monitored.children(recursive=True)]:
                        try:
                            rss += child.memory_info().rss
                        except psutil.NoSuchProcess:
                            pass
                except psutil.NoSuchProcess:
                    pass
                peak = max(peak, rss)
                if time.perf_counter() - start > 180:
                    os.killpg(process.pid, signal.SIGKILL)
                    process.wait()
                    raise TimeoutError(path)
                time.sleep(0.01)
            if process.returncode:
                raise RuntimeError(f"Extraction failed: {path}")
            result = json.loads(result_path.read_text())
            recall = []
            for page, expected in zip(result.get("pages", []), record["pages"]):
                wanted = Counter(expected.split())
                actual = Counter(page["text"].split())
                recall.append(sum((wanted & actual).values()) / sum(wanted.values()))
            return {
                "file": record["file"],
                "seconds": time.perf_counter() - start,
                "peak_rss": peak,
                "recall": recall,
                "result": result,
            }

    samples = []
    for iteration in range(7):
        with ThreadPoolExecutor(max_workers=jobs) as pool:
            for sample in pool.map(run, records):
                samples.append({"iteration": iteration, **sample})
    output.write_text(
        json.dumps(
            {
                "engine": engine,
                "jobs": jobs,
                "pythonpath": os.environ.get("PYTHONPATH"),
                "samples": samples,
            },
            indent=2,
        ),
        encoding="utf-8",
    )


def summarize(path: Path, output: Path) -> None:
    """Summarize measured runs and downstream chunk/citation examples."""
    from langchain_core.documents import Document
    from knowledge_flow_backend.core.processors.output.vectorization_processor.semantic_splitter import (
        SemanticSplitter,
    )

    data = json.loads(path.read_text())
    files = {}
    for name in sorted({sample["file"] for sample in data["samples"]}):
        samples = [
            sample
            for sample in data["samples"]
            if sample["file"] == name and sample["iteration"] >= 2
        ]
        result = samples[0]["result"]
        entry = {
            "median_seconds": statistics.median(
                sample["seconds"] for sample in samples
            ),
            "peak_process_tree_rss": max(sample["peak_rss"] for sample in samples),
            "minimum_raw_token_recall": min(
                [value for sample in samples for value in sample["recall"]] or [1]
            ),
        }
        if "text" in result:
            chunks = SemanticSplitter().split(
                Document(page_content=result["text"], metadata={"document_uid": name})
            )
            entry["chunk_count"] = len(chunks)
            entry["source_pages"] = [page["number"] for page in result["pages"]]
            entry["citation_example"] = (
                {"text": chunks[0].page_content, "metadata": chunks[0].metadata}
                if chunks
                else None
            )
        else:
            entry["images"] = result["images"]
        files[name] = entry
    output.write_text(
        json.dumps({"source": str(path), "files": files}, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "action", choices=["generate", "measure", "extract", "summarize"]
    )
    parser.add_argument("path", type=Path)
    parser.add_argument(
        "--engine", choices=["lite", "pymupdf", "render"], default="lite"
    )
    parser.add_argument("--output", type=Path, default=Path("pdf-results.json"))
    parser.add_argument("--jobs", type=int, default=1)
    args = parser.parse_args()
    if args.action == "generate":
        generate(args.path)
    elif args.action == "summarize":
        summarize(args.path, args.output)
    elif args.action == "extract":
        extract(args.path, args.engine, args.output)
    else:
        measure(args.path, args.engine, args.output, args.jobs)
