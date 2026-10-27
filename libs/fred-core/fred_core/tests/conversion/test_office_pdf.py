# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Offline tests for the shared Office→PDF conversion helper.

These do not invoke the real ``soffice`` binary: they stub ``shutil.which`` and
the ``_run_soffice`` runner so the timeout / missing-binary / success behaviours are asserted
deterministically without any external dependency. A separate ``integration``-marked
test exercises the real LibreOffice path when the binary is present.
"""

from __future__ import annotations

import os
import shutil
import subprocess  # nosec B404: only used to construct result/exception stubs in tests
import sys
import threading
import time
from pathlib import Path

import pytest
from fred_core.conversion import convert_office_bytes_to_pdf, convert_office_file_to_pdf
from fred_core.conversion.office_pdf import _run_soffice

_MODULE = "fred_core.conversion.office_pdf"


@pytest.mark.asyncio
async def test_convert_bytes_returns_pdf_on_success(monkeypatch) -> None:
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")

    # Emulate soffice: write a PDF next to the input in the requested --outdir.
    def fake_run(cmd, **kwargs):
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        src = Path(cmd[-1])
        (outdir / src.with_suffix(".pdf").name).write_bytes(b"%PDF-1.5 fake")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    monkeypatch.setattr(f"{_MODULE}._run_soffice", fake_run)

    pdf = await convert_office_bytes_to_pdf(b"fake pptx bytes", suffix=".pptx")

    assert pdf is not None
    assert pdf.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_convert_bytes_hands_the_source_suffix_to_libreoffice(
    monkeypatch,
) -> None:
    """LibreOffice picks its import filter from the source extension, so a Word
    document has to reach ``soffice`` as ``.docx``, not as some fixed deck suffix."""
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")
    seen_suffixes: list[str] = []

    def fake_run(cmd, **kwargs):
        src = Path(cmd[-1])
        seen_suffixes.append(src.suffix)
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        (outdir / src.with_suffix(".pdf").name).write_bytes(b"%PDF-1.5 fake")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    monkeypatch.setattr(f"{_MODULE}._run_soffice", fake_run)

    pdf = await convert_office_bytes_to_pdf(b"fake docx bytes", suffix=".docx")

    assert pdf is not None
    assert seen_suffixes == [".docx"]


@pytest.mark.parametrize(
    ("suffix", "expected_filter"),
    [
        (".pptx", "pdf:impress_pdf_Export:EmbedStandardFonts=True,SelectPdfVersion=1"),
        (".ppt", "pdf:impress_pdf_Export:EmbedStandardFonts=True,SelectPdfVersion=1"),
        (".docx", "pdf:writer_pdf_Export:EmbedStandardFonts=True,SelectPdfVersion=1"),
        (".odt", "pdf:writer_pdf_Export:EmbedStandardFonts=True,SelectPdfVersion=1"),
        (".PPTX", "pdf:impress_pdf_Export:EmbedStandardFonts=True,SelectPdfVersion=1"),
        (".xlsx", "pdf:calc_pdf_Export:EmbedStandardFonts=True,SelectPdfVersion=1"),
        (".rtf", "pdf"),
    ],
)
def test_export_filter_follows_the_source_application(
    monkeypatch, tmp_path, suffix, expected_filter
) -> None:
    """Writer's PDF filter on an Impress deck is a build-dependent coin-flip: some
    LibreOffice builds fall back to the right application, others exit 0 and write
    nothing. A format the map does not know is left to soffice's own choice."""
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")
    filters: list[str] = []

    def fake_run(cmd, **kwargs):
        filters.append(cmd[cmd.index("--convert-to") + 1])
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        src = Path(cmd[-1])
        (outdir / src.with_suffix(".pdf").name).write_bytes(b"%PDF-1.5 fake")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    monkeypatch.setattr(f"{_MODULE}._run_soffice", fake_run)
    src = tmp_path / f"source{suffix}"
    src.write_bytes(b"fake")

    assert convert_office_file_to_pdf(src) is not None
    assert filters == [expected_filter]


def test_each_conversion_gets_its_own_libreoffice_profile(
    monkeypatch, tmp_path
) -> None:
    """`soffice` locks its user profile exclusively: two runs sharing the default one
    make the loser exit 0 having written nothing. Conversions are request-driven, so
    overlapping runs are normal and each needs a private profile."""
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")
    profiles: list[str] = []

    def fake_run(cmd, **kwargs):
        profiles.extend(arg for arg in cmd if arg.startswith("-env:UserInstallation="))
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        src = Path(cmd[-1])
        (outdir / src.with_suffix(".pdf").name).write_bytes(b"%PDF-1.5 fake")
        return subprocess.CompletedProcess(cmd, 0, b"", b"")

    monkeypatch.setattr(f"{_MODULE}._run_soffice", fake_run)

    first = tmp_path / "a.docx"
    first.write_bytes(b"fake")
    second = tmp_path / "b.docx"
    second.write_bytes(b"fake")
    convert_office_file_to_pdf(first)
    convert_office_file_to_pdf(second)

    assert len(profiles) == 2
    # A bare path is silently ignored by soffice — the option takes a file URI.
    assert all(p.startswith("-env:UserInstallation=file://") for p in profiles)
    assert profiles[0] != profiles[1]


@pytest.mark.asyncio
async def test_convert_bytes_returns_none_on_timeout(monkeypatch) -> None:
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")

    def hung_run(cmd, **kwargs):
        # Mirror a soffice call that exceeds the deadline: to_thread surfaces the
        # TimeoutExpired the helper is expected to swallow into a None result.
        raise subprocess.TimeoutExpired(cmd, kwargs.get("timeout", 0))

    monkeypatch.setattr(f"{_MODULE}._run_soffice", hung_run)

    pdf = await convert_office_bytes_to_pdf(
        b"fake pptx bytes", suffix=".pptx", timeout_seconds=0.01
    )

    assert pdf is None


@pytest.mark.asyncio
async def test_convert_bytes_returns_none_when_soffice_missing(monkeypatch) -> None:
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: None)

    def unexpected_run(cmd, **kwargs):  # pragma: no cover - must not be reached
        raise AssertionError("soffice must not run when it is absent")

    monkeypatch.setattr(f"{_MODULE}._run_soffice", unexpected_run)

    pdf = await convert_office_bytes_to_pdf(b"fake pptx bytes", suffix=".pptx")

    assert pdf is None


@pytest.mark.asyncio
async def test_convert_bytes_returns_none_when_no_pdf_produced(monkeypatch) -> None:
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")

    # soffice exits 0 but writes nothing — a real failure mode worth guarding.
    monkeypatch.setattr(
        f"{_MODULE}._run_soffice",
        lambda cmd, **kwargs: subprocess.CompletedProcess(cmd, 0, b"", b""),
    )

    pdf = await convert_office_bytes_to_pdf(b"fake pptx bytes", suffix=".pptx")

    assert pdf is None


def test_convert_file_returns_none_on_called_process_error(
    monkeypatch, tmp_path
) -> None:
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")

    def failing_run(cmd, **kwargs):
        raise subprocess.CalledProcessError(1, cmd, stderr=b"boom")

    monkeypatch.setattr(f"{_MODULE}._run_soffice", failing_run)

    src = tmp_path / "deck.pptx"
    src.write_bytes(b"fake")

    assert convert_office_file_to_pdf(src) is None


def _gone(pid: int) -> bool:
    """Dead: no process, or a zombie only its new parent has left to reap."""
    try:
        state = Path(f"/proc/{pid}/stat").read_text().rsplit(")", 1)[1].split()[0]
    except FileNotFoundError:
        return True
    return state in ("Z", "X")


@pytest.mark.skipif(
    not hasattr(os, "killpg") or not Path("/proc").is_dir(),
    reason="needs POSIX process groups and /proc",
)
def test_a_timeout_kills_the_whole_process_group(tmp_path) -> None:
    """soffice forks soffice.bin: killing the direct child alone left it running."""
    pid_file = tmp_path / "grandchild.pid"
    script = (
        "import subprocess, sys, time\n"
        "child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])\n"
        f"open({str(pid_file)!r}, 'w').write(str(child.pid))\n"
        "time.sleep(60)\n"
    )
    started = time.monotonic()
    with pytest.raises(subprocess.TimeoutExpired):
        _run_soffice([sys.executable, "-c", script], timeout=1.5)
    assert time.monotonic() - started < 10
    grandchild = int(pid_file.read_text())
    deadline = time.monotonic() + 5
    while not _gone(grandchild) and time.monotonic() < deadline:
        time.sleep(0.05)
    assert _gone(grandchild)


@pytest.mark.skipif(not hasattr(os, "getsid"), reason="needs POSIX sessions")
def test_soffice_leads_its_own_group_inside_the_callers_session(tmp_path) -> None:
    """A new session would escape Knowledge Flow's extraction sweep."""
    ids_file = tmp_path / "ids"
    script = f"import os; open({str(ids_file)!r}, 'w').write(f'{{os.getpid()}} {{os.getpgrp()}} {{os.getsid(0)}}')"
    _run_soffice([sys.executable, "-c", script], timeout=10)
    pid, pgid, sid = map(int, ids_file.read_text().split())
    assert pgid == pid
    assert sid == os.getsid(0)


def test_a_failing_command_raises_like_a_checked_run() -> None:
    with pytest.raises(subprocess.CalledProcessError) as excinfo:
        _run_soffice(
            [sys.executable, "-c", "import sys; sys.stderr.write('boom'); sys.exit(3)"],
            timeout=10,
        )
    assert excinfo.value.returncode == 3 and excinfo.value.stderr == b"boom"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("background", "pool"), [(False, "office2pdf_"), (True, "office2pdf-background")]
)
async def test_background_conversions_queue_on_their_own_worker(
    monkeypatch, background, pool
) -> None:
    monkeypatch.setattr(f"{_MODULE}.shutil.which", lambda _: "/usr/bin/soffice")
    threads: list[str] = []

    def fake_run(cmd, **kwargs):
        threads.append(threading.current_thread().name)
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        (outdir / Path(cmd[-1]).with_suffix(".pdf").name).write_bytes(b"%PDF")
        return None

    monkeypatch.setattr(f"{_MODULE}._run_soffice", fake_run)
    pdf = await convert_office_bytes_to_pdf(b"x", suffix=".pptx", background=background)
    assert pdf == b"%PDF"
    assert threads[0].startswith(pool)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_convert_bytes_with_real_libreoffice() -> None:
    """End-to-end conversion via the real ``soffice`` binary (skipped when absent)."""
    if shutil.which("soffice") is None:
        pytest.skip("soffice not installed")
    pytest.importorskip("pptx")

    import io

    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # A default template always carries slide dimensions; the fallbacks keep the
    # call typed without asserting on something the fixture does not test.
    width = prs.slide_width or Emu(9144000)
    height = prs.slide_height or Emu(6858000)
    box = slide.shapes.add_textbox(Emu(0), Emu(0), width, height)
    box.text_frame.text = "Hello preview"
    buf = io.BytesIO()
    prs.save(buf)

    pdf = await convert_office_bytes_to_pdf(buf.getvalue(), suffix=".pptx")

    assert pdf is not None
    assert pdf.startswith(b"%PDF")
