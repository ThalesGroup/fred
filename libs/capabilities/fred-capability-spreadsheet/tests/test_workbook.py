# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Workbook builder: every generated file is read back with openpyxl."""

from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime
from typing import Any

import pytest
from fred_capability_spreadsheet.workbook import (
    FORMULA_REJECTED,
    LEAKED_ROW_REJECTED,
    MAX_CELL_CHARS,
    MAX_PREVIEW_ROWS,
    CreateSpreadsheetInput,
    SpreadsheetInputError,
    build_preview,
    build_workbook,
    derived_column_hints,
    manual_total_hints,
    prepare,
    sanitize_file_name,
)
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def _request(
    *sheets: dict[str, Any], language: str | None = None
) -> CreateSpreadsheetInput:
    return CreateSpreadsheetInput.model_validate(
        {"file_name": "test", "sheets": list(sheets), "language": language}
    )


def _xlsx(*sheets: dict[str, Any], language: str | None = None) -> bytes:
    return build_workbook(prepare(_request(*sheets, language=language)))


def _sheet(name: str = "Data", columns=None, rows=None) -> dict[str, Any]:
    return {
        "name": name,
        "columns": columns or [{"header": "Item"}],
        "rows": rows if rows is not None else [],
    }


def _build(*sheets: dict[str, Any], language: str | None = None) -> Any:
    return load_workbook(io.BytesIO(_xlsx(*sheets, language=language)))


def _ws(*sheets: dict[str, Any], language: str | None = None) -> Worksheet:
    sheet = _build(*sheets, language=language).worksheets[0]
    assert isinstance(sheet, Worksheet)
    return sheet


def _error(*sheets: dict[str, Any]) -> str:
    with pytest.raises(SpreadsheetInputError) as info:
        _xlsx(*sheets)
    return str(info.value)


TYPED_COLUMNS = [
    {"header": "Name"},
    {"header": "Qty", "type": "integer"},
    {"header": "Price", "type": "number"},
    {"header": "Rate", "type": "percent"},
    {"header": "Amount", "type": "currency", "currency": "EUR"},
    {"header": "Due", "type": "date"},
]


def test_typed_columns_get_values_and_number_formats():
    ws = _ws(
        _sheet(
            columns=TYPED_COLUMNS,
            rows=[["Bolt", 3, "1 234,5", "15%", 1234.5, "2026-10-21"]],
        )
    )
    assert [c.value for c in ws[1]] == ["Name", "Qty", "Price", "Rate", "Amount", "Due"]
    name, qty, price, rate, amount, due = ws[2]
    assert name.value == "Bolt"
    assert (qty.value, qty.number_format) == (3, "#,##0")
    assert (price.value, price.number_format) == (1234.5, "#,##0.00")
    assert (rate.value, rate.number_format) == (0.15, "0%")
    assert amount.value == 1234.5
    assert amount.number_format == '#,##0.00 "€"'
    assert due.value == datetime(2026, 10, 21)
    # Built-in short date (numFmtId 14): Excel follows the reader's locale.
    assert due.number_format == "mm-dd-yy"


def test_dates_use_the_builtin_short_date_format():
    data = _xlsx(
        _sheet(columns=[{"header": "Due", "type": "date"}], rows=[["2026-10-21"]])
    )
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        styles = archive.read("xl/styles.xml").decode()
    assert 'numFmtId="14"' in styles
    assert "mm-dd-yy" not in styles  # no custom format entry


def test_percent_keeps_needed_decimals():
    ws = _ws(_sheet(columns=[{"header": "Rate", "type": "percent"}], rows=[[0.125]]))
    assert ws["A2"].number_format == "0.0%"


@pytest.mark.parametrize(
    ("currency", "fmt"),
    [
        ("USD", '"$"#,##0.00'),
        ("gbp", '"£"#,##0.00'),
        ("JPY", '"¥"#,##0'),
        (None, "#,##0.00"),
    ],
)
def test_currency_formats(currency, fmt):
    column = {"header": "Amount", "type": "currency", "currency": currency}
    ws = _ws(_sheet(columns=[column], rows=[[10]]))
    assert ws["A2"].number_format == fmt


def test_enum_values_are_case_insensitive():
    request = _request(
        _sheet(columns=[{"header": "N", "type": "Number", "total": "SUM"}], rows=[[1]])
    )
    column = request.sheets[0].columns[0]
    assert (column.type, column.total) == ("number", "sum")


def test_header_is_bold_frozen_and_filtered():
    ws = _ws(_sheet(columns=TYPED_COLUMNS[:2], rows=[["a", 1], ["b", 2]]))
    assert all(cell.font.b for cell in ws[1])
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"


def test_widths_are_declared_or_computed_and_capped():
    columns = [
        {"header": "Short"},
        {"header": "Fixed", "width": 30},
        {"header": "Long"},
    ]
    ws = _ws(_sheet(columns=columns, rows=[["abc", "x", "y" * 200]]))
    assert ws.column_dimensions["A"].width == len("Short") + 2
    assert ws.column_dimensions["B"].width == 30
    assert ws.column_dimensions["C"].width == 60


def test_totals_row_holds_formulas_with_column_formats():
    columns = [
        {"header": "Item"},
        {"header": "Qty", "type": "integer", "total": "average"},
        {"header": "Amount", "type": "currency", "currency": "EUR", "total": "sum"},
        {"header": "Min", "type": "number", "total": "min"},
        {"header": "Max", "type": "number", "total": "max"},
        {"header": "Label", "total": "count"},
    ]
    rows = [[f"i{n}", n, n * 10, n, n, "x"] for n in range(1, 11)]
    ws = _ws(_sheet(columns=columns, rows=rows))
    total = [c.value for c in ws[12]]
    assert total == [
        "Moyenne / Total / Min / Max / Nombre",
        "=AVERAGE(B2:B11)",
        "=SUM(C2:C11)",
        "=MIN(D2:D11)",
        "=MAX(E2:E11)",
        "=COUNTA(F2:F11)",
    ]
    assert ws["B12"].number_format == "#,##0.00"
    assert ws["C12"].number_format == '#,##0.00 "€"'
    assert ws["F12"].number_format == "#,##0"
    assert ws["A12"].font.b
    # The totals row stays outside the filter range.
    assert ws.auto_filter.ref == "A1:F11"


def test_no_totals_row_without_data():
    ws = _ws(_sheet(columns=[{"header": "N", "type": "number", "total": "sum"}]))
    assert ws.max_row == 1


def test_formula_like_text_is_stored_as_literal_text():
    payload = "=== Section ==="
    ws = _ws(
        _sheet(columns=[{"header": "Note"}], rows=[[payload], ["+1"], ["@SUM(A1)"]])
    )
    for row, expected in ((2, payload), (3, "+1"), (4, "@SUM(A1)")):
        cell = ws.cell(row=row, column=1)
        assert cell.value == expected
        assert cell.data_type == "s"
        assert cell.quotePrefix


def test_multiple_sheets_and_empty_cells():
    workbook = _build(
        _sheet("Q1", rows=[["a"], [None], [""]]),
        _sheet("Q2", columns=[{"header": "N", "type": "number"}], rows=[[1]]),
    )
    assert workbook.sheetnames == ["Q1", "Q2"]
    q1 = workbook["Q1"]
    assert q1["A3"].value is None and q1["A4"].value is None
    assert workbook.properties.creator == "Fred"


def test_text_column_renders_non_strings():
    ws = _ws(_sheet(rows=[[12.0], [True], [3.5]]))
    assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == ["12", "TRUE", "3.5"]


@pytest.mark.parametrize(
    ("sheets", "fragment"),
    [
        ([_sheet(name=" ")], "sheet name is empty"),
        ([_sheet(name="x" * 32)], "limit is 31"),
        ([_sheet(name="a/b")], "cannot contain"),
        ([_sheet(name="'quoted'")], "cannot contain"),
        ([_sheet("Data"), _sheet("data")], "used twice"),
        ([{"name": "Data", "columns": [], "rows": []}], "has no columns"),
    ],
)
def test_structure_errors(sheets, fragment):
    assert fragment in _error(*sheets)


def test_row_length_mismatch_names_sheet_and_row():
    message = _error(
        _sheet("Sales", columns=TYPED_COLUMNS[:3], rows=[["a", 1, 2], ["b", 1, 2, 3]])
    )
    assert "Sheet 'Sales', row 3: expected 3 values (one per column), got 4" in message


@pytest.mark.parametrize(
    ("column", "value"),
    [
        ({"header": "Amount", "type": "number"}, "=SUM(B2:B6)"),
        ({"header": "Share", "type": "percent"}, "=B2/SUM($B$2:$B$6)"),
        ({"header": "Cost", "type": "currency"}, "=1+1"),
        ({"header": "Qty", "type": "integer"}, " =B2*2"),
        ({"header": "Note"}, '=HYPERLINK("http://example.test","x")'),
        ({"header": "Note"}, "= sum(B2:B6)"),
        ({"header": "Note"}, "=$B$2"),
        ({"header": "Note"}, "=B2/B3"),
    ],
)
def test_formulas_are_rejected_with_an_actionable_message(column, value):
    message = _error(
        _sheet("Sales", columns=[{"header": "Name"}, column], rows=[["a", value]])
    )
    assert f"Sheet 'Sales', row 2, column 2 ('{column['header']}')" in message
    assert FORMULA_REJECTED in message


@pytest.mark.parametrize(
    ("column", "value", "fragment"),
    [
        ({"header": "Qty", "type": "integer"}, 1.5, "a whole number"),
        ({"header": "Due", "type": "date"}, "21/10/2026", "YYYY-MM-DD"),
        ({"header": "Price", "type": "number"}, "cheap", "a number"),
        ({"header": "Rate", "type": "percent"}, "high", "0.15"),
        ({"header": "Price", "type": "number"}, True, "a number"),
    ],
)
def test_conversion_errors_name_sheet_row_and_column(column, value, fragment):
    message = _error(
        _sheet("Sales", columns=[{"header": "Name"}, column], rows=[["a", value]])
    )
    assert f"Sheet 'Sales', row 2, column 2 ('{column['header']}')" in message
    assert fragment in message


def test_numeric_total_on_text_column_is_rejected():
    message = _error(_sheet(columns=[{"header": "Name", "total": "sum"}], rows=[["a"]]))
    assert "needs a numeric column" in message


def test_limits():
    assert "limit is 10 sheets" in _error(*[_sheet(f"S{i}") for i in range(11)])
    wide = [{"header": f"c{i}"} for i in range(51)]
    assert "limit is 50 columns" in _error(_sheet(columns=wide))
    tall = _sheet(rows=[["x"]] * 5001)
    assert "limit is 5000 rows" in _error(tall)
    cells = _sheet(
        columns=[{"header": f"c{i}"} for i in range(20)], rows=[["x"] * 20] * 2600
    )
    assert "limit is 50000 cells" in _error(cells)
    assert f"limit is {MAX_CELL_CHARS}" in _error(
        _sheet(rows=[["x" * (MAX_CELL_CHARS + 1)]])
    )


def test_limits_are_inclusive():
    _xlsx(_sheet(rows=[["x"]] * 5000))


def test_limit_errors_ask_the_model_to_tell_the_user_and_split():
    message = _error(_sheet(rows=[["x"]] * 5001))
    assert "Tell the user plainly about this limit" in message
    assert "splitting" in message


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Q3 budget", "Q3 budget"),
        ("../etc/passwd", "passwd"),
        ("report.xlsx", "report"),
        ('a<b>:"c', "abc"),
        ("Q3 #1 100%", "Q3 1 100"),
        ("   ", "spreadsheet"),
    ],
)
def test_sanitize_file_name(raw, expected):
    assert sanitize_file_name(raw) == expected


def _preview(*sheets: dict[str, Any], language: str | None = None) -> dict[str, Any]:
    return json.loads(build_preview(prepare(_request(*sheets, language=language))))


def test_preview_carries_typed_values_and_display_formats():
    columns = [
        *TYPED_COLUMNS[:4],
        {"header": "Amount", "type": "currency", "currency": "JPY"},
        {"header": "Due", "type": "date"},
    ]
    rows = [["Pen", "3", "1,5", "12.5%", 1200, "2026-01-31T10:00:00"]]
    sheet = _preview(_sheet(columns=columns, rows=rows))["sheets"][0]

    assert sheet["name"] == "Data"
    assert sheet["rows"] == [["Pen", 3, 1.5, 0.125, 1200, "2026-01-31"]]
    assert sheet["row_count"] == 1
    assert sheet["totals"] is None
    assert sheet["columns"] == [
        {"header": "Name", "type": "text"},
        {"header": "Qty", "type": "integer", "decimals": 0},
        {"header": "Price", "type": "number", "decimals": 2},
        {"header": "Rate", "type": "percent", "decimals": 1},
        {"header": "Amount", "type": "currency", "decimals": 0, "currency": "JPY"},
        {"header": "Due", "type": "date"},
    ]


def test_preview_totals_mirror_the_excel_functions():
    columns = [
        {"header": "Line"},
        {"header": "Sum", "type": "currency", "currency": "EUR", "total": "sum"},
        {"header": "Avg", "type": "integer", "total": "average"},
        {"header": "Min", "type": "number", "total": "min"},
        {"header": "Max", "type": "percent", "total": "max"},
        {"header": "Count", "type": "text", "total": "count"},
    ]
    rows = [
        ["a", 1200, 1, 4.5, 0.1, "x"],
        ["b", 350.5, 2, None, "25%", None],
        ["c", None, None, -2, None, "y"],
    ]
    sheet = _preview(_sheet(columns=columns, rows=rows))["sheets"][0]
    assert sheet["totals"] == [
        "Total / Moyenne / Min / Max / Nombre",
        1550.5,
        1.5,
        -2,
        0.25,
        2,
    ]
    assert sheet["columns"][1]["total"] == "sum"


def test_preview_total_of_an_empty_numeric_column():
    columns = [
        {"header": "Sum", "type": "number", "total": "sum"},
        {"header": "Avg", "type": "number", "total": "average"},
    ]
    sheet = _preview(_sheet(columns=columns, rows=[[None, None]]))["sheets"][0]
    assert sheet["totals"] == [0, None]


def test_preview_caps_rows_but_totals_cover_every_row():
    count = MAX_PREVIEW_ROWS + 20
    columns = [{"header": "N", "type": "integer", "total": "sum"}]
    rows = [[1] for _ in range(count)]
    sheet = _preview(_sheet(columns=columns, rows=rows))["sheets"][0]
    assert len(sheet["rows"]) == MAX_PREVIEW_ROWS
    assert sheet["row_count"] == count
    assert sheet["totals"] == [count]


def test_preview_keeps_every_sheet_and_formula_like_text():
    preview = _preview(_sheet("One", rows=[["=1+1"]]), _sheet("Two", rows=[["été"]]))
    assert preview["version"] == 1
    assert [s["name"] for s in preview["sheets"]] == ["One", "Two"]
    assert preview["sheets"][0]["rows"] == [["=1+1"]]
    assert preview["sheets"][1]["rows"] == [["été"]]


@pytest.mark.parametrize(
    ("language", "total", "label"),
    [
        (None, "sum", "Total"),
        ("fr", "average", "Moyenne"),
        ("fr", "count", "Nombre"),
        ("en", "average", "Average"),
        ("en-GB", "count", "Count"),
        ("EN", "max", "Max"),
        ("de", "min", "Min"),
    ],
)
def test_totals_label_follows_the_function_and_the_language(language, total, label):
    columns = [{"header": "Item"}, {"header": "N", "type": "number", "total": total}]
    sheet = _sheet(columns=columns, rows=[["a", 1], ["b", 2]])
    assert _ws(sheet, language=language)["A4"].value == label
    assert _preview(sheet, language=language)["sheets"][0]["totals"][0] == label


def test_control_characters_are_stripped_from_headers_and_cells():
    columns = [{"header": "Na\x0bme"}, {"header": "Note"}]
    rows = [["Bo\x00lt\x1f", "\x08"], ["ok\ttab", "line\nbreak"]]
    ws = _ws(_sheet(columns=columns, rows=rows))
    assert ws["A1"].value == "Name"
    assert ws["A2"].value == "Bolt"
    assert ws["B2"].value is None  # nothing left once stripped
    assert (ws["A3"].value, ws["B3"].value) == ("ok\ttab", "line\nbreak")
    preview = _preview(_sheet(columns=columns, rows=rows))["sheets"][0]
    assert preview["columns"][0]["header"] == "Name"
    assert preview["rows"][0] == ["Bolt", None]


def test_control_characters_in_a_sheet_name_are_rejected():
    assert "control characters" in _error(_sheet(name="Q\x013"))


def test_history_sheet_name_is_reserved():
    assert "reserved by Excel" in _error(_sheet(name="history"))
    assert "reserved by Excel" in _error(_sheet(name="History"))


def test_header_longer_than_a_cell_is_rejected():
    message = _error(_sheet("Sales", columns=[{"header": "h" * (MAX_CELL_CHARS + 1)}]))
    assert "Sheet 'Sales', column 1: the header has" in message
    assert f"limit is {MAX_CELL_CHARS}" in message


@pytest.mark.parametrize(
    ("raw", "column_type"),
    [
        ("1,234", "number"),
        ("$1,500", "currency"),
        ("-12,000", "integer"),
        ("1,234%", "percent"),
    ],
)
def test_a_lone_thousands_like_comma_is_ambiguous(raw, column_type):
    column = {"header": "N", "type": column_type}
    message = _error(_sheet(columns=[column], rows=[[raw]]))
    assert "is ambiguous" in message
    assert "JSON numbers" in message


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("1,5", 1.5), ("12,34", 12.34), ("1.234,5", 1234.5), ("1,234.5", 1234.5)],
)
def test_unambiguous_separators_still_parse(raw, expected):
    ws = _ws(_sheet(columns=[{"header": "N", "type": "number"}], rows=[[raw]]))
    assert ws["A2"].value == expected


@pytest.mark.parametrize(
    "raw", [float("nan"), float("inf"), float("-inf"), "nan", "Infinity"]
)
def test_non_finite_numbers_are_rejected(raw):
    for column_type in ("number", "currency", "percent"):
        message = _error(
            _sheet(columns=[{"header": "N", "type": column_type}], rows=[[raw]])
        )
        assert "row 2, column 1" in message


def test_conversion_happens_once_for_file_and_preview(monkeypatch):
    from fred_capability_spreadsheet import workbook

    calls = 0
    original = workbook._convert

    def counting(*args):
        nonlocal calls
        calls += 1
        return original(*args)

    monkeypatch.setattr(workbook, "_convert", counting)
    prepared = prepare(_request(_sheet(rows=[["a"], ["b"]])))
    build_workbook(prepared)
    build_preview(prepared)
    assert calls == 2


# --- computed columns ---------------------------------------------------------

AMOUNTS = [
    {"header": "Poste"},
    {"header": "Budget", "type": "currency", "currency": "EUR", "total": "sum"},
    {"header": "Réel", "type": "currency", "currency": "EUR", "total": "sum"},
]


def _computed(*extra: dict[str, Any], rows=None, name: str = "Data") -> dict[str, Any]:
    return _sheet(
        name,
        AMOUNTS + list(extra),
        rows if rows is not None else [["A", 100, 120], ["B", 0, 30], ["C", None, 50]],
    )


@pytest.mark.parametrize(
    ("compute", "formula", "values"),
    [
        (
            {"op": "subtract", "left": "Réel", "right": "Budget"},
            "=C2-B2",
            [20.0, 30.0, 50.0],
        ),
        ({"op": "add", "left": "Budget", "right": 10}, "=B2+10", [110.0, 10.0, 10.0]),
        (
            {"op": "multiply", "left": "Réel", "right": 1.2},
            "=C2*1.2",
            [144.0, 36.0, 60.0],
        ),
        (
            {"op": "multiply", "left": "Réel", "right": -2},
            "=C2*(-2)",
            [-240.0, -60.0, -100.0],
        ),
        (
            {"op": "divide", "left": "Réel", "right": "Budget"},
            '=IFERROR(C2/B2,"")',
            [1.2, None, None],
        ),
        (
            {"op": "divide", "left": "Réel", "right": 3, "round": 1},
            '=IFERROR(ROUND(C2/3,1),"")',
            [40.0, 10.0, 16.7],
        ),
        (
            {"op": "share_of_total", "column": "Réel"},
            '=IFERROR(C2/SUM($C$2:$C$4),"")',
            [0.6, 0.15, 0.25],
        ),
        ({"op": "*", "left": "réel", "right": "2"}, "=C2*2", [240.0, 60.0, 100.0]),
        (
            {"op": "multiply", "left": "Réel", "right": "15%"},
            "=C2*0.15",
            [18.0, 4.5, 7.5],
        ),
    ],
)
def test_computed_column_formulas_and_preview_values(compute, formula, values):
    sheet = _computed({"header": "Calc", "type": "number", "compute": compute})
    assert _ws(sheet)["D2"].value == formula
    computed = [row[3] for row in _preview(sheet)["sheets"][0]["rows"]]
    assert computed == [pytest.approx(v) if v is not None else None for v in values]


def test_computed_formulas_use_each_row_and_sheet_addresses():
    first = _computed(
        {
            "header": "Écart",
            "compute": {"op": "subtract", "left": "Réel", "right": "Budget"},
        }
    )
    second = _sheet(
        "Other",
        [
            {"header": "Qty", "type": "integer"},
            {"header": "Share", "compute": {"op": "share_of_total", "column": "Qty"}},
        ],
        [[1], [3]],
    )
    book = _build(first, second)
    ws1, ws2 = book["Data"], book["Other"]
    assert [ws1[f"D{r}"].value for r in (2, 3, 4)] == ["=C2-B2", "=C3-B3", "=C4-B4"]
    assert ws2["B2"].value == '=IFERROR(A2/SUM($A$2:$A$3),"")'
    assert ws2["B3"].value == '=IFERROR(A3/SUM($A$2:$A$3),"")'
    # Untyped computed columns are numeric; share_of_total defaults to percent.
    assert ws1["D2"].number_format == "#,##0.00"
    assert ws2["B2"].number_format == "0%"


def test_totals_and_preview_totals_cover_computed_columns():
    sheet = _computed(
        {
            "header": "Écart",
            "compute": {"op": "subtract", "left": "Réel", "right": "Budget"},
            "total": "sum",
        },
        {
            "header": "Ratio",
            "compute": {"op": "divide", "left": "Réel", "right": "Budget"},
            "total": "count",
        },
        {
            "header": "Part",
            "compute": {"op": "share_of_total", "column": "Réel"},
            "total": "sum",
        },
    )
    ws = _ws(sheet)
    assert ws["D5"].value == "=SUM(D2:D4)"
    assert ws["E5"].value == "=COUNT(E2:E4)"
    assert ws["F5"].value == "=SUM(F2:F4)"
    preview = _preview(sheet)["sheets"][0]
    assert preview["totals"][3:] == [100.0, 1, pytest.approx(1.0)]
    assert preview["columns"][5]["type"] == "percent"


def test_computed_columns_chain_and_propagate_empty_results():
    sheet = _computed(
        {
            "header": "Ratio",
            "compute": {"op": "divide", "left": "Réel", "right": "Budget"},
        },
        {
            "header": "Double",
            "compute": {"op": "multiply", "left": "Ratio", "right": 2},
        },
        {"header": "Plus", "compute": {"op": "add", "left": "Réel", "right": 1}},
        {
            "header": "Écart",
            "compute": {"op": "subtract", "left": "Plus", "right": "Réel"},
        },
    )
    ws = _ws(sheet)
    # A failing source makes the dependent formula fail too, so it is wrapped.
    assert ws["E2"].value == '=IFERROR(D2*2,"")'
    assert ws["G2"].value == "=F2-C2"
    rows = _preview(sheet)["sheets"][0]["rows"]
    assert [row[4] for row in rows] == [2.4, None, None]
    assert [row[6] for row in rows] == [1.0, 1.0, 1.0]


def test_rows_may_send_null_placeholders_or_omit_computed_columns():
    compute = {"op": "multiply", "left": "Réel", "right": 2}
    column = {"header": "x2", "compute": compute}
    with_nulls = _computed(column, rows=[["A", 1, 2, None]])
    without = _computed(column, rows=[["A", 1, 2]])
    # A value sent for a computed column (even a formula) is ignored.
    stale = _computed(column, rows=[["A", 1, 2, "=C2*2"]])
    for sheet in (with_nulls, without, stale):
        assert _ws(sheet)["D2"].value == "=C2*2"
        assert _preview(sheet)["sheets"][0]["rows"] == [["A", 1, 2, 4.0]]
    assert "null for an empty cell and for a computed column" in _error(
        _computed(column, rows=[["A", 1]])
    )


def test_round_matches_excel_half_away_from_zero():
    sheet = _sheet(
        "Data",
        [
            {"header": "V", "type": "number"},
            {
                "header": "R",
                "compute": {"op": "add", "left": "V", "right": 0, "round": 2},
            },
        ],
        [[2.675], [-0.125], [1.005]],
    )
    assert _ws(sheet)["B2"].value == "=ROUND(A2+0,2)"
    assert [row[1] for row in _preview(sheet)["sheets"][0]["rows"]] == [
        2.68,
        -0.13,
        1.01,
    ]


@pytest.mark.parametrize(
    ("compute", "fragment"),
    [
        (
            {"op": "power", "left": "Réel", "right": 2},
            "compute op 'power' is not supported",
        ),
        (
            {"op": "add", "left": "Montant", "right": 1},
            "'Montant', which is not a column header",
        ),
        ({"op": "add", "left": "Poste", "right": 1}, "'Poste', a text column"),
        ({"op": "divide", "left": "Réel", "right": 0}, "divide by 0 is not possible"),
        ({"op": "divide", "left": "Réel", "right": "0"}, "divide by 0 is not possible"),
        ({"op": "add", "left": "Réel"}, "add needs `left` and `right`"),
        ({"op": "add", "left": 1, "right": 2}, "at least one column header"),
        ({"op": "share_of_total"}, "share_of_total needs `column`"),
        ({"op": "share_of_total", "column": "12"}, "share_of_total needs `column`"),
        ({"op": "add", "left": "Calc", "right": 1}, "refers to its own column"),
        (
            {"op": "add", "left": "Réel", "right": 1, "round": 11},
            "round must be between 0 and 10",
        ),
    ],
)
def test_compute_errors_name_sheet_column_and_fix(compute, fragment):
    message = _error(_computed({"header": "Calc", "compute": compute}))
    assert message.startswith("Sheet 'Data', column 4 ('Calc'): ")
    assert fragment in message


def test_compute_rejects_forward_references_text_type_and_duplicate_headers():
    forward = _computed(
        {"header": "A2", "compute": {"op": "add", "left": "B2", "right": 1}},
        {"header": "B2", "compute": {"op": "add", "left": "Réel", "right": 1}},
    )
    assert "'B2', a computed column placed after it" in _error(forward)
    text = _computed(
        {
            "header": "T",
            "type": "text",
            "compute": {"op": "add", "left": "Réel", "right": 1},
        }
    )
    assert "must be number, integer, percent or currency, not 'text'" in _error(text)
    duplicate = _sheet(
        "Data",
        [
            {"header": "V", "type": "number"},
            {"header": "V", "type": "number"},
            {"header": "W", "compute": {"op": "add", "left": "V", "right": 1}},
        ],
        [[1, 2]],
    )
    assert "used by several columns" in _error(duplicate)


def test_free_formula_refusal_points_to_compute():
    message = _error(_sheet("Data", [{"header": "V", "type": "number"}], [["=B2*2"]]))
    assert "set `compute` on the column for a per-row calculation" in message


# Live case: a stray quote closed "Loyer" early, so this cell swallowed the rest
# of its row and the next row's start, and the next row lost its label.
LEAKED_ROWS = [
    ['Loyer", 1200, 1150, null, null, null], [', 8000, 8400, None],
    ["Matériel", 0, 500, None],
]


def test_a_text_cell_holding_leaked_row_json_is_rejected():
    sheet = _computed(
        {
            "header": "Écart",
            "compute": {"op": "subtract", "left": "Réel", "right": "Budget"},
        },
        rows=LEAKED_ROWS,
    )
    message = _error(sheet)
    assert "row 2, column 1 ('Poste')" in message
    assert LEAKED_ROW_REJECTED in message


@pytest.mark.parametrize(
    "raw",
    [
        'Matériel", 0, 500, null]',
        'A", "B", null], ["C',
        'x",null],[',
    ],
)
def test_leaked_row_fragments_are_caught(raw):
    assert LEAKED_ROW_REJECTED in _error(_sheet(rows=[[raw]]))


@pytest.mark.parametrize(
    "raw",
    [
        'Le "Total", 3 lignes',
        'Projet "Alpha", phase 2 [voir annexe]',
        "Liste [a, b], [c]",
        "null",
    ],
)
def test_ordinary_text_with_quotes_and_brackets_is_kept(raw):
    assert _ws(_sheet(rows=[[raw]]))["A2"].value == raw


def test_preview_carries_formula_templates_and_total_formulas():
    sheet = _computed(
        {
            "header": "Taux",
            "type": "percent",
            "compute": {"op": "divide", "left": "Réel", "right": "Budget", "round": 2},
            "total": "average",
        },
        {"header": "Part", "compute": {"op": "share_of_total", "column": "Réel"}},
    )
    columns = _preview(sheet)["sheets"][0]["columns"]
    ws = _ws(sheet)
    assert [c.get("formula") for c in columns] == [
        None,
        None,
        None,
        '=IFERROR(ROUND(C{row}/B{row},2),"")',
        '=IFERROR(C{row}/SUM($C$2:$C$4),"")',
    ]
    # The template with the row number put in is exactly the file's formula.
    for row in (2, 3, 4):
        assert columns[3]["formula"].replace("{row}", str(row)) == ws[f"D{row}"].value
        assert columns[4]["formula"].replace("{row}", str(row)) == ws[f"E{row}"].value
    assert columns[3]["total_formula"] == ws["D5"].value == "=AVERAGE(D2:D4)"
    assert columns[1]["total_formula"] == "=SUM(B2:B4)"
    assert "total_formula" not in columns[0] and "formula" not in columns[0]


def test_preview_total_formulas_cover_every_row_and_count_kinds():
    sheet = _sheet(
        columns=[
            {"header": "Name", "total": "count"},
            {"header": "Qty", "type": "integer", "total": "sum"},
        ],
        rows=[["x", 1]] * (MAX_PREVIEW_ROWS + 3),
    )
    columns = _preview(sheet)["sheets"][0]["columns"]
    last = MAX_PREVIEW_ROWS + 4
    assert columns[0]["total_formula"] == f"=COUNTA(A2:A{last})"
    assert columns[1]["total_formula"] == f"=SUM(B2:B{last})"


# --- literal derived columns ---------------------------------------------------

# The live-test request: every derived column sent as rounded literals.
_BUDGET_COLUMNS = [
    {"header": "Poste"},
    {"header": "Budget", "type": "currency", "total": "sum"},
    {"header": "Réel", "type": "currency", "total": "sum"},
    {"header": "Écart", "type": "currency", "total": "sum"},
    {"header": "Taux de réalisation", "type": "percent"},
    {"header": "Part du total du réel", "type": "percent", "total": "sum"},
]
_BUDGET_ROWS = [
    ["Loyer", 1200, 1150, -50, 0.9583, 0.1144],
    ["Salaires", 8000, 8400, 400, 1.05, 0.8358],
    ["Matériel", 0, 500, 500, None, 0.0498],
]


def _hints(*sheets: dict[str, Any]) -> list[str]:
    return derived_column_hints(prepare(_request(*sheets)))


def test_literal_difference_ratio_and_share_are_named_with_their_compute():
    hints = _hints(_sheet(columns=_BUDGET_COLUMNS, rows=_BUDGET_ROWS))

    assert hints == [
        "'Écart' -> compute {op: 'subtract', left: 'Réel', right: 'Budget'}",
        "'Taux de réalisation' -> compute {op: 'divide', left: 'Réel', right: 'Budget'}",
        "'Part du total du réel' -> compute {op: 'share_of_total', column: 'Réel'}",
    ]


def test_computed_and_independent_columns_get_no_hint():
    columns = [
        {"header": "Item"},
        {"header": "Qty", "type": "integer"},
        {"header": "Price", "type": "currency"},
        {
            "header": "Amount",
            "type": "currency",
            "compute": {"op": "*", "left": "Qty", "right": "Price"},
        },
        {"header": "Weight", "type": "number"},
    ]
    rows = [["a", 2, 10, None, 3.7], ["b", 3, 15, None, 1.2], ["c", 1, 7, None, 9]]

    assert _hints(_sheet(columns=columns, rows=rows)) == []


def test_a_single_matching_row_or_a_copied_column_is_no_derivation():
    columns = [
        {"header": "A", "type": "number"},
        {"header": "Zero", "type": "number"},
        {"header": "B", "type": "number"},
    ]

    assert _hints(_sheet(columns=columns, rows=[[4, 0, 4], [5, 0, 5]])) == []
    assert _hints(_sheet(columns=columns, rows=[[4, 1, 3], [5, 0, None]])) == []


# --- manual total rows ----------------------------------------------------------

_SUM_COLUMNS = [
    {"header": "Poste"},
    {"header": "HT", "type": "currency", "total": "sum"},
    {"header": "TTC", "type": "currency", "total": "sum"},
]


def _totals(columns: list[dict[str, Any]], rows: list[list[Any]]) -> list[str]:
    return manual_total_hints(prepare(_request(_sheet(columns=columns, rows=rows))))


def test_an_empty_manual_total_row_is_named():
    rows = [["Loyer", 100, 120], ["Pay", 200, 240], ["Total", None, None]]

    assert _totals(_SUM_COLUMNS, rows) == ["sheet 'Data', row 4 ('Total')"]


def test_filled_total_rows_equal_to_the_totals_are_named_once_per_sheet():
    rows = [
        ["Loyer", 100, 120],
        ["Pay", 200, 240],
        ["Total HT", 300, None],
        ["  TOTAL  ttc :", None, 360],
    ]

    assert _totals(_SUM_COLUMNS, rows) == ["sheet 'Data', row 4 ('Total HT')"]


def test_a_total_row_with_other_values_or_without_total_columns_is_data():
    rows = [["Total", 50, 60], ["TotalEnergies", 100, 120], ["Pay", 200, 240]]
    plain = [{"header": "Company"}, {"header": "Revenue", "type": "currency"}]

    assert _totals(_SUM_COLUMNS, rows) == []
    assert _totals(_SUM_COLUMNS, [["TotalEnergies", None, None], ["a", 1, 2]]) == []
    assert _totals(plain, [["Total", None], ["TotalEnergies", 5]]) == []
