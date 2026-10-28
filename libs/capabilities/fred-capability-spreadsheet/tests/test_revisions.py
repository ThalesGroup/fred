# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Targeted revisions: changes applied to the delivered request, then rebuilt."""

from __future__ import annotations

import io
import json
from typing import Any

import pytest
from fred_capability_spreadsheet import capability
from fred_capability_spreadsheet.capability import SpreadsheetCapability
from fred_capability_spreadsheet.revisions import NO_BASE, ChangeError, apply_changes
from fred_capability_spreadsheet.store import (
    MemoryWorkbookSpecStore,
    clear_workbook_spec_store_provider,
    set_workbook_spec_store_provider,
)
from fred_capability_spreadsheet.workbook import (
    CreateSpreadsheetInput,
    build_workbook,
    prepare,
)
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
)
from fred_sdk.contracts.runtime import RuntimeServices
from langchain_core.tools import StructuredTool
from openpyxl import load_workbook
from test_capability import FakeWorkspace

BUDGET: dict[str, Any] = {
    "file_name": "Budget",
    "language": "fr",
    "sheets": [
        {
            "name": "Budget",
            "columns": [
                {"header": "Poste"},
                {"header": "Budget", "type": "currency", "total": "sum"},
                {"header": "Réel", "type": "currency", "total": "sum"},
            ],
            "rows": [
                ["Loyer", 1200, 1150],
                ["Voyages", 300, 350],
                ["Matériel", 500, 0],
            ],
        },
        {
            "name": "Synthèse",
            "columns": [
                {"header": "Indicateur"},
                {"header": "Valeur", "type": "number"},
            ],
            "rows": [
                ["Budget total", {"formula": "=SUM(Budget!B2:B4)"}],
                ["Loyer réel", {"formula": "=Budget!C2"}],
            ],
        },
    ],
}


def _apply(
    *changes: dict[str, Any], base: dict[str, Any] | None = None
) -> dict[str, Any]:
    return apply_changes(base or BUDGET, list(changes))


def _refused(*changes: dict[str, Any], base: dict[str, Any] | None = None) -> str:
    with pytest.raises(ChangeError) as info:
        _apply(*changes, base=base)
    return str(info.value)


def _sheets(spec: dict[str, Any]) -> dict[str, Any]:
    return {s["name"]: s for s in spec["sheets"]}


def _cell_formula(spec: dict[str, Any], sheet: str, row: int, column: int) -> str:
    return _sheets(spec)[sheet]["rows"][row][column]["formula"]


def test_add_a_computed_column_keeps_everything_else():
    spec = _apply(
        {
            "op": "add_column",
            "sheet": "Budget",
            "column": {
                "header": "Écart",
                "type": "currency",
                "compute": {"op": "subtract", "left": "Réel", "right": "Budget"},
            },
        }
    )
    budget = _sheets(spec)["Budget"]
    assert [c["header"] for c in budget["columns"]] == [
        "Poste",
        "Budget",
        "Réel",
        "Écart",
    ]
    assert budget["rows"][0] == ["Loyer", 1200, 1150, None]
    ws = load_workbook(
        io.BytesIO(build_workbook(prepare(CreateSpreadsheetInput.model_validate(spec))))
    )["Budget"]
    assert ws["D2"].value == "=C2-B2"


def test_appended_rows_extend_ranges_that_ended_on_the_last_data_row():
    spec = _apply(
        {"op": "add_rows", "sheet": "Budget", "rows": [["Formation", 100, 80]]}
    )
    assert _sheets(spec)["Budget"]["rows"][-1] == ["Formation", 100, 80]
    assert _cell_formula(spec, "Synthèse", 0, 1) == "=SUM(Budget!B2:B5)"
    assert _cell_formula(spec, "Synthèse", 1, 1) == "=Budget!C2"


def test_rows_inserted_inside_a_range_widen_it_and_shift_cells():
    spec = _apply(
        {"op": "add_rows", "sheet": "Budget", "rows": [["Taxes", 10, 10]], "after": 2}
    )
    assert _sheets(spec)["Budget"]["rows"][1][0] == "Taxes"
    assert _cell_formula(spec, "Synthèse", 0, 1) == "=SUM(Budget!B2:B5)"
    assert _cell_formula(spec, "Synthèse", 1, 1) == "=Budget!C2"
    # Above the first row of a range, as in Excel, the range moves down.
    spec = _apply(
        {"op": "add_rows", "sheet": "Budget", "rows": [["Taxes", 10, 10]], "after": 1}
    )
    assert _cell_formula(spec, "Synthèse", 0, 1) == "=SUM(Budget!B3:B5)"
    assert _cell_formula(spec, "Synthèse", 1, 1) == "=Budget!C3"


def test_rows_as_objects_by_header():
    spec = _apply(
        {"op": "add_rows", "sheet": "Budget", "rows": [{"Poste": "Taxes", "Réel": 5}]}
    )
    assert _sheets(spec)["Budget"]["rows"][-1] == ["Taxes", None, 5]


def test_split_a_table_into_two_sheets():
    base = {
        "file_name": "Ventes",
        "sheets": [
            {
                "name": "Ventes",
                "columns": [
                    {"header": "Client"},
                    {"header": "Région"},
                    {"header": "CA", "type": "number", "total": "sum"},
                ],
                "rows": [
                    ["A", "Nord", 1],
                    ["B", "Sud", 2],
                    ["C", "Nord", 3],
                    ["D", "sud", 4],
                ],
            }
        ],
    }
    spec = _apply(
        {
            "op": "move_rows",
            "sheet": "Ventes",
            "where": {"column": "Région", "equals": "Sud"},
            "to_sheet": "Sud",
        },
        {"op": "rename_sheet", "sheet": "Ventes", "name": "Nord"},
        base=base,
    )
    sheets = _sheets(spec)
    assert list(sheets) == ["Nord", "Sud"]
    assert sheets["Nord"]["rows"] == [["A", "Nord", 1], ["C", "Nord", 3]]
    assert sheets["Sud"]["rows"] == [["B", "Sud", 2], ["D", "sud", 4]]
    assert sheets["Sud"]["columns"][2]["total"] == "sum"
    prepared = prepare(CreateSpreadsheetInput.model_validate(spec))
    assert [s.totals[2] for s in prepared.sheets if s.totals] == [4, 6]


def test_move_to_an_existing_sheet_maps_columns_by_header():
    base = json.loads(json.dumps(BUDGET))
    base["sheets"].append(
        {
            "name": "Archive",
            "columns": [
                {"header": "Réel", "type": "currency"},
                {"header": "Poste"},
                {"header": "Budget", "type": "currency"},
            ],
            "rows": [],
        }
    )
    # The summary sums the moved row: refused until that formula is edited too.
    move = {"op": "move_rows", "sheet": "Budget", "rows": [4], "to_sheet": "Archive"}
    assert "move to 'Archive'" in _refused(move, base=base)
    spec = _apply(
        move,
        {
            "op": "edit_cell",
            "sheet": "Synthèse",
            "cell": "B2",
            "value": {"formula": "=SUM(Budget!B2:B3)"},
        },
        base=base,
    )
    assert _sheets(spec)["Archive"]["rows"] == [[0, "Matériel", 500]]


def test_rename_a_sheet_rewrites_every_reference():
    spec = _apply({"op": "rename_sheet", "sheet": "Budget", "name": "Données 2025"})
    assert _cell_formula(spec, "Synthèse", 0, 1) == "=SUM('Données 2025'!B2:B4)"
    assert _cell_formula(spec, "Synthèse", 1, 1) == "='Données 2025'!C2"


def test_removing_a_referenced_cell_names_the_formula():
    message = _refused({"op": "remove_rows", "sheet": "Budget", "rows": [2]})
    assert "cell B3" in message and "reads row 2 of sheet 'Budget'" in message
    # Inside a range only, the range shrinks as in Excel.
    spec = _apply(
        {"op": "remove_rows", "sheet": "Budget", "rows": [3]},
    )
    assert _cell_formula(spec, "Synthèse", 0, 1) == "=SUM(Budget!B2:B3)"


def test_removing_a_column_a_computed_column_uses_is_refused():
    base = _apply(
        {
            "op": "add_column",
            "sheet": "Budget",
            "column": {
                "header": "Écart",
                "compute": {"op": "subtract", "left": "Réel", "right": "Budget"},
            },
        }
    )
    message = _refused(
        {"op": "remove_column", "sheet": "Synthèse", "column": "Valeur"},
        {"op": "remove_column", "sheet": "Budget", "column": "Réel"},
        base=base,
    )
    assert "computed column 'Écart'" in message


def test_renaming_a_header_follows_compute_operands():
    base = _apply(
        {
            "op": "add_column",
            "sheet": "Budget",
            "column": {
                "header": "Écart",
                "compute": {"op": "subtract", "left": "Réel", "right": "Budget"},
            },
        }
    )
    spec = _apply(
        {"op": "edit_column", "sheet": "Budget", "column": "Réel", "header": "Réalisé"},
        base=base,
    )
    assert _sheets(spec)["Budget"]["columns"][3]["compute"]["left"] == "Réalisé"


def test_errors_are_listed_at_once_with_the_fix():
    message = _refused(
        {"op": "remove_rows", "sheet": "Budget", "rows": [40]},
        {"op": "edit_rows", "sheet": "Budget", "rows": [2], "values": {"Montant": 1}},
        {"op": "explode", "sheet": "Budget"},
    )
    assert "changes[0] (remove_rows): row 40 does not exist" in message
    assert "data rows 2 to 4" in message
    assert (
        "changes[1] (edit_rows)" in message and "'Poste', 'Budget', 'Réel'" in message
    )
    assert "changes[2] (explode): unknown op" in message


def test_edit_cell_and_edit_rows_by_where():
    spec = _apply(
        {"op": "edit_cell", "sheet": "Budget", "cell": "C3", "value": 360},
        {
            "op": "edit_rows",
            "sheet": "Budget",
            "where": {"column": "Poste", "equals": "loyer"},
            "values": {"Réel": {"formula": "=B2*0.9"}},
        },
    )
    rows = _sheets(spec)["Budget"]["rows"]
    assert rows[1][2] == 360 and rows[0][2] == {"formula": "=B2*0.9"}


def test_conflicting_changes_are_refused():
    assert "already removed or moved" in _refused(
        {"op": "remove_rows", "sheet": "Budget", "rows": [2]},
        {"op": "edit_cell", "sheet": "Budget", "cell": "B2", "value": 1},
    )
    assert "removed by another change" in _refused(
        {"op": "remove_sheet", "sheet": "Synthèse"},
        {"op": "rename_sheet", "sheet": "Synthèse", "name": "X"},
    )


def test_add_and_remove_sheets():
    spec = _apply(
        {"op": "remove_sheet", "sheet": "Synthèse"},
        {
            "op": "add_sheet",
            "sheet": {
                "name": "Notes",
                "columns": [{"header": "Note"}],
                "rows": [["ok"]],
            },
        },
    )
    assert [s["name"] for s in spec["sheets"]] == ["Budget", "Notes"]
    message = _refused({"op": "remove_sheet", "sheet": "Budget"})
    assert "reads sheet 'Budget'" in message


def test_nothing_changed_is_an_error():
    message = _refused(
        {"op": "edit_cell", "sheet": "Budget", "cell": "B2", "value": 1200}
    )
    assert "nothing changed" in message and "Do not tell the user" in message


# --- through the tool -------------------------------------------------------------


def _tool(workspace: FakeWorkspace, instance: str | None = "agent-1"):
    ctx = CapabilityContext(
        identity=CapabilityIdentity(
            user_id="u-1", session_id="s-1", agent_instance_id=instance
        ),
        config=EmptyModel(),
        turn_options=EmptyModel(),
        services=RuntimeServices(workspace_fs=workspace),
    )
    (tool,) = SpreadsheetCapability().tools(ctx)
    assert isinstance(tool, StructuredTool) and tool.coroutine is not None
    return tool.coroutine


@pytest.fixture
def store():
    memory = MemoryWorkbookSpecStore()
    set_workbook_spec_store_provider(lambda: memory)
    yield memory
    clear_workbook_spec_store_provider()


@pytest.mark.asyncio
async def test_changes_apply_to_the_stored_base_in_a_later_turn(store):
    workspace = FakeWorkspace()
    text, _ = await _tool(workspace)(**BUDGET)
    assert "To change it later, call create_spreadsheet with changes" in text
    assert "'Budget' columns Poste, Budget, Réel; data rows 2-4, totals row 5" in text
    assert text.endswith("Do not create it again; reply now.")
    stored = store.rows[("s-1", "agent-1")]
    assert stored.version == 1
    # A new turn: a new tool instance, no history, the stored base.
    later = _tool(workspace)
    text, artifact = await later(
        changes=[
            {"op": "add_rows", "sheet": "Budget", "rows": [["Formation", 100, 80]]}
        ]
    )
    assert artifact.is_error is False, text
    assert store.rows[("s-1", "agent-1")].version == 2
    assert store.rows[("s-1", "agent-1")].workbook_id == stored.workbook_id
    xlsx = [w for w in workspace.writes if w[0].endswith(".xlsx")][-1][1]
    cached = load_workbook(io.BytesIO(xlsx), data_only=True)
    assert cached["Budget"]["A5"].value == "Formation"
    assert cached["Synthèse"]["B2"].value == 2100


@pytest.mark.asyncio
async def test_no_base_and_both_modes_are_refused(store):
    del store
    text, artifact = await _tool(FakeWorkspace())(
        changes=[{"op": "remove_sheet", "sheet": "X"}]
    )
    assert artifact.is_error and text == NO_BASE
    text, artifact = await _tool(FakeWorkspace())(**BUDGET, changes=[{"op": "x"}])
    assert artifact.is_error and "not both" in text


@pytest.mark.asyncio
async def test_bad_changes_keep_the_base(store):
    workspace = FakeWorkspace()
    tool = _tool(workspace)
    await tool(**BUDGET)
    text, artifact = await tool(
        changes=[{"op": "remove_rows", "sheet": "Budget", "rows": [9]}]
    )
    assert artifact.is_error and "row 9 does not exist" in text
    assert store.rows[("s-1", "agent-1")].version == 1
    text, artifact = await tool(
        changes=[{"op": "remove_rows", "sheet": "Budget", "rows": [4]}]
    )
    assert artifact.is_error is False, text


@pytest.mark.asyncio
async def test_identical_changes_in_a_turn_apply_once(store):
    del store
    workspace = FakeWorkspace()
    tool = _tool(workspace)
    await tool(**BUDGET)
    change = [{"op": "add_rows", "sheet": "Budget", "rows": [["Taxes", 1, 1]]}]
    first, _ = await tool(changes=change)
    second, artifact = await tool(changes=change)
    xlsx = [w for w in workspace.writes if w[0].endswith(".xlsx")]
    assert len(xlsx) == 2 and second == first and artifact.ui_parts == ()


@pytest.mark.asyncio
async def test_without_agent_identity_the_base_lives_in_the_turn_only(store):
    workspace = FakeWorkspace()
    tool = _tool(workspace, instance=None)
    await tool(**BUDGET)
    text, artifact = await tool(changes=[{"op": "remove_sheet", "sheet": "Synthèse"}])
    assert artifact.is_error is False, text
    assert store.rows == {}


@pytest.mark.asyncio
async def test_a_too_large_request_is_delivered_but_not_kept(store, monkeypatch):
    monkeypatch.setattr(capability, "MAX_STORED_SPEC_BYTES", 10)
    workspace = FakeWorkspace()
    tool = _tool(workspace)
    text, artifact = await tool(**BUDGET)
    assert artifact.is_error is False and "too large to keep" in text
    text, artifact = await tool(changes=[{"op": "remove_sheet", "sheet": "Synthèse"}])
    assert artifact.is_error and "too large to keep" in text


@pytest.mark.asyncio
async def test_graph_assembled_tool_takes_changes_through_a_tool_call(store):
    from fred_runtime.capabilities.assembly import build_capability_agent_block
    from fred_runtime.capabilities.registry import CapabilityRegistry

    del store
    registry = CapabilityRegistry()
    registry.register(SpreadsheetCapability())
    ctx = CapabilityContext(
        identity=CapabilityIdentity(
            user_id="u-1", session_id="s-1", agent_instance_id="graph-1"
        ),
        config=EmptyModel(),
        turn_options=EmptyModel(),
        services=RuntimeServices(workspace_fs=FakeWorkspace()),
    )
    (tool,) = build_capability_agent_block(registry, {"spreadsheet": ctx}).tools
    await tool.ainvoke(
        {"type": "tool_call", "id": "c-1", "name": tool.name, "args": BUDGET}
    )
    message = await tool.ainvoke(
        {
            "type": "tool_call",
            "id": "c-2",
            "name": tool.name,
            "args": {"changes": [{"op": "remove_sheet", "sheet": "Synthèse"}]},
        }
    )
    assert message.artifact.is_error is False, message.content
    assert message.artifact.ui_parts[0].type == "spreadsheet_preview"
