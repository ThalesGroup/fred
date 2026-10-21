# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Registration, manifest and `create_spreadsheet` wiring with a fake workspace port."""

from __future__ import annotations

import io
from importlib.metadata import entry_points

import pytest
from fred_capability_spreadsheet.capability import (
    TOOL_NAME,
    XLSX_CONTENT_TYPE,
    SpreadsheetCapability,
)
from fred_runtime.capabilities.assembly import build_capability_agent_block
from fred_runtime.capabilities.registry import CapabilityRegistry
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
    TeamScopePolicy,
)
from fred_sdk.contracts.context import (
    BoundRuntimeContext,
    LinkKind,
    LinkPart,
    PublishedArtifact,
)
from fred_sdk.contracts.runtime import RuntimeServices, WorkspaceFsPort
from langchain_core.tools import StructuredTool
from langchain_core.utils.function_calling import convert_to_openai_tool
from openpyxl import load_workbook


class FakeWorkspace(WorkspaceFsPort):
    def __init__(self, fail: bool = False) -> None:
        self.writes: list[tuple[str, bytes, str | None, str | None]] = []
        self.fail = fail

    def bind(self, binding: BoundRuntimeContext) -> None:
        del binding

    async def write(self, path, content, *, content_type=None, title=None):
        if self.fail:
            raise OSError("storage down")
        self.writes.append((path, content, content_type, title))
        return PublishedArtifact(
            key=f"teams/t/{path}",
            file_name=path.rsplit("/", 1)[-1],
            size=len(content),
            href="https://fred.test/download/1",
            document_uid="doc-1",
            mime=content_type,
            title=title,
        )


def _ctx(workspace: WorkspaceFsPort | None, session_id: str | None = "s-1"):
    return CapabilityContext(
        identity=CapabilityIdentity(user_id="u-1", session_id=session_id),
        config=EmptyModel(),
        turn_options=EmptyModel(),
        services=RuntimeServices(workspace_fs=workspace),
    )


def _tool(workspace: WorkspaceFsPort | None, session_id: str | None = "s-1"):
    (tool,) = SpreadsheetCapability().tools(_ctx(workspace, session_id))
    assert isinstance(tool, StructuredTool) and tool.coroutine is not None
    return tool.coroutine


ARGS = {
    "file_name": "Q3 budget",
    "sheets": [
        {
            "name": "Budget",
            "columns": [
                {"header": "Line"},
                {
                    "header": "Amount",
                    "type": "currency",
                    "currency": "EUR",
                    "total": "sum",
                },
            ],
            "rows": [["Rent", 1200], ["Travel", 350.5]],
        }
    ],
}


def test_registers_and_validates_clean():
    registry = CapabilityRegistry()
    registry.register(SpreadsheetCapability())
    registry.validate()
    assert "spreadsheet" in registry.ids()


def test_entry_point_is_declared():
    names = {ep.name: ep.value for ep in entry_points(group="fred.capabilities")}
    assert (
        names["spreadsheet"]
        == "fred_capability_spreadsheet.capability:SpreadsheetCapability"
    )


def test_manifest_is_tool_only_admin_gated_and_graph_compatible():
    manifest = SpreadsheetCapability.manifest
    assert manifest.id == "spreadsheet"
    assert manifest.icon == "table_chart"
    assert manifest.team_scope == TeamScopePolicy.ADMIN_GATED
    assert manifest.execution_models == ("react", "graph")
    assert manifest.chat_parts == []
    assert manifest.side_panels == []
    assert manifest.router is None
    assert manifest.tables == []
    assert manifest.config_fields == []
    assert SpreadsheetCapability.migrations_location() is None


def test_tool_schema_exposes_only_llm_arguments():
    schema = convert_to_openai_tool(
        SpreadsheetCapability().tools(_ctx(FakeWorkspace()))[0]
    )["function"]
    assert schema["name"] == TOOL_NAME
    assert set(schema["parameters"]["properties"]) == {"file_name", "sheets"}
    assert "$defs" not in str(schema)
    assert "never paste a Markdown table" in schema["description"]


def test_graph_assembly_exposes_the_tool():
    registry = CapabilityRegistry()
    registry.register(SpreadsheetCapability())
    block = build_capability_agent_block(
        registry, {"spreadsheet": _ctx(FakeWorkspace())}
    )
    assert [tool.name for tool in block.tools] == [TOOL_NAME]


@pytest.mark.asyncio
async def test_writes_xlsx_and_returns_download_chip():
    workspace = FakeWorkspace()
    content, artifact = await _tool(workspace)(**ARGS)

    assert "download button is shown" in content
    assert artifact.is_error is False
    (path, data, content_type, title) = workspace.writes[0]
    assert path.startswith("spreadsheets/s-1/Q3 budget-") and path.endswith(".xlsx")
    assert content_type == XLSX_CONTENT_TYPE
    assert title == "Q3 budget.xlsx"
    (link,) = artifact.ui_parts
    assert isinstance(link, LinkPart)
    assert link.kind == LinkKind.download
    assert link.href == "https://fred.test/download/1"
    assert link.file_name == "Q3 budget.xlsx"
    assert link.mime == XLSX_CONTENT_TYPE
    ws = load_workbook(io.BytesIO(data))["Budget"]
    assert ws["B4"].value == "=SUM(B2:B3)"


@pytest.mark.asyncio
async def test_path_without_session():
    workspace = FakeWorkspace()
    await _tool(workspace, session_id=None)(**ARGS)
    assert workspace.writes[0][0].startswith("spreadsheets/Q3 budget-")


@pytest.mark.asyncio
async def test_invalid_input_returns_error_and_writes_nothing():
    workspace = FakeWorkspace()
    args = {**ARGS, "sheets": [{**ARGS["sheets"][0], "rows": [["Rent"]]}]}
    content, artifact = await _tool(workspace)(**args)
    assert artifact.is_error is True
    assert "No file was written" in content
    assert "Sheet 'Budget', row 2" in content
    assert workspace.writes == []


@pytest.mark.asyncio
async def test_missing_workspace_port_is_an_error():
    content, artifact = await _tool(None)(**ARGS)
    assert artifact.is_error is True
    assert "not available" in content


@pytest.mark.asyncio
async def test_storage_failure_is_an_error():
    content, artifact = await _tool(FakeWorkspace(fail=True))(**ARGS)
    assert artifact.is_error is True
    assert "could not be stored" in content


@pytest.mark.asyncio
async def test_tool_call_through_langchain_keeps_the_artifact():
    workspace = FakeWorkspace()
    message = (
        await SpreadsheetCapability()
        .tools(_ctx(workspace))[0]
        .ainvoke({"type": "tool_call", "id": "call-1", "name": TOOL_NAME, "args": ARGS})
    )
    assert message.artifact.ui_parts[0].kind == LinkKind.download
