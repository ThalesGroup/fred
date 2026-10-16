# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Key-figure trends: an arrow before the figure, in the figure's own font
and colour, never a colour that claims good or bad news."""

from __future__ import annotations

import io
from typing import Any, Dict, List, Optional, Tuple, cast

import pytest
from builder_decks import build_slides
from fred_capability_ppt_filler.builder import tools
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import assign
from fred_capability_ppt_filler.builder.builtin import builtin_template
from fred_capability_ppt_filler.builder.changes import SlideChange, apply_changes
from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanItem
from fred_capability_ppt_filler.builder.render import RenderResult, render_deck
from pptx import Presentation

COVER = {"title": "Bilan", "subtitle": "2025"}
FIGURES = [
    {"number": "42 %", "heading": "Marge", "trend": "up"},
    {"number": "12", "heading": "Clients", "trend": "baisse"},
    {"number": "3,5 M€", "heading": "CA", "trend": "stable"},
]


@pytest.fixture(scope="module")
def builtin():
    return builtin_template()


def _render(builtin, slide: Dict[str, Any]) -> RenderResult:
    template, inventory = builtin
    plan = DeckPlan.model_validate({"slides": [COVER, slide]})
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems, [p.as_dict() for p in assignment.problems]
    return render_deck(template, inventory, plan, assignment.patterns)


def _figures(result: RenderResult) -> List[List[Tuple[str, Optional[str]]]]:
    """Each figure's runs as (text, size) on the second slide."""
    slide = Presentation(io.BytesIO(result.pptx)).slides[1]
    return [
        [
            (run.text, run._r.rPr.get("sz") if run._r.rPr is not None else None)
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
        ]
        for shape in cast(Any, slide.shapes)
        if shape.name.endswith("_value")
    ]


def _notes(result: RenderResult) -> List[Tuple[str, Optional[str]]]:
    return [(n.code, n.field) for n in result.report.built[1].notes if n.field]


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("up", "up"),
        ("Hausse", "up"),
        ("increasing", "up"),
        ("↑", "up"),
        ("en baisse", "down"),
        ("decline", "down"),
        ("stable", "flat"),
        ("unchanged", "flat"),
    ],
)
def test_trends_are_read_from_loose_spellings(raw, expected):
    assert PlanItem.model_validate({"number": "4", "trend": raw}).trend == expected


def test_a_word_naming_no_direction_is_kept_for_the_note():
    item = PlanItem.model_validate({"number": "4", "evolution": "sideways"})
    assert (item.trend, item.asked_trend) == (None, "sideways")
    assert PlanItem.model_validate(item.model_dump()).asked_trend == "sideways"


def test_each_figure_leads_with_its_arrow_in_a_smaller_run(builtin):
    result = _render(
        builtin, {"title": "Chiffres", "kind": "figures", "items": FIGURES}
    )
    figures = _figures(result)
    assert [[text for text, _ in runs] for runs in figures] == [
        ["▲ ", "42\xa0%"],
        ["▼ ", "12"],
        ["► ", "3,5\xa0M€"],
    ]
    for (_, glyph), (_, figure) in figures:
        assert glyph is not None and figure is not None
        assert int(glyph) == int(int(figure) * 0.6)
    slide = Presentation(io.BytesIO(result.pptx)).slides[1]
    xml = "".join(s._element.xml for s in slide.shapes if s.name.endswith("_value"))
    assert "srgbClr" not in xml and "schemeClr" not in xml
    assert _notes(result) == []


def test_figures_without_trend_are_drawn_as_before(builtin):
    items = [{k: v for k, v in f.items() if k != "trend"} for f in FIGURES]
    result = _render(builtin, {"title": "Chiffres", "kind": "figures", "items": items})
    assert [len(runs) for runs in _figures(result)] == [1, 1, 1]


def test_a_trend_without_a_figure_zone_is_told(builtin):
    cards = [
        {"heading": "Marge", "text": "En progression continue", "trend": "up"},
        {"heading": "Clients", "text": "Base stable sur l'année"},
        {"heading": "CA", "text": "Porté par l'export"},
    ]
    result = _render(builtin, {"title": "Points", "kind": "cards", "items": cards})
    assert _notes(result) == [("intent_not_shown", "trend")]
    assert any("shows no figure for its items" in n for n in result.report.notes)


def test_a_trend_without_its_number_or_direction_is_told(builtin):
    items = [
        {"heading": "Marge", "trend": "up"},
        {"number": "12", "heading": "Clients", "trend": "sideways"},
        {"number": "3", "heading": "CA"},
    ]
    result = _render(builtin, {"title": "Chiffres", "kind": "figures", "items": items})
    assert _notes(result) == [("intent_not_shown", "trend")]
    notes = "\n".join(result.report.notes)
    assert "items[0].trend needs the item's number" in notes
    assert "items[1].trend 'sideways' is not a direction" in notes
    assert [len(runs) for runs in _figures(result)][1:] == [1, 1]


def test_a_trend_is_edited(builtin):
    _, inventory = builtin
    base = DeckPlan.model_validate(
        {"slides": [COVER, {"title": "Chiffres", "kind": "figures", "items": FIGURES}]}
    )
    edit = {"op": "edit", "slide": 2, "field": "items[1].trend", "value": "up"}
    plan = apply_changes(base, [SlideChange.model_validate(edit)])
    assert plan.slides[1].items[1].trend == "up"
    assert not assign(inventory, plan, "off").problems


def test_trend_is_offered_only_where_items_show_a_figure(builtin):
    _, inventory = builtin
    schema = tools._args_schema(inventory).model_json_schema()
    assert "trend" in schema["$defs"]["BuildItem"]["properties"]
    assert "trend (up, down or flat" in tools._description(inventory, "off", 2)

    def bullets(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Constats"
        slide.placeholders[1].text_frame.text = "Un point"

    plain = analyze_template(build_slides([bullets])).inventory
    assert plain is not None
    schema = tools._args_schema(plain).model_json_schema()
    item = schema.get("$defs", {}).get("BuildItem", {}).get("properties", {})
    assert "trend" not in item
    assert "trend (up" not in tools._description(plain, "off", 2)
