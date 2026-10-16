# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Template slides holding two to four example charts: analysed as chart
slides, chosen for a plan slide with exactly as many `charts`, filled frame by
frame. Built on a template shaped like the owner's (1, 2, 3 and 4 charts)."""

from __future__ import annotations

import io
from typing import Any, Dict, List, cast

import pytest
from builder_decks import add_charts_slide, add_main_chart_slide, build_slides
from fred_capability_ppt_filler.builder import checks, story, tools
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import Assignment, assign
from fred_capability_ppt_filler.builder.changes import (
    ChangeError,
    SlideChange,
    apply_changes,
)
from fred_capability_ppt_filler.builder.models import Pattern, TemplateInventory
from fred_capability_ppt_filler.builder.plan import DeckPlan, zone_summary
from fred_capability_ppt_filler.builder.render import (
    RenderError,
    RenderResult,
    render_deck,
)
from fred_capability_ppt_filler.builder.story import Conversation
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

COVER = {"title": "Quarterly review", "subtitle": "Sales"}
TAKEAWAY = "Revenue grows every quarter while the margin holds, so the plan stands"


def _chart(name: str, values: List[float], kind: str = "column") -> Dict[str, Any]:
    return {
        "type": kind,
        "title": name,
        "categories": ["Q1", "Q2", "Q3", "Q4"][: len(values)],
        "series": [{"name": name, "values": values}],
    }


REVENUE = _chart("Revenue", [12, 15, 18, 21])
MARGIN = _chart("Margin", [8, 9, 9, 10])
SHARE = {
    "type": "pie",
    "title": "Revenue by region",
    "categories": ["Europe", "Americas", "Asia"],
    "series": [{"name": "Share", "values": [45, 35, 20]}],
}


def _pie_then_column(presentation):
    return add_charts_slide(
        presentation, 2, [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.PIE]
    )


def _three(presentation):
    return add_charts_slide(presentation, 3)


def _four(presentation):
    return add_charts_slide(presentation, 4)


def _template(*builders) -> tuple[bytes, TemplateInventory]:
    template = build_slides(list(builders))
    inventory = analyze_template(template).inventory
    assert inventory is not None
    return template, inventory


@pytest.fixture(scope="module")
def owner():
    """One-chart, two-chart, three-chart and four-chart example slides."""
    return _template(
        lambda p: add_main_chart_slide(p, XL_CHART_TYPE.COLUMN_CLUSTERED),
        _pie_then_column,
        _three,
        _four,
    )


def _slides(inventory: TemplateInventory) -> List[Pattern]:
    return [p for p in inventory.patterns if p.source == "slide"]


def _frames(pattern: Pattern) -> int:
    return sum(z.role == "chart" for z in pattern.zones)


def _plan(*slides: Dict[str, Any]) -> DeckPlan:
    return DeckPlan.model_validate({"slides": [COVER, *slides]})


def _multi(*charts: Dict[str, Any], **extra: Any) -> Dict[str, Any]:
    return {"title": "Sales", "takeaway": TAKEAWAY, "charts": list(charts), **extra}


def _assigned(inventory: TemplateInventory, *slides: Dict[str, Any]) -> Assignment:
    return assign(inventory, _plan(*slides), "off")


def _render(template: bytes, inventory: TemplateInventory, *slides) -> RenderResult:
    plan = _plan(*slides)
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems, [p.as_dict() for p in assignment.problems]
    return render_deck(template, inventory, plan, assignment.patterns)


def _delivered_charts(result: RenderResult, slide: int = 1) -> List[Any]:
    deck = Presentation(io.BytesIO(result.pptx))
    shapes = [s for s in deck.slides[slide].shapes if getattr(s, "has_chart", False)]
    return sorted(shapes, key=lambda s: (s.top // Inches(1), s.left))


# --- analysis ---------------------------------------------------------------------


def test_example_slides_with_two_to_four_charts_are_chart_slides(owner):
    _, inventory = owner
    slides = _slides(inventory)
    assert [(p.kind, p.included, _frames(p)) for p in slides] == [
        ("chart", True, 1),
        ("chart", True, 2),
        ("chart", True, 3),
        ("chart", True, 4),
    ]
    pair, _, grid = slides[1:]
    # Reading order, whatever the drawing order: left first, then the 2 × 2 grid.
    assert [z.chart_type for z in pair.zones if z.role == "chart"] == ["column", "pie"]
    boxes = [(z.box.x, z.box.y) for z in grid.zones if z.role == "chart"]
    assert boxes == sorted(boxes, key=lambda b: (b[1] // 150, b[0]))
    assert boxes[0][0] < boxes[1][0] and boxes[0][1] < boxes[2][1]
    assert "charts×4 (column, " in zone_summary(grid)
    assert "chart (column, " in zone_summary(slides[0])


def test_five_charts_or_a_tiny_one_keep_the_slide_out():
    def five(presentation):
        slide = add_charts_slide(presentation, 4)
        data = CategoryChartData()
        data.categories = ["A"]
        data.add_series("S", (1,))
        slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, Inches(0.2), Inches(6.3), Inches(3), Inches(1.2), data
        )
        return slide

    def tiny(presentation):
        slide = add_charts_slide(presentation, 2)
        frame = next(s for s in slide.shapes if s.has_chart)
        frame.width, frame.height = Inches(2), Inches(1.1)
        return slide

    _, inventory = _template(five, tiny)
    assert [(p.included, p.exclusion_reason) for p in _slides(inventory)] == [
        (False, "unsupported_parts"),
        (False, "unsupported_parts"),
    ]


# --- plan and storytelling -----------------------------------------------------------


def test_chart_and_charts_together_are_refused(owner):
    _, inventory = owner
    slide = _multi(REVENUE, MARGIN, chart=REVENUE)
    (problem,) = _assigned(inventory, slide).problems
    assert "both `chart` and `charts` sent" in problem.problem
    assert problem.options == [
        "send `chart` for one chart or `charts` for several, not both"
    ]


def test_a_chart_of_several_is_named_in_its_problems(owner):
    _, inventory = owner
    broken = {**MARGIN, "series": [{"name": "Margin", "values": [8, 9]}]}
    (problem,) = _assigned(inventory, _multi(REVENUE, broken)).problems
    assert problem.problem.startswith("charts[1].series[0] 'Margin' has 2 values")


def _story(*slides: Dict[str, Any]) -> Dict[int, str]:
    plan = _plan(*slides)
    users = ["CA 12 15 18 21, marge 8 % 9 % 9 % 10 %, 3 M€"]
    talk = Conversation(users=users, sources=users)
    result = story.check(plan, range(1, len(plan.slides) + 1), talk, set())
    return {p.index: p.problem for p in result.problems}


@pytest.mark.storytelling
def test_each_chart_shows_one_measure_and_has_a_title():
    mixed = {**MARGIN, "series": [{"name": "Margin", "values": ["8 %", "3 M€"]}]}
    untitled = {**REVENUE, "title": None}
    problems = _story(
        _multi(REVENUE, {**mixed, "categories": ["Q1", "Q2"]}, source="user"),
        _multi(untitled, MARGIN, source="user"),
    )
    assert "charts[1]: chart mixes units (%, m€)" in problems[2]
    assert "charts[0] has no title" in problems[3]
    assert "charts[1]" not in problems[3]


# --- assignment ----------------------------------------------------------------------


def test_k_charts_take_the_slide_with_k_frames_and_one_chart_its_own(owner):
    _, inventory = owner
    assignment = _assigned(
        inventory,
        {"title": "Revenue", "takeaway": TAKEAWAY, "chart": REVENUE},
        _multi(REVENUE, MARGIN),
        _multi(REVENUE, MARGIN, REVENUE),
        _multi(REVENUE, MARGIN, REVENUE, MARGIN),
    )
    assert not assignment.problems
    assert [_frames(p) for p in assignment.patterns[1:] if p] == [1, 2, 3, 4]


def test_variety_never_offers_a_slide_with_another_number_of_charts(owner):
    _, inventory = owner
    slides = [_multi(REVENUE, MARGIN, title=f"Sales {n}") for n in range(3)]
    assignment = assign(inventory, _plan(*slides), "template_slide", 2)
    assert not assignment.problems
    (repeat,) = assignment.repeats
    assert (repeat.slides, repeat.unused, repeat.revisable) == ([2, 3, 4], [], False)


def test_a_count_no_slide_holds_is_refused_with_the_counts_offered():
    _, inventory = _template(
        lambda p: add_main_chart_slide(p, XL_CHART_TYPE.COLUMN_CLUSTERED),
        _pie_then_column,
        _four,
    )
    (problem,) = _assigned(inventory, _multi(REVENUE, MARGIN, REVENUE)).problems
    assert problem.problem == (
        "charts: 3 given; this template offers slides with 1, 2 or 4 charts"
    )
    assert problem.options[0].startswith("split the slide into slides of 1, 2 or 4")
    assert problem.options[1] == "or merge or drop charts to send 1 or 2"


def test_each_frame_measures_its_own_chart(owner):
    _, inventory = owner
    long = {
        **MARGIN,
        "categories": [f"Category number {n}" for n in range(12)],
        "series": [{"name": "Margin", "values": list(range(1, 13))}],
    }
    (problem,) = _assigned(inventory, _multi(REVENUE, MARGIN, REVENUE, long)).problems
    assert "charts[3]" in problem.problem
    assert "charts[0]" not in problem.problem


# --- rendering ------------------------------------------------------------------------


def test_two_charts_refill_their_examples_in_place(owner):
    template, inventory = owner
    result = _render(template, inventory, _multi(REVENUE, SHARE))
    left, right = _delivered_charts(result)
    assert list(left.chart.plots[0].series[0].values) == [12, 15, 18, 21]
    assert list(right.chart.plots[0].categories) == ["Europe", "Americas", "Asia"]
    assert right.chart.chart_type == XL_CHART_TYPE.PIE
    assert left.chart.chart_title.text_frame.text == "Revenue"
    parts = {left.chart.part.partname, right.chart.part.partname}
    assert len(parts) == 2
    for frame in (left, right):
        assert "srgbClr" not in frame.chart._chartSpace.xml


def test_four_charts_fill_the_grid_in_reading_order(owner):
    template, inventory = owner
    charts = [_chart(f"KPI {n}", [n, n + 1, n + 2]) for n in range(1, 5)]
    result = _render(template, inventory, _multi(*charts))
    delivered = _delivered_charts(result)
    assert [list(c.chart.plots[0].series[0].values)[0] for c in delivered] == [
        1,
        2,
        3,
        4,
    ]
    assert len({c.chart.part.partname for c in delivered}) == 4


def test_a_text_box_above_a_chart_holds_its_title():
    template, inventory = _template(lambda p: add_charts_slide(p, 2, headings=True))
    (pattern,) = _slides(inventory)
    # The boxes take the charts' titles, never the slide's text.
    assert zone_summary(pattern).startswith("title ≤4 words (≈33 chars); charts×2 (")
    result = _render(template, inventory, _multi(REVENUE, MARGIN))
    deck = Presentation(io.BytesIO(result.pptx))
    shapes = [cast(Any, s) for s in deck.slides[1].shapes]
    texts = [s.text_frame.text for s in shapes if s.has_text_frame]
    assert "Revenue" in texts and "Margin" in texts
    assert not any(c.chart.has_title for c in _delivered_charts(result))


def test_a_lost_chart_stops_the_build(owner, monkeypatch):
    template, inventory = owner
    real = checks.chart_count
    monkeypatch.setattr(checks, "chart_count", lambda slide: max(0, real(slide) - 1))
    with pytest.raises(RenderError, match="shows 1 of its 2 chart"):
        _render(template, inventory, _multi(REVENUE, MARGIN))


# --- edits and the tool ----------------------------------------------------------------


def _delivered(inventory: TemplateInventory, *slides: Dict[str, Any]) -> DeckPlan:
    plan = _plan(*slides)
    deck, patterns, _ = assign(inventory, plan, "off").deck(plan)
    return deck.model_copy(
        update={
            "slides": [
                s.model_copy(update={"built_on": p.id if p else None})
                for s, p in zip(deck.slides, patterns)
            ]
        }
    )


def _changed(base: DeckPlan, change: Dict[str, Any]) -> DeckPlan:
    return apply_changes(base, [SlideChange.model_validate(change)])


def test_one_value_of_the_second_chart_is_edited_in_place(owner):
    template, inventory = owner
    base = _delivered(inventory, _multi(REVENUE, SHARE))
    edit = {"op": "edit", "slide": 2, "field": "charts[1].series[0].values[2]"}
    plan = _changed(base, {**edit, "value": 42})
    assignment = assign(inventory, plan, "off")
    assert assignment.patterns[1] is not None
    assert assignment.patterns[1].id == base.slides[1].built_on
    result = render_deck(template, inventory, plan, assignment.patterns)
    _, pie = _delivered_charts(result)
    assert list(pie.chart.plots[0].series[0].values) == [45, 35, 42]


def test_chart_paths_on_a_slide_of_several_are_refused(owner):
    _, inventory = owner
    base = _delivered(inventory, _multi(REVENUE, MARGIN))
    with pytest.raises(ChangeError) as caught:
        _changed(
            base, {"op": "edit", "slide": 2, "field": "chart.type", "value": "bar"}
        )
    assert "edit 'charts[i].type' (i from 0 to 1)" in str(caught.value)


def test_four_charts_sent_in_an_edit_move_the_slide(owner):
    _, inventory = owner
    base = _delivered(inventory, _multi(REVENUE, MARGIN))
    edit = {"op": "edit", "slide": 2, "field": "charts"}
    plan = _changed(base, {**edit, "value": [REVENUE, MARGIN, REVENUE, MARGIN]})
    pattern = assign(inventory, plan, "off").patterns[1]
    assert pattern is not None and _frames(pattern) == 4


def test_charts_is_offered_only_with_a_slide_of_several(owner):
    _, inventory = owner
    _, single = _template(lambda p: add_main_chart_slide(p))
    slide = tools._args_schema(inventory).model_json_schema()["$defs"]["BuildSlide"]
    assert "exactly N charts" in slide["properties"]["charts"]["description"]
    lone = tools._args_schema(single).model_json_schema()["$defs"]["BuildSlide"]
    assert "charts" not in lone["properties"]
    described = tools._description(inventory, "off", 2)
    assert "two to four related measures side by side" in described
    assert "charts×2" in described and "charts×4" in described
    assert "two to four" not in tools._description(single, "off", 2)
    # The units written beside the values reach the one-measure check.
    mixed = {**MARGIN, "series": [{"name": "M", "values": ["8 %", "3 M€"]}]}
    sent = tools._args_schema(inventory)(slides=[_multi(REVENUE, mixed)])
    plan = tools._to_plan(sent)
    assert plan.slides[0].charts[1].series[0].units == ["%", "M€"]
