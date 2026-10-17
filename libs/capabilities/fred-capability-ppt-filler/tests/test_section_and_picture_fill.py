# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Built-in deck: a section slide's number zone and filled picture zones."""

from __future__ import annotations

import io
from typing import Any, Dict, List, cast

import pytest
from deck_builders import png_bytes
from fred_capability_ppt_filler.builder.assign import assign
from fred_capability_ppt_filler.builder.builtin import builtin_template
from fred_capability_ppt_filler.builder.plan import DeckPlan
from fred_capability_ppt_filler.builder.render import RenderResult, render_deck
from fred_capability_ppt_filler.images import prepare_picture
from pptx import Presentation

COVER = {"title": "Quarterly review", "subtitle": "Sales"}


@pytest.fixture(scope="module")
def builtin():
    return builtin_template()


# Twice as wide as high: a frame of another shape must crop it, never letterbox it.
WIDE = {f"img-{n}": prepare_picture(png_bytes(400, 200, "blue")) for n in range(4)}


def _render(builtin, *slides: Dict[str, Any], cover=COVER) -> RenderResult:
    template, inventory = builtin
    plan = DeckPlan.model_validate({"slides": [cover, *slides]})
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems, [p.as_dict() for p in assignment.problems]
    return render_deck(template, inventory, plan, assignment.patterns, images=WIDE)


def _numbers(result: RenderResult) -> List[str]:
    """The text of each section slide's number zone, in deck order."""
    deck = Presentation(io.BytesIO(result.pptx))
    found = []
    for slide in deck.slides:
        if not slide.slide_layout.name.startswith("Fred Section"):
            continue
        texts = [
            cast(Any, s).text_frame.text
            for s in slide.placeholders
            if s.placeholder_format.idx == 103 and s.has_text_frame
        ]
        found.append(texts[0] if texts else "")
    return found


def test_section_slides_are_numbered_in_deck_order(builtin):
    body = {"title": "Pipeline", "bullets": ["Two new accounts", "One renewal"]}
    result = _render(
        builtin,
        {"title": "Market", "kind": "section"},
        body,
        {"title": "Team", "kind": "section", "number": "07"},
        body,
        {"title": "Next steps", "kind": "section"},
    )
    # Zero-padded as the example's "01"; a number the plan gives is kept.
    assert _numbers(result) == ["01", "07", "03"]
