# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Chart slides: the plan's data, the zone the template offers for it, and a
native chart drawn in the theme's colours. A few tests are adapted from
``fred-capability-pptx-generator``, by Shuigami."""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any, Dict, List, cast

import pytest
from builder_decks import (
    add_chart_slide,
    add_main_chart_slide,
    build_layouts_only_template,
    build_slides,
)
from fred_capability_ppt_filler.builder import tools
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import Assignment, assign, infer_kinds
from fred_capability_ppt_filler.builder.builtin import builtin_template
from fred_capability_ppt_filler.builder.changes import SlideChange, apply_changes
from fred_capability_ppt_filler.builder.labelling import kind_fits_slots
from fred_capability_ppt_filler.builder.models import TemplateInventory
from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanChart
from fred_capability_ppt_filler.builder.render import RenderResult, render_deck
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Inches

COVER = {"title": "Quarterly review", "subtitle": "Sales"}
REVENUE = {
    "type": "column",
    "categories": ["Q1", "Q2", "Q3", "Q4"],
    "series": [{"name": "Revenue", "values": [12, 15, 18, 21]}],
}


@pytest.fixture(scope="module")
def builtin():
    return builtin_template()


def _plan(*slides: Dict[str, Any]) -> DeckPlan:
    return DeckPlan.model_validate({"slides": [COVER, *slides]})


def _chart_slide(chart: Dict[str, Any], **extra: Any) -> Dict[str, Any]:
    return {"title": "Revenue", "chart": chart, **extra}


def _assigned(inventory: TemplateInventory, *slides: Dict[str, Any]) -> Assignment:
    return assign(inventory, _plan(*slides), "off")


def _render(template: bytes, inventory: TemplateInventory, *slides) -> RenderResult:
    plan = _plan(*slides)
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems, [p.as_dict() for p in assignment.problems]
    return render_deck(template, inventory, plan, assignment.patterns)


def _charts(result: RenderResult, slide: int = 1) -> List[Any]:
    deck = Presentation(io.BytesIO(result.pptx))
    return [s for s in deck.slides[slide].shapes if getattr(s, "has_chart", False)]


def _chart(result: RenderResult, slide: int = 1) -> Any:
    (frame,) = _charts(result, slide)
    return frame.chart


def _problem(assignment: Assignment) -> Dict[str, Any]:
    (problem,) = assignment.problems
    return problem.as_dict()


# --- analysis ---------------------------------------------------------------------


def test_builtin_chart_slides_are_chart_patterns(builtin):
    _, inventory = builtin
    charts = [p for p in inventory.patterns if p.kind == "chart"]
    assert [(p.name, p.included, p.source) for p in charts] == [
        ("Fred Chart", True, "slide"),
        ("Fred Chart Wide", True, "slide"),
    ]
    takeaway, wide = charts
    zone = next(z for z in takeaway.zones if z.role == "chart")
    assert (zone.key, zone.chart_type, zone.max_lines) == ("ph:101", "column", 12)
    assert [z.role for z in wide.zones] == ["title", "chart"]
    assert next(z for z in wide.zones if z.role == "chart").box.w == 910
    assert kind_fits_slots("chart", takeaway)
    assert not kind_fits_slots("bullets", takeaway)


def test_an_example_slides_main_chart_is_its_zone():
    inventory = analyze_template(build_slides([add_main_chart_slide])).inventory
    assert inventory is not None
    (pattern,) = [p for p in inventory.patterns if p.source == "slide"]
    assert (pattern.kind, pattern.included) == ("chart", True)
    zone = next(z for z in pattern.zones if z.role == "chart")
    assert zone.key.startswith("sp:") and zone.chart_type == "pie"


def test_two_charts_or_a_small_chart_keep_the_slide_out():
    def small_chart(presentation):
        slide = add_main_chart_slide(presentation)
        frame = next(s for s in slide.shapes if s.has_chart)
        frame.width, frame.height = Inches(2), Inches(1.5)
        return slide

    inventory = analyze_template(build_slides([add_chart_slide, small_chart])).inventory
    assert inventory is not None
    slides = [p for p in inventory.patterns if p.source == "slide"]
    assert [p.exclusion_reason for p in slides] == ["chart", "unsupported_parts"]


def test_only_a_large_content_placeholder_can_hold_a_chart(builtin):
    inventory = analyze_template(build_layouts_only_template()).inventory
    assert inventory is not None
    content = next(p for p in inventory.patterns if p.name == "Title and Content")
    assert [z.chart_capable for z in content.zones] == [False, True]
    assert not kind_fits_slots("chart", content)
    # The built-in deck types its text bodies as such: only its table layout's
    # content placeholder could host a chart.
    _, fred = builtin
    capable = [p.name for p in fred.patterns if any(z.chart_capable for z in p.zones)]
    assert capable == ["Fred Table"]


# --- the plan -----------------------------------------------------------------------


def test_a_chart_is_read_from_what_small_models_write():
    chart = PlanChart.model_validate(
        {
            "chart_type": "Donut chart",
            "labels": ["Europe", "Asie"],
            "data": [{"label": "Parts", "values": ["45 %", "55,5 %"]}],
        }
    )
    assert (chart.type, chart.categories, chart.unit) == (
        "doughnut",
        ["Europe", "Asie"],
        "%",
    )
    assert chart.series[0].name == "Parts" and chart.series[0].values == [45.0, 55.5]
    pairs = PlanChart.model_validate({"series": {"Q1": "1 200", "Q2": "€3.5"}})
    assert pairs.categories == ["Q1", "Q2"] and pairs.series[0].values == [1200, 3.5]
    # Units that disagree are not a unit; the values are never rescaled.
    assert pairs.unit is None


def test_an_unknown_type_is_built_as_columns_and_said():
    chart = PlanChart.model_validate({**REVENUE, "type": "waterfall"})
    assert (chart.type, chart.asked_type) == ("column", "waterfall")
    again = PlanChart.model_validate(chart.model_dump())
    assert again.asked_type == "waterfall"


def test_design_keys_are_set_aside_and_unknown_keys_are_content():
    chart = PlanChart.model_validate(
        {
            **REVENUE,
            "colors": ["#ff0000"],
            "series": [{"name": "R", "values": [1, 2, 3, 4], "color": "red"}],
            "footnote": "Source: finance",
        }
    )
    assert chart.style_fields == ["colors", "series.color"]
    assert chart.ignored_fields == ["footnote 'Source: finance'"]


def test_a_chart_slide_infers_the_chart_kind_first():
    plan = _plan(_chart_slide(REVENUE)).slides[1]
    assert infer_kinds(plan) == ["chart", "bullets", "two_columns", "image_text"]


# --- validation -----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("chart", "problem", "option"),
    [
        (
            {**REVENUE, "series": [{"name": "Costs", "values": [1, 2, 3]}]},
            "chart.series[0] 'Costs' has 3 values for 4 categories",
            "give one value per category",
        ),
        (
            {**REVENUE, "series": [{"name": "R", "values": [1, "n/a", 3, 4]}]},
            "chart.series[0].values[1] 'n/a' is not a number",
            "drop that category",
        ),
        (
            {
                **REVENUE,
                "type": "pie",
                "series": [
                    {"name": "A", "values": [1, 2, 3, 4]},
                    {"name": "B", "values": [1, 2, 3, 4]},
                ],
            },
            "a 'pie' chart shows one series, 2 given",
            "use type 'column'",
        ),
        (
            {
                **REVENUE,
                "series": [{"name": f"S{n}", "values": [1, 2, 3, 4]} for n in range(6)],
            },
            "chart: 6 series; at most 4 stay readable",
            "keep the 4 that matter",
        ),
        (
            {
                "type": "pie",
                "categories": [f"C{n}" for n in range(8)],
                "series": [{"name": "Share", "values": [1] * 8}],
            },
            "a 'pie' chart: 8 slices; at most 6 stay readable",
            "group the smallest as 'Other'",
        ),
        (
            {"type": "column", "categories": [], "series": []},
            "chart: no data",
            "give `chart`",
        ),
    ],
)
def test_bad_chart_data_is_a_named_revision(builtin, chart, problem, option):
    _, inventory = builtin
    found = _problem(_assigned(inventory, _chart_slide(chart)))
    assert problem in found["problem"]
    assert any(option in o for o in found["options"]), found["options"]


def test_too_much_for_the_zone_is_measured_never_split(builtin):
    _, inventory = builtin
    many = {
        "type": "column",
        "categories": [f"M{n}" for n in range(20)],
        "series": [{"name": "Sales", "values": list(range(20))}],
    }
    found = _problem(_assigned(inventory, _chart_slide(many)))
    assert "chart: 20 categories; this chart holds 12" in found["problem"]
    label = "Second quarter of the fiscal year, after the reorganisation of all teams"
    long = {**REVENUE, "categories": ["Q1", label, "Q3", "Q4"]}
    found = _problem(_assigned(inventory, _chart_slide(long)))
    assert (
        f"chart.categories[1] is {len(label)} chars; each label holds"
        in found["problem"]
    )
    assert any(o.startswith("shorten every category label") for o in found["options"])


# --- where the chart goes ---------------------------------------------------------------


def _chartless(inventory: TemplateInventory) -> TemplateInventory:
    """The template without its chart slides nor any content area able to host one."""
    patterns = [
        p.model_copy(
            update={
                "zones": [
                    z.model_copy(update={"chart_capable": False}) for z in p.zones
                ]
            }
        )
        for p in inventory.patterns
        if p.kind != "chart"
    ]
    return inventory.model_copy(update={"patterns": patterns})


def test_a_takeaway_goes_beside_the_chart_and_a_bare_chart_spans_the_slide(builtin):
    _, inventory = builtin
    assignment = _assigned(
        inventory,
        _chart_slide(REVENUE, bullets=["Q4 is 75 % above Q1"]),
        _chart_slide(REVENUE),
    )
    assert not assignment.problems
    labels = [p.label for p in assignment.patterns if p is not None]
    assert labels[1:] == [
        "Charts stay native and editable",
        "One chart across the slide",
    ]


def _three_chart_slides(inventory: TemplateInventory) -> TemplateInventory:
    """A centred chart and two chart slides with a text box beside it, as a
    customer template draws them."""
    takeaway, wide = [p for p in inventory.patterns if p.kind == "chart"]
    others = [p for p in inventory.patterns if p.kind != "chart"]
    side = [
        takeaway.model_copy(update={"id": f"slide:{n}", "label": label})
        for n, label in ((901, "Left"), (902, "Right"))
    ]
    centred = wide.model_copy(update={"label": "Centred"})
    return inventory.model_copy(update={"patterns": [*others, centred, *side]})


def test_a_bare_chart_keeps_the_centred_slide_and_text_alternates_the_sides(builtin):
    inventory = _three_chart_slides(builtin[1])
    bare = [_chart_slide(REVENUE, title=f"Revenue {n}") for n in range(6)]
    assignment = assign(inventory, _plan(*bare), "template_slide", 1)
    assert not assignment.problems
    # Fewer empty zones beats variety: never a side slide with its text box empty.
    assert [p.label for p in assignment.patterns[1:] if p] == ["Centred"] * 6
    assert not [r for r in assignment.repeats if r.revisable]
    told = [
        _chart_slide(REVENUE, title=f"Revenue {n}", bullets=["Q4 leads"])
        for n in range(4)
    ]
    labels = [
        p.label
        for p in assign(inventory, _plan(*told), "template_slide", 1).patterns[1:]
        if p
    ]
    assert set(labels) == {"Left", "Right"}
    assert all(a != b for a, b in zip(labels, labels[1:]))


def test_a_content_area_hosts_the_chart_when_no_chart_slide_exists():
    template = build_layouts_only_template()
    inventory = analyze_template(template).inventory
    assert inventory is not None
    assignment = _assigned(inventory, _chart_slide(REVENUE))
    assert not assignment.problems
    pattern = assignment.patterns[1]
    assert pattern is not None and pattern.name == "Title and Content"
    assert any(
        "instead of 'chart', the chart in its content area" in n
        for n in assignment.notes
    )
    result = _render(template, inventory, _chart_slide(REVENUE))
    (frame,) = _charts(result)
    zone = next(z for z in pattern.zones if z.chart_capable)
    assert round(frame.left * 1000 / inventory.slide_width) == zone.box.x
    assert round(frame.width * 1000 / inventory.slide_width) == zone.box.w
    built = result.report.built[1]
    assert built.kind == "bullets"
    assert [(n.code, n.kind) for n in built.notes] == [("other_kind", "chart")]
    # The content area holds the chart, not the "Click to add text" prompt.
    deck = Presentation(io.BytesIO(result.pptx))
    assert len(deck.slides[1].shapes) == 2


def test_bullets_in_the_only_content_area_leave_no_place_for_the_chart():
    found_inventory = analyze_template(build_layouts_only_template()).inventory
    assert found_inventory is not None
    kept = ("Title Slide", "Title and Content")
    inventory = found_inventory.model_copy(
        update={"patterns": [p for p in found_inventory.patterns if p.name in kept]}
    )
    found = _problem(_assigned(inventory, _chart_slide(REVENUE, bullets=["Why"])))
    assert "chart: this slide has no place for it" in found["problem"]
    assert "give the chart a slide of its own (title and chart)" in found["options"]


def test_a_template_without_any_chart_place_asks_for_another_form(builtin):
    inventory = _chartless(builtin[1])
    found = _problem(_assigned(inventory, _chart_slide(REVENUE)))
    assert found["problem"] == "chart: this template has no slide for a chart"
    assert "give the 4 values as items of a 'figures' slide" not in " ".join(
        found["options"]
    )  # the built-in figures slide holds 3
    assert found["options"][-1] == "keep the key numbers as bullets"
    small = {**REVENUE, "categories": ["Q1", "Q2", "Q3"]}
    small["series"] = [{"name": "Revenue", "values": [1, 2, 3]}]
    found = _problem(_assigned(inventory, _chart_slide(small)))
    assert any("as items of a 'figures' slide" in o for o in found["options"])


def test_a_chart_slide_needs_a_chart(builtin):
    _, inventory = builtin
    assignment = _assigned(
        inventory, {"title": "Revenue", "bullets": ["Up"], "kind": "chart"}
    )
    assert all(p is None or p.kind != "chart" for p in assignment.patterns)


def test_chart_slides_are_never_asked_to_become_another_kind(builtin):
    _, inventory = builtin
    slides = [_chart_slide(REVENUE, title=f"Revenue {n}") for n in range(4)]
    assignment = assign(inventory, _plan(*slides), "kind", 2)
    repeats = [r for r in assignment.repeats if r.kind == "chart"]
    assert repeats and all(not r.suggested and not r.others for r in repeats)


# --- the drawing ------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("kind", "xl_type"),
    [
        ("column", XL_CHART_TYPE.COLUMN_CLUSTERED),
        ("bar", XL_CHART_TYPE.BAR_CLUSTERED),
        ("line", XL_CHART_TYPE.LINE_MARKERS),
        ("pie", XL_CHART_TYPE.PIE),
        ("doughnut", XL_CHART_TYPE.DOUGHNUT),
        ("stacked_column", XL_CHART_TYPE.COLUMN_STACKED),
        ("area", XL_CHART_TYPE.AREA),
    ],
)
def test_each_type_builds_a_native_chart(builtin, kind, xl_type):
    template, inventory = builtin
    result = _render(template, inventory, _chart_slide({**REVENUE, "type": kind}))
    chart = _chart(result)
    assert chart.chart_type == xl_type
    assert list(chart.plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]
    assert list(chart.plots[0].series[0].values) == [12, 15, 18, 21]


def test_colours_and_fonts_come_from_the_theme(builtin):
    template, inventory = builtin
    three = {
        **REVENUE,
        "series": [{"name": f"S{n}", "values": [1, 2, 3, 4]} for n in range(3)],
    }
    result = _render(template, inventory, _chart_slide(three))
    xml = _chart(result)._chartSpace.xml
    for accent in ("accent1", "accent2", "accent3"):
        assert f'<a:schemeClr val="{accent}"/>' in xml
    assert "srgbClr" not in xml
    assert '<a:latin typeface="+mn-lt"/>' in xml


def test_the_chart_takes_the_zone_and_no_example_frame_is_left(builtin):
    template, inventory = builtin
    result = _render(template, inventory, _chart_slide(REVENUE, bullets=["Up"]))
    (frame,) = _charts(result)
    pattern = next(p for p in inventory.patterns if p.label.startswith("Charts stay"))
    zone = next(z for z in pattern.zones if z.role == "chart")
    assert round(frame.top * 1000 / inventory.slide_height) == zone.box.y
    assert round(frame.width * 1000 / inventory.slide_width) == zone.box.w
    deck = Presentation(io.BytesIO(result.pptx))
    xml = deck.slides[1].shapes._spTree.xml
    assert xml.count("<c:chart ") == 1 and "Charts use theme colours" not in xml


def test_native_chart_slide_adds_editable_chart_with_title(builtin):
    # Adapted from fred-capability-pptx-generator.
    template, inventory = builtin
    result = _render(template, inventory, _chart_slide({**REVENUE, "title": "Revenue"}))
    assert _chart(result).chart_title.text_frame.text == "Revenue"
    untitled = _chart(_render(template, inventory, _chart_slide(REVENUE)))
    assert not untitled.has_title


@pytest.mark.parametrize(
    ("kind", "series_count", "expect_legend"),
    [
        ("pie", 1, True),
        ("doughnut", 1, True),
        ("column", 1, False),
        ("column", 2, True),
    ],
)
def test_native_chart_legend_sits_below_when_shown(
    builtin, kind, series_count, expect_legend
):
    # Adapted from fred-capability-pptx-generator.
    template, inventory = builtin
    series = [{"name": f"S{n}", "values": [45, 35, 20]} for n in range(series_count)]
    chart = {"type": kind, "categories": ["Europe", "Intl", "France"], "series": series}
    built = _chart(_render(template, inventory, _chart_slide(chart)))
    assert built.has_legend is expect_legend
    if expect_legend:
        assert built.legend.position == XL_LEGEND_POSITION.BOTTOM
        assert built.legend.include_in_layout is False


def test_values_carry_their_unit_and_pies_their_percentages(builtin):
    template, inventory = builtin
    money = {
        **REVENUE,
        "unit": "M€",
        "series": [{"name": "R", "values": [1.5, 2, 3, 4]}],
    }
    built = _chart(_render(template, inventory, _chart_slide(money)))
    assert built.value_axis.tick_labels.number_format == '#,##0.0" M€"'
    assert built.plots[0].data_labels.number_format == '#,##0.0" M€"'
    pie = _chart(_render(template, inventory, _chart_slide({**REVENUE, "type": "pie"})))
    assert pie.plots[0].data_labels.show_percentage


PIE = {
    "type": "pie",
    "categories": ["Ventes", "R&D", "Support"],
    "series": [{"name": "Effectif", "values": [27, 15, 58]}],
}


@pytest.mark.parametrize(
    "key, raw, mode",
    [
        ("show", "value", "value"),
        ("show", "Values", "value"),
        ("display", "numbers", "value"),
        ("labels_as", "absolute", "value"),
        ("show", "pourcentages", "share"),
        ("show_as", "%", "share"),
        ("show", "both", "both"),
        ("display", "les deux", "both"),
    ],
)
def test_pie_labels_mode_is_read_from_loose_names(key, raw, mode):
    chart = PlanChart.model_validate({**PIE, key: raw})
    assert (chart.show, chart.asked_show, chart.ignored_fields) == (mode, None, [])
    assert PlanChart.model_validate(chart.model_dump()).show == mode


def test_pie_labels_mode_defaults_to_shares_and_keeps_labels_as_categories():
    chart = PlanChart.model_validate(
        {**PIE, "categories": None, "labels": ["A", "B", "C"]}
    )
    assert (chart.show, chart.categories) == (None, ["A", "B", "C"])
    schema = PlanChart.model_json_schema()["properties"]["show"]
    assert "set show: 'value'" in schema["description"]
    assert set(schema["anyOf"][0]["enum"]) == {"share", "value", "both"}


def test_an_unknown_pie_labels_mode_is_a_named_revision(builtin):
    _, inventory = builtin
    chart = PlanChart.model_validate({**PIE, "show": "bubbles"})
    assert (chart.show, chart.asked_show) == (None, "bubbles")
    assert PlanChart.model_validate(chart.model_dump()).asked_show == "bubbles"
    found = _problem(_assigned(inventory, _chart_slide({**PIE, "show": "bubbles"})))
    assert "chart.show 'bubbles' is not a label mode" in found["problem"]


def _labels_xml(built: Any) -> Any:
    (labels,) = built._chartSpace.iter(qn("c:dLbls"))
    return labels


@pytest.mark.parametrize(
    "show, value, percent, fmt, linked, separator",
    [
        (None, "0", "1", "0%", "0", None),
        ("share", "0", "1", "0%", "0", None),
        ("value", "1", "0", '#,##0" collaborateurs"', "0", None),
        ("both", "1", "1", '#,##0" collaborateurs"', "1", "\n"),
    ],
)
def test_pie_labels_show_what_the_plan_asks(
    builtin, show, value, percent, fmt, linked, separator
):
    template, inventory = builtin
    chart = {**PIE, "unit": "collaborateurs", **({"show": show} if show else {})}
    built = _chart(_render(template, inventory, _chart_slide(chart)))
    labels = _labels_xml(built)
    assert labels.find(qn("c:showVal")).get("val") == value
    assert labels.find(qn("c:showPercent")).get("val") == percent
    number = labels.find(qn("c:numFmt"))
    assert (number.get("formatCode"), number.get("sourceLinked")) == (fmt, linked)
    found = labels.find(qn("c:separator"))
    assert (found.text if found is not None else None) == separator
    # Linked, the value takes the series' own format: the unit, never "0%".
    (cache,) = built._chartSpace.iter(qn("c:formatCode"))
    assert cache.text == '#,##0" collaborateurs"'


def test_a_unit_hidden_by_percentages_is_told(builtin):
    template, inventory = builtin
    result = _render(
        template, inventory, _chart_slide({**PIE, "unit": "collaborateurs"})
    )
    assert (
        "Slide 2: the pie's labels show percentages, so unit 'collaborateurs' is not "
        "shown. To show the numbers with their unit, set chart.show to 'value' "
        "('both' for the number and its percentage)." in result.report.notes
    )
    built = result.report.built[1]
    assert [n.code for n in built.notes] == ["chart_shows_shares"]
    shares = _render(template, inventory, _chart_slide({**PIE, "unit": "%"}))
    assert not any("chart.show" in n for n in shares.report.notes)
    values = _render(
        template,
        inventory,
        _chart_slide({**PIE, "unit": "collaborateurs", "show": "value"}),
    )
    assert not any("chart.show" in n for n in values.report.notes)
    assert values.report.built[1].notes == []


def test_a_labels_mode_on_another_chart_type_is_told(builtin):
    template, inventory = builtin
    result = _render(template, inventory, _chart_slide({**REVENUE, "show": "share"}))
    assert (
        "Slide 2: chart.show 'share' was not applied: it sets the slice labels of a "
        "pie or doughnut only; a 'column' chart shows its values. Nothing to do; drop "
        "chart.show, or set chart.type to 'pie' or 'doughnut'." in result.report.notes
    )
    assert result.report.built[1].notes == []
    plain = _render(template, inventory, _chart_slide(REVENUE))
    assert not any("chart.show" in n for n in plain.report.notes)


def test_a_type_fallback_and_design_keys_are_told(builtin):
    template, inventory = builtin
    result = _render(
        template,
        inventory,
        _chart_slide({**REVENUE, "type": "waterfall", "color": "red"}),
    )
    assert _chart(result).chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    notes = " ".join(result.report.notes)
    assert "chart type 'waterfall' is not offered; built as 'column'" in notes
    assert "the chart's color were ignored" in notes
    built = result.report.built[1]
    assert [(n.code, n.kind) for n in built.notes] == [
        ("chart_type_changed", "waterfall")
    ]


# --- the template's own example chart (D11) ---------------------------------------------

# The "Fred Chart" example slide's column chart: one series in a literal colour,
# with its own value labels.
DESIGNER_FILL = '<a:srgbClr val="4851C3"/>'


def _parts(result: RenderResult) -> List[str]:
    with zipfile.ZipFile(io.BytesIO(result.pptx)) as archive:
        return archive.namelist()


def test_the_example_chart_keeps_its_design_and_takes_the_plans_data(builtin):
    template, inventory = builtin
    money = {
        **REVENUE,
        "unit": "M€",
        "series": [{"name": "R", "values": [1.5, 2, 3, 4]}],
    }
    result = _render(template, inventory, _chart_slide(money, bullets=["Q4 leads"]))
    built = _chart(result)
    assert built.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert list(built.plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]
    assert list(built.plots[0].series[0].values) == [1.5, 2, 3, 4]
    xml = built._chartSpace.xml
    assert DESIGNER_FILL in xml and "schemeClr" not in xml.split("<c:ser>")[1]
    # Its own series labels stay, in the plan's unit.
    labels = built.plots[0]._element.find(f"{qn('c:ser')}/{qn('c:dLbls')}")
    assert labels.find(qn("c:numFmt")).get("formatCode") == '#,##0.0" M€"'
    assert not built.has_legend and not built.has_title
    assert built.value_axis.tick_labels.number_format == '#,##0.0" M€"'


def test_another_type_on_the_example_slide_is_a_fresh_theme_chart(builtin):
    template, inventory = builtin
    pie = {**REVENUE, "type": "pie"}
    result = _render(template, inventory, _chart_slide(pie, bullets=["Q4 leads"]))
    (frame,) = _charts(result)
    xml = frame.chart._chartSpace.xml
    assert frame.chart.chart_type == XL_CHART_TYPE.PIE
    assert "srgbClr" not in xml and '<a:schemeClr val="accent4"/>' in xml


def test_an_example_used_twice_gets_its_own_parts_each_time(builtin):
    template, inventory = builtin
    result = _render(
        template,
        inventory,
        _chart_slide(REVENUE, bullets=["Up"]),
        _chart_slide(
            {**REVENUE, "series": [{"name": "R", "values": [4, 3, 2, 1]}]},
            bullets=["Down"],
        ),
    )
    names = _parts(result)
    assert len(names) == len(set(names))
    charts = [n for n in names if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]
    books = [n for n in names if n.startswith("ppt/embeddings/")]
    assert len(charts) == 2 and len(books) == 2
    deck = Presentation(io.BytesIO(result.pptx))
    frames = _charts(result, 1) + _charts(result, 2)
    assert len({f.chart.part.partname for f in frames}) == 2
    assert len({f.chart.part.chart_workbook.xlsx_part.partname for f in frames}) == 2
    assert [list(f.chart.plots[0].series[0].values) for f in frames] == [
        [12, 15, 18, 21],
        [4, 3, 2, 1],
    ]
    for i in (1, 2):
        # Each slide relates its one chart once.
        rels = deck.slides[i].part.rels
        assert sum(r.reltype.endswith("/chart") for r in rels.values()) == 1


def test_series_beyond_the_examples_take_theme_accents(builtin):
    template, inventory = builtin
    two = {
        **REVENUE,
        "series": [
            {"name": "2025", "values": [1, 2, 3, 4]},
            {"name": "2026", "values": [2, 3, 4, 5]},
        ],
    }
    built = _chart(_render(template, inventory, _chart_slide(two, bullets=["Up"])))
    first, second = [s._element.xml for s in built.plots[0].series]
    assert DESIGNER_FILL in first and '<a:schemeClr val="accent2"/>' in second
    # Several series: a legend, and no value labels.
    assert built.has_legend and "<c:dLbls>" not in built._chartSpace.xml


def test_an_intent_on_the_example_chart_keeps_the_designers_colour(builtin):
    template, inventory = builtin
    chart = {**REVENUE, "highlight": "Q3", "sort": "desc"}
    result = _render(template, inventory, _chart_slide(chart, bullets=["Q4 leads"]))
    built = _chart(result)
    assert list(built.plots[0].categories) == ["Q4", "Q3", "Q2", "Q1"]
    ser = built.plots[0]._element.find(qn("c:ser"))
    muted = [d.find(qn("c:idx")).get("val") for d in ser.findall(qn("c:dPt"))]
    assert muted == ["0", "2", "3"] and DESIGNER_FILL in built._chartSpace.xml
    (label,) = ser.findall(f"{qn('c:dLbls')}/{qn('c:dLbl')}")
    assert label.find(qn("c:idx")).get("val") == "1"
    assert [n.code for n in result.report.built[1].notes if n.field] == []


_STYLE_REL = "http://schemas.microsoft.com/office/2011/relationships/chartStyle"


def _with_chart_style(template: bytes) -> bytes:
    """The template with a chart style part on each example chart, as PowerPoint saves."""
    deck = Presentation(io.BytesIO(template))
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                part = cast(Any, shape).chart.part
                name = deck.part.package.next_partname("/ppt/charts/style%d.xml")
                style = Part(
                    name,
                    "application/vnd.ms-office.chartstyle+xml",
                    part.package,
                    b"<cs/>",
                )
                part.relate_to(style, _STYLE_REL)
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_a_chart_style_part_is_copied_with_its_chart(builtin):
    template, inventory = builtin
    result = _render(
        _with_chart_style(template),
        inventory,
        _chart_slide(REVENUE, bullets=["Up"]),
        _chart_slide(REVENUE, bullets=["Up"]),
    )
    styles = {
        rel.target_part.partname
        for i in (1, 2)
        for frame in _charts(result, i)
        for rel in frame.chart.part.rels.values()
        if rel.reltype == _STYLE_REL
    }
    assert len(styles) == 2


def _with_linked_workbook(template: bytes) -> bytes:
    """The template with its example chart's data linked to an outside workbook,
    which phase B does not copy."""
    deck = Presentation(io.BytesIO(template))
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                cast(Any, shape).chart.part.relate_to(
                    "file:///data.xlsx", RT.OLE_OBJECT, is_external=True
                )
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_a_chart_that_does_not_copy_falls_back_to_a_fresh_one(builtin):
    template, inventory = builtin
    result = _render(
        _with_linked_workbook(template),
        inventory,
        _chart_slide(REVENUE, bullets=["Up"]),
    )
    (frame,) = _charts(result)
    xml = frame.chart._chartSpace.xml
    assert list(frame.chart.plots[0].series[0].values) == [12, 15, 18, 21]
    assert "srgbClr" not in xml and '<a:schemeClr val="accent1"/>' in xml
    names = _parts(result)
    assert len([n for n in names if n.startswith("ppt/charts/chart")]) == 1


# --- changes and the tool ---------------------------------------------------------------


def _delivered(inventory: TemplateInventory, *slides: Dict[str, Any]) -> DeckPlan:
    plan = _plan(*slides)
    assignment = assign(inventory, plan, "off")
    deck, patterns, _ = assignment.deck(plan)
    return deck.model_copy(
        update={
            "delivered": True,
            "slides": [
                s.model_copy(update={"built_on": p.id if p else None})
                for s, p in zip(deck.slides, patterns)
            ],
        }
    )


def _changed(base: DeckPlan, *changes: Dict[str, Any], unplaced=None, cleared=None):
    return apply_changes(
        base, [SlideChange.model_validate(c) for c in changes], unplaced, cleared
    )


def test_a_chart_value_is_edited_on_the_same_template_slide(builtin):
    template, inventory = builtin
    base = _delivered(inventory, _chart_slide(REVENUE))
    edit = {"op": "edit", "slide": 2, "field": "chart.series[0].values[2]", "value": 30}
    plan = _changed(base, edit)
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems
    pattern = assignment.patterns[1]
    assert pattern is not None and pattern.id == base.slides[1].built_on
    result = render_deck(template, inventory, plan, assignment.patterns)
    assert list(_chart(result).plots[0].series[0].values) == [12, 15, 30, 21]


def test_a_pie_is_edited_to_show_its_numbers(builtin):
    template, inventory = builtin
    base = _delivered(inventory, _chart_slide({**PIE, "unit": "collaborateurs"}))
    edit = {"op": "edit", "slide": 2, "field": "chart.show", "value": "value"}
    plan = _changed(base, edit)
    assert plan.slides[1].chart is not None and plan.slides[1].chart.show == "value"
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems
    result = render_deck(template, inventory, plan, assignment.patterns)
    labels = _chart(result).plots[0].data_labels
    assert (labels.show_value, labels.show_percentage) == (True, False)
    assert labels.number_format == '#,##0" collaborateurs"'
    assert not any("chart.show" in n for n in result.report.notes)


def test_a_category_added_without_its_values_is_named(builtin):
    _, inventory = builtin
    base = _delivered(inventory, _chart_slide(REVENUE))
    plan = _changed(
        base, {"op": "edit", "slide": 2, "field": "chart.categories[4]", "value": "Q5"}
    )
    found = _problem(assign(inventory, plan, "off"))
    assert "chart.series[0] 'Revenue' has 4 values for 5 categories" in found["problem"]


def test_a_kind_edit_clears_the_chart_the_new_kind_cannot_show(builtin):
    _, inventory = builtin
    base = _delivered(inventory, _chart_slide(REVENUE, bullets=["Up"]))
    cleared: Dict[int, List[str]] = {}

    def unplaced(slide):
        from fred_capability_ppt_filler.builder.assign import unplaced as found

        return found(inventory, slide, slide.kind)

    plan = _changed(
        base,
        {"op": "edit", "slide": 2, "field": "kind", "value": "bullets"},
        unplaced=unplaced,
        cleared=cleared,
    )
    assert plan.slides[1].chart is None and cleared == {2: ["chart"]}
    plan = _changed(base, {"op": "edit", "slide": 2, "field": "chart", "value": None})
    assert plan.slides[1].chart is None


def test_the_tool_offers_charts_only_where_the_template_has_a_place(builtin):
    _, fred = builtin
    schema = tools._args_schema(fred).model_json_schema()
    assert "chart" in schema["$defs"]["BuildSlide"]["properties"]
    description = tools._description(fred, "template_slide", 2)
    assert (
        "chart {type, categories, series [{name, values}], title, unit, show}"
        in description
    )
    assert "to display numbers instead of percentages on a pie, set show: 'value'" in (
        description
    )
    assert '"chart.show"' in description
    assert "categories max 12" in tools.catalogue(fred)
    none = _chartless(fred)
    schema = tools._args_schema(none).model_json_schema()
    assert "chart" not in schema["$defs"]["BuildSlide"]["properties"]
    assert "chart {" not in tools._description(none, "template_slide", 2)
    content = analyze_template(build_layouts_only_template()).inventory
    assert content is not None
    assert "or a chart in place of the bullets" in tools.catalogue(content)


# --- more chart types: stacked bar, radar, scatter, bubble ----------------------------------

SCATTER = {
    "type": "scatter",
    "x_title": "Price",
    "x_unit": "€",
    "y_title": "Units sold",
    "series": [
        {
            "name": "Products",
            "points": [{"x": 10, "y": 120}, {"x": 15, "y": 90}, {"x": 22, "y": 60}],
        }
    ],
}
BUBBLE = {
    "type": "bubble",
    "x_title": "Growth",
    "x_unit": "%",
    "y_title": "Share",
    "unit": "%",
    "size_title": "Revenue",
    "size_unit": "M€",
    "series": [
        {
            "name": "Markets",
            "points": [
                {"x": 5, "y": 20, "size": 100, "label": "France"},
                {"x": 8, "y": 12, "size": 40, "label": "Spain"},
                {"x": 2, "y": 30, "size": 250, "label": "Italy"},
            ],
        }
    ],
}
RADAR = {
    "type": "radar",
    "categories": ["Price", "Speed", "Support", "Fit"],
    "series": [
        {"name": "Offer A", "values": [3, 4, 5, 2]},
        {"name": "Offer B", "values": [4, 2, 3, 5]},
    ],
}
STACKED_BAR = {
    "type": "stacked_bar",
    "categories": ["Sales", "R&D", "Support"],
    "series": [
        {"name": "2025", "values": [4, 3, 2]},
        {"name": "2026", "values": [5, 3, 3]},
    ],
}


def _points(chart: PlanChart) -> List[tuple]:
    return [(p.x, p.y, p.size, p.label) for s in chart.series for p in s.points]


@pytest.mark.parametrize(
    "written",
    [
        {"series": [{"name": "P", "points": [[1, 2], [3, 4]]}]},
        {"series": [{"name": "P", "points": [{"x": 1, "y": 2}, {"x": "3", "y": "4"}]}]},
        {"series": [{"name": "P", "data": [{"x": 1, "y": 2}, {"x": 3, "y": 4}]}]},
        {"series": [{"name": "P", "x": [1, 3], "y": [2, 4]}]},
        {"x": [1, 3], "y": [2, 4]},
        {"x": [1, 3], "values": [2, 4]},
        {"points": [[1, 2], [3, 4]]},
        {"categories": ["1", "3"], "series": [{"name": "P", "values": [2, 4]}]},
    ],
)
def test_scatter_points_are_read_from_what_small_models_write(written):
    chart = PlanChart.model_validate({"type": "scatter", **written})
    assert chart.categories == []
    assert _points(chart) == [(1, 2, None, None), (3, 4, None, None)]
    assert chart.ignored_fields == [] and chart.style_fields == []


def test_bubble_sizes_labels_and_units_are_read_per_axis():
    chart = PlanChart.model_validate(
        {
            "type": "bulles",
            "x_label": "Croissance",
            "y_axis": "Part",
            "series": [
                {
                    "name": "Marchés",
                    "points": [
                        {"x": "5 %", "y": "20 €", "radius": "100", "name": "France"},
                        ["8 %", "12 €", 40, "Spain"],
                    ],
                }
            ],
        }
    )
    assert chart.type == "bubble"
    assert (chart.x_title, chart.y_title) == ("Croissance", "Part")
    assert (chart.x_unit, chart.unit, chart.size_unit) == ("%", "€", None)
    assert _points(chart) == [(5, 20, 100, "France"), (8, 12, 40, "Spain")]
    # Sizes as a list are data on a bubble, a design key elsewhere.
    listed = PlanChart.model_validate(
        {"type": "bubble", "x": [1, 2], "y": [3, 4], "size": [5, 6]}
    )
    assert _points(listed) == [(1, 3, 5, None), (2, 4, 6, None)]
    assert listed.style_fields == []
    styled = PlanChart.model_validate(
        {**REVENUE, "size": "large", "x_axis": {"min": 0}}
    )
    assert styled.style_fields == ["size", "x_axis"]


def test_x_is_the_categories_of_a_column_chart_and_new_names_are_read():
    column = PlanChart.model_validate(
        {"type": "column", "x": ["Q1", "Q2"], "values": [1, 2]}
    )
    assert column.categories == ["Q1", "Q2"] and column.series[0].values == [1, 2]
    names = {
        "stacked bar": "stacked_bar",
        "Horizontal stacked bar chart": "stacked_bar",
        "spider": "radar",
        "toile d'araignée": "column",
        "araignée": "radar",
        "nuage de points": "scatter",
        "XY": "scatter",
        "bubbles": "bubble",
        "stacked column": "stacked_column",
    }
    for written, kind in names.items():
        assert PlanChart.model_validate({**REVENUE, "type": written}).type == kind


def test_a_points_chart_survives_its_dump_with_its_unread_values():
    chart = PlanChart.model_validate(
        {**SCATTER, "series": [{"name": "P", "points": [{"x": "n/a", "y": 1}]}]}
    )
    again = PlanChart.model_validate(chart.model_dump())
    assert again.series[0].points[0].unread == {"x": "n/a"}
    assert (again.x_title, again.x_unit) == ("Price", "€")


def _with_points(*points: Any, **chart: Any) -> Dict[str, Any]:
    return {**SCATTER, "series": [{"name": "P", "points": list(points)}], **chart}


@pytest.mark.parametrize(
    ("chart", "problem", "option"),
    [
        (
            {
                **SCATTER,
                "categories": ["Q1", "Q2", "Q3"],
                "series": [{"values": [1, 2, 3]}],
            },
            "a 'scatter' chart takes points; series[0] has values over categories 'Q1'",
            "use type 'line' or 'column'",
        ),
        (
            {**RADAR, "series": [{"name": "A", "points": [[1, 2], [2, 3], [3, 4]]}]},
            "a 'radar' chart takes categories and values, not points",
            "use type 'scatter' for points",
        ),
        (
            _with_points([1, 2], ["n/a", 3], [3, 4]),
            "chart.series[0].points[1].x 'n/a' is not a number",
            "drop that point",
        ),
        (
            _with_points([1, 2], {"x": 2}, [3, 4]),
            "chart.series[0].points[1] has no y",
            "give each point {x, y}",
        ),
        (
            {**SCATTER, "x_title": None},
            "chart: no x_title; say what the horizontal axis measures ('Price')",
            "set chart.x_title and chart.y_title",
        ),
        (
            {**BUBBLE, "series": [{"name": "M", "points": [[1, 2, 3], [2, 3, 0]]}]},
            "chart.series[0].points[1].size is 0; a bubble needs a positive size",
            "give a size above 0",
        ),
        (
            {**BUBBLE, "series": [{"name": "M", "points": [[1, 2, 3], [2, 3]]}]},
            "chart.series[0].points[1] has no size",
            "positive size",
        ),
        (
            _with_points([1, 2, 5], [2, 3, 6], [3, 4, 7]),
            "points carry size: use type 'bubble'",
            "set chart.type to 'bubble'",
        ),
        (
            _with_points([1, 2], [2, 3]),
            "a 'scatter' chart: 2 points; it needs 3 to 50",
            "give more points",
        ),
        (
            _with_points(*[[i, i] for i in range(51)]),
            "a 'scatter' chart: 51 points; at most 50 stay readable",
            "keep the 50 that matter",
        ),
        (
            {
                **BUBBLE,
                "series": [{"name": "M", "points": [[i, i, 1] for i in range(16)]}],
            },
            "a 'bubble' chart: 16 points; at most 15 stay readable",
            "use type 'scatter' (no size)",
        ),
        (
            {
                **RADAR,
                "categories": ["A", "B"],
                "series": [{"name": "S", "values": [1, 2]}],
            },
            "radar: 2 criteria; a radar needs 3 to 8",
            "use type 'column' or 'bar'",
        ),
        (
            {
                **RADAR,
                "series": [{"name": f"S{n}", "values": [1, 2, 3, 4]} for n in range(4)],
            },
            "a 'radar' chart: 4 series; at most 3 stay readable",
            "keep the 3 that matter",
        ),
        (
            _with_points(
                {"x": 1, "y": 2, "label": "Very long product name"}, [2, 3], [3, 4]
            ),
            "chart.series[0].points[0].label 'Very long product name' is 4 words",
            "shorten each point label to 2 words",
        ),
    ],
)
def test_bad_points_and_radar_data_are_named_revisions(builtin, chart, problem, option):
    _, inventory = builtin
    found = _problem(_assigned(inventory, _chart_slide(chart)))
    assert problem in found["problem"]
    assert any(option in o for o in found["options"]), found["options"]


def test_a_long_axis_title_is_measured_with_its_unit(builtin):
    _, inventory = builtin
    long = {**SCATTER, "x_title": "Average selling price per unit " * 8}
    found = _problem(_assigned(inventory, _chart_slide(long)))
    assert (
        "chart.x_title is" in found["problem"] and "with its unit" in found["problem"]
    )


@pytest.mark.parametrize(
    ("chart", "xl_type"),
    [
        (STACKED_BAR, XL_CHART_TYPE.BAR_STACKED),
        ({**STACKED_BAR, "stacked_as": "percent"}, XL_CHART_TYPE.BAR_STACKED_100),
        (RADAR, XL_CHART_TYPE.RADAR_MARKERS),
        (SCATTER, XL_CHART_TYPE.XY_SCATTER),
        (BUBBLE, XL_CHART_TYPE.BUBBLE),
    ],
)
def test_each_new_type_builds_a_native_theme_chart(builtin, chart, xl_type):
    template, inventory = builtin
    result = _render(template, inventory, _chart_slide(chart))
    built = _chart(result)
    assert built.chart_type == xl_type
    xml = built._chartSpace.xml
    assert '<a:schemeClr val="accent1"' in xml and "srgbClr" not in xml
    assert '<a:latin typeface="+mn-lt"/>' in xml
    assert not result.report.notes


def test_a_scatter_titles_its_axes_and_names_few_points(builtin):
    template, inventory = builtin
    labelled = {
        **SCATTER,
        "series": [
            {
                "name": "Products",
                "points": [
                    {"x": 10, "y": 120, "label": "Alpha"},
                    {"x": 15, "y": 90},
                    {"x": 22.5, "y": 60, "label": "Gamma"},
                ],
            }
        ],
    }
    built = _chart(_render(template, inventory, _chart_slide(labelled)))
    x_axis, y_axis = built.category_axis, built.value_axis
    assert x_axis.axis_title.text_frame.text == "Price (€)"
    assert y_axis.axis_title.text_frame.text == "Units sold"
    assert x_axis.tick_labels.number_format == '#,##0.0" €"'
    assert y_axis.tick_labels.number_format == "#,##0"
    assert list(built.plots[0].series[0].values) == [120, 90, 60]
    points = built.plots[0].series[0].points
    assert [points[i].data_label.text_frame.text for i in (0, 2)] == ["Alpha", "Gamma"]
    assert not built.has_legend
    many = _with_points(*[{"x": i, "y": i, "label": f"P{i}"} for i in range(13)])
    crowded = _chart(_render(template, inventory, _chart_slide(many)))
    assert "<c:dLbl>" not in crowded._chartSpace.xml


def test_a_bubble_says_what_its_size_measures(builtin):
    template, inventory = builtin
    built = _chart(
        _render(template, inventory, _chart_slide({**BUBBLE, "title": "Markets"}))
    )
    assert built.chart_title.text_frame.text == "Markets\nBubble size: Revenue (M€)"
    assert built.plots[0].bubble_scale == 50
    assert built.value_axis.axis_title.text_frame.text == "Share (%)"


def test_intents_on_the_new_types(builtin):
    template, inventory = builtin
    result = _render(
        template,
        inventory,
        _chart_slide({**BUBBLE, "highlight": "Italy", "sort": "desc"}),
        _chart_slide({**RADAR, "sort": "desc", "highlight": "Offer B"}),
        _chart_slide({**STACKED_BAR, "sort": "desc"}),
    )
    bubble = _chart(result, 1)
    ser = bubble.plots[0]._element.find(qn("c:ser"))
    assert [d.find(qn("c:idx")).get("val") for d in ser.findall(qn("c:dPt"))] == [
        "0",
        "1",
        "2",
    ]
    assert "a 'bubble' chart keeps its own order" in " ".join(result.report.notes)
    assert "a 'radar' chart keeps its own order" in " ".join(result.report.notes)
    # A stacked bar sorts by its totals, the largest on top.
    assert list(_chart(result, 3).plots[0].categories) == ["Support", "R&D", "Sales"]


def test_axis_fields_a_type_does_not_draw_are_told(builtin):
    template, inventory = builtin
    result = _render(
        template, inventory, _chart_slide({**REVENUE, "x_title": "Quarter"})
    )
    assert any(
        "chart.x_title not drawn on a 'column' chart" in n for n in result.report.notes
    )


@pytest.mark.parametrize(
    ("example", "chart"),
    [
        (XL_CHART_TYPE.XY_SCATTER, SCATTER),
        (XL_CHART_TYPE.RADAR_MARKERS, RADAR),
        (XL_CHART_TYPE.BAR_STACKED_100, STACKED_BAR),
    ],
)
def test_an_example_chart_of_a_new_type_is_reused(example, chart):
    template = build_slides([lambda p: add_main_chart_slide(p, example)])
    inventory = analyze_template(template).inventory
    assert inventory is not None
    pattern = next(p for p in inventory.patterns if p.source == "slide")
    zone = next(z for z in pattern.zones if z.role == "chart")
    assert zone.chart_type == chart["type"]
    result = _render(template, inventory, _chart_slide(chart, bullets=["Look"]))
    built = _chart(result)
    xml = built._chartSpace.xml
    # The example's own part, refilled: never a second chart, never 100 % stacking.
    assert "percentStacked" not in xml
    assert len([n for n in _parts(result) if n.startswith("ppt/charts/chart")]) == 1
    if chart is SCATTER:
        assert list(built.plots[0].series[0].values) == [120, 90, 60]
        assert built.category_axis.axis_title.text_frame.text == "Price (€)"


def test_a_point_is_edited_and_a_type_change_across_shapes_is_named(builtin):
    template, inventory = builtin
    base = _delivered(inventory, _chart_slide(SCATTER))
    edit = {
        "op": "edit",
        "slide": 2,
        "field": "chart.series[0].points[2].y",
        "value": 75,
    }
    plan = _changed(base, edit)
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems
    result = render_deck(template, inventory, plan, assignment.patterns)
    assert list(_chart(result).plots[0].series[0].values) == [120, 90, 75]
    title = _changed(
        base, {"op": "edit", "slide": 2, "field": "chart.x_title", "value": "Cost"}
    )
    assert title.slides[1].chart is not None and title.slides[1].chart.x_title == "Cost"
    column = _changed(
        base, {"op": "edit", "slide": 2, "field": "chart.type", "value": "column"}
    )
    found = _problem(assign(inventory, column, "off"))
    assert (
        "a 'column' chart takes categories and values, not points" in found["problem"]
    )


def test_the_tool_guides_the_model_to_the_new_types(builtin):
    _, fred = builtin
    description = tools._description(fred, "template_slide", 2)
    assert "two measures per item → chart scatter (bubble for a third)" in description
    assert "items scored on several criteria → chart radar" in description
    assert "points [{x, y, size, label}]}] with x_title, x_unit, y_title" in description
    assert '"chart.series[0].points[2].y"' in description
    catalogue = tools.catalogue(fred)
    assert "stacked_bar, radar, scatter, bubble" in catalogue
    schema = tools._args_schema(fred).model_json_schema()
    kinds = schema["$defs"]["PlanChart"]["properties"]["type"]
    assert kinds["enum"][-4:] == ["stacked_bar", "radar", "scatter", "bubble"]
    assert "radar (3-8 criteria compared for 1-3 items)" in kinds["description"]
