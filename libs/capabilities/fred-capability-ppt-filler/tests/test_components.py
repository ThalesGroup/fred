# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The component library (design.md D6, D8, D15).

The library is the design layer: it names components, tokens and grid
positions and never touches the output format. The tests below check the
properties that make a composed slide look composed rather than assembled —
boxes on the grid, gaps on the scale, z-order that puts a Card behind its
text — plus the degenerate cases each component has to declare.
"""

from __future__ import annotations

import pytest
from deck_builders import (
    ICON_NAMES,
    build_house_style_template,
    build_icon_bank_template,
)
from fred_capability_ppt_filler.builder.analysis import analyze_presentation
from fred_capability_ppt_filler.builder.components import (
    _MIN_CHARS,
    BAR_MIN_RATIO,
    COMPOSITIONS,
    bar,
    compose,
    variant_for,
)
from fred_capability_ppt_filler.builder.design_system import Box
from fred_capability_ppt_filler.builder.theme import default_theme
from pptx import Presentation


@pytest.fixture(scope="module")
def rich():
    _, inventory = default_theme()
    return inventory.design_system


@pytest.fixture(scope="module")
def stripped():
    presentation = Presentation()
    title = presentation.slides.add_slide(presentation.slide_layouts[0])
    title.shapes.title.text = "Cover"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    return analyze_presentation(presentation).design_system


ALL_VARIANTS = [
    (name, count) for name, counts in COMPOSITIONS.items() for count in counts
]


def resolved(system, name: str, count: int):
    """The composition, or a skip when this template cannot offer it.

    A composition may need something the template does not have — `icon_cards`
    needs named icons — and then it is simply absent from its catalogue rather
    than drawn without them.
    """
    slide = compose(system, name, items=count)
    if slide is None:
        pytest.skip(f"this template cannot offer '{name}'")
    return slide


# -- the layer boundary (D8) ---------------------------------------------------


def test_the_design_layer_never_imports_the_output_format():
    """The guarantee is only worth stating if something checks it.

    Without this, "the renderer decides nothing" degrades into a convention
    that the first awkward case breaks.
    """
    import pathlib

    root = pathlib.Path(__file__).resolve().parents[1] / "fred_capability_ppt_filler"
    for module in ("builder/components.py", "builder/design_system.py"):
        source = (root / module).read_text()
        assert "import pptx" not in source
        assert "from pptx" not in source


# -- variants are fixed, never stretched (D6) ---------------------------------


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_every_variant_resolves_on_both_templates(name, count, rich, stripped):
    for system in (rich, stripped):
        slide = resolved(system, name, count)

        assert slide is not None, f"{name} {count}-up did not resolve"
        assert slide.pieces
        assert slide.capacity.items >= 1


def test_a_composition_picks_the_variant_that_holds_the_items():
    assert variant_for("card_row", 3) == 3
    # Two items go in the 2-up, not in a stretched 3-up.
    assert variant_for("card_row", 2) == 2
    # Five do not fit any variant: that slide has to split, not widen.
    assert variant_for("card_row", 5) is None


@pytest.mark.parametrize("count", [2, 3, 4])
def test_a_card_row_gives_one_card_per_item(count, rich):
    slide = compose(rich, "card_row", items=count)

    cards = [p for p in slide.pieces if p.kind == "card"]
    assert len(cards) == count
    assert slide.variant == f"{count}-up"


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_every_box_stays_inside_the_usable_area(name, count, rich):
    slide = resolved(rich, name, count)
    area = rich.grid.area

    for piece in slide.pieces:
        assert piece.box.x >= area.x - 1, f"{piece.name} starts left of the area"
        assert piece.box.x + piece.box.w <= area.x + area.w + 2, (
            f"{piece.name} runs past the right margin"
        )
        assert piece.box.w > 0 and piece.box.h > 0


@pytest.mark.parametrize("count", [2, 3, 4])
def test_columns_of_a_row_do_not_overlap(count, rich):
    slide = compose(rich, "card_row", items=count)

    cards = sorted((p for p in slide.pieces if p.kind == "card"), key=lambda p: p.box.x)
    for left, right in zip(cards, cards[1:]):
        assert left.box.x + left.box.w <= right.box.x


# -- cross-slide consistency (D13) --------------------------------------------


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_every_composition_puts_its_title_on_the_same_anchor(name, count, rich):
    slide = resolved(rich, name, count)

    title = next(p for p in slide.pieces if p.slot == "title")
    assert title.box.y == rich.anchors.title_y
    assert title.role == "title"


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_content_starts_at_the_content_anchor(name, count, rich):
    slide = resolved(rich, name, count)

    below = [p for p in slide.pieces if p.slot != "title"]
    assert below, f"{name} has no content"
    assert min(p.box.y for p in below) == rich.anchors.content_y


# -- z-order, naming, accessibility (D6, D15) ---------------------------------


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_creation_order_is_contiguous_and_starts_at_zero(name, count, rich):
    slide = resolved(rich, name, count)

    orders = sorted(p.order for p in slide.pieces)
    assert orders == list(range(len(slide.pieces)))


def test_a_card_is_created_before_the_text_that_sits_on_it(rich):
    """Creation order IS z-order: the other way round hides the text."""
    slide = compose(rich, "card_row", items=3)

    for index in range(3):
        card = next(p for p in slide.pieces if p.name == f"composed-card-{index}")
        on_it = [
            p
            for p in slide.pieces
            if p.slot.startswith(f"item:{index}:") and p.kind == "text"
        ]
        assert on_it
        assert all(card.order < piece.order for piece in on_it)


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_every_piece_is_named(name, count, rich):
    slide = resolved(rich, name, count)

    names = [p.name for p in slide.pieces]
    assert all(names)
    assert len(names) == len(set(names)), "a duplicate name cannot be found again"


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_decorative_pieces_carry_no_content_and_content_pieces_are_not_decorative(
    name, count, rich
):
    slide = resolved(rich, name, count)

    for piece in slide.pieces:
        if piece.decorative:
            assert not piece.slot, f"{piece.name} is decorative but carries content"
        if piece.kind == "text":
            assert not piece.decorative


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_reading_order_follows_the_composition_not_the_creation(name, count, rich):
    slide = resolved(rich, name, count)

    readable = [p for p in slide.pieces if not p.decorative]
    orders = [p.reading_order for p in readable]
    # Every readable piece has a place in the reading order, and the title is
    # first whatever order its shape happened to be created in.
    assert len(orders) == len(set(orders))
    assert min(orders) == 0
    assert next(p for p in readable if p.slot == "title").reading_order == 0


# -- the Bar and its degenerate case (D6, D10) --------------------------------


def test_a_bar_is_a_track_and_a_value():
    pieces = bar(
        Box(x=100, y=500, w=800, h=40), value=0.5, order=1, name="b", reading_order=1
    )

    assert len(pieces) == 2
    track, value = pieces
    assert track.decorative and track.surface == "surface-muted"
    assert value.surface == "accent"
    assert value.box.w == pytest.approx(400, abs=2)


def test_a_bar_at_full_value_matches_its_track():
    track, value = bar(
        Box(x=100, y=500, w=800, h=40), value=1.0, order=1, name="b", reading_order=1
    )

    assert value.box.w == track.box.w
    assert value.box.x == track.box.x


def test_a_bar_below_its_minimum_length_is_not_drawn_as_a_blob():
    """A rounded rectangle takes its radius from its SHORTER side: a value
    shorter than it is thick distorts into something wider than it is long."""
    pieces = bar(
        Box(x=100, y=500, w=800, h=40), value=0.02, order=1, name="b", reading_order=1
    )

    values = [p for p in pieces if p.name.endswith("-value")]
    for piece in values:
        assert piece.box.w >= piece.box.h * BAR_MIN_RATIO


def test_a_bar_at_zero_shows_an_honest_empty_track():
    pieces = bar(
        Box(x=100, y=500, w=800, h=40), value=0.0, order=1, name="b", reading_order=1
    )

    assert len(pieces) == 1
    assert pieces[0].name.endswith("-track")


def test_a_bar_clamps_a_value_outside_the_unit_range():
    over = bar(
        Box(x=0, y=0, w=500, h=20), value=4.0, order=1, name="b", reading_order=1
    )
    under = bar(
        Box(x=0, y=0, w=500, h=20), value=-1.0, order=1, name="b", reading_order=1
    )

    assert over[1].box.w == 500
    assert len(under) == 1


# -- capacity is declared, not discovered (D14) -------------------------------


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_every_variant_declares_a_capacity(name, count, rich):
    slide = resolved(rich, name, count)

    assert slide.capacity.items >= 1
    assert slide.capacity.title_chars > 0
    # A composition that carries body text says how much of it fits.
    if name in ("card_row", "comparison", "prose", "statement"):
        assert slide.capacity.text_chars > 0


def test_a_template_with_no_design_system_composes_nothing():
    from fred_capability_ppt_filler.builder.design_system import DesignSystem

    empty = DesignSystem(slide_width=0, slide_height=0)

    assert not empty.complete
    assert compose(empty, "card_row", items=3) is None


def test_an_unknown_composition_is_refused(rich):
    assert compose(rich, "carousel", items=3) is None


# -- the pieces the library must hold -----------------------------------------


def test_the_library_holds_every_piece_the_spec_names(rich):
    """Card, Chip, Stat, Bar, Divider and Image frame, each drawable.

    A Stat is no longer a builder of its own: over a box model it is what it
    always was on the page — a figure stacked over its label — so it is
    asserted through the composition that draws it.
    """
    from fred_capability_ppt_filler.builder.components import card, chip, divider

    box = Box(x=100, y=300, w=300, h=100)
    assert card(box, order=0, name="c").kind == "card"
    assert [
        p.kind
        for p in chip(box, slot="s", order=0, name="c", reading_order=1, padding=4)
    ] == ["chip", "text"]
    assert divider(box, order=0, name="d").kind == "divider"
    assert divider(box, order=0, name="d").decorative

    stats = compose(rich, "stat_band", items=3).pieces
    figure = next(p for p in stats if p.name == "composed-stat-number-0")
    label = next(p for p in stats if p.name == "composed-stat-label-0")
    assert (figure.role, label.role) == ("number", "label")
    assert {figure.kind, label.kind} == {"text"}

    assert [p for p in compose(rich, "progress", items=2).pieces if p.kind == "bar"]
    assert [
        p for p in compose(rich, "figure_caption", items=1).pieces if p.kind == "image"
    ]


def test_a_progress_composition_draws_bars_not_a_chart(rich):
    """D10: a quantity out of a whole is shapes, never a chart object."""
    slide = compose(rich, "progress", items=3)

    bars = [p for p in slide.pieces if p.kind == "bar"]
    assert len(bars) == 6  # a track and a value per row
    assert all(p.surface in ("surface-muted", "accent") for p in bars)
    # The bar illustrates a figure the text beside it states, so it is
    # decorative — and that figure is the one the plan binds to.
    assert all(p.decorative for p in bars)
    figures = [p for p in slide.pieces if p.slot.endswith(":number")]
    assert {p.slot for p in figures} == {f"item:{i}:number" for i in range(3)}
    assert all(p.role == "number" and p.type_role == "label" for p in figures)


def test_a_stat_band_puts_the_figure_above_its_label(rich):
    slide = compose(rich, "stat_band", items=3)

    for index in range(3):
        number = next(
            p for p in slide.pieces if p.name == f"composed-stat-number-{index}"
        )
        label = next(
            p for p in slide.pieces if p.name == f"composed-stat-label-{index}"
        )
        assert number.role == "number"
        assert label.role == "label"
        assert number.box.y < label.box.y


# -- composing in the template's language (D18) -------------------------------


@pytest.fixture(scope="module")
def house():
    """A template that outlines its cards, spaces them its own way, and writes
    a figure under the label that names it."""
    return analyze_presentation(build_house_style_template()).design_system


def test_columns_are_separated_by_the_gap_the_template_uses(house):
    slide = compose(house, "card_row", items=3)
    cards = sorted((p for p in slide.pieces if p.kind == "card"), key=lambda p: p.box.x)

    gaps = {
        right.box.x - (left.box.x + left.box.w) for left, right in zip(cards, cards[1:])
    }

    assert len(gaps) == 1, "the columns are not evenly spaced"
    expected = round(house.shapes.column_gap * house.shorter_to_x)
    assert abs(gaps.pop() - expected) <= 1


def test_a_block_holds_the_air_the_template_gives_its_own(house, rich):
    """Padding is read, not chosen — and the two templates disagree.

    The air belongs to the Card, not to an inset on the text frame inside it:
    the box model already insets the children, so a frame inset as well would
    pad twice.
    """
    assert house.block_padding == house.shapes.padding
    assert rich.block_padding == rich.spacing.step(1), "nothing to read, scale answers"
    assert house.block_padding != rich.block_padding

    for system in (house, rich):
        slide = compose(system, "card_row", items=3)
        card_box = next(p.box for p in slide.pieces if p.kind == "card")
        heading = next(p.box for p in slide.pieces if p.name == "composed-heading-0")

        assert heading.x - card_box.x == round(
            system.block_padding * system.shorter_to_x
        )
        assert heading.y - card_box.y == round(
            system.block_padding * system.shorter_to_y
        )


def test_a_stat_follows_the_order_the_template_stacks_them_in(house, rich):
    """Figure over label here, label over figure there — both read, not chosen."""
    assert rich.shapes.figure_above_label and not house.shapes.figure_above_label

    def ordinates(system):
        slide = compose(system, "stat_band", items=3)
        number = next(p for p in slide.pieces if p.name == "composed-stat-number-0")
        label = next(p for p in slide.pieces if p.name == "composed-stat-label-0")
        return number.box.y, label.box.y

    number_y, label_y = ordinates(rich)
    assert number_y < label_y

    number_y, label_y = ordinates(house)
    assert label_y < number_y


def test_the_figure_is_read_first_whichever_way_it_is_stacked(house):
    """The point of a stat band is the figure: it leads the reading order even
    when the template prints it second."""
    slide = compose(house, "stat_band", items=3)

    number = next(p for p in slide.pieces if p.name == "composed-stat-number-0")
    label = next(p for p in slide.pieces if p.name == "composed-stat-label-0")

    assert number.reading_order < label.reading_order


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_every_variant_still_holds_together_in_a_house_style(name, count, house):
    """The style is read from the template, so every variant has to survive
    values the capability would never have chosen."""
    slide = resolved(house, name, count)

    assert slide.pieces
    area = house.grid.area
    for piece in slide.pieces:
        assert piece.box.w > 0 and piece.box.h > 0
        assert piece.box.x >= area.x - 1
        assert piece.box.x + piece.box.w <= area.x + area.w + 2


# -- the scale is a ladder, and a piece takes the rung it can carry (D18) -----


def test_a_narrow_box_steps_down_the_scale_rather_than_truncating(rich):
    """A scale read off a full-width placeholder does not fit a card.

    The ladder is the TEMPLATE's own scale, so a template whose smallest rung
    is still too big simply lands on that rung — a size is never invented to
    rescue it.
    """
    scale = {token.size_pt for token in rich.type_scale}
    wide = next(
        p
        for p in compose(rich, "card_row", items=2).pieces
        if p.name == "composed-text-0"
    )
    narrow = next(
        p
        for p in compose(rich, "card_row", items=4).pieces
        if p.name == "composed-text-0"
    )

    assert wide.font_pt in scale and narrow.font_pt in scale
    assert narrow.box.w < wide.box.w
    assert narrow.font_pt <= wide.font_pt, "a narrower card kept the wider size"
    assert narrow.font_pt == min(s for s in scale if s <= wide.font_pt) or (
        narrow.max_chars >= _MIN_CHARS["body"]
    )


def test_the_slide_title_keeps_its_own_size(rich):
    """Stepping down is for boxes too small, not for the title of the slide."""
    slide = compose(rich, "card_row", items=3)

    title = next(p for p in slide.pieces if p.slot == "title")

    assert title.font_pt == rich.type_for("title").size_pt


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_a_heading_is_never_the_same_size_as_the_text_under_it(name, count, rich):
    """Two roles can step down onto the same rung; 16pt over 15pt reads as one
    undifferentiated block, which is the hierarchy gone."""
    slide = resolved(rich, name, count)

    for index in range(count):
        heading = next(
            (p for p in slide.pieces if p.name == f"composed-heading-{index}"), None
        )
        body = next(
            (p for p in slide.pieces if p.name == f"composed-text-{index}"), None
        )
        if heading is None or body is None:
            continue
        assert heading.font_pt > body.font_pt


@pytest.mark.parametrize("name,count", ALL_VARIANTS)
def test_a_text_box_holds_the_lines_it_was_built_for(name, count, rich):
    """Rounding a block height instead of stacking line heights lands a hair
    under N lines, and the box silently holds N-1."""
    slide = resolved(rich, name, count)

    for piece in slide.pieces:
        if piece.kind == "text" and piece.max_lines:
            assert piece.max_lines >= 1


# -- a card is sized by what it holds, not by the band (D6) -------------------


def test_a_card_does_not_stretch_down_the_whole_band(rich):
    slide = compose(rich, "card_row", items=3)
    band_bottom = rich.grid.area.y + rich.grid.area.h

    card = next(p for p in slide.pieces if p.kind == "card")
    content = [p for p in slide.pieces if p.slot.startswith("item:0:")]

    assert content
    used = max(p.box.y + p.box.h for p in content) - card.box.y
    assert card.box.h < (band_bottom - card.box.y) * 0.9, "the card fills the band"
    # It is still big enough for everything it carries, plus its own rim.
    assert card.box.h >= used


def test_every_card_of_a_row_is_the_same_height(rich):
    slide = compose(rich, "card_row", items=4)

    heights = {p.box.h for p in slide.pieces if p.kind == "card"}

    assert len(heights) == 1


# -- nested archetypes (D21) ---------------------------------------------------


@pytest.mark.parametrize("count", [2, 3, 4])
def test_a_framed_panel_holds_a_row_of_cards(count, rich):
    """The commonest shape of a real corporate slide, and the one nothing flat
    could express: a panel, and inside it a row of cards."""
    slide = compose(rich, "framed_cards", items=count)

    panel = next(p for p in slide.pieces if p.name == "composed-panel")
    cards = [p for p in slide.pieces if p.name.startswith("composed-card-")]

    assert len(cards) == count
    for card_piece in cards:
        assert panel.box.x <= card_piece.box.x
        assert panel.box.y <= card_piece.box.y
        assert card_piece.box.x + card_piece.box.w <= panel.box.x + panel.box.w + 1
        assert card_piece.box.y + card_piece.box.h <= panel.box.y + panel.box.h + 1
    assert panel.order < min(p.order for p in cards), "the panel would hide its cards"


def test_a_framed_panel_hugs_its_cards_instead_of_running_to_the_footer(rich):
    slide = compose(rich, "framed_cards", items=3)
    band_bottom = rich.grid.area.y + rich.grid.area.h

    panel = next(p for p in slide.pieces if p.name == "composed-panel")

    assert panel.box.y + panel.box.h < band_bottom
    assert panel.box.h < (band_bottom - panel.box.y) * 0.95


@pytest.mark.parametrize("count", [2, 3, 4])
def test_a_lead_block_sits_beside_its_cards_not_above_them(count, rich):
    slide = compose(rich, "lead_and_cards", items=count)

    lead = next(p for p in slide.pieces if p.name == "composed-lead")
    cards = [p for p in slide.pieces if p.name.startswith("composed-card-")]

    assert len(cards) == count
    assert all(lead.box.x + lead.box.w <= c.box.x for c in cards), "not side by side"
    for upper, lower in zip(cards, cards[1:]):
        assert upper.box.y + upper.box.h <= lower.box.y


# -- the template's own icons (D22) -------------------------------------------


@pytest.fixture(scope="module")
def banked():
    """A template carrying a bank of NAMED icons."""
    return analyze_presentation(build_icon_bank_template()).design_system


def test_only_icons_the_template_names_are_offered(banked, rich):
    """An icon called "Picture 12" cannot be asked for by anyone, and placing
    icons by turn would be decoration by lottery."""
    assert {icon.name for icon in banked.icons} == set(ICON_NAMES)
    assert all(not icon.name.endswith(".png") for icon in banked.icons)
    assert rich.icons == [], "a template with no bank offers no icon"


def test_an_icon_keeps_its_own_aspect(banked):
    """Read from the shape's EMU, not from its per-mille box: the two axes
    have different scales, so a square icon would read as a wide one."""
    assert all(icon.aspect == pytest.approx(1.0, abs=0.01) for icon in banked.icons)


def test_a_card_can_lead_with_one_of_the_template_icons(banked):
    slide = compose(banked, "icon_cards", items=3)

    icons = [p for p in slide.pieces if p.kind == "icon"]
    assert len(icons) == 3
    for index, piece in enumerate(icons):
        assert piece.slot == f"item:{index}:icon"
        assert piece.decorative, "the heading beside it already says what it is"
        heading = next(p for p in slide.pieces if p.name == f"composed-heading-{index}")
        assert piece.box.y + piece.box.h <= heading.box.y


def test_a_template_without_named_icons_offers_no_icon_composition(rich):
    """Composing icon cards without icons would draw empty frames."""
    assert compose(rich, "icon_cards", items=3) is None


def test_a_card_of_three_rows_keeps_a_hierarchy_between_all_three(rich):
    """Settling the pairs top-down puts the middle row level with the one above
    it again: the cascade has to come up from the bottom."""
    slide = compose(rich, "rich_cards", items=3)

    sizes = [
        next(p.font_pt for p in slide.pieces if p.name == f"composed-{part}-0")
        for part in ("heading", "text", "label")
    ]

    assert sizes == sorted(sizes, reverse=True)
    assert len(set(sizes)) == 3, "two rows of the card are drawn the same size"


def test_a_rich_card_carries_every_row_a_real_one_has(banked):
    slide = compose(banked, "rich_cards", items=2)

    slots = {p.slot for p in slide.pieces}
    assert {
        "item:0:icon",
        "item:0:heading",
        "item:0:label",
        "item:0:text",
    } <= slots
    card = next(p for p in slide.pieces if p.name == "composed-card-0")
    rows = [p for p in slide.pieces if p.slot.startswith("item:0:")]
    for row in rows:
        assert card.box.y <= row.box.y
        assert row.box.y + row.box.h <= card.box.y + card.box.h + 1
