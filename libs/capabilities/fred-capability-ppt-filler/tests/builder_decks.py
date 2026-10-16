# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""In-code synthetic templates for the builder tests (no checked-in binaries).

Each builder returns ``.pptx`` bytes. ``build_showcase_template`` assembles the
example-slide cases the analysis rules are written for: an agenda of numbered
groups, KPI cards, a team roster, a process, sample-text boxes, shared
footers, an icon bank (a library slide), a chart, a table, declared shape names, an ``intent:``
note and a purely decorative divider.
"""

from __future__ import annotations

import io
from typing import Any, List, Optional, cast

from deck_builders import png_bytes
from fred_capability_ppt_filler.builder.models import Box, Pattern, RepeatGroup, Zone
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# python-pptx default template layout indices.
(
    L_TITLE,
    L_CONTENT,
    L_SECTION,
    L_TWO,
    L_COMPARISON,
    L_TITLE_ONLY,
    L_BLANK,
    L_CAPTION,
    L_PICTURE,
) = range(9)

FOOTER = "Confidential - internal use only"
FONT_HEADING = Pt(20)
FONT_BODY = Pt(12)
FONT_NUMBER = Pt(48)


def _save(presentation) -> bytes:
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _textbox(
    shapes,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size=None,
    name: Optional[str] = None,
):
    box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.text = text
    if size is not None:
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = size
    if name:
        box.name = name
    return box


def _footer(slide) -> None:
    _textbox(slide.shapes, 0.5, 7.0, 6.0, 0.3, FOOTER, Pt(8))


def _set_title(slide, title: str) -> None:
    cast(Any, slide.shapes.title).text = title


def _titled(presentation, title: str):
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
    _set_title(slide, title)
    return slide


def add_agenda_slide(presentation, count: int = 4):
    """Title + N groups of (number, heading) laid out on a grid."""
    slide = _titled(presentation, "Agenda")
    for index in range(count):
        col, row = index % 2, index // 2
        group = slide.shapes.add_group_shape()
        x, y = 0.6 + col * 4.6, 1.8 + row * 2.4
        _textbox(group.shapes, x, y, 1.0, 0.6, f"0{index + 1}", FONT_NUMBER)
        _textbox(
            group.shapes,
            x,
            y + 0.8,
            4.0,
            0.9,
            "Titre du chapitre sur une a deux lignes",
            FONT_HEADING,
        )
    _footer(slide)
    return slide


def add_kpi_slide(presentation, count: int = 4):
    """Title + N groups of (value, label); the icon of each card sits over the
    card but outside its group, as designers do."""
    slide = _titled(presentation, "Key figures")
    for index in range(count):
        group = slide.shapes.add_group_shape()
        x = 0.5 + index * 2.3
        _textbox(group.shapes, x, 1.6, 1.6, 0.2, "", None)
        _textbox(group.shapes, x, 2.6, 1.6, 1.0, f"{60 + index * 10}%", FONT_NUMBER)
        _textbox(group.shapes, x, 3.7, 1.6, 0.35, f"KPI #{index + 1}", FONT_BODY)
        slide.shapes.add_picture(
            io.BytesIO(png_bytes(12, 12)),
            Inches(x + 0.2),
            Inches(1.9),
            Inches(0.4),
            Inches(0.4),
        )
    _footer(slide)
    return slide


def add_steps_slide(presentation, count: int = 3):
    slide = _titled(presentation, "Process")
    for index in range(count):
        group = slide.shapes.add_group_shape()
        x = 0.5 + index * 3.1
        _textbox(group.shapes, x, 1.6, 0.8, 1.2, str(index + 1), FONT_NUMBER)
        _textbox(group.shapes, x, 3.0, 2.8, 0.5, f"Etape {index + 1}", FONT_HEADING)
        _textbox(group.shapes, x, 3.6, 2.8, 2.5, "Changez ce texte", FONT_BODY)
    return slide


def add_team_slide(presentation, count: int = 6):
    slide = _titled(presentation, "Team")
    for index in range(count):
        col, row = index % 3, index // 3
        group = slide.shapes.add_group_shape()
        x, y = 0.5 + col * 3.1, 1.8 + row * 2.6
        _textbox(group.shapes, x, y, 2.0, 0.3, "Poste en une ligne", Pt(9))
        _textbox(group.shapes, x, y + 0.35, 2.0, 0.3, "Nom", Pt(9))
    return slide


def add_free_text_slide(presentation):
    """Free boxes: one sample text (zone), one real content label (not a zone)."""
    slide = _titled(presentation, "Free text")
    _textbox(
        slide.shapes, 0.5, 1.5, 8.0, 4.0, "Cliquez ici pour ajouter du texte", Pt(18)
    )
    _textbox(slide.shapes, 0.5, 6.2, 3.0, 0.4, "Source: annual report 2025", Pt(10))
    _footer(slide)
    return slide


def add_declared_slide(presentation):
    """Shape names from the Selection pane declare roles; a `keep` name hides a box."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    _textbox(slide.shapes, 0.5, 0.5, 9.0, 1.0, "Quarterly review", Pt(28), name="title")
    _textbox(
        slide.shapes,
        0.5,
        1.8,
        9.0,
        3.5,
        "Some real looking paragraph of content.",
        Pt(14),
        name="body_main",
    )
    _textbox(slide.shapes, 0.5, 6.0, 4.0, 0.5, "Do not touch me", Pt(10), name="keep")
    slide.notes_slide.notes_text_frame.text = "intent: statement"
    return slide


def add_icon_bank_slide(presentation, pictures: int = 24):
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    _textbox(slide.shapes, 0.3, 3.0, 2.0, 0.5, "1/3", Pt(24))
    for index in range(pictures):
        col, row = index % 8, index // 8
        slide.shapes.add_picture(
            io.BytesIO(png_bytes(8, 8)),
            Inches(2.5 + col * 0.9),
            Inches(1.0 + row * 1.2),
            Inches(0.5),
            Inches(0.5),
        )
    return slide


def add_chart_slide(presentation):
    """Five large charts share the slide, more than one slide fills (four at
    most): it stays excluded."""
    slide = _titled(presentation, "Chart")
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("S", (1, 2))
    for left in (0.2, 1.6, 3.0, 4.4, 5.8):
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(left),
            Inches(1.5),
            Inches(2.5),
            Inches(5),
            data,
        )
    return slide


# Frames of the owner's example slides, in inches on a 10 × 7.5 slide: two
# charts of 24 %, three of 15 % and a 2 × 2 grid of 9.3 %.
CHART_GRIDS = {
    2: [(1.0, 1.4, 3.7, 4.8), (5.3, 1.4, 3.7, 4.8)],
    3: [(0.5, 1.5, 2.7, 4.2), (3.65, 1.5, 2.7, 4.2), (6.8, 1.5, 2.7, 4.2)],
    4: [
        (1.7, 1.25, 3.03, 2.31),
        (5.2, 1.25, 3.03, 2.31),
        (1.7, 3.95, 3.03, 2.31),
        (5.2, 3.95, 3.03, 2.31),
    ],
}


def add_charts_slide(presentation, count: int = 2, types=None, headings: bool = False):
    """An example slide of ``count`` charts (``CHART_GRIDS``), shuffled in the
    drawing order; ``headings`` puts a free text box right above each chart."""
    slide = _titled(presentation, f"Graphiques ×{count}")
    kinds = types or [XL_CHART_TYPE.COLUMN_CLUSTERED] * count
    frames = list(zip(CHART_GRIDS[count], kinds))
    for (x, y, w, h), kind in frames[::-1]:
        data = CategoryChartData()
        data.categories = ["A", "B", "C"]
        data.add_series("Exemple", (3, 2, 1))
        slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
        if headings:
            _textbox(slide.shapes, x, y - 0.5, w, 0.4, "Titre du graphique", FONT_BODY)
    return slide


def add_main_chart_slide(presentation, chart_type=XL_CHART_TYPE.PIE):
    """One free chart over the slide's main area, under a title and a takeaway."""
    slide = _titled(presentation, "Parts de marché")
    data: Any = CategoryChartData()
    if chart_type == XL_CHART_TYPE.XY_SCATTER:
        data = XyChartData()
        points = data.add_series("Parts")
        for x, y in ((1, 45), (2, 35), (3, 20)):
            points.add_data_point(x, y)
    else:
        data.categories = ["Europe", "Amériques", "Asie"]
        data.add_series("Parts", (45, 35, 20))
    slide.shapes.add_chart(
        chart_type, Inches(0.5), Inches(1.5), Inches(6), Inches(5), data
    )
    _textbox(slide.shapes, 6.8, 2.0, 2.8, 3.0, "Changez ce texte", FONT_BODY)
    return slide


def add_table_slide(presentation):
    """A grid with a merged cell: a drawing, not data, so it stays excluded."""
    slide = _titled(presentation, "Table")
    frame = slide.shapes.add_table(3, 3, Inches(1), Inches(1.5), Inches(8), Inches(4))
    table = cast(Any, frame).table
    table.cell(0, 0).merge(table.cell(0, 1))
    return slide


_A16 = "http://schemas.microsoft.com/office/drawing/2014/main"


def add_data_table_slide(presentation, rows: int = 12, cols: int = 5):
    """The Paintress table slide: an empty title over a header-row table of
    empty cells, each row carrying its ``a16:rowId``."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
    frame = slide.shapes.add_table(
        rows, cols, Inches(0.5), Inches(1.1), Inches(9), Emu(383378 * rows)
    )
    table = cast(Any, frame).table
    for number, row in enumerate(table.rows):
        row.height = Emu(383378)
        tr = row._tr
        ext_list = tr.makeelement(qn("a:extLst"), {})
        ext = ext_list.makeelement(
            qn("a:ext"), {"uri": "{0D108BD9-81ED-4DB2-BD59-A6C34878D82A}"}
        )
        ext.append(ext.makeelement(f"{{{_A16}}}rowId", {"val": str(1000 + number)}))
        ext_list.append(ext)
        tr.append(ext_list)
    return slide


def add_gantt_slide(presentation):
    """A month grid whose bars are cell fills: a drawing, never data."""
    slide = _titled(presentation, "Timeline 1 an")
    frame = slide.shapes.add_table(4, 6, Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    table = cast(Any, frame).table
    for column, month in enumerate(("", "JAN", "FEV", "MAR", "AVR", "MAI")):
        table.cell(0, column).text = month
    for row in range(1, 4):
        table.cell(row, 0).text = f"Titre {row}"
        for column in range(row, row + 2):
            table.cell(row, column).fill.solid()
    return slide


def add_divider_slide(presentation):
    """Title-only instance of a title+body layout, over a large graphic."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_CONTENT])
    _set_title(slide, "Motif")
    body = cast(Any, slide.placeholders[1])
    body._element.getparent().remove(body._element)
    for top in (4.0, 5.2):
        slide.shapes.add_picture(
            io.BytesIO(png_bytes(40, 4)),
            Inches(0),
            Inches(top),
            Inches(10),
            Inches(0.8),
        )
    return slide


def add_pure_instance_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_CONTENT])
    _set_title(slide, "Cliquez ici pour modifier le titre")
    cast(
        Any, slide.placeholders[1]
    ).text_frame.text = "Cliquez ici pour ajouter du texte"
    return slide


def add_big_picture_slide(presentation):
    slide = _titled(presentation, "Photo")
    slide.shapes.add_picture(
        io.BytesIO(png_bytes(40, 40)), Inches(0.5), Inches(1.5), Inches(9), Inches(5)
    )
    return slide


def add_scaled_kpi_slide(presentation):
    """KPI cards whose groups were resized after authoring (ext = chExt / 2),
    the first label box narrower than its siblings."""
    slide = _titled(presentation, "Scaled figures")
    for index in range(3):
        group = slide.shapes.add_group_shape()
        x = 0.5 + index * 3.0
        _textbox(group.shapes, x, 2.0, 1.6, 1.0, f"{10 + index}%", FONT_NUMBER)
        _textbox(
            group.shapes,
            x,
            3.2,
            0.7 if index == 0 else 1.6,
            0.35,
            f"KPI #{index + 1}",
            FONT_BODY,
        )
        # Shrink the group's placement without touching its child space.
        group.width = int(group.width / 2)
        group.height = int(group.height / 2)
    return slide


def add_notes_without_body_slide(presentation):
    """A notes slide whose body placeholder was deleted (no text frame)."""
    slide = _titled(presentation, "Silent notes")
    _textbox(
        slide.shapes, 0.5, 1.5, 8.0, 3.0, "Cliquez ici pour ajouter du texte", Pt(18)
    )
    notes = slide.notes_slide
    for shape in list(notes.placeholders):
        if (
            shape.placeholder_format.type is not None
            and shape.placeholder_format.type.name == "BODY"
        ):
            shape._element.getparent().remove(shape._element)
    return slide


def add_two_group_columns_slide(presentation):
    """Two grouped columns (heading + body): two siblings, not a list."""
    slide = _titled(presentation, "Two columns")
    for index in range(2):
        group = slide.shapes.add_group_shape()
        x = 0.5 + index * 4.6
        _textbox(group.shapes, x, 1.6, 4.0, 0.5, "Changez ce titre", FONT_HEADING)
        _textbox(group.shapes, x, 2.3, 4.0, 3.5, "Changez ce texte", FONT_BODY)
    return slide


def add_loose_figures_slide(presentation, count: int = 3, labels: bool = True):
    """N loose (ungrouped) figure boxes in a row, each with a label box under
    it: the shape of a placeholder-based "key figures" layout."""
    slide = _titled(presentation, "Loose figures")
    for index in range(count):
        x = 0.5 + index * 2.3
        _textbox(slide.shapes, x, 2.0, 1.6, 1.0, f"{20 + index}%", FONT_NUMBER)
        if labels:
            _textbox(slide.shapes, x, 3.1, 1.6, 0.4, f"Indicateur {index}", FONT_BODY)
    return slide


def add_three_column_bodies_slide(presentation):
    """Three tall body boxes side by side: columns, not a list of items."""
    slide = _titled(presentation, "Three columns")
    for index in range(3):
        _textbox(
            slide.shapes,
            0.5 + index * 3.1,
            1.6,
            2.9,
            4.8,
            "Une colonne de texte assez longue pour un paragraphe entier",
            FONT_BODY,
        )
    return slide


def _glyph(shapes, x: float, y: float, size: float = 0.4):
    """A small freeform icon (custom geometry), as Slidesgo draws its icons."""
    builder = shapes.build_freeform(Inches(x), Inches(y + size), scale=1.0)
    builder.add_line_segments(
        [
            (Inches(x + size / 2), Inches(y)),
            (Inches(x + size), Inches(y + size)),
        ]
    )
    return builder.convert_to_shape()


def add_icon_sheet_slide(presentation, title: str = "Nature icons", icons: int = 30):
    """A resource sheet: grouped freeform icons under one title."""
    slide = _titled(presentation, title)
    for index in range(icons):
        col, row = index % 10, index // 10
        group = slide.shapes.add_group_shape()
        _glyph(group.shapes, 0.5 + col * 0.9, 2.0 + row * 1.2)
        _glyph(group.shapes, 0.6 + col * 0.9, 2.1 + row * 1.2, 0.2)
    return slide


def add_fonts_sheet_slide(presentation):
    """A font and palette sheet: font names, then hex swatches."""
    slide = _titled(presentation, "Fonts")
    _textbox(slide.shapes, 0.5, 1.6, 4.0, 0.5, "Titles: Montserrat", FONT_BODY)
    _textbox(slide.shapes, 0.5, 2.2, 4.0, 0.5, "Body: Open Sans", FONT_BODY)
    for index, colour in enumerate(("#083c77", "#ebfcff", "#93c4e0", "#ffffff")):
        swatch = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.5 + index * 1.5), Inches(4.0), Inches(1), Inches(1)
        )
        swatch.text_frame.text = colour
    return slide


def add_icon_cards_slide(presentation):
    """A design: four icons each over its own text, never a library."""
    slide = _titled(presentation, "Our strengths")
    for index in range(4):
        _glyph(slide.shapes, 0.8 + index * 2.3, 1.8)
        _textbox(
            slide.shapes,
            0.5 + index * 2.3,
            2.6,
            2.0,
            1.2,
            "Une force de l'entreprise en une phrase",
            FONT_BODY,
        )
    return slide


def add_dotted_timeline_slide(presentation):
    """A design: a row of 24 small preset dots under dated labels."""
    slide = _titled(presentation, "Milestones")
    for index in range(24):
        slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(0.5 + index * 0.38),
            Inches(3.5),
            Inches(0.15),
            Inches(0.15),
        )
    for index, year in enumerate(("2021", "2022", "2023")):
        _textbox(slide.shapes, 0.5 + index * 3.0, 2.8, 1.5, 0.5, year, FONT_BODY)
        _textbox(
            slide.shapes,
            0.5 + index * 3.0,
            4.0,
            2.5,
            1.0,
            "Un jalon atteint",
            FONT_BODY,
        )
    return slide


def add_attribution_closing_slide(presentation):
    """A closing slide asking to keep an attribution: still a closing slide."""
    slide = _titled(presentation, "Thank you")
    _textbox(slide.shapes, 0.5, 5.5, 6.0, 0.5, "Please keep this slide for attribution")
    return slide


# --- Google Slides exports: every text box a TITLE or SUBTITLE placeholder -------------


def _placeholder(slide, kind: str, idx: int, x, y, w, h, text: str, size) -> None:
    """A ``kind`` (title/subTitle) placeholder at an explicit position, as
    Slidesgo templates type their card headings and texts."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = size
    nv_pr = shape._element.find(qn("p:nvSpPr")).find(qn("p:nvPr"))
    ph = nv_pr.makeelement(qn("p:ph"), {"type": kind, "idx": str(idx)})
    nv_pr.insert(0, ph)
    c_nv = shape._element.find(qn("p:nvSpPr")).find(qn("p:cNvSpPr"))
    c_nv.attrib.pop("txBox", None)


def _blank_titled(presentation, title: str):
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    _placeholder(slide, "title", 90, 0.3, 0.3, 9.4, 0.9, title, Pt(28))
    return slide


def add_placeholder_cards_slide(presentation):
    """Three (TITLE heading, SUBTITLE text) pairs, heading left of its text."""
    slide = _blank_titled(presentation, "Three ideas")
    for index, (x, y) in enumerate(((0.3, 1.8), (5.2, 1.8), (5.2, 4.6))):
        _placeholder(
            slide, "title", 10 + index, x, y, 1.7, 1.7, "Pourquoi Venus", Pt(16)
        )
        _placeholder(
            slide,
            "subTitle",
            20 + index,
            x + 1.8,
            y,
            2.5,
            1.7,
            "Venus porte un beau nom et reste la deuxieme planete",
            Pt(12),
        )
    return slide


def add_placeholder_two_ideas_slide(presentation):
    """Two (heading, text) pairs stacked: the two columns of a two_columns slide."""
    slide = _blank_titled(presentation, "Two ideas")
    for index, y in enumerate((2.0, 4.6)):
        _placeholder(
            slide, "subTitle", 10 + index, 0.3, y, 1.6, 1.4, "Sur Saturne", Pt(16)
        )
        _placeholder(
            slide,
            "subTitle",
            20 + index,
            2.0,
            y,
            3.8,
            2.0,
            "Saturne est une geante gazeuse entouree de plusieurs anneaux",
            Pt(12),
        )
    return slide


def add_big_figures_slide(presentation):
    """Two big figures, the first captioned above, the second below."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    _placeholder(
        slide, "subTitle", 1, 0.7, 2.7, 4.2, 1.1, "Masse du Soleil rapportee", Pt(12)
    )
    _placeholder(slide, "title", 2, 0.7, 4.4, 4.2, 1.3, "333,000", Pt(44))
    _placeholder(slide, "title", 3, 5.1, 2.7, 4.2, 1.3, "386,000 km", Pt(44))
    _placeholder(
        slide, "subTitle", 4, 5.1, 4.6, 4.2, 1.1, "Distance Terre Lune", Pt(12)
    )
    return slide


def add_big_number_slide(presentation):
    """A lone stand-in figure and its caption: a "big number" design."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    _placeholder(slide, "title", 1, 1.4, 2.8, 7.2, 1.2, "xx%", Pt(60))
    _placeholder(slide, "subTitle", 2, 1.4, 4.1, 7.2, 0.6, "", Pt(16))
    return slide


def add_numbered_grid_slide(presentation):
    """Six numbered texts in two columns: cards, not a process."""
    slide = _blank_titled(presentation, "Six ideas")
    for index in range(6):
        x, y = 0.3 + (index // 3) * 4.7, 2.4 + (index % 3) * 1.3
        _placeholder(
            slide, "title", 10 + index, x, y, 0.6, 1.0, f"0{index + 1}", Pt(20)
        )
        _placeholder(
            slide,
            "subTitle",
            20 + index,
            x + 0.5,
            y + 0.15,
            3.3,
            1.2,
            "Mercure est la plus petite planete du systeme solaire",
            Pt(12),
        )
    return slide


TIMELINE_ENTRY = "Saturne est une geante gazeuse avec des anneaux"


def add_staircase_timeline_slide(presentation):
    """Five "20XX" dates each over its text, climbing left to right."""
    slide = _titled(presentation, "Our history")
    for index in range(5):
        x, y = 0.4 + index * 1.8, 5.7 - index * 1.0
        _textbox(slide.shapes, x, y, 1.6, 0.8, "20XX", Pt(16))
        text = TIMELINE_ENTRY if index == 1 else f"Etape numero {index} du parcours"
        _textbox(slide.shapes, x, y + 0.75, 1.6, 0.7, text, Pt(12))
    return slide


def add_reused_entry_slide(presentation):
    """Another slide repeating a timeline entry's text."""
    slide = _titled(presentation, "Elsewhere")
    _textbox(slide.shapes, 0.5, 5.0, 4.0, 0.7, TIMELINE_ENTRY, Pt(12))
    return slide


def add_paragraph_subtitle_slide(presentation):
    """A SUBTITLE placeholder sized for paragraphs: a text slide, not a cover."""
    slide = _blank_titled(presentation, "Let's have a look")
    _placeholder(
        slide,
        "subTitle",
        1,
        0.3,
        2.5,
        6.9,
        4.6,
        "Mercure est la planete la plus proche du Soleil et la plus petite",
        Pt(14),
    )
    return slide


def add_graphic_sheet_slide(presentation):
    """Editable graphics: grouped shapes, no text, no placeholder."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    for index in range(10):
        group = slide.shapes.add_group_shape()
        x, y = 0.5 + (index % 5) * 1.8, 1.0 + (index // 5) * 3.0
        group.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(1.5), Inches(1.5)
        )
        group.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 1.6), Inches(1.5), Inches(0.5)
        )
    return slide


def add_promo_slide(presentation):
    """A linked banner picture alone: a promotion, not a design."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    picture = slide.shapes.add_picture(
        io.BytesIO(png_bytes(64, 16)),
        Inches(3.8),
        Inches(3.5),
        Inches(2.4),
        Inches(0.5),
    )
    picture.click_action.hyperlink.address = "http://bit.ly/example"
    return slide


def add_empty_slide(presentation):
    presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])


def build_library_template() -> bytes:
    return build_slides(
        [
            add_icon_sheet_slide,
            add_fonts_sheet_slide,
            add_icon_cards_slide,
            add_dotted_timeline_slide,
            add_attribution_closing_slide,
        ]
    )


def build_showcase_template() -> bytes:
    presentation = Presentation()
    add_agenda_slide(presentation)
    add_kpi_slide(presentation)
    add_steps_slide(presentation)
    add_team_slide(presentation)
    add_free_text_slide(presentation)
    add_declared_slide(presentation)
    add_icon_bank_slide(presentation)
    add_chart_slide(presentation)
    add_table_slide(presentation)
    add_divider_slide(presentation)
    add_pure_instance_slide(presentation)
    add_big_picture_slide(presentation)
    add_scaled_kpi_slide(presentation)
    add_notes_without_body_slide(presentation)
    add_two_group_columns_slide(presentation)
    return _save(presentation)


def add_cover_photo_slide(presentation):
    """A full-bleed photo with the title written over it: the picture is drawn
    first, so replacing it must not bring it in front of the text."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
    slide.shapes.add_picture(
        io.BytesIO(png_bytes(64, 36)),
        Inches(0),
        Inches(0),
        Inches(10),
        Inches(7.5),
    )
    _textbox(slide.shapes, 0.8, 2.5, 6.0, 1.2, "Changez ce titre", Pt(36))
    return slide


def build_layouts_only_template() -> bytes:
    return _save(Presentation())


def build_spaced_template(
    space_before_pt: int = 24, no_wrap_title: bool = True
) -> bytes:
    """A deck whose layout declares body line spacing and space before each
    paragraph — the shape of a corporate master — plus a figure that does not
    wrap. Both are read off the layout, never off the slide's own text.
    """
    presentation = Presentation()
    layout = presentation.slide_layouts[L_CONTENT]
    for placeholder in layout.placeholders:
        if placeholder.placeholder_format.idx == 0:
            body = placeholder.text_frame._txBody.find(qn("a:bodyPr"))
            if no_wrap_title:
                # Sized to its text, the way a designer sets a figure.
                body.set("wrap", "none")
            else:
                # Wraps inside its width and grows downwards - a cover title.
                body.set("wrap", "square")
                body.append(body.makeelement(qn("a:spAutoFit"), {}))
            continue
        body = placeholder.text_frame._txBody
        lst = body.makeelement(qn("a:lstStyle"), {})
        lvl = body.makeelement(qn("a:lvl1pPr"), {})
        spacing = body.makeelement(qn("a:lnSpc"), {})
        spacing.append(body.makeelement(qn("a:spcPct"), {"val": "114000"}))
        before = body.makeelement(qn("a:spcBef"), {})
        before.append(
            body.makeelement(qn("a:spcPts"), {"val": str(space_before_pt * 100)})
        )
        lvl.append(spacing)
        lvl.append(before)
        lst.append(lvl)
        body.insert(1, lst)

    slide = presentation.slides.add_slide(layout)
    _set_title(slide, "Changez ce titre")
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx != 0:
            cast(Any, placeholder).text_frame.text = "Changez ce texte"
    _footer(slide)
    return _save(presentation)


def build_slides(builders: List) -> bytes:
    presentation = Presentation()
    for builder in builders:
        builder(presentation)
    return _save(presentation)


# --- variety: the plan of a real first deck -------------------------------------------------

# Five `bullets` slides (3, 5, 7, 9, 10) whose lines suit other kinds: labelled
# points (5), dated events (7), two then three 'Label :' groups (9, 10).
EVIDENCE_SLIDES: List[dict] = [
    {"title": "Thales en Espagne et en Italie", "subtitle": "Activités"},
    {"title": "Sommaire", "items": [{"heading": "Espagne"}, {"heading": "Italie"}]},
    {
        "title": "Introduction",
        "kind": "bullets",
        "bullets": [
            "Thales : acteur mondial de la défense",
            "Présence historique en Europe du Sud : 1980",
            "Objectif : explorer les activités locales",
        ],
    },
    {
        "title": "Espagne",
        "columns": [{"bullets": ["Défense"]}, {"bullets": ["Spatial"]}],
    },
    {
        "title": "Espagne : enjeux",
        "kind": "bullets",
        "bullets": [
            "Direction : IA souveraine et spatial",
            "Objectif : hub exportateur",
            "Enjeux majeurs :",
            "- Chaînes de valeur de défense",
            "- Décarbonation des transports",
        ],
    },
    {"title": "Italie", "columns": [{"bullets": ["Cyber"]}, {"bullets": ["Défense"]}]},
    {
        "title": "Italie : réalisations",
        "kind": "bullets",
        "bullets": [
            "Événements clés récents :",
            "- Janvier 2025 : accord national cyber",
            "- Juin 2025 : radars pour la Marine",
            "- Mars 2026 : pôle spatial à Turin",
            "Perspectives :",
            "- Croissance portée par le PNRR",
        ],
    },
    {"title": "Chiffres", "items": [{"number": "2 800", "label": "experts"}] * 2},
    {
        "title": "Synthèse",
        "kind": "bullets",
        "bullets": [
            "Espagne :",
            "- Hub exportateur",
            "- IA souveraine",
            "Italie :",
            "- Moteur de croissance",
            "- Cyber et défense",
        ],
    },
    {
        "title": "Recommandations",
        "kind": "bullets",
        "bullets": [
            "Pour l'Espagne :",
            "- Attirer les talents cyber",
            "Pour l'Italie :",
            "- Co-investissements",
            "Recommandation globale :",
            "- Soutien actif du COMEX",
        ],
    },
    {"title": "Conclusion", "kind": "section"},
]


def with_item_slides(
    inventory: Any, kinds: Any = ("timeline", "cards"), count: int = 6
):
    """``inventory`` plus one example slide per kind holding ``count`` heading +
    text items: what the variety advice may name. Never rendered."""

    def zone(key: str, role: Any, x: int) -> Zone:
        box = Box(x=x, y=2_000_000, w=1_000_000, h=2_000_000)
        return Zone(
            key=key, role=role, box=box, font_pt=14.0, chars_per_line=20, max_lines=4
        )

    extra = []
    for n, kind in enumerate(kinds):
        pid = f"slide:{900 + n}"
        items = [
            [
                zone(f"{pid}/{i}/h", "heading", i * 1_200_000),
                zone(f"{pid}/{i}/b", "body", i * 1_200_000),
            ]
            for i in range(count)
        ]
        group = RepeatGroup(
            key=f"{pid}/g",
            count=count,
            item_zones=["heading", "body"],
            items=items,
            item_keys=[f"{pid}/{i}" for i in range(count)],
        )
        title = Zone(
            key="ph:0",
            role="title",
            box=Box(x=0, y=0, w=9_000_000, h=900_000),
            font_pt=32.0,
            chars_per_line=40,
            max_lines=1,
        )
        extra.append(
            Pattern(
                id=pid,
                source="slide",
                source_index=900 + n,
                name=kind,
                label=f"Synthetic {kind}",
                kind=kind,
                zones=[title],
                groups=[group],
            )  # type: ignore[arg-type]
        )
    return inventory.model_copy(update={"patterns": inventory.patterns + extra})


__all__ = [name for name in globals() if name.startswith(("add_", "build_"))] + [
    "FOOTER",
    "Emu",
    "EVIDENCE_SLIDES",
    "with_item_slides",
]
