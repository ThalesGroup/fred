# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Table intents: the model says which row is a total and which row or
column matters; the table's own style shows it, and what the style cannot
show is drawn in bold and told."""

from __future__ import annotations

import io
from typing import Any, Callable, List, Optional, Tuple

import pytest
from fred_capability_ppt_filler.builder import tools
from fred_capability_ppt_filler.builder.assign import assign
from fred_capability_ppt_filler.builder.builtin import builtin_template
from fred_capability_ppt_filler.builder.changes import SlideChange, apply_changes
from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanTable
from fred_capability_ppt_filler.builder.render import RenderResult, render_deck
from fred_capability_ppt_filler.builder.table_intents import PREDEFINED
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

COVER = {"title": "Choix du CRM", "subtitle": "Comité"}
OFFERS = {
    "header": ["Offre", "Prix", "Délai"],
    "rows": [["Offre A", "10", "3"], ["Offre B", "12", "2"], ["Total", "22", "5"]],
}
ACCENT_2 = "21E4AEA4-8DFA-4A89-87EB-49C32662AFE0"
CUSTOM = "11111111-2222-3333-4444-555555555555"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


@pytest.fixture(scope="module")
def builtin():
    return builtin_template()


def _restyled(content: bytes, edit: Callable[[Any, Any], None]) -> bytes:
    """The built-in template with its table's properties (and the file's table
    styles) changed by ``edit(tblPr, presentation)``."""
    presentation = Presentation(io.BytesIO(content))
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                edit(shape._element.find(f".//{qn('a:tblPr')}"), presentation)
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def _style_id(guid: str) -> Callable[[Any, Any], None]:
    def edit(tblPr: Any, _: Any) -> None:
        tblPr.find(qn("a:tableStyleId")).text = "{" + guid + "}"

    return edit


def _custom_style(fill: str, last_row: bool) -> Callable[[Any, Any], None]:
    """A style the file defines, its whole table filled with ``fill``."""

    def edit(tblPr: Any, presentation: Any) -> None:
        tblPr.find(qn("a:tableStyleId")).text = "{" + CUSTOM + "}"
        part = presentation.part.part_related_by(RT.TABLE_STYLES)
        last = "<a:lastRow><a:tcTxStyle b='on'/></a:lastRow>" if last_row else ""
        part._blob = (
            f"<a:tblStyleLst xmlns:a='{_A}' def='{{{CUSTOM}}}'>"
            f"<a:tblStyle styleId='{{{CUSTOM}}}' styleName='House'>"
            f"<a:wholeTbl><a:tcStyle><a:fill><a:solidFill><a:schemeClr val='{fill}'/>"
            "</a:solidFill></a:fill></a:tcStyle></a:wholeTbl>"
            "<a:firstRow><a:tcStyle><a:fill><a:solidFill><a:schemeClr val='accent3'/>"
            f"</a:solidFill></a:fill></a:tcStyle></a:firstRow>{last}"
            "</a:tblStyle></a:tblStyleLst>"
        ).encode()

    return edit


def _plan(table: Any) -> DeckPlan:
    return DeckPlan.model_validate(
        {"slides": [COVER, {"title": "Comparatif", "table": table}]}
    )


def _render(
    builtin, table: Any, content: Optional[bytes] = None
) -> Tuple[RenderResult, Any]:
    template, inventory = builtin
    plan = _plan(table)
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems, [p.as_dict() for p in assignment.problems]
    deck, patterns, numbers = assignment.deck(plan)
    result = render_deck(content or template, inventory, deck, patterns, None, numbers)
    slide = Presentation(io.BytesIO(result.pptx)).slides[1]
    frame = next(s for s in slide.shapes if getattr(s, "has_table", False))
    return result, frame._element.find(f".//{qn('a:tbl')}")


def _rows(tbl: Any) -> List[Any]:
    return tbl.findall(qn("a:tr"))


def _bold(tc: Any) -> bool:
    runs = list(tc.iter(qn("a:r")))
    return bool(runs) and all(r.find(qn("a:rPr")).get("b") == "1" for r in runs)


def _tint(tc: Any) -> Optional[str]:
    """The accent of a cell's own fill, when it is the light tint."""
    colour = tc.find(f"{qn('a:tcPr')}/{qn('a:solidFill')}/{qn('a:schemeClr')}")
    if colour is None:
        return None
    mods = {child.tag.split("}")[1]: child.get("val") for child in colour}
    assert mods == {"lumMod": "20000", "lumOff": "80000"}
    return colour.get("val")


def _notes(result: RenderResult) -> List[Tuple[str, Optional[str]]]:
    return [(n.code, n.field) for n in result.report.built[1].notes if n.field]


# --- plan ------------------------------------------------------------------------------


def test_table_intents_are_read_from_loose_spellings():
    table = PlanTable.model_validate(
        {
            **OFFERS,
            "total": "yes",
            "highlighted_row": ["Offre B"],
            "highlight_column": 1.0,
        }
    )
    assert (table.total_row, table.highlight_row, table.highlight_column) == (
        True,
        "Offre B",
        1,
    )
    assert PlanTable.model_validate({**OFFERS, "total_row": "dernière"}).total_row
    assert not PlanTable.model_validate({**OFFERS, "total_row": "no"}).total_row
    again = PlanTable.model_validate(table.model_dump())
    assert again == table


def test_a_table_without_intents_is_drawn_as_before(builtin):
    _, tbl = _render(builtin, OFFERS)
    assert tbl.find(qn("a:tblPr")).get("lastRow") is None
    assert not any(_tint(tc) for tc in tbl.iter(qn("a:tc")))


# --- total row ----------------------------------------------------------------------------


def test_a_total_row_takes_the_predefined_style_s_own_formatting(builtin):
    result, tbl = _render(builtin, {**OFFERS, "total_row": True})
    assert tbl.find(qn("a:tblPr")).get("lastRow") == "1"
    last = _rows(tbl)[-1]
    assert [tc.xpath("string(.)") for tc in last.findall(qn("a:tc"))] == [
        "Total",
        "22",
        "5",
    ]
    assert not _bold(last.findall(qn("a:tc"))[0])
    assert _notes(result) == []


def test_a_template_total_row_is_cleared_when_the_plan_has_none(builtin):
    template, _ = builtin

    def flagged(tblPr: Any, _: Any) -> None:
        tblPr.set("lastRow", "1")

    _, tbl = _render(builtin, OFFERS, _restyled(template, flagged))
    assert tbl.find(qn("a:tblPr")).get("lastRow") is None


def test_a_style_without_total_formatting_shows_the_total_in_bold(builtin):
    template, _ = builtin
    content = _restyled(template, _custom_style("lt1", last_row=False))
    result, tbl = _render(builtin, {**OFFERS, "total_row": True}, content)
    assert tbl.find(qn("a:tblPr")).get("lastRow") is None
    assert all(_bold(tc) for tc in _rows(tbl)[-1].findall(qn("a:tc")))
    assert _notes(result) == [("intent_shown_otherwise", "total_row")]
    assert any("total row is shown in bold" in n for n in result.report.notes)


def test_a_custom_style_with_total_formatting_takes_the_flag(builtin):
    template, _ = builtin
    content = _restyled(template, _custom_style("lt1", last_row=True))
    result, tbl = _render(builtin, {**OFFERS, "total_row": True}, content)
    assert tbl.find(qn("a:tblPr")).get("lastRow") == "1"
    assert _notes(result) == []


def test_a_total_of_one_row_is_told(builtin):
    result, tbl = _render(
        builtin, {"header": ["Offre", "Prix"], "rows": [["A", "1"]], "total_row": True}
    )
    assert tbl.find(qn("a:tblPr")).get("lastRow") is None
    assert _notes(result) == [("intent_not_shown", "total_row")]


# --- highlighted row and column -------------------------------------------------------------


def test_a_row_named_by_its_first_cell_is_bold_and_tinted_with_the_style_accent(
    builtin,
):
    result, tbl = _render(builtin, {**OFFERS, "highlight_row": "offre b"})
    header, first, second, total = _rows(tbl)
    assert all(
        _bold(tc) and _tint(tc) == "accent1" for tc in second.findall(qn("a:tc"))
    )
    assert not any(
        _tint(tc) for tc in first.findall(qn("a:tc")) + header.findall(qn("a:tc"))
    )
    assert "srgbClr" not in tbl.xml and _notes(result) == []


def test_the_tint_follows_the_style_accent(builtin):
    template, _ = builtin
    content = _restyled(template, _style_id(ACCENT_2))
    _, tbl = _render(builtin, {**OFFERS, "highlight_row": 0}, content)
    assert {_tint(tc) for tc in _rows(tbl)[1].findall(qn("a:tc"))} == {"accent2"}


def test_a_custom_style_tints_with_its_header_accent(builtin):
    template, _ = builtin
    content = _restyled(template, _custom_style("lt1", last_row=True))
    _, tbl = _render(builtin, {**OFFERS, "highlight_row": 1}, content)
    assert {_tint(tc) for tc in _rows(tbl)[2].findall(qn("a:tc"))} == {"accent3"}


def test_a_column_named_by_its_header_tints_its_body_cells(builtin):
    _, tbl = _render(builtin, {**OFFERS, "highlight_column": "Prix"})
    header, *body = _rows(tbl)
    assert _tint(header.findall(qn("a:tc"))[1]) is None
    assert all(_tint(tr.findall(qn("a:tc"))[1]) == "accent1" for tr in body)
    assert not any(_tint(tr.findall(qn("a:tc"))[0]) for tr in body)


def test_references_that_match_nothing_are_told_and_built_without(builtin):
    result, tbl = _render(
        builtin, {**OFFERS, "highlight_row": "Offre Z", "highlight_column": 7}
    )
    assert not any(_tint(tc) for tc in tbl.iter(qn("a:tc")))
    assert _notes(result) == [
        ("intent_not_shown", "highlight_row"),
        ("intent_not_shown", "highlight_column"),
    ]
    assert any(
        "table.highlight_row 'Offre Z' matches no row (first cells: 'Offre A', "
        "'Offre B', 'Total')" in n
        for n in result.report.notes
    )


def test_a_first_row_moved_into_the_template_header_shifts_the_highlight(builtin):
    rows = [["Offre", "Prix"], ["A", "1"], ["B", "2"]]
    result, tbl = _render(builtin, {"rows": rows, "highlight_row": 2})
    assert {_tint(tc) for tc in _rows(tbl)[2].findall(qn("a:tc"))} == {"accent1"}
    result, _ = _render(builtin, {"rows": rows, "highlight_row": 0})
    assert _notes(result) == [("intent_not_shown", "highlight_row")]
    assert any("shown as the table's header" in n for n in result.report.notes)


def test_a_dark_style_shows_the_highlight_in_bold_only(builtin):
    template, _ = builtin
    dark = next(g for g, s in PREDEFINED.items() if s.dark)
    content = _restyled(template, _style_id(dark))
    result, tbl = _render(builtin, {**OFFERS, "highlight_row": "Offre A"}, content)
    cells = _rows(tbl)[1].findall(qn("a:tc"))
    assert all(_bold(tc) and _tint(tc) is None for tc in cells)
    assert _notes(result) == [("intent_shown_otherwise", "highlight_row")]


def test_every_predefined_style_is_listed_once():
    assert len(PREDEFINED) == 74
    assert PREDEFINED["5C22544A-7EE6-4342-B048-85BDC9FD1C3A"].accent == 1
    assert PREDEFINED[ACCENT_2].accent == 2


# --- split tables, edits, the tool ---------------------------------------------------------


def test_a_split_table_shows_its_total_and_highlight_on_the_right_part(builtin):
    template, inventory = builtin
    rows = [[f"Ligne {n}", str(n)] for n in range(16)] + [["Total", "120"]]
    plan = _plan(
        {
            "header": ["Ligne", "Valeur"],
            "rows": rows,
            "total_row": True,
            "highlight_row": 1,
        }
    )
    assignment = assign(inventory, plan, "off")
    deck, patterns, numbers = assignment.deck(plan)
    first, second = (s.table for s in deck.slides[1:])
    assert first is not None and second is not None
    assert (first.total_row, first.highlight_row) == (False, 1)
    assert (second.total_row, second.highlight_row) == (True, None)
    render_deck(template, inventory, deck, patterns, None, numbers)


def test_table_intents_are_edited(builtin):
    _, inventory = builtin
    base = _plan(OFFERS)
    edits = [
        {"op": "edit", "slide": 2, "field": "table.total_row", "value": True},
        {"op": "edit", "slide": 2, "field": "table.highlight_row", "value": "Offre A"},
        {"op": "edit", "slide": 2, "field": "table.highlight_column", "value": 2},
    ]
    plan = apply_changes(base, [SlideChange.model_validate(e) for e in edits])
    table = plan.slides[1].table
    assert table is not None
    assert (table.total_row, table.highlight_row, table.highlight_column) == (
        True,
        "Offre A",
        2,
    )
    assert not assign(inventory, plan, "off").problems


def test_the_tool_offers_the_table_intents(builtin):
    _, inventory = builtin
    schema = tools._args_schema(inventory).model_json_schema()
    fields = schema["$defs"]["PlanTable"]["properties"]
    assert {"total_row", "highlight_row", "highlight_column"} <= set(fields)
    assert "total_row: true when the last row" in tools._description(
        inventory, "template_slide", 2
    )
