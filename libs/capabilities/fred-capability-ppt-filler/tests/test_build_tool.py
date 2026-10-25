# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The ``build_presentation`` and image-listing tools, offline on the port
fakes: schema and description, delivery, the needs-revision error and the
repair loop."""

from __future__ import annotations

import asyncio
import io
import itertools
import json
from pathlib import Path
from typing import Any, cast

import pytest
from builder_decks import EVIDENCE_SLIDES, build_showcase_template, with_item_slides
from deck_builders import png_bytes
from fred_capability_ppt_filler.builder import analysis_job
from fred_capability_ppt_filler.builder import tools as tools_mod
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.capability import (
    PPT_BUILDER_TEMPLATE_KEY,
    PptBuilderCapability,
    PptBuilderConfig,
    build_instructions,
)
from fred_capability_ppt_filler.builder.checks import SlideReport
from fred_capability_ppt_filler.builder.models import PatternOverride
from fred_capability_ppt_filler.builder.plan import DeckPlan, SlidePlan
from fred_capability_ppt_filler.builder.tools import (
    BUILD_TOOL_NAME,
    LIST_IMAGES_TOOL_NAME,
    _build_notes,
)
from fred_capability_ppt_filler.fill import BuiltSlide, PptPreviewPart
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
)
from fred_sdk.contracts.runtime import RuntimeServices
from langchain_core.messages import AIMessage, HumanMessage, ToolMessage
from port_fakes import FakeAssets, FakeDocs, FakeFolders, FakeWorkspace
from pptx import Presentation

TEMPLATE = build_showcase_template()
COVER = {"title": "Kick-off", "subtitle": "Team"}


@pytest.fixture
def preview_union():
    """Register ``PptPreviewPart`` in the runtime ``UiPart`` union, as the
    capability registry does at pod boot."""
    from fred_sdk.contracts.ui_part_union import rebuild_ui_part_union

    rebuild_ui_part_union([PptPreviewPart])
    yield
    rebuild_ui_part_union(())


async def _pdf(_bytes: bytes, *, suffix: str, **_options):
    return b"%PDF-fake"


async def _no_pdf(_bytes: bytes, *, suffix: str, **_options):
    return None


def _ctx(
    config: PptBuilderConfig, *, assets=None, docs=None, folders=None, workspace=None
):
    return CapabilityContext(
        identity=CapabilityIdentity(
            user_id="u-1",
            session_id="sess-9",
            agent_instance_id="00000000-0000-4000-8000-000000000007",
        ),
        config=config,
        turn_options=EmptyModel(),
        services=RuntimeServices(
            agent_assets=assets,
            document_content=docs,
            document_folders=folders,
            workspace_fs=workspace if workspace is not None else FakeWorkspace(),
        ),
    )


def _tools(ctx) -> dict[str, Any]:
    (middleware,) = PptBuilderCapability().middleware(ctx)
    return {tool.name: cast(Any, tool) for tool in cast(Any, middleware).tools}


def _configured(
    extend=None, variety: str = "template_slide", template: bytes = TEMPLATE
) -> tuple[PptBuilderConfig, FakeAssets]:
    """A template whose analysis succeeded and is in this pod's cache."""
    inventory = analyze_template(template).inventory
    assert inventory is not None
    if extend is not None:
        inventory = extend(inventory)
    sha = "sha-test" if template is TEMPLATE else "sha-other"
    config = PptBuilderConfig.model_validate(
        {"template_sha256": sha, "variety_check": variety}
    )
    analysis_job.remember_inventory(
        "00000000-0000-4000-8000-000000000007",
        sha,
        config.labelling_key(),
        inventory,
    )
    analysis_job._templates.set(sha, template)
    assets = FakeAssets({PPT_BUILDER_TEMPLATE_KEY: template})
    return config, assets


async def _build(
    monkeypatch,
    *,
    workspace=None,
    docs=None,
    extend=None,
    variety="template_slide",
    **kwargs,
):
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured(extend, variety)
    tools = _tools(_ctx(config, assets=assets, docs=docs, workspace=workspace))
    return await tools[BUILD_TOOL_NAME].coroutine(**kwargs)


# --- schema and description ---------------------------------------------------------------


def test_schema_and_description_offer_kinds_never_pattern_ids():
    config, assets = _configured()
    excluded = PatternOverride(
        pattern_id="layout:0", included=False, label="Hidden cover"
    )
    config = config.model_copy(update={"overrides": [excluded]})
    tools = _tools(_ctx(config, assets=assets))
    assert set(tools) == {BUILD_TOOL_NAME, LIST_IMAGES_TOOL_NAME}
    description = tools[BUILD_TOOL_NAME].description
    assert "slide:" not in description and "layout:" not in description
    assert "Hidden cover" not in description
    process = next(
        line for line in description.splitlines() if line.startswith("- process (")
    )
    assert ": title" in process  # variant label first, then its zones
    assert "items max 3 [number, heading, text" in description
    assert "≈chars per line × lines" in description
    # The only cover was excluded: no cover kind, and no cover rule announced.
    assert "- cover:" not in description and "is the cover" not in description
    properties = tools[BUILD_TOOL_NAME].args_schema.model_json_schema()["$defs"][
        "BuildSlide"
    ]["properties"]
    assert "pattern" not in properties
    # The takeaway first, then the content, `kind` last, listing only what this
    # template offers.
    assert list(properties)[:2] == ["takeaway", "title"]
    assert list(properties)[-1] == "kind"
    # The only enum anywhere in the field: the type's full list stays out.
    kind_schema = properties["kind"]
    assert "enum" not in kind_schema
    (kinds,) = [b["enum"] for b in kind_schema["anyOf"] if "enum" in b]
    assert "process" in kinds and "cover" not in kinds and "cards" not in kinds


def test_a_title_states_the_conclusion_and_the_takeaway_stays_in_the_notes():
    """'title (short)' after the takeaway gave topical titles ("Chronologie clé")."""
    inventory = analyze_template(TEMPLATE).inventory
    assert inventory is not None
    rules = build_instructions(False)
    for text in (tools_mod._description(inventory, "template_slide", 2), rules):
        assert "a claim, not a topic" in text
        assert "title (short)" not in text and "not the takeaway itself" not in text
    takeaway = SlidePlan.model_fields["takeaway"].description or ""
    assert "shown itself only when the user explicitly asks" in takeaway


def test_source_is_asked_only_for_the_user_s_documents_or_data():
    """A history deck from general knowledge was labelled 'Source : l'utilisateur'."""
    inventory = analyze_template(TEMPLATE).inventory
    assert inventory is not None
    source = SlidePlan.model_fields["source"].description or ""
    for text in (
        tools_mod._description(inventory, "template_slide", 2),
        build_instructions(False),
        source,
    ):
        assert "general knowledge" in text
    assert "required" not in source


def test_each_text_field_states_its_room_where_the_model_writes_it():
    """The cover subtitle and figure labels failed most first calls while their
    sizes sat only in the catalogue."""
    inventory = analyze_template(TEMPLATE).inventory
    assert inventory is not None
    schema = tools_mod._args_schema(inventory)
    defs = schema.model_json_schema()["$defs"]
    slide = defs["BuildSlide"]["properties"]
    # The most common room, then each kind that differs; a single kind is named.
    title = slide["title"]["description"]
    assert title.startswith(
        "On a content slide, its conclusion as a short sentence (a claim, not a topic). "
        "Room: ≤4 words (≈33 chars); "
        "cover ≤8 words (≈29 chars × 2 lines)"
    )
    assert (
        "Room: section ≤27 words (≈66 chars × 3 lines); "
        "cover ≤12 words (≈33 chars × 3 lines)." in (slide["subtitle"]["description"])
    )
    heading = defs["BuildColumn"]["properties"]["heading"]["description"]
    assert heading.endswith("Room: comparison ≤4 words (≈30 chars).")
    # Hints only: an over-long text must reach the tool, not fail at the provider.
    assert "maxLength" not in json.dumps(schema.model_json_schema())
    # Equal templates share one cached schema, and the plan still coerces.
    assert tools_mod._args_schema(inventory.model_copy(deep=True)) is schema
    plan = tools_mod._to_plan(schema(slides=[{"title": "A", "items": ["x"]}]))
    assert plan.slides[0].items[0].heading == "x"


def test_the_tool_tells_the_agent_one_call_builds_the_whole_deck():
    config, assets = _configured()
    description = _tools(_ctx(config, assets=assets))[BUILD_TOOL_NAME].description
    assert (
        "A NEW deck: write out every slide in `slides`, in order, in ONE call"
        in description
    )
    assert "To CHANGE the last deck of this conversation" in description
    assert (
        '{"op": "insert", "after": N (1 = right after the cover; 0 only for a new '
        "first slide placed before the cover)"
    ) in description
    assert (
        "To change text inside a slide, use `edit`: the slide keeps its design"
        in description
    )
    assert "Never split a deck across several calls" in description
    assert "A first slide of only a title and subtitle is the cover." in description
    # Each offered kind says what it is for, on its own catalogue line only.
    assert (
        "- figures (a few big numbers, each with a short label): " in description
        and "- process (ordered steps or a causal chain; 1 look): " in description
    )
    assert "timeline" not in description
    assert "Use any variant below at most 2 times." in description


def test_the_choice_guide_states_the_authors_max():
    config, assets = _configured()

    def guide(**update: Any) -> str:
        ctx = _ctx(config.model_copy(update=update), assets=assets)
        return _tools(ctx)[BUILD_TOOL_NAME].description

    assert "Each slide must look different." in guide(variety_max_uses=1)
    assert "Use any kind below on at most 3 slides," in guide(
        variety_check="kind", variety_max_uses=3
    )
    off = guide(variety_check="off", variety_max_uses=1)
    assert "at most" not in off.split("Slide kinds of this template")[0]
    assert "look different" not in off


# --- delivery -------------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_build_returns_preview_part_and_uploads_deck(monkeypatch, preview_union):
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _pdf)
    config, assets = _configured()
    workspace = FakeWorkspace()
    tools = _tools(_ctx(config, assets=assets, workspace=workspace))
    text, result = await tools[BUILD_TOOL_NAME].coroutine(
        slides=[COVER, {"title": "Plan", "bullets": ["one", "- two"]}],
        file_name="Kick off 2026",
    )
    assert not result.is_error, text
    (part,) = result.ui_parts
    assert isinstance(part, PptPreviewPart)
    assert part.file_name == "Kick off 2026.pptx"
    # The card opens this capability's pane: an agent may not have ppt_filler.
    assert part.capability_id == "ppt_builder"
    assert "2 slides" in text
    assert tools_mod.deck_delivered(_turn(text))  # a later turn sees it
    # Another call is allowed: it rebuilds the whole deck.
    assert "do not call again" not in text
    assert "send only `changes` numbered on the file's slides (1 to 2" in text
    # Keyed by the deck's name: a model that splits one deck across calls, or
    # retries one, lands on the same key and the reader sees one presentation.
    paths = [w["path"] for w in workspace.writes]
    assert paths == [
        "outputs/sess-9/Kick off 2026.pptx",
        "outputs/sess-9/Kick off 2026.pdf",
    ]
    deck = Presentation(io.BytesIO(workspace.writes[0]["content"]))
    assert [cast(Any, s.shapes.title).text for s in deck.slides] == ["Kick-off", "Plan"]


@pytest.mark.asyncio
async def test_build_without_template_uses_the_built_in_template(monkeypatch):
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    workspace = FakeWorkspace()
    tools = _tools(_ctx(PptBuilderConfig(), workspace=workspace))
    text, result = await tools[BUILD_TOOL_NAME].coroutine(
        slides=[
            COVER,
            {"title": "KPIs", "bullets": ["Margin up to 42%", "Churn down"]},
            {"title": "Merci"},
        ]
    )
    assert not result.is_error, text
    assert "only the download" in text
    assert result.ui_parts[0].kind.value == "download"
    deck = Presentation(io.BytesIO(workspace.writes[0]["content"]))
    assert len(deck.slides) == 3


@pytest.mark.asyncio
async def test_several_builds_of_one_answer_land_on_one_preview(
    monkeypatch, preview_union
):
    """They share a key, so the reader is shown one deck."""
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _pdf)
    config, assets = _configured()
    tools = _tools(_ctx(config, assets=assets, workspace=FakeWorkspace()))
    previews = []
    for title in ("Part one", "Part two", "Part three"):
        _, result = await tools[BUILD_TOOL_NAME].coroutine(slides=[{"title": title}])
        previews.append(cast(Any, result.ui_parts[0]).preview_id)
    assert len(set(previews)) == 1, previews


# --- malformed plans --------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_empty_plan_is_an_error_the_user_can_read(monkeypatch):
    text, result = await _build(monkeypatch, slides=[{"kind": "bullets"}])
    assert result.is_error and "every slide of the plan is empty" in text
    assert "and call build_presentation again" in text
    # The text travels in the result so the runtime shows it, not a generic failure.
    assert result.blocks and result.blocks[0].text == text


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "slide",
    [
        {"title": "Line one\nline two", "bullets": ["a"]},
        {"title": "T", "bullets": ["first\r\nsecond", "b\u2028c"]},
        {
            "title": "Cards",
            "items": [
                {"title": "A\nB", "text": "x\ny"},
                {"title": "C", "text": "z"},
                {"title": "D", "text": "w"},
            ],
        },
    ],
)
async def test_line_breaks_in_the_text_still_build(monkeypatch, slide):
    workspace = FakeWorkspace()
    text, result = await _build(
        monkeypatch, workspace=workspace, variety="off", slides=[COVER, slide]
    )
    assert not result.is_error, text
    assert workspace.writes


@pytest.mark.asyncio
async def test_a_checker_bug_is_a_tool_error_not_an_exception(monkeypatch):
    def _broken(*_args, **_kwargs):
        raise RuntimeError("secret plan text")

    monkeypatch.setattr(tools_mod, "_evaluate", _broken)
    text, result = await _build(monkeypatch, slides=[COVER])
    assert result.is_error and "internal problem" in text
    assert "secret" not in text


@pytest.mark.asyncio
async def test_malformed_plan_from_a_small_model_still_builds(monkeypatch):
    """The shape of a real failing call: the old `pattern` key holding ids or
    legacy names, stray strings in `slides`, bullets as one string, numbers as
    integers, an unknown top-level key."""
    workspace = FakeWorkspace()
    text, result = await _build(
        monkeypatch,
        workspace=workspace,
        notes="stray top-level key",
        slides=[
            {"pattern": "title", "title": "Kick-off", "subtitle": "Team"},
            {
                "pattern": "slide:999",
                "title": "Options",
                "columns": [
                    {"heading": "A", "bullets": "fast\ncheap"},
                    {"heading": "B", "bullets": ["robust"]},
                ],
            },
            {
                "kind": "Figures",
                "title": "KPIs",
                "items": [{"number": 3, "label": "Sites"}, "Loose item"],
            },
            {"title": "Plan", "bullets": "one\ntwo"},
            "notes**: ",
            "A long stray note that is not a slide",
        ],
    )
    assert not result.is_error, text
    deck = Presentation(io.BytesIO(workspace.writes[0]["content"]))
    assert len(deck.slides) == 4
    assert "4 slides" in text
    assert "Entry 5 of the outline was not a slide and was left out." in text


# --- no silent fallback, repair loop ----------------------------------------------------------


@pytest.mark.asyncio
async def test_a_slide_that_fits_nothing_builds_nothing_and_says_what_to_change(
    monkeypatch,
):
    """Seven steps where the largest process slide holds three: the old tool
    dropped four of them silently."""
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, _ = _configured()
    # No template in the store: a fetch would turn the answer into a load error.
    tools = _tools(_ctx(config, assets=FakeAssets(), workspace=workspace))
    steps = [{"number": str(i), "heading": f"Step {i}", "text": "x"} for i in range(7)]
    text, result = await tools[BUILD_TOOL_NAME].coroutine(
        slides=[COVER, {"title": "Method", "kind": "process", "items": steps}]
    )
    # A revision, not a failure: the runtime hands it to the model, never ends the turn.
    assert not result.is_error and result.blocks[0].text == text
    assert "Your very next message MUST be a new build_presentation call" in text
    assert "never write the call, its JSON" in text
    assert workspace.writes == []
    payload = json.loads(text[text.index("{") :])
    assert payload["status"] == "needs_revision"
    (problem,) = payload["slides"]
    assert problem["slide"] == 2 and problem["kind"] == "process"
    assert "items: 7 given; the largest 'process' slide holds 3" in problem["problem"]
    assert any(o.startswith("trim to ") for o in problem["options"])
    assert (
        "or keep 3 here and move the rest to slides of your own" in problem["options"]
    )
    assert "repeated" not in problem

    # The same slide in another call is flagged as resent, not answered word for word.
    text, _ = await tools[BUILD_TOOL_NAME].coroutine(
        slides=[COVER, {"title": "Method", "kind": "process", "items": steps}],
        file_name="Method",
    )
    (again,) = json.loads(text[text.index("{") :])["slides"]
    assert again["repeated"].startswith("you resent this slide unchanged")
    text, _ = await tools[BUILD_TOOL_NAME].coroutine(
        slides=[COVER, {"title": "Method v2", "kind": "process", "items": steps}]
    )
    (changed,) = json.loads(text[text.index("{") :])["slides"]
    assert changed["repeated"] == (
        "the slide changed but this problem is still there: apply one of the options"
    )


@pytest.mark.asyncio
async def test_the_corrected_plan_is_built_as_a_whole(monkeypatch):
    workspace = FakeWorkspace()
    steps = [{"number": str(i), "heading": f"Step {i}", "text": "x"} for i in range(7)]
    _, first = await _build(
        monkeypatch,
        workspace=workspace,
        slides=[COVER, {"title": "Method", "kind": "process", "items": steps}],
    )
    assert not first.is_error and workspace.writes == []
    text, second = await _build(
        monkeypatch,
        workspace=workspace,
        slides=[
            COVER,
            {"title": "Method 1/2", "kind": "process", "items": steps[:3]},
            {"title": "Method 2/2", "kind": "process", "items": steps[3:6]},
        ],
    )
    assert not second.is_error, text
    deck = Presentation(io.BytesIO(workspace.writes[-1]["content"]))
    assert len(deck.slides) == 3


@pytest.mark.asyncio
async def test_a_parallel_copy_of_a_call_is_not_built_twice(monkeypatch):
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    plan = [COVER, {"title": "Agenda", "bullets": ["One", "Two"]}]
    # Called directly (no middleware), the copy waits and finds the deck delivered.
    (text_a, a), (text_b, b) = await asyncio.gather(
        build.coroutine(slides=plan), build.coroutine(slides=plan)
    )
    assert text_a.startswith(tools_mod._DELIVERED) and not a.is_error
    assert text_b.startswith("No deck was built: unchanged") and not b.is_error
    assert len(workspace.writes) == 1
    # A different plan sent in parallel is built too, one after the other.
    (_, c), (_, d) = await asyncio.gather(
        build.coroutine(slides=[COVER]), build.coroutine(slides=plan)
    )
    assert not c.is_error and not d.is_error and len(workspace.writes) == 3


# --- variety, asked once on the first deck --------------------------------------------------


def _steps(n: int) -> dict:
    items = [
        {"number": "1", "heading": "a", "text": "b"},
        {"number": "2", "heading": "c", "text": "d"},
    ]
    return {"title": f"Step {n}", "items": items}


_REPEATED = [COVER] + [_steps(n) for n in range(4)]


@pytest.mark.asyncio
async def test_the_first_deck_is_asked_to_vary_once_and_a_resend_is_built(monkeypatch):
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error and workspace.writes == []
    assert "Your very next message MUST be a new build_presentation call" in text
    assert "The limit counts within each section" in text
    payload = json.loads(text[text.index("{") :])
    assert payload["status"] == "needs_revision"
    (repeat,) = payload["variety"]
    assert repeat["slides"] == [2, 3, 4, 5] and repeat["kind"] == "process"
    assert repeat["at_most"] == 2
    # Two figures variants: two slides suggested, and no design untested on the content.
    assert repeat["suggested"] == [
        "slides 2, 3 (short figures) → figures holds items max 4 [number, label]"
    ]
    assert "options" not in repeat
    assert "send the same call again: it will be built as is" in text
    # The same plan again: built, the repetition no longer told.
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error, text
    assert "all use the template slide" not in text


@pytest.mark.asyncio
async def test_a_malformed_first_call_does_not_use_up_the_variety_ask(monkeypatch):
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    text, result = await build.coroutine(slides=[{"kind": "bullets"}])
    assert result.is_error and "every slide of the plan is empty" in text
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error and '"variety": [' in text


_TOO_MANY = {
    "title": "Method",
    "kind": "process",
    "items": [
        {"number": str(i), "heading": f"Step {i}", "text": "x"} for i in range(7)
    ],
}


@pytest.mark.asyncio
async def test_no_variety_ask_after_a_fit_repair(monkeypatch):
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    text, result = await build.coroutine(slides=[COVER, _TOO_MANY])
    assert not result.is_error and "variety" not in text
    # The repaired deck repeats a kind: built, the repetition not told.
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error, text
    assert "all use the template slide" not in text


@pytest.mark.asyncio
async def test_a_first_call_with_fit_problems_carries_the_variety_ask(monkeypatch):
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    text, result = await build.coroutine(slides=_REPEATED + [_TOO_MANY])
    assert not result.is_error and workspace.writes == []
    payload = json.loads(text[text.index("{") :])
    assert [p["slide"] for p in payload["slides"]] == [6]
    (repeat,) = payload["variety"]
    # The slide to fix counts on its closest fit, the same process slide.
    assert repeat["slides"] == [2, 3, 4, 5, 6]
    assert "the next deck that fits is built as is" in text
    # Fixed, kinds unchanged: one round was enough, the deck is built.
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error, text


def _labelled_steps(n: int, label: bool = True) -> dict:
    steps = _steps(n)
    if label:
        steps["items"] = [{**i, "label": "note"} for i in steps["items"]]
    return steps


@pytest.mark.asyncio
async def test_slides_to_fix_that_repeat_are_asked_to_vary_in_the_same_round(
    monkeypatch,
):
    """The live call: both repeated slides had a fit problem, so the first
    answer named no repeat and the repaired deck broke the max of 1."""
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    config = config.model_copy(update={"variety_max_uses": 1})
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    text, result = await build.coroutine(
        slides=[COVER, _labelled_steps(1), _labelled_steps(2)]
    )
    assert not result.is_error and workspace.writes == []
    payload = json.loads(text[text.index("{") :])
    assert [p["slide"] for p in payload["slides"]] == [2, 3]
    assert all(
        "closest fit, a 'process' slide" in p["problem"] for p in payload["slides"]
    )
    (repeat,) = payload["variety"]
    assert repeat["slides"] == [2, 3] and repeat["at_most"] == 1
    assert "several slides would look the same" in text
    # Fixed, kinds unchanged: built as is, the repeat not told.
    fixed = [COVER, _labelled_steps(1, False), _labelled_steps(2, False)]
    text, result = await build.coroutine(slides=fixed)
    assert not result.is_error, text
    assert "all use the template slide" not in text


class _Request:
    system_prompt = ""

    def __init__(self, messages: list, tools: list) -> None:
        self.messages, self.tools = messages, tools

    def override(self, **_update: Any) -> "_Request":
        return self


async def _later_turn(
    monkeypatch, history: list, extend=None, variety="template_slide"
):
    """A new turn's build tool, after the model call that was sent ``history``."""
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured(extend, variety)
    (middleware,) = PptBuilderCapability().middleware(_ctx(config, assets=assets))
    middleware = cast(Any, middleware)

    async def _handler(request):
        return request

    await middleware.awrap_model_call(_Request(history, middleware.tools), _handler)
    return {t.name: cast(Any, t) for t in middleware.tools}[BUILD_TOOL_NAME]


_CALLS = itertools.count(1)


def _turn(text: str, name: str = BUILD_TOOL_NAME) -> list:
    call = {"id": f"call-{next(_CALLS)}", "name": name, "args": {}}
    return [
        HumanMessage(content="Make me a deck"),
        AIMessage(content="", tool_calls=[call]),
        ToolMessage(content=text, tool_call_id=call["id"], name=name),
    ]


def _rehydrated(history: list) -> list:
    """What a capability middleware is sent on a later turn: the platform's
    checkpoint hygiene strips every message name; content may come as blocks."""
    return [
        m.model_copy(
            update={"name": None, "content": [{"type": "text", "text": m.content}]}
            if isinstance(m, ToolMessage)
            else {"name": None}
        )
        for m in history
    ]


@pytest.mark.asyncio
async def test_a_deck_after_the_first_one_is_built_as_is(monkeypatch):
    """A delivered deck in the history: the user may have asked for these slides."""
    delivered, result = await _build(monkeypatch, slides=[COVER])
    assert not result.is_error
    build = await _later_turn(monkeypatch, _turn(delivered))
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error, text
    assert '"variety"' not in text and "all use the template slide" not in text


@pytest.mark.asyncio
async def test_a_deck_delivered_in_an_earlier_turn_is_seen_with_names_stripped(
    monkeypatch,
):
    """Live: the variety ask came back on turn 3 for slides delivered on turn 1,
    because the model call's tool messages carry no name."""
    delivered, result = await _build(monkeypatch, slides=[COVER])
    assert not result.is_error
    history = _rehydrated(_turn(delivered))
    assert tools_mod.deck_delivered(history)
    build = await _later_turn(monkeypatch, history)
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error, text
    # Another tool's delivery-like text, names stripped, still delivers nothing.
    assert not tools_mod.deck_delivered(_rehydrated(_turn(delivered, "other")))


@pytest.mark.asyncio
async def test_a_history_without_a_delivered_deck_is_still_a_first_deck(monkeypatch):
    asked, _ = await _build(monkeypatch, slides=_REPEATED)
    assert '"variety": [' in asked
    unread, _ = await _build(monkeypatch, slides=[{"kind": "bullets"}])
    # A needs_revision, an error, or another tool's text never delivered a deck.
    history = _turn(asked) + _turn(unread) + _turn("Delivered: 'x.pptx'", "other")
    build = await _later_turn(monkeypatch, history)
    text, result = await build.coroutine(slides=_REPEATED)
    assert not result.is_error and '"variety": [' in text


def test_both_success_texts_count_as_a_delivered_deck():
    assert tools_mod.deck_delivered(_turn("Delivered: 'a.pptx', 3 slides."))
    assert tools_mod.deck_delivered(
        _turn("The presentation 'a.pptx' (3 slides) is ready")
    )
    assert not tools_mod.deck_delivered(
        _turn("The presentation template could not be loaded right now.")
    )
    assert not tools_mod.deck_delivered([])


@pytest.mark.asyncio
async def test_one_kind_on_different_slides_is_asked_once_with_content_advice(
    monkeypatch,
):
    """The real first deck: five bullets slides on five template slides."""
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured(with_item_slides, "kind")
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    text, result = await build.coroutine(slides=EVIDENCE_SLIDES)
    assert not result.is_error and workspace.writes == []
    (repeat,) = json.loads(text[text.index("{") :])["variety"]
    assert repeat["kind"] == "bullets" and repeat["slides"] == [3, 5, 7, 9, 10]
    suggested = repeat["suggested"]
    assert any(
        s.startswith(
            "slide 7 (dated items) → timeline holds items max 6 [heading ≤12 words (≈20"
        )
        for s in suggested
    )
    assert any(s.startswith("slide 5 (parallel") for s in suggested)
    text, result = await build.coroutine(slides=EVIDENCE_SLIDES)
    assert not result.is_error, text
    assert "are all 'bullets' slides" not in text


@pytest.mark.asyncio
async def test_a_later_deck_repeating_a_kind_is_built_as_is(monkeypatch):
    delivered, _ = await _build(monkeypatch, slides=[COVER])
    build = await _later_turn(
        monkeypatch, _turn(delivered), extend=with_item_slides, variety="kind"
    )
    text, result = await build.coroutine(slides=EVIDENCE_SLIDES)
    assert not result.is_error, text
    assert '"variety"' not in text and "are all 'bullets' slides" not in text


@pytest.mark.asyncio
async def test_each_variety_check_mode_asks_for_its_own_repeat(monkeypatch):
    """Five bullets slides on five template slides: a kind repeat, no template
    slide used three times; the steps reuse one template slide."""
    for variety, plan, asked in [
        ("template_slide", EVIDENCE_SLIDES, False),
        ("kind", EVIDENCE_SLIDES, True),
        ("template_slide", _REPEATED, True),
        ("off", _REPEATED, False),
        ("off", EVIDENCE_SLIDES, False),
    ]:
        # The added item slides are never rendered: only the evidence plan needs them.
        extend = with_item_slides if plan is EVIDENCE_SLIDES else None
        text, result = await _build(
            monkeypatch, extend=extend, variety=variety, slides=plan
        )
        assert not result.is_error, (variety, text)
        assert asked is ('"variety": [' in text), (variety, text)
        # A delivered deck never talks layout reuse.
        assert "all use the template slide" not in text and "are all" not in text


def test_the_variety_check_defaults_and_never_triggers_a_re_analysis():
    stored = PptBuilderConfig.model_validate({"template_sha256": "sha"})
    assert stored.variety_check == "template_slide" and stored.variety_max_uses == 2
    keys = {
        PptBuilderConfig(
            template_sha256="sha", variety_check=v, variety_max_uses=m
        ).labelling_key()
        for v in ("template_slide", "kind", "off")
        for m in (1, 5)
    }
    assert keys == {stored.labelling_key()}
    with pytest.raises(ValueError):
        PptBuilderConfig(variety_max_uses=6)
    fields = {f.key: f for f in PptBuilderCapability.manifest.config_fields}
    check, most = fields["variety_check"], fields["variety_max_uses"]
    assert check.enum == ["template_slide", "kind", "off"]
    assert check.default == "template_slide" and most.default == 2
    assert (most.min, most.max) == (1, 5)
    # The plugin widget owns both: one sentence, not two generic fields.
    assert check.ui.widget == most.ui.widget == "ppt_builder_template"


@pytest.mark.asyncio
async def test_structural_repeats_are_never_a_revision(monkeypatch):
    sections = [COVER] + [{"title": f"Part {n}", "kind": "section"} for n in range(4)]
    text, result = await _build(monkeypatch, slides=sections)
    assert not result.is_error, text


@pytest.mark.asyncio
async def test_a_heading2_column_is_returned_for_revision_not_dropped(monkeypatch):
    workspace = FakeWorkspace()
    text, result = await _build(
        monkeypatch,
        workspace=workspace,
        slides=[
            COVER,
            {
                "kind": "two_columns",
                "title": "Implantation",
                "columns": {
                    "bullets": ["2 800 experts"],
                    "heading": "Espagne",
                    "heading2": "Italie",
                },
            },
        ],
    )
    assert not result.is_error and workspace.writes == []
    (problem,) = json.loads(text[text.index("{") :])["slides"]
    assert problem["slide"] == 2
    assert "heading2 'Italie' is not a column field" in problem["problem"]


def test_the_success_notes_name_what_was_cut_per_slide():
    """Read from the render's after-check report, one line per slide."""
    plan = DeckPlan.model_validate({"slides": [COVER, "junk"]})
    report = SlideReport(1, removed=["Lorem"], truncated=["title"], lost=["subtitle"])
    notes = _build_notes(plan, [report], [])
    assert "Entry 2 of the outline was not a slide and was left out." in notes
    assert (
        "Slide 1: text cut to fit in title; not placed: subtitle (no room on this "
        "slide)." in notes
    )
    assert "Lorem" not in notes
    assert "Tell the user (they can ask for a fix in their next message):" in notes
    assert _build_notes(plan, [SlideReport(1, removed=["Lorem"])], []).count("\n") == 2
    assert _build_notes(DeckPlan.model_validate({"slides": [COVER]}), [], []) == ""


@pytest.mark.asyncio
async def test_a_renderer_finding_reaches_the_model_in_the_review_notes(monkeypatch):
    """An agenda announcing parts the deck never covers was only logged; the
    model now reads it, with its remedy, after the unchanged success prefix."""
    agenda = {
        "title": "Agenda",
        "kind": "agenda",
        "items": [{"heading": f"Part {i}"} for i in range(4)],
    }
    text, result = await _build(monkeypatch, slides=[COVER, agenda])
    assert not result.is_error, text
    assert tools_mod.deck_delivered(_turn(text))
    review = text[text.index("\nTell the user") :]
    assert (
        "Slide 2: the agenda announces 4 parts but only 0 slide(s) follow it. Add the "
        "missing slides or remove those parts from the agenda."
    ) in review


def _points(n: int) -> list:
    return [{"title": f"Point {i}", "bullets": ["a", "b"]} for i in range(n)]


def _agenda(n: int) -> dict:
    return {
        "title": "Agenda",
        "kind": "agenda",
        "items": [{"heading": f"Part {i}"} for i in range(n)],
    }


@pytest.mark.asyncio
async def test_a_split_agenda_is_one_agenda_for_the_gap_check(monkeypatch):
    """Six parts over two agenda slides of 3: the slides that follow are
    counted after the second part, the parts announced once."""
    text, result = await _build(monkeypatch, slides=[COVER, _agenda(6), *_points(2)])
    assert not result.is_error, text
    review = text[text.index("\nTell the user") :]
    assert "Slide 2: its 6 items continue on a second 'agenda' slide (3 + 3" in review
    assert (
        "Slide 2: the agenda announces 6 parts but only 2 slide(s) follow it." in review
    )
    text, _ = await _build(monkeypatch, slides=[COVER, _agenda(6), *_points(3)])
    assert "announces" not in text


@pytest.mark.asyncio
async def test_notes_after_a_split_name_the_files_slide_numbers(monkeypatch):
    """The process list of slide 2 takes two slides; the agenda after it is
    slide 4 of the file, and every note says so."""
    steps = [{"heading": f"Step {i}", "text": "x"} for i in range(5)]
    method = {"title": "Method", "kind": "process", "items": steps}
    text, result = await _build(
        monkeypatch, slides=[COVER, method, _agenda(4), *_points(1)]
    )
    assert not result.is_error, text
    assert "(5 slides)" in text
    assert "Slide 2: its 5 items continue on a second 'process' slide (3 + 2" in text
    assert (
        "Slide 4: the agenda announces 4 parts but only 1 slide(s) follow it." in text
    )


def _relabel(**labels: str):
    def extend(inventory):
        patterns = [
            p.model_copy(update={"label": labels[p.id]}) if p.id in labels else p
            for p in inventory.patterns
        ]
        return inventory.model_copy(update={"patterns": patterns})

    return extend


@pytest.mark.asyncio
async def test_the_delivered_text_names_each_file_slides_design(monkeypatch):
    """A "same layout as slide 5" request is about the template slide, not the kind: the
    model reads which design each file slide has, a split list's parts included."""
    steps = [{"heading": f"Step {i}", "text": "x"} for i in range(5)]
    method = {"title": "Method", "kind": "process", "items": steps}
    text, result = await _build(
        monkeypatch, slides=[COVER, method, _agenda(4), *_points(3)]
    )
    assert not result.is_error, text
    (designs,) = [line for line in text.splitlines() if line.startswith("Designs: ")]
    assert designs.startswith(
        "Designs: 1 'Title Slide' · 2 'Process' · 3 'Process' · 4 'Agenda' · "
    )
    # Three bullets slides on three template slides: three names.
    names = designs.split(". Slides")[0].split(" · ")[4:]
    assert [n.split(" ", 1)[0] for n in names] == ["5", "6", "7"]
    assert len({n.split(" ", 1)[1] for n in names}) == 3
    assert '"field": "like", "value": M}' in designs
    assert "slide:" not in designs and "layout:" not in designs


@pytest.mark.asyncio
async def test_two_template_slides_sharing_a_label_get_distinct_names(monkeypatch):
    extend = _relabel(**{"slide:269": "Points"}, **{"slide:266": "Points"})
    text, result = await _build(monkeypatch, extend=extend, slides=[COVER, *_points(3)])
    assert not result.is_error, text
    (designs,) = [line for line in text.splitlines() if line.startswith("Designs: ")]
    assert "'Points'" in designs and "'Points #2'" in designs


def test_a_long_design_name_is_cut():
    inventory = analyze_template(TEMPLATE).inventory
    assert inventory is not None
    inventory = _relabel(**{"slide:260": "A very long name for a template slide"})(
        inventory
    )
    names = tools_mod._design_names(inventory)
    assert names["slide:260"] == "A very long name for a temp…"
    assert len(names["slide:260"]) <= tools_mod._DESIGN_NAME_CHARS
    # Two labels cut to the same text are still two designs.
    inventory = _relabel(**{"slide:269": "A very long name for a template layout"})(
        inventory
    )
    names = tools_mod._design_names(inventory)
    assert {names["slide:260"], names["slide:269"]} == {
        "A very long name for a temp…",
        "A very long name for a temp… #2",
    }


@pytest.mark.asyncio
async def test_the_delivered_part_summarises_what_was_built(monkeypatch, preview_union):
    """The dialog beside the download reads the tool's record, not the reply:
    a split list, a close-kind fallback and a picture that did not load."""
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _pdf)
    config, assets = _configured()
    docs = FakeDocs({"txt-1": (b"hello", "text/plain")})
    tools = _tools(_ctx(config, assets=assets, docs=docs))
    points = {"title": "Points", "bullets": ["a", "b"], "kind": "timeline"}
    site = {"title": "Site", "image": "txt-1"}
    text, result = await tools[BUILD_TOOL_NAME].coroutine(
        slides=[COVER, _agenda(6), points, site]
    )
    assert not result.is_error, text
    (part,) = result.ui_parts
    summary = part.summary
    assert [(s.number, s.part, s.parts) for s in summary] == [
        (1, 1, 1),
        (2, 1, 2),
        (3, 2, 2),
        (4, 1, 1),
        (5, 1, 1),
    ]
    cover, first, second, fallback, picture = summary
    assert (cover.title, cover.kind, cover.notes) == ("Kick-off", "cover", [])
    assert first.title == "Agenda (1/2)" and second.title == "Agenda (2/2)"
    assert first.kind == second.kind == "agenda"
    assert (
        [n.code for n in first.notes]
        == [n.code for n in second.notes]
        == ["list_continued"]
    )
    assert fallback.kind == "bullets"
    assert [(n.code, n.kind) for n in fallback.notes] == [("other_kind", "timeline")]
    assert [(n.code, n.count) for n in picture.notes] == [("image_unusable", 1)]
    # Codes and counts only: no English prose reaches the user.
    dumped = json.dumps(part.model_dump(mode="json")["summary"])
    assert "Nothing to do" not in dumped and "call again" not in dumped
    # A layout kept as the user asked, or not, is never silent.
    built = [BuiltSlide(number=n, kind="bullets", template="A") for n in (1, 2, 3)]
    sent = DeckPlan(slides=[SlidePlan(title=str(n)) for n in (1, 2, 3)])
    noted = tools_mod._summary(built, sent, [], [1, 2, 3], {2: True, 3: False})
    assert [[n.code for n in s.notes] for s in noted] == [
        [],
        ["layout_repeated"],
        ["layout_not_repeated"],
    ]
    # ppt_filler delivers the same part without one: no button there.
    assert (
        PptPreviewPart(
            preview_id="k", title="t", pdf_download_url="/p", version="v"
        ).summary
        is None
    )


# --- images -----------------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_build_places_images_from_documents_and_reports_bad_ones(monkeypatch):
    docs = FakeDocs(
        {"img-1": (png_bytes(20, 10), "image/png"), "txt-1": (b"hello", "text/plain")}
    )

    def plan(image: str) -> list:
        return [COVER, {"title": "Site", "image": image}]

    text, result = await _build(monkeypatch, docs=docs, slides=plan("img-1"))
    assert not result.is_error, text
    assert docs.fetch_uids == ["img-1"]
    # A bad image reference costs the picture, never the deck.
    text, result = await _build(monkeypatch, docs=docs, slides=plan("txt-1"))
    assert not result.is_error and "image 'txt-1'" in text
    text, result = await _build(monkeypatch, docs=docs, slides=plan("missing"))
    assert not result.is_error and "image 'missing'" in text


@pytest.mark.asyncio
async def test_list_images_resolves_the_folder_at_call_time():
    folders = FakeFolders(
        {"Images/Sites": "tag-1"},
        {"tag-1": [("d1", "bordeaux.png"), ("d2", "notes.docx")]},
    )
    config, assets = _configured()
    tools = _tools(_ctx(config, assets=assets, folders=folders))
    listed = await tools[LIST_IMAGES_TOOL_NAME].coroutine(folder="/Images/Sites/")
    assert '"d1"' in listed and "d2" not in listed
    missing = await tools[LIST_IMAGES_TOOL_NAME].coroutine(folder="Nope")
    assert "No folder named 'Nope'" in missing


# --- parallel calls of one model round, identical resends --------------------------------


def _round(*args: dict) -> tuple[AIMessage, list[dict]]:
    """One assistant message asking for several tool calls at once."""
    calls = [
        {"id": f"call-{next(_CALLS)}", "name": BUILD_TOOL_NAME, "args": a} for a in args
    ]
    other = {"id": f"call-{next(_CALLS)}", "name": "other", "args": {}}
    return AIMessage(content="", tool_calls=[other, *calls]), calls


async def _through_middleware(middleware, message: AIMessage, call: dict):
    """One tool call as the agent loop runs it: the middleware, then the tool."""
    from langgraph.prebuilt.tool_node import ToolCallRequest

    build = {t.name: t for t in middleware.tools}[BUILD_TOOL_NAME]
    ran: list[str] = []

    async def _handler(request):
        ran.append(request.tool_call["id"])
        text, artifact = await build.coroutine(**request.tool_call["args"])
        return ToolMessage(
            content=text, artifact=artifact, tool_call_id=request.tool_call["id"]
        )

    request = ToolCallRequest(
        tool_call=cast(Any, call),
        tool=build,
        state={"messages": [HumanMessage(content="deck"), message]},
        runtime=cast(Any, None),
    )
    return await middleware.awrap_tool_call(request, _handler), ran


@pytest.mark.asyncio
async def test_parallel_build_calls_of_one_round_build_only_the_first(monkeypatch):
    """Live: a small model sent one deck as several different calls at once;
    the cover+agenda call was delivered and the fragments were refused for
    not starting with a cover, which led it to split further."""
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    (middleware,) = PptBuilderCapability().middleware(
        _ctx(config, assets=assets, workspace=workspace)
    )
    middleware = cast(Any, middleware)
    message, calls = _round(
        {"slides": [COVER, {"title": "Agenda", "bullets": ["One", "Two"]}]},
        {"slides": [{"title": "Method", "bullets": ["Plan", "Do"]}]},
        {"slides": [{"title": "Results", "bullets": ["Up", "Down"]}]},
    )
    results = await asyncio.gather(
        *(_through_middleware(middleware, message, call) for call in calls)
    )
    (first, ran_first), *others = results
    assert ran_first == [calls[0]["id"]]
    assert first.content.startswith(tools_mod._DELIVERED), first.content
    assert len(workspace.writes) == 1
    for position, (refused, ran) in enumerate(others, start=2):
        # Not a failure: the round still runs its first call.
        assert ran == [] and not refused.artifact.is_error
        assert refused.tool_call_id == calls[position - 1]["id"]
        assert refused.content == (
            "No deck was built: build_presentation was called 3 times at once in this "
            "answer; only the first call (the one whose first slide is 'Kick-off') is "
            f"processed and this one (call {position} of 3) was ignored. A deck is never "
            "split across calls: send ONE call with every slide in `slides`, the cover "
            "first. To add or change slides of a delivered deck, send only `changes` "
            '({"op": "insert", "after": N, "slides": [...]}, {"op": "replace", "slide": '
            'N, "slides": [...]}, {"op": "edit", ...}). Read the first call\'s result '
            "before calling again, and do not describe this call to the user as a deck."
        )


@pytest.mark.asyncio
async def test_a_single_build_call_or_a_copy_in_parallel_is_handled_once(monkeypatch):
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    (middleware,) = PptBuilderCapability().middleware(
        _ctx(config, assets=assets, workspace=workspace)
    )
    middleware = cast(Any, middleware)
    plan = {"slides": [COVER, {"title": "Agenda", "bullets": ["One", "Two"]}]}
    message, (alone,) = _round(plan)
    reply, ran = await _through_middleware(middleware, message, alone)
    assert ran == [alone["id"]] and not reply.artifact.is_error
    # The same call twice at once: the copy is answered without a build.
    message, (a, b) = _round(plan, {**plan, "file_name": "copy"})
    (_, ran_a), (copy, ran_b) = await asyncio.gather(
        _through_middleware(middleware, message, a),
        _through_middleware(middleware, message, b),
    )
    assert ran_a == [a["id"]] and ran_b == [] and not copy.artifact.is_error
    assert len(workspace.writes) == 1
    # A call missing from the state (another path) runs as before.
    assert tools_mod.round_refusal([], "call-x") is None


@pytest.mark.asyncio
async def test_a_deck_without_a_cover_is_built_as_sent(monkeypatch):
    """No cover refusal: the deck holds exactly what the model sends."""
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    build = _tools(_ctx(config, assets=assets, workspace=FakeWorkspace()))[
        BUILD_TOOL_NAME
    ]
    alone = [{"title": "Method", "columns": [{"heading": "A", "bullets": ["x"]}] * 2}]
    text, result = await build.coroutine(slides=alone)
    assert text.startswith(tools_mod._DELIVERED), text
    assert "Designs: 1 'Two columns'." in text


@pytest.mark.asyncio
async def test_resending_the_delivered_deck_rebuilds_nothing(monkeypatch):
    """Live: the model resent the deck it had just been delivered, which was
    rebuilt for ~14 s for nothing and used up one of the turn's tool calls."""
    workspace = FakeWorkspace()
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, assets = _configured()
    build = _tools(_ctx(config, assets=assets, workspace=workspace))[BUILD_TOOL_NAME]
    plan = [COVER, {"title": "Agenda", "bullets": ["One", "Two"]}]
    text, _ = await build.coroutine(slides=plan, file_name="Kick-off")
    assert text.startswith(tools_mod._DELIVERED), text

    rendered: list[int] = []
    render = tools_mod._prepare_and_render

    def _counted(*args: Any, **kwargs: Any):
        rendered.append(1)
        return render(*args, **kwargs)

    monkeypatch.setattr(tools_mod, "_prepare_and_render", _counted)
    text, result = await build.coroutine(
        slides=json.loads(json.dumps(plan)), file_name="Kick-off"
    )
    assert not result.is_error and not result.ui_parts
    assert text == (
        "No deck was built: unchanged, deck already delivered as 'Kick-off.pptx' "
        "(2 slides). This plan is exactly that deck, so nothing was rebuilt and the "
        "user's file is the same. Do not tell the user it was changed or rebuilt; to "
        "change it, send only `changes`."
    )
    assert rendered == [] and len(workspace.writes) == 1
    # Another file name, or other content, is a new deck.
    text, _ = await build.coroutine(slides=plan, file_name="Other")
    assert text.startswith(tools_mod._DELIVERED), text
    text, _ = await build.coroutine(slides=[COVER], file_name="Kick-off")
    assert text.startswith(tools_mod._DELIVERED), text
    assert len(rendered) == 2


def test_deck_fields_written_inside_changes_are_taken_to_the_top_level():
    """A model's broken quoting leaves `file_name': 'X` among the changes: the
    field is unambiguous there, so it is read rather than the call refused."""
    edit = {"op": "edit", "slide": 3, "field": "title", "value": "Coûts"}
    sent = {
        "changes": [
            edit,
            "file_name': 'Modernisation_SI",
            "audience': 'Comité de pilotage",
            "main_message': ",
            "}```json",
            {"main_message": "Réduire les coûts"},
        ]
    }
    hoisted, twice = tools_mod._hoisted(sent)
    assert twice is None
    assert hoisted == {
        "changes": [edit],
        "file_name": "Modernisation_SI",
        "audience": "Comité de pilotage",
        "main_message": "Réduire les coûts",
    }
    same, twice = tools_mod._hoisted({**sent, "audience": "Comité de pilotage"})
    assert twice is None and same["changes"] == [edit]
    _, twice = tools_mod._hoisted({**sent, "file_name": "Autre"})
    assert twice == (
        "No deck was built: `file_name` is given twice, 'Autre' and "
        "'Modernisation_SI' inside `changes`; send it once, next to `changes`."
    )
    other = {"changes": [edit, "not a change"]}
    assert tools_mod._hoisted(other) == (other, None)


async def _called(build: Any, args: dict[str, Any]) -> tuple[str, Any]:
    """As the runtime calls it: the args schema parses the call first."""
    message = await build.ainvoke(
        {"type": "tool_call", "id": "call-1", "name": build.name, "args": args}
    )
    return message.content, message.artifact


def _replay(name: str) -> dict[str, Any]:
    return json.loads((Path(__file__).parent / "replays" / name).read_text())


@pytest.mark.asyncio
async def test_a_deck_field_written_inside_a_slide_is_read_as_the_decks(monkeypatch):
    """Live (mistral-medium): `file_name` sat inside the cover slide and the
    whole deck was refused, telling the model to put it in the slide's text."""
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    # The session's own template is not in the repo: the built-in one, whose
    # variety note that template did not raise, checked apart.
    config = PptBuilderConfig.model_validate({"variety_check": "off"})
    build = _tools(_ctx(config))[BUILD_TOOL_NAME]
    sent = _replay("a611e51e_first_call.json")
    assert sent["slides"][0]["file_name"] == "Seconde_Guerre_mondiale_ados"
    text, result = await _called(build, sent)
    assert text.startswith(tools_mod._DELIVERED), text
    assert (
        cast(Any, result.ui_parts[0]).file_name == "Seconde_Guerre_mondiale_ados.pptx"
    )
    text, _ = await _called(build, {**sent, "file_name": "Autre"})
    assert text == (
        "No deck was built: `file_name` is given twice, 'Autre' and "
        "'Seconde_Guerre_mondiale_ados' inside `slides`; send it once, next to "
        "`slides`."
    )


@pytest.mark.asyncio
async def test_a_deck_field_inside_a_parsed_slide_is_never_a_slide_problem(
    monkeypatch,
):
    """Live (0f2506ea, mistral-medium): the cover's `file_name` came back as
    'a deck field, not a slide field' once the runtime had parsed the slides."""
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    build = _tools(_ctx(PptBuilderConfig()))[BUILD_TOOL_NAME]
    sent = _replay("0f2506ea_first_build.json")
    assert "file_name" in sent["slides"][0] and "file_name" not in sent
    text, result = await _called(build, sent)
    assert "deck field" not in text and "file_name" not in text, text
    if text.startswith(tools_mod._DELIVERED):
        assert cast(Any, result.ui_parts[0]).file_name.startswith("Valteris_")


def test_the_designs_line_names_look_alike_template_slides_once():
    inventory = analyze_template(build_showcase_template()).inventory
    assert inventory is not None
    first, second = [p for p in inventory.patterns if p.included][:2]
    twin = second.model_copy(update={"design": first.id})
    inventory = inventory.model_copy(
        update={"patterns": [twin if p is second else p for p in inventory.patterns]}
    )
    line = tools_mod._designs_line(
        inventory, [SlidePlan(title="a"), SlidePlan(title="b")], [first, twin]
    )
    assert f"1 '{first.label}' · 2 '{first.label}'." in line


@pytest.mark.asyncio
async def test_an_unreadable_plan_is_logged_without_its_text(monkeypatch, caplog):
    from pydantic import BaseModel

    class _Slide(BaseModel):
        title: int

    def _unreadable(_validated):
        return _Slide.model_validate({"title": "confidential figures"})

    monkeypatch.setattr(tools_mod, "_to_plan", _unreadable)
    text, result = await _build(monkeypatch, slides=[COVER])
    assert result.is_error, text
    logged = [r.getMessage() for r in caplog.records if "unreadable plan" in r.msg]
    assert logged == ["[PPTBUILD][TOOL] unreadable plan: title: int_parsing"]
    assert "confidential" not in " ".join(r.getMessage() for r in caplog.records)


@pytest.mark.asyncio
async def test_a_ragged_table_is_sent_back_with_its_rows_named(monkeypatch):
    table = {"columns": ["A", "B", "C"], "rows": [["1"], ["1", "2", "3", "4"]]}
    text, result = await _build(
        monkeypatch, variety="off", slides=[COVER, {"title": "T", "table": table}]
    )
    assert not result.is_error and "No deck was built" in text
    assert "table.rows[0] has 1 cells for 3 columns" in text
    assert "table.rows[1] has 4 cells for 3 columns" in text


@pytest.mark.asyncio
async def test_characters_xml_refuses_are_dropped_and_the_deck_builds(monkeypatch):
    workspace = FakeWorkspace()
    slide = {"title": "bad\x00\x0b title \ud800", "bullets": ["￾ x", "a\x08b"]}
    text, result = await _build(
        monkeypatch, workspace=workspace, variety="off", slides=[COVER, slide]
    )
    assert not result.is_error and text.startswith(tools_mod._DELIVERED), text
    deck = Presentation(io.BytesIO(workspace.writes[0]["content"]))
    title = cast(Any, deck.slides[1].shapes.title).text
    assert title.split() == ["bad", "title"]


@pytest.mark.asyncio
async def test_a_picture_too_large_is_left_out_with_a_note(monkeypatch):
    from deck_builders import declared_png

    docs = FakeDocs({"huge-1": (declared_png(9_000, 9_000), "image/png")})
    text, result = await _build(
        monkeypatch,
        docs=docs,
        variety="off",
        slides=[COVER, {"title": "Site", "image": "huge-1"}],
    )
    assert not result.is_error, text
    assert "image 'huge-1' was left out: the picture is too large" in text


@pytest.mark.asyncio
async def test_a_preview_stuck_in_the_queue_leaves_the_deck_download_only(
    monkeypatch, preview_union
):
    async def _stuck(_bytes: bytes, *, suffix: str, **_options):
        await asyncio.sleep(30)

    config, assets = _configured()
    workspace = FakeWorkspace()
    tools = _tools(_ctx(config, assets=assets, workspace=workspace))
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _stuck)
    monkeypatch.setattr(tools_mod, "PREVIEW_WAIT_SECONDS", 0.05)
    text, result = await tools[BUILD_TOOL_NAME].coroutine(slides=[COVER])
    assert not result.is_error and "only the download" in text
    assert result.ui_parts[0].kind.value == "download"
    assert [w["path"][-5:] for w in workspace.writes] == [".pptx"]
