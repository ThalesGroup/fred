# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The design system read from a template (design.md D2, D11, D12, D13).

Two templates all the way through: the built-in theme, which stands for a rich
corporate one, and a stripped deck reduced to a title layout, a blank layout
and a closing slide — the shape that produced nine cover slides.
"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Inches

from fred_capability_ppt_filler.builder.analysis import (
    analyze_presentation,
    theme_colors,
)
from fred_capability_ppt_filler.builder.design_system import (
    MIN_CONTRAST,
    Box,
    ColorToken,
    best_pair,
    contrast_ratio,
    free_area,
    relative_luminance,
)
from fred_capability_ppt_filler.builder.theme import default_theme
from tests.deck_builders import (
    HOUSE_CORNER,
    HOUSE_GAP_IN,
    HOUSE_INSET_IN,
    add_fixture,
    build_cover_only_template,
    build_house_style_template,
)


@pytest.fixture(scope="module")
def rich():
    _, inventory = default_theme()
    return inventory.design_system


@pytest.fixture(scope="module")
def stripped():
    """A template with nothing to fill: one title slide, one blank, one closing."""
    presentation = Presentation()
    title = presentation.slides.add_slide(presentation.slide_layouts[0])
    title.shapes.title.text = "Cover"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    closing = presentation.slides.add_slide(presentation.slide_layouts[5])
    closing.shapes.title.text = "Thank you"
    return analyze_presentation(presentation).design_system


# -- what every template must yield -------------------------------------------


@pytest.mark.parametrize("name", ["rich", "stripped"])
def test_every_template_yields_a_complete_design_system(name, request):
    system = request.getfixturevalue(name)

    assert system is not None
    assert system.complete, "a template with no design system cannot be composed on"
    assert system.pairs
    assert system.type_scale
    assert system.spacing.steps
    assert system.grid is not None
    assert system.anchors is not None


@pytest.mark.parametrize("name", ["rich", "stripped"])
def test_the_type_scale_covers_every_text_role(name, request):
    system = request.getfixturevalue(name)

    roles = {token.role for token in system.type_scale}
    assert roles == {
        "title",
        "subtitle",
        "heading",
        "body",
        "caption",
        "number",
        "label",
    }
    # A figure is large because it is a figure: that is where hierarchy comes
    # from now that the plan cannot ask for a size.
    number = system.type_for("number")
    body = system.type_for("body")
    assert number is not None and body is not None
    assert number.size_pt > body.size_pt


@pytest.mark.parametrize("name", ["rich", "stripped"])
def test_the_usable_area_keeps_a_margin(name, request):
    system = request.getfixturevalue(name)
    area = system.grid.area

    # One full-bleed decorative band must not erase the margins every other
    # slide respects, so the area is never the whole slide.
    assert area.x > 0
    assert area.x + area.w < 1000
    assert area.w > 200


@pytest.mark.parametrize("name", ["rich", "stripped"])
def test_the_spacing_is_a_scale_not_one_unit(name, request):
    system = request.getfixturevalue(name)

    assert len(system.spacing.steps) >= 3
    assert system.spacing.steps == sorted(system.spacing.steps)
    assert system.spacing.step(0) < system.spacing.step(4)
    # Snapping a measured gap is how a component stays on the rhythm.
    assert system.spacing.nearest(system.spacing.step(2) + 1) == system.spacing.step(2)


# -- colour pairs and contrast (D12) ------------------------------------------


@pytest.mark.parametrize("name", ["rich", "stripped"])
def test_colours_come_in_usable_pairs(name, request):
    system = request.getfixturevalue(name)

    assert system.pairs
    for pair in system.pairs:
        if pair.usable:
            assert pair.contrast >= MIN_CONTRAST, (
                f"{pair.name} is offered as a fill but its ink is unreadable on it"
            )


def test_a_surface_with_no_readable_ink_is_not_offered():
    # Every candidate ink is as dark as the surface: nothing can be read on it.
    surface = ColorToken(ref="accent1", rgb="303030")
    inks = [ColorToken(ref="dk1", rgb="2A2A2A"), ColorToken(ref="dk2", rgb="383838")]

    pair = best_pair("accent", surface, inks)

    assert not pair.usable
    # The fallback never invents: the ink stays one the template contains.
    assert pair.on_surface.rgb in {"2A2A2A", "383838"}


def test_the_most_readable_ink_wins():
    surface = ColorToken(ref="lt1", rgb="FFFFFF")
    inks = [ColorToken(ref="lt2", rgb="EEEEEE"), ColorToken(ref="dk1", rgb="101010")]

    pair = best_pair("surface", surface, inks)

    assert pair.on_surface.rgb == "101010"
    assert pair.usable


def test_contrast_ratio_matches_the_wcag_bounds():
    assert contrast_ratio("FFFFFF", "000000") == pytest.approx(21.0, abs=0.01)
    assert contrast_ratio("FFFFFF", "FFFFFF") == pytest.approx(1.0, abs=0.01)
    assert relative_luminance("FFFFFF") > relative_luminance("000000")
    # A malformed value must not raise here: extraction is best effort.
    assert contrast_ratio("nope", "000000") > 0


def test_a_template_misusing_its_theme_slots_still_yields_a_readable_pair():
    """Slot names lie; what the master actually paints with does not."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Dark on light"

    system = analyze_presentation(presentation).design_system

    surface = system.usable_pair("surface")
    assert surface is not None
    assert surface.contrast >= MIN_CONTRAST


# -- theme references (D11) ---------------------------------------------------


def test_theme_slots_are_read_as_references():
    presentation = Presentation()

    slots = theme_colors(presentation)

    assert "lt1" in slots and "accent1" in slots
    assert all(len(value) == 6 for value in slots.values())


@pytest.mark.parametrize("name", ["rich", "stripped"])
def test_pairs_carry_the_slot_they_came_from(name, request):
    system = request.getfixturevalue(name)

    referenced = [p for p in system.pairs if p.surface.referenced]
    assert referenced, "a themed template must give something to point at"
    for pair in referenced:
        assert pair.surface.ref in theme_colors(Presentation()) or pair.surface.ref


def test_a_pair_that_would_write_a_literal_is_recorded():
    """D11: a literal is allowed, but never silently — the author is told."""
    presentation = Presentation()

    system = analyze_presentation(presentation).design_system

    # Whatever the template gave us, the two statements below must agree: a
    # pair writes a literal if and only if it is named in `unreferenced`.
    writes_literal = {
        pair.name
        for pair in system.pairs
        if not pair.surface.referenced or not pair.on_surface.referenced
    }
    assert writes_literal == set(system.unreferenced)


def test_a_fully_themed_template_reports_nothing_unreferenced():
    _, inventory = default_theme()

    system = inventory.design_system

    assert system.unreferenced == []
    assert all(pair.surface.referenced for pair in system.pairs)


# -- optical anchors (D13) ----------------------------------------------------


@pytest.mark.parametrize("name", ["rich", "stripped"])
def test_anchors_are_read_and_ordered(name, request):
    system = request.getfixturevalue(name)
    anchors = system.anchors

    assert anchors is not None
    assert 0 <= anchors.title_y < anchors.content_y <= 1000
    assert anchors.source in ("slides", "master")


def test_a_template_with_its_own_slides_takes_their_convention():
    """The slides win over the master: that is what the reader recognises."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Title"

    system = analyze_presentation(presentation).design_system

    assert system.anchors.source == "slides"


def test_a_template_with_no_slide_falls_back_to_its_master():
    presentation = Presentation()

    system = analyze_presentation(presentation).design_system

    assert system.anchors.source == "master"


# -- the grid follows the real slide ------------------------------------------


def test_the_grid_follows_a_non_16_9_template():
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Four by three"

    system = analyze_presentation(presentation).design_system

    assert system.slide_width == Inches(10)
    assert system.slide_height == Inches(7.5)
    # A gap of one step is optically equal in both directions, which on a 4:3
    # slide means different per-mille values.
    assert system.gap_x(2) != system.gap_y(2)


def test_a_column_span_stays_inside_the_usable_area():
    _, inventory = default_theme()
    system = inventory.design_system
    grid = system.grid

    x, w = grid.span(0, grid.columns, shorter_to_x=system.shorter_to_x)

    assert x >= grid.area.x
    assert x + w <= grid.area.x + grid.area.w + 1


def test_column_spans_do_not_overlap():
    _, inventory = default_theme()
    system = inventory.design_system
    grid = system.grid

    left = grid.span(0, 4, shorter_to_x=system.shorter_to_x)
    right = grid.span(4, 4, shorter_to_x=system.shorter_to_x)

    assert left[0] + left[1] <= right[0]


# -- the template's own fixtures ----------------------------------------------
#
# A logo on a master or a layout is not copied onto each slide: PowerPoint
# INHERITS it. So a composed slide already shows it — and must keep off it.


def _with_logo(*, suppress: bool = False) -> Presentation:
    """The common corporate shape: a logo on the master, bottom right."""
    presentation = Presentation()
    add_fixture(
        presentation,
        presentation.slide_masters[0],
        Inches(8.2),
        Inches(6.6),
        Inches(1.4),
        Inches(0.6),
    )
    if suppress:
        presentation.slide_layouts[6]._element.set("showMasterSp", "0")
    presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = "T"
    return presentation


def test_a_logo_on_the_master_is_reserved_and_kept_clear():
    system = analyze_presentation(_with_logo()).design_system

    assert system.reserved, "the master's logo was not seen"
    logo = system.reserved[0]
    area = system.grid.area
    assert not area.overlaps(logo), "the grid would put content over the logo"
    # Bottom right: stopping above it costs less room than giving up the width.
    assert area.y + area.h <= logo.y
    assert area.w > 500


def test_a_layout_that_switches_the_master_off_inherits_no_logo():
    """Then there is nothing to keep clear, and the whole area is usable."""
    system = analyze_presentation(_with_logo(suppress=True)).design_system

    assert system.reserved == []
    assert system.grid.area.h > 900


def test_the_background_is_not_a_fixture():
    """Reserving the full-bleed rectangle would leave nothing to compose on."""
    presentation = Presentation()
    add_fixture(
        presentation,
        presentation.slide_masters[0],
        0,
        0,
        presentation.slide_width,
        presentation.slide_height,
    )
    presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = "T"

    system = analyze_presentation(presentation).design_system

    assert system.reserved == []
    assert system.complete


def test_the_canvas_is_the_layout_that_carries_the_template_furniture():
    """Not the emptiest one: a bare canvas throws the charter away.

    The layout's own shapes come with it, so composing on the dressed layout is
    what makes a composed slide look like the rest of the deck.
    """
    presentation = _with_logo()
    dressed = presentation.slide_layouts[5]
    for placeholder in list(dressed.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    add_fixture(presentation, dressed, 0, 0, presentation.slide_width, Inches(0.4))

    inventory = analyze_presentation(presentation)

    assert inventory.canvas_layout == 5, (
        "the bare layout was preferred to the dressed one"
    )
    assert len(inventory.design_system.reserved) == 2  # the band and the logo
    canvas = next(
        p for p in inventory.patterns if p.source == "layout" and p.source_index == 5
    )
    assert canvas.exclusion_reason == "canvas"


def test_a_layout_whose_furniture_crowds_the_slide_is_not_the_canvas():
    """A block across the middle is a section divider, not a canvas.

    Nothing can be composed round it without splitting the page in two, so the
    barer layout is the right answer even though it carries less of the charter.
    """
    presentation = _with_logo()
    crowded = presentation.slide_layouts[5]
    for placeholder in list(crowded.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    add_fixture(
        presentation, crowded, 0, Inches(2.5), presentation.slide_width, Inches(2.5)
    )

    inventory = analyze_presentation(presentation)

    assert inventory.canvas_layout == 6
    assert inventory.design_system.complete


# -- trimming the area (pure arithmetic) --------------------------------------


def test_free_area_leaves_an_area_nothing_touches_alone():
    area = Box(x=50, y=40, w=900, h=900)

    assert free_area(area, [Box(x=0, y=960, w=200, h=40)]) == area


def test_free_area_trims_the_side_that_costs_least():
    area = Box(x=50, y=40, w=900, h=900)

    # A thin footer band: stopping above it keeps the full width.
    trimmed = free_area(area, [Box(x=0, y=900, w=1000, h=100)])

    assert trimmed.x == 50 and trimmed.w == 900
    assert trimmed.y + trimmed.h == 900


def test_free_area_trims_a_side_column_off_the_width():
    area = Box(x=50, y=40, w=900, h=900)

    trimmed = free_area(area, [Box(x=0, y=0, w=250, h=1000)])

    assert trimmed.x == 250
    assert trimmed.w == 700
    assert trimmed.h == 900


# -- the component style the template demonstrates (D18) ----------------------
#
# A template teaches more than its palette: whether it paints its blocks or
# outlines them, how round they are, how much air they hold and how it stacks a
# figure and its label. Composing in its language means reading that back.


def test_the_house_style_is_read_off_the_blocks_the_template_draws():
    system = analyze_presentation(build_house_style_template()).design_system
    style = system.shapes

    assert style.corner_radius == pytest.approx(HOUSE_CORNER, abs=0.01)
    assert style.filled is False, "outlined cards were read as painted ones"
    assert style.stroked is True
    assert style.figure_above_label is False
    # Gaps and insets are per-mille of the SHORTER side, like the spacing scale.
    shorter = min(system.slide_width, system.slide_height)
    assert style.column_gap == pytest.approx(
        Inches(HOUSE_GAP_IN) * 1000 / shorter, abs=2
    )
    assert style.padding == pytest.approx(
        Inches(HOUSE_INSET_IN) * 1000 / shorter, abs=2
    )


def test_the_grid_takes_its_gutter_from_the_gap_the_template_uses():
    system = analyze_presentation(build_house_style_template()).design_system

    assert system.grid.gutter == system.shapes.column_gap
    # And it is the template's, not a step of the scale that happens to be near.
    assert system.grid.gutter not in system.spacing.steps


def test_a_template_that_shows_none_of_it_keeps_the_defaults():
    """Every field has to have a defensible answer for a bare template."""
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = "T"

    style = analyze_presentation(presentation).design_system.shapes

    assert style.filled is True
    assert style.figure_above_label is True
    assert style.column_gap is None and style.padding is None


def test_a_bare_text_box_does_not_vote_on_how_blocks_are_painted():
    """It paints nothing, so counting it would read every plain template as one
    that outlines rather than fills."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Words only"
    for index in range(3):
        box = slide.shapes.add_textbox(
            Inches(0.5 + index * 3.0), Inches(3.0), Inches(2.6), Inches(1.0)
        )
        box.text_frame.text = f"Column {index}"

    style = analyze_presentation(presentation).design_system.shapes

    assert style.filled is True
    # The gap between them is still a gap, and still worth reading.
    assert style.column_gap is not None


def test_a_cover_only_template_does_not_read_its_content_band_off_its_cover():
    """A cover's title sits in a narrow column BY DESIGN.

    Reading the content band off it gave a band 41% of the slide wide, in
    which a card held eight characters — the deck of title slides.
    """
    system = analyze_presentation(build_cover_only_template()).design_system
    area = system.grid.area

    assert area.w >= 600, "the content band was read off a cover again"
    # Nothing in this template says where content goes, so nothing is invented:
    # the anchors fall back to the master rather than to the cover's ordinates.
    assert system.anchors.source == "master"
    assert system.complete, "the template must still be composable"


def test_a_content_band_measured_too_narrow_is_refused_rather_than_used():
    """The backstop under the rule above: a template whose one content layout
    IS a narrow column still has to yield a band a slide fits in."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    master = presentation.slide_masters[0]
    keep = {presentation.slide_layouts[1].name, presentation.slide_layouts[6].name}
    id_list = master._element.get_or_add_sldLayoutIdLst()
    for entry, layout in zip(list(id_list), list(master.slide_layouts)):
        if layout.name not in keep:
            id_list.remove(entry)
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    for shape in (slide.shapes.title, slide.placeholders[1]):
        shape.left, shape.width = Inches(9.0), Inches(3.5)
    slide.shapes.title.text = "Narrow"
    slide.placeholders[1].text_frame.text = "one line"

    area = analyze_presentation(presentation).design_system.grid.area

    assert area.w >= 600, "a 26%-wide column was taken for the content band"
