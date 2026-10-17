# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The deck assigner: kind inference, the fit check, the deck rules and the
problems it reports instead of substituting a slide."""

from __future__ import annotations

from typing import Any, List

from builder_decks import EVIDENCE_SLIDES, build_showcase_template, with_item_slides
from deck_builders import png_bytes
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import (
    _advice,
    assign,
    content_shape,
    infer_kinds,
    misfit,
)
from fred_capability_ppt_filler.builder.models import (
    Box,
    Pattern,
    PatternOverride,
    RepeatGroup,
    TemplateInventory,
    Zone,
)
from fred_capability_ppt_filler.builder.overrides import apply_overrides
from fred_capability_ppt_filler.builder.plan import (
    DeckPlan,
    PlanColumn,
    PlanItem,
    SlidePlan,
    capacity_text,
    zone_summary,
)
from fred_capability_ppt_filler.builder.render import _assign_text
from fred_capability_ppt_filler.builder.tools import _assign_loadable
from pptx.util import Inches

_analysed = analyze_template(build_showcase_template()).inventory
assert _analysed is not None
INVENTORY: TemplateInventory = _analysed
COVER = SlidePlan(title="Kick-off", subtitle="Team")


def _pattern(label: str) -> Pattern:
    return next(p for p in INVENTORY.patterns if p.label == label)


def _plan(*slides: Any) -> DeckPlan:
    return DeckPlan(
        slides=[
            s if isinstance(s, SlidePlan) else SlidePlan.model_validate(s)
            for s in slides
        ]
    )


def _labels(inventory: TemplateInventory, *slides: Any) -> List[str]:
    assignment = assign(inventory, _plan(*slides))
    assert not assignment.problems, [p.as_dict() for p in assignment.problems]
    return [p.label if p else "-" for p in assignment.patterns]


def _title_zone() -> Zone:
    return Zone(
        key="ph:0",
        role="title",
        box=Box(x=0, y=0, w=900, h=100),
        font_pt=32.0,
        chars_per_line=60,
        max_lines=1,
    )


def _body_zone(key: str = "ph:1", chars: int = 300) -> Zone:
    return Zone(
        key=key,
        role="body",
        box=Box(x=0, y=200, w=900, h=600),
        font_pt=18.0,
        chars_per_line=chars // 8,
        max_lines=8,
    )


def _inventory(*patterns: Pattern) -> TemplateInventory:
    return TemplateInventory(
        slide_width=Inches(13.333).emu,
        slide_height=Inches(7.5).emu,
        patterns=list(patterns),
    )


def _bullets(pid: str, label: str, source: str = "layout") -> Pattern:
    return Pattern(
        id=pid,
        source=source,  # type: ignore[arg-type]
        source_index=int(pid.split(":")[1]),
        name=label,
        label=label,
        kind="bullets",
        zones=[_title_zone(), _body_zone()],
    )


# --- kind inference -------------------------------------------------------------------------


def test_kind_is_inferred_from_the_shape_of_the_content():
    def first(**fields: Any) -> str:
        return infer_kinds(SlidePlan(**fields))[0]

    assert (
        first(title="KPIs", items=[PlanItem(number="42%", label="Margin")]) == "figures"
    )
    assert (
        first(title="Plan", items=[PlanItem(number="1", heading="Go", text="x")])
        == "process"
    )
    assert (
        first(title="History", items=[PlanItem(number="2024", heading="x")])
        == "timeline"
    )
    assert (
        first(title="Topics", items=[PlanItem(heading="A"), PlanItem(heading="B")])
        == "agenda"
    )
    assert (
        first(title="Options", columns=[PlanColumn(bullets=["a"])] * 2) == "two_columns"
    )
    assert (
        first(
            title="A vs B", columns=[PlanColumn(heading="A"), PlanColumn(heading="B")]
        )
        == "comparison"
    )
    assert first(title="Site", image="doc-1") == "image_text"
    assert first(title="Points", bullets=["a"]) == "bullets"
    assert first(title="Part one") == "section"
    assert first(title="Merci") == "closing"
    # A title naming a kind breaks the tie when the content suits it.
    assert first(title="Sommaire", bullets=["a"]) == "bullets"


# --- fit check: the pattern is chosen for the content it is given ---------------------------


def _cards_inventory() -> TemplateInventory:
    """Two cards patterns: the first shows only a description, the second a
    heading too - the shape of the real corporate template."""

    def group(roles, key) -> RepeatGroup:
        return RepeatGroup(
            key=key,
            count=3,
            item_zones=roles,
            item_keys=[f"{key}/{i}" for i in range(3)],
            items=[
                [
                    Zone(
                        key=f"{key}/{i}/{role}",
                        role=role,
                        box=Box(x=i * 300, y=300, w=280, h=200),
                        font_pt=14.0,
                        chars_per_line=26,
                        max_lines=3,
                    )
                    for role in roles
                ]
                for i in range(3)
            ],
        )

    def cards(pid: str, label: str, roles: list) -> Pattern:
        return Pattern(
            id=pid,
            source="slide",
            source_index=int(pid.split(":")[1]),
            name=label,
            label=label,
            kind="cards",
            zones=[_title_zone()],
            groups=[group(roles, f"g{pid}")],
        )

    return _inventory(
        cards("slide:1", "Text cards", ["body"]),
        cards("slide:2", "Titled cards", ["heading", "body"]),
    )


def test_cards_with_headings_skip_the_pattern_that_cannot_show_them():
    slide = SlidePlan(
        title="Three axes",
        kind="cards",
        items=[PlanItem(heading="Near", text="Close to clients")] * 3,
    )
    assert _labels(_cards_inventory(), slide) == ["Titled cards"]


def test_a_pattern_without_headings_still_wins_when_the_plan_has_none():
    slide = SlidePlan(
        title="Three axes", kind="cards", items=[PlanItem(text="Close")] * 3
    )
    assert _labels(_cards_inventory(), slide) == ["Text cards"]


def test_an_image_pattern_is_not_chosen_for_a_plan_without_images():
    """Every unfilled picture zone is a hole on the finished slide."""
    plain = _bullets("layout:1", "Plain")
    imaged = plain.model_copy(
        update={
            "id": "slide:99",
            "source": "slide",
            "label": "Photos",
            "zones": plain.zones
            + [
                Zone(
                    key=f"sp:{i}", role="image", box=Box(x=i * 300, y=400, w=280, h=300)
                )
                for i in range(4)
            ],
        }
    )
    inventory = _inventory(imaged, plain)
    slide = SlidePlan(
        title="Method", bullets=["Frame", "Build", "Measure"], kind="bullets"
    )
    assert _labels(inventory, slide) == ["Plain"]
    assert _labels(inventory, slide.model_copy(update={"image": "doc-1"})) == ["Photos"]


def test_a_pattern_that_would_truncate_is_not_a_candidate():
    plain = _bullets("layout:1", "Plain")
    narrow = _bullets("slide:98", "Narrow", source="slide")
    narrow.zones[1].chars_per_line = 3
    slide = SlidePlan(title="Detail", bullets=["A sentence far longer than twenty"] * 4)
    # The example slide would win on preference, but it would cut the text.
    assert _labels(_inventory(narrow, plain), slide) == ["Plain"]


def test_an_example_slide_is_preferred_over_a_layout_that_fits_as_well():
    slide = SlidePlan(title="Points", bullets=["a", "b"], kind="bullets")
    layout = _bullets("layout:1", "Layout")
    example = _bullets("slide:7", "Example", source="slide")
    assert _labels(_inventory(layout, example), slide) == ["Example"]


# --- deck rules -------------------------------------------------------------------------------


def test_the_cover_is_used_once_and_only_first():
    labels = _labels(
        INVENTORY,
        COVER,
        {"title": "Part one", "subtitle": "Why"},
        {"title": "Points", "bullets": ["a"]},
    )
    assert labels[0] == "Title Slide"
    assert "Title Slide" not in labels[1:]


def test_the_first_slide_must_fit_the_cover():
    assignment = assign(
        INVENTORY,
        _plan(
            {"title": "Topics", "items": [{"heading": "A"}, {"heading": "B"}]}, COVER
        ),
    )
    (problem,) = assignment.problems
    assert problem.index == 1 and assignment.patterns == []
    assert "always the template's cover" in problem.problem
    assert any("add a cover slide" in o for o in problem.options)


def test_a_later_cover_is_built_from_its_content_and_the_swap_is_named():
    assignment = assign(INVENTORY, _plan(COVER, {"title": "Again", "kind": "cover"}))
    assert not assignment.problems
    assert [p.kind for p in assignment.patterns if p] == ["cover", "section"]
    assert assignment.notes == [
        "Slide 2: built as 'section' instead of 'cover': only the first slide can be "
        "the cover. Nothing to do."
    ]


def _with_closing() -> TemplateInventory:
    closing = Pattern(
        id="slide:90",
        source="slide",
        source_index=90,
        name="Merci",
        label="Merci",
        kind="closing",
    )
    return INVENTORY.model_copy(update={"patterns": INVENTORY.patterns + [closing]})


def test_closing_is_used_at_most_once_and_only_last():
    inventory = _with_closing()
    labels = _labels(
        inventory, COVER, {"title": "Points", "bullets": ["a"]}, {"title": "Merci"}
    )
    assert labels[-1] == "Merci"
    # A closing word in the middle of the deck is never the closing slide.
    labels = _labels(
        inventory, COVER, {"title": "Questions"}, {"title": "Points", "bullets": ["a"]}
    )
    assert "Merci" not in labels

    assignment = assign(
        inventory,
        _plan(
            COVER,
            {"kind": "closing", "title": "Merci"},
            {"title": "Points", "bullets": ["a"]},
        ),
    )
    assert not assignment.problems
    assert assignment.patterns[1] and assignment.patterns[1].label != "Merci"
    assert assignment.notes[0].endswith(
        "a closing slide can only end the deck. To make it the closing slide, move it "
        "to the end; otherwise nothing to do."
    )


def test_a_fixed_closing_slide_takes_no_content_so_a_close_kind_holds_it():
    assignment = assign(
        _with_closing(),
        _plan(
            COVER, {"kind": "closing", "title": "Contact", "bullets": ["jane@x.com"]}
        ),
    )
    assert not assignment.problems
    assert assignment.patterns[1] and assignment.patterns[1].kind == "section"
    assert assignment.notes == [
        "Slide 2: built as 'section' instead of 'closing': the 'closing' slide does not "
        "fit (this slide is shown as is and takes no content). To get the 'closing' "
        "slide, call again and send only the title 'Merci'; otherwise nothing to do."
    ]


def test_no_three_identical_slides_in_a_row_when_an_alternative_fits():
    inventory = _inventory(_bullets("layout:1", "A"), _bullets("layout:2", "B"))
    slides = [SlidePlan(title=f"P{i}", bullets=["x"]) for i in range(5)]
    labels = _labels(inventory, *slides)
    assert all(not (labels[i] == labels[i + 1] == labels[i + 2]) for i in range(3)), (
        labels
    )


def test_repetition_is_allowed_and_reported_on_a_one_archetype_template():
    inventory = _inventory(_bullets("layout:1", "Only"))
    assignment = assign(
        inventory, _plan(*[SlidePlan(title=f"P{i}", bullets=["x"]) for i in range(4)])
    )
    assert not assignment.problems
    assert [p.label for p in assignment.patterns if p] == ["Only"] * 4
    (note,) = assignment.notes
    assert note.startswith("Slides 1, 2, 3, 4 use the template slide 'Only' in a row.")
    assert "Nothing to do: no other slide of this template suits them" in note
    # With the check off, the passive note of three in a row is kept.
    off = assign(
        inventory,
        _plan(*[SlidePlan(title=f"P{i}", bullets=["x"]) for i in range(4)]),
        "off",
    )
    assert "Slides 1 to 3" in off.notes[0]


def test_invalid_entries_get_no_pattern_and_do_not_shift_the_rules():
    plan = DeckPlan(
        slides=[
            SlidePlan.model_validate("junk"),
            COVER,
            SlidePlan(title="P", bullets=["x"]),
        ]
    )
    assignment = assign(INVENTORY, plan)
    assert not assignment.problems
    assert assignment.patterns[0] is None
    assert assignment.patterns[1] is not None and assignment.patterns[1].kind == "cover"


# --- no silent fallback -----------------------------------------------------------------------


def test_items_beyond_two_slides_are_a_problem_with_trim_options():
    steps = [PlanItem(number=str(i), heading=f"S{i}", text="x") for i in range(7)]
    assignment = assign(
        INVENTORY, _plan(COVER, SlidePlan(title="How", kind="process", items=steps))
    )
    (problem,) = assignment.problems
    assert problem.as_dict()["slide"] == 2
    # The team slide holds more, but it is not a close kind of 'process'.
    assert problem.problem == (
        "items: 7 given; the largest 'process' slide holds 3, 2 slides hold 6"
    )
    assert problem.options[:2] == [
        "trim to 6 items (beyond 3 they continue on a second slide)",
        "or keep 3 here and move the rest to slides of your own",
    ]


def test_a_kind_the_template_lacks_is_replaced_by_a_close_kind():
    slide = SlidePlan(title="Points", bullets=["a", "b"], kind="timeline")
    assignment = assign(INVENTORY, _plan(COVER, slide))
    assert not assignment.problems
    assert assignment.patterns[1] and assignment.patterns[1].kind == "bullets"
    assert assignment.notes == [
        "Slide 2: built as 'bullets' instead of 'timeline': this template has no "
        "'timeline' slide. Nothing to do."
    ]


def _items_pattern(
    pid: str, kind: Any, count: int, roles: List[Any], sample: str = ""
) -> Pattern:
    """An example slide of ``count`` items, each with one zone per role."""
    group = RepeatGroup(
        key=f"g{pid}",
        count=count,
        item_zones=roles,
        item_keys=[f"g{pid}/{i}" for i in range(count)],
        items=[
            [
                Zone(
                    key=f"g{pid}/{i}/{role}",
                    role=role,
                    box=Box(x=i * 100, y=300, w=90, h=200),
                    font_pt=14.0,
                    chars_per_line=20,
                    max_lines=2,
                    sample=sample if role == "number" else "",
                )
                for role in roles
            ]
            for i in range(count)
        ],
    )
    return Pattern(
        id=pid,
        source="slide",
        source_index=int(pid.split(":")[1]),
        name=kind,
        label=kind,
        kind=kind,
        zones=[_title_zone()],
        groups=[group],
    )


def _sommaire(n: int, numbered: bool = True) -> SlidePlan:
    return SlidePlan(
        title="Sommaire",
        kind="agenda",
        items=[
            PlanItem(number=f"{i:02d}" if numbered else None, heading=f"Topic {i}")
            for i in range(1, n + 1)
        ],
    )


_SOMMAIRE = _sommaire(17)


def test_an_agenda_too_long_is_never_built_on_an_unrelated_kind():
    """Live Paintress case, beyond what a continuation slide adds: 17 agenda
    items, an agenda slide of 8 and a team slide of 30."""
    inventory = _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "agenda", 8, ["number", "heading"]),
        _items_pattern("slide:3", "team", 30, ["heading", "label"]),
    )
    assignment = assign(
        inventory, _plan(SlidePlan(title="Intro", bullets=["a"]), _SOMMAIRE)
    )
    assert assignment.patterns == []
    (problem,) = assignment.problems
    assert problem.index == 2
    assert problem.problem == (
        "items: 17 given; the largest 'agenda' slide holds 8, 2 slides hold 16"
    )
    assert problem.options[:2] == [
        "trim to 16 items (beyond 8 they continue on a second slide)",
        "or keep 8 here and move the rest to slides of your own",
    ]
    assert any(o.startswith("or rewrite it as kind 'bullets'") for o in problem.options)
    assert not any("team" in o for o in problem.options)


def test_an_agenda_too_long_goes_to_a_close_kind_that_holds_it():
    inventory = _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "agenda", 8, ["number", "heading"]),
        _items_pattern("slide:3", "team", 30, ["heading", "label"]),
        _items_pattern("slide:4", "process", 20, ["number", "heading"]),
    )
    assignment = assign(
        inventory, _plan(SlidePlan(title="Intro", bullets=["a"]), _SOMMAIRE)
    )
    assert not assignment.problems and not assignment.parts
    assert assignment.patterns[1] and assignment.patterns[1].kind == "process"
    assert assignment.notes == [
        "Slide 2: built as 'process' instead of 'agenda': the 'agenda' slide does not "
        "fit (items: 17 given, max 8). To get the 'agenda' slide, call again and trim "
        "to 8 items or split them over two slides; otherwise nothing to do."
    ]


# --- a list continued on a second slide -------------------------------------------------------


def _agenda_inventory(*extra: Pattern) -> TemplateInventory:
    return _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "agenda", 8, ["number", "heading"], sample="01"),
        _items_pattern("slide:5", "agenda", 6, ["number", "heading"], sample="01"),
        _items_pattern("slide:3", "team", 15, ["heading", "label"]),
        _items_pattern("slide:4", "process", 10, ["number", "heading"]),
        *extra,
    )


def test_an_agenda_one_entry_too_long_continues_on_a_second_agenda_slide():
    """Live Paintress case: 9 entries, an agenda slide of 8. A same-kind split
    beats the process slide that would hold all nine."""
    plan = _plan(SlidePlan(title="Intro", bullets=["a"]), _sommaire(9), _points(2))
    assignment = assign(_agenda_inventory(), plan)
    assert not assignment.problems
    deck, patterns, numbers = assignment.deck(plan)
    assert [s.title for s in deck.slides] == [
        "Intro",
        "Sommaire (1/2)",
        "Sommaire (2/2)",
        "Points 2",
    ]
    assert [len(s.items) for s in deck.slides[1:3]] == [5, 4]
    # The same template slide for both parts, the numbers as sent.
    first, second = patterns[1:3]
    assert first and second and first.kind == "agenda" and first.id == second.id
    assert [i.number for i in deck.slides[2].items] == ["06", "07", "08", "09"]
    assert numbers == [1, 2, 2, 3]
    assert assignment.notes == [
        "Slide 2: its 9 items continue on a second 'agenda' slide (5 + 4, titles marked "
        "(1/2) and (2/2)): the largest 'agenda' slide holds 8. Nothing to do; to keep "
        "one slide, send at most 8 items."
    ]
    # The model's plan itself is untouched.
    assert plan.slides[1].title == "Sommaire" and len(plan.slides[1].items) == 9


def test_items_without_numbers_continue_the_count_on_the_second_slide():
    """Numbered at render from ``starts``: the stored parts carry no number of
    their own, so a later change to part 1 renumbers part 2."""
    plan = _plan(SlidePlan(title="Intro", bullets=["a"]), _sommaire(9, numbered=False))
    assignment = assign(_agenda_inventory(), plan)
    deck, _, _ = assignment.deck(plan)
    assert [i.number for i in deck.slides[2].items] == [None] * 4
    assert assignment.starts == [0, 0, 5]
    # Rebuilt from the stored parts: the continuation still starts after part 1.
    shorter = deck.model_copy(
        update={
            "slides": [
                deck.slides[0],
                deck.slides[1].model_copy(update={"items": deck.slides[1].items[1:]}),
                deck.slides[2],
            ]
        }
    )
    assert assign(_agenda_inventory(), shorter).starts == [0, 0, 4]


def test_steps_numbered_from_six_need_no_number_zone_on_a_continuation():
    team = _items_pattern("slide:3", "team", 5, ["heading", "label"])
    items = [PlanItem(number=str(i), heading=f"S{i}") for i in range(6, 10)]
    assert not misfit(team, SlidePlan(title="Who", items=items), (1, 1))[0]
    years = [PlanItem(number=str(y), heading="x") for y in range(2024, 2027)]
    assert misfit(team, SlidePlan(title="When", items=years), (1, 1))[0]


def test_a_list_without_a_requested_kind_splits_on_its_first_inferred_kind():
    agenda = _sommaire(9).model_copy(update={"kind": None})
    plan = _plan(SlidePlan(title="Intro", bullets=["a"]), agenda)
    inventory = _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "agenda", 8, ["number", "heading"]),
    )
    assignment = assign(inventory, plan)
    assert not assignment.problems
    assert [p.kind for _, p in assignment.parts[2]] == ["agenda", "agenda"]


def test_a_list_whose_kind_holds_it_whole_is_never_split():
    plan = _plan(SlidePlan(title="Intro", bullets=["a"]), _sommaire(8))
    assignment = assign(_agenda_inventory(), plan)
    assert not assignment.parts and not assignment.notes


def test_a_text_too_long_is_never_split():
    long = _sommaire(9)
    long.items[0] = PlanItem(number="01", heading="word " * 30)
    inventory = _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "agenda", 8, ["number", "heading"]),
    )
    assignment = assign(inventory, _plan(SlidePlan(title="Intro", bullets=["a"]), long))
    assert not assignment.parts
    (problem,) = assignment.problems
    assert problem.problem.startswith(
        "items: 9 given; the largest 'agenda' slide holds 8"
    )
    assert "items[0].heading" in problem.problem


def test_bullets_are_never_split():
    slide = SlidePlan(title="Many", bullets=["word " * 12] * 30)
    assignment = assign(_inventory(_bullets("layout:1", "Plain")), _plan(slide))
    assert not assignment.parts and assignment.problems


def test_a_split_list_alone_never_asks_for_variety():
    """Both parts on one template slide count as one slide, even with max 1."""
    plan = _plan(SlidePlan(title="Intro", bullets=["a"]), _sommaire(9))
    agenda = _items_pattern("slide:2", "agenda", 8, ["number", "heading"])
    inventory = _inventory(_bullets("layout:1", "Text"), agenda)
    for variety in ("template_slide", "kind"):
        assignment = assign(inventory, plan, variety, 1)  # type: ignore[arg-type]
        assert not assignment.repeats
        assert len(assignment.notes) == 1  # the split note only


def test_a_continuation_stays_on_its_first_parts_slide_even_at_max_one():
    inventory = _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "cards", 4, ["heading", "body"]),
        _items_pattern("slide:3", "cards", 4, ["heading", "body"]),
    )
    items = [PlanItem(heading=f"H{i}", text="t") for i in range(7)]
    plan = _plan(SlidePlan(title="Offers", kind="cards", items=items))
    assignment = assign(inventory, plan, "template_slide", 1)
    assert not assignment.problems and not assignment.repeats
    first, second = (p for _, p in assignment.parts[1])
    assert first.id == second.id


def _offers() -> TemplateInventory:
    return _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "cards", 4, ["heading", "body"]),
        _items_pattern("slide:3", "cards", 4, ["heading", "body"]),
    )


def test_a_delivered_split_rebuilds_as_it_was_and_still_counts_once():
    """The stored deck holds both parts as slides: rebuilt at max one, they are
    neither split again nor a repeat, and stay on one template slide."""
    items = [PlanItem(heading=f"H{i}", text="t") for i in range(7)]
    plan = _plan(
        SlidePlan(title="Intro", bullets=["a"]),
        SlidePlan(title="Offers", items=items),
    )
    deck, _, numbers = assign(_offers(), plan, "template_slide", 1).deck(plan)
    assert [s.title for s in deck.slides] == ["Intro", "Offers (1/2)", "Offers (2/2)"]
    assert numbers == [1, 2, 2] and deck.slides[2].kind == "cards"
    again = assign(_offers(), deck, "template_slide", 1)
    assert not again.problems and not again.repeats and not again.parts
    assert not [n for n in again.notes if "continue on" in n]
    first, second = again.patterns[1:]
    assert first and second and first.id == second.id


def test_a_delivered_slide_keeps_its_template_slide_while_it_fits_its_kind():
    items = [PlanItem(heading=f"H{i}", text="t") for i in range(3)]
    for built_on in ("slide:2", "slide:3"):
        slide = SlidePlan(title="Offers", items=items, built_on=built_on)
        (pattern,) = assign(_offers(), _plan(slide)).patterns
        assert pattern and pattern.id == built_on
    # Another kind asked, or content that no longer fits there: assigned anew.
    asked = SlidePlan(title="Offers", bullets=["a"], kind="bullets", built_on="slide:3")
    (pattern,) = assign(_offers(), _plan(asked)).patterns
    assert pattern and pattern.id == "layout:1"
    many = [PlanItem(heading=f"H{i}", text="t") for i in range(5)]
    split = SlidePlan(title="Offers", items=many, built_on="slide:3")
    assert assign(_offers(), _plan(split)).parts


def test_a_substituted_kind_keeps_its_template_slide():
    """A timeline built on a process slide (no timeline in the template) stays
    on that process slide when the deck is rebuilt for an unrelated change."""
    inventory = _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "process", 3, ["number", "heading"]),
        _items_pattern("slide:3", "process", 3, ["number", "heading"]),
    )
    items = [PlanItem(number=str(i), heading=f"Step {i}") for i in range(1, 4)]
    for built_on in ("slide:2", "slide:3"):
        slide = SlidePlan(title="Plan", kind="timeline", items=items, built_on=built_on)
        (pattern,) = assign(inventory, _plan(slide)).patterns
        assert pattern and pattern.id == built_on


def test_delivered_notes_name_the_files_slide_numbers():
    """A list split on slide 2 makes the next slide the file's slide 4."""
    items = [PlanItem(heading=f"H{i}", text="t") for i in range(7)]
    plan = _plan(
        SlidePlan(title="Intro", bullets=["a"]),
        SlidePlan(title="Offers", kind="cards", items=items),
        SlidePlan(title="Points", bullets=["a", "b"], kind="timeline"),
    )
    notes = assign(_offers(), plan).notes
    assert any(
        n.startswith("Slide 4: built as 'bullets' instead of 'timeline'") for n in notes
    )
    assert any(n.startswith("Slide 2: its 7 items continue on a second") for n in notes)


def test_a_split_list_still_counts_once_beside_other_slides_of_its_template_slide():
    cards = _items_pattern("slide:2", "cards", 4, ["heading", "body"])
    inventory = _inventory(_bullets("layout:1", "Text"), cards)

    def card(title: str, n: int) -> SlidePlan:
        items = [PlanItem(heading=f"H{i}", text="t") for i in range(n)]
        return SlidePlan(title=title, kind="cards", items=items)

    split = _plan(card("A", 6), card("B", 2))
    assignment = assign(inventory, split, "template_slide", 2)
    assert not assignment.problems and not assignment.repeats
    assert list(assignment.parts) == [1]
    # Three slides of their own on that template slide are a repeat (cap 2).
    three = _plan(card("A", 6), card("B", 2), card("C", 2), card("D", 2))
    (repeat,) = assign(inventory, three, "template_slide", 2).repeats
    assert repeat.slides == [1, 2, 3, 4]


def test_an_agenda_written_as_bullets_goes_to_a_bullets_slide():
    inventory = _inventory(
        _bullets("layout:1", "Text"),
        _items_pattern("slide:2", "agenda", 8, ["number", "heading"]),
        _items_pattern("slide:3", "team", 15, ["heading", "label"]),
    )
    agenda = SlidePlan(
        title="Sommaire", kind="agenda", bullets=[f"Topic {i}" for i in range(9)]
    )
    assignment = assign(
        inventory, _plan(SlidePlan(title="Intro", bullets=["a"]), agenda)
    )
    assert not assignment.problems
    assert assignment.patterns[1] and assignment.patterns[1].label == "Text"
    assert assignment.notes[0] == (
        "Slide 2: built as 'bullets' instead of 'agenda': the 'agenda' slide does not "
        "fit (bullets: this slide has no place for it). To get the 'agenda' slide, "
        "call again and give the bullets as items (at most 8); otherwise nothing to do."
    )


def test_a_requested_kind_that_fits_is_kept():
    slide = SlidePlan(title="Points", bullets=["a", "b"], kind="two_columns")
    assignment = assign(INVENTORY, _plan(COVER, slide))
    assert assignment.patterns[1] and assignment.patterns[1].kind == "two_columns"
    assert not assignment.notes


def test_a_text_too_long_everywhere_names_the_element_and_its_measured_limit():
    slide = SlidePlan(title="Detail", bullets=["short", "word " * 199 + "word"])
    (problem,) = assign(
        _inventory(_bullets("layout:1", "Plain")), _plan(slide)
    ).problems
    assert problem.problem == (
        "closest fit, a 'bullets' slide: bullets: 2 entries wrap to 23 lines; "
        "this zone holds 11 (longest: bullets[1], 999 chars)"
    )
    assert any(o.startswith("keep each bullets entry") for o in problem.options)


def test_a_forced_cover_misfit_reads_once_without_a_close_kind_lead():
    slide = SlidePlan(title="T" * 400, kind="bullets", bullets=["a"])
    (problem,) = assign(INVENTORY, _plan(slide)).problems
    assert problem.problem == (
        "the first slide is always the template's cover; closest fit, a 'cover' "
        "slide: title is 400 chars; this zone holds ≤8 words (≈29 chars × 2 lines)"
    )


def test_a_close_kind_misfit_names_why_and_the_close_kind_once():
    slide = SlidePlan(
        title="Detail", kind="timeline", bullets=["short", "word " * 199 + "word"]
    )
    (problem,) = assign(
        _inventory(_bullets("layout:1", "Plain")), _plan(slide)
    ).problems
    assert problem.problem == (
        "this template has no 'timeline' slide; closest close kind, a 'bullets' "
        "slide: bullets: 2 entries wrap to 23 lines; this zone holds 11 "
        "(longest: bullets[1], 999 chars)"
    )


def test_a_room_is_quoted_as_words_then_its_measured_chars():
    """One rule for every template: 7 chars per word, words never span lines."""
    assert capacity_text(12, 1) == "1 word (≈12 chars)"
    assert capacity_text(32, 1) == "≤4 words (≈32 chars)"
    assert capacity_text(71, 2) == "≤20 words (≈71 chars × 2 lines)"
    assert capacity_text(4, 1) == "1 short word (≈4 chars)"
    assert capacity_text(4, 3) == "1 short word (≈4 chars × 3 lines)"


def _list_pattern(kind: str, chars: int, lines: int, bodies: int = 1) -> Pattern:
    """A title and ``bodies`` side-by-side body zones of ``lines`` lines each."""
    zones = [_title_zone()] + [
        Zone(
            key=f"ph:{i + 1}",
            role="body",
            box=Box(x=i * 450, y=200, w=440, h=100 * lines),
            font_pt=18.0,
            chars_per_line=chars,
            max_lines=lines,
        )
        for i in range(bodies)
    ]
    return Pattern(
        id="layout:1",
        source="layout",
        source_index=1,
        name=kind,
        label=kind,
        kind=kind,  # type: ignore[arg-type]
        zones=zones,
    )


_SIZE = (Inches(13.333).emu, Inches(7.5).emu)


def test_too_many_entries_for_a_one_line_zone_name_the_entry_limit_not_chars():
    """Live case: 3 short bullets refused by a 1-line zone; quoting a char
    budget they already met sent the agent shortening instead of merging."""
    slide = SlidePlan(
        title="Jalons",
        kind="bullets",
        bullets=["Accord Cyber avec l'ACN", "Radars pour la Marine", "Pôle à Turin"],
    )
    (reason,), _ = misfit(_list_pattern("bullets", 101, 1), slide, _SIZE)
    assert reason.text == (
        "bullets: 3 entries; this zone holds 1 line (≈101 chars): at most 1 entry"
    )
    assert reason.option == "merge bullets into at most 1 entry, or use another kind"


def test_too_many_entries_for_a_two_line_zone():
    slide = SlidePlan(title="Jalons", kind="bullets", bullets=["a", "b", "c"])
    (reason,), _ = misfit(_list_pattern("bullets", 100, 2), slide, _SIZE)
    assert reason.text == (
        "bullets: 3 entries; this zone holds 2 lines (≈100 chars each): "
        "at most 2 entries"
    )
    assert reason.option == "merge bullets into at most 2 entries, or use another kind"


def test_too_many_column_entries_name_the_column_list():
    columns = [PlanColumn(bullets=["a", "b", "c", "d"]), PlanColumn(bullets=["e"])]
    slide = SlidePlan(title="Options", kind="two_columns", columns=columns)
    reasons, _ = misfit(_list_pattern("two_columns", 20, 1, bodies=2), slide, _SIZE)
    assert [r.text for r in reasons] == [
        "columns[0].bullets: 4 entries; this zone holds 1 line (≈20 chars): "
        "at most 1 entry"
    ]


def test_entries_few_enough_but_too_long_get_a_length_message():
    slide = SlidePlan(title="Detail", kind="bullets", bullets=["a", "word " * 60])
    (reason,), _ = misfit(_list_pattern("bullets", 40, 2), slide, _SIZE)
    assert reason.text.startswith("bullets: 2 entries wrap to ")
    assert reason.text.endswith(
        "lines; this zone holds 2 (longest: bullets[1], 299 chars)"
    )
    assert reason.option == (
        "keep each bullets entry to one line (≈40 chars), or drop entries"
    )


def test_a_list_field_room_quotes_its_lines_and_entry_limit():
    assert zone_summary(_list_pattern("bullets", 101, 1)) == (
        "title ≤8 words (≈60 chars); "
        "bullets ≤14 words (≈101 chars × 1 line, max 1 entry)"
    )
    assert "columns.bullets×2 ≤9 words (≈20 chars × 3 lines, max 4 entries)" in (
        zone_summary(_list_pattern("two_columns", 20, 3, bodies=2))
    )


def test_one_title_too_long_is_named_with_the_room_it_has():
    slide = SlidePlan(title="T" * 400, subtitle="Team")
    (problem,) = assign(INVENTORY, _plan(slide)).problems
    assert "title is 400 chars; this zone holds ≤8 words (≈29 chars × 2 lines)" in (
        problem.problem
    )
    assert "shorten title to ≤8 words (≈29 chars × 2 lines)" in problem.options


def test_columns_beyond_every_pattern_are_a_problem_not_a_lost_column():
    columns = [PlanColumn(heading=f"C{i}", bullets=["text"]) for i in range(9)]
    (problem,) = assign(
        INVENTORY, _plan(COVER, SlidePlan(title="T", columns=columns))
    ).problems
    assert "columns: 9 given, max 2" in problem.problem
    assert any("columns.bullets×2" in o for o in problem.options)


def test_excluded_and_rekinded_patterns_follow_the_overrides():
    process = _pattern("Process")
    effective, _ = apply_overrides(
        INVENTORY,
        [
            PatternOverride(pattern_id=process.id, included=False),
            PatternOverride(pattern_id=_pattern("Agenda").id, kind="process"),
        ],
    )
    steps = [PlanItem(number=str(i), heading=f"S{i}") for i in range(1, 4)]
    labels = _labels(
        effective, COVER, SlidePlan(title="How", kind="process", items=steps)
    )
    assert labels[1] == "Agenda"


def test_an_other_slide_given_a_kind_in_the_review_is_assigned():
    """The analysis marks an 'other' slide one-off; the author's kind wins."""
    other = _bullets("slide:5", "Motif", source="slide").model_copy(
        update={"kind": "other", "reusable": False}
    )
    inventory = _inventory(other)
    slide = SlidePlan(title="Points", bullets=["a", "b"], kind="bullets")
    assert assign(inventory, _plan(slide)).problems
    effective, _ = apply_overrides(
        inventory, [PatternOverride(pattern_id=other.id, kind="bullets")]
    )
    assert _labels(effective, slide) == ["Motif"]
    # Re-kinded to 'other', or to the kind it already has, it stays unused.
    back, _ = apply_overrides(
        effective, [PatternOverride(pattern_id=other.id, kind="other")]
    )
    assert not back.pattern(other.id).reusable  # type: ignore[union-attr]
    one_off = other.model_copy(update={"kind": "bullets"})
    kept, _ = apply_overrides(
        _inventory(one_off), [PatternOverride(pattern_id=one_off.id, kind="bullets")]
    )
    assert not kept.pattern(one_off.id).reusable  # type: ignore[union-attr]


def test_an_empty_slide_in_the_middle_is_a_problem():
    (problem,) = assign(
        INVENTORY, _plan(COVER, {}, {"title": "P", "bullets": ["x"]})
    ).problems
    assert problem.index == 2 and problem.problem == "the slide is empty"


def test_a_subtitle_without_its_zone_leads_the_body_instead_of_failing():
    inventory = _inventory(_bullets("layout:1", "Plain"))
    slide = SlidePlan(title="Why", subtitle="Three reasons", bullets=["a", "b"])
    assert _labels(inventory, slide) == ["Plain"]
    lines = _assign_text(inventory.patterns[0], slide)["ph:1"]
    assert lines[0].text == "Three reasons" and lines[0].heading
    assert [line.source for line in lines] == ["subtitle", "bullets[0]", "bullets[1]"]


def test_step_numbers_need_no_number_zone_but_figures_do():
    process = _pattern("Process")
    process = process.model_copy(
        update={
            "groups": [
                g.model_copy(
                    update={
                        "item_zones": [r for r in g.item_zones if r != "number"],
                        "items": [
                            [z for z in item if z.role != "number"] for item in g.items
                        ],
                    }
                )
                for g in process.groups
            ]
        }
    )
    steps = [PlanItem(number=f"0{i}", heading=f"S{i}") for i in range(1, 4)]
    assert not misfit(process, SlidePlan(title="How", items=steps), (1, 1))[0]
    figures = [PlanItem(number="42%", heading="Margin")] * 3
    reasons, _ = misfit(process, SlidePlan(title="KPIs", items=figures), (1, 1))
    assert [r.text for r in reasons] == [
        "items' number: this slide has no place for it"
    ]


def test_a_cover_with_only_a_heading_is_described_as_taking_the_title():
    """The renderer writes the title there: calling it `columns.heading` sent
    models sending columns to the cover."""
    heading = _title_zone().model_copy(update={"role": "heading"})
    cover = Pattern(
        id="slide:1",
        source="slide",
        source_index=1,
        name="Cover",
        label="Cover",
        kind="cover",
        zones=[heading],
    )
    assert zone_summary(cover) == "title ≤8 words (≈60 chars)"
    assert _labels(_inventory(cover), SlidePlan(title="Kick-off")) == ["Cover"]


# --- deck-wide variety --------------------------------------------------------------------------


def _points(n: int, kind: Any = "bullets") -> SlidePlan:
    return SlidePlan(title=f"Points {n}", bullets=["a", "b"], kind=kind)


def test_a_b_a_b_is_avoided_when_a_third_pattern_fits():
    three = _inventory(
        _bullets("slide:1", "A", source="slide"),
        _bullets("slide:2", "B", source="slide"),
        _bullets("slide:3", "C", source="slide"),
    )
    labels = _labels(three, *(_points(n) for n in range(4)))
    assert labels[:3] == ["A", "B", "C"]
    assert labels[3] != labels[1]  # no A B A B: the reuse goes to the oldest slide


def test_a_first_reuse_beats_another_kind_but_a_second_does_not():
    def photo(pid: str, label: str) -> Pattern:
        pattern = _bullets(pid, label, source="slide")
        image = Zone(key="pic", role="image", box=Box(x=500, y=200, w=400, h=600))
        return pattern.model_copy(
            update={"kind": "image_text", "zones": pattern.zones + [image]}
        )

    inventory = _inventory(
        _bullets("slide:1", "Only bullets", source="slide"),
        _bullets("slide:2", "Two cols", source="slide").model_copy(
            update={"kind": "two_columns"}
        ),
        photo("slide:3", "Photo 1"),
        photo("slide:4", "Photo 2"),
    )
    # Inferred, bullets then two_columns, one _KIND_STEP apart.
    points = SlidePlan(title="Points", bullets=["a", "b"])
    shot = SlidePlan(title="Site", bullets=["a"], image="doc")
    assert _labels(inventory, points, shot, points, shot, points) == [
        "Only bullets",
        "Photo 1",
        "Only bullets",  # first reuse: still cheaper than another kind
        "Photo 2",
        "Two cols",  # second reuse: another kind wins
    ]


def _max_uses_inventory(*extra: Pattern) -> TemplateInventory:
    """One bullets example slide, one two-columns slide, and what else is given."""
    two = _bullets("slide:2", "Two cols", source="slide")
    return _inventory(
        _bullets("slide:1", "A", source="slide"),
        two.model_copy(update={"kind": "two_columns"}),
        *extra,
    )


def _pictured_layout() -> Pattern:
    # A layout with an empty picture: 9 in fit costs, above one reuse (8).
    pattern = _bullets("layout:3", "Pic")
    image = Zone(key="pic", role="image", box=Box(x=500, y=200, w=400, h=600))
    return pattern.model_copy(update={"zones": pattern.zones + [image]})


def test_max_uses_one_takes_another_variant_of_the_same_kind():
    plan = [_points(1), _points(2, "two_columns"), _points(3)]

    def labels(inventory: TemplateInventory, variety: Any, max_uses: int) -> List[str]:
        found = assign(inventory, _plan(*plan), variety, max_uses).patterns
        return [p.label if p else "-" for p in found]

    plain = _max_uses_inventory(_bullets("layout:3", "B"))
    assert labels(plain, "template_slide", 1) == ["A", "Two cols", "B"]
    assert labels(plain, "kind", 1) == ["A", "Two cols", "B"]
    # A variant leaving a zone unfilled is never taken for variety alone.
    pictured = _max_uses_inventory(_pictured_layout())
    for variety, max_uses in (("template_slide", 2), ("off", 1), ("template_slide", 1)):
        assert labels(pictured, variety, max_uses) == ["A", "Two cols", "A"]
    assert labels(pictured, "kind", 1) == ["A", "Two cols", "A"]
    # One bullets variant: the requested kind is kept, the reuse only priced.
    alone = assign(_max_uses_inventory(), _plan(*plan), "template_slide", 1)
    assert [p.label for p in alone.patterns if p] == ["A", "Two cols", "A"]


def test_max_uses_is_raised_to_what_the_template_allows():
    plan = [_points(1), _points(2, "two_columns"), _points(3), _points(4)]
    inventory = _max_uses_inventory(_bullets("layout:3", "B"))
    # Four content slides on three template slides: one is used twice anyway.
    assert assign(inventory, _plan(*plan), "template_slide", 1).repeats == []
    # By kind, two kinds for four slides: two each; bullets on three is one too many.
    (repeat,) = assign(inventory, _plan(*plan), "kind", 1).repeats
    assert repeat.kind == "bullets" and repeat.slides == [1, 3, 4]
    assert repeat.at_most == 2
    assert assign(inventory, _plan(*plan[:3]), "kind", 1).repeats == []
    # Structural slides never count: a cover adds nothing to the bullets' cap.
    cover = _bullets("slide:9", "Cover", source="slide").model_copy(
        update={"kind": "cover"}
    )
    covered = _max_uses_inventory(_bullets("layout:3", "B"), cover)
    with_cover = [SlidePlan(title="Deck", subtitle="x")] + plan[:3]
    assert assign(covered, _plan(*with_cover), "kind", 1).repeats == []


def test_a_slide_used_three_times_is_reported_with_the_unused_variants():
    narrow = _bullets("slide:3", "Narrow", source="slide")
    narrow.zones[1].chars_per_line, narrow.zones[1].max_lines = 2, 1
    inventory = _inventory(
        _bullets("slide:1", "A", source="slide"),
        _bullets("slide:2", "B", source="slide"),
        narrow,
        *_cards_inventory().patterns,
    )
    slides = [
        SlidePlan(title=f"P{n}", bullets=["a long enough line"], kind="bullets")
        for n in range(5)
    ]
    (note,) = [n for n in assign(inventory, _plan(*slides)).notes if "all use" in n]
    assert note.startswith("Slides 1, 3, 5 all use the template slide 'A'.")
    assert (
        "Narrow: title ≤8 words (≈60 chars); bullets 1 short word (≈2 chars × 1 line, max 1 entry)"
        in note
    )
    # By kind, the two template slides make one repeat.
    assignment = assign(inventory, _plan(*slides), "kind")
    (note,) = [n for n in assignment.notes if "are all" in n]
    assert note.startswith(
        "Slides 1, 2, 3, 4, 5 are all 'bullets' slides (template slides 'A', 'B')."
    )
    # The unused variants of that kind, with their sizes; never another kind's.
    assert (
        "Narrow: title ≤8 words (≈60 chars); bullets 1 short word (≈2 chars × 1 line, max 1 entry)"
        in note
    )
    assert "cards" not in note
    assert note.endswith("To vary them, recast one as above; otherwise nothing to do.")


def test_one_kind_on_five_template_slides_is_a_repeat_with_content_advice():
    """The real first deck: five lists on five template slides look alike;
    their lines name the kind each one suits."""
    inventory = with_item_slides(INVENTORY)
    assignment = assign(inventory, _plan(*EVIDENCE_SLIDES), "kind")
    (repeat,) = assignment.repeats
    assert repeat.kind == "bullets" and repeat.slides == [3, 5, 7, 9, 10]
    # 'Free text' would leave a zone unfilled where the others fill all of theirs.
    assert len(repeat.labels) == 4 and "Free text" not in repeat.labels
    assert repeat.revisable
    # Each suggestion says what the template's slide of that kind holds.
    items = "items max 6 [heading ≤12 words (≈20 chars × 4 lines), text ≤12 words (≈20 chars × 4 lines)]"
    # One cards and one timeline variant; the one two_columns variant already
    # holds two slides: slides 9 and 10 get nothing left to suggest.
    assert repeat.suggested == [
        f"slide 5 (parallel 'Label : text' points) → cards holds {items}",
        f"slide 7 (dated items) → timeline holds {items}",
    ]
    (note,) = [n for n in assignment.notes if "are all" in n]
    assert note.startswith("Slides 3, 5, 7, 9, 10 are all 'bullets' slides")
    assert f"slide 7 (dated items) → timeline holds {items}" in note


def test_advice_names_only_kinds_the_template_can_hold():
    # No timeline nor cards slide in the showcase, and its two_columns slide is used.
    (repeat,) = assign(INVENTORY, _plan(*EVIDENCE_SLIDES), "kind").repeats
    assert repeat.suggested == []
    # Three dated lines need a timeline of three items or more.
    small = with_item_slides(INVENTORY, kinds=("timeline",), count=2)
    (repeat,) = assign(small, _plan(*EVIDENCE_SLIDES), "kind").repeats
    assert not any("timeline" in s for s in repeat.suggested)


def test_advice_quotes_the_favoured_variant_and_the_size_to_shorten_to():
    inventory = with_item_slides(INVENTORY, kinds=("timeline",))
    timeline = inventory.patterns[-1]
    # A roomier layout of the same kind: quoted only when the example slide is too small.
    roomy = timeline.model_copy(
        deep=True, update={"id": "layout:950", "source": "layout"}
    )
    for item in roomy.groups[0].items:
        for zone in item:
            zone.chars_per_line = 40
    inventory.patterns.append(roomy)

    def advice(words: int) -> List[str]:
        long = "Janvier 2025 : " + "accord national cyber " * words
        dated = {
            "title": "Jalons",
            "kind": "bullets",
            "bullets": [long, "Juin 2025 : radars", "Mars 2026 : pôle spatial"],
        }
        slides = [COVER, dated] + [_points(n) for n in range(3)]
        (repeat,) = assign(inventory, _plan(*slides), "kind").repeats
        return repeat.suggested

    lead = "slide 2 (dated items) → timeline holds items max 6 ["
    # The example slide when the entries fit it.
    assert advice(2) == [
        f"{lead}heading ≤12 words (≈20 chars × 4 lines), text ≤12 words (≈20 chars × 4 lines)]"
    ]
    # Only the layout holds the long entry: it is quoted.
    assert advice(10) == [
        f"{lead}heading ≤20 words (≈40 chars × 4 lines), "
        "text ≤20 words (≈40 chars × 4 lines)]"
    ]
    # Nothing holds it: the example slide, with the size to shorten to.
    assert advice(20) == [
        f"{lead}heading ≤12 words (≈20 chars × 4 lines), text ≤12 words (≈20 chars × 4 lines)], "
        "shorten each entry to ≤23 words (≈160 chars)"
    ]


def test_advice_prefers_room_for_the_entries_over_an_unmeasured_field():
    base = with_item_slides(INVENTORY, kinds=("cards",)).patterns[-1]
    # Headings unmeasured, texts of 6 chars: no endless room for the heading.
    tight = base.model_copy(deep=True, update={"id": "slide:951"})
    for heading, text in tight.groups[0].items:
        heading.chars_per_line = None
        text.chars_per_line, text.max_lines = 6, 1
    # Room for the entries but a picture per item: only a cost, so it wins.
    pictured = base.model_copy(deep=True, update={"id": "slide:952"})
    for n, item in enumerate(pictured.groups[0].items):
        item.append(Zone(key=f"pic{n}", role="image", box=item[0].box))
    shape = content_shape(
        SlidePlan(title="T", bullets=["Coût : élevé et durable", "Délai : court"])
    )
    assert shape is not None
    advice = _advice([tight, pictured], "cards", shape)
    assert advice is not None and advice.startswith("cards holds items max 6 [heading")
    assert "image" in advice and "shorten" not in advice


def test_repeats_are_found_beside_slides_that_fit_nothing():
    too_long = {"title": "T", "bullets": ["x" * 2000], "kind": "bullets"}
    slides = [_points(n) for n in range(3)] + [too_long]
    # A second kind, so three slides on one template slide exceed the max.
    quote = _bullets("slide:2", "Q", source="slide").model_copy(
        update={"kind": "quote"}
    )
    inventory = _inventory(_bullets("slide:1", "A", source="slide"), quote)
    assignment = assign(inventory, _plan(*slides))
    assert [p.index for p in assignment.problems] == [4]
    assert assignment.patterns == []
    (repeat,) = assignment.repeats
    # The slide that fits nothing counts on its closest fit, the same 'A'.
    assert repeat.kind == "bullets" and repeat.slides == [1, 2, 3, 4]


def test_slides_that_fit_nothing_repeat_on_their_closest_fit():
    """The live Paintress call: a label on the items of two cards slides, max 1.
    The repeat is asked with the fit problems, not left to the repaired deck."""
    inventory = with_item_slides(INVENTORY, kinds=("cards", "timeline"))
    labelled = [
        {"heading": "Coût", "text": "bas", "label": "x"},
        {"heading": "Délai", "text": "court", "label": "y"},
    ]
    cards = [
        {"title": f"Atouts {n}", "kind": "cards", "items": labelled} for n in (1, 2)
    ]
    slides = [COVER, cards[0], {"title": "P", "bullets": ["a", "b"]}, cards[1]]
    assignment = assign(inventory, _plan(*slides), "template_slide", 1)
    assert [p.index for p in assignment.problems] == [2, 4]
    assert all("closest fit, a 'cards' slide" in p.problem for p in assignment.problems)
    (repeat,) = assignment.repeats
    assert repeat.slides == [2, 4] and repeat.labels == ["Synthetic cards"]
    assert repeat.revisable
    # The timeline slide has no place for a label either: never offered.
    assert not any(o.startswith("'timeline'") for o in repeat.others)
    assert any(o.startswith("'figures'") for o in repeat.others)


def test_content_shapes_are_read_from_text_patterns():
    def shape(*bullets: str) -> Any:
        found = content_shape(SlidePlan(title="T", bullets=list(bullets)))
        return found and (found[0], found[1])

    assert shape("1. Frame", "2. Build", "3. Ship") == (("process",), 3)
    assert shape("2024 : centre cyber", "Q2 2025 : radar", "2027 : usine") == (
        ("timeline",),
        3,
    )
    assert shape("2 800 experts", "9 sites", "45 ans") == (("figures",), 3)
    assert shape("Coût : élevé", "Délai : court") == (("cards",), 2)
    # A plain list, or a long lead before the colon, is no other shape.
    assert shape("Frame the need", "Build it", "Ship it") is None
    assert shape("Présence historique en Europe du Sud : 1980", "Objectif : x") is None


def test_structural_kinds_count_by_template_slide_only():
    parts = [SlidePlan(title=f"Part {n}", kind="section") for n in range(4)]
    assignment = assign(INVENTORY, _plan(*parts))
    assert not [r for r in assignment.repeats if r.revisable]


# The live mistral-small first call on Paintress: the agenda, too long for the
# agenda slide, lands on a bullets slide beside two bullets slides.
SMALL_MODEL_PLAN = [
    {
        "kind": "cover",
        "title": "Thales en Espagne et en Italie",
        "subtitle": "Activités",
    },
    {
        "kind": "agenda",
        "title": "Sommaire",
        "bullets": ["Contexte", "Espagne", "Italie", "Chiffres", "Synthèse", "Suite"],
    },
    {"kind": "section", "title": "Espagne"},
    {
        "kind": "bullets",
        "title": "Activités en Espagne",
        "bullets": ["Défense", "Spatial"],
    },
    {
        "kind": "figures",
        "title": "Chiffres Espagne",
        "items": [
            {"number": "3", "label": "Sites"},
            {"number": "1200", "label": "Staff"},
        ],
    },
    {"kind": "section", "title": "Italie"},
    {"kind": "bullets", "title": "Activités en Italie", "bullets": ["Radars", "Rail"]},
    {
        "kind": "figures",
        "title": "Chiffres Italie",
        "items": [
            {"number": "3", "label": "Sites"},
            {"number": "950", "label": "Staff"},
        ],
    },
]


def test_a_structural_slide_never_counts_towards_a_repeat():
    plan = _plan(*SMALL_MODEL_PLAN)
    assignment = assign(INVENTORY, plan, "kind")
    assert assignment.patterns[1] is not None
    assert assignment.patterns[1].kind == "bullets"  # the agenda, by fallback
    assert not [r for r in assignment.repeats if r.revisable]
    assert not [r for r in assign(INVENTORY, plan).repeats if r.revisable]


def _figure_lines(n: int) -> dict:
    return {
        "title": f"Chiffres {n}",
        "kind": "bullets",
        "bullets": [f"{n}2 sites", f"{n}00 experts", f"{n}5 ans"],
    }


def test_a_kind_is_suggested_no_more_than_the_template_has_variants():
    # Paintress-like: a single figures slide in the template.
    single = INVENTORY.model_copy(
        update={"patterns": [p for p in INVENTORY.patterns if p.id != "slide:268"]}
    )
    slides = [COVER] + [_figure_lines(n) for n in range(1, 5)]
    (repeat,) = assign(single, _plan(*slides), "kind").repeats
    assert sum("→ figures" in s for s in repeat.suggested) == 1
    assert len(repeat.suggested[0].split("(")[0].split(",")) == 1  # one slide
    assert not any(o.startswith("'figures'") for o in repeat.others)
    # A figures slide already in the deck uses that variant: none left to suggest.
    figures = {"title": "KPI", "items": [{"number": "3", "label": "Sites"}] * 2}
    (repeat,) = assign(single, _plan(*slides, figures), "kind").repeats
    assert repeat.revisable and not repeat.suggested
    assert not any(o.startswith("'figures'") for o in repeat.others)


def test_suggestions_spread_across_the_kinds_a_slide_suits():
    inventory = with_item_slides(INVENTORY, kinds=("cards",))
    cards = inventory.patterns[-1]
    wide = cards.model_copy(deep=True, update={"id": "slide:902"})
    for item in wide.groups[0].items:
        item[1].chars_per_line = 30
    second = _pattern("Two Content").model_copy(deep=True, update={"id": "layout:903"})
    for zone in second.zones:
        zone.chars_per_line = 30 if zone.role == "body" else zone.chars_per_line
    inventory.patterns += [wide, second]
    groups = ["Espagne :", "- défense", "- spatial", "Italie :", "- radars", "- rail"]
    slides = [COVER] + [
        {"title": f"Pays {n}", "kind": "bullets", "bullets": groups} for n in range(3)
    ]
    (repeat,) = assign(inventory, _plan(*slides), "kind").repeats
    kinds = {s.split("→ ")[1].split(" ")[0]: s.split(" (")[0] for s in repeat.suggested}
    # Two variants each: two_columns, then cards, then two_columns again.
    assert kinds == {"two_columns": "slides 2, 4", "cards": "slide 3"}


def test_empty_item_pictures_cost_like_empty_slide_pictures():
    inventory = _cards_inventory()
    text = inventory.patterns[0]
    imaged = text.model_copy(
        update={
            "id": "slide:3",
            "label": "Image cards",
            "groups": [
                text.groups[0].model_copy(
                    update={
                        "item_zones": ["image", "body"],
                        "items": [
                            [
                                Zone(
                                    key=f"pic{i}",
                                    role="image",
                                    box=Box(x=0, y=0, w=1, h=1),
                                )
                            ]
                            + item
                            for i, item in enumerate(text.groups[0].items)
                        ],
                    }
                )
            ],
        }
    )
    plain = SlidePlan(title="Axes", items=[PlanItem(text="Close")] * 3)
    size = (inventory.slide_width, inventory.slide_height)
    assert misfit(imaged, plain, size)[1] > misfit(text, plain, size)[1]
    pictured = SlidePlan(title="Axes", items=[PlanItem(text="Close", image="d")] * 3)
    assert misfit(imaged, pictured, size)[1] == misfit(text, plain, size)[1]
    # Without pictures, the icon-like text cards win over the image cards.
    assert _labels(_inventory(imaged, text), plain) == ["Text cards"]


def test_a_framed_picture_is_chosen_only_with_a_picture():
    """Paintress '3 colonnes d'image + texte': its layout draws a band under the
    pictures, so a card left without one shows a grey band."""
    text = _cards_inventory().patterns[0]
    steps = text.model_copy(
        update={"id": "slide:4", "label": "Steps", "kind": "process"}
    )
    framed = text.model_copy(
        deep=True,
        update={"id": "slide:3", "label": "Image cards"},
    )
    group = framed.groups[0]
    group.item_zones = ["image", "body"]
    for i, item in enumerate(group.items):
        item.insert(
            0,
            Zone(key=f"pic{i}", role="image", box=Box(x=0, y=0, w=1, h=1), framed=True),
        )
    plain = SlidePlan(title="Axes", kind="cards", items=[PlanItem(text="Close")] * 3)
    pictured = plain.model_copy(
        update={"items": [PlanItem(text="Close", image=f"d{i}") for i in range(3)]}
    )
    assert _labels(_inventory(framed, steps), plain) == ["Steps"]
    assert _labels(_inventory(framed, steps), pictured) == ["Image cards"]
    # Only the items used need a picture: the others leave with their items.
    two = pictured.model_copy(update={"items": pictured.items[:2]})
    assert _labels(_inventory(framed, steps), two) == ["Image cards"]
    (problem,) = assign(_inventory(framed), _plan(plain)).problems
    assert "images: this slide's picture area shows a frame" in problem.problem
    assert "image (required)" in zone_summary(framed)


def test_an_image_id_that_does_not_load_leaves_a_framed_picture_empty():
    """The build tool assigns with the pictures that load: an invalid id named
    by the plan must not choose a framed slide that will show its frame."""
    text = _cards_inventory().patterns[0]
    steps = text.model_copy(
        update={"id": "slide:4", "label": "Steps", "kind": "process"}
    )
    framed = text.model_copy(
        deep=True, update={"id": "slide:3", "label": "Image cards"}
    )
    framed.groups[0].item_zones = ["image", "body"]
    for i, item in enumerate(framed.groups[0].items):
        item.insert(
            0,
            Zone(key=f"pic{i}", role="image", box=Box(x=0, y=0, w=1, h=1), framed=True),
        )
    inventory = _inventory(framed, steps)
    plan = _plan(
        SlidePlan(
            title="Axes",
            kind="cards",
            items=[PlanItem(text="Close", image=f"d{i}") for i in range(3)],
        )
    )
    good = png_bytes(4, 4)
    fetched: dict[str, object] = {"d0": KeyError("d0"), "d1": b"text", "d2": b"<svg"}
    assignment, kept, bad = _assign_loadable(inventory, plan, fetched, "off", 2)
    assert bad == ["d0", "d1", "d2"] and not assignment.problems
    assert [p.label for p in assignment.patterns if p] == ["Steps"]
    assert [i.image for i in kept.slides[0].items] == [None, None, None]
    loaded = {f"d{i}": good for i in range(3)}
    assignment, _, bad = _assign_loadable(inventory, plan, loaded, "off", 2)
    assert bad == [] and [p.label for p in assignment.patterns if p] == ["Image cards"]


def test_a_heading_and_text_item_fits_a_card_with_one_text_zone():
    text_cards = _cards_inventory().patterns[0]
    item = SlidePlan(title="Axes", items=[PlanItem(heading="Near", text="Close")] * 3)
    size = (_cards_inventory().slide_width, _cards_inventory().slide_height)
    reasons, cost = misfit(text_cards, item, size)
    assert reasons == []
    # A designed heading zone still reads better than a bold lead.
    titled = _cards_inventory().patterns[1]
    assert misfit(titled, item, size)[1] < cost
    assert "heading+text" in zone_summary(text_cards)


def test_a_heading_that_does_not_fit_beside_the_text_is_reported():
    text_cards = _cards_inventory().patterns[0]
    long_text = "word " * 24  # fits the 26 × 3 zone alone, not under a heading
    item = SlidePlan(
        title="Axes", items=[PlanItem(heading="Near", text=long_text.strip())] * 3
    )
    size = (_cards_inventory().slide_width, _cards_inventory().slide_height)
    reasons, _ = misfit(text_cards, item, size)
    assert any("items' heading" in r.text for r in reasons)


# --- text the deck would drop ---------------------------------------------------------------


def _dropped(slide: dict) -> str:
    assignment = assign(INVENTORY, _plan(COVER, slide))
    (problem,) = assignment.problems
    assert problem.index == 2 and not assignment.patterns
    return problem.problem + " | " + " | ".join(problem.options)


def test_a_second_column_written_as_heading2_is_a_problem_not_a_lost_heading():
    """The real call: one `columns` object holding `heading2` and no second
    bullets. The coercion used to drop `heading2` and the deck showed one column."""
    text = _dropped(
        {
            "kind": "two_columns",
            "title": "Implantation et effectifs",
            "columns": {
                "bullets": ["45 ans d'implantation", "2 800 experts"],
                "heading": "Espagne",
                "heading2": "Italie",
            },
        }
    )
    assert "columns[0]: heading2 'Italie' is not a column field" in text
    assert "one {heading, bullets} per column" in text


def test_columns_sent_as_one_object_are_a_problem_not_one_column():
    """The eval call: `columns` one object on a two-column slide rendered one
    column; the tool cannot tell one column meant from several."""
    text = _dropped(
        {
            "kind": "two_columns",
            "title": "Les enjeux du projet",
            "columns": {"heading": "Objectifs", "bullets": ["Gagner", "Réduire"]},
        }
    )
    assert "columns sent as one object, not a list" in text
    assert '[{"heading": ..., "bullets": [...]}, {"heading"' in text
    # The marker survives the dump the tool re-validates.
    slide = SlidePlan.model_validate({"title": "T", "columns": {"bullets": ["x"]}})
    assert SlidePlan.model_validate(slide.model_dump()).single_columns


def test_a_half_written_column_or_item_is_a_problem():
    text = _dropped(
        {
            "title": "Implantation",
            "columns": [
                {"heading": "Espagne", "bullets": ["2 800 experts"]},
                {"heading": "Italie"},
            ],
        }
    )
    assert "columns[1] has heading 'Italie' but no bullets" in text
    text = _dropped(
        {
            "title": "Axes",
            "items": [{"heading": "Cloud", "text": "Migrate"}, {"text": "Train"}],
        }
    )
    assert "items[1] has text but no heading" in text
    # Headings alone (an agenda) are a whole list, not half-written entries.
    assert not assign(
        INVENTORY,
        _plan(
            COVER, {"title": "Agenda", "items": [{"heading": "A"}, {"heading": "B"}]}
        ),
    ).problems


def test_a_field_the_model_invented_is_a_problem_naming_its_text():
    text = _dropped(
        {"title": "Results", "bullets": ["Up"], "footnote": "Source: internal"}
    )
    assert "footnote 'Source: internal': not a slide field" in text
    text = _dropped(
        {"title": "Axes", "items": [{"heading": "Cloud", "detail": "Migrate"}] * 2}
    )
    assert "items[0]: detail 'Migrate' is not an item field" in text


def test_the_coercion_keeps_every_text_it_maps():
    column = PlanColumn.model_validate({"title": "A", "text": "x", "bullets": ["y"]})
    assert (column.heading, column.bullets, column.ignored_fields) == (
        "A",
        ["y", "x"],
        [],
    )
    slide = SlidePlan.model_validate(
        {"title": "T", "bullets": [{"heading": "Cloud", "text": "migrate"}]}
    )
    assert slide.bullets == ["Cloud: migrate"]
    # The tool re-validates a dump: the ignored fields survive the round trip.
    again = SlidePlan.model_validate(
        SlidePlan.model_validate({"title": "T", "extra": "kept"}).model_dump()
    )
    assert again.ignored_fields == ["extra 'kept'"]


def test_an_item_led_by_its_figure_is_not_half_written():
    figures = {
        "title": "CO2",
        "items": [
            {"number": "18%", "text": "less since 2022"},
            {"number": "32%", "label": "women"},
            {"heading": "Net zero", "text": "by 2040"},
        ],
    }
    assert not [
        p
        for p in assign(INVENTORY, _plan(COVER, figures)).problems
        if p.index == 2 and "incomplete" in p.problem
    ]


# --- parallel sections, alternatives the content can fill ----------------------------------


def _section_inventory() -> TemplateInventory:
    """A section slide, one bullets slide, one quote slide and three slides no
    plan below uses: the cap stays at the author's max (ceil(n / 5))."""
    section = Pattern(
        id="slide:50",
        source="slide",
        source_index=50,
        name="S",
        label="Section",
        kind="section",
        zones=[_title_zone()],
    )
    others = [
        _bullets(f"slide:{60 + n}", f"Cols {n}", source="slide").model_copy(
            update={"kind": "two_columns"}
        )
        for n in range(3)
    ]
    quote = _bullets("slide:2", "Q", source="slide").model_copy(
        update={"kind": "quote"}
    )
    return _inventory(section, _bullets("slide:1", "A", source="slide"), quote, *others)


def _quoted(n: int) -> SlidePlan:
    return SlidePlan(title=f"Quote {n}", bullets=["one key line"], kind="quote")


def _part(name: str) -> SlidePlan:
    return SlidePlan(title=name, kind="section")


def _revisable(plan: List[SlidePlan], max_uses: int = 2) -> List[Any]:
    found = assign(_section_inventory(), _plan(*plan), "template_slide", max_uses)
    return [r for r in found.repeats if r.revisable]


def test_the_limit_counts_within_each_section():
    """Twice in Italy (places 2 and 4), twice in Spain (places 1 and 3), max 2."""
    italy = [_part("Italie"), _quoted(1), _points(1), _quoted(2), _points(2)]
    spain = [_part("Espagne"), _points(3), _quoted(3), _points(4)]
    assignment = assign(
        _section_inventory(), _plan(*italy, *spain), "template_slide", 2
    )
    assert not assignment.problems and _revisable(italy + spain) == []
    labels = [p.label if p else "-" for p in assignment.patterns]
    assert labels == ["Section", "Q", "A", "Q", "A", "Section", "A", "Q", "A"]
    # Without the section slides, the deck is one section: four uses.
    flat = [s for s in italy + spain if s.kind != "section"]
    (repeat,) = [r for r in _revisable(flat) if r.kind == "bullets"]
    assert repeat.slides == [2, 4, 5, 7]


def test_three_uses_inside_one_section_are_still_revised():
    italy = [_part("Italie"), _points(1), _quoted(1), _points(2), _quoted(2)]
    italy += [_points(3)]
    (repeat,) = _revisable(italy + [_part("Espagne"), _points(4)])
    assert repeat.slides == [2, 4, 6] and repeat.at_most == 2 and not repeat.in_a_row


def _across_a_boundary(asked: bool) -> List[SlidePlan]:
    """Agenda parts, so no section slide sits between Italy's end and Spain."""
    agenda = SlidePlan(title="Sommaire", kind="agenda", bullets=["Italie", "Espagne"])
    spain = _points(3).model_copy(update={"title": "Espagne", "repeat_layout": asked})
    italy = _quoted(0).model_copy(update={"title": "Italie"})
    return [agenda, italy, _points(1), _points(2), spain, _quoted(1)]


def test_three_identical_slides_in_a_row_are_revised_across_sections():
    (repeat,) = _revisable(_across_a_boundary(asked=False))
    assert repeat.in_a_row and repeat.slides == [3, 4, 5]
    assert repeat.labels == ["A"] and repeat.at_most == 2


def test_identical_slides_the_user_asked_for_are_built_and_noted():
    slides = _across_a_boundary(asked=True)
    assignment = assign(_section_inventory(), _plan(*slides), "template_slide", 2)
    assert not assignment.problems and _revisable(slides) == []
    assert [p.label if p else "-" for p in assignment.patterns][2:5] == ["A"] * 3
    assert assignment.repeated == {5: True}
    assert "Slide 5 keeps slide 4's template slide, as the user asked" in " ".join(
        assignment.notes
    )


def test_repeat_layout_keeps_the_previous_template_slide_outside_the_cap():
    inventory = _inventory(
        _bullets("slide:1", "A", source="slide"),
        _bullets("slide:2", "B", source="slide"),
    )
    slides = [_points(1), _points(2), _points(3)]
    assert _labels(inventory, *slides) == ["A", "B", "A"]
    slides[1] = slides[1].model_copy(update={"repeat_layout": True})
    assignment = assign(inventory, _plan(*slides), "template_slide", 1)
    # Slide 2 keeps slide 1's design; slide 3, not asked, is not a third in a row.
    assert [p.label for p in assignment.patterns if p] == ["A", "A", "B"]
    assert assignment.repeated == {2: True} and assignment.repeats == []


def test_repeat_layout_the_content_cannot_keep_is_noted():
    inventory = _section_inventory()
    slides = [_quoted(1), _points(1).model_copy(update={"repeat_layout": True})]
    assignment = assign(inventory, _plan(*slides), "template_slide", 2)
    assert [p.label for p in assignment.patterns if p] == ["Q", "A"]
    assert assignment.repeated == {2: False}
    assert any(
        n.startswith("Slide 2 has `repeat_layout`, but its content does not fit")
        for n in assignment.notes
    )


def test_an_agenda_splits_the_deck_into_sections_without_section_slides():
    inventory = _section_inventory()
    agenda = SlidePlan(
        title="Sommaire", kind="agenda", bullets=["1. Italie", "2. Espagne", "France"]
    )
    slides = [agenda]
    for n, name in enumerate(("Italie", "Espagne", "France")):
        slides += [
            SlidePlan(title=name, bullets=["a", "b"], kind="bullets"),
            _quoted(n),
        ]

    def revisable(plan: List[SlidePlan]) -> List[Any]:
        found = assign(inventory, _plan(*plan), "template_slide", 1).repeats
        return [r for r in found if r.revisable]

    assert revisable(slides) == []
    # Titles that name no agenda entry: one deck, a repeat.
    renamed = [agenda] + [
        s.model_copy(update={"title": f"{s.title} x"}) for s in slides[1:]
    ]
    assert revisable(renamed)


def _framed_cards(pid: str, label: str) -> Pattern:
    """A cards slide whose items each show a framed picture: it needs pictures."""
    cards = with_item_slides(INVENTORY, kinds=("cards",)).patterns[-1]
    pattern = cards.model_copy(deep=True, update={"id": pid, "label": label})
    for n, item in enumerate(pattern.groups[0].items):
        item.append(Zone(key=f"pic{n}", role="image", box=item[0].box, framed=True))
    return pattern


def test_variety_never_offers_a_variant_needing_pictures_the_slides_lack():
    labelled = SlidePlan(title="T", bullets=["Coût : élevé", "Délai : court"])
    slides = [labelled.model_copy(update={"title": f"T{n}"}) for n in range(3)]
    photo_cards = _framed_cards("slide:970", "Photo cards")
    inventory = _inventory(_bullets("slide:1", "A", source="slide"), photo_cards)
    (repeat,) = assign(inventory, _plan(*slides), "template_slide", 1).repeats
    assert not any("Photo cards" in o for o in repeat.others)
    assert not any("cards" in s for s in repeat.suggested)
    # A cards slide without pictures is offered and suggested instead.
    plain = with_item_slides(INVENTORY, kinds=("cards",)).patterns[-1]
    inventory = _inventory(_bullets("slide:1", "A", source="slide"), photo_cards, plain)
    (repeat,) = assign(inventory, _plan(*slides), "template_slide", 1).repeats
    assert not any("Photo cards" in o for o in repeat.others)
    (advice,) = [s for s in repeat.suggested if "cards holds" in s]
    assert "image" not in advice


def test_a_fit_remedy_never_names_a_kind_needing_a_missing_picture():
    photo = _bullets("slide:3", "Photo", source="slide")
    image = Zone(
        key="pic", role="image", box=Box(x=500, y=200, w=400, h=600), framed=True
    )
    photo = photo.model_copy(
        update={"kind": "image_text", "zones": photo.zones + [image]}
    )
    two = _bullets("slide:2", "Two cols", source="slide").model_copy(
        update={"kind": "two_columns"}
    )
    inventory = _inventory(_bullets("slide:1", "A", source="slide"), two, photo)
    too_long = SlidePlan(title="T", bullets=["x" * 2000], kind="bullets")
    (problem,) = assign(inventory, _plan(too_long)).problems
    assert not any("'image_text'" in o for o in problem.options), problem.options
    pictured = too_long.model_copy(update={"image": "doc"})
    (problem,) = assign(inventory, _plan(pictured)).problems
    assert any("'image_text'" in o for o in problem.options), problem.options
