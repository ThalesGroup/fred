# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""A small upload that declares a huge or abnormally compressed member is
refused before python-pptx inflates anything."""

from __future__ import annotations

import copy
import io
import re
import struct
import zipfile
from typing import Any

import pytest
from builder_decks import build_showcase_template
from deck_builders import build_deck
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.capability import (
    PptBuilderCapability,
    PptBuilderConfig,
)
from fred_capability_ppt_filler.capability import (
    TEMPLATE_SLOT,
    PptFillerCapability,
    PptFillerConfig,
)
from fred_capability_ppt_filler.pptx_archive import (
    UnsafeArchiveError,
    check_archive,
    sanitize_template,
)
from fred_sdk.contracts.capability import (
    CapabilityIdentity,
    SaveContext,
    UploadedFile,
)
from fred_sdk.contracts.runtime import RuntimeServices
from port_fakes import FakeAssets
from pptx import Presentation
from pptx.util import Inches


def _declaring_huge_member() -> bytes:
    """A real deck whose central directory claims one member inflates to ~4 GB."""
    content = bytearray(build_showcase_template())
    entry = content.find(b"PK\x01\x02")
    struct.pack_into("<I", content, entry + 24, 0xF0000000)
    return bytes(content)


def _ratio_bomb() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", b"\0" * (12 * 1024 * 1024))
    return buffer.getvalue()


def _uploads(content: bytes):
    return {TEMPLATE_SLOT: [UploadedFile(filename="t.pptx", content=content)]}


def _ctx() -> SaveContext:
    return SaveContext(
        identity=CapabilityIdentity(
            user_id="u",
            team_id="t",
            agent_instance_id="00000000-0000-4000-8000-000000000001",
        ),
        services=RuntimeServices(agent_assets=FakeAssets()),
    )


def test_declared_sizes_and_ratios_are_checked():
    bomb = _declaring_huge_member()
    assert len(bomb) < 1024 * 1024
    with pytest.raises(UnsafeArchiveError, match="inflates"):
        check_archive(bomb)
    with pytest.raises(UnsafeArchiveError, match="abnormally compressed"):
        check_archive(_ratio_bomb())
    check_archive(build_showcase_template())  # an ordinary deck passes


def test_the_analysis_reports_a_bomb_as_an_unreadable_template():
    result = analyze_template(_declaring_huge_member())
    assert result.inventory is None and "inflates" in result.errors[0].message


@pytest.mark.asyncio
async def test_both_capabilities_refuse_a_bomb_at_save():
    with pytest.raises(ValueError, match="inflates"):
        await PptBuilderCapability().validate_config(
            PptBuilderConfig(), _uploads(_declaring_huge_member()), _ctx()
        )
    with pytest.raises(ValueError, match="inflates"):
        await PptFillerCapability().validate_config(
            PptFillerConfig(), _uploads(_declaring_huge_member()), _ctx()
        )


_IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
    b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01"
    b"\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _linked_template(base: bytes | None = None) -> bytes:
    """Linked pictures on a slide and a layout, a hyperlink, and a macro
    project, as a .pptm saved under a .pptx name."""
    presentation = Presentation(io.BytesIO(base or build_showcase_template()))
    slide = presentation.slides[0]
    layout = slide.slide_layout
    picture = slide.shapes.add_picture(io.BytesIO(_PNG), Inches(1), Inches(1))
    element: Any = picture._element
    blip = element.blipFill.find("{*}blip")
    del blip.attrib[_R + "embed"]
    blip.set(
        _R + "link", slide.part.rels.get_or_add_ext_rel(_IMAGE_REL, "http://x/s.png")
    )
    on_layout = copy.deepcopy(element)
    on_layout.blipFill.find("{*}blip").set(
        _R + "link", layout.part.rels.get_or_add_ext_rel(_IMAGE_REL, "http://x/l.png")
    )
    layout.shapes._spTree.append(on_layout)
    box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "site"
    run.hyperlink.address = "https://example.org"
    buffer = io.BytesIO()
    presentation.save(buffer)
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml."
                    b"presentation.main+xml",
                    b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
                ).replace(
                    b"</Types>",
                    b'<Default Extension="bin" '
                    b'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
            elif info.filename == "ppt/_rels/presentation.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/'
                    b'office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
                    b"</Relationships>",
                )
            target.writestr(info, data)
        target.writestr("ppt/vbaProject.bin", b"macro code")
    return out.getvalue()


def _parts(content: bytes) -> dict[str, bytes]:
    archive = zipfile.ZipFile(io.BytesIO(content))
    return {name: archive.read(name) for name in archive.namelist()}


def test_a_template_loses_what_reaches_outside_it_and_keeps_its_links():
    original = _linked_template()
    before = _parts(original)
    assert any(b'TargetMode="External"' in v for v in before.values())
    parts = _parts(sanitize_template(original))
    rels = b"".join(v for k, v in parts.items() if k.endswith(".rels"))
    external = re.findall(rb"<Relationship [^>]*TargetMode=\"External\"[^>]*>", rels)
    assert len(external) == 1 and b"/hyperlink" in external[0]
    assert not any(b"r:link=" in v for v in parts.values())
    assert "ppt/vbaProject.bin" not in parts and b"vbaProject" not in rels
    assert b"macroEnabled" not in parts["[Content_Types].xml"]
    assert b"ms-office.vbaProject" not in parts["[Content_Types].xml"]
    # Still a template: python-pptx opens it and the analysis reads it.
    assert analyze_template(sanitize_template(original)).inventory is not None


def test_a_clean_template_is_left_byte_for_byte():
    clean = build_showcase_template()
    assert sanitize_template(clean) == clean


@pytest.mark.asyncio
async def test_the_builder_stores_and_analyses_the_sanitised_template(
    monkeypatch, analysis_db
):
    from fred_capability_ppt_filler.builder import capability as builder

    launched: list[bytes] = []

    async def _launch(_config, content, _ctx):
        launched.append(content)
        return "sha"

    monkeypatch.setattr(builder.PptBuilderCapability, "_launch", staticmethod(_launch))
    ctx = _ctx()
    await PptBuilderCapability().validate_config(
        PptBuilderConfig(), _uploads(_linked_template()), ctx
    )
    assets = ctx.services.agent_assets
    assert isinstance(assets, FakeAssets)
    (stored,) = assets.blobs.values()
    assert launched == [stored]
    assert b"http://x/" not in _parts(stored)["ppt/slides/_rels/slide1.xml.rels"]


@pytest.mark.asyncio
async def test_the_filler_stores_the_sanitised_template():
    ctx = _ctx()
    deck = build_deck([("Hello {{name}}", "{{name}}:\nThe name")])
    await PptFillerCapability().validate_config(
        PptFillerConfig(), _uploads(_linked_template(deck)), ctx
    )
    assets = ctx.services.agent_assets
    assert isinstance(assets, FakeAssets)
    (stored,) = assets.blobs.values()
    assert b"http://x/" not in _parts(stored)["ppt/slides/_rels/slide1.xml.rels"]
