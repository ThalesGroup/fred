# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The build tool on the built-in template, for a turn whose model call was
sent a given history: the storytelling and plan-validation tests share it."""

from __future__ import annotations

import io
import itertools
from typing import Any, Dict, List, Optional, cast

from fred_capability_ppt_filler.builder import tools as tools_mod
from fred_capability_ppt_filler.builder.capability import (
    PptBuilderCapability,
    PptBuilderConfig,
)
from fred_capability_ppt_filler.builder.tools import BUILD_TOOL_NAME
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
)
from fred_sdk.contracts.runtime import RuntimeServices
from langchain_core.messages import AIMessage, HumanMessage, ToolMessage
from port_fakes import FakeWorkspace
from pptx import Presentation

COVER = {"title": "Bilan 2025", "subtitle": "Direction des ressources humaines"}
USER = "Fais une présentation de nos effectifs : Italie 1 200, Espagne 600."
POINTS = {
    "takeaway": "L'Italie porte l'essentiel de nos effectifs européens",
    "title": "Où sont nos équipes",
    "bullets": ["Italie en tête", "Espagne en retrait"],
    "notes": "Insister sur l'Italie.",
}
CHART = {
    "takeaway": "L'Italie pèse deux fois l'Espagne, le recrutement doit commencer là",
    "title": "Effectifs par pays",
    "chart": {
        "type": "column",
        "categories": ["Italie", "Espagne"],
        "series": [{"name": "Effectif", "values": [1200, 600]}],
    },
    "source": "rapport RH 2025",
}
DECK = [COVER, POINTS, CHART]

_CALLS = itertools.count(1)


class _Request:
    system_prompt = ""

    def __init__(self, messages: list, tools: list) -> None:
        self.messages, self.tools = messages, tools

    def override(self, **_update: Any) -> "_Request":
        return self


async def _pdf(_bytes: bytes, *, suffix: str, timeout_seconds: float = 0):
    return b"%PDF-fake"


async def build_tool(
    monkeypatch,
    history: List[Any],
    *,
    validate_plan: bool = False,
    workspace: Optional[FakeWorkspace] = None,
) -> Any:
    """A new turn's build tool on the built-in template."""
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _pdf)
    config = PptBuilderConfig(validate_plan=validate_plan, variety_check="off")
    ctx = CapabilityContext(
        identity=CapabilityIdentity(
            user_id="u-1",
            session_id="sess-9",
            agent_instance_id="00000000-0000-4000-8000-000000000007",
        ),
        config=config,
        turn_options=EmptyModel(),
        services=RuntimeServices(
            workspace_fs=workspace if workspace is not None else FakeWorkspace()
        ),
    )
    (middleware,) = PptBuilderCapability().middleware(ctx)
    middleware = cast(Any, middleware)

    async def _handler(request):
        return request

    await middleware.awrap_model_call(_Request(history, middleware.tools), _handler)
    return {t.name: cast(Any, t) for t in middleware.tools}[BUILD_TOOL_NAME].coroutine


def answered(text: str, user: str = USER) -> List[Any]:
    """A turn: the user's message, a build call and its answer."""
    call = {"id": f"call-{next(_CALLS)}", "name": BUILD_TOOL_NAME, "args": {}}
    return [
        HumanMessage(content=user),
        AIMessage(content="", tool_calls=[call]),
        ToolMessage(content=text, tool_call_id=call["id"], name=BUILD_TOOL_NAME),
    ]


def deck(workspace: FakeWorkspace) -> Any:
    pptx = [w for w in workspace.writes if w["path"].endswith(".pptx")]
    return Presentation(io.BytesIO(pptx[-1]["content"]))


def notes(slide: Any) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def texts(slide: Any) -> List[str]:
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def summary_codes(result: Any) -> Dict[int, List[str]]:
    (part,) = result.ui_parts
    return {s.number: [n.code for n in s.notes] for s in part.summary}
