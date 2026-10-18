# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Rendering a plan on the synthetic showcase template (and on the reference
corporate template when `PPT_BUILDER_REFERENCE_TEMPLATE` points at it)."""

from __future__ import annotations

import io
import math
import os
import re
from typing import Any, cast

import pytest
from builder_decks import (
    FOOTER,
    add_cover_photo_slide,
    add_loose_figures_slide,
    build_showcase_template,
    build_slides,
    build_spaced_template,
)
from deck_builders import png_bytes
from fred_capability_ppt_filler import text_metrics
from fred_capability_ppt_filler.builder import builtin
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import assign, misfit
from fred_capability_ppt_filler.builder.lexicon import is_sample_text
from fred_capability_ppt_filler.builder.models import (
    Box,
    Pattern,
    TemplateInventory,
    Zone,
)
from fred_capability_ppt_filler.builder.plan import (
    DeckPlan,
    PlanColumn,
    PlanItem,
    SlidePlan,
)
from fred_capability_ppt_filler.builder.preview import build_patterns_deck
from fred_capability_ppt_filler.builder.render import (
    MIN_FONT_SCALE,
    RenderError,
    _assign_text,
    _covered,
    _fit,
    _Line,
    _unused_item_fields,
    _write_lines,
    render_deck,
)
from fred_capability_ppt_filler.builder.tools import _issue_line, catalogue
from fred_capability_ppt_filler.images import prepare_picture
from fred_capability_ppt_filler.traversal import walk_shapes
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

TEMPLATE = build_showcase_template()
_analysed = analyze_template(TEMPLATE).inventory
assert _analysed is not None
INVENTORY: TemplateInventory = _analysed


def _pattern(label: str):
    # A sample-text copy of a layout stands for it, labelled "<layout> (<n>)".
    return next(
        p
        for p in INVENTORY.patterns
        if p.label == label or p.label.startswith(f"{label} (")
    )


def _title(slide) -> str:
    return cast(Any, slide.shapes.title).text


def _texts(slide) -> list[str]:
    out = []
    for placed in walk_shapes(slide.shapes):
        shape = cast(Any, placed.shape)
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text and text != FOOTER:
                out.append(text)
    return out


def _findings(result) -> list[str]:
    """What the agent reads of a render: one line per slide, then the notes."""
    return [_issue_line(r) for r in result.report.issues()] + result.report.notes


def _pictures(slide) -> int:
    return sum(
        1
        for p in walk_shapes(slide.shapes)
        if cast(Any, p.shape).shape_type == MSO_SHAPE_TYPE.PICTURE
    )


Placed = tuple[Pattern | None, SlidePlan]


def _on(pattern: Pattern, **fields: Any) -> Placed:
    """One plan slide rendered on a given pattern, as the assigner would hand it."""
    return pattern, SlidePlan(**fields)


def _deck(*slides: Placed | SlidePlan) -> tuple[DeckPlan, list[Pattern | None]]:
    """A bare `SlidePlan` is an entry with no pattern (not a slide object)."""
    pairs = [s if isinstance(s, tuple) else (None, s) for s in slides]
    return DeckPlan(slides=[s for _, s in pairs]), [p for p, _ in pairs]


def _render(
    *slides: Placed | SlidePlan, images=None, template=TEMPLATE, inventory=INVENTORY
):
    plan, patterns = _deck(*slides)
    result = render_deck(template, inventory, plan, patterns, images)
    return result, Presentation(io.BytesIO(result.pptx))


def test_layout_pattern_fills_title_and_levelled_bullets():
    result, deck = _render(
        _on(
            _pattern("Title and Content"),
            title="Quarterly results",
            bullets=[
                "Revenue up **12 %**",
                "- Europe +8 %",
                "- Asia +20 %",
                "Costs flat",
            ],
        )
    )
    assert _findings(result) == []
    (slide,) = deck.slides
    assert _title(slide) == "Quarterly results"
    body = cast(Any, slide.placeholders[1]).text_frame
    assert [p.text for p in body.paragraphs] == [
        "Revenue up 12 %",
        "Europe +8 %",
        "Asia +20 %",
        "Costs flat",
    ]
    assert [p.level for p in body.paragraphs] == [0, 1, 1, 0]
    bold_runs = [r for p in body.paragraphs for r in p.runs if r.font.bold]
    assert [r.text for r in bold_runs] == ["12 %"]


def test_example_slides_are_duplicated_in_plan_order_and_originals_dropped():
    result, deck = _render(
        _on(
            _pattern("Key figures"),
            title="KPIs",
            items=[
                PlanItem(number="42%", label="Margin"),
                PlanItem(number="7", label="Sites"),
            ],
        ),
        _on(
            _pattern("Agenda"),
            title="Agenda",
            items=[
                PlanItem(heading="Context"),
                PlanItem(heading="Plan"),
                PlanItem(heading="Next steps"),
            ],
        ),
        _on(
            _pattern("Title and Content"),
            title="Closing",
            bullets=["Thanks"],
        ),
    )
    assert result.slide_count == 3
    assert len(deck.slides) == 3
    kpi, agenda, closing = deck.slides
    assert _title(kpi) == "KPIs"
    assert set(_texts(kpi)) == {"KPIs", "42%", "Margin", "7", "Sites"}
    # Two of the four cards were not used: their groups are gone, and so are the
    # loose icons drawn over them; the icons of the used cards stay.
    assert sum(1 for s in kpi.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP) == 2
    assert _pictures(kpi) == 2
    assert set(_texts(agenda)) == {
        "Agenda",
        "01",
        "Context",
        "02",
        "Plan",
        "03",
        "Next steps",
    }
    assert _title(closing) == "Closing"


def test_no_sample_text_survives_rendering():
    _, deck = _render(
        _on(
            _pattern("Process"),
            title="How",
            items=[PlanItem(heading="Discover")],
        ),
        _on(_pattern("Free text"), title="Notes"),
        _on(
            _pattern("Team"),
            title="Who",
            items=[PlanItem(heading="CTO", label="Ada")],
        ),
    )
    for slide in deck.slides:
        for text in _texts(slide):
            assert not is_sample_text(text), text


def test_shared_image_parts_keep_the_file_small():
    result, _ = _render(
        *[
            _on(
                _pattern("Key figures"),
                title=f"KPI {i}",
                items=[PlanItem(number="1", label="x")] * 4,
            )
            for i in range(10)
        ]
    )
    assert len(result.pptx) < len(TEMPLATE) * 1.2
    deck = Presentation(io.BytesIO(result.pptx))
    assert len(deck.slides) == 10
    assert all(_pictures(slide) == 4 for slide in deck.slides)


def test_narrow_item_box_is_widened_in_group_units():
    scaled = _pattern("Scaled figures")
    _, deck = _render(
        _on(
            scaled,
            title="Scaled",
            items=[PlanItem(number="1%", label="Wide label now")] * 3,
        )
    )
    labels = [
        p
        for p in walk_shapes(deck.slides[0].shapes)
        if cast(Any, p.shape).has_text_frame
        and cast(Any, p.shape).text_frame.text == "Wide label now"
    ]
    assert len(labels) == 3
    widths = sorted(p.width for p in labels)
    # Absolute widths within 5 % of each other: the narrow box caught up.
    assert widths[-1] - widths[0] < widths[-1] * 0.05


def test_notes_are_skipped_when_the_notes_slide_has_no_body():
    silent = _pattern("Silent notes")
    result, deck = _render(_on(silent, title="Quiet", bullets=["a"], notes="ignored"))
    assert len(deck.slides) == 1
    assert _findings(result) == []


def test_columns_without_enough_bodies_are_reported():
    result, _ = _render(
        _on(
            _pattern("Title and Content"),
            title="Three columns",
            columns=[
                PlanColumn(bullets=["a"]),
                PlanColumn(bullets=["b"]),
                PlanColumn(bullets=["c"]),
            ],
        )
    )
    assert any("not placed: column 2, column 3" in w for w in _findings(result))


def test_columns_fill_two_content_and_comparison_headings():
    result, deck = _render(
        _on(
            _pattern("Comparison"),
            title="Before / after",
            columns=[
                PlanColumn(heading="Before", bullets=["slow"]),
                PlanColumn(heading="After", bullets=["fast"]),
            ],
        ),
        _on(
            _pattern("Two Content"),
            title="Options",
            columns=[PlanColumn(bullets=["a", "b"]), PlanColumn(bullets=["c", "d"])],
        ),
    )
    comparison, two = deck.slides
    texts = _texts(comparison)
    assert {"Before", "After", "slow", "fast", "Before / after"} <= set(texts)
    bodies = [
        cast(Any, s) for s in two.placeholders if s.placeholder_format.idx in (1, 2)
    ]
    assert sorted(len(b.text_frame.paragraphs) for b in bodies) == [2, 2]


def test_overflow_shrinks_then_truncates_with_a_warning():
    long_bullets = [
        "This is a deliberately verbose bullet that keeps going on and on " * 3
    ] * 8
    result, deck = _render(
        _on(_pattern("Title and Content"), title="Long", bullets=long_bullets)
    )
    body = cast(Any, deck.slides[0].placeholders[1])
    # The body inherits 32 pt; the shrink is written as explicit run sizes.
    sizes = {run.font.size.pt for p in body.text_frame.paragraphs for run in p.runs}
    assert sizes and max(sizes) < 32

    very_long = ["word " * 200] * 40
    result, deck = _render(
        _on(
            _pattern("Title and Content"),
            title="Very long",
            bullets=very_long,
        )
    )
    assert any("text cut to fit" in w for w in _findings(result))
    body = cast(Any, deck.slides[0].placeholders[1])
    assert len(body.text_frame.paragraphs) < 40


def test_image_zone_filled_from_prepared_bytes_or_removed_when_empty():
    photo = _pattern("Photo")
    prepared = {"doc-1": prepare_picture(png_bytes(30, 20, "blue"))}
    result, deck = _render(_on(photo, title="Site", image="doc-1"), images=prepared)
    assert _findings(result) == []
    assert _pictures(deck.slides[0]) == 1

    caption = _pattern("Picture with Caption")
    _, deck = _render(_on(caption, title="No picture", bullets=["text only"]))
    slide = deck.slides[0]
    assert all(p.placeholder_format.type.name != "PICTURE" for p in slide.placeholders)

    result, deck = _render(
        _on(photo, title="Site", image="doc-missing"),
        images=prepared,
    )
    assert any("not placed: image 'doc-missing'" in w for w in _findings(result))
    assert _pictures(deck.slides[0]) == 1  # the template's own picture stays


def test_the_user_summary_records_each_rendered_slide_and_what_changed():
    """Codes, not prose: the chat dialog words them in the user's language."""
    very_long = ["word " * 200] * 40
    result, _ = _render(
        _on(_pattern("Title and Content"), title="**Very** long", bullets=very_long),
        _on(_pattern("Photo"), title="Site", image="doc-missing", kind="image_text"),
        images={},
    )
    first, photo = result.report.built
    assert (first.number, first.part, first.parts) == (1, 1, 1)
    assert first.title == "Very long"  # inline Markdown stripped
    assert first.template == _pattern("Title and Content").label
    assert [n.code for n in first.notes] == ["text_shortened"]
    assert [(n.code, n.count) for n in photo.notes] == [("image_unusable", 1)]
    assert photo.kind == _pattern("Photo").kind

    grouped = next(p for p in INVENTORY.patterns if p.groups and p.reusable)
    room = grouped.groups[0].count
    steps = [{"heading": f"Step {i}"} for i in range(room + 2)]
    result, _ = _render(_on(grouped, title="Steps", items=steps))
    (built,) = result.report.built
    assert ("items_left_out", 2) in [(n.code, n.count) for n in built.notes]


def test_non_slide_entries_are_skipped_and_an_all_invalid_plan_is_an_error():
    result, deck = _render(
        _on(_pattern("Title and Content"), title="Kept"),
        SlidePlan.model_validate("junk"),
    )
    assert len(deck.slides) == 1
    # The tool's notes name the skipped entry; the render only leaves it out.
    assert _findings(result) == []
    with pytest.raises(RenderError, match="no usable slide"):
        _render(SlidePlan.model_validate("junk"))


def test_figures_never_break_and_long_words_shrink_instead_of_splitting():
    kpi = _pattern("Key figures")
    _, deck = _render(
        _on(
            kpi,
            title="KPIs",
            items=[PlanItem(number="1 200", label="Collaborateurs")],
        )
    )
    figures = [
        cast(Any, p.shape).text_frame.text
        for p in walk_shapes(deck.slides[0].shapes)
        if cast(Any, p.shape).has_text_frame
        and "200" in cast(Any, p.shape).text_frame.text
    ]
    assert figures == ["1\u00a0200"]


def test_fixed_pattern_is_copied_verbatim_and_notes_are_written():
    divider = _pattern("Motif")
    _, deck = _render(_on(divider, title="ignored", notes="Say hello"))
    slide = deck.slides[0]
    assert _title(slide) == "Motif"
    assert cast(Any, slide.notes_slide.notes_text_frame).text == "Say hello"


def test_content_the_pattern_cannot_show_is_reported():
    # The agent only learns what the deck lost from this report.
    result, _ = _render(
        _on(_pattern("Title Only"), title="Spain", bullets=["55 years"]),
        _on(
            _pattern("Key figures"),
            title="Figures",
            items=[PlanItem(number="1 600", label="People", text="Across 5 sites")],
        ),
        _on(_pattern("Motif"), title="Merci"),
        _on(_pattern("Motif"), title="Motif"),
        _on(
            _pattern("Free text"),
            title="No list",
            items=[PlanItem(heading="x")],
        ),
    )
    findings = " ".join(_findings(result))
    assert "Slide 1: not placed: bullets" in findings
    assert "Slide 2: not placed: items' text" in findings
    assert result.report.notes == [
        "Slide 3: 'Motif' is a fixed slide shown as is, so its content was left out. "
        "If it matters, move it to another slide."
    ]
    # A fixed slide named by its own title lost nothing.
    assert "Slide 4:" not in findings
    assert "Slide 5: not placed: items" in findings


def test_zone_role_names_used_as_plan_keys_still_fill_the_zones():
    steps = SlidePlan.model_validate(
        {
            "title": "Plan",
            "items": [{"number": "1", "heading": "Start", "body": "Kick-off meeting"}],
        }
    )
    content = SlidePlan.model_validate({"title": "Plan", "body": ["one", "two"]})
    result, deck = _render(
        (_pattern("Process"), steps), (_pattern("Title and Content"), content)
    )
    assert _findings(result) == []
    assert "Kick-off meeting" in _texts(deck.slides[0])
    assert any("one" in text and "two" in text for text in _texts(deck.slides[1]))
    # A field under its name and an alias keeps both texts, never first-wins.
    assert PlanItem.model_validate({"text": "kept", "body": "alias"}).text == (
        "kept\nalias"
    )
    assert PlanItem.model_validate({"text": "same", "body": "same"}).text == "same"


def test_one_line_heading_shrinks_then_asks_to_be_shortened():
    process = _pattern("Process")
    result, deck = _render(
        _on(
            process,
            title="Plan",
            items=[
                PlanItem(heading="Kick-off", text="a"),
                PlanItem(heading="A heading far too long for its one line", text="b"),
            ],
        )
    )
    (long_heading,) = [n for n in result.report.notes if "one line" in n]
    assert long_heading.startswith(
        "Slide 1: the heading 'A heading far too long for its one line' is too long for "
        "one line and runs into the text below. Shorten it to under "
    )

    # A label with nothing under it may take a second line.
    figures, _ = _render(
        _on(
            _pattern("Key figures"),
            title="Figures",
            items=[
                PlanItem(number="12", label="A label far too long for its one line")
            ],
        )
    )
    assert not [w for w in _findings(figures) if "text cut to fit" in w]


def test_truncation_drops_the_lines_the_measured_capacity_cannot_hold():
    # Capitals run ~20 % wider than the average character: costing the kept
    # lines in characters shipped a slide with twice its text in it.
    shouting = ["MODERNISATION DES SYSTEMES DE COMMANDEMENT"] * 12
    result, deck = _render(
        _on(_pattern("Title and Content"), title="Overflow", bullets=shouting)
    )
    assert any("text cut to fit" in w for w in _findings(result))
    body = cast(Any, deck.slides[0].placeholders[1]).text_frame
    kept = [p.text for p in body.paragraphs if p.text.strip()]
    assert 0 < len(kept) < len(shouting)

    # What is left must fit the box at the smallest font the renderer uses,
    # measured here rather than counted in characters.
    zone = next(z for z in _pattern("Title and Content").zones if z.role == "body")
    assert zone.font_pt and zone.max_lines
    line_pt = zone.box.w / 1000 * INVENTORY.slide_width / 914400 * 72
    em_per_line = line_pt / (zone.font_pt * MIN_FONT_SCALE)
    rows = sum(math.ceil(text_metrics.width_em(text) / em_per_line) for text in kept)
    assert rows <= zone.max_lines / MIN_FONT_SCALE


def test_long_item_text_in_a_one_line_box_wraps_without_a_warning():
    # Only headings and labels are held to one line; a description may wrap.
    process = _pattern("Process")
    one_line = process.model_copy(deep=True)
    for item in one_line.groups[0].items:
        for zone in item:
            if zone.role == "body":
                zone.max_lines = 1
    inventory = INVENTORY.model_copy(
        update={
            "patterns": [
                one_line if p.id == process.id else p for p in INVENTORY.patterns
            ]
        }
    )
    result, _ = _render(
        _on(
            one_line,
            title="Plan",
            items=[
                PlanItem(
                    heading="Start",
                    text="A description written in full sentences, " * 2,
                )
            ],
        ),
        inventory=inventory,
    )
    assert not [w for w in _findings(result) if "text cut to fit" in w]


def _zone(key: str, role: str, y: int) -> Zone:
    return Zone(key=key, role=role, box=Box(x=50, y=y, w=900, h=60))  # type: ignore[arg-type]


def test_contact_bullets_skip_zones_already_filled():
    contact = Pattern(
        id="layout:0",
        source="layout",
        source_index=0,
        name="Contact",
        label="Contact",
        kind="closing",
        zones=[
            _zone("ph:0", "title", 50),
            _zone("ph:1", "subtitle", 150),
            _zone("ph:2", "body", 250),
            _zone("ph:3", "body", 350),
        ],
    )
    lost: list[str] = []
    out = _assign_text(
        contact,
        SlidePlan(title="Jane", subtitle="CEO", bullets=["jane@x.io", "+33 1"]),
        lost,
    )
    assert lost == []
    assert [out[k][0].text for k in ("ph:1", "ph:2", "ph:3")] == [
        "CEO",
        "jane@x.io",
        "+33 1",
    ]


def test_a_lost_subtitle_is_reported_as_the_subtitle():
    lost: list[str] = []
    _assign_text(
        _pattern("Two columns"),
        SlidePlan(
            title="T",
            subtitle="S",
            columns=[PlanColumn(bullets=["a"]), PlanColumn(bullets=["b"])],
        ),
        lost,
    )
    assert lost == ["subtitle"]


def test_items_shown_by_another_group_of_the_slide_are_not_reported():
    figures = _pattern("Key figures").groups[0]  # number, label
    steps = _pattern("Process").groups[0]  # number, heading, body
    item = PlanItem(number="1", label="People", heading="Start", text="Kick-off")
    plan = SlidePlan(items=[item])
    assert (
        _unused_item_fields(
            [figures, steps], plan, (INVENTORY.slide_width, INVENTORY.slide_height)
        )
        == []
    )
    assert _unused_item_fields(
        [figures], plan, (INVENTORY.slide_width, INVENTORY.slide_height)
    ) == ["items' heading", "items' text"]


def test_preview_copy_widens_one_line_text_the_delivered_deck_keeps():
    # LibreOffice breaks a "do not wrap" box once its substitute font runs wide.
    template = Presentation(io.BytesIO(TEMPLATE))
    layout = next(lay for lay in template.slide_layouts if lay.name == "Title Only")
    scratch = template.slides.add_slide(layout)
    box = scratch.shapes.add_textbox(Inches(4), Inches(3), Inches(2), Inches(1))
    box.text_frame.text = "Merci"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(60)
    body = box.text_frame._txBody.find(qn("a:bodyPr"))
    body.set("wrap", "none")
    box.text_frame.paragraphs[0]._p.get_or_add_pPr().set("algn", "ctr")
    layout.shapes._spTree.append(box._element)
    buffer = io.BytesIO()
    template.save(buffer)

    result, _ = _render(
        _on(_pattern("Title Only"), title="End"), template=buffer.getvalue()
    )

    def merci_box(pptx: bytes):
        deck = Presentation(io.BytesIO(pptx))
        shape = next(
            s
            for s in deck.slides[0].slide_layout.shapes
            if s.has_text_frame and cast(Any, s).text_frame.text == "Merci"
        )
        return shape.left, shape.width

    assert merci_box(result.pptx) == (Inches(4), Inches(2))
    left, width = merci_box(result.preview_pptx)
    assert width > Inches(2)
    assert left + width // 2 == Inches(5)  # still centred where the designer put it


def test_preview_copy_draws_wingdings_3_bullets_the_delivered_deck_keeps():
    """The Thales master's sub-bullet is Wingdings 3, which LibreOffice cannot
    recode: the preview showed missing-glyph boxes."""
    template = Presentation(io.BytesIO(TEMPLATE))
    styles = template.slide_masters[0].element.find(f".//{qn('p:bodyStyle')}")
    level = styles.find(qn("a:lvl2pPr"))
    for tag in ("a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        for old in level.findall(qn(tag)):
            level.remove(old)
    font = level.makeelement(
        qn("a:buFont"), {"typeface": "Wingdings 3", "charset": "2"}
    )
    char = level.makeelement(qn("a:buChar"), {"char": "ê"})
    level.insert(0, char)
    level.insert(0, font)
    buffer = io.BytesIO()
    template.save(buffer)

    result, _ = _render(
        _on(_pattern("Title Only"), title="End"), template=buffer.getvalue()
    )

    def bullet(pptx: bytes):
        master = Presentation(io.BytesIO(pptx)).slide_masters[0].element
        found = master.find(f".//{qn('p:bodyStyle')}/{qn('a:lvl2pPr')}")
        return (
            found.find(qn("a:buFont")).get("typeface"),
            found.find(qn("a:buChar")).get("char"),
        )

    assert bullet(result.pptx) == ("Wingdings 3", "ê")
    assert bullet(result.preview_pptx) == ("Arial", "►")


def test_a_closing_holds_its_bullets_and_on_a_section_slide_takes_no_number():
    """The eval deck: three closing bullets found no place on the built-in
    closing (one line per box), so it showed as section '01 Merci'."""
    content, inventory = builtin.builtin_template()
    by_id = {p.id: p for p in inventory.patterns}
    size = (inventory.slide_width, inventory.slide_height)
    lines = ["Un projet clé", "Objectif : performance", "Prochaine étape : lancement"]
    closing = SlidePlan(title="Merci – questions", kind="closing", bullets=lines)
    assert misfit(by_id["slide:297"], closing, size) == ([], 0.0)
    plan = DeckPlan(slides=[closing.model_copy(update={"bullets": []})])
    result = render_deck(content, inventory, plan, [by_id["slide:261"]])
    (slide,) = Presentation(io.BytesIO(result.pptx)).slides
    assert not any(re.fullmatch(r"0?1", text) for text in _texts(slide))


def test_a_figures_label_reads_under_its_heading_without_a_label_zone():
    """The eval slide: number, heading and label on the built-in key figures,
    refused over the label, so the agent fell back to plain bullets."""
    content, inventory = builtin.builtin_template()
    figures = next(p for p in inventory.patterns if p.id == "slide:281")
    size = (inventory.slide_width, inventory.slide_height)
    items = [
        PlanItem(number="6,5 M€", heading="Investissement total", label="sur 3 ans"),
        PlanItem(number="2,8 M€", heading="Économies annuelles", label="dès 2027"),
        PlanItem(number="70 %", label="de ressaisies en moins"),
    ]
    slide = SlidePlan(title="Chiffres clés", kind="figures", items=items)
    assert misfit(figures, slide, size)[0] == []
    result = render_deck(content, inventory, DeckPlan(slides=[slide]), [figures])
    (built,) = Presentation(io.BytesIO(result.pptx)).slides
    texts = _texts(built)
    assert "Investissement total\nsur 3 ans" in texts
    assert "de ressaisies en moins" in texts


REFERENCE = os.environ.get("PPT_BUILDER_REFERENCE_TEMPLATE")


@pytest.mark.skipif(
    not REFERENCE or not os.path.exists(REFERENCE or ""),
    reason="reference template not provided",
)
def test_reference_template_renders_a_full_deck():
    template = open(REFERENCE, "rb").read()  # type: ignore[arg-type]
    inventory = analyze_template(template).inventory
    assert inventory is not None
    plan = DeckPlan(
        slides=[
            SlidePlan(title="Plan de transformation", subtitle="Comité de direction"),
            SlidePlan(
                title="Sommaire",
                items=[
                    PlanItem(heading=h)
                    for h in ("Contexte", "Ambition", "Feuille de route")
                ],
            ),
            SlidePlan(title="Contexte", kind="section"),
            SlidePlan(
                title="Constats",
                bullets=["Coûts en hausse", "- +8 % sur un an", "Délais stables"],
            ),
            SlidePlan(
                title="Chiffres clés",
                items=[
                    PlanItem(number="60%", label="Marge"),
                    PlanItem(number="3", label="Sites"),
                ],
            ),
            SlidePlan(
                title="Démarche",
                kind="process",
                items=[
                    PlanItem(heading="Cadrer", text="Deux semaines"),
                    PlanItem(heading="Construire", text="Un trimestre"),
                    PlanItem(heading="Déployer", text="Progressif"),
                ],
            ),
            SlidePlan(
                title="Options",
                columns=[
                    PlanColumn(heading="A", bullets=["rapide"]),
                    PlanColumn(heading="B", bullets=["robuste"]),
                ],
            ),
            SlidePlan(
                title="Nous contacter",
                kind="closing",
                bullets=[
                    "Jane Doe",
                    "Directrice",
                    "+33 1 00 00 00 00",
                    "jane@example.com",
                ],
            ),
        ]
    )
    assignment = assign(inventory, plan)
    assert not assignment.problems, assignment.problems
    kinds = [p.kind for p in assignment.patterns if p is not None]
    assert kinds == [
        "cover",
        "agenda",
        "section",
        "bullets",
        "figures",
        "process",
        # The template's comparison layout has no title zone: the title would be lost.
        "two_columns",
        "closing",
    ]
    result = render_deck(template, inventory, plan, assignment.patterns)
    deck = Presentation(io.BytesIO(result.pptx))
    assert len(deck.slides) == 8
    contact = deck.slides[7]
    assert {"Jane Doe", "Directrice", "+33 1 00 00 00 00", "jane@example.com"} <= set(
        _texts(contact)
    )
    for slide in deck.slides:
        for text in _texts(slide):
            assert not is_sample_text(text), text
    assert len(result.pptx) < len(template) * 1.1


@pytest.mark.skipif(
    not REFERENCE or not os.path.exists(REFERENCE or ""),
    reason="reference template not provided",
)
def test_reference_image_cards_need_their_pictures():
    """The live 'Nos quatre priorités' slide, cards without pictures, showed the
    layout's grey band where the removed pictures were."""
    template = open(REFERENCE, "rb").read()  # type: ignore[arg-type]
    inventory = analyze_template(template).inventory
    assert inventory is not None
    image_cards = next(
        p for p in inventory.patterns if p.label == "3 colonnes d’image + texte"
    )
    pictures = [z for z in image_cards.groups[0].items[0] if z.role == "image"]
    assert pictures and all(z.framed for z in pictures)
    priorities = SlidePlan(
        title="Nos quatre priorités stratégiques",
        kind="cards",
        items=[
            PlanItem(
                heading="Protection", text="Renforcer les infrastructures critiques"
            ),
            PlanItem(
                heading="Sensibilisation",
                text="Former les équipes aux bonnes pratiques",
            ),
            PlanItem(
                heading="Innovation", text="Solutions avancées (IA, automatisation)"
            ),
        ],
    )
    cover = SlidePlan(title="Stratégie Cybersécurité 2027", subtitle="Comité")
    plan = DeckPlan(slides=[cover, priorities])
    chosen = assign(inventory, plan).patterns[1]
    assert chosen is not None and chosen.id != image_cards.id
    pictured = priorities.model_copy(
        update={
            "items": [
                i.model_copy(update={"image": f"d{n}"})
                for n, i in enumerate(priorities.items)
            ]
        }
    )
    plan = DeckPlan(slides=[cover, pictured])
    assignment = assign(inventory, plan)
    assert assignment.patterns[1] is not None
    assert assignment.patterns[1].id == image_cards.id
    images = {f"d{n}": prepare_picture(png_bytes()) for n in range(3)}
    result = render_deck(template, inventory, plan, assignment.patterns, images)
    slide = Presentation(io.BytesIO(result.pptx)).slides[1]
    assert len(slide._element.findall(f".//{qn('p:blipFill')}")) == 3


# --- the room quoted to the agent is the room the renderer enforces ---


def _two_body_pattern(first: int, second: int) -> Pattern:
    return Pattern(
        id="layout:9",
        source="layout",
        source_index=9,
        name="Contact",
        label="Contact",
        kind="closing",
        zones=[
            Zone(
                key=f"ph:{index}",
                role="body",
                box=Box(x=index * 300, y=400, w=280, h=100),
                font_pt=14.0,
                chars_per_line=chars // 2,
                max_lines=2,
            )
            for index, chars in enumerate((first, second))
        ],
    )


def test_catalogue_quotes_the_tightest_zone_not_the_first():
    """Told the roomiest zone's limit, an agent writes text the narrow ones cut."""
    pattern = _two_body_pattern(36, 22)
    line = catalogue(
        TemplateInventory(slide_width=1, slide_height=1, patterns=[pattern])
    )
    assert "≈11 chars × 2 lines" in line
    assert "≈18 chars" not in line


def test_catalogue_names_columns_when_a_pattern_has_several_bodies():
    """`bullets` fills only the first: naming it that way loses the rest. A
    closing slide takes one bullet per zone instead."""
    contact = _two_body_pattern(40, 40)
    columns = contact.model_copy(update={"kind": "two_columns"})
    line = catalogue(
        TemplateInventory(slide_width=1, slide_height=1, patterns=[columns, contact])
    )
    assert "two_columns (two groups of points): Contact: columns.bullets×2" in line
    assert "closing: Contact: bullets, one per line,×2" in line


def test_catalogue_quotes_no_limit_for_a_box_that_grows_with_its_text():
    pattern = _two_body_pattern(2, 2)
    for zone in pattern.zones:
        zone.grows = True
    line = catalogue(
        TemplateInventory(slide_width=1, slide_height=1, patterns=[pattern])
    )
    assert "chars" not in line


def test_shrinking_text_to_the_floor_is_not_reported_as_truncated():
    """A zone that had to shrink still holds every word: telling the agent it
    was cut sends it rewriting content the reader can see in full."""
    zone = Zone(
        key="ph:1",
        role="body",
        box=Box(x=0, y=0, w=400, h=300),
        font_pt=18.0,
        chars_per_line=20,
        max_lines=9,
    )
    lines = [
        _Line(
            "Une ligne de texte nettement plus longue que ce que la zone affiche "
            "sans rétrécir la police au plancher"
        )
    ] * 3
    kept, scale, truncated = _fit(zone, lines, Inches(13.333).emu, Inches(7.5).emu)
    assert [line.text for line in kept] == [line.text for line in lines]
    assert not truncated
    assert scale == MIN_FONT_SCALE


def test_a_truncated_zone_is_never_left_empty():
    """A narrow zone used to blank itself rather than keep a cut first line."""
    zone = Zone(
        key="ph:1",
        role="body",
        box=Box(x=0, y=0, w=30, h=20),
        font_pt=18.0,
        chars_per_line=4,
        max_lines=1,
    )
    kept, _, truncated = _fit(
        zone, [_Line("Innovation durable")], Inches(13.333).emu, Inches(7.5).emu
    )
    assert truncated
    assert kept and kept[0].text.strip()


def test_a_placeholder_that_does_not_wrap_on_its_layout_is_not_truncated():
    """A cloned placeholder carries a bare bodyPr but still draws with the
    layout's rule: reading only the slide's cut text the box grows to hold."""
    inventory = analyze_template(build_spaced_template(space_before_pt=0)).inventory
    assert inventory is not None
    pattern = next(p for p in inventory.patterns if any(z.grows for z in p.zones))
    long_title = "Plan de transformation numérique du groupe à l'horizon 2030"
    result, deck = _render(
        _on(pattern, title=long_title),
        template=build_spaced_template(space_before_pt=0),
        inventory=inventory,
    )
    assert long_title in _texts(deck.slides[0])
    assert not [w for w in _findings(result) if "text cut to fit" in w]


# --- unused item slots leave the slide clean ---


def test_a_group_pattern_given_no_items_keeps_the_slide_design():
    """Sweeping decoration by overlap with nothing kept took the slide's own
    background with it, silently."""
    kpi = _pattern("Key figures")
    result, deck = _render(_on(kpi, title="KPIs", bullets=["Stray"]))
    assert _title(deck.slides[0]) == "KPIs"
    assert result.report.notes == [
        "Slide 1: 'Key figures' shows its content as items and none were given, so it "
        "stays empty. Give that content as `items` (up to 4)."
    ]
    # The icons drawn beside the cards are gone with them, the title stays.
    assert _texts(deck.slides[0]) == ["KPIs"]


def test_a_connector_counts_as_inside_the_area_it_spans():
    """A connector is a line, zero-thick on one axis: every area ratio scores it
    nothing, so the org-chart stubs of dropped slots used to dangle on."""
    area = (0, 0, 1000, 1000)
    horizontal_rule = (100, 500, 900, 500)
    outside = (100, 1500, 900, 1500)
    assert _covered(horizontal_rule, area) == 1.0
    assert _covered(outside, area) == 0.0


def test_the_kept_items_are_recentred_on_the_room_the_group_had():
    kpi = _pattern("Key figures")
    _, deck = _render(
        _on(kpi, title="KPIs", items=[PlanItem(number="42%", label="Margin")])
    )
    groups = [
        p
        for p in walk_shapes(deck.slides[0].shapes)
        if cast(Any, p.shape).shape_type == MSO_SHAPE_TYPE.GROUP
    ]
    assert len(groups) == 1
    centre = groups[0].left + groups[0].width // 2
    slide_centre = INVENTORY.slide_width // 2
    assert abs(centre - slide_centre) < Inches(0.6).emu


def test_unused_loose_items_are_removed_and_the_rest_recentred():
    template = build_slides([add_loose_figures_slide])
    inventory = analyze_template(template).inventory
    assert inventory is not None
    pattern = next(p for p in inventory.patterns if p.groups)
    _, deck = _render(
        _on(
            pattern,
            title="KPIs",
            items=[
                PlanItem(number="42%", label="Marge"),
                PlanItem(number="7", label="Sites"),
            ],
        ),
        template=template,
        inventory=inventory,
    )
    slide = deck.slides[0]
    assert sorted(_texts(slide)) == ["42%", "7", "KPIs", "Marge", "Sites"]
    figures = [
        p
        for p in walk_shapes(slide.shapes)
        if cast(Any, p.shape).text_frame.text in ("42%", "7")
    ]
    left = min(p.left for p in figures)
    right = max(p.left + p.width for p in figures)
    # Three columns spanned 0.5" to 6.7": the two kept ones sit around its middle.
    assert abs((left + right) // 2 - Inches(3.6).emu) < Inches(0.2).emu


def test_builtin_process_with_three_steps_drops_the_fourth():
    content, inventory = builtin.builtin_template()
    process = next(p for p in inventory.patterns if p.kind == "process")
    _, deck = _render(
        _on(
            process,
            title="Trois étapes",
            items=[
                PlanItem(heading=h, text="Une phrase") for h in ("Un", "Deux", "Trois")
            ],
        ),
        template=content,
        inventory=inventory,
    )
    texts = _texts(deck.slides[0])
    # The layout's step tiles come onto the slide so the fourth can go; the
    # slide number field comes with them.
    assert sorted(texts) == ["1", "2", "3", "Deux", "Trois", "Trois étapes", "Un"] + [
        "Une phrase"
    ] * 3 + ["‹#›"]
    assert "Check and deliver" not in texts and "4" not in texts
    # Recentred placeholders keep the height their layout gave them.
    tops = {
        cast(Any, p.shape).text_frame.text: p.top
        for p in walk_shapes(deck.slides[0].shapes)
        if p.shape.has_text_frame
    }
    assert tops["Un"] == tops["Deux"] > Inches(2).emu


def _section_slide_ids(presentation) -> list[str]:
    return [
        el.get("id")
        for el in presentation.part._element.iter(
            "{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId"
        )
    ]


_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def _with_sections(content: bytes) -> bytes:
    """The deck with one section naming all its slides, as PowerPoint saves it."""
    deck = Presentation(io.BytesIO(content))
    root = deck.part._element
    ext_lst = root.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = root.makeelement(qn("p:extLst"), {})
        root.append(ext_lst)
    ext = ext_lst.makeelement(
        qn("p:ext"), {"uri": "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"}
    )
    ext_lst.insert(0, ext)
    section = ext.makeelement(f"{{{_P14}}}sectionLst", {})
    ext.append(section)
    one = section.makeelement(f"{{{_P14}}}section", {"name": "All", "id": "{0}"})
    section.append(one)
    ids = one.makeelement(f"{{{_P14}}}sldIdLst", {})
    one.append(ids)
    for slide in deck.slides:
        ids.append(ids.makeelement(f"{{{_P14}}}sldId", {"id": str(slide.slide_id)}))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_rendered_and_patterns_decks_carry_no_stale_sections():
    """A template's sections name its own slides, which the render removes: a
    section list naming vanished slides makes PowerPoint repair."""
    content, inventory = builtin.builtin_template()
    content = _with_sections(content)
    assert _section_slide_ids(Presentation(io.BytesIO(content)))
    cover = next(p for p in inventory.patterns if p.kind == "cover")
    _, deck = _render(_on(cover, title="Titre"), template=content, inventory=inventory)
    assert _section_slide_ids(deck) == []
    patterns = Presentation(io.BytesIO(build_patterns_deck(content, inventory)))
    assert _section_slide_ids(patterns) == []


def test_an_unfilled_placeholder_is_removed_not_left_prompting():
    """An empty placeholder shows "Click to add text" as soon as the user opens
    the deck to edit it."""
    two = next(p for p in INVENTORY.patterns if p.label == "Two Content")
    _, deck = _render(_on(two, title="One side", bullets=["Only this"]))
    empty = [
        shape
        for shape in deck.slides[0].placeholders
        if not cast(Any, shape).text_frame.text.strip()
    ]
    assert not empty


def _filled_panel_slide(presentation):
    """A body box and, beside it, a second box drawn with its own fill."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    cast(Any, slide.shapes.title).text = "Panel slide"
    boxes = [
        slide.shapes.add_textbox(Inches(x), Inches(1.5), Inches(w), Inches(4))
        for x, w in ((0.5, 5.5), (6.3, 3.2))
    ]
    for box in boxes:
        box.text_frame.text = "Cliquez ici pour ajouter du texte"
    box = boxes[1]
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x1E, 0x2A, 0x5A)


def test_an_unfilled_text_box_is_removed_with_its_own_fill():
    template = build_slides([_filled_panel_slide])
    inventory = analyze_template(template).inventory
    assert inventory is not None
    pattern = next(p for p in inventory.patterns if p.source == "slide")
    _, deck = _render(
        _on(pattern, title="Panel slide", columns=[PlanColumn(bullets=["Only this"])]),
        template=template,
        inventory=inventory,
    )
    # Failed before: the box was emptied but kept, its navy panel still drawn.
    boxes = [s for s in deck.slides[0].shapes if not s.is_placeholder]
    assert [cast(Any, s).text_frame.text for s in boxes] == ["Only this"]


# --- the remaining ways content or design was quietly lost ---


def test_a_replaced_picture_keeps_its_place_in_the_drawing_order():
    """Appending the new picture put it in front of everything, so a replaced
    cover photo hid the title drawn over it."""
    template = build_slides([add_cover_photo_slide])
    inventory = analyze_template(template).inventory
    assert inventory is not None
    pattern = next(p for p in inventory.patterns if p.source == "slide")
    _, deck = _render(
        _on(pattern, title="Cover", image="doc-1"),
        images={"doc-1": prepare_picture(png_bytes(400, 300))},
        template=template,
        inventory=inventory,
    )
    tree = list(deck.slides[0].shapes._spTree)
    pictures = [i for i, el in enumerate(tree) if el.tag == qn("p:pic")]
    texts = [
        i
        for i, el in enumerate(tree)
        if el.tag == qn("p:sp")
        and "".join(t.text or "" for t in el.iter(qn("a:t"))).strip()
    ]
    assert pictures and texts
    assert max(pictures) < min(texts), "the replaced photo is in front of the title"


def test_the_hero_image_is_not_dropped_when_images_is_also_given():
    pattern = _pattern("Photo")
    result, deck = _render(
        _on(pattern, title="Cover", image="doc-1", images=["doc-2"]),
        images={
            "doc-1": prepare_picture(png_bytes(400, 300)),
            "doc-2": prepare_picture(png_bytes(300, 400)),
        },
    )
    assert _pictures(deck.slides[0]) >= 1
    assert not [w for w in _findings(result) if "not placed: image" in w]


def test_a_bullet_wrapped_under_an_unknown_key_keeps_its_text():
    plan = DeckPlan.model_validate(
        {
            "slides": [
                {
                    "title": "Results",
                    "bullets": [{"point": "Revenue up"}, "Costs flat"],
                }
            ]
        }
    )
    assert plan.slides[0].bullets == ["Revenue up", "Costs flat"]


def test_sibling_zones_of_one_slide_share_a_font_size():
    """Two headings of the same row at 24 pt and 13 pt read as a mistake, even
    when each was right on its own."""
    steps = _pattern("Process")
    _, deck = _render(
        _on(
            steps,
            title="How",
            items=[
                PlanItem(heading="Go", text="Start"),
                PlanItem(heading="A heading far longer than its siblings", text="Work"),
                PlanItem(heading="End", text="Close"),
            ],
        )
    )
    headings = ("Go", "A heading far longer than its siblings", "End")
    sizes = {}
    for placed in walk_shapes(deck.slides[0].shapes):
        shape = cast(Any, placed.shape)
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text_frame.text.strip()
        if text in headings:
            sizes[text] = {
                int(rPr.get("sz"))
                for rPr in shape.text_frame._txBody.iter(qn("a:rPr"))
                if rPr.get("sz")
            }
    assert set(sizes) == set(headings), sizes
    assert all(sizes.values()), "the shrunk sibling should carry an explicit size"
    assert len({frozenset(v) for v in sizes.values()}) == 1, sizes


def test_a_column_heading_written_into_the_body_drops_the_bullet():
    """Rendered as a bullet it is indistinguishable from the content below it."""
    content = _pattern("Title and Content")
    _, deck = _render(
        _on(
            content,
            title="Compare",
            columns=[PlanColumn(heading="Left", bullets=["a", "b"])],
        )
    )
    body = cast(Any, deck.slides[0].placeholders[1]).text_frame
    first = body.paragraphs[0]._p
    pPr = first.find(qn("a:pPr"))
    assert pPr is not None and pPr.find(qn("a:buNone")) is not None


def test_one_unbreakable_token_does_not_shrink_its_whole_zone_below_the_floor():
    """An e-mail address is one word: it used to set the size for the zone and
    took it down to 7 pt."""
    zone = Zone(
        key="ph:1",
        role="body",
        box=Box(x=0, y=0, w=200, h=300),
        font_pt=18.0,
        chars_per_line=15,
        max_lines=4,
    )
    _, scale, _ = _fit(
        zone,
        [_Line("marie.dupont@entreprise-exemple-tres-longue.fr")],
        Inches(13.333).emu,
        Inches(7.5).emu,
    )
    assert scale is None or scale >= MIN_FONT_SCALE


def test_a_title_that_auto_sizes_in_height_is_fitted_not_widened():
    """Such a box wraps inside its width and grows downwards. Widening it pushes
    a cover title over the photo beside it, and growing unchecked runs the text
    off the bottom of the slide - neither is reported unless it is fitted."""
    template = build_spaced_template(space_before_pt=0, no_wrap_title=False)
    inventory = analyze_template(template).inventory
    assert inventory is not None
    # The example slide standing for the layout given an auto-sizing title.
    pattern = next(p for p in inventory.patterns if p.source == "slide")
    zone = next(z for z in pattern.zones if z.role == "title")
    assert not zone.grows
    designed = zone.box.w
    result, deck = _render(
        _on(
            pattern,
            title="Un titre de couverture nettement plus long que la zone que "
            "le concepteur a dessinee pour lui, au point de sortir du cadre "
            "prevu et de courir sur la photo placee juste a cote de lui",
        ),
        template=template,
        inventory=inventory,
    )
    titles = [
        p
        for p in walk_shapes(deck.slides[0].shapes)
        if cast(Any, p.shape).has_text_frame
        and cast(Any, p.shape).text_frame.text.strip()
    ]
    assert titles
    widest = max(p.width for p in titles)
    assert widest <= designed / 1000 * inventory.slide_width * 1.02, (
        "the box was widened"
    )
    assert [w for w in _findings(result) if "text cut to fit" in w]


def test_an_agenda_announcing_more_than_the_deck_holds_is_reported():
    """A model can plan ten slides and write out two: the deck holds what was
    sent, and without this nothing says so."""
    result, _ = _render(
        _on(_pattern("Title and Content"), title="Cover"),
        _on(
            _pattern("Agenda"),
            title="Agenda",
            items=[PlanItem(heading=f"Topic {i}") for i in range(4)],
        ),
    )
    assert result.report.notes == [
        "Slide 2: the agenda announces 4 parts but only 0 slide(s) follow it. Add the "
        "missing slides or remove those parts from the agenda."
    ]


def test_an_agenda_matched_by_the_deck_is_not_reported():
    slides = [
        _on(
            _pattern("Agenda"),
            title="Agenda",
            items=[PlanItem(heading="One"), PlanItem(heading="Two")],
        )
    ] + [
        _on(_pattern("Title and Content"), title=t, bullets=["x"])
        for t in ("One", "Two")
    ]
    result, _ = _render(*slides)
    assert not [n for n in result.report.notes if "announces" in n]


# --- duplication fidelity and after-render check ---------------------------------

_DARK_MAP = (
    '<p:clrMapOvr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:overrideClrMapping bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2" accent1="accent1"'
    ' accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5"'
    ' accent6="accent6" hlink="hlink" folHlink="folHlink"/></p:clrMapOvr>'
)
_SHARED = "Confidential"


def _linked_template() -> tuple[bytes, TemplateInventory]:
    """A dark agenda slide linking to another slide, with notes and a footer."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml import parse_xml

    prs = Presentation()
    agenda = prs.slides.add_slide(prs.slide_layouts[5])
    target = prs.slides.add_slide(prs.slide_layouts[5])
    cast(Any, agenda.shapes.title).text = "Agenda overview"
    cast(Any, target.shapes.title).text = "Target slide"
    jump = agenda.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    jump.text_frame.text = "Go to target"
    jump.click_action.target_slide = target
    cast(
        Any, jump.text_frame.paragraphs[0].runs[0]
    ).hyperlink.address = "https://example.com/run"
    button = agenda.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(1), Inches(1)
    )
    button.click_action.hyperlink.address = "https://example.com"
    agenda.shapes.add_textbox(
        Inches(1), Inches(6), Inches(3), Inches(0.5)
    ).text_frame.text = _SHARED
    element = agenda._element
    element.remove(element.find(qn("p:clrMapOvr")))
    element.cSld.addnext(parse_xml(_DARK_MAP))
    notes = agenda.notes_slide
    cast(Any, notes.notes_text_frame).text = "intent: agenda"
    # A name only the copied notes slide carries (not the notes master).
    cast(Any, notes.notes_placeholder).name = "Brand note"
    buffer = io.BytesIO()
    prs.save(buffer)
    title = Zone(
        key="ph:0",
        role="title",
        box=Box(x=0, y=0, w=1000, h=150),
        chars_per_line=60,
        max_lines=1,
    )
    pattern = Pattern(
        id=f"slide:{agenda.slide_id}",
        source="slide",
        source_index=1,
        name="Agenda",
        label="Agenda",
        kind="agenda",
        zones=[title],
    )
    inventory = TemplateInventory(
        slide_width=int(prs.slide_width or 0),
        slide_height=int(prs.slide_height or 0),
        patterns=[pattern],
        shared_texts=[_SHARED],
    )
    return buffer.getvalue(), inventory


def _render_linked(notes: str | None = None):
    template, inventory = _linked_template()
    return _render(
        _on(inventory.patterns[0], title="Plan", notes=notes),
        template=template,
        inventory=inventory,
    )


def test_links_to_removed_slides_are_dropped_and_leave_no_orphan():
    import zipfile

    result, deck = _render_linked()
    names = zipfile.ZipFile(io.BytesIO(result.pptx)).namelist()
    assert len([n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]) == 1
    assert len(deck.slides) == 1
    slide = deck.slides[0]
    rels = slide.part.rels
    assert not any(rel.reltype.endswith("/slide") for rel in rels.values())
    external = {rel.target_ref for rel in rels.values() if rel.is_external}
    assert "https://example.com" in external
    # Every relationship reference left in the XML resolves.
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for node in slide._element.iter():
        for attr, value in node.attrib.items():
            if attr.startswith(r_ns) and value:
                assert value in rels, (node.tag, value)
    assert "hlinksldjump" not in slide._element.xml


def test_a_dark_example_slide_keeps_its_colour_map():
    _, deck = _render_linked()
    mapping = deck.slides[0]._element.find(
        f"{qn('p:clrMapOvr')}/{qn('a:overrideClrMapping')}"
    )
    assert mapping is not None
    assert (mapping.get("bg1"), mapping.get("tx1")) == ("dk1", "lt1")


def test_notes_slide_is_copied_and_holds_only_the_plan_notes():
    _, deck = _render_linked(notes="Say hello")
    notes = deck.slides[0].notes_slide
    assert "Brand note" in [shape.name for shape in notes.shapes]
    assert cast(Any, notes.notes_text_frame).text == "Say hello"
    _, deck = _render_linked()
    # The template's own notes (an `intent:` annotation) never reach the deck.
    assert cast(Any, deck.slides[0].notes_slide.notes_text_frame).text == ""


def test_example_text_left_on_a_duplicated_slide_is_removed_and_reported():
    result, deck = _render_linked()
    texts = [
        cast(Any, p.shape).text_frame.text
        for p in walk_shapes(deck.slides[0].shapes)
        if cast(Any, p.shape).has_text_frame
    ]
    assert "Go to target" not in texts
    assert "Plan" in texts and _SHARED in texts
    assert result.report.slide(1).removed == ["Go to target"]
    assert any("'Go to target'" in w for w in _findings(result))


def test_each_level_of_a_body_keeps_its_template_formatting():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    frame = box.text_frame
    frame.text = "Top level"
    first = frame.paragraphs[0]
    cast(Any, first.runs[0]).font.size = Pt(20)
    second = frame.add_paragraph()
    second.text = "Sub level"
    second.level = 1
    cast(Any, second.runs[0]).font.size = Pt(14)
    cast(Any, second._p.get_or_add_pPr()).set("marL", "457200")
    _write_lines(box, [_Line("A"), _Line("a1", 1), _Line("a2", 1), _Line("B")])
    paragraphs = frame.paragraphs
    assert [p.level for p in paragraphs] == [0, 1, 1, 0]
    assert [cast(Any, p.runs[0]).font.size.pt for p in paragraphs] == [20, 14, 14, 20]
    margins = [p._p.find(qn("a:pPr")).get("marL") for p in paragraphs]
    assert margins == [None, "457200", "457200", None]
