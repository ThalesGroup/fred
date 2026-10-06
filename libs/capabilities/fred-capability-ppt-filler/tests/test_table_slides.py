# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Table slides (design.md D14): a data table the template holds is a zone,
filled from the plan's `table`, rows and columns adapted to the content."""

from __future__ import annotations

import io
from typing import Any, List, cast

import pytest
from builder_decks import (
    add_data_table_slide,
    add_gantt_slide,
    add_table_slide,
    build_slides,
)
from fred_capability_ppt_filler.builder import tools
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import assign, infer_kinds, misfit
from fred_capability_ppt_filler.builder.labelling import kind_fits_slots
from fred_capability_ppt_filler.builder.models import TemplateInventory
from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanTable, SlidePlan
from fred_capability_ppt_filler.builder.render import render_deck
from pptx import Presentation
from pptx.oxml.ns import qn

_A16_ROW_ID = "{http://schemas.microsoft.com/office/drawing/2014/main}rowId"


def _bullets_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    cast(Any, slide.shapes.title).text = "Constats"
    cast(Any, slide.placeholders[1]).text_frame.text = "Un point"
    return slide


@pytest.fixture(scope="module")
def template() -> bytes:
    return build_slides([_bullets_slide, add_data_table_slide])


@pytest.fixture(scope="module")
def inventory(template) -> TemplateInventory:
    found = analyze_template(template).inventory
    assert found is not None
    return found


def _table_pattern(inventory: TemplateInventory):
    return next(p for p in inventory.patterns if p.kind == "table")


def _table(rows: int, cols: int = 5, cell: str = "x") -> PlanTable:
    return PlanTable(
        header=[f"Col {c}" for c in range(cols)],
        rows=[[f"{cell} {r}.{c}" for c in range(cols)] for r in range(rows)],
    )


# The template's layouts include a cover, which always opens the deck.
_COVER = SlidePlan(title="Choix du CRM", subtitle="Comité")


def _deck(*slides: SlidePlan) -> DeckPlan:
    return DeckPlan(slides=[_COVER, *slides])


def _render(template, inventory, table: PlanTable):
    plan = _deck(SlidePlan(title="Comparatif", table=table))
    assignment = assign(inventory, plan)
    assert not assignment.problems, assignment.problems
    result = render_deck(template, inventory, plan, assignment.patterns)
    slide = Presentation(io.BytesIO(result.pptx)).slides[1]
    frame: Any = next(s for s in slide.shapes if getattr(s, "has_table", False))
    return result, frame


def _rows(frame) -> List[List[str]]:
    return [[cell.text for cell in row.cells] for row in frame.table.rows]


def _row_ids(frame) -> List[str]:
    return [node.get("val") for node in frame._element.iter(_A16_ROW_ID)]


# --- analysis ---------------------------------------------------------------------


def test_a_data_table_is_a_table_zone(inventory):
    pattern = _table_pattern(inventory)
    assert pattern.included and pattern.reusable
    zone = next(z for z in pattern.zones if z.role == "table")
    assert zone.table is not None and zone.table.header
    assert len(zone.table.columns) == 5
    # 12 template rows: the header plus 11 one-line body rows at full size.
    assert zone.max_lines == 11
    assert kind_fits_slots("table", pattern)
    assert not kind_fits_slots("bullets", pattern)
    assert not kind_fits_slots("other", pattern)


def test_merged_and_gantt_grids_stay_excluded():
    inventory = analyze_template(
        build_slides([add_table_slide, add_gantt_slide])
    ).inventory
    assert inventory is not None
    slides = [p for p in inventory.patterns if p.source == "slide"]
    assert [p.exclusion_reason for p in slides] == ["table", "table"]
    assert all(not p.included for p in slides)


def test_catalogue_quotes_the_table(inventory):
    line = next(
        line
        for line in tools.catalogue(inventory).splitlines()
        if line.startswith("- table")
    )
    assert "table 2 to 5 columns" in line and "header + ≈11 lines of rows" in line
    schema = tools._args_schema(inventory).model_json_schema()
    assert "table" in schema["$defs"]["BuildSlide"]["properties"]


# --- plan ---------------------------------------------------------------------------


def test_plan_table_coercion_pads_ragged_rows_and_reads_objects():
    ragged = PlanTable.model_validate(
        {"header": ["A", "B", "C"], "rows": [["1"], ["2", "3"]]}
    )
    assert ragged.rows == [["1", "", ""], ["2", "3", ""]]
    objects = PlanTable.model_validate([{"Offre": "A", "Coût": 3}, {"Offre": "B"}])
    assert objects.header == ["Offre", "Coût"]
    assert objects.rows == [["A", "3"], ["B", ""]]
    assert infer_kinds(SlidePlan(title="T", table=objects)) == ["table"]


# --- fit and assignment ---------------------------------------------------------------


def test_table_content_goes_to_the_table_slide_only(inventory):
    plan = _deck(
        SlidePlan(title="Constats", bullets=["Un", "Deux"]),
        SlidePlan(title="Comparatif", table=_table(4)),
    )
    kinds = [p.kind for p in assign(inventory, plan).patterns if p]
    assert kinds == ["cover", "bullets", "table"]
    reasons, _ = misfit(
        _table_pattern(inventory),
        SlidePlan(title="Constats", bullets=["Un"]),
        (inventory.slide_width, inventory.slide_height),
    )
    assert any("needs `table`" in reason.text for reason in reasons)


def test_too_many_columns_or_one_column_is_a_misfit(inventory):
    size = (inventory.slide_width, inventory.slide_height)
    pattern = _table_pattern(inventory)
    wide, _ = misfit(pattern, SlidePlan(title="T", table=_table(3, cols=7)), size)
    assert wide[0].text == "table: 7 columns given, max 5"
    narrow, _ = misfit(pattern, SlidePlan(title="T", table=_table(3, cols=1)), size)
    assert "at least 2 columns" in narrow[0].text


def test_too_long_table_needs_revision_naming_the_cell(inventory):
    table = _table(14, cell="une cellule bien trop longue pour sa colonne")
    table.rows[3][2] += " et encore plus longue"
    problems = assign(inventory, _deck(SlidePlan(title="T", table=table))).problems
    assert len(problems) == 1
    text = problems[0].problem
    assert "14 rows need" in text and "this table holds ≈11" in text
    assert "at the smallest size" in text and "rows[3][2]" in text
    assert "chars per line" in text


def test_a_table_on_a_template_without_table_slide_is_a_problem():
    inventory = analyze_template(build_slides([_bullets_slide])).inventory
    assert inventory is not None
    problems = assign(inventory, _deck(SlidePlan(title="T", table=_table(3)))).problems
    assert problems and problems[0].problem == "table: this template has no table slide"
    assert (
        "table"
        not in tools._args_schema(inventory).model_json_schema()["$defs"]["BuildSlide"][
            "properties"
        ]
    )


# --- render ------------------------------------------------------------------------------


def test_fewer_rows_remove_rows_and_shrink_the_frame(template, inventory):
    result, frame = _render(template, inventory, _table(5))
    rows = _rows(frame)
    assert rows[0] == [f"Col {c}" for c in range(5)]
    assert rows[1][0] == "x 0.0" and len(rows) == 6
    assert frame.height < 383378 * 12
    assert result.report.issues() == []
    ids = _row_ids(frame)
    assert len(ids) == 6 and len(set(ids)) == 6


def test_more_rows_clone_rows_with_fresh_ids_and_shrink_the_text(template, inventory):
    result, frame = _render(template, inventory, _table(15))
    rows = _rows(frame)
    assert len(rows) == 16 and rows[-1][4] == "x 14.4"
    ids = _row_ids(frame)
    assert len(ids) == 16 and len(set(ids)) == 16
    assert frame.top + frame.height <= frame.top + 383378 * 12 + 6350
    sizes = {
        node.get("sz") for node in frame._element.iter(qn("a:rPr")) if node.get("sz")
    }
    assert len(sizes) == 1  # one font scale for the whole table
    assert result.report.slides[2].shrunk


def test_fewer_columns_drop_trailing_ones_and_widen_the_rest(template, inventory):
    _, frame = _render(template, inventory, _table(3, cols=3))
    widths = [column.width for column in frame.table.columns]
    assert len(widths) == 3 and sum(widths) == 9 * 914400
    assert all(len(row) == 3 for row in _rows(frame))


def test_cells_keep_the_template_style_and_bold(template, inventory):
    table = _table(2)
    table.rows[0][1] = "**Forte**"
    _, frame = _render(template, inventory, table)
    run = frame.table.cell(1, 1).text_frame.paragraphs[0].runs[0]
    assert run.text == "Forte" and run.font.bold
