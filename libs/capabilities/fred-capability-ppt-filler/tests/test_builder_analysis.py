# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Template analysis rules, each on the synthetic slide written for it, plus the
reference corporate template (optional, `PPT_BUILDER_REFERENCE_TEMPLATE`)."""

from __future__ import annotations

import copy
import io
import os
from typing import Any, cast

import pytest
from builder_decks import (
    FOOTER,
    L_BLANK,
    L_CONTENT,
    L_PICTURE,
    L_TITLE,
    L_TITLE_ONLY,
    TIMELINE_ENTRY,
    add_agenda_slide,
    add_big_figures_slide,
    add_big_number_slide,
    add_divider_slide,
    add_empty_slide,
    add_graphic_sheet_slide,
    add_loose_figures_slide,
    add_numbered_grid_slide,
    add_paragraph_subtitle_slide,
    add_placeholder_cards_slide,
    add_placeholder_two_ideas_slide,
    add_promo_slide,
    add_reused_entry_slide,
    add_staircase_timeline_slide,
    add_three_column_bodies_slide,
    build_layouts_only_template,
    build_library_template,
    build_showcase_template,
    build_slides,
    build_spaced_template,
)
from fred_capability_ppt_filler.builder.analysis import (
    ANALYSIS_VERSION,
    FrameMetrics,
    analyze_template,
    estimate_capacity,
)
from fred_capability_ppt_filler.builder.lexicon import (
    is_figure,
    is_library_title,
    is_sample_text,
    role_from_shape_name,
)
from fred_capability_ppt_filler.builder.models import Box, TemplateInventory
from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches


@pytest.fixture(scope="module")
def showcase() -> TemplateInventory:
    result = analyze_template(build_showcase_template())
    assert result.errors == []
    assert result.inventory is not None
    return result.inventory


def _slide(inventory: TemplateInventory, label: str):
    matches = [
        p for p in inventory.patterns if p.source == "slide" and p.label == label
    ]
    assert matches, (
        f"no slide pattern labelled {label!r}: {[p.label for p in inventory.patterns]}"
    )
    return matches[0]


def _layout(inventory: TemplateInventory, name: str):
    return next(
        p for p in inventory.patterns if p.source == "layout" and p.name == name
    )


# --- layouts -----------------------------------------------------------------


def test_default_layouts_classify_by_placeholder_signature():
    inventory = analyze_template(build_layouts_only_template()).inventory
    assert inventory is not None
    kinds = {p.name: p.kind for p in inventory.patterns}
    assert kinds["Title Slide"] == "cover"
    assert kinds["Title and Content"] == "bullets"
    assert kinds["Section Header"] == "section"
    assert kinds["Two Content"] == "two_columns"
    assert kinds["Comparison"] == "comparison"
    assert kinds["Picture with Caption"] == "image_text"
    assert all(p.labelled_by == "heuristic" for p in inventory.patterns)
    blank = _layout(inventory, "Blank")
    assert not blank.included and blank.exclusion_reason == "no_placeholder"


def test_layout_zones_carry_roles_and_capacity():
    inventory = analyze_template(build_layouts_only_template()).inventory
    assert inventory is not None
    content = _layout(inventory, "Title and Content")
    roles = {z.key: z.role for z in content.zones}
    assert roles == {"ph:0": "title", "ph:1": "body"}
    body = next(z for z in content.zones if z.role == "body")
    # Inherited from the master body style (32 pt in the default template).
    assert body.font_pt == 32
    assert body.chars_per_line and body.max_lines
    assert body.chars_per_line * body.max_lines > 200
    assert body.max_lines and body.max_lines >= 5


# --- example slides -----------------------------------------------------------


def test_agenda_groups_become_a_numbered_list(showcase):
    agenda = _slide(showcase, "Agenda")
    assert agenda.kind == "agenda"
    assert [z.role for z in agenda.zones] == ["title"]
    (group,) = agenda.groups
    assert group.count == 4
    assert group.item_zones == ["number", "heading"]
    assert [z.sample for z in group.items[0]] == [
        "01",
        "Titre du chapitre sur une a deux lignes",
    ]
    assert len(group.item_keys) == 4


def test_kpi_groups_are_figures_with_number_and_label(showcase):
    kpi = _slide(showcase, "Key figures")
    assert kpi.kind == "figures"
    (group,) = kpi.groups
    assert group.count == 4
    assert group.item_zones == ["number", "label"]


def test_sequential_numbers_with_text_make_a_process(showcase):
    steps = _slide(showcase, "Process")
    assert steps.kind == "process"
    (group,) = steps.groups
    assert group.item_zones == ["number", "heading", "body"]


def test_same_font_pairs_read_heading_then_label(showcase):
    team = _slide(showcase, "Team")
    assert team.kind == "team"
    (group,) = team.groups
    assert group.count == 6
    assert group.item_zones == ["heading", "label"]


def test_every_free_text_box_is_a_zone_sample_or_not(showcase):
    free = _slide(showcase, "Free text")
    roles = {z.sample: z.role for z in free.zones}
    assert roles["Cliquez ici pour ajouter du texte"] == "body"
    assert roles["Source: annual report 2025"] == "label"
    assert free.kind == "bullets"


def test_shared_footer_is_never_a_zone(showcase):
    assert FOOTER in showcase.shared_texts
    for pattern in showcase.patterns:
        assert all(z.sample != FOOTER for z in pattern.zones)


def test_shape_names_and_kind_note_override_heuristics(showcase):
    declared = _slide(showcase, "Quarterly review")
    roles = {z.role for z in declared.zones}
    assert roles == {"title", "body"}
    assert all(z.sample != "Do not touch me" for z in declared.zones)
    # The legacy `intent: statement` directive still reads as a quote.
    assert declared.kind == "quote"


def test_library_chart_and_table_are_excluded_by_default(showcase):
    by_reason = {
        p.exclusion_reason: p
        for p in showcase.patterns
        if not p.included and p.source == "slide"
    }
    assert set(by_reason) == {"library", "chart", "table"}
    assert by_reason["chart"].label == "Chart"
    assert by_reason["table"].label == "Table"


def test_library_slides_are_excluded_and_designs_with_icons_kept():
    inventory = analyze_template(build_library_template()).inventory
    assert inventory is not None
    for label in ("Nature icons", "Fonts"):
        sheet = _slide(inventory, label)
        assert not sheet.included and sheet.exclusion_reason == "library"
    for label in ("Our strengths", "Milestones"):
        design = _slide(inventory, label)
        assert design.included and design.exclusion_reason is None
    assert _slide(inventory, "Thank you").included


@pytest.mark.parametrize(
    "title, library",
    [
        ("Fonts", True),
        ("Color palette", True),
        ("Attributions", True),
        ("Instructions for use (premium users)", True),
        ("Use our editable graphic resources...", True),
        ("Contents of this template", True),
        ("Resources", True),
        ("Human resources", False),
        ("3 icônes + texte", False),
        ("Nature icons", False),
        ("Thank you", False),
    ],
)
def test_library_titles(title, library):
    assert is_library_title(title) is library


def test_title_only_over_graphics_is_decoration(showcase):
    divider = _slide(showcase, "Motif")
    assert divider.kind == "other" and not divider.reusable


def test_a_slide_copying_its_layout_stands_for_it(showcase):
    named = [p for p in showcase.patterns if p.name == "Title and Content"]
    assert all(p.source == "slide" for p in named)
    copy = _slide(showcase, "Title and Content (11)")
    assert copy.kind == "bullets"
    assert any("Cliquez ici" in z.sample for z in copy.zones)


def test_a_second_copy_of_one_layout_is_dropped():
    def copy(presentation, title: str):
        slide = presentation.slides.add_slide(presentation.slide_layouts[L_CONTENT])
        cast(Any, slide.shapes.title).text = title
        cast(Any, slide.placeholders[1]).text_frame.text = "Premier point"

    inventory = analyze_template(
        build_slides([lambda p: copy(p, "Un"), lambda p: copy(p, "Deux")])
    ).inventory
    assert inventory is not None
    assert [p.label for p in inventory.patterns if p.name == "Title and Content"] == [
        "Un"
    ]


L_PICTURE_CAPTION = 8


def _picture_slide(presentation, title: str = "Site de Lyon"):
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_PICTURE_CAPTION])
    cast(Any, slide.shapes.title).text = title
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.type.name == "BODY":  # type: ignore[union-attr]
            cast(Any, placeholder).text_frame.text = "Une légende"
    return slide


def test_a_slide_holding_its_layout_slots_stands_for_it_despite_decoration():
    def decorated(presentation):
        slide = _picture_slide(presentation)
        slide.shapes.add_shape(1, Inches(0.2), Inches(0.2), Inches(0.5), Inches(0.5))

    inventory = analyze_template(build_slides([decorated])).inventory
    assert inventory is not None
    named = [p for p in inventory.patterns if p.name == "Picture with Caption"]
    # The empty picture placeholder is a slot, the drawn square decoration.
    assert [p.source for p in named] == ["slide"]
    assert any(z.role == "image" for z in named[0].zones)


def test_a_drawing_that_makes_another_kind_keeps_the_layout():
    def agenda_on_title_only(presentation):
        add_agenda_slide(presentation)

    inventory = analyze_template(build_slides([agenda_on_title_only])).inventory
    assert inventory is not None
    assert _slide(inventory, "Agenda").kind == "agenda"
    assert _layout(inventory, "Title Only").included


def test_placeholders_sharing_an_index_get_distinct_keys_and_a_warning():
    def duplicated(presentation):
        slide = _picture_slide(presentation)
        body = next(
            p
            for p in slide.placeholders
            if p.placeholder_format.type.name == "BODY"  # type: ignore[union-attr]
        )
        twin = copy.deepcopy(body._element)
        slide.shapes._spTree.append(twin)

    inventory = analyze_template(build_slides([duplicated])).inventory
    assert inventory is not None
    pattern = _slide(inventory, "Site de Lyon")
    keys = [z.key for z in pattern.zones]
    assert len(keys) == len(set(keys))
    assert any("share the index" in w for w in pattern.warnings)


def test_the_first_slide_without_body_content_is_the_cover():
    def opening(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
        cast(Any, slide.shapes.title).text = "Revue annuelle du programme"
        for y, text in (
            (2.6, "Comité de direction du groupe"),
            (6.4, "Octobre 2026 édition interne"),
        ):
            box = slide.shapes.add_textbox(
                Inches(0.5), Inches(y), Inches(6), Inches(0.5)
            )
            box.text_frame.text = text

    inventory = analyze_template(build_slides([opening, opening])).inventory
    assert inventory is not None
    slides = sorted(
        (p for p in inventory.patterns if p.source == "slide"),
        key=lambda p: p.source_index,
    )
    assert slides[0].kind == "cover" and slides[1].kind != "cover"


def test_a_title_less_quote_takes_its_quote_as_title():
    def quote(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[L_BLANK])
        for y, h, text in (
            (1.5, 3.5, "Une citation inspirante"),
            (5.3, 0.6, "Une autrice"),
        ):
            box = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(8), Inches(h))
            box.text_frame.text = text
        slide.notes_slide.notes_text_frame.text = "kind: quote"

    inventory = analyze_template(build_slides([quote])).inventory
    assert inventory is not None
    pattern = next(p for p in inventory.patterns if p.source == "slide")
    assert pattern.kind == "quote"
    assert sorted(z.role for z in pattern.zones) == ["subtitle", "title"]


def test_a_content_placeholder_named_image_is_an_image_slot():
    def visual(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[L_CONTENT])
        cast(Any, slide.shapes.title).text = "Un visuel"
        slide.placeholders[1].name = "image"

    inventory = analyze_template(build_slides([visual])).inventory
    assert inventory is not None
    pattern = next(p for p in inventory.patterns if p.source == "slide")
    assert [z.role for z in pattern.zones] == ["title", "image"]


# --- rows of loose same-size shapes -----------------------------------------------


def test_a_row_of_loose_figures_with_labels_is_a_repeat_group():
    inventory = analyze_template(build_slides([add_loose_figures_slide])).inventory
    assert inventory is not None
    pattern = _slide(inventory, "Loose figures")
    assert [z.role for z in pattern.zones] == ["title"]
    group = pattern.groups[0]
    assert group.count == 3 and group.item_zones == ["number", "label"]
    assert pattern.kind == "figures"


def test_two_loose_items_need_a_second_row():
    paired = analyze_template(
        build_slides([lambda p: add_loose_figures_slide(p, count=2)])
    ).inventory
    single = analyze_template(
        build_slides([lambda p: add_loose_figures_slide(p, count=2, labels=False)])
    ).inventory
    assert paired is not None and single is not None
    assert _slide(paired, "Loose figures").groups[0].count == 2
    assert not _slide(single, "Loose figures").groups


def test_aligned_rows_of_cards_are_one_grid():
    def grid(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
        cast(Any, slide.shapes.title).text = "Grid"
        for row in range(2):
            for col in range(3):
                x, y = 0.5 + col * 3.1, 1.8 + row * 2.6
                for dy, h, text in ((0, 0.4, "Titre"), (0.5, 1.2, "Un texte de carte")):
                    box = slide.shapes.add_textbox(
                        Inches(x), Inches(y + dy), Inches(2.6), Inches(h)
                    )
                    box.text_frame.text = text

    inventory = analyze_template(build_slides([grid])).inventory
    assert inventory is not None
    group = _slide(inventory, "Grid").groups[0]
    assert group.count == 6 and len(group.item_zones) == 2


def test_column_bodies_stay_columns(showcase):
    inventory = analyze_template(
        build_slides([add_three_column_bodies_slide])
    ).inventory
    assert inventory is not None
    columns = _slide(inventory, "Three columns")
    assert not columns.groups
    assert [z.role for z in columns.zones].count("body") == 3
    for name in ("Two Content", "Comparison"):
        assert not next(p for p in showcase.patterns if p.name == name).groups


def test_large_free_picture_is_a_replaceable_image(showcase):
    photo = _slide(showcase, "Photo")
    assert [z.role for z in photo.zones] == ["title", "image"]
    assert photo.kind == "image_text"


def test_a_picture_over_a_layout_shape_is_framed():
    """A layout band under the picture shows when the picture is left empty;
    a full-slide background is no frame."""

    def picture_slide(presentation, band: bool):
        layout = presentation.slide_layouts[L_PICTURE]
        slide = presentation.slides.add_slide(layout)
        cast(Any, slide.shapes.title).text = "Une photo"
        slide.placeholders[2].text_frame.text = "Légende"
        picture = slide.placeholders[1]
        box = (
            (picture.left, picture.top, picture.width, picture.height)
            if band
            else (0, 0, presentation.slide_width, presentation.slide_height)
        )
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *box)
        layout.shapes._spTree.append(shape._element)  # moved to the layout

    for band in (True, False):
        inventory = analyze_template(
            build_slides([lambda p, band=band: picture_slide(p, band)])
        ).inventory
        assert inventory is not None
        pattern = next(p for p in inventory.patterns if p.source == "slide")
        (image,) = [z for z in pattern.zones if z.role == "image"]
        assert image.framed is band


def test_decorative_slide_is_not_reusable():
    # Used verbatim, a decorative example would otherwise outrank the closing layout.
    def name_blank_layout_merci(presentation):
        presentation.slide_layouts[6]._element.cSld.set("name", "Merci")

    inventory = analyze_template(
        build_slides([name_blank_layout_merci, add_agenda_slide, add_divider_slide])
    ).inventory
    assert inventory is not None
    merci = _layout(inventory, "Merci")
    assert merci.kind == "closing" and merci.included and merci.reusable
    motif = _slide(inventory, "Motif")
    assert motif.kind == "other" and not motif.reusable


def test_example_slide_placeholders_inherit_their_font_size(showcase):
    agenda = _slide(showcase, "Agenda")
    title = next(z for z in agenda.zones if z.role == "title")
    # Inherited from the master title style (44 pt in the default template),
    # so the capacity the catalogue quotes is the real one.
    assert title.font_pt == 44
    assert title.chars_per_line and title.chars_per_line * (title.max_lines or 1) < 40


def test_notes_without_body_placeholder_do_not_abort_the_analysis(showcase):
    silent = _slide(showcase, "Silent notes")
    assert silent.kind == "bullets"


def test_two_sibling_groups_are_columns_not_a_list(showcase):
    two = _slide(showcase, "Two columns")
    assert two.groups == []
    assert two.kind == "comparison"
    assert sorted(z.role for z in two.zones) == [
        "body",
        "body",
        "heading",
        "heading",
        "title",
    ]


def test_summary_counts(showcase):
    summary = showcase.summary()
    assert summary.layouts == 10  # Title and Content stands as its copy
    assert summary.excluded == 4  # Blank layout + library + chart + table
    assert summary.fixed == 1
    assert (
        summary.patterns
        == summary.layouts + summary.slides - summary.excluded - summary.fixed
    )


def test_unreadable_upload_is_one_error():
    result = analyze_template(b"not a pptx")
    assert result.inventory is None
    assert [e.code for e in result.errors] == ["invalid_upload"]


# --- lexicon -----------------------------------------------------------------------


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Cliquez ici pour modifier le titre", True),
        ("CHANGEZ CE TEXTE", True),
        ("Titre du chapitre sur une à deux lignes", True),
        ("Lorem ipsum dolor", True),
        ("Nom", True),
        ("Revenue grew 12 % in 2025", False),
        ("Texte et images", False),
        ("3 icônes + texte", False),
    ],
)
def test_sample_text_lexicon(text, expected):
    assert is_sample_text(text) is expected


@pytest.mark.parametrize(
    "name,expected",
    [
        ("TextBox 12", None),
        ("Title 1", None),
        ("Image 44", None),
        ("Titre 3", None),
        ("title", "title"),
        ("Sous_titre", "subtitle"),
        ("heading_2", "heading"),
        ("image-main", "image"),
        ("keep", "keep"),
    ],
)
def test_shape_name_declarations(name, expected):
    assert role_from_shape_name(name) == expected


# --- reference template -------------------------------------------------------------

REFERENCE = os.environ.get("PPT_BUILDER_REFERENCE_TEMPLATE")

# label -> (kind, group count, item roles) for the reference corporate deck.
REFERENCE_EXPECTATIONS = {
    "SOMMAIRE": ("agenda", 8, ["number", "heading"]),
    "Indicateurs de performance / chiffres clés": ("figures", 4, ["number", "label"]),
    "3 icônes + texte": ("cards", 3, ["body"]),
    "Planche de processus": ("process", 3, ["number", "heading", "body"]),
    "Timeline jusqu’à 6 ans": ("timeline", 6, ["number", "heading", "body"]),
    "Equipe": ("team", 22, ["heading", "label"]),
}

_reference = pytest.mark.skipif(
    not REFERENCE or not os.path.exists(REFERENCE or ""),
    reason="reference template not provided",
)


@_reference
def test_reference_template_free_form_slides():
    inventory = analyze_template(open(REFERENCE, "rb").read()).inventory  # type: ignore[arg-type]
    assert inventory is not None
    by_label = {}
    for pattern in inventory.patterns:
        by_label.setdefault(pattern.label, pattern)
    for label, (kind, count, roles) in REFERENCE_EXPECTATIONS.items():
        pattern = by_label[label]
        assert pattern.kind == kind, label
        assert pattern.groups[0].count == count, label
        assert pattern.groups[0].item_zones == roles, label
    # A layout or the example slide copying it, whichever stands for the design.
    layouts: dict = {}
    for pattern in inventory.patterns:
        if pattern.included:
            layouts.setdefault(pattern.name, pattern.kind)
    assert layouts["Couverture"] == "cover"
    assert layouts["Slide chapitre"] == "section"
    assert layouts["2 colonnes de texte"] == "two_columns"
    assert layouts["Bullet Points"] == "comparison"
    assert layouts["Contact"] == "closing"
    assert layouts["Merci"] == "closing"
    excluded = {p.exclusion_reason for p in inventory.patterns if not p.included}
    assert excluded == {"chart", "table", "library"}


@_reference
def test_reference_template_has_no_title_only_example_with_real_text():
    """Every text shape of an example slide is a slot, so no slide carrying
    text beside its title is reduced to its title."""
    source = open(REFERENCE, "rb").read()  # type: ignore[arg-type]
    inventory = analyze_template(source).inventory
    assert inventory is not None
    shared = set(inventory.shared_texts)
    for slide in Presentation(io.BytesIO(source)).slides:
        pattern = inventory.pattern(f"slide:{slide.slide_id}")
        if pattern is None or not pattern.included:
            continue
        texts = [
            text
            for shape in slide.shapes
            if shape.has_text_frame
            and (text := cast(Any, shape).text_frame.text.strip())
            and text not in shared
        ]
        slots = len([z for z in pattern.zones if z.role != "image"]) + sum(
            g.count for g in pattern.groups
        )
        if len(texts) > 1:
            assert slots > 1, pattern.label
            assert pattern.kind not in ("cover", "section") or not pattern.groups


# --- every text shape is a slot ------------------------------------------------


def _real_text_slide(presentation):
    """Real corporate text on a cover-typed layout: content, not a cover."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE])
    cast(Any, slide.shapes.title).text = "Revenue grew in every region"
    cast(Any, slide.placeholders[1]).text_frame.text = "Fiscal year 2025"
    for index, text in enumerate(
        (
            "Europe up 12 % on the back of the new rail contracts",
            "Asia up 8 % with two new plants opened in the spring",
            "Americas flat after the end of the legacy programme",
        )
    ):
        box = slide.shapes.add_textbox(
            Inches(0.5 + index * 3.1), Inches(4.5), Inches(2.9), Inches(1.5)
        )
        box.text_frame.text = text
    return slide


def _numbered_slide(
    presentation,
    footer: str = "Group strategy 2026",
    text: str = "Margins recover in the second half",
):
    """A footer repeated on two slides and a slide-number field: chrome."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
    cast(Any, slide.shapes.title).text = "Outlook"
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    body.text_frame.text = text
    slide.shapes.add_textbox(
        Inches(0.5), Inches(7.0), Inches(4), Inches(0.3)
    ).text_frame.text = footer
    number = slide.shapes.add_textbox(Inches(9), Inches(7.0), Inches(0.5), Inches(0.3))
    paragraph = number.text_frame.paragraphs[0]._p
    field = etree.SubElement(paragraph, qn("a:fld"), id="{1}", type="slidenum")
    etree.SubElement(field, qn("a:t")).text = "2"
    return slide


def test_real_text_slide_is_all_slots_and_never_a_cover():
    inventory = analyze_template(build_slides([_real_text_slide])).inventory
    assert inventory is not None
    pattern = _slide(inventory, "Revenue grew in every region")
    samples = {z.sample for z in pattern.zones} | {
        z.sample for g in pattern.groups for item in g.items for z in item
    }
    assert "Fiscal year 2025" in samples
    assert "Asia up 8 % with two new plants opened in the spring" in samples
    assert pattern.kind not in ("cover", "section")
    assert pattern.reusable


def test_twice_repeated_text_and_number_fields_are_not_slots():
    texts = ("Margins recover in the second half", "Cash stays above target")
    inventory = analyze_template(
        build_slides([lambda p, t=t: _numbered_slide(p, text=t) for t in texts])
    ).inventory
    assert inventory is not None
    assert "Group strategy 2026" in inventory.shared_texts
    slides = [p for p in inventory.patterns if p.source == "slide"]
    assert len(slides) == 2
    for pattern, text in zip(slides, texts):
        samples = [z.sample for z in pattern.zones]
        assert "Group strategy 2026" not in samples and "2" not in samples
        assert text in samples


def test_a_link_to_another_slide_does_not_block_copying():
    def linked(presentation):
        agenda = add_agenda_slide(presentation)
        target = _numbered_slide(presentation, footer="Only once")
        agenda.part.relate_to(target.part, RT.SLIDE)

    inventory = analyze_template(build_slides([linked])).inventory
    assert inventory is not None
    agenda = _slide(inventory, "Agenda")
    assert agenda.included and agenda.exclusion_reason is None


def test_analysis_version_is_stamped():
    inventory = analyze_template(build_layouts_only_template()).inventory
    assert inventory is not None and inventory.analysis_version == ANALYSIS_VERSION


# --- capacity: the room the agent is promised is the room the renderer has ---


def _capacity(**frame) -> tuple[int, int]:
    return estimate_capacity(
        Box(x=0, y=0, w=400, h=300),
        18.0,
        "body",
        Inches(13.333).emu,
        Inches(7.5).emu,
        frame=FrameMetrics(**frame),
    )


def test_capacity_pays_the_paragraph_spacing_the_template_declares():
    """A body style with 114 % line spacing and 16 pt before each bullet holds
    barely half the bullets a bare line height would promise."""
    bare = _capacity(line_spacing=1.0, space_before_pt=0.0)
    spaced = _capacity(line_spacing=1.14, space_before_pt=16.0)
    assert spaced[1] * 2 <= bare[1]
    assert spaced[0] * spaced[1] < bare[0] * bare[1]


def test_capacity_takes_the_padding_off_the_box():
    flush = _capacity(lr_inset=0, tb_inset=0)
    padded = _capacity(lr_inset=Inches(1).emu, tb_inset=Inches(1).emu)
    assert padded[0] < flush[0]
    assert padded[1] < flush[1]


def test_capacity_does_not_hand_out_the_room_wrapping_wastes():
    """Characters × lines is not reachable: every line break loses part of a word."""
    one_line = estimate_capacity(
        Box(x=0, y=0, w=400, h=40),
        18.0,
        "title",
        Inches(13.333).emu,
        Inches(7.5).emu,
        frame=FrameMetrics(lr_inset=0, tb_inset=0),
    )
    per_line, lines = _capacity(lr_inset=0, tb_inset=0)
    assert lines > 1
    assert per_line < one_line[0]


def test_the_designers_sample_text_calibrates_a_slot_it_overflows_by_our_ruler():
    """The sample fits its box in PowerPoint: a ruler that says otherwise is
    measuring a wider substitute font, so the line is credited more room."""
    box = Box(x=0, y=0, w=400, h=40)
    size = (Inches(13.333).emu, Inches(7.5).emu)
    flat = FrameMetrics(lr_inset=0, tb_inset=0)
    measured, lines = estimate_capacity(box, 18.0, "title", *size, frame=flat)
    prose = "Les présentations montrent des phrases ordinaires, the usual mix. " * 9
    sample = prose[: int(measured * 1.2)]
    calibrated, _ = estimate_capacity(
        box, 18.0, "title", *size, frame=flat, sample=sample
    )
    short, _ = estimate_capacity(box, 18.0, "title", *size, frame=flat, sample="Hi")
    absurd, _ = estimate_capacity(
        box, 18.0, "title", *size, frame=flat, sample=prose[: measured * 3]
    )
    assert lines == 1
    assert measured < calibrated <= measured / 0.75 + 1
    assert short == measured  # a short sample says nothing about the font
    assert absurd == measured  # far over its box: overflowing text, not evidence


def test_analysis_reads_the_line_spacing_a_layout_declares():
    """The spacing lives on the layout's placeholder, not on the slide's text."""
    plain = analyze_template(build_spaced_template(space_before_pt=0)).inventory
    spaced = analyze_template(build_spaced_template(space_before_pt=24)).inventory
    assert plain is not None and spaced is not None

    def body_lines(inventory) -> int:
        lines = [
            z.max_lines
            for p in inventory.patterns
            for z in p.zones
            if z.role == "body" and z.max_lines
        ]
        assert lines
        return sum(lines)

    assert body_lines(spaced) < body_lines(plain)


def test_a_box_that_does_not_wrap_is_marked_as_growing():
    """Its width follows its text, so no character limit can be quoted for it."""
    inventory = analyze_template(build_spaced_template(space_before_pt=0)).inventory
    assert inventory is not None
    zones = [z for p in inventory.patterns for z in p.zones]
    assert [z for z in zones if z.grows], "the no-wrap heading should be marked"
    assert all(z.chars_per_line for z in zones if not z.grows and z.role != "image")


@_reference
def test_reference_template_offers_every_cards_variant_and_icon_cards_win():
    from fred_capability_ppt_filler.builder.assign import assign
    from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanItem, SlidePlan
    from fred_capability_ppt_filler.builder.render import render_deck
    from fred_capability_ppt_filler.builder.tools import catalogue

    source = open(REFERENCE, "rb").read()  # type: ignore[arg-type]
    inventory = analyze_template(source).inventory
    assert inventory is not None
    # Layouts the deck's own slides hold are gone, so each design is listed once.
    names = {p.name for p in inventory.patterns if p.source == "layout"}
    assert not names & {"Images + 3 colonnes de texte", "3 images + message clé"}
    cards = next(
        line
        for line in catalogue(inventory).splitlines()
        if line.startswith("- cards (")
    )
    assert "3 icônes + texte" in cards and "6 icônes + texte" in cards
    assert cards.count(" | ") >= 4
    # Short heading + text items, no pictures: the icon cards, not the photo cards.
    items = [PlanItem(heading="Agilité", text="Décider vite")] * 2
    cover = SlidePlan(title="Stratégie", subtitle="2027")
    plan = DeckPlan(
        slides=[cover, SlidePlan(title="Nos axes", items=items, kind="cards")]
    )
    patterns = assign(inventory, plan).patterns
    assert patterns[1] is not None and patterns[1].label == "3 icônes + texte"
    deck = Presentation(
        io.BytesIO(render_deck(source, inventory, plan, patterns, {}).pptx)
    )
    slide = deck.slides[1]
    icons = [s for s in slide.shapes if s.shape_type == 13]  # top-level pictures
    assert len(icons) == 2  # the dropped card's icon went with it
    texts = [
        run
        for shape in slide.shapes
        if shape.shape_type == 6
        for inner in cast(Any, shape).shapes
        if inner.has_text_frame
        for run in inner.text_frame.paragraphs[0].runs
    ]
    assert any(run.text == "Agilité" and run.font.bold for run in texts)


# --- Google Slides exports: cards, figures and timelines without groups ----------------


def _google_deck():
    inventory = analyze_template(
        build_slides(
            [
                add_placeholder_cards_slide,
                add_placeholder_two_ideas_slide,
                add_big_figures_slide,
                add_big_number_slide,
                add_numbered_grid_slide,
                add_staircase_timeline_slide,
                add_reused_entry_slide,
                add_paragraph_subtitle_slide,
                add_graphic_sheet_slide,
                add_promo_slide,
                add_empty_slide,
            ]
        )
    ).inventory
    assert inventory is not None
    return {p.source_index: p for p in inventory.patterns if p.source == "slide"}


@pytest.fixture(scope="module")
def google() -> dict:
    return _google_deck()


def _texts(group) -> list:
    return [[z.sample for z in item] for item in group.items]


def test_repeated_title_subtitle_pairs_are_cards(google):
    cards = google[1]
    assert cards.kind == "cards"
    assert [z.role for z in cards.zones] == ["title"]
    group = cards.groups[0]
    assert group.count == 3 and group.item_zones == ["heading", "body"]


def test_two_stacked_pairs_are_headed_columns(google):
    columns = google[2]
    assert columns.kind == "two_columns" and not columns.groups
    roles = sorted(z.role for z in columns.zones)
    assert roles == ["body", "body", "heading", "heading", "title"]


def test_big_figures_pair_with_the_caption_in_their_column(google):
    figures = google[3]
    assert figures.kind == "figures"
    assert figures.groups[0].item_zones == ["number", "label"]
    assert _texts(figures.groups[0]) == [
        ["333,000", "Masse du Soleil rapportee"],
        ["386,000 km", "Distance Terre Lune"],
    ]


def test_a_lone_big_number_is_a_one_item_figures_slide(google):
    figures = google[4]
    assert figures.kind == "figures" and figures.groups[0].count == 1


def test_numbered_texts_on_a_grid_are_cards_not_a_process(google):
    grid = google[5]
    assert grid.kind == "cards"
    assert grid.groups[0].count == 6
    assert grid.groups[0].item_zones == ["number", "body"]


def test_a_staircase_of_dates_is_a_timeline_read_left_to_right(google):
    timeline = google[6]
    assert timeline.kind == "timeline"
    group = timeline.groups[0]
    assert group.count == 5 and group.item_zones == ["number", "body"]
    xs = [item[0].box.x for item in group.items]
    assert xs == sorted(xs)
    # The entry repeated on another slide is still one of the row.
    assert _texts(group)[1][1] == TIMELINE_ENTRY


def test_a_paragraph_sized_subtitle_is_a_body_not_a_cover(google):
    text = google[8]
    assert text.kind == "bullets"
    assert sorted(z.role for z in text.zones) == ["body", "title"]


def test_text_free_graphic_and_promo_slides_are_library(google):
    for number in (9, 10):
        assert google[number].exclusion_reason == "library"
    empty = google[11]
    assert empty.included and empty.kind == "other"


def test_pairing_leaves_the_existing_designs_alone(showcase):
    """Grouped agenda, KPI and step designs keep their groups and kinds."""
    kinds = {p.label: p.kind for p in showcase.patterns if p.source == "slide"}
    assert kinds["Agenda"] == "agenda"
    assert kinds["Process"] == "process"
    assert not _slide(showcase, "Two columns").groups


@pytest.mark.parametrize(
    "text, figure",
    [
        ("333,000", True),
        ("386,000 km", True),
        ("xx%", True),
        ("20XX", True),
        ("+12 %", True),
        ("01", True),
        ("Here is the title", False),
        ("2 planètes sur 8 dans le système", False),
    ],
)
def test_figures_lexicon(text, figure):
    assert is_figure(text) is figure


def test_look_alike_template_slides_are_one_design():
    """Same zones on the same backdrop are one design, whatever their names; a
    large drawn shape or a moved zone makes another one."""

    def drawn(presentation) -> None:
        slide = add_placeholder_two_ideas_slide(presentation)
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(2), Inches(3), Inches(4)
        )

    def moved(presentation) -> None:
        slide = add_placeholder_two_ideas_slide(presentation)
        slide.shapes[1].left += Inches(1)

    inventory = analyze_template(
        build_slides(
            [
                add_placeholder_two_ideas_slide,
                add_placeholder_two_ideas_slide,
                drawn,
                moved,
            ]
        )
    ).inventory
    assert inventory is not None
    first, twin, decorated, shifted = [
        p for p in inventory.patterns if p.source == "slide"
    ]
    assert twin.design == first.id and twin.design_key == first.design_key
    assert decorated.design == "" and shifted.design == ""


_CFB = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504


@pytest.mark.parametrize(
    ("content", "said"),
    [
        (_CFB + "EncryptedPackage".encode("utf-16-le"), "password-protected"),
        (_CFB, "older PowerPoint file (.ppt)"),
        (b"not a presentation", "another kind of file or it is damaged"),
    ],
)
def test_an_unopenable_file_says_why_in_the_author_s_words(content, said):
    (error,) = analyze_template(content).errors
    assert error.code == "invalid_upload" and said in error.message


def test_an_analyser_bug_is_logged_and_reported_without_its_text(monkeypatch, caplog):
    from fred_capability_ppt_filler.builder import analysis

    def _broken(_presentation):
        raise RuntimeError("secret slide text")

    monkeypatch.setattr(analysis, "analyze_presentation", _broken)
    (error,) = analyze_template(build_showcase_template()).errors
    assert error.code == "analysis_failed" and "report it" in error.message
    assert "secret" not in error.message
    assert any(r.exc_info for r in caplog.records)
