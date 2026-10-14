# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Chart intents: the model says which category matters, in which order and
whether a stack shows shares; the theme decides how that looks, and an intent
the chart cannot show is told, never blocking."""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, cast

import pytest
from fred_capability_ppt_filler.builder import tools
from fred_capability_ppt_filler.builder.assign import assign
from fred_capability_ppt_filler.builder.builtin import builtin_template
from fred_capability_ppt_filler.builder.changes import (
    SlideChange,
    apply_changes,
    unchanged,
)
from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanChart
from fred_capability_ppt_filler.builder.render import RenderResult, render_deck
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

COVER = {"title": "Quarterly review", "subtitle": "Sales"}
REVENUE = {
    "type": "column",
    "categories": ["Q1", "Q2", "Q3", "Q4"],
    "series": [{"name": "Revenue", "values": [12, 21, 18, 15]}],
}
TWO = {
    "type": "line",
    "categories": ["Q1", "Q2", "Q3", "Q4"],
    "series": [
        {"name": "2024", "values": [10, 11, 12, 13]},
        {"name": "2025", "values": [12, 14, 17, 20]},
    ],
}
MUTED = '<a:schemeClr val="tx1"><a:alpha val="35000"/></a:schemeClr>'
ACCENT = '<a:schemeClr val="accent1"/>'


@pytest.fixture(scope="module")
def builtin():
    return builtin_template()


def _plan(*charts: Dict[str, Any]) -> DeckPlan:
    slides = [{"title": f"Chart {n}", "chart": c} for n, c in enumerate(charts)]
    return DeckPlan.model_validate({"slides": [COVER, *slides]})


def _render(builtin, *charts: Dict[str, Any]) -> RenderResult:
    template, inventory = builtin
    plan = _plan(*charts)
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems, [p.as_dict() for p in assignment.problems]
    return render_deck(template, inventory, plan, assignment.patterns)


def _chart(result: RenderResult, slide: int = 1) -> Any:
    deck = Presentation(io.BytesIO(result.pptx))
    (frame,) = [s for s in deck.slides[slide].shapes if getattr(s, "has_chart", False)]
    return cast(Any, frame).chart


def _xml(node: Any) -> str:
    return re.sub(r">\s+<", "><", str(node.xml)).strip()


def _points(chart: Any) -> Dict[int, str]:
    """Index → spPr of each data point styled on its own, compacted."""
    ser = chart.plots[0]._element.find(qn("c:ser"))
    return {
        int(dpt.find(qn("c:idx")).get("val")): _xml(dpt.find(qn("c:spPr")))
        for dpt in ser.findall(qn("c:dPt"))
    }


def _intent_notes(result: RenderResult, slide: int = 1) -> List[tuple]:
    return [(n.code, n.field) for n in result.report.built[slide].notes if n.field]


# --- plan ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ({"sort": "descending"}, ("desc", None, None)),
        ({"sort": "Décroissant"}, ("desc", None, None)),
        ({"order": "largest first"}, ("desc", None, None)),
        ({"sort": True}, ("desc", None, None)),
        ({"sort": "croissant"}, ("asc", None, None)),
        ({"sort_by": "as given"}, ("given", None, None)),
        ({"stacked_as": "100%"}, (None, "percent", None)),
        ({"stacking": "pourcentage"}, (None, "percent", None)),
        ({"stacked_as": "absolute"}, (None, "value", None)),
        ({"emphasis": "Q3"}, (None, None, "Q3")),
        ({"highlight": ["Q3"]}, (None, None, "Q3")),
    ],
)
def test_intents_are_read_from_loose_spellings(raw, expected):
    chart = PlanChart.model_validate({**REVENUE, **raw})
    assert (chart.sort, chart.stacked_as, chart.highlight) == expected
    assert chart.ignored_fields == [] and chart.asked_intents == {}


def test_words_that_name_no_intent_are_kept_for_the_note():
    chart = PlanChart.model_validate(
        {**REVENUE, "sort": "sideways", "highlight": ["Q1", "Q2"]}
    )
    assert (chart.sort, chart.highlight) == (None, None)
    assert chart.asked_intents == {"sort": "sideways", "highlight": "Q1, Q2"}
    again = PlanChart.model_validate(chart.model_dump())
    assert again.asked_intents == chart.asked_intents


def test_the_schema_offers_the_intents_in_plain_words():
    props = PlanChart.model_json_schema()["properties"]
    assert props["sort"]["anyOf"][0]["enum"] == ["desc", "asc", "given"]
    assert props["stacked_as"]["anyOf"][0]["enum"] == ["value", "percent"]
    assert "never a colour" in props["highlight"]["description"]
    assert "asked_intents" not in props


# --- rendering ---------------------------------------------------------------------------


def test_a_chart_without_intents_is_drawn_as_before(builtin):
    result = _render(builtin, REVENUE, {**REVENUE, "type": "bar"})
    column = _chart(result)
    assert list(column.plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]
    assert _points(column) == {} and "a:alpha" not in _xml(column._chartSpace)
    bar = _chart(result, 2)
    assert bar.category_axis.reverse_order is False
    assert list(bar.plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]
    assert _intent_notes(result) == []


def test_sort_reorders_categories_with_their_values_and_keeps_the_plan(builtin):
    template, inventory = builtin
    plan = _plan({**REVENUE, "sort": "desc"}, {**REVENUE, "sort": "asc"})
    result = render_deck(
        template, inventory, plan, assign(inventory, plan, "off").patterns
    )
    desc, asc = _chart(result), _chart(result, 2)
    assert list(desc.plots[0].categories) == ["Q2", "Q3", "Q4", "Q1"]
    assert list(desc.plots[0].series[0].values) == [21, 18, 15, 12]
    assert list(asc.plots[0].categories) == ["Q1", "Q4", "Q3", "Q2"]
    assert plan.slides[1].chart is not None
    assert plan.slides[1].chart.categories == ["Q1", "Q2", "Q3", "Q4"]


def test_a_sorted_bar_chart_puts_the_first_of_the_order_on_top(builtin):
    bar = _chart(_render(builtin, {**REVENUE, "type": "bar", "sort": "desc"}))
    # A bar chart draws its first category at the bottom: the largest comes last.
    assert list(bar.plots[0].categories) == ["Q1", "Q4", "Q3", "Q2"]
    assert bar.category_axis.reverse_order is False


def test_sort_is_stable_and_uses_the_first_series_or_the_totals(builtin):
    clustered = {
        "type": "column",
        "categories": ["A", "B", "C", "D"],
        "series": [
            {"name": "X", "values": [5, 9, 5, 1]},
            {"name": "Y", "values": [1, 2, 30, 4]},
        ],
        "sort": "desc",
    }
    stacked = {**clustered, "type": "stacked_column"}
    result = _render(builtin, clustered, stacked)
    built = _chart(result)
    assert list(built.plots[0].categories) == ["B", "A", "C", "D"]
    assert list(built.plots[0].series[1].values) == [2, 1, 30, 4]
    assert list(_chart(result, 2).plots[0].categories) == ["C", "B", "A", "D"]


def test_a_line_chart_keeps_its_order_and_says_so_once(builtin):
    line = {**TWO, "sort": "desc"}
    result = _render(builtin, line, {**line, "type": "area"}, line)
    assert list(_chart(result).plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]
    assert _intent_notes(result) == [("intent_not_shown", "sort")]
    sorts = [n for n in result.report.notes if "chart.sort" in n]
    assert sorts == [
        "Slides 2, 4: chart.sort 'desc' was not applied: a line chart keeps its "
        "categories in the order given (usually time). Nothing to do; drop "
        "chart.sort, or use type 'column' or 'bar' to sort.",
        "Slide 3: chart.sort 'desc' was not applied: an area chart keeps its "
        "categories in the order given (usually time). Nothing to do; drop "
        "chart.sort, or use type 'column' or 'bar' to sort.",
    ]


def test_a_highlighted_category_takes_the_accent_and_the_rest_is_muted(builtin):
    built = _chart(_render(builtin, {**REVENUE, "highlight": "q3"}))
    points = _points(built)
    # The template's own column chart: the emphasised bar keeps its series' accent.
    assert list(points) == [0, 1, 3] and all(MUTED in p for p in points.values())
    assert ACCENT in _xml(built.plots[0]._element.find(f"{qn('c:ser')}/{qn('c:spPr')}"))
    assert "srgbClr" not in _xml(built._chartSpace)
    # The emphasised bar's value label is bold; the others keep the shared labels.
    ser = built.plots[0]._element.find(qn("c:ser"))
    (label,) = ser.findall(f"{qn('c:dLbls')}/{qn('c:dLbl')}")
    assert label.find(qn("c:idx")).get("val") == "2"
    assert 'b="1"' in _xml(label) and '<c:showVal val="1"/>' in _xml(label)
    assert '<c:showVal val="1"/>' in _xml(ser.find(qn("c:dLbls")))


def test_a_highlighted_slice_takes_the_accent(builtin):
    pie = {**REVENUE, "type": "pie", "highlight": "Q2"}
    built = _chart(_render(builtin, pie))
    points = _points(built)
    assert list(points) == [0, 1, 2, 3] and ACCENT in points[1]
    assert all("tx1" in p and ACCENT not in p for i, p in points.items() if i != 1)
    assert "accent2" not in _xml(built._chartSpace)


def test_a_highlighted_series_takes_the_accent_and_the_others_thin_out(builtin):
    three = {
        **TWO,
        "series": [*TWO["series"], {"name": "2026", "values": [1, 2, 3, 4]}],
        "highlight": "2025",
    }
    built = _chart(_render(builtin, three))
    first, second, third = [_xml(s._element) for s in built.plots[0].series]
    assert ACCENT in second and 'w="38100"' in second
    assert MUTED in first and 'w="19050"' in first and '<c:symbol val="none"/>' in first
    assert '<a:alpha val="55000"/>' in third
    assert "accent2" not in _xml(built._chartSpace)


@pytest.mark.parametrize(
    ("chart", "expected"),
    [
        (
            {**REVENUE, "highlight": "Nord"},
            "Slide 2: chart.highlight 'Nord' matches no category (Q1, Q2, Q3, Q4); "
            "built without highlight. To set it, edit chart.highlight to one of those "
            "(a name, never a colour).",
        ),
        (
            {**REVENUE, "highlight": "red"},
            "Slide 2: chart.highlight 'red' matches no category (Q1, Q2, Q3, Q4); "
            "built without highlight. To set it, edit chart.highlight to one of those "
            "(a name, never a colour).",
        ),
        (
            {**TWO, "highlight": "Q3"},
            "Slide 2: chart.highlight 'Q3' is a category, and this chart has 2 series: "
            "only a series can be emphasised. Built without highlight; edit "
            "chart.highlight to one of the series (2024, 2025), or keep one series.",
        ),
        (
            {**REVENUE, "highlight": ["Q1", "Q2"]},
            "Slide 2: chart.highlight 'Q1, Q2' names more than one; built without "
            "highlight. To set it, edit chart.highlight to the one name that matters.",
        ),
        (
            {**REVENUE, "type": "area", "highlight": "Q2"},
            "Slide 2: chart.highlight 'Q2' was not applied: an area chart cannot "
            "emphasise one category. Nothing to do; drop chart.highlight, or use type "
            "'column' to emphasise it.",
        ),
        (
            {**REVENUE, "type": "line", "highlight": "Q2"},
            "Slide 2: chart.highlight 'Q2' was not applied: a line chart cannot "
            "emphasise one category. Nothing to do; drop chart.highlight, or use type "
            "'column' to emphasise it.",
        ),
    ],
)
def test_a_highlight_the_chart_cannot_show_is_told_and_built_without(
    builtin, chart, expected
):
    result = _render(builtin, chart)
    assert expected in result.report.notes
    assert _intent_notes(result) == [("intent_not_shown", "highlight")]
    xml = _xml(_chart(result)._chartSpace)
    assert "a:alpha" not in xml and "c:dPt" not in xml


def test_percent_stacking_builds_a_full_height_stack(builtin):
    stacked = {
        "type": "stacked_column",
        "categories": ["A", "B"],
        "series": [
            {"name": "X", "values": [1, 3]},
            {"name": "Y", "values": [3, 1]},
        ],
        "stacked_as": "percent",
    }
    result = _render(builtin, stacked, {**stacked, "type": "line"})
    built = _chart(result)
    assert built.chart_type == XL_CHART_TYPE.COLUMN_STACKED_100
    assert built.value_axis.tick_labels.number_format == "0%"
    assert _intent_notes(result) == []
    line = _chart(result, 2)
    assert line.chart_type == XL_CHART_TYPE.LINE_MARKERS
    assert _intent_notes(result, 2) == [("intent_not_shown", "stacked_as")]
    assert (
        "Slide 3: chart.stacked_as 'percent' was not applied: it needs type "
        "'stacked_column'; a 'line' chart was built. Nothing to do; drop "
        "chart.stacked_as, or set chart.type to 'stacked_column'."
        in result.report.notes
    )


def test_unread_intent_words_are_told(builtin):
    result = _render(builtin, {**REVENUE, "sort": "sideways", "stacked_as": "tall"})
    assert _intent_notes(result) == [
        ("intent_not_shown", "sort"),
        ("intent_not_shown", "stacked_as"),
    ]
    assert (
        "Slide 2: chart.sort 'sideways' is not an order; built in the order given. "
        "To sort, edit chart.sort to 'desc' (largest first) or 'asc'."
        in result.report.notes
    )
    assert (
        "Slide 2: chart.stacked_as 'tall' is not a stacking; built with the values. "
        "Use 'percent' for stacked columns that each fill 100 %, or drop it."
        in result.report.notes
    )


# --- changes and the tool ----------------------------------------------------------------


def _delivered(inventory, *charts: Dict[str, Any]) -> DeckPlan:
    plan = _plan(*charts)
    assignment = assign(inventory, plan, "off")
    deck, patterns, _ = assignment.deck(plan)
    return deck.model_copy(
        update={
            "delivered": True,
            "slides": [
                s.model_copy(update={"built_on": p.id if p else None})
                for s, p in zip(deck.slides, patterns)
            ],
        }
    )


@pytest.mark.parametrize(
    ("field", "value", "attr", "expected"),
    [
        ("chart.highlight", "Q3", "highlight", "Q3"),
        ("chart.sort", "décroissant", "sort", "desc"),
        ("chart.stacked_as", "percent", "stacked_as", "percent"),
    ],
)
def test_an_intent_is_edited_on_the_same_template_slide(
    builtin, field, value, attr, expected
):
    template, inventory = builtin
    base = _delivered(inventory, REVENUE)
    edit = {"op": "edit", "slide": 2, "field": field, "value": value}
    plan = apply_changes(base, [SlideChange.model_validate(edit)])
    chart = plan.slides[1].chart
    assert chart is not None and getattr(chart, attr) == expected
    kinds = {p.id: p.kind for p in inventory.patterns}
    assert not unchanged(base, plan, kinds)
    assignment = assign(inventory, plan, "off")
    assert not assignment.problems
    pattern = assignment.patterns[1]
    assert pattern is not None and pattern.id == base.slides[1].built_on
    result = render_deck(template, inventory, plan, assignment.patterns)
    if attr == "highlight":
        points = _points(_chart(result))
        assert 2 not in points and MUTED in points[0]


def test_the_tool_names_the_intents_where_charts_are_offered(builtin):
    _, inventory = builtin
    description = tools._description(inventory, "template_slide", 2)
    assert "highlight" in description and "never a colour" in description
    assert '"chart.highlight"' in description
