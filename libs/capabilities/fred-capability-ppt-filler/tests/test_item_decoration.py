# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""What an unused item takes with it: the frames, marks and connectors drawn
for it, on the slide or on its layout."""

from __future__ import annotations

import io
from typing import Any, cast

from fred_capability_ppt_filler.builder import builtin
from fred_capability_ppt_filler.builder.models import Pattern
from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanItem, SlidePlan
from fred_capability_ppt_filler.builder.render import _decoration_fate, render_deck
from fred_capability_ppt_filler.traversal import walk_shapes
from pptx import Presentation


def _box(x: int, y: int, w: int, h: int) -> tuple[int, int, int, int]:
    return (x, y, x + w, y + h)


# Three columns of a row, the third unused (per-mille boxes are enough here).
KEPT = [_box(45, 350, 280, 490), _box(360, 350, 280, 490)]
DROPPED = [_box(676, 350, 280, 490)]


def test_a_frame_around_an_unused_item_and_what_it_holds_go():
    frame = _box(660, 330, 310, 530)
    icon = _box(690, 360, 20, 30)
    assert _decoration_fate(KEPT, DROPPED, [frame, icon]) == ["remove", "remove"]


def test_a_mark_just_above_an_item_follows_it():
    # A numbered tile over each column, outside the column's own box.
    tiles = [_box(45, 253, 37, 66), _box(360, 253, 37, 66), _box(676, 253, 37, 66)]
    assert _decoration_fate(KEPT, DROPPED, tiles) == ["move", "move", "remove"]


def test_a_connector_goes_with_the_unused_neighbour_and_moves_between_kept_ones():
    between_kept = _box(330, 580, 20, 30)
    to_dropped = _box(645, 580, 20, 30)
    assert _decoration_fate(KEPT, DROPPED, [between_kept, to_dropped]) == [
        "move",
        "remove",
    ]


def test_a_line_or_background_across_every_item_stays():
    line = _box(45, 404, 909, 4)
    background = _box(0, 0, 1000, 1000)
    footer_mark = _box(45, 930, 9, 17)
    assert _decoration_fate(KEPT, DROPPED, [line, background, footer_mark]) == [
        "stay",
        "stay",
        "stay",
    ]


def _builtin(name: str) -> Pattern:
    _, inventory = builtin.builtin_template()
    return next(p for p in inventory.patterns if p.name == name)


def _rendered(pattern: Pattern, plan: SlidePlan):
    content, inventory = builtin.builtin_template()
    result = render_deck(content, inventory, DeckPlan(slides=[plan]), [pattern])
    return Presentation(io.BytesIO(result.pptx)).slides[0]


def _named(slide, name: str) -> list[Any]:
    return [p for p in walk_shapes(slide.shapes) if cast(Any, p.shape).name == name]


CARDS = [
    PlanItem(heading="Agents", text="Specialised assistants."),
    PlanItem(heading="Knowledge", text="Indexed libraries."),
    PlanItem(heading="Connectors", text="Drive and Git."),
    PlanItem(heading="Decks", text="Ready to edit."),
]


def test_cards_drawn_on_the_layout_leave_with_their_unused_item():
    slide = _rendered(
        _builtin("Fred Four Cards"), SlidePlan(title="Three cards", items=CARDS[:3])
    )
    # The layout's graphics are on the slide now, the layout's own hidden.
    assert slide._element.get("showMasterSp") == "0"
    assert len(_named(slide, "Card")) == 3
    assert len(_named(slide, "Tile")) == 3
    # The slide number and the brand marks the master drew stay.
    assert len(_named(slide, "Brand mark")) == 3
    assert _named(slide, "Slide number")


def test_a_full_set_of_items_leaves_the_layout_as_it_is():
    slide = _rendered(
        _builtin("Fred Four Cards"), SlidePlan(title="Four cards", items=CARDS)
    )
    assert slide._element.get("showMasterSp") is None
    assert not _named(slide, "Card")


def test_a_timeline_keeps_a_dot_per_milestone_and_each_date_in_its_item():
    milestones = [
        PlanItem(number=f"Q{n} 2026", heading=f"Step {n}", text="One sentence.")
        for n in range(1, 5)
    ]
    slide = _rendered(
        _builtin("Fred Timeline"), SlidePlan(title="Roadmap", items=milestones)
    )
    # A number never breaks: its spaces are non-breaking.
    texts = [
        cast(Any, p.shape).text_frame.text.replace("\u00a0", " ")
        for p in walk_shapes(slide.shapes)
        if getattr(p.shape, "has_text_frame", False)
    ]
    assert [f"Q{n} 2026" for n in range(1, 5)] == [t for t in texts if "2026" in t]
    # The line and four of the five dots.
    assert len(_named(slide, "Decoration")) == 4
    assert len(_named(slide, "Line")) == 1
