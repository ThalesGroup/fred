# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Closing slides: a contact design only when the user gave a contact, a
thank-you design otherwise, and a new deck refused unless it ends on one or the
user opted out (synthetic template, plus the reference corporate template when
`PPT_BUILDER_REFERENCE_TEMPLATE` points at it)."""

from __future__ import annotations

import io
import os
from typing import Any, List, cast

import pytest
from builder_decks import L_TITLE_ONLY, _set_title, _textbox, build_slides
from deck_builders import png_bytes
from fred_capability_ppt_filler.builder import lexicon
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import assign
from fred_capability_ppt_filler.builder.models import Pattern, TemplateInventory
from fred_capability_ppt_filler.builder.plan import DeckPlan
from fred_capability_ppt_filler.builder.render import render_deck
from fred_capability_ppt_filler.builder.tools import _unclosed, offered_kinds
from pptx import Presentation
from pptx.util import Inches

COVER = {"title": "Projet X", "subtitle": "Revue de lancement"}
POINTS = {"title": "Le contexte", "bullets": ["Un marché en croissance", "Deux sites"]}
THANKS = {"title": "Merci", "subtitle": "Questions et échanges"}
EMAIL = "jean.dupont@example.com"
PHONE = "+33 6 12 34 56 78"


def _contact_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
    _set_title(slide, "Contact")
    _textbox(slide.shapes, 3.0, 2.6, 4.0, 0.4, "Nom Prénom")
    _textbox(slide.shapes, 3.0, 3.1, 4.0, 0.4, "Intitulé du poste")
    for y, sample in ((3.9, "Numéro de téléphone"), (4.4, "email")):
        _textbox(slide.shapes, 4.4, y, 2.5, 0.3, sample)
        slide.shapes.add_picture(
            io.BytesIO(png_bytes(8, 8)),
            Inches(4.1),
            Inches(y),
            Inches(0.2),
            Inches(0.3),
        )


def _thanks_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
    _set_title(slide, "Merci")
    _textbox(slide.shapes, 2.0, 4.0, 6.0, 0.6, "Questions et échanges")


def _cover_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    _set_title(slide, "Titre de la présentation")
    slide.placeholders[1].text = "Sous-titre"


def _points_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    _set_title(slide, "Titre")
    slide.placeholders[1].text = "Premier point\nSecond point"


def _template(*closings) -> bytes:
    return build_slides([_cover_slide, _points_slide, *closings])


def _inventory(content: bytes) -> TemplateInventory:
    inventory = analyze_template(content).inventory
    assert inventory is not None
    return inventory


def _closing(inventory: TemplateInventory, last: dict) -> Pattern:
    assignment = assign(
        inventory, DeckPlan.model_validate({"slides": [COVER, POINTS, last]})
    )
    assert not assignment.problems
    pattern = assignment.patterns[-1]
    assert pattern is not None
    return pattern


def _contact_roles(pattern: Pattern) -> List[str]:
    return sorted(z.role for z in pattern.zones if z.role.startswith("contact_"))


@pytest.mark.parametrize(
    "sample, role",
    [
        ("email", "contact_email"),
        ("E-mail", "contact_email"),
        ("prenom.nom@thalesgroup.com", "contact_email"),
        ("Numéro de téléphone", "contact_phone"),
        ("Tél.", "contact_phone"),
        ("+33 1 23 45 67 89", "contact_phone"),
        ("01 23 45 67 89", "contact_phone"),
        ("Nom Prénom", None),
        ("2024 - 2026", None),
        ("Questions et échanges", None),
    ],
)
def test_contact_samples_are_recognised(sample, role):
    assert lexicon.contact_kind(sample) == role


def test_the_contact_lines_of_a_closing_slide_get_contact_roles():
    inventory = _inventory(_template(_contact_slide, _thanks_slide))
    closings = [p for p in inventory.patterns if p.kind == "closing"]
    assert [_contact_roles(p) for p in closings] == [
        ["contact_email", "contact_phone"],
        [],
    ]


def test_a_closing_without_contact_never_takes_the_contact_design():
    inventory = _inventory(_template(_contact_slide, _thanks_slide))
    assert _closing(inventory, THANKS).label == "Merci"
    assert _closing(inventory, {"title": "Merci", "kind": "closing"}).label == "Merci"


def test_a_closing_with_a_contact_takes_the_contact_design():
    inventory = _inventory(_template(_contact_slide, _thanks_slide))
    assert _contact_roles(_closing(inventory, {"title": "Merci", "email": EMAIL}))
    assert _contact_roles(_closing(inventory, {"title": "Contact", "phone": PHONE}))


def test_with_only_a_contact_design_a_closing_without_contact_uses_another_design():
    inventory = _inventory(_template(_contact_slide))
    plan = DeckPlan.model_validate({"slides": [COVER, POINTS, THANKS]})
    assignment = assign(inventory, plan)
    assert not assignment.problems
    pattern = assignment.patterns[-1]
    assert pattern is not None and not _contact_roles(pattern)


def test_the_contact_fills_its_lines_and_an_empty_one_loses_its_icon():
    content = _template(_contact_slide, _thanks_slide)
    inventory = _inventory(content)
    last = {"title": "Contact", "bullets": ["Jean Dupont", "Directeur"], "email": EMAIL}
    plan = DeckPlan.model_validate({"slides": [COVER, POINTS, last]})
    assignment = assign(inventory, plan)
    result = render_deck(content, inventory, plan, assignment.patterns)
    slide = Presentation(io.BytesIO(result.pptx)).slides[-1]
    texts = [cast(Any, s).text_frame.text for s in slide.shapes if s.has_text_frame]
    assert texts == ["Contact", "Jean Dupont", "Directeur", EMAIL]
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    # The mail icon stays beside its line; the phone icon left with its line.
    assert len(pictures) == 1
    assert abs(pictures[0].top - Inches(4.4)) < Inches(0.05)


def test_a_closing_without_a_contact_line_lists_the_contact():
    content = _template(_thanks_slide)
    inventory = _inventory(content)
    last = {"title": "Merci", "bullets": ["Jean Dupont"], "email": EMAIL}
    plan = DeckPlan.model_validate({"slides": [COVER, POINTS, last]})
    assignment = assign(inventory, plan)
    assert not assignment.problems
    result = render_deck(content, inventory, plan, assignment.patterns)
    slide = Presentation(io.BytesIO(result.pptx)).slides[-1]
    text = " ".join(
        cast(Any, s).text_frame.text for s in slide.shapes if s.has_text_frame
    )
    assert EMAIL in text


REFERENCE = os.environ.get("PPT_BUILDER_REFERENCE_TEMPLATE")


@pytest.mark.skipif(
    not REFERENCE or not os.path.exists(REFERENCE or ""),
    reason="reference template not provided",
)
def test_reference_template_picks_thanks_or_contact():
    inventory = _inventory(open(REFERENCE, "rb").read())  # type: ignore[arg-type]
    contacts = [p for p in inventory.patterns if _contact_roles(p)]
    assert contacts and all(p.kind == "closing" for p in contacts)
    assert not _contact_roles(_closing(inventory, THANKS))
    chosen = _closing(inventory, {"title": "Contact", "email": EMAIL, "phone": PHONE})
    assert _contact_roles(chosen) == ["contact_email", "contact_phone"]


# --- a new deck must end on a closing slide --------------------------------------


def _missing(content: bytes, slides: List[dict], opted_out: bool = False) -> bool:
    inventory = _inventory(content)
    plan = DeckPlan.model_validate({"slides": slides})
    return _unclosed(plan, offered_kinds(inventory), opted_out)


def test_a_deck_not_ending_on_a_closing_is_refused():
    assert _missing(_template(_contact_slide, _thanks_slide), [COVER, POINTS])


@pytest.mark.parametrize(
    "last",
    [
        THANKS,
        {"title": "Questions ?"},
        {"title": "Fin", "kind": "closing"},
        {"title": "Contact", "bullets": ["Jean Dupont"], "email": EMAIL},
    ],
)
def test_a_deck_ending_on_a_closing_is_accepted(last):
    template = _template(_contact_slide, _thanks_slide)
    assert not _missing(template, [COVER, POINTS, last])


def test_a_user_who_opted_out_needs_no_closing_slide():
    template = _template(_contact_slide, _thanks_slide)
    assert not _missing(template, [COVER, POINTS], opted_out=True)


def test_a_template_without_a_closing_slide_needs_none():
    assert not _missing(_template(), [COVER, POINTS])
