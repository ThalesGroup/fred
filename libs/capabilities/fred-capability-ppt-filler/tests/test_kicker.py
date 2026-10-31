# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""A kicker (eyebrow line above the title), named so by the template, is no
zone: it neither changes the slide's kind nor shows on the delivered slide."""

from __future__ import annotations

import io
from typing import Any, List, cast

from builder_decks import L_TITLE_ONLY, L_TWO, _save, _textbox
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.models import Pattern, TemplateInventory
from fred_capability_ppt_filler.builder.plan import DeckPlan, SlidePlan
from fred_capability_ppt_filler.builder.render import render_deck
from fred_capability_ppt_filler.traversal import walk_shapes
from pptx import Presentation
from pptx.util import Pt

KICKER = "Quarterly review"


def _template(kicker_name: str | None) -> bytes:
    """The default deck: its Two Content layout, its left placeholder (idx 1)
    named ``kicker_name``, and one Title Only slide with a named line above the
    title."""
    presentation = Presentation()
    if kicker_name:
        presentation.slide_layouts[L_TWO].placeholders[1].name = kicker_name
    slide = presentation.slides.add_slide(presentation.slide_layouts[L_TITLE_ONLY])
    cast(Any, slide.shapes.title).text = "Results"
    _textbox(slide.shapes, 0.5, 0.2, 4.0, 0.3, KICKER, Pt(11), name=kicker_name)
    _textbox(slide.shapes, 0.5, 2.0, 9.0, 4.0, "Lorem ipsum dolor sit amet", Pt(18))
    return _save(presentation)


def _inventory(template: bytes) -> TemplateInventory:
    inventory = analyze_template(template).inventory
    assert inventory is not None
    return inventory


def _by_name(
    inventory: TemplateInventory, name: str, source: str = "layout"
) -> Pattern:
    return next(p for p in inventory.patterns if p.name == name and p.source == source)


def _texts(slide) -> List[str]:
    return [
        cast(Any, p.shape).text_frame.text
        for p in walk_shapes(slide.shapes)
        if getattr(p.shape, "has_text_frame", False)
    ]


def test_named_kicker_is_no_zone():
    inventory = _inventory(_template("kicker"))
    two = _by_name(inventory, "Two Content")
    assert [z.key for z in two.zones] == ["ph:0", "ph:2"]
    slide = _by_name(inventory, "Title Only", "slide")
    assert all(z.sample != KICKER for z in slide.zones)


def test_without_a_kicker_name_the_same_shapes_stay_zones():
    inventory = _inventory(_template(None))
    two = _by_name(inventory, "Two Content")
    assert [z.key for z in two.zones] == ["ph:0", "ph:1", "ph:2"]
    slide = _by_name(inventory, "Title Only", "slide")
    assert any(z.sample == KICKER for z in slide.zones)


def test_french_and_eyebrow_names_are_kickers():
    for name in ("Surtitre", "eyebrow", "Sur-titre 2"):
        two = _by_name(_inventory(_template(name)), "Two Content")
        assert [z.key for z in two.zones] == ["ph:0", "ph:2"], name


def test_kicker_is_removed_from_the_delivered_slide_silently():
    template = _template("kicker")
    inventory = _inventory(template)
    two = _by_name(inventory, "Two Content")
    copied = _by_name(inventory, "Title Only", "slide")
    plan = DeckPlan(
        slides=[
            SlidePlan(title="Revenue", bullets=["Up 12 %", "Costs flat"]),
            SlidePlan(title="Outlook", bullets=["Stable demand"]),
        ]
    )
    result = render_deck(template, inventory, plan, [two, copied])
    deck = Presentation(io.BytesIO(result.pptx))
    from_layout, from_slide = deck.slides
    assert all(ph.placeholder_format.idx != 1 for ph in from_layout.placeholders), (
        "the layout's kicker placeholder is not left empty"
    )
    assert KICKER not in _texts(from_slide)
    assert all(not r.removed for r in result.report.issues())
    assert "Revenue" in _texts(from_layout) and "Outlook" in _texts(from_slide)
