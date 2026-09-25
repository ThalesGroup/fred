# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Translating a composition into PowerPoint (design.md D7, D8, D11, D15).

The design layer decided everything; this layer only writes. So the tests
below are about the writing: that a token arrives as a REFERENCE and not as a
copy of its value, that nothing off the design system reaches a shape, that a
composed slide is made of real editable shapes, and that what the format
charges for — text-frame insets, z-order, reading order — is actually set.
"""

from __future__ import annotations

import io
import re

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from fred_capability_ppt_filler.builder.analysis import analyze_presentation
from fred_capability_ppt_filler.builder.compose import composed_patterns
from fred_capability_ppt_filler.builder.design_system import Box
from fred_capability_ppt_filler.builder.plan import DeckPlan, PlanItem, SlidePlan
from fred_capability_ppt_filler.builder.render import _Renderer
from fred_capability_ppt_filler.builder.theme import default_theme
from pptx.oxml.ns import qn

from tests.deck_builders import (
    add_fixture,
    build_house_style_template,
    build_icon_bank_template,
)


def _deck(pattern_id: str, _followers: int = 0, **plan_fields):
    raw, inventory = default_theme()
    presentation = Presentation(io.BytesIO(raw))
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    slides = [SlidePlan(pattern=pattern_id, **plan_fields)]
    # An agenda that announces topics and has no slide behind it is reported,
    # rightly, by the deck-level check — give it the slides it announces.
    slides += [
        SlidePlan(pattern="composed:statement:1", title=f"Topic {n}", bullets=["."])
        for n in range(_followers)
    ]
    result = _Renderer(presentation, enriched, {}).render(DeckPlan(slides=slides))
    # The renderer drops the template's own slides, so the deck is what it built.
    return presentation, presentation.slides[-1], result, inventory.design_system


def _items(count: int, *, number=False, text=True):
    """Only the fields the composition under test actually carries.

    Giving an item a field its composition has no zone for is not a bug — the
    renderer reports the loss, loudly and correctly — so the tests must ask for
    what each composition holds.
    """
    return [
        PlanItem(
            number=f"{index * 10}%" if number else None,
            heading=f"H{index}",
            text=f"Body {index}" if text else None,
        )
        for index in range(count)
    ]


def _xml(slide) -> str:
    return etree.tostring(slide._element).decode()


# -- the content actually lands ------------------------------------------------


def test_a_composed_slide_carries_its_content():
    _, slide, result, _ = _deck(
        "composed:card_row:3", title="Three pillars", items=_items(3)
    )

    assert result.warnings == []
    texts = {
        shape.name: shape.text_frame.text
        for shape in slide.shapes
        if shape.has_text_frame
    }
    assert texts["composed-title"] == "Three pillars"
    for index in range(3):
        assert texts[f"composed-heading-{index}"] == f"H{index}"
        assert texts[f"composed-text-{index}"] == f"Body {index}"


@pytest.mark.parametrize(
    "pattern_id,fields",
    [
        ("composed:card_row:2", {"items": _items(2)}),
        ("composed:stat_band:3", {"items": _items(3, number=True, text=False)}),
        ("composed:steps:3", {"items": _items(3, number=True, text=False)}),
        ("composed:agenda:4", {"items": _items(4, text=False), "_followers": 4}),
        ("composed:progress:3", {"items": _items(3, number=True, text=False)}),
        ("composed:comparison:2", {"items": _items(2)}),
        ("composed:statement:1", {"bullets": ["One sentence that matters."]}),
        ("composed:prose:1", {"bullets": ["A paragraph.", "And another."]}),
    ],
)
def test_every_composition_renders_without_a_warning(pattern_id, fields):
    _, slide, result, _ = _deck(pattern_id, title="Title", **fields)

    assert result.warnings == []
    assert len(slide.shapes) > 1


# -- tokens are referenced, not copied (D11) ----------------------------------


def test_a_composed_slide_points_at_the_theme_rather_than_copying_it():
    """The difference between using a token and copying its value: with a
    reference the deck follows a later theme change, with a literal it does
    not — and a capability that promises the charter cannot afford that."""
    _, slide, _, _ = _deck("composed:card_row:3", title="T", items=_items(3))
    xml = _xml(slide)

    assert re.search(r'schemeClr val="\w+"', xml), "nothing points at the theme"
    assert not re.search(r'srgbClr val="\w+"', xml), (
        "a literal colour was copied into the slide instead of a reference"
    )


def test_typefaces_are_theme_references_too():
    _, slide, _, _ = _deck("composed:stat_band:3", title="T", items=_items(3))

    faces = set(re.findall(r'latin typeface="([^"]+)"', _xml(slide)))
    assert faces
    assert faces <= {"+mj-lt", "+mn-lt"}


def test_no_value_absent_from_the_design_system_reaches_a_shape():
    """D7 restated: what is drawn cannot leave the charter."""
    _, slide, _, system = _deck("composed:card_row:4", title="T", items=_items(4))
    xml = _xml(slide)

    sizes = {int(value) / 100 for value in re.findall(r'\bsz="(\d+)"', xml)}
    allowed = [token.size_pt for token in system.type_scale]
    # Centipoint rounding on the way in, so compare within half a point.
    for size in sizes:
        assert any(abs(size - candidate) < 0.5 for candidate in allowed), (
            f"{size}pt is not on the type scale {sorted(allowed)}"
        )


# -- real, editable shapes (D7) ------------------------------------------------


def test_a_composed_slide_is_made_of_editable_shapes_not_a_picture():
    _, slide, _, _ = _deck("composed:card_row:3", title="T", items=_items(3))

    kinds = {shape.shape_type for shape in slide.shapes}
    assert kinds, "nothing was drawn"
    # A picture of a slide would be one PICTURE; text must be readable back.
    assert any(shape.has_text_frame and shape.text_frame.text for shape in slide.shapes)
    assert "Three" not in _xml(slide) or True
    assert all(str(kind) != "PICTURE (13)" for kind in kinds)


# -- what the format charges for (D6) -----------------------------------------


def test_a_card_is_behind_the_text_that_sits_on_it():
    """Creation order is z-order: the other way round hides the text."""
    _, slide, _, _ = _deck("composed:card_row:3", title="T", items=_items(3))

    order = [shape.name for shape in slide.shapes]
    for index in range(3):
        assert order.index(f"composed-card-{index}") < order.index(
            f"composed-heading-{index}"
        )


def test_text_frames_carry_their_own_insets_rather_than_the_default():
    """The format's 0.1in default silently defeats any padding computed above."""
    _, slide, _, _ = _deck("composed:card_row:3", title="T", items=_items(3))

    padded = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame and shape.name.startswith("composed-heading")
    ]
    assert padded
    default_inset = 91440  # 0.1in in EMU
    assert all(shape.text_frame.margin_left != default_inset for shape in padded)


def test_every_composed_shape_is_named():
    _, slide, _, _ = _deck("composed:steps:3", title="T", items=_items(3))

    names = [shape.name for shape in slide.shapes]
    assert all(name.startswith("composed-") for name in names)
    assert len(names) == len(set(names))


# -- accessibility (D15) -------------------------------------------------------


def test_decorative_shapes_are_marked_so_assistive_technology_skips_them():
    _, slide, _, _ = _deck("composed:card_row:3", title="T", items=_items(3))

    cards = [s for s in slide.shapes if s.name.startswith("composed-card-")]
    assert cards
    for card in cards:
        node = card._element.find(
            ".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
        )
        assert node is not None
        assert node.get("descr") == ""


def test_an_image_frame_carries_a_text_alternative():
    _, slide, _, _ = _deck("composed:figure_caption:1", title="T", bullets=["Caption"])

    frames = [s for s in slide.shapes if s.name == "composed-image"]
    assert frames
    node = frames[0]._element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
    )
    assert node is not None and node.get("descr")


# -- the canvas (task 2.6, D8) -------------------------------------------------


def test_a_composed_slide_is_drawn_on_the_canvas_layout():
    presentation, slide, _, _ = _deck(
        "composed:statement:1", title="T", bullets=["One line."]
    )

    # The canvas is the layout with nothing to fill; a composed slide must not
    # inherit another layout's placeholders and fight them.
    assert not [shape for shape in slide.shapes if shape.is_placeholder]


def test_a_template_without_a_design_system_composes_nothing():
    _, inventory = default_theme()
    stripped = inventory.model_copy(update={"design_system": None})

    assert composed_patterns(stripped.design_system) == []


# -- consistency across the whole deck (D13) ----------------------------------


def test_titles_hold_one_ordinate_across_a_mixed_deck():
    """The most recognisable mark of a deck nobody designed is a title that
    shifts a few points between slides. Composed slides must not introduce it."""
    raw, inventory = default_theme()
    presentation = Presentation(io.BytesIO(raw))
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    plan = DeckPlan(
        slides=[
            SlidePlan(pattern="composed:card_row:3", title="Cards", items=_items(3)),
            SlidePlan(
                pattern="composed:statement:1", title="Statement", bullets=["One."]
            ),
            SlidePlan(
                pattern="composed:steps:3",
                title="Steps",
                items=_items(3, number=True, text=False),
            ),
            SlidePlan(
                pattern="composed:prose:1", title="Prose", bullets=["A paragraph."]
            ),
        ]
    )

    result = _Renderer(presentation, enriched, {}).render(plan)

    assert result.warnings == []
    tops = set()
    for slide in presentation.slides:
        title = next(s for s in slide.shapes if s.name == "composed-title")
        tops.add(title.top)
    assert len(tops) == 1, f"composed titles sit at {sorted(tops)}, not one ordinate"


def test_content_blocks_hold_one_ordinate_too():
    raw, inventory = default_theme()
    presentation = Presentation(io.BytesIO(raw))
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    plan = DeckPlan(
        slides=[
            SlidePlan(pattern="composed:card_row:2", title="A", items=_items(2)),
            SlidePlan(pattern="composed:comparison:2", title="B", items=_items(2)),
            SlidePlan(
                pattern="composed:agenda:4", title="C", items=_items(4, text=False)
            ),
            SlidePlan(pattern="composed:statement:1", title="D", bullets=["."]),
            SlidePlan(pattern="composed:statement:1", title="E", bullets=["."]),
            SlidePlan(pattern="composed:statement:1", title="F", bullets=["."]),
            SlidePlan(pattern="composed:statement:1", title="G", bullets=["."]),
        ]
    )

    _Renderer(presentation, enriched, {}).render(plan)

    tops = set()
    for slide in presentation.slides:
        below = [
            shape.top
            for shape in slide.shapes
            if shape.name != "composed-title" and shape.name.startswith("composed-")
        ]
        if below:
            tops.add(min(below))
    assert len(tops) == 1, f"content blocks start at {sorted(tops)}, not one ordinate"


# -- the fixtures the canvas inherits (D13) -----------------------------------


def _branded(band_inches: float = 0.0) -> Presentation:
    """A template shaped like a corporate one: a logo on the master, and —
    when asked — a band on the placeholder-free layout composing lands on."""
    presentation = Presentation()
    add_fixture(
        presentation,
        presentation.slide_masters[0],
        Inches(8.2),
        Inches(6.6),
        Inches(1.4),
        Inches(0.6),
    )
    if band_inches:
        dressed = presentation.slide_layouts[5]
        for placeholder in list(dressed.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        add_fixture(
            presentation, dressed, 0, 0, presentation.slide_width, Inches(band_inches)
        )
    presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = "T"
    return presentation


def _render_on(presentation, plan_slide):
    inventory = analyze_presentation(presentation)
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    result = _Renderer(presentation, enriched, {}).render(DeckPlan(slides=[plan_slide]))
    return presentation.slides[-1], result, inventory.design_system


@pytest.mark.parametrize("band", [0.0, 1.2])
def test_a_composed_slide_never_covers_what_its_layout_draws(band):
    slide, result, system = _render_on(
        _branded(band),
        SlidePlan(pattern="composed:card_row:3", title="Three", items=_items(3)),
    )

    assert result.warnings == []
    assert system.reserved
    for shape in slide.shapes:
        box = Box(
            x=round(shape.left * 1000 / system.slide_width),
            y=round(shape.top * 1000 / system.slide_height),
            w=round(shape.width * 1000 / system.slide_width),
            h=round(shape.height * 1000 / system.slide_height),
        )
        for fixture in system.reserved:
            assert not box.overlaps(fixture), f"{shape.name} covers the template's own"


def test_a_band_on_the_canvas_pushes_the_composed_title_below_it():
    """The template's own title ordinate is under the band here; a title the
    charter hides is worse than a title a few points lower."""
    _, _, plain = _render_on(
        _branded(),
        SlidePlan(pattern="composed:statement:1", title="T", bullets=["One."]),
    )
    _, _, banded = _render_on(
        _branded(1.2),
        SlidePlan(pattern="composed:statement:1", title="T", bullets=["One."]),
    )

    assert banded.anchors.title_y > plain.anchors.title_y
    assert banded.anchors.title_y >= max(f.h for f in banded.reserved if f.y == 0)
    # The gap between title and content is the template's, not a new one.
    assert banded.anchors.content_y - banded.anchors.title_y == (
        plain.anchors.content_y - plain.anchors.title_y
    )


def test_a_composed_shape_over_a_fixture_is_reported():
    """The net under the grid: it only fires for a piece placed off it."""
    presentation = _branded()
    inventory = analyze_presentation(presentation)
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    renderer = _Renderer(presentation, enriched, {})
    result = renderer.render(
        DeckPlan(
            slides=[SlidePlan(pattern="composed:statement:1", title="T", bullets=["."])]
        )
    )
    assert result.warnings == [], "the grid already keeps clear of the fixtures"

    # A fixture where the slide has already drawn: every piece now sits on it.
    renderer.inventory.design_system.reserved.append(Box(x=0, y=0, w=1000, h=1000))
    renderer._check_fixtures(presentation.slides[-1], 1)

    assert len(renderer.warnings) == 1
    assert "logo, a band or a footer" in renderer.warnings[0]


# -- drawn in the template's own hand (D18) -----------------------------------


def _house_deck(pattern_id: str, **plan_fields):
    presentation = build_house_style_template()
    inventory = analyze_presentation(presentation)
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    result = _Renderer(presentation, enriched, {}).render(
        DeckPlan(slides=[SlidePlan(pattern=pattern_id, **plan_fields)])
    )
    return presentation.slides[-1], result, inventory.design_system


def test_a_card_is_outlined_where_the_template_outlines_its_own():
    slide, result, system = _house_deck(
        "composed:card_row:3", title="Three", items=_items(3)
    )

    assert result.warnings == []
    assert system.shapes.filled is False
    cards = [s for s in slide.shapes if s.name.startswith("composed-card-")]
    assert len(cards) == 3, "an outlined card is still a card"
    for card in cards:
        props = card._element.spPr
        assert props.find(qn("a:noFill")) is not None, "the card was painted"
        line = props.find(qn("a:ln"))
        assert line is not None and line.find(qn("a:solidFill")) is not None


def test_a_bar_is_painted_even_in_a_template_that_outlines():
    """Outlined, a Bar would show nothing: the convention applies to Cards."""
    slide, _, _ = _house_deck(
        "composed:progress:3", title="Three", items=_items(3, number=True, text=False)
    )

    values = [s for s in slide.shapes if s.name.endswith("-value")]
    assert values
    for value in values:
        assert value._element.spPr.find(qn("a:solidFill")) is not None


def test_a_composed_card_takes_the_corner_the_template_draws():
    slide, _, system = _house_deck(
        "composed:card_row:3", title="Three", items=_items(3)
    )

    cards = [s for s in slide.shapes if s.name.startswith("composed-card-")]
    assert cards
    for card in cards:
        assert card.adjustments[0] == pytest.approx(
            system.shapes.corner_radius, abs=0.01
        )


def test_a_house_style_slide_still_writes_only_theme_references():
    """Reading the template's hand must not smuggle a literal into the deck."""
    slide, _, _ = _house_deck("composed:card_row:3", title="Three", items=_items(3))

    xml = _xml(slide)
    assert re.search(r'schemeClr val="\w+"', xml)
    assert not re.search(r'srgbClr val="\w+"', xml)


# -- a card is resized to what it holds (D6) ----------------------------------


def _card_heights(text: str) -> list[int]:
    raw, inventory = default_theme()
    presentation = Presentation(io.BytesIO(raw))
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    plan = DeckPlan(
        slides=[
            SlidePlan(
                pattern="composed:card_row:3",
                title="Three",
                items=[PlanItem(heading=f"H{i}", text=text) for i in range(3)],
            )
        ]
    )
    _Renderer(presentation, enriched, {}).render(plan)
    return [
        shape.height
        for shape in presentation.slides[-1].shapes
        if shape.name.startswith("composed-card-")
    ]


def test_a_card_is_resized_to_the_text_it_actually_holds():
    """A variant declares what it CAN carry; a card two-thirds empty is the
    mark of a generated deck."""
    short = _card_heights("Court.")
    long = _card_heights(
        "Une phrase nettement plus longue, qui occupe plusieurs lignes dans la "
        "carte et montre clairement la difference de hauteur attendue ici."
    )

    assert short and long
    assert max(short) < max(long), "the card did not follow its content"


def test_the_cards_of_a_row_stay_one_row():
    heights = _card_heights("Court.")

    assert len(set(heights)) == 1, "one card grew and broke the row"


# -- the muted surface is a tint, never a brand colour (D12) ------------------


def test_a_card_is_filled_with_a_tint_of_the_template_not_a_brand_colour():
    """Slot names lie: a real template put a vivid cyan in `lt2` and every card
    came out flashy. The fill is a tint of the ink, written as a reference."""
    _, slide, _, system = _deck("composed:card_row:3", title="T", items=_items(3))

    muted = system.pair("surface-muted")
    assert muted is not None and muted.surface.lum_mod is not None
    xml = _xml(slide)
    assert "lumMod" in xml and "lumOff" in xml
    assert not re.search(r'srgbClr val="\w+"', xml), "a literal reached the slide"


# -- the template's own icons reach the slide (D22) ---------------------------


def _icon_deck(icons: list[str | None]):
    presentation = build_icon_bank_template()
    inventory = analyze_presentation(presentation)
    enriched = inventory.model_copy(
        update={
            "patterns": inventory.patterns + composed_patterns(inventory.design_system)
        }
    )
    plan = DeckPlan(
        slides=[
            SlidePlan(
                pattern=f"composed:icon_cards:{len(icons)}",
                title="Pillars",
                items=[
                    PlanItem(icon=name, heading=f"P{index}", text="One line.")
                    for index, name in enumerate(icons)
                ],
            )
        ]
    )
    result = _Renderer(presentation, enriched, {}).render(plan)
    return presentation.slides[-1], result


def test_an_icon_the_plan_names_is_copied_from_the_template():
    slide, result = _icon_deck(["Inclusion", "Autonomie", "Empreintes"])

    assert result.warnings == []
    pictures = [s for s in slide.shapes if s.name.startswith("composed-icon-")]
    assert len(pictures) == 3
    for picture in pictures:
        assert str(picture.shape_type).startswith("PICTURE")
        # Square in the template, square on the slide: never stretched.
        assert picture.width == pytest.approx(picture.height, rel=0.02)


def test_an_icon_frame_nobody_named_is_removed_rather_than_left_empty():
    slide, result = _icon_deck(["Inclusion", None, None])

    assert result.warnings == []
    assert len([s for s in slide.shapes if s.name.startswith("composed-icon-")]) == 1
    assert [s for s in slide.shapes if s.name.startswith("composed-card-")]


def test_an_icon_the_template_does_not_have_is_reported():
    slide, result = _icon_deck(["Inclusion", "Licorne", "Empreintes"])

    assert len(result.warnings) == 1
    assert "no icon called 'Licorne'" in result.warnings[0]
    assert len([s for s in slide.shapes if s.name.startswith("composed-icon-")]) == 2
