# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""``changes`` against the previous plan: the operations, their errors, the
base plan across tool instances, and the whole deck rebuilt as for `slides`."""

from __future__ import annotations

import io
import json
from typing import Any, List

import pytest
from deck_builders import png_bytes
from fred_capability_ppt_filler.builder import deck_plans
from fred_capability_ppt_filler.builder import tools as tools_mod
from fred_capability_ppt_filler.builder.changes import (
    ChangeError,
    SlideChange,
    apply_changes,
    deck_delta,
)
from fred_capability_ppt_filler.builder.deck_plans import (
    MemoryDeckPlanStore,
    SqlDeckPlanStore,
)
from fred_capability_ppt_filler.builder.plan import DeckPlan
from fred_capability_ppt_filler.builder.tools import BUILD_TOOL_NAME
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
)
from fred_sdk.contracts.runtime import RuntimeServices
from port_fakes import FakeDocs, FakeWorkspace
from pptx import Presentation
from test_build_tool import _REPEATED, COVER, _configured, _ctx, _no_pdf, _tools
from test_builder_render import _texts

BASE = DeckPlan.model_validate(
    {
        "slides": [
            COVER,
            {"title": "Agenda", "bullets": ["One", "Two"]},
            {
                "title": "Figures",
                "items": [
                    {"number": "12", "label": "sites"},
                    {"number": "3k", "label": "staff"},
                ],
            },
            {"title": "Thanks", "bullets": ["Questions"]},
        ],
        "file_name": "base",
    }
)


def _apply(*changes: Any) -> DeckPlan:
    return apply_changes(BASE, [SlideChange.model_validate(c) for c in changes])


def _problems(*changes: Any) -> List[str]:
    with pytest.raises(ChangeError) as caught:
        _apply(*changes)
    return caught.value.problems


def _titles(plan: DeckPlan) -> List[Any]:
    return [s.title for s in plan.slides]


# --- operations -------------------------------------------------------------------------


def test_each_operation_numbers_on_the_previous_plan():
    plan = _apply(
        {"op": "insert", "after": 0, "slides": [{"title": "Intro"}]},
        {"op": "replace", "slide": 2, "slides": [{"title": "A"}, {"title": "B"}]},
        {"op": "remove", "slide": 3},
        {"op": "insert", "after": 3, "slides": [{"title": "X"}]},
        {"op": "insert", "after": 3, "slides": [{"title": "Y"}]},
        {"op": "edit", "slide": 4, "field": "title", "value": "Merci"},
    )
    assert _titles(plan) == ["Intro", "Kick-off", "A", "B", "X", "Y", "Merci"]
    # Untouched slides are the base's own objects, byte-identical in the plan.
    assert plan.slides[1] is BASE.slides[0]
    assert plan.file_name == "base"


def test_edit_paths_set_append_and_remove_elements():
    plan = _apply(
        {"op": "edit", "slide": 2, "field": "bullets[1]", "value": "Deux"},
        {"op": "edit", "slide": 2, "field": "bullets[2]", "value": "Trois"},
        {"op": "edit", "slide": 3, "field": "items[1].label", "value": "experts"},
        {"op": "edit", "slide": 4, "field": "bullets[0]", "value": None},
    )
    assert plan.slides[1].bullets == ["One", "Deux", "Trois"]
    assert plan.slides[2].items[1].label == "experts"
    assert plan.slides[2].items[1].number == "3k"
    assert plan.slides[3].bullets == []
    # The edited slide is coerced as any slide: a list sent as one string.
    plan = _apply({"op": "edit", "slide": 4, "field": "bullets", "value": "a\nb"})
    assert plan.slides[3].bullets == ["a", "b"]


def test_a_table_cell_and_a_new_table_are_edited():
    plan = _apply(
        {
            "op": "edit",
            "slide": 2,
            "field": "table",
            "value": {"header": ["A", "B"], "rows": [["1", "2"]]},
        }
    )
    base = plan.model_copy(update={"slides": plan.slides})
    edited = apply_changes(
        base,
        [
            SlideChange.model_validate(
                {"op": "edit", "slide": 2, "field": "table.rows[0][1]", "value": "9"}
            )
        ],
    )
    table = edited.slides[1].table
    assert table is not None and table.rows == [["1", "9"]]


def test_small_model_shapes_are_read():
    plan = _apply(
        {"op": "Delete", "slide": "4"},
        {"op": "set", "slide": 1.0, "field": "title", "value": "K"},
    )
    assert _titles(plan) == ["K", "Agenda", "Figures"]
    change = SlideChange.model_validate(
        {"op": "replace", "slide": 2, "with": {"title": "W"}}
    )
    assert change.slides[0].title == "W" and change.slide == 2
    # A slide object under `slide` is the content; the number is then missing.
    assert _problems({"op": "replace", "slide": {"title": "W"}}) == [
        "changes[0] (replace): replace needs `slide`: the slide's number in the previous deck."
    ]
    # A validated change dumped and read again (the deferred tool) keeps its flags.
    again = SlideChange.model_validate(change.model_dump())
    assert again.has_value is False and again.unknown == []


# --- errors -----------------------------------------------------------------------------


def test_every_bad_change_is_reported_with_its_remedy():
    problems = _problems(
        {"op": "remove", "slide": 14},
        {"op": "move", "slide": 2},
        {"op": "insert", "slides": [{"title": "X"}]},
        {"op": "replace", "slide": 2},
        {"op": "edit", "slide": 2},
        {"op": "edit", "slide": "last", "field": "title", "value": "x"},
        {"op": "remove", "slide": 1, "colour": "red"},
    )
    assert problems == [
        "changes[0] (remove): slide 14 does not exist; the previous deck has 4 slides. "
        "Use a number from 1 to 4.",
        "changes[1] (move): 'move' is not an operation: use replace, insert, remove or edit.",
        "changes[2] (insert): insert needs `after`: the slide number the new slides follow "
        "(0 = first).",
        "changes[3] (replace): replace needs `slides`: the new slide objects.",
        "changes[4] (edit): edit needs `field`, e.g. 'title' or 'items[2].label'.",
        "changes[4] (edit): edit needs `value`: the new text, list or object (null removes).",
        "changes[5] (edit): `slide` 'last' is not a number: give the slide's number.",
        "changes[6] (remove): unknown key 'colour': a change takes op, slide, after, "
        "slides, field and value; drop it.",
    ]


def test_a_new_slide_written_in_the_change_itself_is_read_as_its_slide():
    """Live: the model sent `kind`, `title` and `items` beside `op`."""
    items = [{"heading": "SIRDE", "text": "Radars"}]
    flat = {"op": "replace", "slide": 2, "kind": "cards", "title": "Projets"}
    plan = apply_changes(BASE, [SlideChange.model_validate({**flat, "items": items})])
    replaced = plan.slides[1]
    assert (replaced.kind, replaced.title) == ("cards", "Projets")
    assert replaced.items[0].heading == "SIRDE" and not replaced.ignored_fields
    inserted = SlideChange.model_validate(
        {"op": "insert", "after": 0, "title": "Intro", "body": "Hello"}
    )
    assert [(s.title, s.bullets) for s in inserted.slides] == [("Intro", ["Hello"])]


def test_slide_fields_beside_slides_or_on_another_op_say_where_they_go():
    problems = _problems(
        {"op": "replace", "slide": 2, "title": "A", "slides": [{"title": "B"}]},
        {"op": "edit", "slide": 2, "title": "C"},
        {"op": "remove", "slide": 3, "title": "D"},
    )
    assert problems == [
        "changes[0] (replace): title sent beside `slides`: it is unclear which new slide "
        "they belong to; put them inside that slide's entry of `slides`.",
        "changes[1] (edit): 'title' is a slide field: write it as `field` 'title' with its "
        "`value`.",
        "changes[1] (edit): edit needs `field`, e.g. 'title' or 'items[2].label'.",
        "changes[1] (edit): edit needs `value`: the new text, list or object (null removes).",
        "changes[2] (remove): 'title' is a slide field, which remove does not take: drop it.",
    ]


def test_conflicting_changes_and_bad_paths_are_reported():
    problems = _problems(
        {"op": "replace", "slide": 2, "slides": [{"title": "A"}]},
        {"op": "remove", "slide": 2},
        {"op": "edit", "slide": 2, "field": "title", "value": "x"},
        {"op": "edit", "slide": 3, "field": "items[5].label", "value": "x"},
        {"op": "edit", "slide": 3, "field": "colour", "value": "x"},
        {"op": "edit", "slide": 4, "field": "title[0]", "value": "x"},
    )
    assert problems == [
        "changes[1] (remove): slide 2 is already changed by changes[0]: send one replace "
        "or remove per slide, and put edits in the replacement.",
        "changes[2] (edit): slide 2 is replaced or removed by changes[0]: put the edit in "
        "the replacement slide instead.",
        "changes[3] (edit): slide 3: 'items' has 2 entries (items[0] to items[1]).",
        "changes[4] (edit): slide 3: 'colour' is not a field of a slide (it takes takeaway, "
        "source, title, subtitle, number, bullets, columns, items, image, images, table, chart, charts, notes, kind, repeat_layout).",
        "changes[5] (edit): slide 4: 'title' is not a list: write 'title' without [0].",
    ]
    assert _problems(
        {"op": "edit", "slide": 2, "field": "bullets[5]", "value": "x"}
    ) == [
        "changes[0] (edit): slide 2: 'bullets' has 2 entries (bullets[0] to bullets[1]); "
        "give an index up to 2 (that one adds an entry)."
    ]


# --- the tool ---------------------------------------------------------------------------


_STEPS = [{"number": str(i), "heading": f"Step {i}", "text": "x"} for i in range(7)]
_TOO_LONG = {"title": "Method", "kind": "process", "items": _STEPS}
_DECK = [COVER, {"title": "Agenda", "bullets": ["One", "Two"]}, _TOO_LONG]


class _Turn:
    """One turn's tools (a fresh tool instance), recording what was rendered."""

    def __init__(self, monkeypatch, workspace: FakeWorkspace, docs=None) -> None:
        self.rendered: List[DeckPlan] = []
        real = tools_mod.render_deck

        def recording(template, inventory, plan, patterns, *rest):
            self.rendered.append(plan)
            return real(template, inventory, plan, patterns, *rest)

        monkeypatch.setattr(tools_mod, "render_deck", recording)
        monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
        config, assets = _configured()
        self.build = _tools(
            _ctx(config, assets=assets, docs=docs, workspace=workspace)
        )[BUILD_TOOL_NAME].coroutine


def _payload(text: str) -> dict:
    return json.loads(text[text.index("{") :])


def _dumps(plan: DeckPlan) -> List[str]:
    # `rendered` is the delivered file's hash, set once a deck is delivered.
    return [s.model_dump_json(exclude={"rendered"}) for s in plan.slides]


@pytest.mark.asyncio
async def test_a_partial_fix_after_needs_revision_rebuilds_the_whole_deck(monkeypatch):
    workspace = FakeWorkspace()
    turn = _Turn(monkeypatch, workspace)
    text, result = await turn.build(slides=_DECK)
    assert not result.is_error and _payload(text)["slides"][0]["slide"] == 3
    assert "call build_presentation again with only `changes` for these slides" in text
    split = [
        {"title": "Method 1/2", "kind": "process", "items": _STEPS[:3]},
        {"title": "Method 2/2", "kind": "process", "items": _STEPS[3:6]},
    ]
    text, result = await turn.build(
        changes=[{"op": "replace", "slide": 3, "slides": split}]
    )
    assert not result.is_error, text
    assert text.startswith(tools_mod._DELIVERED)
    (plan,) = turn.rendered
    assert [s.title for s in plan.slides] == [
        "Kick-off",
        "Agenda",
        "Method 1/2",
        "Method 2/2",
    ]
    deck = Presentation(io.BytesIO(workspace.writes[-1]["content"]))
    assert len(deck.slides) == 4


@pytest.mark.asyncio
async def test_a_flattened_replace_is_built_as_its_slide(monkeypatch, deck_plan_store):
    workspace = FakeWorkspace()
    turn = _Turn(monkeypatch, workspace)
    deck = [COVER, {"title": "Agenda", "bullets": ["One", "Two"]}, {"title": "Contact"}]
    text, result = await turn.build(slides=deck)
    assert not result.is_error, text
    text, result = await _Turn(monkeypatch, workspace).build(
        changes=[{"op": "replace", "slide": 2, "title": "Plan", "bullets": ["A", "B"]}]
    )
    assert not result.is_error, text
    stored = DeckPlan.model_validate_json(deck_plan_store.plans[("sess-9", "inst-7")])
    assert (stored.slides[1].title, stored.slides[1].bullets) == ("Plan", ["A", "B"])


def test_a_replacement_inherits_the_replaced_slides_template_slide():
    base = DeckPlan(
        slides=[
            s.model_copy(update={"built_on": f"slide:{n}"})
            for n, s in enumerate(BASE.slides)
        ]
    )
    two = [{"title": "A", "bullets": ["a"]}, {"title": "B", "bullets": ["b"]}]
    plan = apply_changes(
        base, [SlideChange.model_validate({"op": "replace", "slide": 4, "slides": two})]
    )
    assert [s.built_on for s in plan.slides] == [
        "slide:0",
        "slide:1",
        "slide:2",
        "slide:3",
        None,
    ]


@pytest.mark.asyncio
async def test_a_text_replace_keeps_the_replaced_slides_design(
    monkeypatch, deck_plan_store
):
    """Live: a closing slide replaced only to change its text moved to another
    template slide. A replacement of the same kind that fits stays where it was."""
    workspace = FakeWorkspace()
    deck = [
        COVER,
        {"title": "Part A", "kind": "section"},
        {"title": "Points", "bullets": ["a"]},
    ]
    text, result = await _Turn(monkeypatch, workspace).build(slides=deck)
    assert not result.is_error, text
    key = ("sess-9", "inst-7")
    stored = DeckPlan.model_validate_json(deck_plan_store.plans[key])
    built_on = stored.slides[1].built_on
    assert built_on in ("layout:2", "layout:5")
    # Pin it on the other section variant, as an earlier build could have chosen.
    other = "layout:5" if built_on == "layout:2" else "layout:2"
    stored.slides[1].built_on = other
    deck_plan_store.plans[key] = stored.model_dump_json()

    used: List[Any] = []
    real = tools_mod.render_deck

    def recording(template, inventory, plan, patterns, *rest):
        used.append([p.id for p in patterns])
        return real(template, inventory, plan, patterns, *rest)

    monkeypatch.setattr(tools_mod, "render_deck", recording)
    replace = {
        "op": "replace",
        "slide": 2,
        "slides": [{"title": "Part B", "kind": "section"}],
    }
    text, result = await _Turn(monkeypatch, workspace).build(changes=[replace])
    assert not result.is_error, text
    assert used[-1][1] == other
    # Another kind is a different slide: the pin does not hold it.
    text, _ = await _Turn(monkeypatch, workspace).build(
        changes=[{**replace, "slides": [{"title": "Part B", "kind": "quote"}]}]
    )
    assert used[-1][1] not in ("layout:2", "layout:5")


@pytest.mark.asyncio
async def test_a_later_turn_edit_keeps_every_other_slide_identical(
    monkeypatch, deck_plan_store
):
    workspace = FakeWorkspace()
    first = _Turn(monkeypatch, workspace)
    deck = [COVER, {"title": "Agenda", "bullets": ["One", "Two"]}, {"title": "Contact"}]
    text, result = await first.build(slides=deck, file_name="Kick off")
    assert not result.is_error, text
    (delivered,) = first.rendered
    assert list(deck_plan_store.plans) == [("sess-9", "inst-7")]

    # The next turn is another tool instance, possibly on another pod: the base
    # comes from the store.
    second = _Turn(monkeypatch, workspace)
    text, result = await second.build(
        changes=[
            {
                "op": "edit",
                "slide": 3,
                "field": "bullets",
                "value": ["+33 1 23 45 67 89", "contact@acme.fr"],
            }
        ]
    )
    assert not result.is_error, text
    (edited,) = second.rendered
    assert _dumps(edited)[:2] == _dumps(delivered)[:2]
    assert edited.slides[2].bullets == ["+33 1 23 45 67 89", "contact@acme.fr"]
    assert edited.file_name == "Kick off" and "Kick off.pptx" in text


@pytest.mark.asyncio
async def test_numbers_in_a_revision_are_the_deck_the_changes_apply_to(monkeypatch):
    turn = _Turn(monkeypatch, FakeWorkspace())
    await turn.build(slides=_DECK)
    text, _ = await turn.build(
        changes=[
            {"op": "insert", "after": 1, "slides": [{"title": "Why", "bullets": ["a"]}]}
        ]
    )
    (problem,) = _payload(text)["slides"]
    # Slide 4 of the changed plan, numbered 3 as the corrected resend numbers it;
    # still the slide refused before.
    assert problem["slide"] == 3
    assert problem["repeated"].startswith("you resent this slide unchanged")
    # A slide the changes sent is pointed to its change, not to a number.
    text, _ = await turn.build(
        changes=[{"op": "insert", "after": 1, "slides": [_TOO_LONG]}]
    )
    numbers = [p["slide"] for p in _payload(text)["slides"]]
    assert numbers == ["new, sent in changes[0]: fix it there", 3]


@pytest.mark.asyncio
async def test_an_insert_at_the_start_goes_after_the_cover(monkeypatch):
    """The eval turn: `after` 0 put a Sommaire before the cover, refused three
    times as a cover with no place for its items."""
    turn = _Turn(monkeypatch, FakeWorkspace())
    await turn.build(slides=[COVER, {"title": "Agenda", "bullets": ["One", "Two"]}])
    text, result = await turn.build(
        changes=[
            {"op": "insert", "after": 0, "slides": [{"title": "Why", "bullets": ["a"]}]}
        ]
    )
    assert not result.is_error and text.startswith(tools_mod._DELIVERED), text
    assert [s.title for s in turn.rendered[-1].slides] == ["Kick-off", "Why", "Agenda"]
    assert "went after the cover (slide 1)" in text


def test_the_delta_names_changed_added_and_removed_slides():
    after = _apply(
        {"op": "edit", "slide": 2, "field": "title", "value": "Plan"},
        {"op": "remove", "slide": 3},
        {"op": "insert", "after": 4, "slides": [{"title": "Next"}]},
    )
    line = deck_delta(BASE, after, {})
    assert "previous deck (4 slides), slides of this file changed: 2; added: 4" in line
    assert "removed: 3 of the previous deck; 2 unchanged." in line
    assert "none changed: tell the user the deck is the same" in deck_delta(
        BASE, BASE, {}
    )


def test_the_delta_names_the_slide_a_resend_dropped():
    """Seen in the eval: a resend without slide 2 that also retitled slides 3
    and 4 read 'removed: 4' (paired by position), not the slide that went."""
    before = DeckPlan.model_validate(
        {
            "slides": [
                COVER,
                {"title": "Calendar", "bullets": ["Scoping in Q1", "Design in Q2"]},
                {"title": "Budget (1/2)", "bullets": ["Licences 2.4 M", "Training"]},
                {"title": "Budget (2/2)", "bullets": ["Savings 2.8 M", "ROI 3 years"]},
                {"title": "Thanks", "bullets": ["Questions"]},
            ]
        }
    )
    after = before.model_copy(
        update={
            "slides": [
                before.slides[0],
                before.slides[2].model_copy(update={"title": "Budget"}),
                before.slides[3].model_copy(update={"title": "Budget, savings"}),
                before.slides[4],
            ]
        }
    )
    line = deck_delta(before, after, {})
    assert "changed: 2, 3; removed: 2 of the previous deck; 2 unchanged." in line


def test_the_delta_compares_what_the_slides_render():
    """Seen in the eval: chart.show set to what the chart already showed read
    'changed: 9' on an identical file."""
    shown = [
        s.model_copy(update={"rendered": f"h{n}"}) for n, s in enumerate(BASE.slides)
    ]
    before = BASE.model_copy(update={"slides": shown})
    after = before.model_copy(
        update={
            "slides": [*shown[:2], shown[2].model_copy(update={"title": "F"}), shown[3]]
        }
    )
    assert "none changed: tell the user the deck is the same" in deck_delta(
        before, after, {}
    )
    redrawn = after.model_copy(
        update={
            "slides": [
                *shown[:2],
                after.slides[2].model_copy(update={"rendered": "x"}),
                shown[3],
            ]
        }
    )
    assert "changed: 3; 3 unchanged." in deck_delta(before, redrawn, {})


@pytest.mark.asyncio
async def test_a_change_only_the_speaker_notes_show_is_said_so(monkeypatch):
    """Seen in the eval: a takeaway edit on a slide that does not show it was
    described to the user as a new key message on the slide."""
    turn = _Turn(monkeypatch, FakeWorkspace())
    agenda = {"title": "Agenda", "bullets": ["a"], "takeaway": "We start now."}
    await turn.build(slides=[COVER, agenda, {"title": "Next", "bullets": ["b"]}])
    text, _ = await turn.build(
        changes=[
            {"op": "edit", "slide": 2, "field": "takeaway", "value": "We cut costs."}
        ]
    )
    assert (
        "changed in the speaker notes only, the slide itself looks the same: 2; "
        "2 unchanged." in text
    )


@pytest.mark.asyncio
async def test_a_slide_moved_by_a_removal_renders_unchanged(monkeypatch):
    turn = _Turn(monkeypatch, FakeWorkspace())
    deck = [COVER, *({"title": f"Part {n}", "bullets": [f"b{n}"]} for n in range(1, 4))]
    await turn.build(slides=deck)
    text, _ = await turn.build(changes=[{"op": "remove", "slide": 2}])
    assert "slides of this file removed: 2 of the previous deck; 3 unchanged." in text


def test_the_delta_counts_the_points_a_recast_slide_lost():
    """Seen in the eval: a narrative slide recast as a process over two
    slides silently dropped one of its seven points."""
    points = [f"Point {n}" for n in range(1, 8)]
    before = DeckPlan.model_validate(
        {
            "slides": [
                {**COVER, "built_on": "cover"},
                {"title": "Change", "bullets": points, "built_on": "text"},
                {"title": "Same", "bullets": ["a", "b"], "built_on": "text"},
            ]
        }
    )
    steps = [{"heading": p} for p in points[:6]]
    after = DeckPlan.model_validate(
        {
            "slides": [
                before.slides[0].model_dump(),
                {"title": "Change", "items": steps[:3], "built_on": "process"},
                {"title": "Change", "items": steps[3:], "built_on": "process"},
                {"title": "Same", "bullets": ["a"], "built_on": "text"},
            ]
        }
    )
    line = deck_delta(before, after, {}, numbers=[1, 2, 2, 3])
    # Slide 4 keeps its design: one bullet less is an edit, not a recast.
    assert "; fewer points where the design changed: slide 2 (7 → 6)." in line


def test_a_whole_deck_resent_for_a_small_edit_is_pointed_to_changes():
    """Most revisions resent every slide for a one-slide edit: the result says so."""
    after = _apply({"op": "edit", "slide": 2, "field": "title", "value": "Plan"})
    assert "send only `changes` next time" in deck_delta(BASE, after, {}, resent=True)
    assert "next time" not in deck_delta(BASE, after, {})
    rewritten = _apply(
        *(
            {"op": "edit", "slide": n, "field": "title", "value": f"T{n}"}
            for n in (2, 3, 4)
        )
    )
    assert "next time" not in deck_delta(BASE, rewritten, {}, resent=True)


@pytest.mark.asyncio
async def test_a_revision_result_says_which_slides_changed(monkeypatch):
    """Revision replies claimed edits the deck did not hold: the result now
    says which slides differ from the deck the user had."""
    turn = _Turn(monkeypatch, FakeWorkspace())
    text, _ = await turn.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    assert "Compared with the previous deck" not in text
    text, _ = await turn.build(
        changes=[{"op": "edit", "slide": 2, "field": "title", "value": "Plan"}]
    )
    assert "slides of this file changed: 2; 1 unchanged." in text


@pytest.mark.asyncio
async def test_a_small_edit_resent_as_a_whole_deck_is_built_and_pointed_to_changes(
    monkeypatch,
):
    turn = _Turn(monkeypatch, FakeWorkspace())
    deck = [COVER, *({"title": f"Part {n}", "bullets": ["a"]} for n in range(1, 5))]
    await turn.build(slides=deck)
    deck[2] = {"title": "Plan", "bullets": ["a"]}
    text, _ = await turn.build(slides=deck)
    assert text.startswith(tools_mod._DELIVERED), text
    assert "changed: 3; 4 unchanged." in text and "`changes` next time" in text


@pytest.mark.asyncio
async def test_the_same_changes_resent_are_never_applied_twice(
    monkeypatch, deck_plan_store
):
    turn = _Turn(monkeypatch, FakeWorkspace())
    await turn.build(slides=_DECK)
    insert = [
        {"op": "insert", "after": 1, "slides": [{"title": "Why", "bullets": ["a"]}]}
    ]
    await turn.build(changes=insert)
    text, _ = await turn.build(changes=insert, file_name="again")
    (problem,) = _payload(text)["slides"]
    # Applied twice, the refused slide would be slide 4 of the base.
    assert problem["slide"] == 3
    # Nothing was delivered: nothing is stored for a later turn.
    assert deck_plan_store.plans == {}


@pytest.mark.asyncio
async def test_a_refused_plan_is_never_the_next_turns_base(
    monkeypatch, deck_plan_store
):
    workspace = FakeWorkspace()
    first = _Turn(monkeypatch, workspace)
    await first.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    stored = dict(deck_plan_store.plans)
    text, result = await first.build(
        changes=[{"op": "insert", "after": 2, "slides": [_TOO_LONG]}]
    )
    assert not result.is_error and deck_plan_store.plans == stored
    # The next turn changes the delivered deck, not the refused plan.
    second = _Turn(monkeypatch, workspace)
    text, result = await second.build(changes=[{"op": "remove", "slide": 2}])
    assert not result.is_error, text
    assert [s.title for s in second.rendered[-1].slides] == ["Kick-off"]


@pytest.mark.asyncio
async def test_a_change_to_a_delivered_deck_is_never_asked_to_vary(monkeypatch):
    """The later turn's history no longer shows the delivered deck (trimmed):
    its stored base still says the conversation has one, so only the note."""
    workspace = FakeWorkspace()
    first = _Turn(monkeypatch, workspace)
    await first.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    second = _Turn(monkeypatch, workspace)
    add = [{"op": "insert", "after": 2, "slides": _REPEATED[1:]}]
    text, result = await second.build(changes=add)
    assert not result.is_error, text
    assert "all use the template slide 'Process'" in text
    assert len(second.rendered[-1].slides) == 2 + len(_REPEATED) - 1


@pytest.mark.asyncio
async def test_no_base_both_modes_and_bad_changes_build_nothing(
    monkeypatch, deck_plan_store
):
    workspace = FakeWorkspace()
    turn = _Turn(monkeypatch, workspace)
    text, result = await turn.build(changes=[{"op": "remove", "slide": 1}])
    assert result.is_error and text == (
        "No deck was built: there is no previous deck in this conversation to change. "
        "Send the whole deck in `slides`."
    )
    text, _ = await turn.build(slides=[COVER], changes=[{"op": "remove", "slide": 1}])
    assert text.startswith(
        "No deck was built: the call sent both `slides` and `changes`"
    )
    text, _ = await turn.build(changes=[])
    assert text.startswith("No deck was built: `changes` is empty.")
    await turn.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    stored = dict(deck_plan_store.plans)
    text, result = await turn.build(changes=[{"op": "remove", "slide": 9}])
    assert result.is_error and text == (
        "No deck was built: some changes cannot be applied to the previous deck (nothing "
        "was changed). Fix them and call again:\nchanges[0] (remove): slide 9 does not "
        "exist; the previous deck has 2 slides. Use a number from 1 to 2."
    )
    assert deck_plan_store.plans == stored and len(workspace.writes) == 1


@pytest.mark.asyncio
async def test_an_error_after_evaluation_keeps_the_previous_base(
    monkeypatch, deck_plan_store
):
    turn = _Turn(monkeypatch, FakeWorkspace())
    await turn.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    stored = dict(deck_plan_store.plans)
    monkeypatch.setattr(tools_mod, "acquire_heavy_job_slot", lambda: False)
    text, _ = await turn.build(changes=[{"op": "remove", "slide": 2}])
    assert text.startswith("No deck was built: the server is busy")
    assert deck_plan_store.plans == stored


class _BrokenStore(MemoryDeckPlanStore):
    async def put(self, session_id: str, instance_id: str, plan_json: str) -> None:
        raise ConnectionError("database down")


@pytest.mark.asyncio
async def test_a_failed_base_write_still_delivers(monkeypatch):
    broken = _BrokenStore()
    deck_plans.set_deck_plan_store_provider(lambda: broken)
    turn = _Turn(monkeypatch, FakeWorkspace())
    text, result = await turn.build(slides=[COVER])
    assert not result.is_error and text.startswith(tools_mod._DELIVERED)
    # The turn still knows its own last plan.
    text, result = await turn.build(
        changes=[{"op": "insert", "after": 1, "slides": [{"title": "More"}]}]
    )
    assert not result.is_error, text


@pytest.mark.asyncio
async def test_the_base_is_shared_through_the_sql_table(monkeypatch, analysis_db):
    deck_plans.set_deck_plan_store_provider(lambda: SqlDeckPlanStore(analysis_db))
    store = SqlDeckPlanStore(analysis_db)
    assert await store.get("s", "i") is None
    await store.put("s", "i", "one")
    await store.put("s", "i", "two")
    await store.put("s", "other", "three")
    assert await store.get("s", "i") == "two"

    workspace = FakeWorkspace()
    first = _Turn(monkeypatch, workspace)
    await first.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    second = _Turn(monkeypatch, workspace)
    text, result = await second.build(changes=[{"op": "remove", "slide": 2}])
    assert not result.is_error, text
    assert [s.title for s in second.rendered[-1].slides] == ["Kick-off"]


@pytest.mark.asyncio
async def test_a_deck_without_session_keeps_its_base_for_the_turn_only(monkeypatch):
    monkeypatch.setattr(tools_mod, "convert_office_bytes_to_pdf", _no_pdf)
    config, _ = _configured()
    ctx = CapabilityContext(
        identity=CapabilityIdentity(user_id="u-1", agent_instance_id="inst-7"),
        config=config,
        turn_options=EmptyModel(),
        services=RuntimeServices(workspace_fs=FakeWorkspace()),
    )
    build = _tools(ctx)[BUILD_TOOL_NAME].coroutine
    await build(slides=[COVER, _TOO_LONG])
    text, _ = await build(changes=[{"op": "remove", "slide": 2}])
    assert "there is no previous deck" not in text


@pytest.mark.asyncio
async def test_a_corrected_resend_of_refused_changes_replaces_them(
    monkeypatch, deck_plan_store
):
    """Seen in the eval: a refused insert resent corrected, as `insert` again,
    added one more copy per call when applied to the refused plan."""
    workspace = FakeWorkspace()
    first = _Turn(monkeypatch, workspace)
    await first.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    second = _Turn(monkeypatch, workspace)
    text, result = await second.build(
        changes=[{"op": "insert", "after": 1, "slides": [_TOO_LONG]}]
    )
    assert not result.is_error
    assert _payload(text)["slides"][0]["slide"].startswith("new, sent in changes[0]")
    assert "with the same `changes`, corrected (all of them" in text
    fixed = {"title": "Method", "kind": "process", "items": _STEPS[:3]}
    text, result = await second.build(
        changes=[{"op": "insert", "after": 1, "slides": [fixed]}]
    )
    assert not result.is_error, text
    assert [s.title for s in second.rendered[-1].slides] == [
        "Kick-off",
        "Method",
        "Agenda",
    ]
    # A fix numbered on the refused plan is not applied to it: the refused
    # insert left no trace, slide 4 does not exist.
    await second.build(changes=[{"op": "insert", "after": 3, "slides": [_TOO_LONG]}])
    text, result = await second.build(
        changes=[{"op": "replace", "slide": 4, "slides": [fixed]}]
    )
    assert result.is_error and "slide 4 does not exist" in text


@pytest.mark.asyncio
async def test_a_refused_changes_call_leaves_no_trace(monkeypatch):
    """Seen in the eval: an edit sent in a refused call (items emptied) came
    back in the next call of another shape, which was built with it."""
    turn = _Turn(monkeypatch, FakeWorkspace())
    await turn.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    text, _ = await turn.build(
        changes=[
            {"op": "edit", "slide": 2, "field": "bullets", "value": []},
            {"op": "insert", "after": 2, "slides": [_TOO_LONG]},
        ]
    )
    assert _payload(text)["status"] == "needs_revision"
    text, result = await turn.build(
        changes=[{"op": "edit", "slide": 2, "field": "title", "value": "Plan"}]
    )
    assert not result.is_error and text.startswith(tools_mod._DELIVERED), text
    plan = turn.rendered[-1]
    assert [s.title for s in plan.slides] == ["Kick-off", "Plan"]
    assert plan.slides[1].bullets == ["a"]


@pytest.mark.asyncio
async def test_the_same_changes_after_a_delivery_are_a_new_change(monkeypatch):
    """Two inserts of one shape in one turn, the first delivered: the second
    adds to that deck, it never replaces the first insert."""
    turn = _Turn(monkeypatch, FakeWorkspace())
    await turn.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    for title in ("A", "B"):
        text, result = await turn.build(
            changes=[{"op": "insert", "after": 2, "slides": [{"title": title}]}]
        )
        assert not result.is_error, text
    assert [s.title for s in turn.rendered[-1].slides] == [
        "Kick-off",
        "Agenda",
        "B",
        "A",
    ]


@pytest.mark.asyncio
async def test_null_slides_and_a_single_change_object_are_read(monkeypatch):
    turn = _Turn(monkeypatch, FakeWorkspace())
    await turn.build(slides=[COVER, {"title": "Agenda", "bullets": ["a"]}])
    text, result = await turn.build(
        slides=None,
        changes={"op": "edit", "slide": 2, "field": "title", "value": "Plan"},
    )
    assert not result.is_error, text
    assert turn.rendered[-1].slides[1].title == "Plan"


def test_change_slides_do_not_repeat_the_slide_schema():
    """Every model call pays for the tool schema: `changes` names the slide
    fields once, through `slides`, not a second time."""
    from langchain_core.utils.function_calling import convert_to_openai_tool

    config, _ = _configured()
    tool = _tools(_ctx(config))[BUILD_TOOL_NAME]
    params = convert_to_openai_tool(tool)["function"]["parameters"]
    changes = json.dumps(params["properties"]["changes"])
    assert "Room:" in json.dumps(params["properties"]["slides"])
    assert "Room:" not in changes and len(changes) < 1650


@pytest.mark.asyncio
async def test_a_split_list_is_stored_as_the_delivered_deck(
    monkeypatch, deck_plan_store
):
    """Five steps over two process slides: the stored plan is the deck as
    delivered, so a later change numbers on the file's slides, and the parts
    are rebuilt as they are, never split or suffixed again."""
    workspace = FakeWorkspace()
    first = _Turn(monkeypatch, workspace)
    method = {"title": "Method", "kind": "process", "items": _STEPS[:5]}
    deck = [COVER, method, {"title": "Contact", "bullets": ["a"]}]
    text, result = await first.build(slides=deck)
    assert not result.is_error, text
    assert "Slide 2: its 5 items continue on a second 'process' slide" in text
    (rendered,) = first.rendered
    titles = ["Kick-off", "Method (1/2)", "Method (2/2)", "Contact"]
    assert [s.title for s in rendered.slides] == titles
    stored = DeckPlan.model_validate_json(deck_plan_store.plans[("sess-9", "inst-7")])
    assert _dumps(stored) == _dumps(rendered)
    assert [s.kind for s in stored.slides[1:3]] == ["process", "process"]

    second = _Turn(monkeypatch, workspace)
    text, result = await second.build(
        changes=[{"op": "edit", "slide": 4, "field": "title", "value": "Contacts"}]
    )
    assert not result.is_error, text
    assert "continue on" not in text
    (edited,) = second.rendered
    assert [s.title for s in edited.slides] == titles[:3] + ["Contacts"]
    assert [i.number for i in edited.slides[2].items] == ["3", "4"]
    assert _dumps(edited)[:3] == _dumps(rendered)[:3]


@pytest.mark.asyncio
async def test_a_continuation_is_numbered_at_render_from_its_first_part(
    monkeypatch, deck_plan_store
):
    """Steps sent without numbers: the stored parts carry none, so removing a
    step from part 1 on a later turn renumbers part 2."""
    workspace = FakeWorkspace()

    def numbers(slide: int) -> List[str]:
        deck = Presentation(io.BytesIO(workspace.writes[-1]["content"]))
        return [t for t in _texts(deck.slides[slide - 1]) if t.isdigit()]

    steps = [{"heading": f"Step {i}", "text": "x"} for i in range(5)]
    method = {"title": "Method", "kind": "process", "items": steps}
    text, result = await _Turn(monkeypatch, workspace).build(slides=[COVER, method])
    assert not result.is_error, text
    assert (numbers(2), numbers(3)) == (["1", "2", "3"], ["4", "5"])
    stored = DeckPlan.model_validate_json(deck_plan_store.plans[("sess-9", "inst-7")])
    assert [i.number for i in stored.slides[2].items] == [None, None]
    text, result = await _Turn(monkeypatch, workspace).build(
        changes=[{"op": "edit", "slide": 2, "field": "items[0]", "value": None}]
    )
    assert not result.is_error, text
    assert (numbers(2), numbers(3)) == (["1", "2"], ["3", "4"])


# --- `like`: the design of another slide -------------------------------------------------

PINNED = DeckPlan(
    slides=[
        s.model_copy(update={"built_on": f"slide:{n}"})
        for n, s in enumerate(BASE.slides, start=1)
    ]
)


def _like(*changes: Any) -> DeckPlan:
    return apply_changes(PINNED, [SlideChange.model_validate(c) for c in changes])


def test_repeat_layout_is_an_edit_with_loose_names_and_values():
    plan = _apply(
        {"op": "edit", "slide": 2, "field": "repeat_layout", "value": True},
        {"op": "edit", "slide": 3, "same_layout": "yes"},
        {"op": "edit", "slide": 4, "field": "keep_layout", "value": "true"},
    )
    assert [s.repeat_layout for s in plan.slides] == [False, True, True, True]
    plan = _apply({"op": "set", "slide": 2, "field": "repeat_layout", "value": 0})
    assert not any(s.repeat_layout for s in plan.slides)
    # Written in a new slide, under the field or an alias.
    plan = _apply(
        {
            "op": "insert",
            "after": 3,
            "slides": [{"title": "Again", "same_layout": True}],
        }
    )
    assert plan.slides[3].repeat_layout


def test_like_asks_for_the_referenced_slides_delivered_design():
    plan = _like(
        {"op": "edit", "slide": 4, "field": "like", "value": 2},
        # The referenced slide changed in the same call: its delivered design counts.
        {"op": "replace", "slide": 2, "slides": [{"title": "New", "kind": "quote"}]},
    )
    assert (plan.slides[3].asked_design, plan.slides[3].asked_from) == ("slide:2", 2)
    assert plan.slides[3].kind is None and plan.slides[3].title == "Thanks"
    # The written shorthand, a kind set by the same call, and a replaced target.
    plan = _like(
        {"op": "edit", "slide": 4, "like": "2"},
        {"op": "edit", "slide": 4, "field": "kind", "value": "bullets"},
        {"op": "replace", "slide": 3, "slides": [{"title": "A"}, {"title": "B"}]},
        {"op": "edit", "slide": 3, "field": "like", "value": 1},
    )
    assert [s.asked_design for s in plan.slides] == [
        None,
        None,
        "slide:1",
        None,
        "slide:2",
    ]
    assert plan.slides[4].kind == "bullets"


def test_bad_like_references_say_why_and_what_to_send():
    assert _problems_on(
        PINNED,
        {"op": "edit", "slide": 4, "field": "like", "value": "the agenda"},
        {"op": "edit", "slide": 4, "field": "like", "value": 9},
        {"op": "edit", "slide": 2, "field": "like", "value": 2},
        {"op": "remove", "slide": 3},
        {"op": "edit", "slide": 3, "field": "like", "value": 1},
        {"op": "edit", "slide": 1, "field": "like", "value": 2},
        {"op": "edit", "slide": 1, "field": "like", "value": 4},
    ) == [
        "changes[0] (edit): `like` takes as `value` the number of the slide whose design "
        "to use (e.g. 5), not 'the agenda'.",
        "changes[1] (edit): `like` 9: slide 9 does not exist; the previous deck has 4 "
        "slides. Use a number from 1 to 4.",
        "changes[2] (edit): `like` 2 is slide 2 itself, which already has its design: give "
        "the number of the slide whose design it should take.",
        "changes[4] (edit): slide 3 is removed by changes[3].",
        "changes[6] (edit): slide 1 already takes slide 2's design.",
    ]
    # A plan never delivered has no design to copy.
    (problem,) = _problems_on(BASE, {"op": "edit", "slide": 4, "like": 2})
    assert "slide 2 has no design yet" in problem and "build the deck first" in problem


def _problems_on(base: DeckPlan, *changes: Any) -> List[str]:
    with pytest.raises(ChangeError) as caught:
        apply_changes(base, [SlideChange.model_validate(c) for c in changes])
    return caught.value.problems


def _bullets(n: int) -> dict:
    return {"title": f"S{n}", "bullets": [f"Point {n}", "Detail"]}


_SPREAD = [COVER] + [_bullets(n) for n in range(2, 8)]


async def _delivered_patterns(monkeypatch, workspace, **kwargs) -> tuple:
    used: List[List[str]] = []
    real = tools_mod.render_deck

    def recording(template, inventory, plan, patterns, *rest):
        used.append([p.id for p in patterns])
        return real(template, inventory, plan, patterns, *rest)

    turn = _Turn(monkeypatch, workspace)
    monkeypatch.setattr(tools_mod, "render_deck", recording)
    text, result = await turn.build(**kwargs)
    return text, result, used


@pytest.mark.asyncio
async def test_a_picture_design_asked_for_a_slide_without_picture_is_refused(
    monkeypatch, deck_plan_store
):
    """Seen in the eval: `like` a photo slide on a slide with no picture
    delivered the photo design with its picture half empty."""
    workspace = FakeWorkspace()
    photo = {"title": "Site", "image": "img-1", "bullets": ["Lyon"]}
    docs = FakeDocs({"img-1": (png_bytes(), "image/png")})
    turn = _Turn(monkeypatch, workspace, docs=docs)
    text, result = await turn.build(slides=[COVER, photo, _bullets(3)])
    assert not result.is_error and text.startswith(tools_mod._DELIVERED), text
    text, result = await turn.build(
        changes=[{"op": "edit", "slide": 3, "field": "like", "value": 2}]
    )
    (problem,) = _payload(text)["slides"]
    assert "its picture area would stay empty" in problem["problem"]
    assert "give a picture from list_images_in_folder" in problem["options"]
    assert len(workspace.writes) == 1


@pytest.mark.asyncio
async def test_a_refused_like_never_pins_a_later_edit(monkeypatch, deck_plan_store):
    """Seen in the eval: after a refused `like`, every later edit of the slide
    stayed pinned to that design and was told to drop a like it never sent."""
    workspace = FakeWorkspace()
    turn = _Turn(monkeypatch, workspace)
    await turn.build(slides=_SPREAD)
    text, _ = await turn.build(
        changes=[{"op": "edit", "slide": 3, "field": "like", "value": 1}]
    )
    assert "drop the `like` change" in text
    text, result = await turn.build(
        changes=[{"op": "edit", "slide": 3, "field": "title", "value": "Plan"}]
    )
    assert not result.is_error and text.startswith(tools_mod._DELIVERED), text
    assert turn.rendered[-1].slides[2].title == "Plan"


@pytest.mark.asyncio
async def test_like_gives_a_slide_the_design_of_another(monkeypatch, deck_plan_store):
    """Live: bullets slides spread over several template slides; "slide 7 like
    slide 5" was sent as a no-op kind edit. `like` makes them share."""
    workspace = FakeWorkspace()
    text, result, used = await _delivered_patterns(
        monkeypatch, workspace, slides=_SPREAD
    )
    assert not result.is_error, text
    assert used[-1][4] != used[-1][6]
    text, result, used = await _delivered_patterns(
        monkeypatch,
        workspace,
        changes=[{"op": "edit", "slide": 7, "field": "like", "value": 5}],
    )
    assert not result.is_error and text.startswith(tools_mod._DELIVERED), text
    assert used[-1][6] == used[-1][4]
    stored = DeckPlan.model_validate_json(deck_plan_store.plans[("sess-9", "inst-7")])
    assert stored.slides[6].built_on == stored.slides[4].built_on
    assert stored.slides[6].asked_design is None and stored.slides[6].asked_from is None


@pytest.mark.asyncio
async def test_a_design_the_slide_cannot_take_is_a_revision(
    monkeypatch, deck_plan_store
):
    workspace = FakeWorkspace()
    text, result, _ = await _delivered_patterns(monkeypatch, workspace, slides=_SPREAD)
    assert not result.is_error, text
    items = [{"heading": "One"}, {"heading": "Two"}]
    text, result, _ = await _delivered_patterns(
        monkeypatch,
        workspace,
        changes=[
            {
                "op": "replace",
                "slide": 7,
                "slides": [{"title": "Plan", "items": items}],
            },
            {"op": "edit", "slide": 7, "field": "like", "value": 5},
        ],
    )
    assert not result.is_error and len(workspace.writes) == 1
    (problem,) = _payload(text)["slides"]
    assert problem["slide"] == 7
    assert problem["problem"].startswith("the slide does not fit slide 5's design")
    assert "items: this slide has no place for it" in problem["problem"]
    assert "give the items as bullets" in problem["options"]
    assert problem["options"][-1].startswith("or drop the `like` change")
    # A kind the design does not hold, set in the same call.
    text, _, _ = await _delivered_patterns(
        monkeypatch,
        workspace,
        changes=[
            {"op": "edit", "slide": 7, "field": "kind", "value": "quote"},
            {"op": "edit", "slide": 7, "field": "like", "value": 5},
        ],
    )
    (problem,) = _payload(text)["slides"]
    assert (
        "holds 'bullets' content, but the slide asks for kind 'quote'"
        in (problem["problem"])
    )
    assert problem["options"][0].startswith("set its kind to 'bullets'")


def test_a_design_asked_for_is_never_a_variety_ask():
    from builder_decks import build_showcase_template
    from fred_capability_ppt_filler.builder.analysis import analyze_template
    from fred_capability_ppt_filler.builder.assign import assign

    inventory = analyze_template(build_showcase_template()).inventory
    assert inventory is not None
    slides = [_bullets(n) for n in range(2, 7)]
    plan = DeckPlan.model_validate(
        {
            "slides": [COVER]
            + [{**s, "asked_design": "slide:266", "asked_from": 2} for s in slides]
        }
    )
    assignment = assign(inventory, plan, "template_slide", max_uses=1)
    assert not assignment.problems
    assert [p.id for p in assignment.patterns if p][1:] == ["slide:266"] * 5
    assert assignment.repeats == []


@pytest.mark.asyncio
async def test_design_fields_sent_by_the_model_are_ignored(
    monkeypatch, deck_plan_store
):
    """`built_on` and `asked_*` are the tool's own: written by a model, they
    would pin a slide on a template slide the checks never chose."""
    workspace = FakeWorkspace()
    pin = {"built_on": "layout:5", "asked_design": "layout:5", "asked_from": 1}
    text, result, used = await _delivered_patterns(
        monkeypatch, workspace, slides=[COVER, {**_bullets(2), **pin}]
    )
    assert not result.is_error, text
    assert used[-1][1] != "layout:5"
    replace = {"op": "replace", "slide": 2, "slides": [{**_bullets(3), **pin}]}
    text, result, used = await _delivered_patterns(
        monkeypatch, workspace, changes=[replace]
    )
    assert not result.is_error, text
    assert used[-1][1] != "layout:5"
    stored = DeckPlan.model_validate_json(deck_plan_store.plans[("sess-9", "inst-7")])
    assert stored.slides[1].built_on == used[-1][1]


# --- changes that change nothing ---------------------------------------------------------


@pytest.mark.asyncio
async def test_changes_that_leave_the_deck_as_is_build_nothing(
    monkeypatch, deck_plan_store
):
    """Live: the model sent slide 7's kind and bullets as they were, the deck
    was rebuilt identical and it told the user the design had changed."""
    workspace = FakeWorkspace()
    text, result, _ = await _delivered_patterns(monkeypatch, workspace, slides=_SPREAD)
    assert not result.is_error, text
    stored = dict(deck_plan_store.plans)
    same = _bullets(7)["bullets"]
    for changes in (
        [
            {"op": "edit", "slide": 7, "field": "kind", "value": "bullets"},
            {"op": "edit", "slide": 7, "field": "bullets", "value": same},
        ],
        [{"op": "replace", "slide": 3, "slides": [_bullets(3)]}],
    ):
        text, result, used = await _delivered_patterns(
            monkeypatch, workspace, changes=changes
        )
        assert not result.is_error and text.startswith(
            "No deck was built: nothing changed"
        )
        assert not text.startswith(tools_mod._DELIVERED)
        assert (
            '"field": "like"' in text and "Do not tell the user it was changed" in text
        )
        assert used == [] and len(workspace.writes) == 1
        assert deck_plan_store.plans == stored
    # A `like` naming the design the slide already has changes nothing either.
    text, _, _ = await _delivered_patterns(
        monkeypatch,
        workspace,
        changes=[{"op": "edit", "slide": 7, "field": "like", "value": 5}],
    )
    assert text.startswith(tools_mod._DELIVERED), text
    text, result, _ = await _delivered_patterns(
        monkeypatch,
        workspace,
        changes=[{"op": "edit", "slide": 7, "field": "like", "value": 5}],
    )
    assert not result.is_error and text.startswith("No deck was built: nothing changed")
    # A new file name is a change.
    text, result, _ = await _delivered_patterns(
        monkeypatch,
        workspace,
        changes=[{"op": "edit", "slide": 2, "field": "title", "value": "S2"}],
        file_name="Renamed",
    )
    assert not result.is_error, text


# --- a kind edit, dropping a field, an identical resend -------------------------------------

_COUNTRIES = [
    {"heading": "Espagne", "bullets": ["Partenariats locaux", "R&D à Madrid"]},
    {"heading": "Italie", "bullets": ["Spatial", "1,5 Md€ CA", "Cyber"]},
]
_PERSPECTIVES = {
    "title": "Perspectives 2025",
    "bullets": ["Espagne : partenariats", "Italie : spatial"],
    "kind": "bullets",
}
# The live session's call: the kind and the columns sent, the bullets left as they were.
_TO_TWO_COLUMNS = [
    {"op": "edit", "field": "kind", "slide": 2, "value": "two_columns"},
    {"op": "edit", "field": "title", "slide": 2, "value": "Perspectives 2025 par pays"},
    {"op": "edit", "field": "columns", "slide": 2, "value": _COUNTRIES},
]


@pytest.mark.asyncio
async def test_a_kind_edit_clears_the_fields_the_new_kind_cannot_hold(
    monkeypatch, deck_plan_store
):
    """Live: the old bullets stayed on a slide edited to two_columns, and the
    model resent the same call until the turn's tool-call cap."""
    workspace = FakeWorkspace()
    deck = [COVER, _PERSPECTIVES, {"title": "Merci"}]
    text, result = await _Turn(monkeypatch, workspace).build(slides=deck)
    assert not result.is_error, text
    turn = _Turn(monkeypatch, workspace)
    text, result = await turn.build(changes=_TO_TWO_COLUMNS)
    assert not result.is_error, text
    slide = turn.rendered[-1].slides[1]
    assert (slide.kind, slide.bullets, slide.title) == (
        "two_columns",
        [],
        "Perspectives 2025 par pays",
    )
    assert "Slide 2: bullets cleared, as kind 'two_columns' has no place for it" in text


def test_a_kind_edit_keeps_what_the_same_call_sets_and_what_the_kind_holds():
    def unplaced(slide: Any) -> List[str]:
        return ["bullets"]

    base = DeckPlan.model_validate({"slides": [COVER, _PERSPECTIVES]})
    sent = [SlideChange.model_validate(c) for c in _TO_TWO_COLUMNS[::2]]
    cleared: dict = {}
    plan = apply_changes(base, sent, unplaced, cleared)
    assert plan.slides[1].bullets == [] and plan.slides[1].title == "Perspectives 2025"
    assert cleared == {2: ["bullets"]}
    # Bullets sent by the same call stay; the same kind resent clears nothing.
    bullets = {"op": "edit", "field": "bullets", "slide": 2, "value": ["a"]}
    plan = apply_changes(base, sent + [SlideChange.model_validate(bullets)], unplaced)
    assert plan.slides[1].bullets == ["a"]
    same = [SlideChange.model_validate({**_TO_TWO_COLUMNS[0], "value": "bullets"})]
    plan = apply_changes(base, same, unplaced)
    assert plan.slides[1].bullets == _PERSPECTIVES["bullets"]


@pytest.mark.asyncio
async def test_a_drop_option_gives_its_exact_change(monkeypatch, deck_plan_store):
    workspace = FakeWorkspace()
    two = {**_PERSPECTIVES, "kind": "two_columns", "columns": _COUNTRIES}
    turn = _Turn(monkeypatch, workspace)
    text, result = await turn.build(slides=[COVER, two])
    (problem,) = _payload(text)["slides"]
    # A whole deck is fixed in `slides`: the option stays in words.
    assert "drop bullets" in problem["options"]
    text, result = await turn.build(slides=[COVER, _PERSPECTIVES])
    assert not result.is_error, text
    columns = {"op": "edit", "field": "columns", "slide": 2, "value": _COUNTRIES}
    text, result = await turn.build(changes=[columns])
    (problem,) = _payload(text)["slides"]
    drop = {"op": "edit", "slide": 2, "field": "bullets", "value": []}
    assert f"drop bullets: {json.dumps(drop)}" in problem["options"]
    # The op as given clears the field, and the deck builds.
    text, result = await turn.build(changes=[columns, drop])
    assert not result.is_error, text
    assert turn.rendered[-1].slides[1].bullets == []


@pytest.mark.asyncio
async def test_a_call_identical_to_a_refused_one_is_not_evaluated_again(monkeypatch):
    turn = _Turn(monkeypatch, FakeWorkspace())
    evaluated: List[int] = []
    real = tools_mod.assign

    def counting(*args: Any) -> Any:
        evaluated.append(1)
        return real(*args)

    monkeypatch.setattr(tools_mod, "assign", counting)
    await turn.build(slides=_DECK)
    text, result = await turn.build(slides=_DECK)
    # The loop breaker: answered without a second evaluation, still not a failure.
    assert not result.is_error and len(evaluated) == 1
    assert text.startswith("No deck was built: this call is identical to one already")
    assert "slide 3: " in text and '"op": "replace", "slide": N' in text
    assert "leave that slide out" in text
    insert = [
        {"op": "insert", "after": 1, "slides": [{"title": "Why", "bullets": ["a"]}]}
    ]
    await turn.build(changes=insert)
    text, _ = await turn.build(changes=insert)
    assert len(evaluated) == 2 and "drop the change that causes it" in text
    # The slide sent whole breaks the loop, numbered on the refused `slides`
    # plan the refused insert left as it was.
    fixed = {"title": "Method", "kind": "process", "items": _STEPS[:3]}
    text, result = await turn.build(
        changes=[{"op": "replace", "slide": 3, "slides": [fixed]}]
    )
    assert not result.is_error, text


def test_an_edit_of_a_whole_list_with_one_object_is_refused():
    """The eval call: two `columns` edits, each one object; the second replaced
    the first and two points were lost without a word."""
    column = {"heading": "Objectifs", "bullets": ["Gagner"]}
    (problem,) = _problems(
        {"op": "edit", "slide": 3, "field": "items", "value": column}
    )
    assert "'items' is a list: send the whole list" in problem
    assert "'items[i]'" in problem
    # One entry by its index, or the whole list, stays an edit.
    edited = _apply({"op": "edit", "slide": 3, "field": "items[0].label", "value": "x"})
    assert edited.slides[2].items[0].label == "x"


def test_two_edits_of_one_field_on_one_slide_are_refused():
    problems = _problems(
        {"op": "edit", "slide": 2, "field": "bullets", "value": ["A"]},
        {"op": "edit", "slide": 2, "field": "bullets", "value": ["B"]},
    )
    (problem,) = problems
    assert problem.startswith("changes[1] (edit): slide 2: changes[0] already sets")
    # Different fields of one slide still go together.
    edited = _apply(
        {"op": "edit", "slide": 2, "field": "title", "value": "Plan"},
        {"op": "edit", "slide": 2, "field": "bullets[0]", "value": "Un"},
    )
    assert (edited.slides[1].title, edited.slides[1].bullets[0]) == ("Plan", "Un")
