# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Opt-in deck-quality evaluation of ``ppt_builder``.

Prompts x templates through the real analysis, labelling (vision or rules),
the real ``build_presentation`` tool and a real LLM acting as the agent, with
the capability's own prompt fragment and tool. Writes one JSON per run, a
markdown summary and, on request, PNG contact sheets. Run with ``make eval``;
never part of ``make test``. Usage: ``python eval/run_eval.py --help``.
"""

from __future__ import annotations

import argparse
import asyncio
import contextvars
import io
import json
import os
import re
import statistics
import sys
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple, cast, get_args

import yaml
from dotenv import dotenv_values
from fred_capability_ppt_filler.builder import analysis_job, checks, deck_plans
from fred_capability_ppt_filler.builder import tools as tools_mod
from fred_capability_ppt_filler.builder.analysis import analyze_template
from fred_capability_ppt_filler.builder.assign import (
    _STRUCTURAL,
    VarietyCheck,
    assign,
    assignable,
    misfit,
)
from fred_capability_ppt_filler.builder.builtin import builtin_template
from fred_capability_ppt_filler.builder.capability import (
    _BUILD_INSTRUCTIONS,
    PptBuilderConfig,
)
from fred_capability_ppt_filler.builder.models import Pattern, TemplateInventory
from fred_capability_ppt_filler.builder.plan import DeckPlan, SlidePlan, has_content
from fred_capability_ppt_filler.builder.render import RenderResult
from fred_capability_ppt_filler.fill import PptPreviewPart
from fred_core.common import ModelConfiguration
from fred_core.model.factory import get_model
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
)
from fred_sdk.contracts.context import PublishedArtifact
from fred_sdk.contracts.runtime import (
    CapabilityModelPort,
    CapabilityModelUnavailableError,
    RuntimeServices,
)
from fred_sdk.contracts.ui_part_union import rebuild_ui_part_union
from langchain_core.language_models import BaseChatModel
from langchain_core.messages import (
    AIMessage,
    BaseMessage,
    HumanMessage,
    SystemMessage,
    ToolMessage,
)
from langchain_core.tools import BaseTool
from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from prompts import EDITS, PROMPTS  # noqa: E402

REPO = Path(__file__).resolve().parents[4]
DEFAULT_ENV = REPO / "apps/fred-agents/config/.env"
DEFAULT_CATALOG = REPO / "apps/fred-agents/config/models_catalog.yaml"
MAX_REPAIRS = 3
# Stays under the package's four heavy-job slots (renders and labelling).
RUN_CONCURRENCY = 3
# Stand-in for the agent's own prompt; the capability fragment is appended to it.
_BASE_PROMPT = "You are a helpful assistant for a company team. Answer in the language of the user."


# --- model access ----------------------------------------------------------------------


def _load_secrets(env_file: Path) -> None:
    """Only API keys reach the process env, never printed; a corporate proxy
    from the same file would break calls made off that network."""
    for key, value in dotenv_values(env_file).items():
        if key.endswith("_API_KEY") and value and not os.environ.get(key):
            os.environ[key] = value


def _merge(base: Dict[str, Any], top: Dict[str, Any]) -> Dict[str, Any]:
    merged = dict(base)
    for key, value in top.items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = _merge(merged[key], value)
        else:
            merged[key] = value
    return merged


@dataclass
class Profile:
    profile_id: str
    name: str
    supports_vision: bool
    model: BaseChatModel


def load_profile(catalog: Path, profile_id: Optional[str]) -> Profile:
    """A catalogue profile built as fred-runtime builds it (common settings
    merged under the profile's; reasoning stays off, the admin default)."""
    data = yaml.safe_load(catalog.read_text(encoding="utf-8"))
    wanted = profile_id or data["default_profile_by_capability"]["chat"]
    profile = next(p for p in data["profiles"] if p["profile_id"] == wanted)
    settings = _merge(
        data.get("common_model_settings") or {}, profile["model"].get("settings") or {}
    )
    settings.pop("reasoning_effort", None)
    config = ModelConfiguration(
        provider=profile["model"]["provider"],
        name=profile["model"]["name"],
        settings=settings,
    )
    return Profile(
        wanted,
        profile["model"]["name"],
        bool(profile.get("supports_vision")),
        get_model(config),
    )


class EvalModelPort(CapabilityModelPort):
    """The turn-free port over one catalogue profile, failing closed like the
    runtime adapter when vision is required and the profile lacks it."""

    def __init__(self, profile: Profile) -> None:
        self.profile = profile

    async def ainvoke(
        self,
        messages: Sequence[BaseMessage],
        *,
        profile_id: Optional[str] = None,
        require_vision: bool = False,
    ) -> AIMessage:
        if require_vision and not self.profile.supports_vision:
            raise CapabilityModelUnavailableError(
                f"{self.profile.profile_id} has no vision"
            )
        answer = await self.profile.model.ainvoke(list(messages))
        assert isinstance(answer, AIMessage)
        return answer


# --- templates and analysis ------------------------------------------------------------


@dataclass
class Variant:
    template: str
    mode: str  # "vision" | "rules"
    template_bytes: bytes
    inventory: TemplateInventory
    analysis_seconds: float
    labelled: Dict[str, int]


def _labelled(inventory: TemplateInventory) -> Dict[str, int]:
    counts: Dict[str, int] = {}
    for pattern in inventory.patterns:
        if pattern.source == "slide" and pattern.included:
            counts[pattern.labelled_by] = counts.get(pattern.labelled_by, 0) + 1
    return counts


async def analyse(
    name: str, content: bytes, modes: Sequence[str], port: Optional[EvalModelPort]
) -> List[Variant]:
    started = time.monotonic()
    result = await asyncio.to_thread(analyze_template, content)
    if result.inventory is None:
        raise SystemExit(f"{name}: analysis failed: {result.errors}")
    rules = result.inventory
    rules_seconds = time.monotonic() - started
    variants = []
    if "rules" in modes:
        variants.append(
            Variant(name, "rules", content, rules, rules_seconds, _labelled(rules))
        )
    if "vision" in modes and port is not None:
        # The background job's own pipeline: analysis, patterns PDF, labelling.
        started = time.monotonic()
        labelled, pdf = await analysis_job._analyse(
            content, port, analysis_job.labelling_key(True, None)
        )
        if pdf is None:
            print(f"{name}: no patterns PDF (LibreOffice/TMPDIR?): rules labels only")
        seconds = time.monotonic() - started
        variants.append(
            Variant(name, "vision", content, labelled, seconds, _labelled(labelled))
        )
    return variants


# --- the agent loop --------------------------------------------------------------------


class MemoryWorkspace:
    """The write half of the workspace port, kept in memory."""

    def __init__(self) -> None:
        self.files: Dict[str, bytes] = {}

    async def write(self, path: str, content: bytes, **_: Any) -> Any:
        self.files[path] = bytes(content)
        name = path.rsplit("/", 1)[-1]
        return PublishedArtifact(
            key=path,
            file_name=name,
            size=len(content),
            href=f"memory://{path}",
            document_uid=path,
            mime=None,
        )


_renders: contextvars.ContextVar[List[RenderResult]] = contextvars.ContextVar("renders")
# The whole plan each render was given: a `changes` call's arguments are not one.
_plans: contextvars.ContextVar[List[DeckPlan]] = contextvars.ContextVar("plans")
_real_render = tools_mod.render_deck


def _recording_render(*args: Any, **kwargs: Any) -> RenderResult:
    # Observes only: the worker thread inherits the run's context.
    result = _real_render(*args, **kwargs)
    _renders.get().append(result)
    _plans.get().append(args[2])
    return result


tools_mod.render_deck = _recording_render


@dataclass
class Attempt:
    seconds: float
    tool_called: bool
    status: str  # delivered | needs_revision | error | no_tool_call
    message: str
    slides: int = 0
    problems: List[Dict[str, Any]] = field(default_factory=list)
    # The call sent `changes` rather than the whole deck.
    changes: bool = False


@dataclass
class EditResult:
    """The second user turn: a change to the delivered deck."""

    request: str
    success: bool = False
    failure: Optional[str] = None
    calls: int = 0
    calls_with_changes: int = 0
    first_call: bool = False
    input_tokens: int = 0
    output_tokens: int = 0
    slides_before: int = 0
    slides_after: int = 0
    # Slides of the first deck found unchanged in the edited one.
    kept: int = 0
    attempts: List[Attempt] = field(default_factory=list)


@dataclass
class RunResult:
    template: str
    mode: str
    prompt: str
    success: bool = False
    failure: Optional[str] = None
    repairs: int = 0
    wall_seconds: float = 0.0
    slides: int = 0
    patterns: List[str] = field(default_factory=list)
    kinds: List[str] = field(default_factory=list)
    distinct_ratio: float = 0.0
    cover_once_first: Optional[bool] = None
    closing_ok: Optional[bool] = None
    triples: int = 0
    # Deck-wide variety: the most uses of one pattern, A B A B returns, and
    # reuses while an unused pattern of the same kind would have held the slide.
    max_reuse: int = 0
    # Most slides of one content kind (cover, section, agenda, closing aside).
    max_kind_repeat: int = 0
    alternations: int = 0
    avoidable_reuses: int = 0
    # Tables: given in the plan, put on a table slide, rendered
    # with the plan's rows and columns, and kept inside the slide.
    tables_given: int = 0
    tables_used: int = 0
    tables_shape_kept: int = 0
    tables_overflow: int = 0
    leftovers: List[str] = field(default_factory=list)
    removed_by_check: int = 0
    truncated: int = 0
    # Text strings of the raw tool arguments absent from the delivered deck:
    # catches what the coercion drops before the renderer's report sees it.
    lost: int = 0
    lost_texts: List[str] = field(default_factory=list)
    variety_asks: int = 0
    # Model tokens over the run, for cost comparisons between models.
    input_tokens: int = 0
    output_tokens: int = 0
    # The answer the user reads after delivery; its tokens stay out of the build's.
    reply: Optional[str] = None
    reply_output_tokens: int = 0
    # The model called the tool again instead of answering (a wasted rebuild).
    reply_tool_call: bool = False
    attempts: List[Attempt] = field(default_factory=list)
    plan: Optional[Dict[str, Any]] = None
    edit: Optional[EditResult] = None


def _build_tool(
    variant: Variant,
    workspace: MemoryWorkspace,
    variety: VarietyCheck,
    max_uses: int,
    session: str = "eval",
    history: Any = None,
) -> BaseTool:
    """The tool the capability's middleware registers once the analysis is in
    the pod's cache, with the template read from memory; one per turn."""
    ctx = CapabilityContext(
        identity=CapabilityIdentity(
            user_id="eval", session_id=session, agent_instance_id="eval-instance"
        ),
        config=PptBuilderConfig(
            template_sha256="eval", variety_check=variety, variety_max_uses=max_uses
        ),
        turn_options=EmptyModel(),
        services=RuntimeServices(workspace_fs=cast(Any, workspace)),
    )
    content = variant.template_bytes

    async def load() -> bytes:
        return content

    if history is None:
        return tools_mod.build_builder_tools(ctx, load, variant.inventory)[0]
    return tools_mod.build_builder_tools(ctx, load, variant.inventory, history)[0]


def _whole(plan: DeckPlan) -> Dict[str, Any]:
    """A rendered plan as the arguments of a full call, for the metrics."""
    return {
        "slides": [
            s.model_dump(exclude={"ignored_fields", "invalid_entry"})
            for s in plan.slides
        ],
        "file_name": plan.file_name,
    }


async def run_one(
    variant: Variant,
    prompt_id: str,
    profile: Profile,
    variety: VarietyCheck,
    max_uses: int,
    agent_prompt: str = _BASE_PROMPT,
    edit_turn: bool = False,
) -> RunResult:
    _renders.set([])
    _plans.set([])
    workspace = MemoryWorkspace()
    session = f"eval-{variant.template}-{variant.mode}-{prompt_id}"
    messages: List[BaseMessage] = [
        SystemMessage(content=f"{agent_prompt}\n\n{_BUILD_INSTRUCTIONS}"),
        HumanMessage(content=PROMPTS[prompt_id]),
    ]
    tool = _build_tool(variant, workspace, variety, max_uses, session, lambda: messages)
    model = profile.model.bind_tools([tool])
    run = RunResult(variant.template, variant.mode, prompt_id)
    started = time.monotonic()
    last_args: Optional[Dict[str, Any]] = None
    tool_rounds = 0
    try:
        for _round in range(MAX_REPAIRS + 1):
            t0 = time.monotonic()
            answer = await model.ainvoke(messages)
            assert isinstance(answer, AIMessage)
            messages.append(answer)
            usage = answer.usage_metadata or {}
            run.input_tokens += usage.get("input_tokens", 0)
            run.output_tokens += usage.get("output_tokens", 0)
            if not any(c["name"] == tool.name for c in answer.tool_calls):
                run.attempts.append(
                    Attempt(time.monotonic() - t0, False, "no_tool_call", _text(answer))
                )
                break
            tool_rounds += 1
            # Every call gets its answer, as in a turn; the last one counts.
            status = ""
            for call in answer.tool_calls:
                if call["name"] != tool.name:
                    messages.append(
                        ToolMessage(content="Unknown tool.", tool_call_id=call["id"])
                    )
                    continue
                last_args = call["args"]
                result = await tool.ainvoke({**call, "type": "tool_call"})
                messages.append(result)
                text = str(result.content)
                status = _status(text)
                run.variety_asks += '"variety": [' in text
                run.attempts.append(
                    Attempt(
                        time.monotonic() - t0,
                        True,
                        status,
                        text[:2000],
                        problems=_problems(text),
                        changes=bool(call["args"].get("changes")),
                    )
                )
            if status == "delivered":
                run.success = True
                # A `changes` call's arguments are not the deck: measure what was rendered.
                if last_args is not None and last_args.get("changes") and _plans.get():
                    last_args = _whole(_plans.get()[-1])
                await _reply(run, model, messages)
                break
        run.repairs = max(0, tool_rounds - 1)
        if not run.success:
            last = run.attempts[-1].status if run.attempts else "no_attempt"
            run.failure = "repairs_exhausted" if tool_rounds > MAX_REPAIRS else last
    except Exception as exc:  # noqa: BLE001 - recorded as a failed run
        run.failure = f"exception: {type(exc).__name__}: {exc}"[:300]
    run.wall_seconds = round(time.monotonic() - started, 1)
    if last_args is not None:
        run.plan = last_args
    if run.success and last_args is not None:
        _measure(run, variant, last_args, workspace, _renders.get(), variety, max_uses)
        if edit_turn and _plans.get():
            run.edit = await _edit_turn(
                variant,
                prompt_id,
                model,
                profile,
                messages,
                run,
                variety,
                max_uses,
                session,
            )
    return run


async def _edit_turn(
    variant: Variant,
    prompt_id: str,
    bound: Any,
    profile: Profile,
    messages: List[BaseMessage],
    run: RunResult,
    variety: VarietyCheck,
    max_uses: int,
    session: str,
) -> EditResult:
    """The user's next turn asks to change the deck: a new tool instance on the
    same conversation and base-plan store, as on another pod."""
    before = _plans.get()[-1]
    edit = EditResult(EDITS[prompt_id], slides_before=len(before.slides))
    messages.append(AIMessage(content=run.reply or "Done."))
    messages.append(HumanMessage(content=EDITS[prompt_id]))
    workspace = MemoryWorkspace()
    tool = _build_tool(variant, workspace, variety, max_uses, session, lambda: messages)
    model = profile.model.bind_tools([tool])
    rendered = len(_plans.get())
    try:
        for _round in range(MAX_REPAIRS + 1):
            t0 = time.monotonic()
            answer = await model.ainvoke(messages)
            assert isinstance(answer, AIMessage)
            messages.append(answer)
            usage = answer.usage_metadata or {}
            edit.input_tokens += usage.get("input_tokens", 0)
            edit.output_tokens += usage.get("output_tokens", 0)
            calls = [c for c in answer.tool_calls if c["name"] == tool.name]
            if not calls:
                edit.attempts.append(
                    Attempt(time.monotonic() - t0, False, "no_tool_call", _text(answer))
                )
                break
            status = ""
            for call in answer.tool_calls:
                if call["name"] != tool.name:
                    messages.append(
                        ToolMessage(content="Unknown tool.", tool_call_id=call["id"])
                    )
                    continue
                edit.calls += 1
                edit.calls_with_changes += bool(call["args"].get("changes"))
                result = await tool.ainvoke({**call, "type": "tool_call"})
                messages.append(result)
                text = str(result.content)
                status = _status(text)
                edit.attempts.append(
                    Attempt(
                        time.monotonic() - t0,
                        True,
                        status,
                        text[:2000],
                        problems=_problems(text),
                        changes=bool(call["args"].get("changes")),
                    )
                )
            if status == "delivered":
                edit.success = True
                break
        if not edit.success:
            edit.failure = edit.attempts[-1].status if edit.attempts else "no_attempt"
    except Exception as exc:  # noqa: BLE001 - recorded as a failed edit
        edit.failure = f"exception: {type(exc).__name__}: {exc}"[:300]
    edit.first_call = edit.success and edit.calls == 1
    if edit.success and len(_plans.get()) > rendered:
        after = _plans.get()[-1]
        edit.slides_after = len(after.slides)
        left = [s.model_dump_json() for s in after.slides]
        for slide in before.slides:
            dumped = slide.model_dump_json()
            if dumped in left:
                left.remove(dumped)
                edit.kept += 1
    return edit


async def _reply(run: RunResult, model: Any, messages: List[BaseMessage]) -> None:
    try:
        answer = await model.ainvoke(messages)
    except Exception as exc:  # noqa: BLE001 - the deck is delivered either way
        run.reply = f"exception: {type(exc).__name__}"
        return
    assert isinstance(answer, AIMessage)
    run.reply = _text(answer)
    run.reply_tool_call = bool(answer.tool_calls)
    run.reply_output_tokens = (answer.usage_metadata or {}).get("output_tokens", 0)


def _text(message: AIMessage) -> str:
    content = message.content
    return content if isinstance(content, str) else json.dumps(content)[:2000]


def _status(text: str) -> str:
    if '"status": "needs_revision"' in text:
        return "needs_revision"
    if text.startswith(("Delivered:", "The presentation '")):
        return "delivered"
    return "error"


def _problems(text: str) -> List[Dict[str, Any]]:
    start = text.find("{")
    if '"needs_revision"' not in text or start < 0:
        return []
    try:
        return json.loads(text[start:]).get("slides", [])
    except ValueError:
        return []


# --- metrics ---------------------------------------------------------------------------


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def _plan_texts(slide: SlidePlan) -> set[str]:
    texts = [slide.title, slide.subtitle, *slide.bullets]
    for column in slide.columns:
        texts += [column.heading, *column.bullets]
    for item in slide.items:
        texts += [item.number, item.heading, item.text, item.label]
    return {_norm(t.lstrip("- ")) for t in texts if t}


def _measure(
    run: RunResult,
    variant: Variant,
    args: Dict[str, Any],
    workspace: MemoryWorkspace,
    renders: List[RenderResult],
    variety: VarietyCheck,
    max_uses: int,
) -> None:
    inventory = variant.inventory
    schema = tools_mod._args_schema(inventory)
    plan: DeckPlan = tools_mod._to_plan(schema(**args))
    # Deterministic: the same assignment, with the run's settings, the tool rendered.
    assigned = assign(inventory, plan, variety, max_uses).patterns
    chosen = [p for p in assigned if p is not None]
    slides = [s for s in plan.slides if s.invalid_entry is None]
    run.slides = len(chosen)
    run.patterns = [p.id for p in chosen]
    run.kinds = [p.kind for p in chosen]
    run.distinct_ratio = round(len(set(run.patterns)) / max(1, len(chosen)), 2)
    offered = {p.kind for p in inventory.patterns if p.included and p.reusable}
    if "cover" in offered:
        run.cover_once_first = run.kinds.count("cover") == 1 and run.kinds[0] == "cover"
    closings = [i for i, k in enumerate(run.kinds) if k == "closing"]
    run.closing_ok = len(closings) <= 1 and (
        not closings or closings[0] == len(run.kinds) - 1
    )
    run.triples = sum(
        1
        for i in range(2, len(chosen))
        if chosen[i].id == chosen[i - 1].id == chosen[i - 2].id
    )
    _variety(run, inventory, slides, chosen)
    if renders:
        report = renders[-1].report
        for slide_report in report.slides.values():
            run.removed_by_check += len(slide_report.removed)
            run.truncated += len(slide_report.truncated)
    deck = next((b for k, b in workspace.files.items() if k.endswith(".pptx")), None)
    if deck is not None:
        run.leftovers = _leftovers(variant, slides, chosen, deck)
        _tables(run, slides, chosen, deck)
        run.lost_texts = _lost_texts(args, assigned, deck)
        run.lost = len(run.lost_texts)


_NOT_TEXT = {"kind", "type", "pattern", "image", "images", "notes", "file_name"}


def _raw_texts(value: Any, key: Optional[str] = None) -> List[str]:
    """Every text line of the model's raw arguments, under whatever key."""
    if key in _NOT_TEXT or isinstance(value, bool) or value is None:
        return []
    if isinstance(value, dict):
        return [t for k, v in value.items() for t in _raw_texts(v, k)]
    if isinstance(value, list):
        return [t for v in value for t in _raw_texts(v, key)]
    return str(value).splitlines()


def _bare(text: str) -> str:
    """Compared without the inline markers the renderer turns into formatting."""
    return _norm(re.sub(r"[*_`]", "", text)).removeprefix("- ")


def _lost_texts(
    args: Dict[str, Any], patterns: Sequence[Optional[Pattern]], deck: bytes
) -> List[str]:
    """Raw-argument text that is not on its slide of the deck; fixed slides
    (shown as is) and step numbers (the item order shows them) are not content."""
    shown = iter(
        "\n".join(
            _bare("".join(t.text or "" for t in p.iter(qn("a:t"))))
            for p in slide._element.iter(qn("a:p"))
        )
        for slide in Presentation(io.BytesIO(deck)).slides
    )
    raw = args.get("slides") or []
    lost = []
    for entry, pattern in zip(raw if isinstance(raw, list) else [], patterns):
        if pattern is None:  # not a slide: nothing rendered for it
            continue
        on_slide = next(shown, "")
        if not isinstance(entry, dict) or not pattern.fillable:
            continue
        for text in _raw_texts(entry):
            key = _bare(text)
            if key and not (key.isdigit() and len(key) <= 2) and key not in on_slide:
                lost.append(text[:60])
    return lost


def _tables(
    run: RunResult, slides: List[SlidePlan], chosen: List[Pattern], deck: bytes
) -> None:
    presentation = Presentation(io.BytesIO(deck))
    for plan, pattern, slide in zip(slides, chosen, presentation.slides):
        if plan.table is None or not plan.table.has_content():
            continue
        run.tables_given += 1
        if pattern.kind != "table":
            continue
        run.tables_used += 1
        frame = next((s for s in slide.shapes if getattr(s, "has_table", False)), None)
        if frame is None:
            continue
        table = cast(Any, frame).table
        header = 1 if any(plan.table.header) else 0
        if (len(table.rows), len(table.columns)) == (
            len(plan.table.rows) + header,
            plan.table.width,
        ):
            run.tables_shape_kept += 1
        if frame.top + frame.height > (presentation.slide_height or 0):
            run.tables_overflow += 1


def _variety(
    run: RunResult,
    inventory: TemplateInventory,
    slides: List[SlidePlan],
    chosen: List[Pattern],
) -> None:
    ids = [p.id for p in chosen]
    run.max_reuse = max((ids.count(i) for i in ids), default=0)
    content = [p.kind for p in chosen if p.kind not in _STRUCTURAL]
    run.max_kind_repeat = max((content.count(k) for k in content), default=0)
    run.alternations = sum(
        1 for i in range(2, len(ids)) if ids[i] == ids[i - 2] != ids[i - 1]
    )
    size = (inventory.slide_width, inventory.slide_height)
    pool = assignable(inventory)
    used = set(ids)
    seen: set[str] = set()
    for slide, pattern in zip(slides, chosen):
        if pattern.id in seen and any(
            p.kind == pattern.kind
            and p.id not in used
            and not misfit(p, slide, size)[0]
            for p in pool
        ):
            run.avoidable_reuses += 1
        seen.add(pattern.id)


def _leftovers(
    variant: Variant, slides: List[SlidePlan], chosen: List[Pattern], deck: bytes
) -> List[str]:
    """Original non-shared text of a filled example slide still in the output."""
    source = Presentation(io.BytesIO(variant.template_bytes))
    by_id = {s.slide_id: s for s in source.slides}
    shared = {_norm(t) for t in variant.inventory.shared_texts}
    found: List[str] = []
    output = Presentation(io.BytesIO(deck)).slides
    for position, (plan, pattern, slide) in enumerate(
        zip(slides, chosen, output), start=1
    ):
        filled = pattern.reusable and pattern.fillable and has_content(plan)
        if pattern.source != "slide" or not filled:
            continue
        original = by_id[int(pattern.id.split(":", 1)[1])]
        before = {_norm(t) for t in checks.original_texts(original).values()}
        wanted = _plan_texts(plan)
        for text in checks.original_texts(slide).values():
            key = _norm(text)
            # Step numbers the renderer writes ("1", "2"...) equal the samples.
            if key.rstrip(".").isdigit():
                continue
            if key in before and key not in shared and key not in wanted:
                found.append(f"{position}: {text[:60]}")
    return found


# --- output ----------------------------------------------------------------------------


def _summary(results: List[RunResult], variants: List[Variant]) -> str:
    lines = [
        "| template | mode | runs | success | first call | avg repairs | cover once@1 "
        "| 3-in-a-row | distinct/slides | max reuse | max kind repeat (max / mean) | A-B-A | avoidable reuses "
        "| patterns used | tables used/given | table shape kept | table overflow "
        "| leftovers | removed by check | truncated | lost | variety asks | avg wall s | tokens in / out (k) | reply chars (mean) | tool call instead of reply |",
        "|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|",
    ]
    for variant in variants:
        runs = [
            r
            for r in results
            if r.template == variant.template and r.mode == variant.mode
        ]
        ok = [r for r in runs if r.success]
        covers = [r.cover_once_first for r in ok if r.cover_once_first is not None]
        across = {p for r in ok for p in r.patterns}

        def mean(values: List[float]) -> str:
            return f"{statistics.mean(values):.2f}" if values else "-"

        lines.append(
            f"| {variant.template} | {variant.mode} | {len(runs)} | {len(ok)} "
            f"| {sum(1 for r in ok if r.repairs == 0)} "
            f"| {mean([float(r.repairs) for r in runs])} "
            f"| {f'{100 * sum(covers) / len(covers):.0f}%' if covers else 'n/a'} "
            f"| {sum(r.triples for r in ok)} "
            f"| {mean([r.distinct_ratio for r in ok])} "
            f"| {max((r.max_reuse for r in ok), default=0)} "
            f"| {max((r.max_kind_repeat for r in ok), default=0)} / "
            f"{mean([float(r.max_kind_repeat) for r in ok])} "
            f"| {sum(r.alternations for r in ok)} | {sum(r.avoidable_reuses for r in ok)} "
            f"| {len(across)}/{len(assignable(variant.inventory))} "
            f"| {sum(r.tables_used for r in ok)}/{sum(r.tables_given for r in ok)} "
            f"| {sum(r.tables_shape_kept for r in ok)} "
            f"| {sum(r.tables_overflow for r in ok)} "
            f"| {sum(len(r.leftovers) for r in ok)} "
            f"| {sum(r.removed_by_check for r in ok)} "
            f"| {sum(r.truncated for r in ok)} | {sum(r.lost for r in ok)} "
            f"| {sum(r.variety_asks for r in runs)} "
            f"| {mean([r.wall_seconds for r in runs])} "
            f"| {sum(r.input_tokens for r in runs) / 1000:.0f} / "
            f"{sum(r.output_tokens for r in runs) / 1000:.0f} "
            f"| {mean([float(len(r.reply)) for r in ok if r.reply and not r.reply_tool_call])} "
            f"| {sum(r.reply_tool_call for r in ok)} |"
        )
    lines += [
        "",
        "| template | mode | analysis s | example slides labelled |",
        "|---|---|---|---|",
    ]
    for v in variants:
        lines.append(
            f"| {v.template} | {v.mode} | {v.analysis_seconds:.0f} | {v.labelled} |"
        )
    revisions = [a for r in results for a in r.attempts[1:] if a.tool_called]
    lines += [
        "",
        f"Revision calls with `changes`: {sum(a.changes for a in revisions)}/{len(revisions)}",
    ]
    edits = [r for r in results if r.edit is not None]
    if edits:
        lines += [
            "",
            "| edit turn | runs | success | first call | calls | calls with changes "
            "| slides kept identical / before | tokens in / out (k) |",
            "|---|---|---|---|---|---|---|---|",
        ]
        done = [r.edit for r in edits if r.edit is not None]
        lines.append(
            f"| all | {len(done)} | {sum(e.success for e in done)} "
            f"| {sum(e.first_call for e in done)} | {sum(e.calls for e in done)} "
            f"| {sum(e.calls_with_changes for e in done)} "
            f"| {sum(e.kept for e in done if e.success)} / "
            f"{sum(e.slides_before for e in done if e.success)} "
            f"| {sum(e.input_tokens for e in done) / 1000:.0f} / "
            f"{sum(e.output_tokens for e in done) / 1000:.0f} |"
        )
        lines += [
            f"- {r.prompt}: {'ok' if r.edit and r.edit.success else r.edit and r.edit.failure} "
            f"calls={r.edit.calls if r.edit else 0} changes={r.edit.calls_with_changes if r.edit else 0} "
            f"kept={r.edit.kept if r.edit else 0}/{r.edit.slides_before if r.edit else 0} "
            f"slides {r.edit.slides_before if r.edit else 0}→{r.edit.slides_after if r.edit else 0}"
            for r in edits
        ]
    failures = [r for r in results if not r.success]
    if failures:
        lines += ["", "Failures:"]
        lines += [f"- {r.template}/{r.mode}/{r.prompt}: {r.failure}" for r in failures]
    return "\n".join(lines) + "\n"


async def contact_sheet(pdf: bytes, target: Path) -> None:
    """Every page of the preview PDF, tiled four across."""
    import pypdfium2
    from PIL import Image

    document = pypdfium2.PdfDocument(pdf)
    # pypdfium2 types `scale` as int; it takes any positive float.
    pages = [
        document[i].render(scale=0.35).to_pil()  # pyright: ignore[reportArgumentType]
        for i in range(len(document))
    ]
    if not pages:
        return
    width, height = pages[0].size
    columns = 4
    rows = (len(pages) + columns - 1) // columns
    sheet = Image.new("RGB", (columns * width, rows * height), "white")
    for index, page in enumerate(pages):
        sheet.paste(page, ((index % columns) * width, (index // columns) * height))
    sheet.save(target)


# --- main ------------------------------------------------------------------------------


def _templates(
    values: Sequence[str], builtin: bool
) -> List[Tuple[str, Optional[Path]]]:
    found: List[Tuple[str, Optional[Path]]] = [("builtin", None)] if builtin else []
    for value in values:
        name, _, path = value.rpartition("=")
        found.append((name or Path(path).stem, Path(path)))
    return found


async def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--template",
        action="append",
        default=[t for t in os.environ.get("PPT_EVAL_TEMPLATES", "").split(",") if t],
        help="[name=]path.pptx, repeatable (env PPT_EVAL_TEMPLATES, comma-separated). "
        "The built-in template runs too (its stored rules analysis) unless --no-builtin.",
    )
    parser.add_argument("--no-builtin", action="store_true")
    parser.add_argument("--modes", default="vision,rules")
    parser.add_argument("--prompts", default=",".join(PROMPTS))
    parser.add_argument("--profile", default=os.environ.get("PPT_EVAL_PROFILE"))
    parser.add_argument(
        "--vision-profile", default=os.environ.get("PPT_EVAL_VISION_PROFILE")
    )
    parser.add_argument("--env-file", type=Path, default=DEFAULT_ENV)
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument(
        "--out", type=Path, default=Path(os.environ.get("PPT_EVAL_OUT", "target/eval"))
    )
    parser.add_argument("--contact-sheets", action="store_true")
    parser.add_argument(
        "--variety", choices=get_args(VarietyCheck), default="template_slide"
    )
    parser.add_argument("--max-uses", type=int, choices=range(1, 6), default=2)
    parser.add_argument(
        "--agent-prompt",
        type=Path,
        help="The agent author's prompt, in place of the short stand-in "
        "(e.g. eval/agent_prompts/agent_prompt_fr.txt).",
    )
    parser.add_argument(
        "--edit-turn",
        action="store_true",
        help="After a delivered deck, a second user turn asks to change it (prompts.EDITS).",
    )
    args = parser.parse_args()
    agent_prompt = (
        args.agent_prompt.read_text(encoding="utf-8").strip()
        if args.agent_prompt
        else _BASE_PROMPT
    )

    rebuild_ui_part_union([PptPreviewPart])
    # Base plans in memory: each run is its own conversation (session id).
    plan_store = deck_plans.MemoryDeckPlanStore()
    deck_plans.set_deck_plan_store_provider(lambda: plan_store)
    _load_secrets(args.env_file)
    agent = load_profile(args.catalog, args.profile)
    vision = load_profile(args.catalog, args.vision_profile or agent.profile_id)
    port = EvalModelPort(vision)
    modes = [m for m in args.modes.split(",") if m]
    prompt_ids = [p for p in args.prompts.split(",") if p]
    args.out.mkdir(parents=True, exist_ok=True)
    print(
        f"agent model {agent.profile_id}, vision model {vision.profile_id}, "
        f"variety check {args.variety}, max uses {args.max_uses}, "
        f"agent prompt {args.agent_prompt or 'stand-in'}"
    )

    variants: List[Variant] = []
    for name, path in _templates(args.template, not args.no_builtin):
        if path is None:
            content, inventory = builtin_template()
            variants.append(Variant(name, "rules", content, inventory, 0.0, {}))
            continue
        found = await analyse(name, path.read_bytes(), modes, port)
        for variant in found:
            print(
                f"analysed {name}/{variant.mode} in {variant.analysis_seconds:.0f}s: "
                f"{variant.labelled}"
            )
            (args.out / f"inventory-{name}-{variant.mode}.json").write_text(
                variant.inventory.model_dump_json(indent=1), encoding="utf-8"
            )
        variants += found

    semaphore = asyncio.Semaphore(RUN_CONCURRENCY)

    async def bounded(variant: Variant, prompt_id: str) -> RunResult:
        async with semaphore:
            result = await run_one(
                variant,
                prompt_id,
                agent,
                args.variety,
                args.max_uses,
                agent_prompt,
                args.edit_turn,
            )
        print(
            f"{variant.template}/{variant.mode}/{prompt_id}: "
            f"{'ok' if result.success else result.failure} repairs={result.repairs} "
            f"slides={result.slides} {result.wall_seconds}s"
        )
        return result

    results = await asyncio.gather(
        *(asyncio.create_task(bounded(v, p)) for v in variants for p in prompt_ids)
    )
    runs_dir = args.out / "runs"
    runs_dir.mkdir(exist_ok=True)
    for result in results:
        stem = f"{result.template}-{result.mode}-{result.prompt}"
        (runs_dir / f"{stem}.json").write_text(
            json.dumps(asdict(result), ensure_ascii=False, indent=1), encoding="utf-8"
        )
    summary = _summary(list(results), variants)
    (args.out / "summary.md").write_text(summary, encoding="utf-8")
    print(summary)
    if args.contact_sheets:
        await _sheets(
            list(results), variants, args.out / "sheets", args.variety, args.max_uses
        )


async def _sheets(
    results: List[RunResult],
    variants: List[Variant],
    out: Path,
    variety: VarietyCheck,
    max_uses: int,
) -> None:
    out.mkdir(exist_ok=True)
    by_key = {(v.template, v.mode): v for v in variants}
    for result in results:
        if not result.success or result.plan is None:
            continue
        variant = by_key[(result.template, result.mode)]
        workspace = MemoryWorkspace()
        # The run's settings, so the assigner picks the same slides; a first
        # call may be a variety ask, and the resend is built as sent.
        tool = _build_tool(
            variant, workspace, variety, max_uses, "sheets-" + result.prompt
        )
        _renders.set([])
        _plans.set([])
        await tool.ainvoke(result.plan)
        if not workspace.files:
            await tool.ainvoke(result.plan)
        pdf = next((b for k, b in workspace.files.items() if k.endswith(".pdf")), None)
        if pdf is not None:
            stem = f"{result.template}-{result.mode}-{result.prompt}"
            await contact_sheet(pdf, out / f"{stem}.png")


if __name__ == "__main__":
    asyncio.run(main())
