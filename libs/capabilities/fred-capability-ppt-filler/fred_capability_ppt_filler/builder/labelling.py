# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Vision labelling of example slides (design.md D2): a vision model looks at
each slide with its slots numbered and refines the heuristic kind, label,
reusable flag and slot purposes. Geometry and capacities never change, and
any failure leaves that slide's heuristic label in place."""

from __future__ import annotations

import asyncio
import base64
import io
import json
import logging
import time
from typing import Dict, List, Optional, Set, Tuple

from fred_sdk.contracts.runtime import (
    CapabilityModelPort,
    CapabilityModelUnavailableError,
)
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from pptx import Presentation
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from fred_capability_ppt_filler.builder import lexicon
from fred_capability_ppt_filler.builder.models import (
    SLIDE_KINDS,
    Pattern,
    SlideKind,
    TemplateInventory,
)
from fred_capability_ppt_filler.builder.preview import (
    marked_pattern_images,
    numbered_slots,
)

logger = logging.getLogger(__name__)

# One slide per call keeps the slot numbers unambiguous and a bad answer local;
# four in flight label ~40 slides in a few minutes without hogging the gateway.
LABELLING_CONCURRENCY = 4
LABELLING_CALL_TIMEOUT_SECONDS = 90.0

_KIND_GUIDE = {
    "cover": "title slide opening the deck",
    "agenda": "list of the deck's parts",
    "section": "divider introducing a part",
    "bullets": "title and a text body",
    "two_columns": "two parallel text blocks",
    "comparison": "two or more options compared side by side",
    "cards": "repeated cards or tiles, each with a heading and text",
    "figures": "key numbers or KPIs with captions",
    "process": "ordered steps",
    "timeline": "dated milestones",
    "team": "people with names and roles",
    "quote": "a quotation and its author",
    "image_text": "a picture next to text",
    "closing": "thank-you, contact or end slide",
    "other": "none of the above",
}

_SYSTEM_PROMPT = (
    "You classify one slide of a corporate PowerPoint template so it can be reused as an "
    "archetype for new content. Pink numbered boxes mark its fillable slots. Answer with one "
    "JSON object and nothing else:\n"
    '{"kind": <one kind>, "label": <short name of the layout, at most 6 words, e.g. '
    '"Four KPI cards">, "slots": {"<slot number>": <what the slot holds on this slide, a few '
    'words, e.g. "slide title", "KPI value", "step description">}, "max_items": <number of '
    'repeated items (cards, steps, entries) the layout holds, or null>, "reusable": <false '
    "when the slide is a one-off that cannot carry other content, such as a specific org chart, "
    "a legal notice or a fixed decorative visual; otherwise true>}\n"
    "Kinds: " + "; ".join(f"{kind}: {text}" for kind, text in _KIND_GUIDE.items()) + "."
)


class SlideLabel(BaseModel):
    """The model's answer for one slide, validated before any merge."""

    model_config = ConfigDict(extra="ignore")

    kind: SlideKind
    label: str = Field(min_length=1, max_length=80)
    slots: Dict[int, str] = Field(default_factory=dict)
    max_items: Optional[int] = Field(default=None, ge=0, le=50)
    reusable: bool


def _text_of(message: AIMessage) -> str:
    content = message.content
    if isinstance(content, str):
        return content
    return "".join(
        part if isinstance(part, str) else str(part.get("text", "")) for part in content
    )


def parse_label(message: AIMessage, slot_count: int) -> Optional[SlideLabel]:
    """The validated answer, or ``None`` when it is not usable as is."""
    text = _text_of(message)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        return None
    try:
        answer = SlideLabel.model_validate(json.loads(text[start : end + 1]))
    except (ValueError, ValidationError):
        return None
    purposes = {number: " ".join(p.split()) for number, p in answer.slots.items()}
    if any(not 1 <= n <= slot_count or not p for n, p in purposes.items()):
        return None
    if any(len(p) > 80 for p in purposes.values()):
        return None
    return answer.model_copy(update={"slots": purposes, "label": answer.label.strip()})


def _messages(pattern: Pattern, png: bytes) -> List[SystemMessage | HumanMessage]:
    lines = [f"Slide {pattern.source_index}. Slots:"]
    for number, zone in enumerate(numbered_slots(pattern), start=1):
        sample = " ".join(zone.sample.split())[:80]
        lines.append(f"{number}: {zone.role}" + (f' - "{sample}"' if sample else ""))
    url = "data:image/png;base64," + base64.b64encode(png).decode("ascii")
    return [
        SystemMessage(content=_SYSTEM_PROMPT),
        HumanMessage(
            content=[
                {"type": "text", "text": "\n".join(lines)},
                {"type": "image_url", "image_url": {"url": url}},
            ]
        ),
    ]


def _annotated(template_bytes: bytes) -> Dict[str, Tuple[bool, bool]]:
    """Per slide pattern id: whether its notes declare a (kind, label), which
    outrank the vision answer as they outrank the heuristics."""
    declared: Dict[str, Tuple[bool, bool]] = {}
    for slide in Presentation(io.BytesIO(template_bytes)).slides:
        frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        notes = frame.text if frame is not None else ""
        match = lexicon.KIND_NOTE.search(notes)
        kind = match.group(1).lower() if match else ""
        kind = lexicon.LEGACY_KINDS.get(kind, kind)
        declared[f"slide:{slide.slide_id}"] = (
            kind in SLIDE_KINDS,
            lexicon.LABEL_NOTE.search(notes) is not None,
        )
    return declared


def _prepare(
    template_bytes: bytes,
    patterns_pdf: bytes,
    inventory: TemplateInventory,
    ids: Set[str],
) -> Tuple[Dict[str, bytes], Dict[str, Tuple[bool, bool]]]:
    return (
        marked_pattern_images(patterns_pdf, inventory, ids),
        _annotated(template_bytes),
    )


def _merge(
    pattern: Pattern, answer: SlideLabel, declared: Tuple[bool, bool]
) -> Pattern:
    merged = pattern.model_copy(deep=True)
    if not declared[0]:
        merged.kind = answer.kind
    if not declared[1]:
        merged.label = answer.label[:60]
    merged.reusable = answer.reusable
    for number, zone in enumerate(numbered_slots(merged), start=1):
        zone.purpose = answer.slots.get(number)
    merged.labelled_by = "vision"
    return merged


async def label_inventory(
    inventory: TemplateInventory,
    template_bytes: bytes,
    patterns_pdf: Optional[bytes],
    port: Optional[CapabilityModelPort],
    *,
    profile_id: Optional[str] = None,
    use_vision: bool = True,
) -> TemplateInventory:
    """Refine the included example slides with a vision model.

    ``patterns_pdf`` is the patterns deck converted to PDF (page N = pattern N).
    No call is made when vision is off, no port or PDF is available; the port
    refusing (no vision-capable model) stops all calls. Unlabelled slides keep
    their heuristic label and ``labelled_by="heuristic"``.
    """
    targets = [p for p in inventory.patterns if p.source == "slide" and p.included]
    if not use_vision or port is None or not patterns_pdf or not targets:
        return inventory
    started = time.monotonic()
    images, declared = await asyncio.to_thread(
        _prepare, template_bytes, patterns_pdf, inventory, {p.id for p in targets}
    )
    semaphore = asyncio.Semaphore(LABELLING_CONCURRENCY)
    refused: List[str] = []

    async def label_one(pattern: Pattern) -> Optional[SlideLabel]:
        async with semaphore:
            if refused:
                return None
            try:
                message = await asyncio.wait_for(
                    port.ainvoke(
                        _messages(pattern, images[pattern.id]),
                        profile_id=profile_id,
                        require_vision=True,
                    ),
                    LABELLING_CALL_TIMEOUT_SECONDS,
                )
            except CapabilityModelUnavailableError as exc:
                refused.append(str(exc))
                return None
            except Exception:  # provider/network errors only cost this slide
                logger.warning(
                    "Vision labelling of %s failed", pattern.id, exc_info=True
                )
                return None
        answer = parse_label(message, len(numbered_slots(pattern)))
        if answer is None:
            logger.warning("Vision labelling of %s: invalid answer", pattern.id)
        return answer

    labelled = [p for p in targets if p.id in images]
    answers = await asyncio.gather(*(label_one(p) for p in labelled))
    merged = {
        pattern.id: _merge(pattern, answer, declared.get(pattern.id, (False, False)))
        for pattern, answer in zip(labelled, answers)
        if answer is not None
    }
    if refused:
        logger.info("Vision labelling skipped: %s", refused[0])
    logger.info(
        "Vision labelled %d/%d example slides in %.1fs",
        len(merged),
        len(targets),
        time.monotonic() - started,
    )
    if not merged:
        return inventory
    patterns = [merged.get(p.id) or p.model_copy() for p in inventory.patterns]
    # Same rule as the analysis: kinds may have moved, so `preferred` may too.
    layout_kinds = {p.kind for p in patterns if p.source == "layout" and p.included}
    for pattern in patterns:
        if pattern.source == "slide":
            pattern.preferred = (
                pattern.included and pattern.reusable and pattern.kind in layout_kinds
            )
    return inventory.model_copy(update={"patterns": patterns})
