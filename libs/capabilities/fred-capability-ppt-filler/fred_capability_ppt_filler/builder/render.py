# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Render a :class:`DeckPlan` on a template into ``.pptx`` bytes.

Synchronous, CPU-bound: the caller runs it off the event loop under the
package's heavy-job slot. Works on a copy of the template: layout patterns
are instantiated, example-slide patterns are duplicated (shapes deep-copied,
picture/media relationships re-created on the new slide), every zone is
filled by role, unused zones are emptied or removed, example text left on a
duplicated slide is swept (`checks`), and the template's own
slides are dropped at the end so the output holds the plan's slides only, in
plan order.
"""

from __future__ import annotations

import copy
import io
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import (
    Any,
    Collection,
    Dict,
    List,
    Mapping,
    Optional,
    Sequence,
    Tuple,
    cast,
)

from pptx.oxml.ns import qn
from pptx.parts.chart import ChartPart
from pptx.util import Emu

from fred_capability_ppt_filler import text_metrics
from fred_capability_ppt_filler.builder import (
    charts,
    checks,
    lexicon,
    notes,
    table_intents,
)
from fred_capability_ppt_filler.builder.analysis import (
    _cell_text,
    _font_name_of,
    estimate_capacity,
    theme_fonts,
)
from fred_capability_ppt_filler.builder.models import (
    TABLE_SLACK_PT,
    Pattern,
    RepeatGroup,
    TemplateInventory,
    Zone,
)
from fred_capability_ppt_filler.builder.plan import (
    MIN_FONT_SCALE,
    TREND_GLYPHS,
    DeckPlan,
    PlanChart,
    PlanItem,
    PlanTable,
    SlidePlan,
    chart_headings,
    has_content,
    shown_charts,
    slide_charts,
)
from fred_capability_ppt_filler.fill import BuildNoteCode, BuiltSlide
from fred_capability_ppt_filler.images import fit_inside_box
from fred_capability_ppt_filler.inline_markdown import parse_inline_markdown
from fred_capability_ppt_filler.pptx_archive import open_presentation
from fred_capability_ppt_filler.traversal import PlacedShape, walk_shapes

# Sub-bullets get narrower lines: this factor per level in the capacity estimate.
_LEVEL_WIDTH_FACTOR = 0.85
# Slack over the measured width when sizing a box that must not wrap: the
# renderer's font is not always the one measured here (see `text_metrics`).
_NO_WRAP_MARGIN = 1.06

_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_MEDIA_RELS = ("/image", "/media", "/video", "/audio", "/hdphoto")
_HYPERLINK_REL = "/hyperlink"
# A chart's axis fields, in the order a scatter (first three) and a bubble draw them.
_AXIS_FIELDS = ("x_title", "x_unit", "y_title", "size_title", "size_unit")
# A title in the user-facing summary: enough to recognise the slide.
_SUMMARY_TITLE_CHARS = 120
_SECTIONS_EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
# Elements meaningless without their uncopied target (a link to another slide).
_DROPPED_LINKS = frozenset(
    qn(tag) for tag in ("a:hlinkClick", "a:hlinkHover", "p:tags")
)

PreparedImage = Tuple[bytes, int, int]


@dataclass
class RenderResult:
    pptx: bytes
    slide_count: int = 0
    # The same deck adjusted for the PDF preview only (see `_widen_no_wrap_text`).
    preview_pptx: bytes = b""
    # The same findings per slide, for a structured result message.
    report: checks.DeckReport = field(default_factory=checks.DeckReport)


_FILLS = ("a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill")


def _is_bare_text_box(sp) -> bool:
    """A text box drawing nothing but its text: widening it moves no visible edge."""
    non_visual = sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvSpPr')}")
    props = sp.find(qn("p:spPr"))
    if non_visual is None or non_visual.get("txBox") != "1" or props is None:
        return False
    if sp.find(qn("p:style")) is not None:
        return False
    if any(props.find(qn(tag)) is not None for tag in _FILLS):
        return False
    line = props.find(qn("a:ln"))
    return line is None or line.find(qn("a:noFill")) is not None


def _widen_no_wrap_text(presentation) -> int:
    """Give every one-line text box the room it would take in PowerPoint.

    LibreOffice ignores "do not wrap" and breaks such a box as soon as its
    substitute font runs wider than the designer's ("Merc|i"). Widening it in
    the preview copy, anchored on its alignment, changes nothing a reader of the
    .pptx would see; the delivered file is never touched.
    """
    width = presentation.slide_width
    widened = 0
    parts = [slide for slide in presentation.slides]
    for master in presentation.slide_masters:
        parts.append(master)
        parts.extend(master.slide_layouts)
    for part in parts:
        for sp in part.shapes._spTree.iter(qn("p:sp")):
            body = sp.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
            xfrm = sp.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
            if (
                body is None
                or body.get("wrap") != "none"
                or body.get("vert") not in (None, "horz")
                or xfrm is None
                or xfrm.get("rot") not in (None, "0")
                or sp.find(f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}") is not None
                or not _is_bare_text_box(sp)
                or not "".join(t.text or "" for t in sp.iter(qn("a:t"))).strip()
            ):
                continue
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            if off is None or ext is None:
                continue
            low, high = 0, width
            parent = sp.getparent()
            if parent is not None and parent.tag == qn("p:grpSp"):
                ch_off = parent.find(
                    f"{qn('p:grpSpPr')}/{qn('a:xfrm')}/{qn('a:chOff')}"
                )
                ch_ext = parent.find(
                    f"{qn('p:grpSpPr')}/{qn('a:xfrm')}/{qn('a:chExt')}"
                )
                if ch_off is None or ch_ext is None:
                    continue
                low = int(ch_off.get("x", "0"))
                high = low + int(ch_ext.get("cx", "0"))
            left, box = int(off.get("x", "0")), int(ext.get("cx", "0"))
            ppr = sp.find(f"{qn('p:txBody')}/{qn('a:p')}/{qn('a:pPr')}")
            align = ppr.get("algn") if ppr is not None else None
            if align == "ctr":
                center = left + box // 2
                half = min(center - low, high - center)
                new_left, new_box = center - half, 2 * half
            elif align == "r":
                new_left, new_box = low, left + box - low
            else:
                new_left, new_box = left, high - left
            if new_box > box:
                off.set("x", str(new_left))
                ext.set("cx", str(new_box))
                widened += 1
    return widened


@dataclass
class _Line:
    text: str
    level: int = 0
    # A column heading written into the body for want of a heading zone. It
    # drops the bullet glyph, so the reader can still tell it from the content.
    heading: bool = False
    # The plan element it comes from ("bullets[2]", "items[0].text"), so a
    # misfit can name what to shorten.
    source: str = ""


_TEMPLATE_CHANGED = (
    "The template changed since its analysis: tell the user that the agent's owner must "
    "save the template again in the agent settings."
)


class RenderError(ValueError):
    """The plan cannot be rendered on this template (the agent's correctable mistake)."""


# --- slide creation -------------------------------------------------------------


def _all_layouts(presentation) -> list:
    return [
        layout
        for master in presentation.slide_masters
        for layout in master.slide_layouts
    ]


def _copy_relationships(
    source_part, dest_part, shared: Sequence[str] = _MEDIA_RELS
) -> Dict[str, str]:
    """Re-create the media and external-link relationships on ``dest_part``.

    Targets are shared, so an image copied ten times stays one part. Links to
    other slides are not copied: they would keep removed template slides alive.
    """
    rid_map: Dict[str, str] = {}
    for rId, rel in source_part.rels.items():
        if rel.is_external:
            if rel.reltype.endswith(_HYPERLINK_REL):
                rid_map[rId] = dest_part.rels.get_or_add_ext_rel(
                    rel.reltype, rel.target_ref
                )
        elif rel.reltype.endswith(tuple(shared)):
            rid_map[rId] = dest_part.relate_to(rel.target_part, rel.reltype)
    return rid_map


def _copy_shapes(source_tree, dest_tree, rid_map: Mapping[str, str]) -> None:
    for element in source_tree:
        if element.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        clone = copy.deepcopy(element)
        for node in list(clone.iter()):
            # Every relationship-namespace attribute, not a fixed list.
            for attr, old in list(node.attrib.items()):
                if not attr.startswith(_R_NS) or not old:
                    continue
                if old in rid_map:
                    node.set(attr, rid_map[old])
                elif node.tag in _DROPPED_LINKS:
                    parent = node.getparent()
                    if parent is not None:
                        parent.remove(node)
                    break
                else:
                    del node.attrib[attr]
        dest_tree.append(clone)


def _copy_notes(source, dest) -> None:
    if not source.has_notes_slide:
        return
    src = source.notes_slide
    tree = dest.notes_slide.shapes._spTree
    for element in list(tree):
        if element.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(element)
    _copy_shapes(
        src.shapes._spTree, tree, _copy_relationships(src.part, dest.notes_slide.part)
    )


def _clone_part(part: Any, package: Any) -> Any:
    """A copy of ``part`` under the next free name of its own series."""
    template = re.sub(r"\d*(\.\w+)$", r"%d\1", str(part.partname))
    name = package.next_partname(template)
    return type(part).load(name, part.content_type, package, part.blob)


def _clone_charts(source_part: Any, dest_part: Any) -> Dict[str, str]:
    """Each chart of ``source_part`` copied for ``dest_part`` with its workbook and
    style parts, so the copy can take new data. A chart that does not copy cleanly
    is left out: its frame loses its link and is redrawn or removed."""
    rid_map: Dict[str, str] = {}
    for rId, rel in source_part.rels.items():
        if rel.is_external or not rel.reltype.endswith("/chart"):
            continue
        chart = rel.target_part
        inner = list(chart.rels.items())
        # A linked workbook or a drawing with its own parts: not copied.
        if not isinstance(chart, ChartPart) or any(
            r.is_external or r.target_part.rels for _, r in inner
        ):
            continue
        clone = _clone_part(chart, dest_part.package)
        new_rid = dest_part.relate_to(clone, rel.reltype)
        renamed = {
            old: clone.relate_to(_clone_part(r.target_part, clone.package), r.reltype)
            for old, r in inner
        }
        for node in clone._element.iter():
            for attr, old in list(node.attrib.items()):
                if attr.startswith(_R_NS) and old in renamed:
                    node.set(attr, renamed[old])
        rid_map[rId] = new_rid
    return rid_map


def _duplicate_slide(presentation, source, preview: bool = False):
    """Append a copy of ``source``: layout, background, colour map, shapes,
    media, charts and notes. A ``preview`` copy shares its charts' parts: fine
    for a thumbnail, never for a delivered deck."""
    dest = presentation.slides.add_slide(source.slide_layout)
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)

    background = source._element.cSld.find(qn("p:bg"))
    if background is not None:
        dest._element.cSld.insert(0, copy.deepcopy(background))
    # A dark example slide maps its text and background colours here.
    override = source._element.find(qn("p:clrMapOvr"))
    if override is not None:
        current = dest._element.find(qn("p:clrMapOvr"))
        if current is not None:
            dest._element.remove(current)
        dest._element.cSld.addnext(copy.deepcopy(override))

    if preview:
        rid_map = _copy_relationships(source.part, dest.part, _MEDIA_RELS + ("/chart",))
    else:
        rid_map = _copy_relationships(source.part, dest.part)
        rid_map.update(_clone_charts(source.part, dest.part))
    _copy_shapes(source.shapes._spTree, dest.shapes._spTree, rid_map)
    _copy_notes(source, dest)
    return dest


def _remove_slides(presentation, slide_ids: Sequence[int]) -> None:
    """Drop the template's own slides, and its sections with them."""
    wanted = set(slide_ids)
    id_list = presentation.slides._sldIdLst
    for sldId in list(id_list):
        if sldId.id in wanted:
            presentation.part.drop_rel(sldId.rId)
            id_list.remove(sldId)
    # Sections list the template's slides, never the copies: kept, they make
    # PowerPoint repair the file, and their names mean nothing in a new deck.
    for ext in presentation.part._element.findall(f"{qn('p:extLst')}/{qn('p:ext')}"):
        if ext.get("uri") == _SECTIONS_EXT_URI:
            ext.getparent().remove(ext)


# --- shape addressing -----------------------------------------------------------


def _shape_index(slide) -> Dict[str, PlacedShape]:
    """Zone key → shape for one slide (placeholder idx and shape-id path keys)."""
    index: Dict[str, PlacedShape] = {}
    for placed in walk_shapes(slide.shapes):
        shape = cast(Any, placed.shape)
        index["sp:" + "/".join(str(p) for p in placed.path)] = placed
        if shape.is_placeholder:
            index.setdefault(f"ph:{shape.placeholder_format.idx}", placed)
    return index


def _box_of(placed: PlacedShape) -> Tuple[int, int, int, int]:
    return (
        placed.left,
        placed.top,
        placed.left + placed.width,
        placed.top + placed.height,
    )


def _axis_cover(low: int, high: int, other_low: int, other_high: int) -> float:
    span = high - low
    if span <= 0:
        # A connector is a line: zero-thick on one axis, and any area ratio
        # scores it 0 however squarely it sits inside the other shape.
        return 1.0 if other_low <= low <= other_high else 0.0
    return max(0, min(high, other_high) - max(low, other_low)) / span


def _covered(box: Tuple[int, int, int, int], other: Tuple[int, int, int, int]) -> float:
    """Share of ``box`` that lies inside ``other``, lines included."""
    return _axis_cover(box[0], box[2], other[0], other[2]) * _axis_cover(
        box[1], box[3], other[1], other[3]
    )


def _own_geometry(shape) -> None:
    """Give a placeholder that inherits its position its own: python-pptx
    creates the `xfrm` on the first write and zeroes what was not written."""
    if shape.is_placeholder and shape._element.spPr.find(qn("a:xfrm")) is None:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        shape.left, shape.top, shape.width, shape.height = left, top, width, height


def _union_box(shapes: Sequence[PlacedShape]) -> Tuple[int, int, int, int]:
    boxes = [_box_of(p) for p in shapes]
    return (
        min(b[0] for b in boxes),
        min(b[1] for b in boxes),
        max(b[2] for b in boxes),
        max(b[3] for b in boxes),
    )


def _is_container(key: str, item: Sequence[Zone]) -> bool:
    """A grouped item is addressed by its group; a row item by its own zones."""
    return all(zone.key != key for zone in item)


def _item_shapes(
    group: RepeatGroup, index: Dict[str, PlacedShape]
) -> List[List[PlacedShape]]:
    """The shapes making up each item of ``group``, in item order."""
    out: List[List[PlacedShape]] = []
    for key, item in zip(group.item_keys, group.items):
        keys = [key] if _is_container(key, item) else [z.key for z in item]
        out.append([p for p in (index.get(k) for k in keys) if p])
    return out


# Below this the shift is not worth the risk of nudging a deliberate alignment.
_RECENTRE_MIN_EMU = 91440


def _remove(placed: PlacedShape) -> None:
    element = cast(Any, placed.shape)._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


# --- text -----------------------------------------------------------------------------


def _lines_from_bullets(bullets: Sequence[str], source: str = "") -> List[_Line]:
    """``source`` names the element; a ``{}`` in it takes the bullet's index."""
    lines: List[_Line] = []
    for index, raw in enumerate(bullets):
        text = raw.rstrip()
        if not text.strip():
            continue
        level = 0
        stripped = text.lstrip()
        indent = len(text) - len(stripped)
        while stripped.startswith(("- ", "* ", "• ")):
            level += 1
            stripped = stripped[2:].lstrip()
        if level == 0 and indent >= 2:
            level = 1
        lines.append(
            _Line(stripped, min(level, 2), source=source.replace("{}", str(index)))
        )
    return lines


def _paragraph_style(paragraph) -> Tuple[int, Any, Any]:
    """``(level, pPr, rPr)`` of a template paragraph; ``rPr`` falls back on
    the end-run properties of an empty paragraph."""
    pPr = paragraph.find(qn("a:pPr"))
    level = int(pPr.get("lvl", "0")) if pPr is not None else 0
    run = paragraph.find(qn("a:r"))
    rPr = run.find(qn("a:rPr")) if run is not None else None
    if rPr is None:
        rPr = paragraph.find(qn("a:endParaRPr"))
    if rPr is not None:
        rPr = copy.deepcopy(rPr)
        rPr.tag = qn("a:rPr")
    return level, copy.deepcopy(pPr) if pPr is not None else None, rPr


def _write_lines(shape, lines: Sequence[_Line]) -> None:
    """Replace a text frame's content, inheriting the template's style per level.

    The first paragraph of each level gives the ``pPr``/``rPr`` every new line
    of that level clones; a level the template lacks takes the first
    paragraph's. Inline ``**bold**``/``*italic*`` is honoured per span.
    """
    _write_paragraphs(shape.text_frame._txBody, lines)


def _write_paragraphs(txBody, lines: Sequence[_Line]) -> None:
    """:func:`_write_lines` on a bare ``txBody`` (a table cell has no shape)."""
    paragraphs = txBody.findall(qn("a:p"))
    styles: Dict[int, Tuple[Any, Any]] = {}
    base: Tuple[Any, Any] = (None, None)
    for index, paragraph in enumerate(paragraphs):
        level, pPr, rPr = _paragraph_style(paragraph)
        if index == 0:
            base = (pPr, rPr)
        styles.setdefault(level, (pPr, rPr))
    for paragraph in paragraphs:
        txBody.remove(paragraph)

    for line in lines or [_Line("")]:
        base_pPr, base_rPr = styles.get(line.level, base)
        p = txBody.add_p()
        if base_pPr is not None:
            p.insert(0, copy.deepcopy(base_pPr))
        if line.level or (base_pPr is not None and base_pPr.get("lvl")):
            p.get_or_add_pPr().set("lvl", str(line.level))
        if line.heading:
            pPr = p.get_or_add_pPr()
            for existing in pPr.findall(qn("a:buChar")) + pPr.findall(
                qn("a:buAutoNum")
            ):
                pPr.remove(existing)
            pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        for span in parse_inline_markdown(line.text):
            r = p.add_r()
            if base_rPr is not None:
                r.insert(0, copy.deepcopy(base_rPr))
            r.text = span.text
            if span.bold or span.italic:
                rPr = r.get_or_add_rPr()
                if span.bold:
                    rPr.set("b", "1")
                if span.italic:
                    rPr.set("i", "1")


# The arrow before a key figure, against the figure's own size.
_GLYPH_SCALE = 0.6


def _shrink_trend_glyph(shape, base_pt: Optional[float]) -> None:
    """Split a figure's leading arrow into its own run, smaller, in the
    figure's own font and colour."""
    run = shape.text_frame._txBody.find(f"{qn('a:p')}/{qn('a:r')}")
    text = run.text if run is not None else ""
    if not text or text[0] not in TREND_GLYPHS.values():
        return
    glyph = copy.deepcopy(run)
    glyph.text, run.text = text[:2], text[2:]
    run.addprevious(glyph)
    rPr = glyph.get_or_add_rPr()
    size = int(rPr.get("sz")) / 100 if rPr.get("sz") else base_pt
    if size:
        rPr.set("sz", str(max(600, int(size * _GLYPH_SCALE * 100))))


def _clear_text(shape) -> None:
    _write_lines(shape, [])


def _set_font_scale(shape, scale: float, base_pt: Optional[float] = None) -> None:
    """Shrink a text frame. With a known base size every run gets an explicit
    scaled size, which every renderer honours; otherwise fall back to
    ``normAutofit`` (``fontScale`` in 1/1000 %)."""
    txBody = shape.text_frame._txBody
    if base_pt:
        _scale_runs(txBody, scale, base_pt)
        return
    bodyPr = txBody.find(qn("a:bodyPr"))
    if bodyPr is None:
        return
    for child in list(bodyPr):
        if child.tag in (qn("a:noAutofit"), qn("a:spAutoFit"), qn("a:normAutofit")):
            bodyPr.remove(child)
    autofit = bodyPr.makeelement(qn("a:normAutofit"), {})
    autofit.set("fontScale", str(int(scale * 100000)))
    autofit.set("lnSpcReduction", "10000")
    bodyPr.append(autofit)


def _scale_runs(element, scale: float, base_pt: float) -> None:
    """Give every run under ``element`` an explicit size, scaled."""
    for tag in ("a:rPr", "a:endParaRPr"):
        for rPr in element.iter(qn(tag)):
            current = rPr.get("sz")
            size = int(current) / 100 if current else base_pt
            rPr.set("sz", str(max(600, int(size * scale * 100))))
    for run in element.iter(qn("a:r")):
        if run.find(qn("a:rPr")) is None:
            rPr = run.makeelement(
                qn("a:rPr"), {"sz": str(max(600, int(base_pt * scale * 100)))}
            )
            run.insert(0, rPr)


# PowerPoint's default text inset, used when a box declares none.
_DEFAULT_INSET_EMU = 91440


def _text_insets(shape) -> int:
    """Left plus right padding of a text box: the text is laid inside them."""
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    if body is None:
        return 2 * _DEFAULT_INSET_EMU
    return sum(int(body.get(side, _DEFAULT_INSET_EMU)) for side in ("lIns", "rIns"))


def _grows_with_text(shape, inherited=None) -> bool:
    """A box that does not wrap: its width follows its text, so it is sized
    rather than fitted.

    Auto-sizing the *height* is a different thing and must not come here: such a
    box still wraps inside its width, and widening it pushes the text out of the
    design (a cover title over the photo beside it). It grows downwards without
    complaint, which is exactly why its line budget has to be enforced — off the
    bottom of the slide is not "fitted".

    A placeholder cloned from a layout carries a bare ``bodyPr`` and draws with
    the layout's rule, so the layout's own is what decides.
    """
    for bodyPr in (shape.text_frame._txBody.find(qn("a:bodyPr")), inherited):
        if bodyPr is not None and bodyPr.get("wrap") == "none":
            return True
    return False


def _sibling_key(zone: Zone, pattern: Pattern) -> str:
    """Zones that play the same part on one slide: the same role, and for a
    repeated group the same position inside each item."""
    for group in pattern.groups:
        for item in group.items:
            for position, member in enumerate(item):
                if member.key == zone.key:
                    return f"{group.key}/{position}/{zone.role}"
    return f"slide/{zone.role}"


def _layout_body_props(slide) -> Dict[str, Any]:
    """``bodyPr`` of each layout placeholder, by the zone key it answers to."""
    props: Dict[str, Any] = {}
    try:
        placeholders = slide.slide_layout.placeholders
    except (AttributeError, KeyError):
        return props
    for shape in placeholders:
        if not shape.has_text_frame:
            continue
        body = shape.text_frame._txBody.find(qn("a:bodyPr"))
        if body is not None:
            props[f"ph:{shape.placeholder_format.idx}"] = body
    return props


def zone_capacity(
    zone: Zone,
    width: int,
    height: int,
    font_name: Optional[str] = None,
    bold: bool = False,
) -> Tuple[int, int]:
    """(chars_per_line, max_lines) the analysis stored; only an inventory saved
    without one is measured here, so the agent, the fit check and the renderer
    read the same numbers."""
    if zone.chars_per_line and zone.max_lines:
        return zone.chars_per_line, zone.max_lines
    return estimate_capacity(
        zone.box, zone.font_pt, zone.role, width, height, font_name, bold
    )


def _line_cost(
    line: _Line, room_em: float, font_name: Optional[str], bold: bool
) -> int:
    """Lines one entry wraps to: its own paragraph, so a line at least."""
    room = room_em * _LEVEL_WIDTH_FACTOR**line.level
    return max(1, math.ceil(text_metrics.width_em(line.text, font_name, bold) / room))


def floor_lines(
    zone: Zone,
    lines: List[_Line],
    width: int,
    height: int,
    font_name: Optional[str] = None,
    bold: bool = False,
) -> List[int]:
    """Lines each entry wraps to at the smallest font, as `_fit` counts them."""
    chars_per_line, _ = zone_capacity(zone, width, height, font_name, bold)
    room_em = chars_per_line * text_metrics.char_em(font_name, bold) / MIN_FONT_SCALE
    return [_line_cost(line, room_em, font_name, bold) for line in lines]


def _fit(
    zone: Zone,
    lines: List[_Line],
    width: int,
    height: int,
    font_name: Optional[str] = None,
    bold: bool = False,
) -> Tuple[List[_Line], Optional[float], bool]:
    """Fit ``lines`` into ``zone``: returns ``(lines, font scale or None, truncated)``."""
    chars_per_line, max_lines = zone_capacity(zone, width, height, font_name, bold)
    # How many em of text one line holds: each line is measured in the real
    # font, so a line of capitals counts for what it draws.
    line_em = chars_per_line * text_metrics.char_em(font_name, bold)

    def needed(scale: float) -> int:
        return sum(_line_cost(line, line_em / scale, font_name, bold) for line in lines)

    # A word longer than a line breaks in the middle: shrink until it fits.
    longest = max(
        (
            text_metrics.width_em(word, font_name, bold)
            for line in lines
            for word in line.text.split()
        ),
        default=0.0,
    )
    word_scale = min(1.0, line_em / longest) if longest else 1.0
    if needed(1.0) <= max_lines and word_scale >= 1.0:
        return lines, None, False
    scale = min(max_lines / needed(1.0), word_scale)
    if scale >= MIN_FONT_SCALE:
        return lines, round(scale, 2), False
    if needed(MIN_FONT_SCALE) <= max_lines:
        # A word longer than the line (a URL, an address) would drag the whole
        # zone below the floor on its own. It breaks instead: illegible text
        # helps nobody, and the break shows the agent what to shorten.
        return lines, MIN_FONT_SCALE, False

    capacity = max_lines / MIN_FONT_SCALE
    room_em = line_em / MIN_FONT_SCALE
    per_char = text_metrics.char_em(font_name, bold) or text_metrics.FALLBACK_CHAR_EM
    kept: List[_Line] = []
    used = 0.0
    dropped = False
    for line in lines:
        cost = _line_cost(line, room_em, font_name, bold)
        if used + cost > capacity:
            room = int((capacity - used) * room_em / per_char) - 1
            # Never leave the zone blank: a cut first line beats an empty box.
            if not kept:
                kept.append(_Line(_cut(line.text, max(room, 1)), line.level))
            dropped = True
            break
        kept.append(line)
        used += cost
    # Everything fitted once shrunk to the floor: shrinking is not truncating,
    # and saying it is sends the agent rewriting text that was never cut.
    return kept, MIN_FONT_SCALE, dropped


def _cut(text: str, room: int) -> str:
    if room >= len(text):
        return text
    cut = text[:room]
    if " " in cut:
        cut = cut[: cut.rfind(" ")]
    return cut.rstrip(" ,;:") + "…"


def _auto_number(sample: str, index: int) -> str:
    if re.fullmatch(r"0\d", sample.strip()):
        return f"{index + 1:02d}"
    if sample.strip().isdigit():
        return str(index + 1)
    return ""


# --- tables -------------------------------------------------------------------------------

_TABLE_SCALES = (1.0, 0.95, 0.9, 0.85, 0.8, 0.75, MIN_FONT_SCALE)
_A16 = "{http://schemas.microsoft.com/office/drawing/2014/main}"


@dataclass
class TableFit:
    """How a plan table sits in a template table: the fit check and the
    renderer both read it. ``problem`` is ``(what, fix)`` when it does not fit."""

    header: Optional[List[str]] = None
    rows: List[List[str]] = field(default_factory=list)
    widths: List[int] = field(default_factory=list)
    scale: Optional[float] = None
    heights: List[float] = field(default_factory=list)  # pt, header row first
    problem: Optional[Tuple[str, str]] = None
    # Plan row i is fit row i - shift: 1 when the plan's first row became the
    # header, -1 when its header went into the body.
    shift: int = 0


def _cell_lines(widths_em: Sequence[float], room_em: float) -> int:
    return sum(max(1, math.ceil(width / room_em)) for width in widths_em)


def fit_table(zone: Zone, table: PlanTable) -> TableFit:
    """Columns dropped from the right (others widened), rows grown or removed,
    and the one font scale at which every row fits the template's height.
    Cells are never cut: what does not fit is a problem."""
    spec = zone.table
    if spec is None:
        return TableFit(problem=("table: this slide has no table", "drop table"))
    count, most = table.width, len(spec.columns)
    header = list(table.header) if any(table.header) else None
    rows, first = [list(r) for r in table.rows], 0
    if spec.header and header is None and len(rows) >= 2:
        header, rows, first = rows[0], rows[1:], 1
    elif not spec.header and header is not None:
        rows, header, first = [header] + rows, None, -1
    if count > most:
        return TableFit(
            problem=(
                f"table: {count} columns given, max {most}",
                f"keep at most {most} columns or split the table over two slides",
            )
        )
    if count < 2:
        return TableFit(
            problem=("table: a table needs at least 2 columns", "write it as bullets")
        )
    if not rows:
        return TableFit(
            problem=("table: no rows under the header", "add rows, or use bullets")
        )
    if spec.header and header is None:
        header = [""] * count
    kept = spec.widths[:count]
    widths = [round(w * sum(spec.widths) / sum(kept)) for w in kept]
    widths[-1] += sum(spec.widths) - sum(widths)
    columns = [
        max(1, int(c * max(1, w2 - spec.margin) / max(1, w - spec.margin)))
        for c, w, w2 in zip(spec.columns, kept, widths)
    ]
    char = text_metrics.char_em()
    # Measured once: only the room per line changes with the scale.
    measured = [
        [[text_metrics.width_em(line) for line in t.splitlines() or [""]] for t in row]
        for row in ([header] if header is not None else []) + rows
    ]

    def lines(scale: float) -> List[int]:
        return [
            max(_cell_lines(t, c * char / scale) for t, c in zip(row, columns))
            for row in measured
        ]

    for scale in _TABLE_SCALES:
        found = [spec.row_height(n, scale) for n in lines(scale)]
        if sum(found) <= spec.height_pt + TABLE_SLACK_PT:
            return TableFit(
                header,
                rows,
                widths,
                None if scale == 1.0 else scale,
                found,
                None,
                first,
            )

    full, floor = lines(1.0), lines(MIN_FONT_SCALE)
    head = 1 if header is not None else 0
    holds = spec.body_lines(1.0, full[0] if head else 1)
    holds_floor = spec.body_lines(MIN_FONT_SCALE, floor[0] if head else 1)
    r, c = max(
        ((r, c) for r in range(len(rows)) for c in range(count)),
        key=lambda rc: len(rows[rc[0]][rc[1]]),
    )
    return TableFit(
        problem=(
            f"table: {len(rows)} rows need ≈{sum(full[head:])} lines; this table holds "
            f"≈{holds} (≈{holds_floor} at the smallest size); longest cell "
            f"rows[{r + first}][{c}] {len(rows[r][c])} chars, its column holds "
            f"≈{columns[c]} chars per line",
            f"shorten the longest cells or keep about {holds_floor} one-line rows; "
            "split the table over two slides",
        )
    )


def _clear_cell(tc) -> None:
    txBody = tc.find(qn("a:txBody"))
    if txBody is not None:
        checks.clear_paragraphs(txBody)


def _fresh_row_ids(tbl, rows: Sequence[Any]) -> None:
    """A cloned row keeps its source's ``a16:rowId``; PowerPoint repairs a
    table holding two rows with one id."""
    ids = [node for node in tbl.iter(f"{_A16}rowId")]
    taken = {int(node.get("val", "0")) for node in ids}
    for row in rows:
        for node in row.iter(f"{_A16}rowId"):
            new = (max(taken) + 1) % 2**32
            taken.add(new)
            node.set("val", str(new))


def fill_table(frame, zone: Zone, fit: TableFit) -> List[str]:
    """Write ``fit`` into a duplicated table frame; returns the template cell
    text the check found left over (and cleared)."""
    tbl = frame._element.find(f".//{qn('a:tbl')}")
    grid = tbl.find(qn("a:tblGrid"))
    originals = {_cell_text(tc) for tc in tbl.iter(qn("a:tc"))} - {""}
    columns = grid.findall(qn("a:gridCol"))
    for column in columns[len(fit.widths) :]:
        grid.remove(column)
    for column, width in zip(columns, fit.widths):
        column.set("w", str(width))
    trs = tbl.findall(qn("a:tr"))
    for tr in trs:
        for tc in tr.findall(qn("a:tc"))[len(fit.widths) :]:
            tr.remove(tc)
    head = trs[:1] if zone.table is not None and zone.table.header else []
    body = trs[len(head) :]
    for tr in body[len(fit.rows) :]:
        tbl.remove(tr)
    added = []
    for _ in range(len(fit.rows) - len(body)):
        clone = copy.deepcopy(body[-1])
        for tc in clone.findall(qn("a:tc")):
            _clear_cell(tc)
        tbl.findall(qn("a:tr"))[-1].addnext(clone)
        added.append(clone)
    _fresh_row_ids(tbl, added)
    values = ([fit.header] if head and fit.header is not None else []) + fit.rows
    for tr, row, height in zip(tbl.findall(qn("a:tr")), values, fit.heights):
        tr.set("h", str(round(height * 12700)))
        for tc, text in zip(tr.findall(qn("a:tc")), row):
            _clear_cell(tc)
            if text:
                _write_paragraphs(
                    tc.find(qn("a:txBody")), [_Line(t) for t in text.splitlines()]
                )
    if fit.scale is not None and zone.font_pt:
        _scale_runs(tbl, fit.scale, zone.font_pt)
    frame.height = Emu(round(sum(fit.heights) * 12700))
    wanted = {t for row in values for t in row}
    left = []
    for tc in tbl.iter(qn("a:tc")):
        text = _cell_text(tc)
        if text in originals and text not in wanted:
            left.append(text)
            _clear_cell(tc)
    return left


# --- zone mapping -----------------------------------------------------------------------


def _reading_order(zones: Sequence[Zone]) -> List[Zone]:
    return sorted(zones, key=lambda z: (z.box.y // 150, z.box.x))


def _column_order(zones: Sequence[Zone]) -> List[Zone]:
    return sorted(zones, key=lambda z: (z.box.x // 150, z.box.y))


class _ZoneMap:
    def __init__(self, pattern: Pattern) -> None:
        titled = {z.key for z in chart_headings(pattern).values()}
        zones = [z for z in pattern.zones if z.role != "keep" and z.key not in titled]
        self.titles = [z for z in zones if z.role == "title"]
        self.subtitles = [z for z in zones if z.role == "subtitle"]
        self.headings = _column_order([z for z in zones if z.role == "heading"])
        self.bodies = _column_order([z for z in zones if z.role == "body"])
        self.captions = _column_order([z for z in zones if z.role == "caption"])
        self.numbers = _column_order([z for z in zones if z.role == "number"])
        self.labels = _column_order([z for z in zones if z.role == "label"])
        self.images = _reading_order([z for z in zones if z.role == "image"])
        self.all_text = [z for z in zones if z.role != "image"]

    def body_like(self) -> List[Zone]:
        return self.bodies or self.captions or self.subtitles

    def stacked(self) -> List[Zone]:
        """Every non-title text line in reading order (contact cards)."""
        return sorted(
            (z for z in self.all_text if z.role not in ("title", "number")),
            key=lambda z: (z.box.y, z.box.x),
        )


def _assign_text(
    pattern: Pattern, plan: SlidePlan, lost: Optional[List[str]] = None
) -> Dict[str, List[_Line]]:
    """Decide which lines go into which text zone (by zone key). Plan content
    no zone receives is named in ``lost``, so the agent hears about it."""
    out = _place_text(pattern, plan, lost)
    headings = chart_headings(pattern)
    for zone, (at, chart) in zip(_chart_frames(pattern), shown_charts(plan)):
        box = headings.get(zone.key)
        if box is not None and chart.title:
            out[box.key] = _lines_from_bullets([chart.title], f"{at}.title")
    return out


def _place_text(
    pattern: Pattern, plan: SlidePlan, lost: Optional[List[str]] = None
) -> Dict[str, List[_Line]]:
    zmap = _ZoneMap(pattern)
    out: Dict[str, List[_Line]] = {}
    lost = lost if lost is not None else []

    def put(zone: Optional[Zone], lines: List[_Line], what: str) -> None:
        if not lines:
            return
        if zone is None or zone.key in out:
            lost.append(what)
            return
        out[zone.key] = lines

    title = _lines_from_bullets([plan.title], "title") if plan.title else []
    subtitle = _lines_from_bullets([plan.subtitle], "subtitle") if plan.subtitle else []
    bullets = _lines_from_bullets(plan.bullets, "bullets[{}]")

    if zmap.titles:
        put(zmap.titles[0], title, "title")
    elif zmap.headings and title and not plan.columns:
        put(zmap.headings[0], title, "title")
    else:
        put(None, title, "title")

    if pattern.kind in ("cover", "section", "quote"):
        # Covers put their subtitle in whatever sits under the title.
        target = zmap.subtitles or zmap.body_like()
        put(
            target[0] if target else None,
            subtitle or bullets,
            "subtitle" if subtitle else "bullets",
        )
        if subtitle and bullets:
            lost.append("bullets")
        if plan.columns:
            lost.append("columns")
        return out

    if zmap.subtitles:
        put(zmap.subtitles[0], subtitle, "subtitle")
    elif subtitle and not bullets and zmap.bodies and not plan.columns:
        bullets = subtitle
    elif subtitle and zmap.bodies and not plan.columns:
        # No subtitle zone: it leads the body, unbulleted, like a column heading.
        lead = subtitle[0]
        bullets = [_Line(lead.text, heading=True, source=lead.source)] + bullets
    elif subtitle:
        lost.append("subtitle")

    if plan.columns:
        for index, column in enumerate(plan.columns):
            body = zmap.bodies[index] if index < len(zmap.bodies) else None
            heading = zmap.headings[index] if index < len(zmap.headings) else None
            lines = _lines_from_bullets(
                column.bullets, f"columns[{index}].bullets[{{}}]"
            )
            if column.heading and heading is not None:
                put(
                    heading,
                    _lines_from_bullets([column.heading], f"columns[{index}].heading"),
                    f"column {index + 1}",
                )
            elif column.heading:
                lines = [
                    _Line(
                        f"**{column.heading}**",
                        heading=True,
                        source=f"columns[{index}].heading",
                    )
                ] + lines
            put(body, lines, f"column {index + 1}")
        if bullets:
            lost.append("bullets")
        return out

    if pattern.kind == "closing" and bullets:
        zones = [z for z in zmap.stacked() if z.key not in out]
        for zone, line in zip(zones, bullets):
            put(zone, [line], "bullets")
        if len(bullets) > len(zones):
            lost.append(f"{len(bullets) - len(zones)} bullet(s)")
        return out
    bodies = zmap.body_like()
    if not bullets:
        return out
    if not bodies:
        lost.append("bullets")
        return out
    if len(bodies) > 1 and pattern.kind in ("two_columns", "comparison"):
        half = math.ceil(len(bullets) / 2)
        put(bodies[0], bullets[:half], "bullets")
        put(bodies[1], bullets[half:], "bullets")
    else:
        put(bodies[0], bullets, "bullets")
    return out


def _lead_lines(item: PlanItem, at: str) -> List[_Line]:
    """An item's heading in bold, unbulleted, above its text."""
    lead = _Line(f"**{item.heading}**", heading=True, source=at + "heading")
    return [lead] + _lines_from_bullets([item.text or ""], at + "text")


def _heading_lead(
    zones: Sequence[Zone], item: PlanItem, at: str, size: Tuple[int, int]
) -> Optional[Zone]:
    """The single text zone an item's heading leads when no zone shows it (as a
    subtitle leads a body); None when another zone shows it or it does not fit."""
    roles = {z.role for z in zones}
    if not item.heading or roles & {"heading", "title", "subtitle"}:
        return None
    if "label" in roles and not item.label:
        return None  # the label zone shows the heading
    bodies = [z for z in zones if z.role in ("body", "caption")]
    if len(bodies) != 1:
        return None
    zone = bodies[0]
    if not zone.grows and _fit(zone, _lead_lines(item, at), *size)[2]:
        return None
    return zone


def _unused_item_fields(
    groups: Sequence[RepeatGroup], plan: SlidePlan, size: Tuple[int, int]
) -> List[str]:
    """Item fields the plan fills that none of the slide's groups has a zone for
    (every group is fed the same items)."""
    roles = {role for group in groups for role in group.item_zones}
    lost: Dict[str, None] = {}
    numbers = [i.number.strip().rstrip(".") for i in plan.items if i.number]
    # Step numbers 1, 2, 3… (or 6, 7, 8… on a list's continuation slide) say what
    # the item order already shows; years are never step numbers.
    start = int(numbers[0]) if numbers and numbers[0].isdigit() else 1
    start = start if start < 100 else 1
    counting = len(numbers) == len(plan.items) and numbers == [
        f"{n:0{len(numbers[0])}d}" for n in range(start, start + len(numbers))
    ]
    for index, item in enumerate(
        plan.items[: max((g.count for g in groups), default=0)]
    ):
        heading_shown = (
            bool(roles & {"heading", "title", "subtitle"})
            or ("label" in roles and not item.label)
            or any(
                _heading_lead(g.items[index], item, "", size) is not None
                for g in groups
                if index < g.count
            )
        )
        checks = (
            ("number", item.number, "number" in roles or counting),
            ("heading", item.heading, heading_shown),
            ("text", item.text, bool(roles & {"body", "caption"})),
            ("label", item.label, "label" in roles),
            ("image", item.image, "image" in roles),
        )
        for name, value, shown in checks:
            if value and not shown:
                lost[name] = None
    return [f"items' {name}" for name in lost]


def _assign_group(
    group: RepeatGroup, plan: SlidePlan, size: Tuple[int, int], start: int = 0
) -> Tuple[Dict[str, List[_Line]], Dict[str, Optional[str]], int]:
    """Lines and image ids per item zone key; returns ``(text, images, used items)``.
    ``start`` is the entries before a list's continuation: its count goes on."""
    text: Dict[str, List[_Line]] = {}
    images: Dict[str, Optional[str]] = {}
    used = min(len(plan.items), group.count)
    for index in range(used):
        item = plan.items[index]
        zones = group.items[index]
        has_heading = any(z.role == "heading" for z in zones)
        at = f"items[{index}]."
        lead = _heading_lead(zones, item, at, size)
        for zone in zones:
            if zone is lead:
                text[zone.key] = _lead_lines(item, at)
            elif zone.role == "number":
                value = item.number or _auto_number(zone.sample, start + index)
                if item.trend and item.number:
                    # A thin space after the arrow: it reads as part of the figure.
                    value = f"{TREND_GLYPHS[item.trend]}\u2009{value}"
                text[zone.key] = (
                    _lines_from_bullets([value], at + "number") if value else []
                )
            elif zone.role == "heading":
                text[zone.key] = _lines_from_bullets(
                    [item.heading or ""], at + "heading"
                )
            elif zone.role in ("body", "caption"):
                text[zone.key] = _lines_from_bullets([item.text or ""], at + "text")
            elif zone.role == "label":
                value = item.label or (item.heading if not has_heading else None) or ""
                field = "label" if item.label else "heading"
                text[zone.key] = _lines_from_bullets([value], at + field)
            elif zone.role == "image":
                images[zone.key] = item.image
            elif zone.role in ("title", "subtitle"):
                text[zone.key] = _lines_from_bullets(
                    [item.heading or ""], at + "heading"
                )
    return text, images, used


# --- charts -------------------------------------------------------------------------------


def _chart_frames(pattern: Pattern) -> List[Zone]:
    return _reading_order([z for z in pattern.zones if z.role == "chart"])


def chart_hosts(pattern: Pattern, text_keys: Collection[str]) -> List[Zone]:
    """Where the slide's charts go, in reading order: its chart zones, else one
    content placeholder able to hold a chart that receives no text."""
    frames = _chart_frames(pattern)
    if frames:
        return frames
    zones = _reading_order(pattern.zones)
    return [z for z in zones if z.chart_capable and z.key not in text_keys][:1]


# --- rendering ------------------------------------------------------------------------------


def _has_text_below(zone: Zone, siblings: Sequence[Zone]) -> bool:
    """Whether another text zone of the same item starts under this one: only
    then does a second line run into something."""
    bottom = zone.box.y + zone.box.h
    return any(
        other.key != zone.key
        and other.role not in ("image", "keep")
        and other.box.y >= zone.box.y + zone.box.h / 2
        and other.box.y <= bottom + zone.box.h
        and other.box.x < zone.box.x + zone.box.w
        and zone.box.x < other.box.x + other.box.w
        for other in siblings
    )


def _only_names_the_slide(plan: SlidePlan, pattern: Pattern) -> bool:
    """A fixed slide given only the title it already shows ("Merci" on "Merci",
    or any closing word on a closing slide) lost nothing."""
    title = plan.title or ""
    if not title or has_content(plan.model_copy(update={"title": None})):
        return False
    names = {lexicon.normalize(pattern.label), lexicon.normalize(pattern.name)}
    closing = lexicon.name_matches(
        pattern.label, lexicon.CLOSING_NAMES
    ) or lexicon.name_matches(pattern.name, lexicon.CLOSING_NAMES)
    return lexicon.normalize(title) in names or (
        closing and lexicon.name_matches(title, lexicon.CLOSING_NAMES)
    )


class _Renderer:
    def __init__(
        self,
        presentation,
        inventory: TemplateInventory,
        images: Mapping[str, PreparedImage],
    ) -> None:
        self.prs = presentation
        self.inventory = inventory
        self.images = images
        self.width = presentation.slide_width
        self.height = presentation.slide_height
        self.layouts = _all_layouts(presentation)
        self.theme = theme_fonts(presentation)
        self.slides_by_id = {slide.slide_id: slide for slide in presentation.slides}
        self.report = checks.DeckReport()
        self._originals: Dict[int, Dict[str, str]] = {}
        # An intent note's text → (its line in the report, the slides it concerns).
        self._intents: Dict[str, Tuple[int, List[int]]] = {}
        self.language = "en"

    def render(
        self,
        plan: DeckPlan,
        patterns: Sequence[Optional[Pattern]],
        numbers: Optional[Sequence[int]] = None,
        starts: Optional[Sequence[int]] = None,
    ) -> RenderResult:
        original_ids = list(self.slides_by_id)
        rendered = 0
        self.language = notes.plan_language(plan)
        # [source, rendered index of its last part, topics, number] of each agenda:
        # a deck announcing more topics than it has slides left is one the agent
        # meant to write and did not. The parts of a split agenda are one agenda.
        announced: List[List[int]] = []
        parts_of = Counter(numbers or [])
        seen: Counter = Counter()
        for position, slide_plan in enumerate(plan.slides, start=1):
            # Messages name the deck's slide; a split list's parts share a source.
            source = numbers[position - 1] if numbers else position
            if slide_plan.invalid_entry is not None:  # the tool's notes name it
                continue
            pattern = patterns[position - 1]
            if pattern is None:
                raise RenderError(
                    f"slide {position} has no template slide (internal). Call again with "
                    "the same plan; if it persists, tell the user."
                )
            rendered += 1
            if pattern.kind == "agenda" and slide_plan.items:
                if announced and announced[-1][0] == source:
                    announced[-1][1] = rendered
                    announced[-1][2] += len(slide_plan.items)
                else:
                    announced.append(
                        [source, rendered, len(slide_plan.items), position]
                    )
            slide, pattern = self._create(pattern)
            seen[source] += 1
            self._record(
                slide_plan,
                pattern,
                position,
                seen[source],
                parts_of[source] or 1,
                rendered,
            )
            # Nothing to say on it (a "Thank you" slide): keep the template's text.
            filled = pattern.reusable and pattern.fillable and has_content(slide_plan)
            if filled:
                start = starts[position - 1] if starts else 0
                self._fill(slide, pattern, slide_plan, position, start)
                if pattern.source == "slide":
                    self._sweep(slide, pattern, position)
            elif has_content(slide_plan) and not _only_names_the_slide(
                slide_plan, pattern
            ):
                self.report.note("fixed_slide")
                self._tell(
                    f"Slide {position}: '{pattern.label}' is a fixed slide shown as is, so its "
                    "content was left out. If it matters, move it to another slide; otherwise "
                    "nothing to do."
                )
            # A copied notes slide keeps its formatting, never the template's
            # own notes (author guidance, `intent:` annotations).
            spoken = notes.speaker_notes(slide_plan, self.language)
            if spoken or slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None:
                    frame.text = spoken or ""
            # A slide shown as is never keeps its example chart's made-up data.
            checks.settle_charts(slide, keep=filled)
            planned = len(slide_charts(slide_plan)) if filled else 0
            delivered = checks.chart_count(slide)
            if delivered != planned:
                raise RenderError(
                    f"slide {position} shows {delivered} of its {planned} chart(s) after "
                    "rendering, an internal problem. Tell the user and suggest trying again."
                )
        if rendered == 0:
            raise RenderError(
                "the plan holds no usable slide. Write each entry of `slides` as an "
                "object with its content and call again."
            )
        for _, index, topics, position in announced:
            remaining = rendered - index
            # Only a wide gap: a deck may well cover five topics in three
            # slides, but not in none.
            if topics >= 3 and remaining * 2 < topics:
                self._tell(
                    f"Slide {position}: the agenda announces {topics} parts but only {remaining} "
                    "slide(s) follow it. Add the missing slides or remove those parts from the "
                    "agenda; otherwise nothing to do."
                )
        _remove_slides(self.prs, original_ids)
        buffer = io.BytesIO()
        self.prs.save(buffer)
        preview = buffer
        try:
            widened = _widen_no_wrap_text(self.prs)
        except (ValueError, TypeError):  # unexpected geometry: preview as delivered
            widened = 0
        if widened:
            preview = io.BytesIO()
            self.prs.save(preview)
        return RenderResult(
            pptx=buffer.getvalue(),
            slide_count=rendered,
            preview_pptx=preview.getvalue(),
            report=self.report,
        )

    def _record(
        self,
        slide_plan: SlidePlan,
        pattern: Pattern,
        number: int,
        part: int,
        parts: int,
        rendered: int,
    ) -> None:
        """Open the user-facing record of the slide about to be rendered."""
        shown = pattern.reusable and pattern.fillable and has_content(slide_plan)
        title = slide_plan.title if shown and slide_plan.title else ""
        title = (
            "".join(span.text for span in parse_inline_markdown(title)) if title else ""
        )
        self.report.built.append(
            BuiltSlide(
                number=number,
                part=part,
                parts=parts,
                title=title[:_SUMMARY_TITLE_CHARS],
                kind=pattern.kind,
                template=pattern.label,
            )
        )
        if parts > 1:
            self.report.note("list_continued")
        drawn = slide_charts(slide_plan) if shown else []
        # A chart in a content area is a substitution even when no kind was asked.
        asked = slide_plan.kind or ("chart" if drawn else None)
        # The first slide is the cover by rule: not a substitution.
        if (
            asked
            and asked != pattern.kind
            and not (rendered == 1 and pattern.kind == "cover")
        ):
            self.report.note("other_kind", kind=asked)
        for chart in drawn:
            if chart.asked_type:
                self.report.note("chart_type_changed", kind=chart.asked_type)
        if part == 1:
            if notes.is_illustrative(slide_plan):
                self.report.note("illustrative_data")
            for code in slide_plan.story_notes:
                self.report.note(code)  # type: ignore[arg-type]

    def _tell(self, note: str) -> None:
        """A finding the after-check report does not carry: sent to the agent."""
        self.report.notes.append(note)

    def _tell_intent(
        self,
        position: int,
        name: str,
        note: str,
        code: BuildNoteCode = "intent_not_shown",
    ) -> None:
        """An intent built without (or drawn otherwise): one agent line per
        text, naming every slide."""
        self.report.note(code, field=name)
        line, slides = self._intents.get(note) or (len(self.report.notes), [])
        if not slides:
            self.report.notes.append("")
            self._intents[note] = (line, slides)
        if position not in slides:
            slides.append(position)
        named = ", ".join(str(n) for n in slides)
        self.report.notes[line] = (
            f"Slide{'s' if len(slides) > 1 else ''} {named}: {note}"
        )

    def _create(self, pattern: Pattern) -> Tuple[Any, Pattern]:
        """The slide, and the pattern whose zones address its actual shapes."""
        if pattern.source == "layout":
            if pattern.source_index >= len(self.layouts):
                raise RenderError(
                    f"layout {pattern.source_index} no longer exists in the template. "
                    + _TEMPLATE_CHANGED
                )
            return self.prs.slides.add_slide(
                self.layouts[pattern.source_index]
            ), pattern
        slide_id = int(pattern.id.split(":", 1)[1])
        source = self.slides_by_id.get(slide_id)
        if source is None:
            raise RenderError(
                f"slide {slide_id} no longer exists in the template. "
                + _TEMPLATE_CHANGED
            )
        return _duplicate_slide(self.prs, source), pattern

    def _sweep(self, slide, pattern: Pattern, position: int) -> None:
        """Remove the example text the fill left on a duplicated slide."""
        slide_id = int(pattern.id.split(":", 1)[1])
        if slide_id not in self._originals:
            self._originals[slide_id] = checks.original_texts(
                self.slides_by_id[slide_id]
            )
        index = _shape_index(slide)
        zone_keys = [z.key for z in pattern.zones] + [
            z.key for g in pattern.groups for item in g.items for z in item
        ]
        handled = [
            cast(Any, index[key].shape)._element for key in zone_keys if key in index
        ]
        removed = checks.remove_leftovers(
            slide,
            self._originals[slide_id],
            handled,
            set(self.inventory.shared_texts),
        )
        if removed:
            self.report.slide(position).removed.extend(removed)

    def _fill(
        self, slide, pattern: Pattern, plan: SlidePlan, position: int, start: int = 0
    ) -> None:
        index = _shape_index(slide)
        report = self.report.slide(position)
        lost: List[str] = []
        text = _assign_text(pattern, plan, lost)
        image_ids: Dict[str, Optional[str]] = {}
        zmap = _ZoneMap(pattern)
        # Both fields are offered, so a model names the hero picture in `image`
        # and the rest in `images`; taking one or the other dropped the hero.
        plan_images = list(
            dict.fromkeys(([plan.image] if plan.image else []) + list(plan.images))
        )
        for zone, image_id in zip(zmap.images, plan_images + [None] * len(zmap.images)):
            image_ids[zone.key] = image_id
        if len(plan_images) > len(zmap.images):
            lost.append(f"{len(plan_images) - len(zmap.images)} image(s)")
        if plan.items and not pattern.groups:
            lost.append("items")
        tables = [z for z in pattern.zones if z.role == "table"]
        if plan.table is not None and plan.table.has_content() and not tables:
            lost.append("table")
        lost.extend(
            _unused_item_fields(pattern.groups, plan, (self.width, self.height))
        )
        hosted = self._fill_charts(slide, pattern, plan, index, text, position, lost)
        self._tell_trends(pattern, plan, position)

        for group in pattern.groups:
            group_text, group_images, used = _assign_group(
                group, plan, (self.width, self.height), start
            )
            text.update(group_text)
            image_ids.update(group_images)
            if len(plan.items) > group.count:
                report.lost.append(f"{len(plan.items) - group.count} item(s)")
                self.report.note("items_left_out", len(plan.items) - group.count)
            items = _item_shapes(group, index)
            kept, dropped = items[:used], [i for i in items[used:] if i]
            if dropped and kept:
                self._remove_decoration(dropped, kept, index, pattern)
            elif dropped:
                # Nothing was kept, so nothing anchors the decoration: sweeping
                # by overlap here takes the slide's own background with it.
                self.report.note("left_empty")
                self._tell(
                    f"Slide {position}: '{pattern.label}' shows its content as items and none "
                    f"were given, so it stays empty. Give that content as `items` (up to "
                    f"{group.count})."
                )
            for placed in (p for item in dropped for p in item):
                _remove(placed)
            if kept and dropped:
                self._recentre_items(kept, dropped, index, pattern)
            self._even_out_item_widths(group, index)
        report.lost.extend(lost)
        if lost:
            self.report.note("content_left_out")
        if pattern.groups:
            # Widths and removals changed the slide: re-read absolute geometry.
            index = _shape_index(slide)

        # A row of loose shapes has no container: its zones reach the slide edge.
        containers = {
            z.key: index.get(key) if _is_container(key, item) else None
            for g in pattern.groups
            for key, item in zip(g.item_keys, g.items)
            for z in item
        }
        siblings: Dict[str, Sequence[Zone]] = {
            z.key: pattern.zones for z in pattern.zones
        }
        siblings.update(
            {z.key: item for g in pattern.groups for item in g.items for z in item}
        )
        layout_body = _layout_body_props(slide)
        scaled: Dict[str, List[Tuple[Any, Zone, float]]] = {}
        for zone in self._fillable_text_zones(pattern):
            if zone.key in hosted:
                continue
            placed = index.get(zone.key)
            shape = cast(Any, placed.shape) if placed is not None else None
            if (
                placed is None
                or shape is None
                or not getattr(shape, "has_text_frame", False)
            ):
                continue
            lines = text.get(zone.key, [])
            if not lines:
                # An empty placeholder shows "Click to add text" once the deck
                # is edited: drop it (the layout can always put it back).
                if cast(Any, shape).is_placeholder:
                    _remove(placed)
                else:
                    _clear_text(shape)
                continue
            if zone.role == "number":
                # "1 200", "62 %", "4,2 M€" never break across lines.
                lines = [
                    _Line(line.text.replace(" ", " "), line.level) for line in lines
                ]
            font_name, bold = _font_name_of(
                shape, self.theme, is_title=zone.role == "title"
            )
            one_line_room: Optional[int] = None
            if _grows_with_text(shape, layout_body.get(zone.key)):
                scale, one_line_room = self._no_wrap_scale(
                    zone,
                    lines,
                    placed,
                    containers.get(zone.key),
                    _has_text_below(zone, siblings.get(zone.key, ())),
                    font_name,
                    bold,
                )
                truncated = False
            else:
                lines, scale, truncated = _fit(
                    zone, lines, self.width, self.height, font_name, bold
                )
            _write_lines(shape, lines)
            if zone.role == "number":
                _shrink_trend_glyph(shape, zone.font_pt)
            # Registered even when it needs no shrinking: a sibling that does is
            # what decides the size of the set.
            scaled.setdefault(_sibling_key(zone, pattern), []).append(
                (shape, zone, 1.0 if scale is None else scale)
            )
            if one_line_room is not None:
                self.report.note("text_overflows")
                self._tell(
                    f"Slide {position}: the {zone.role} '{lines[0].text}' is too long for one line "
                    f"and runs into the text below. Shorten it to under {one_line_room} characters."
                )
            if truncated:
                report.truncated.append(zone.role)
                self.report.note("text_shortened")

        # Zones that play the same part on one slide are read as a set: two
        # column headings at 24 pt and 13 pt look like a mistake, even when each
        # was right on its own. They all take the smallest of their scales.
        for group_scales in scaled.values():
            smallest = min(scale for _, _, scale in group_scales)
            if smallest >= 1.0:
                continue
            for shape, zone, _ in group_scales:
                _set_font_scale(shape, smallest, zone.font_pt)

        for zone in tables:
            placed = index.get(zone.key)
            if placed is None or plan.table is None:
                continue
            self._fill_table(placed, zone, plan.table, position)

        for zone in [z for z in pattern.zones if z.role == "image"] + [
            z
            for g in pattern.groups
            for item in g.items
            for z in item
            if z.role == "image"
        ]:
            placed = index.get(zone.key)
            if placed is None:
                continue
            image_id = image_ids.get(zone.key)
            if not image_id:
                if cast(Any, placed.shape).is_placeholder:
                    _remove(placed)
                continue
            prepared = self.images.get(image_id)
            if prepared is None:
                # A bad image reference costs the picture, never the deck.
                report.lost.append(f"image '{image_id}'")
                self.report.note("image_unusable", 1)
                if cast(Any, placed.shape).is_placeholder:
                    _remove(placed)
                continue
            self._place_image(slide, placed, prepared)

    def _fill_charts(
        self,
        slide,
        pattern: Pattern,
        plan: SlidePlan,
        index: Dict[str, PlacedShape],
        text: Dict[str, List[_Line]],
        position: int,
        lost: List[str],
    ) -> set[str]:
        """Draw each chart of the plan in its host, in reading order; an
        unfilled chart zone goes. Returns the hosts' keys."""
        shown = shown_charts(plan)
        hosts = chart_hosts(pattern, text) if shown else []
        if len(shown) > len(hosts):
            lost.append("chart")
        filled = hosts[: len(shown)]
        for zone in (z for z in pattern.zones if z.role == "chart" and z not in filled):
            placed = index.get(zone.key)
            if placed is not None:
                _remove(placed)
        headings = chart_headings(pattern)
        for host, (at, chart) in zip(filled, shown):
            placed = index.get(host.key)
            if placed is None:
                continue
            box = headings.get(host.key)
            if notes.is_illustrative(plan):
                title = notes.caption(chart, self.language)
                chart = chart.model_copy(update={"title": title})
                if box is not None:
                    text[box.key] = _lines_from_bullets([title], f"{at}.title")
            if box is not None:
                # Its heading box above shows the title, never twice.
                chart = chart.model_copy(update={"title": None})
            size_line = charts.axis_title(chart.size_title, chart.size_unit)
            if chart.type == "bubble" and size_line:
                title = notes.with_size_line(chart, self.language, size_line)
                chart = chart.model_copy(update={"title": title})
            self._draw(slide, placed, host, chart)
            self._tell_chart(chart, position, at)
        return {host.key for host in filled}

    def _draw(self, slide, placed: PlacedShape, host: Zone, chart: PlanChart) -> None:
        """The designer's own chart when it already shows the type asked, else a
        fresh one in its place."""
        example = (
            charts.example_frame(placed, chart)
            if host.role == "chart" and host.chart_type == chart.type
            else None
        )
        if example is not None:
            charts.refill(example, host, chart)
        else:
            charts.draw(slide, placed, host, chart)

    def _tell_chart(self, chart: PlanChart, position: int, at: str) -> None:
        """What the chart could not show as sent, naming it on a slide of several."""
        where = f"Slide {position}" + ("" if at == "chart" else f", {at}")
        if chart.asked_type:
            self._tell(
                f"{where}: chart type '{chart.asked_type}' is not offered; built as "
                f"'column'. Nothing to do; to change it, edit {at}.type to one of "
                f"{', '.join(charts.XL_TYPES)}."
            )
        self._tell_labels(chart, where, at)
        self._tell_axes(chart, where, at)
        for name, note in charts.intent_notes(chart):
            self._tell_intent(position, name, note)
        if chart.style_fields:
            self._tell(
                f"{where}: the chart's {', '.join(chart.style_fields)} "
                "were ignored: the template sets a chart's colours, fonts and place. "
                "Nothing to do; never send them."
            )

    def _tell_trends(self, pattern: Pattern, plan: SlidePlan, position: int) -> None:
        """Trends the slide cannot draw: no figure zone, no figure, or a word
        naming no direction."""
        shown = plan.items[: max((g.count for g in pattern.groups), default=0)]
        figures = any("number" in g.item_zones for g in pattern.groups)
        for index, item in enumerate(shown):
            at = f"items[{index}].trend"
            if item.asked_trend:
                self._tell_intent(
                    position,
                    "trend",
                    f"{at} '{item.asked_trend}' is not a direction; built without "
                    "its arrow. Use 'up', 'down' or 'flat', or drop it.",
                )
            elif item.trend and figures and not item.number:
                self._tell_intent(
                    position,
                    "trend",
                    f"{at} needs the item's number: the arrow goes before a figure. "
                    "Give items[].number, or drop the trend.",
                )
        if not figures and any(item.trend for item in shown):
            self._tell_intent(
                position,
                "trend",
                "items' trend was not shown: this slide shows no figure for its items. "
                "Nothing to do; drop items[].trend, or use a slide showing figures.",
            )

    def _tell_axes(self, chart: PlanChart, where: str, at: str) -> None:
        """Axis fields the chart type does not draw: said, never dropped quietly."""
        drawn = _AXIS_FIELDS if chart.type == "bubble" else ()
        if chart.type == "scatter":
            drawn = _AXIS_FIELDS[:3]
        sent = [
            f"{at}.{n}" for n in _AXIS_FIELDS if getattr(chart, n) and n not in drawn
        ]
        if sent:
            self._tell(
                f"{where}: {', '.join(sent)} not drawn on a '{chart.type}' "
                "chart (axis titles: scatter and bubble only; size: bubble only). "
                "Nothing to do; drop them."
            )

    def _tell_labels(self, chart: PlanChart, where: str, at: str) -> None:
        """A label setting the chart cannot show: said, never dropped quietly."""
        unit = (chart.unit or "").strip()
        if chart.type in charts.ROUND:
            if chart.show in (None, "share") and unit and unit != "%":
                self.report.note("chart_shows_shares")
                self._tell(
                    f"{where}: the {chart.type}'s labels show percentages, so "
                    f"unit '{unit}' is not shown. To show the numbers with their unit, "
                    f"set {at}.show to 'value' ('both' for the number and its percentage)."
                )
            return
        show = chart.show or chart.asked_show
        if show:
            self._tell(
                f"{where}: {at}.show '{show}' was not applied: it sets the slice "
                f"labels of a pie or doughnut only; a '{chart.type}' chart shows its values. "
                f"Nothing to do; drop {at}.show, or set {at}.type to 'pie' or 'doughnut'."
            )

    def _fill_table(
        self,
        placed: PlacedShape,
        zone: Zone,
        table: PlanTable,
        position: int,
    ) -> None:
        report = self.report.slide(position)
        fit = fit_table(zone, table)
        if fit.problem is not None:
            # Its own line: "not placed: table" would name the wrong remedy.
            self.report.note("table_left_out")
            self._tell(
                f"Slide {position}: the table was left out "
                f"({fit.problem[0].removeprefix('table: ')}). To show it, "
                f"{fit.problem[1]}."
            )
            return
        frame = cast(Any, placed.shape)
        removed = fill_table(frame, zone, fit)
        if removed:
            report.removed.extend(removed)
        tbl = frame._element.find(f".//{qn('a:tbl')}")
        head = zone.table is not None and zone.table.header
        for name, code, note in table_intents.emphasise(
            self.prs, tbl, table, bool(head), fit.shift
        ):
            self._tell_intent(position, name, note, code)
        if fit.scale is not None:
            report.shrunk.append(f"table text at {round(fit.scale * 100)} % to fit")
            self.report.note("table_text_smaller")
        if placed.top + frame.height > self.height:
            report.lost.append("table rows below the slide")
            self.report.note("content_left_out")

    @staticmethod
    def _loose_decoration(
        index: Dict[str, PlacedShape], pattern: Pattern
    ) -> List[PlacedShape]:
        """Top-level shapes that carry no content: the icons, card frames and
        connectors a designer draws beside a group rather than inside it."""
        zone_keys = {z.key for z in pattern.zones}
        zone_keys.update(
            z.key for g in pattern.groups for item in g.items for z in item
        )
        # An item's own container is content, not decoration drawn beside it.
        zone_keys.update(key for g in pattern.groups for key in g.item_keys)
        loose = []
        for key, placed in index.items():
            if len(placed.path) != 1 or key in zone_keys or key.startswith("ph:"):
                continue
            shape = cast(Any, placed.shape)
            if (
                getattr(shape, "has_text_frame", False)
                and shape.text_frame.text.strip()
            ):
                continue
            loose.append(placed)
        return loose

    def _remove_decoration(
        self,
        dropped: List[List[PlacedShape]],
        kept: List[List[PlacedShape]],
        index: Dict[str, PlacedShape],
        pattern: Pattern,
    ) -> None:
        """Drop the loose decoration drawn with the unused items.

        A shape goes when it sits on one of them, or anywhere inside the area
        they occupied — a connector between two dropped slots overlaps neither
        by much and used to be left dangling over an empty half-slide. Anything
        that also serves a kept item stays, and so does a full-slide background:
        it reaches beyond the dropped area and touches what was kept.
        """
        dropped_boxes = [_union_box(item) for item in dropped]
        kept_boxes = [_union_box(item) for item in kept if item]
        removed = [p for item in dropped for p in item]
        area = (
            min(b[0] for b in dropped_boxes),
            min(b[1] for b in dropped_boxes),
            max(b[2] for b in dropped_boxes),
            max(b[3] for b in dropped_boxes),
        )
        for other in self._loose_decoration(index, pattern):
            if other in removed:
                continue
            box = _box_of(other)
            on_dropped = any(_covered(box, b) >= 0.5 for b in dropped_boxes)
            if not on_dropped and _covered(box, area) < 0.9:
                continue
            if any(_covered(box, b) >= 0.2 for b in kept_boxes):
                continue
            _remove(other)

    def _recentre_items(
        self,
        kept: List[List[PlacedShape]],
        dropped: List[List[PlacedShape]],
        index: Dict[str, PlacedShape],
        pattern: Pattern,
    ) -> None:
        """Slide the used items back to the middle of the room the group had.

        Four cards of a six-card row left at the designer's coordinates hug the
        left edge with a third of the slide empty beside them. A translation
        only: the corporate spacing and sizing are untouched.
        """
        used = [_union_box(item) for item in kept if item]
        whole = used + [_union_box(item) for item in dropped]
        shift_x = ((min(b[0] for b in whole) + max(b[2] for b in whole)) // 2) - (
            (min(b[0] for b in used) + max(b[2] for b in used)) // 2
        )
        shift_y = ((min(b[1] for b in whole) + max(b[3] for b in whole)) // 2) - (
            (min(b[1] for b in used) + max(b[3] for b in used)) // 2
        )
        if abs(shift_x) < _RECENTRE_MIN_EMU and abs(shift_y) < _RECENTRE_MIN_EMU:
            return
        # The icons and frames drawn over a kept item belong to it and move with
        # it; anything else on the slide stays where the designer put it.
        movers = [p for item in kept for p in item]
        for other in self._loose_decoration(index, pattern):
            box = _box_of(other)
            if any(_covered(box, b) >= 0.5 for b in used):
                movers.append(other)
        for placed in movers:
            shape = cast(Any, placed.shape)
            try:
                _own_geometry(shape)
                shape.left = Emu(shape.left + shift_x)
                shape.top = Emu(shape.top + shift_y)
            except (TypeError, ValueError):
                continue

    def _even_out_item_widths(
        self, group: RepeatGroup, index: Dict[str, PlacedShape]
    ) -> None:
        """Widen a text box narrower than its siblings of the same role: designers
        size sample boxes to their sample text, and the narrow one would wrap."""
        for position in range(len(group.item_zones)):
            zones = [item[position] for item in group.items if position < len(item)]
            if not zones or zones[0].role == "image":
                continue
            widest = max(z.box.w for z in zones)
            for zone in zones:
                if zone.box.w < widest * 0.8:
                    placed = index.get(zone.key)
                    if placed is not None and placed.width > 0:
                        # `placed.width` is absolute, `.width` is in the group's
                        # own units: scale the target through that ratio.
                        target = widest / 1000 * self.width
                        shape = cast(Any, placed.shape)
                        _own_geometry(shape)
                        shape.width = Emu(round(target * shape.width / placed.width))

    def _no_wrap_scale(
        self,
        zone: Zone,
        lines: List[_Line],
        placed: PlacedShape,
        container: Optional[PlacedShape],
        text_below: bool = False,
        font_name: Optional[str] = None,
        bold: bool = False,
    ) -> Tuple[Optional[float], Optional[int]]:
        """Keep a box that grows with its text (a KPI figure, its caption)
        inside its card, identically in PowerPoint and in the PDF preview.

        The room runs to the card's right edge (or the slide's). A box too
        narrow for its text is widened to that room. A figure stays on one line
        and shrinks as needed; a one-line heading or label with text right below
        it shrinks down to the overflow floor, then wraps and returns how many
        characters would fit, so the agent shortens it; other text wraps between words and only shrinks
        for a word longer than the room. Returns ``(font scale, one-line room)``.
        """
        if not zone.font_pt or placed.width <= 0:
            return None, None
        right = (
            container.left + container.width
            if container is not None and container.width > 0
            else int(self.width * 0.95)
        )
        room = int((right - placed.left) * 0.95)
        if room <= 0:
            return None, None
        shape_width = getattr(cast(Any, placed.shape), "width", 0) or placed.width
        # `placed.width` is absolute, the insets are in the shape's own units:
        # inside a scaled group the two are not the same ruler.
        insets = _text_insets(cast(Any, placed.shape)) * placed.width / shape_width
        inner = placed.width - insets
        emu_per_em = zone.font_pt / 72 * 914400
        texts = [line.text for line in lines]
        needed = (
            max((text_metrics.width_em(t, font_name, bold) for t in texts), default=0.0)
            * emu_per_em
        )
        longest_word = (
            max(
                (
                    text_metrics.width_em(w, font_name, bold)
                    for t in texts
                    for w in t.split()
                ),
                default=0.0,
            )
            * emu_per_em
        )
        if needed * _NO_WRAP_MARGIN <= inner:
            return None, None
        shape = cast(Any, placed.shape)
        # The margin absorbs the difference between these metrics and the
        # renderer's own: without it a box sized to the exact text wraps anyway.
        target = min(needed * _NO_WRAP_MARGIN + insets, room)
        # Only ever widen. The room is capped below the container, so a box the
        # designer drew wider than that cap used to be cut back — the opposite
        # of what a text too wide for its box needs.
        if target > placed.width and shape.width:
            _own_geometry(shape)
            shape.width = Emu(round(target * shape.width / placed.width))
        fitted = needed * _NO_WRAP_MARGIN
        if zone.role == "number":
            return (round(max(0.4, room / fitted), 2) if fitted > room else None), None
        if fitted <= room:
            return None, None
        if (
            zone.role in ("heading", "label")
            and text_below
            and (zone.max_lines or 1) <= 1
            and len(texts) == 1
        ):
            if room / fitted >= MIN_FONT_SCALE:
                return round(room / fitted, 2), None
            shrunk = MIN_FONT_SCALE
            per_char = (
                text_metrics.char_em(font_name, bold) or text_metrics.FALLBACK_CHAR_EM
            ) * emu_per_em
            fits = int(room / (per_char * shrunk * _NO_WRAP_MARGIN))
        else:
            shrunk, fits = None, None
        bodyPr = shape.text_frame._txBody.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("wrap", "square")
        word_scale = room / (longest_word * _NO_WRAP_MARGIN) if longest_word else 1.0
        if word_scale < (shrunk or 1.0):
            # One unbreakable token (an e-mail, a URL) must not take its whole
            # zone down to 7 pt: below the floor it overflows its measured box,
            # still inside the card, and the text around it stays readable.
            floor = 0.4 if zone.role == "number" else MIN_FONT_SCALE
            shrunk = round(max(floor, word_scale), 2)
        return shrunk, fits

    @staticmethod
    def _fillable_text_zones(pattern: Pattern) -> List[Zone]:
        zones = [
            z
            for z in pattern.zones
            if z.role not in ("image", "keep", "table", "chart")
        ]
        for group in pattern.groups:
            for item in group.items:
                zones.extend(z for z in item if z.role not in ("image", "keep"))
        return zones

    def _place_image(self, slide, placed: PlacedShape, prepared: PreparedImage) -> None:
        image_bytes, img_w, img_h = prepared
        shape = cast(Any, placed.shape)
        if shape.is_placeholder and hasattr(shape, "insert_picture"):
            shape.insert_picture(io.BytesIO(image_bytes))
            return
        left, top, width, height = fit_inside_box(
            img_w, img_h, placed.left, placed.top, placed.width, placed.height
        )
        picture = slide.shapes.add_picture(
            io.BytesIO(image_bytes), Emu(left), Emu(top), Emu(width), Emu(height)
        )
        # `add_picture` appends, which puts the new picture in front of
        # everything: a replaced cover photo then hides the title drawn over it.
        # Put it back where the one it replaces sat in the drawing order.
        element = cast(Any, placed.shape)._element
        anchor = element
        while anchor.getparent() is not None and anchor.getparent().tag != qn(
            "p:spTree"
        ):
            anchor = anchor.getparent()
        tree = anchor.getparent()
        if tree is not None:
            new_element = picture._element
            new_element.getparent().remove(new_element)
            tree.insert(list(tree).index(anchor), new_element)
        _remove(placed)


def render_deck(
    template_bytes: bytes,
    inventory: TemplateInventory,
    plan: DeckPlan,
    patterns: Sequence[Optional[Pattern]],
    images: Optional[Mapping[str, PreparedImage]] = None,
    numbers: Optional[Sequence[int]] = None,
    starts: Optional[Sequence[int]] = None,
) -> RenderResult:
    """Render ``plan`` on the template, slide ``i`` on ``patterns[i]`` as
    chosen by :func:`assign.assign` (CPU-bound; call off the event loop).
    ``numbers`` gives each slide's number in the sent plan: a split list's parts
    share it. Messages name the slide's own position in the deck. ``starts``
    (``Assignment.starts``) numbers a list's continuation on from its first part.

    Raises :class:`RenderError` for a pattern the template no longer holds;
    anything else propagates.
    """
    presentation = open_presentation(template_bytes)
    try:
        return _Renderer(presentation, inventory, images or {}).render(
            plan, patterns, numbers, starts
        )
    except ValueError as exc:
        raise RenderError(
            f"the template could not be filled ({exc}). Call again with a simpler plan; "
            "if it persists, tell the user."
        ) from exc
