# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Template analysis: from a plain ``.pptx`` to a :class:`TemplateInventory`.

Deterministic and offline. One pass over each layout and slide collects a flat
list of shape records (kind, geometry, text, font size, placeholder type);
every rule below works on those records, never on the XML again.

On an example slide every text-bearing shape is a slot, except text repeated
on several slides and date/footer/slide-number fields; roles come from
placeholder types, repeated groups, geometry and font rank. The slide kind is
a heuristic the vision pass refines. Optional annotations
(shape names, ``kind:`` notes) override the corresponding heuristic.
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import (
    TYPE_CHECKING,
    Any,
    Dict,
    List,
    Mapping,
    Optional,
    Sequence,
    Tuple,
    Union,
    cast,
)

from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from fred_capability_ppt_filler import text_metrics
from fred_capability_ppt_filler.builder import charts, lexicon
from fred_capability_ppt_filler.builder.models import (
    SLIDE_KINDS,
    AnalysisError,
    AnalyzeResult,
    Box,
    Pattern,
    RepeatGroup,
    SlideKind,
    TableSpec,
    TemplateInventory,
    Zone,
    ZoneRole,
)
from fred_capability_ppt_filler.builder.plan import (
    MAX_CHARTS_PER_SLIDE,
    chart_headings,
)
from fred_capability_ppt_filler.pptx_archive import open_presentation
from fred_capability_ppt_filler.traversal import KEY_PATTERN, walk_shapes

if TYPE_CHECKING:  # pragma: no cover - typing only
    from pptx.shapes.base import BaseShape

CODE_INVALID_UPLOAD = "invalid_upload"

# A relationship of these kinds does not block duplicating a slide (slide-to-
# slide links are dropped, not copied); any other part (chart, SmartArt, OLE)
# makes the slide uncopyable.
COPYABLE_REL_SUFFIXES = (
    "/slide",
    "/image",
    "/audio",
    "/video",
    "/media",
    "/hyperlink",
    "/slideLayout",
    "/notesSlide",
    "/tags",
    "/customXml",
)

# Default font sizes per role when the run carries none (inherited from the
# layout/master), used only for the capacity estimate.
_DEFAULT_PT: Dict[str, float] = {
    "title": 32.0,
    "subtitle": 20.0,
    "heading": 20.0,
    "body": 16.0,
    "caption": 12.0,
    "number": 40.0,
    "label": 12.0,
}

# Text repeated on this many slides is decoration (footer, confidentiality).
_SHARED_TEXT_MIN_SLIDES = 2
# A free picture covering this share of the slide is content, not decoration.
_REPLACEABLE_PICTURE_AREA = 0.25
# A chart/table covering this share of the slide is the slide's content.
_MAIN_CONTENT_AREA = 0.15
# A chart frame this large is filled; several share a slide (a 2×2 grid).
_GRID_CHART_AREA = 0.05
# A content placeholder this large may host a chart in place of its text.
_CHART_MIN_AREA = 0.25
# A library sheet lays out this many small text-free icons (pictures, groups,
# freeform glyphs); a design with a few icons beside its text is far below.
_LIBRARY_ITEMS = 20
_LIBRARY_ITEM_AREA = 0.02
# A layout/master shape under this share of the slide is a frame, not a background.
_FRAME_MAX_AREA = 0.5
# Loose text shapes read as one row of items: per-mille limits
# past which a shape is a column body, and the tolerances of "same size, aligned".
_ITEM_MAX_W = 350
_ITEM_MAX_H = 400
_SAME_HEIGHT = 0.10
# Designers size a loose box to its sample text: widths vary more than heights.
_SAME_WIDTH = 0.25
_SAME_BAND = 20
_SAME_COLUMN = 30
_FIELD_GAP = 120
_MAX_FIELDS = 4
# A SUBTITLE placeholder this large (share of the slide) holds paragraphs.
_SUBTITLE_BODY_AREA = 0.2
# A (heading, text) pair: the text sits beside or under its heading, overlapping
# it by half the smaller side, at most this far (per mille; overlaps allowed).
_PAIR_GAP = 120
_PAIR_OVERLAP = -40
# A figure placeholder alone with its caption is a one-item figures slide.
_FIGURE_MIN_PT = 40
# A text-free slide drawing this many shapes is a graphic resource sheet.
_RESOURCE_SHAPES = 8


@dataclass
class _Rec:
    key: str
    path: Tuple[int, ...]
    kind: str  # text | picture | table | chart | group | other
    ph_type: Optional[str]
    ph_idx: Optional[int]
    box: Box
    area: float
    text: str
    font_pt: Optional[float]
    font_name: Optional[str]
    bold: bool
    lines: int
    name: str
    shape: "BaseShape" = field(repr=False)
    frame: "FrameMetrics" = field(default_factory=lambda: FrameMetrics(), repr=False)
    # Holds a slide-number or date field: chrome, never a slot.
    is_field: bool = False

    @property
    def words(self) -> int:
        return len(self.text.split())


def _text_of(shape) -> str:
    return shape.text_frame.text.replace("\x0b", "\n").strip()


def _font_pt_of(shape) -> Optional[float]:
    best: Optional[float] = None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            size = run.font.size
            if size is not None:
                pt = size.pt
                if best is None or pt > best:
                    best = pt
    return best


_THEME_FONTS = {"+mj-lt": 0, "+mn-lt": 1}
# Placeholders drawn with the theme's heading font ("SUBTITLE" is not one).
_TITLE_PLACEHOLDERS = ("TITLE", "CENTER_TITLE", "VERTICAL_TITLE")


# Bumped whenever the rules below change what an inventory holds, so agents
# saved with an older analysis are recomputed from their stored template.
ANALYSIS_VERSION = 17


def theme_fonts(presentation) -> Tuple[Optional[str], Optional[str]]:
    """The theme's (heading, body) latin typefaces, or ``(None, None)``."""
    try:
        master_part = presentation.slide_masters[0].part
        theme_part = next(
            rel.target_part
            for rel in master_part.rels.values()
            if not rel.is_external and rel.reltype.endswith("/theme")
        )
        scheme = etree.fromstring(theme_part.blob).find(f".//{qn('a:fontScheme')}")
    except Exception:  # noqa: BLE001 - a theme-less deck just measures by default
        return None, None
    if scheme is None:
        return None, None

    def typeface(kind: str) -> Optional[str]:
        node = scheme.find(f"{qn(kind)}/{qn('a:latin')}")
        return node.get("typeface") if node is not None else None

    return typeface("a:majorFont"), typeface("a:minorFont")


def _font_name_of(
    shape, theme: Tuple[Optional[str], Optional[str]], is_title: bool
) -> Tuple[Optional[str], bool]:
    """The font a shape's text is drawn with, and whether it is bold. Runs
    first, then any declared typeface, then the theme's heading/body font."""
    name: Optional[str] = None
    bold = False
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            name = name or run.font.name
            bold = bold or bool(run.font.bold)
    if name is None:
        latin = shape._element.find(f".//{qn('a:latin')}")
        name = latin.get("typeface") if latin is not None else None
    if not bold:
        bold = any(
            node.get("b") in ("1", "true", "on")
            for node in shape._element.iter(qn("a:rPr"))
        )
    major, minor = theme
    if name in _THEME_FONTS:
        name = (major, minor)[_THEME_FONTS[name]]
    return (name or (major if is_title else minor)), bold


def _lst_style_pt(element) -> Optional[float]:
    """First-level default size declared in an element's ``a:lstStyle``/``a:lvl1pPr``."""
    if element is None:
        return None
    node = element.find(f".//{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
    if node is None:
        return None
    size = node.get("sz")
    return int(size) / 100 if size else None


def _inherited_from(fmt, holder):
    """The placeholder of ``holder`` a placeholder inherits from: same idx,
    else the first of the same type (a master has one per type)."""
    if holder is None:
        return None
    candidates = list(holder.placeholders)
    for candidate in candidates:
        if candidate.placeholder_format.idx == fmt.idx:
            return candidate
    if fmt.type is None:
        return None
    return next((c for c in candidates if c.placeholder_format.type == fmt.type), None)


def _placeholder_pt(shape, layout, master) -> Optional[float]:
    """Font size a placeholder inherits when its runs carry none: its own list
    style, then the layout's and master's matching placeholder, then the
    master's title/body text styles. Enough for a capacity estimate."""
    own = _lst_style_pt(shape._element.find(f".//{qn('a:lstStyle')}"))
    if own:
        return own
    fmt = shape.placeholder_format
    for holder in (layout, master):
        candidate = _inherited_from(fmt, holder)
        if candidate is not None:
            pt = _font_pt_of(candidate) if candidate.has_text_frame else None
            pt = pt or _lst_style_pt(candidate._element.find(f".//{qn('a:lstStyle')}"))
            if pt:
                return pt
    if master is None:
        return None
    styles = master._element.find(qn("p:txStyles"))
    if styles is None:
        return None
    is_title = fmt.type is not None and fmt.type.name in ("TITLE", "CENTER_TITLE")
    return _lst_style_pt(styles.find(qn("p:titleStyle" if is_title else "p:bodyStyle")))


@dataclass(frozen=True)
class FrameMetrics:
    """What a text box costs in room before a single glyph is drawn.

    PowerPoint lays text inside the box's padding, and advances each paragraph
    by its own line spacing plus the space declared around it. A template that
    sets ``lnSpc 114%`` and ``spcBef 16pt`` on its body style fits barely half
    the bullets a bare "height / (size × 1.2)" would promise.
    """

    lr_inset: int = 2 * text_metrics.DEFAULT_INSET_EMU
    tb_inset: int = 2 * text_metrics.DEFAULT_V_INSET_EMU
    line_spacing: float = 1.0
    space_before_pt: float = 0.0
    space_after_pt: float = 0.0
    # Does not wrap: the box widens with its text instead of breaking a line.
    # Auto-sizing the *height* does not count — such a box still wraps, and the
    # width limit that follows from it is the one that matters.
    grows: bool = False


def _pPr_chain(shape, layout, master) -> List[Any]:
    """Paragraph-property elements that apply to a shape's first level, most
    specific first: its own text, its list style, then the ones it inherits."""
    chain: List[Any] = []
    element = shape._element
    own = element.find(f".//{qn('a:p')}/{qn('a:pPr')}")
    if own is not None:
        chain.append(own)
    lvl = element.find(f".//{qn('a:lstStyle')}/{qn('a:lvl1pPr')}")
    if lvl is not None:
        chain.append(lvl)
    fmt = shape.placeholder_format if shape.is_placeholder else None
    if fmt is not None:
        for holder in (layout, master):
            candidate = _inherited_from(fmt, holder)
            if candidate is not None:
                node = candidate._element.find(
                    f".//{qn('a:lstStyle')}/{qn('a:lvl1pPr')}"
                )
                if node is not None:
                    chain.append(node)
    if master is not None:
        styles = master._element.find(qn("p:txStyles"))
        if styles is not None:
            is_title = (
                fmt is not None
                and fmt.type is not None
                and fmt.type.name
                in (
                    "TITLE",
                    "CENTER_TITLE",
                )
            )
            if fmt is None:
                names = ("p:otherStyle", "p:bodyStyle")
            elif is_title:
                names = ("p:titleStyle",)
            else:
                names = ("p:bodyStyle",)
            for name in names:
                style = styles.find(qn(name))
                node = None if style is None else style.find(qn("a:lvl1pPr"))
                if node is not None:
                    chain.append(node)
    return chain


def _spc_pt(pPr, tag: str, font_pt: Optional[float]) -> Optional[float]:
    """``spcBef``/``spcAft`` of a paragraph, in points."""
    element = pPr.find(qn(tag))
    if element is None:
        return None
    points = element.find(qn("a:spcPts"))
    if points is not None and points.get("val"):
        return int(points.get("val")) / 100
    percent = element.find(qn("a:spcPct"))
    if percent is not None and percent.get("val") and font_pt:
        return int(percent.get("val")) / 100000 * font_pt * text_metrics.LINE_HEIGHT
    return None


def _line_ratio(pPr, font_pt: Optional[float]) -> Optional[float]:
    """``lnSpc`` as a multiple of the font's own line height."""
    element = pPr.find(qn("a:lnSpc"))
    if element is None:
        return None
    percent = element.find(qn("a:spcPct"))
    if percent is not None and percent.get("val"):
        return int(percent.get("val")) / 100000
    points = element.find(qn("a:spcPts"))
    if points is not None and points.get("val") and font_pt:
        return int(points.get("val")) / 100 / (font_pt * text_metrics.LINE_HEIGHT)
    return None


def _frame_metrics(shape, layout, master, font_pt: Optional[float]) -> FrameMetrics:
    """Padding and paragraph spacing a text box applies, inherited included."""
    # A placeholder's empty `bodyPr` inherits its layout's and master's.
    bodies = [shape.text_frame._txBody.find(qn("a:bodyPr"))]
    if shape.is_placeholder:
        for holder in (layout, master):
            source = _inherited_from(shape.placeholder_format, holder)
            if source is not None and source.has_text_frame:
                bodies.append(source.text_frame._txBody.find(qn("a:bodyPr")))

    def attribute(name: str) -> Optional[str]:
        return next(
            (b.get(name) for b in bodies if b is not None and b.get(name) is not None),
            None,
        )

    lr = sum(
        int(attribute(side) or text_metrics.DEFAULT_INSET_EMU)
        for side in ("lIns", "rIns")
    )
    tb = sum(
        int(attribute(side) or text_metrics.DEFAULT_V_INSET_EMU)
        for side in ("tIns", "bIns")
    )
    spacing: Optional[float] = None
    before: Optional[float] = None
    after: Optional[float] = None
    for pPr in _pPr_chain(shape, layout, master):
        if spacing is None:
            spacing = _line_ratio(pPr, font_pt)
        if before is None:
            before = _spc_pt(pPr, "a:spcBef", font_pt)
        if after is None:
            after = _spc_pt(pPr, "a:spcAft", font_pt)
        if spacing is not None and before is not None and after is not None:
            break
    return FrameMetrics(
        lr_inset=max(0, lr),
        tb_inset=max(0, tb),
        line_spacing=spacing or 1.0,
        space_before_pt=max(0.0, before or 0.0),
        space_after_pt=max(0.0, after or 0.0),
        grows=attribute("wrap") == "none",
    )


_CHROME_FIELDS = ("slidenum", "datetime")


def _has_chrome_field(shape) -> bool:
    return any(
        (node.get("type") or "").lower().startswith(_CHROME_FIELDS)
        for node in shape._element.iter(qn("a:fld"))
    )


def _placeholder_key(shape) -> str:
    return f"ph:{shape.placeholder_format.idx}"


def _shape_key(path: Tuple[int, ...]) -> str:
    return "sp:" + "/".join(str(part) for part in path)


def _records(
    shapes,
    width: int,
    height: int,
    layout=None,
    master=None,
    theme: Tuple[Optional[str], Optional[str]] = (None, None),
) -> List[_Rec]:
    """Flatten ``shapes`` (groups included) into records with per-mille boxes.

    ``layout``/``master`` let a placeholder without explicit run sizes report
    the size it inherits, which the capacity estimate needs.
    """
    slide_area = float(width * height) or 1.0
    records: List[_Rec] = []
    for placed in walk_shapes(shapes):
        shape = cast(Any, placed.shape)
        box = Box(
            x=round(placed.left * 1000 / width),
            y=round(placed.top * 1000 / height),
            w=round(placed.width * 1000 / width),
            h=round(placed.height * 1000 / height),
        )
        area = (placed.width * placed.height) / slide_area
        ph_type = ph_idx = None
        if shape.is_placeholder:
            fmt = shape.placeholder_format
            ph_type = fmt.type.name if fmt.type is not None else None
            ph_idx = fmt.idx
        shape_type = shape.shape_type
        text = ""
        font_pt = None
        font_name = None
        bold = False
        lines = 0
        frame = FrameMetrics()
        is_field = False
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            kind = "group"
        elif getattr(shape, "has_chart", False):
            kind = "chart"
        elif getattr(shape, "has_table", False):
            kind = "table"
        elif shape_type == MSO_SHAPE_TYPE.PICTURE or ph_type == "PICTURE":
            kind = "picture"
        elif shape.has_text_frame:
            kind = "text"
            text = _text_of(shape)
            font_pt = _font_pt_of(shape)
            if font_pt is None and ph_idx is not None:
                font_pt = _placeholder_pt(shape, layout, master)
            font_name, bold = _font_name_of(
                shape, theme, is_title=ph_type in _TITLE_PLACEHOLDERS
            )
            frame = _frame_metrics(shape, layout, master, font_pt)
            lines = text.count("\n") + 1 if text else 0
            is_field = _has_chrome_field(shape)
        else:
            kind = "other"
        key = _placeholder_key(shape) if ph_idx is not None else _shape_key(placed.path)
        records.append(
            _Rec(
                key=key,
                path=placed.path,
                kind=kind,
                ph_type=ph_type,
                ph_idx=ph_idx,
                box=box,
                area=area,
                text=text,
                font_pt=font_pt,
                font_name=font_name,
                bold=bold,
                lines=lines,
                name=shape.name or "",
                shape=shape,
                frame=frame,
                is_field=is_field,
            )
        )
    return records


# Roles whose text is a list: every line is its own paragraph, so it pays the
# paragraph spacing again. A title wraps within one paragraph and pays it once.
_BULLET_ROLES = frozenset({"body", "caption"})


def estimate_capacity(
    box: Box,
    font_pt: Optional[float],
    role: str,
    width_emu: int,
    height_emu: int,
    font_name: Optional[str] = None,
    bold: bool = False,
    frame: Optional[FrameMetrics] = None,
    sample: str = "",
) -> Tuple[int, int]:
    """(chars_per_line, max_lines) of a box, as the renderer will really fill it.

    Both dimensions go through the box's own formatting: the padding comes off
    the usable width and height, and a line advances by the paragraph's line
    spacing plus the space declared around it. Widths are measured with the
    real font (see ``text_metrics``); greedy wrapping wastes about half a word
    at every line break, so each line is credited a little less than it holds.
    The designer's ``sample`` calibrates the width (``text_metrics.calibrate``).
    """
    pt = font_pt or _DEFAULT_PT.get(role, 16.0)
    metrics = frame or FrameMetrics()
    width_emu_box = box.w / 1000 * width_emu - metrics.lr_inset
    height_emu_box = box.h / 1000 * height_emu - metrics.tb_inset
    width_pt = max(0.0, width_emu_box) / 914400 * 72
    height_pt = max(0.0, height_emu_box) / 914400 * 72

    pitch = pt * text_metrics.LINE_HEIGHT * metrics.line_spacing
    around = metrics.space_before_pt + metrics.space_after_pt
    if role in _BULLET_ROLES:
        lines = int(height_pt / (pitch + around)) if pitch + around > 0 else 1
    else:
        lines = int((height_pt - around) / pitch) if pitch > 0 else 1
    lines = max(1, lines)

    per_line = text_metrics.chars_per_line(width_pt, pt, font_name, bold)
    waste = text_metrics.WRAP_WASTE_CHARS * (lines - 1)
    chars = max(1, int((per_line * lines - waste) * text_metrics.SAFETY / lines))
    return text_metrics.calibrate(sample, chars, lines, font_name, bold), lines


def _zone(rec: _Rec, role: ZoneRole, width: int, height: int) -> Zone:
    chars_per_line = max_lines = None
    if role not in ("image", "keep", "chart"):
        chars_per_line, max_lines = estimate_capacity(
            rec.box,
            rec.font_pt,
            role,
            width,
            height,
            rec.font_name,
            rec.bold,
            rec.frame,
            rec.text,
        )
    return Zone(
        key=rec.key,
        role=role,
        box=rec.box,
        sample=rec.text[:80],
        font_pt=rec.font_pt,
        chars_per_line=chars_per_line,
        max_lines=max_lines,
        grows=rec.frame.grows,
        chart_capable=rec.ph_type == "OBJECT"
        and role == "body"
        and rec.area >= _CHART_MIN_AREA,
    )


def _is_top_band(rec: _Rec) -> bool:
    return rec.box.y < 120 and rec.box.h < 80


def _promote_headings(zones: List[Zone]) -> None:
    """A short body placeholder sitting right above another body in the same
    column is that column's heading (the Comparison layout shape)."""
    bodies = [z for z in zones if z.role == "body"]
    for zone in bodies:
        if zone.box.h >= 100:
            continue
        for other in bodies:
            # Only a tall body below counts as a column: contact cards stack
            # short lines and those are all bodies.
            if other is zone or other.box.y <= zone.box.y or other.box.h < 150:
                continue
            overlap = min(zone.box.x + zone.box.w, other.box.x + other.box.w) - max(
                zone.box.x, other.box.x
            )
            if (
                overlap > zone.box.w * 0.5
                and other.box.y - (zone.box.y + zone.box.h) < 60
            ):
                zone.role = "heading"
                break


def _placeholder_role(rec: _Rec) -> Optional[ZoneRole]:
    """Role implied by a placeholder type; None for the chrome placeholders."""
    ph = rec.ph_type
    if ph in ("TITLE", "CENTER_TITLE"):
        return "title"
    if ph == "SUBTITLE":
        # Google Slides exports type every text box SUBTITLE: one sized for
        # paragraphs is a body, not the line under a title.
        return "body" if rec.area >= _SUBTITLE_BODY_AREA else "subtitle"
    if ph in ("BODY", "OBJECT"):
        # A body placeholder on the top band is a column heading (some
        # templates type their headings as BODY), not a bullet area.
        return "heading" if _is_top_band(rec) and rec.box.w < 600 else "body"
    if ph == "PICTURE":
        return "image"
    if ph == "CHART":
        return "chart"
    return None


def _declared_role(rec: _Rec) -> Optional[str]:
    return lexicon.role_from_shape_name(rec.name)


_MERGES = ("gridSpan", "rowSpan", "hMerge", "vMerge")
# OOXML defaults for a table cell's margins, in EMU.
_CELL_MARGIN_LR = 91440
_CELL_MARGIN_TB = 45720


def _default_text_pt(presentation) -> float:
    """The presentation's default text size: what a table cell draws with."""
    node = presentation.part._element.find(
        f"{qn('p:defaultTextStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}"
    )
    size = node.get("sz") if node is not None else None
    return int(size) / 100 if size else 18.0


def _cell_text(tc) -> str:
    return "\n".join(
        "".join(t.text or "" for t in p.iter(qn("a:t"))) for p in tc.iter(qn("a:p"))
    ).strip()


def _cell_fill(tc) -> str:
    """The cell's own fill, serialised, so body cells can be compared."""
    props = tc.find(qn("a:tcPr"))
    if props is None:
        return ""
    fills = [
        child
        for child in props
        if child.tag in (qn("a:solidFill"), qn("a:gradFill"), qn("a:pattFill"))
    ]
    return b"".join(etree.tostring(f) for f in fills).decode()


def _cell_pt(tc) -> Optional[float]:
    for tag in ("a:rPr", "a:endParaRPr"):
        for node in tc.iter(qn(tag)):
            if node.get("sz"):
                return int(node.get("sz")) / 100
    return None


def table_spec(
    tbl, frame_height: int, default_pt: float, font_name: Optional[str]
) -> Optional[Tuple[TableSpec, float]]:
    """(spec, font size) of a data table: at least 2 columns and 2 rows, no
    merged cells, body cells empty or sample text, one fill for every body
    cell. Anything else is a drawing (a Gantt grid) and stays excluded."""
    rows = tbl.findall(qn("a:tr"))
    widths = [int(c.get("w", "0")) for c in tbl.iter(qn("a:gridCol"))]
    if len(widths) < 2 or len(rows) < 2 or min(widths) <= 0:
        return None
    if any(tc.get(attr) for tc in tbl.iter(qn("a:tc")) for attr in _MERGES):
        return None
    if any(len(tr.findall(qn("a:tc"))) != len(widths) for tr in rows):
        return None
    props = tbl.find(qn("a:tblPr"))
    header = props is not None and props.get("firstRow") in ("1", "true")
    body = [tc for tr in rows[1 if header else 0 :] for tc in tr.findall(qn("a:tc"))]
    if any(
        (text := _cell_text(tc)) and not lexicon.is_sample_text(text) for tc in body
    ):
        return None
    if len({_cell_fill(tc) for tc in body}) > 1:
        return None
    first = rows[-1].find(qn("a:tc"))
    cell = first.find(qn("a:tcPr")) if first is not None else None

    def margin(name: str, default: int) -> int:
        value = cell.get(name) if cell is not None else None
        return int(value) if value else default

    lr = margin("marL", _CELL_MARGIN_LR) + margin("marR", _CELL_MARGIN_LR)
    tb = margin("marT", _CELL_MARGIN_TB) + margin("marB", _CELL_MARGIN_TB)
    pt = next((p for tc in body if (p := _cell_pt(tc))), None) or default_pt
    columns = [
        text_metrics.chars_per_line(max(0, w - lr) / 914400 * 72, pt, font_name, False)
        for w in widths
    ]
    row_heights = [int(tr.get("h", "0")) for tr in rows]
    line_pt = pt * text_metrics.LINE_HEIGHT
    pad_pt = tb / 914400 * 72
    row_pt = max(row_heights[-1] / 914400 * 72, line_pt + pad_pt)
    height_pt = max(frame_height, sum(row_heights)) / 914400 * 72
    spec = TableSpec(
        header=header,
        columns=columns,
        widths=widths,
        margin=lr,
        line_pt=round(line_pt, 2),
        pad_pt=round(pad_pt, 2),
        row_pt=round(row_pt, 2),
        height_pt=round(height_pt, 2),
    )
    return spec, pt


class _Analyzer:
    def __init__(self, presentation) -> None:
        self.prs = presentation
        self.width = presentation.slide_width
        self.height = presentation.slide_height
        self.shared_texts: set[str] = set()
        self.theme = theme_fonts(presentation)
        self.default_pt = _default_text_pt(presentation)
        self._frames: Dict[Any, List[Box]] = {}

    def frames(self, layout) -> List[Box]:
        """Shapes every slide on ``layout`` shows and none can remove: the
        layout's and (unless it hides them) the master's drawn shapes without
        text, backgrounds aside."""
        if layout.part not in self._frames:
            shapes = list(layout.shapes)
            if layout._element.get("showMasterSp") != "0":
                shapes += list(layout.slide_master.shapes)
            self._frames[layout.part] = [
                rec.box
                for rec in _records(shapes, self.width, self.height)
                if rec.ph_idx is None
                and not rec.text
                and rec.kind != "group"
                and rec.area < _FRAME_MAX_AREA
            ]
        return self._frames[layout.part]

    # -- entry -----------------------------------------------------------------

    def run(self) -> TemplateInventory:
        slides = list(self.prs.slides)
        slide_records = [
            _records(
                slide.shapes,
                self.width,
                self.height,
                slide.slide_layout,
                slide.slide_layout.slide_master,
                self.theme,
            )
            for slide in slides
        ]
        self.shared_texts = self._shared_texts(slide_records)

        patterns: List[Pattern] = []
        # Layout part → (its pattern, the role of each placeholder idx holding a slot).
        layout_slots: Dict[Any, Tuple[Pattern, Dict[int, ZoneRole]]] = {}
        # One running index over every master, matching the renderer's view.
        layouts = [
            layout
            for master in self.prs.slide_masters
            for layout in master.slide_layouts
        ]
        for index, layout in enumerate(layouts):
            pattern, slots = self._layout_pattern(layout, index)
            patterns.append(pattern)
            if pattern.included:
                layout_slots[layout.part] = (pattern, slots)

        copied: set[Any] = set()
        for number, (slide, records) in enumerate(zip(slides, slide_records), start=1):
            pattern = self._slide_pattern(slide, number, records)
            part = slide.slide_layout.part
            layout = layout_slots.get(part)
            draws = _draws(records)
            if (
                layout is None
                or not pattern.included
                or not _covers(records, layout[1])
                or (draws and pattern.kind != layout[0].kind)
            ):
                patterns.append(pattern)
                continue
            # A slide holding its layout's slots stands for it (sample text,
            # notes, position) unless its drawing makes it another kind; a later
            # bare copy of the same layout is dropped.
            if part not in copied:
                patterns[patterns.index(layout[0])] = pattern
            elif draws:
                patterns.append(pattern)
            copied.add(part)

        return TemplateInventory(
            analysis_version=ANALYSIS_VERSION,
            slide_width=self.width,
            slide_height=self.height,
            patterns=patterns,
            shared_texts=sorted(self.shared_texts),
        )

    # -- shared decoration -------------------------------------------------------

    @staticmethod
    def _shared_texts(slide_records: Sequence[Sequence[_Rec]]) -> set[str]:
        counter: Counter[str] = Counter()
        for records in slide_records:
            seen = {
                rec.text
                for rec in records
                if rec.kind == "text" and rec.ph_idx is None and len(rec.text) >= 3
            }
            counter.update(seen)
        return {
            text for text, count in counter.items() if count >= _SHARED_TEXT_MIN_SLIDES
        }

    # -- layouts -----------------------------------------------------------------

    def _layout_pattern(
        self, layout, index: int
    ) -> Tuple[Pattern, Dict[int, ZoneRole]]:
        records = _records(
            layout.placeholders,
            self.width,
            self.height,
            None,
            layout.slide_master,
            self.theme,
        )
        slotted = [rec for rec in records if _placeholder_role(rec) is not None]
        groups, in_rows = self._row_group(
            [
                rec
                for rec in slotted
                if _placeholder_role(rec) in ("body", "heading", "image")
            ]
        )
        paired: Dict[str, ZoneRole] = {}
        if not groups:
            groups, in_rows, paired = self._pair_group(
                [rec for rec in slotted if _placeholder_role(rec) in _PAIR_ROLES]
            )
        zones: List[Zone] = []
        for rec in slotted:
            if rec.key in in_rows:
                continue
            role = paired.get(rec.key) or cast(ZoneRole, _placeholder_role(rec))
            zones.append(_zone(rec, role, self.width, self.height))
        _promote_headings(zones)
        self._chart_zones(zones, records)
        name = layout.name or f"Layout {index + 1}"
        pattern = Pattern(
            id=f"layout:{index}",
            source="layout",
            source_index=index,
            name=name,
            label=name,
            kind="other",
            zones=zones,
            groups=groups,
        )
        pattern.kind = _kind(pattern, records, name, layout._element.get("type"))
        if not zones and not groups and pattern.kind != "closing":
            pattern.included = False
            pattern.exclusion_reason = "no_placeholder"
        _quote_roles(pattern)
        mark_takeaway(pattern, {rec.key: [rec.name] for rec in records})
        _mark_framed(pattern, self.frames(layout))
        pattern.reusable = pattern.kind != "other"
        return pattern, {
            rec.ph_idx: cast(ZoneRole, _placeholder_role(rec))
            for rec in slotted
            if rec.ph_idx is not None
        }

    # -- slides --------------------------------------------------------------------

    def _slide_pattern(self, slide, number: int, records: List[_Rec]) -> Pattern:
        layout_name = slide.slide_layout.name or ""
        zones: List[Zone] = []
        warnings: List[str] = []

        # 1. Repeated sibling groups, else rows of loose same-size shapes.
        groups, grouped_paths = self._repeat_groups(records)
        loose = [
            rec
            for rec in records
            if rec.kind == "text"
            and rec.ph_idx is None
            and rec.text
            and not rec.is_field
            and rec.path[0] not in grouped_paths
        ]
        slots = [
            rec
            for rec in loose
            if rec.text not in self.shared_texts or lexicon.is_sample_text(rec.text)
        ]
        # Text repeated on other slides is chrome, unless it is one of a row of
        # same-size slots here (a timeline entry reused elsewhere).
        free_text = [
            rec
            for rec in loose
            if rec in slots or sum(_same_class(rec, slot) for slot in slots) >= 2
        ]
        in_rows: set[str] = set()
        paired: Dict[str, ZoneRole] = {}
        if not groups:
            groups, in_rows = self._row_group(
                [
                    rec
                    for rec in records
                    if rec.ph_idx is not None
                    and _slot_role(rec)
                    in ("body", "heading", "label", "caption", "image")
                ]
                + [rec for rec in free_text if _declared_role(rec) != "keep"]
            )
        if not groups:
            groups, in_rows, paired = self._pair_group(
                [
                    rec
                    for rec in records
                    if rec.ph_idx is not None
                    and rec.kind == "text"
                    and _slot_role(rec) in _PAIR_ROLES
                ]
                + [rec for rec in free_text if _declared_role(rec) != "keep"]
            )

        # 2. Typed placeholders keep their meaning. A designer who duplicated
        # a placeholder left two shapes with one idx: the second is addressed
        # by its shape path so both stay fillable.
        title_text = ""
        used_keys: set[str] = set()
        for rec in records:
            if rec.ph_idx is None or rec.key in in_rows:
                continue
            if _declared_role(rec) == "keep":
                continue
            role = paired.get(rec.key) or _slot_role(rec)
            if role is None:
                continue
            if role == "title" and not title_text:
                title_text = rec.text
            zone = _zone(rec, role, self.width, self.height)
            if zone.key in used_keys:
                warnings.append(
                    f"Two placeholders share the index {rec.ph_idx}; the second is "
                    "addressed by its shape."
                )
                zone.key = _shape_key(rec.path)
            used_keys.add(zone.key)
            zones.append(zone)
        _promote_headings(zones)

        # 3. Every other text box is a slot, but chrome: text repeated across
        # slides (a repeated sample box still counts) and date/number fields.
        has_title = any(z.role == "title" for z in zones)
        zones.extend(
            self._free_text_zones(
                [rec for rec in free_text if rec.key not in in_rows], has_title, paired
            )
        )
        if not title_text:
            title_text = next((z.sample for z in zones if z.role == "title"), "")

        # 4. Free pictures large enough to be content.
        for rec in records:
            if (
                rec.kind == "picture"
                and rec.ph_idx is None
                and rec.path[0] not in grouped_paths
            ):
                declared = _declared_role(rec)
                if declared == "keep":
                    continue
                if declared == "image" or rec.area >= _REPLACEABLE_PICTURE_AREA:
                    zones.append(_zone(rec, "image", self.width, self.height))

        # 4b. The slide's one main table, when it is data to fill.
        tables = [
            rec
            for rec in records
            if rec.kind == "table" and rec.area >= _MAIN_CONTENT_AREA
        ]
        if len(tables) == 1:
            zone = self._table_zone(tables[0])
            if zone is not None:
                zones.append(zone)

        # 4c. The slide's chart frames, up to four, when no chart placeholder
        # holds one; a chart left out keeps the slide excluded (step 5).
        frames = [
            rec
            for rec in records
            if rec.kind == "chart"
            and rec.ph_idx is None
            and len(rec.path) == 1
            and rec.area >= _GRID_CHART_AREA
            and _declared_role(rec) != "keep"
        ]
        if len(frames) <= MAX_CHARTS_PER_SLIDE and not any(
            z.role == "chart" for z in zones
        ):
            zones.extend(
                _zone(rec, "chart", self.width, self.height)
                for rec in sorted(frames, key=lambda r: (r.box.y // 150, r.box.x))
            )
        self._chart_zones(zones, records)

        pattern = Pattern(
            id=f"slide:{slide.slide_id}",
            source="slide",
            source_index=number,
            name=layout_name,
            label=self._slide_label(title_text, layout_name, number),
            kind="other",
            zones=zones,
            groups=groups,
            warnings=warnings,
        )

        # 5. Exclusions.
        self._apply_exclusions(pattern, slide, records)

        # 6. Label then kind, with the notes directives on top.
        notes = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text if frame is not None else ""
        label = lexicon.LABEL_NOTE.search(notes)
        if label:
            pattern.label = label.group(1)[:60]
        pattern.kind = _kind(
            pattern,
            records,
            layout_name,
            slide.slide_layout._element.get("type"),
            slide_number=number,
        )
        match = lexicon.KIND_NOTE.search(notes)
        if match:
            declared = match.group(1).lower()
            declared = lexicon.LEGACY_KINDS.get(declared, declared)
            if declared in SLIDE_KINDS:
                pattern.kind = cast(SlideKind, declared)
        _quote_roles(pattern)
        mark_takeaway(pattern, self._names(slide, records))
        if slide._element.get("showMasterSp") != "0":
            _mark_framed(pattern, self.frames(slide.slide_layout))
        pattern.reusable = pattern.kind != "other"
        return pattern

    @staticmethod
    def _names(slide, records: List[_Rec]) -> Dict[str, List[str]]:
        """Each record's shape name, with its layout placeholder's for a
        placeholder: the designer may have named either."""
        layout = {
            ph.placeholder_format.idx: ph.name or ""
            for ph in slide.slide_layout.placeholders
        }
        return {
            rec.key: [rec.name]
            + ([layout.get(rec.ph_idx, "")] if rec.ph_idx is not None else [])
            for rec in records
        }

    def _chart_zones(self, zones: List[Zone], records: List[_Rec]) -> None:
        """A chart zone draws its text from the presentation's default size and
        stores its column-chart room; an example chart's type is kept too."""
        by_key = {rec.key: rec for rec in records}
        for zone in zones:
            if zone.role != "chart":
                continue
            zone.font_pt = self.default_pt
            size = (self.width, self.height)
            zone.max_lines = charts.max_categories(zone, "column", size)
            shown = min(charts.QUOTED_CATEGORIES, zone.max_lines)
            zone.chars_per_line = charts.label_room(zone, "column", size, shown)[0]
            rec = by_key.get(zone.key)
            if rec is not None and rec.kind == "chart":
                zone.chart_type = charts.type_of(cast(Any, rec.shape).chart)

    def _table_zone(self, rec: _Rec) -> Optional[Zone]:
        tbl = rec.shape._element.find(f".//{qn('a:tbl')}")
        found = (
            table_spec(tbl, int(rec.shape.height), self.default_pt, self.theme[1])
            if tbl is not None
            else None
        )
        if found is None:
            return None
        spec, pt = found
        return Zone(
            key=rec.key,
            role="table",
            box=rec.box,
            font_pt=pt,
            max_lines=max(1, spec.body_lines()),
            table=spec,
        )

    def _slide_label(self, title_text: str, layout_name: str, number: int) -> str:
        if (
            title_text
            and not lexicon.is_sample_text(title_text)
            and title_text not in self.shared_texts
        ):
            return " ".join(title_text.split())[:60]
        return f"{layout_name or 'Slide'} ({number})"

    # -- repeated groups -------------------------------------------------------------

    def _repeat_groups(self, records: List[_Rec]) -> Tuple[List[RepeatGroup], set[int]]:
        by_group: Dict[int, List[_Rec]] = defaultdict(list)
        group_recs: Dict[int, _Rec] = {}
        for rec in records:
            if rec.kind == "group" and len(rec.path) == 1:
                group_recs[rec.path[0]] = rec
            elif len(rec.path) > 1 and rec.kind in ("text", "picture"):
                by_group[rec.path[0]].append(rec)

        signatures: Dict[tuple, List[int]] = defaultdict(list)
        for group_id, children in by_group.items():
            texts = [c for c in children if c.kind == "text" and c.text]
            if not texts:
                continue
            # Structure only (font sizes, picture count): sibling cards differ
            # in box sizes by a few per cent, never in structure.
            signature = tuple(sorted(round(c.font_pt or 0) for c in texts)) + (
                sum(1 for c in children if c.kind == "picture"),
            )
            signatures[signature].append(group_id)

        groups: List[RepeatGroup] = []
        grouped: set[int] = set()
        for signature, group_ids in signatures.items():
            if len(group_ids) < 3:
                continue
            ordered = sorted(
                group_ids,
                key=lambda gid: (group_recs[gid].box.y // 40, group_recs[gid].box.x),
            )
            item_roles: Optional[List[ZoneRole]] = None
            items: List[List[Zone]] = []
            for gid in ordered:
                children = sorted(
                    (c for c in by_group[gid] if c.kind == "text" and c.text),
                    key=lambda c: (c.box.y, c.box.x),
                )
                roles = self._item_roles(children)
                if item_roles is None:
                    item_roles = roles
                elif roles != item_roles:
                    # Same signature but a different reading: keep the first
                    # item's roles so every item has the same shape.
                    roles = item_roles
                items.append(
                    [
                        _zone(c, r, self.width, self.height)
                        for c, r in zip(children, roles)
                        if r != "keep"
                    ]
                )
            grouped.update(ordered)
            groups.append(
                RepeatGroup(
                    key=_shape_key((ordered[0],)),
                    count=len(ordered),
                    item_zones=[r for r in (item_roles or []) if r != "keep"],
                    items=items,
                    item_keys=[_shape_key((gid,)) for gid in ordered],
                )
            )
        groups.sort(key=lambda g: -g.count)
        return groups, grouped

    def _row_group(self, recs: List[_Rec]) -> Tuple[List[RepeatGroup], set[str]]:
        """The largest row of loose same-size shapes as a repeat group (one at
        most: every group of a slide is fed the same items), and its keys."""
        best = max(
            _rows_of_items(recs),
            key=lambda items: len(items) * len(items[0]),
            default=[],
        )
        # Contact cards stay stacked lines: a closing slide takes one per line.
        if not best or lexicon.has_contact_words(
            [rec.text for item in best for rec in item]
        ):
            return [], set()
        texts = [rec for rec in best[0] if rec.kind == "text"]
        text_roles = iter(self._item_roles(texts))
        roles: List[ZoneRole] = [
            "image" if rec.kind != "text" else next(text_roles) for rec in best[0]
        ]
        items_zones = [
            [
                _zone(rec, role, self.width, self.height)
                for rec, role in zip(item, roles)
                if role != "keep"
            ]
            for item in best
        ]
        group = RepeatGroup(
            key=best[0][0].key,
            count=len(best),
            item_zones=[r for r in roles if r != "keep"],
            items=items_zones,
            item_keys=[item[0].key for item in best],
        )
        return [group], {rec.key for item in best for rec in item}

    def _pair_group(
        self, recs: List[_Rec]
    ) -> Tuple[List[RepeatGroup], set[str], Dict[str, ZoneRole]]:
        """Repeated (heading, text) pairs, whatever their placeholder types, as
        one repeat group, its keys, and roles for zones outside any group: two
        plain pairs stay the columns' headings and bodies."""
        if lexicon.has_contact_words([rec.text for rec in recs]):
            return [], set(), {}
        items = _pairs(recs) or _lone_figure(recs)
        if not items:
            return [], set(), {}
        heads = [item[0].text for item in items]
        texts = [item[1] for item in items]
        roles: List[ZoneRole] = ["heading", "body"]
        if all(lexicon.is_figure(text) for text in heads):
            if _sequential(heads) or all(lexicon.is_date_sample(t) for t in heads):
                # Numbered entries and dates: a short text names the entry.
                short = max(rec.words for rec in texts) <= 4
                roles = ["number", "heading" if short else "body"]
            else:
                # A figure's caption: a label shows the plan's label or heading.
                roles = ["number", "label"]
        elif len(items) == 2:
            return (
                [],
                set(),
                {rec.key: role for item in items for rec, role in zip(item, roles)},
            )
        group = RepeatGroup(
            key=items[0][0].key,
            count=len(items),
            item_zones=roles,
            items=[
                [
                    _zone(rec, role, self.width, self.height)
                    for rec, role in zip(item, roles)
                ]
                for item in items
            ],
            item_keys=[item[0].key for item in items],
        )
        return [group], {rec.key for item in items for rec in item}, {}

    @staticmethod
    def _item_roles(children: List[_Rec]) -> List[ZoneRole]:
        roles: List[Optional[ZoneRole]] = [None] * len(children)
        for i, rec in enumerate(children):
            declared = _declared_role(rec)
            if declared:
                roles[i] = declared  # type: ignore[assignment]
            elif lexicon.is_numberish(rec.text):
                roles[i] = "number"
        rest = [i for i, r in enumerate(roles) if r is None]
        if not rest:
            return [r or "body" for r in roles]
        has_number = any(r == "number" for r in roles)
        ranked = sorted(rest, key=lambda i: -(children[i].font_pt or 0))
        if len(rest) == 1:
            rec = children[rest[0]]
            if has_number:
                roles[rest[0]] = (
                    "label" if rec.words <= 4 and rec.box.h <= 60 else "heading"
                )
            else:
                roles[rest[0]] = (
                    "label" if rec.box.h <= 60 and rec.words <= 4 else "body"
                )
        else:
            roles[ranked[0]] = "heading"
            for i in ranked[1:]:
                rec = children[i]
                roles[i] = "label" if rec.box.h <= 40 and rec.words <= 4 else "body"
        return [r or "body" for r in roles]

    # -- free text ---------------------------------------------------------------------

    def _free_text_zones(
        self,
        free_text: List[_Rec],
        has_title: bool,
        paired: Optional[Dict[str, ZoneRole]] = None,
    ) -> List[Zone]:
        zones: List[Zone] = []
        ranked = sorted(free_text, key=lambda r: -(r.font_pt or 0))
        for rank, rec in enumerate(ranked):
            declared = _declared_role(rec) or (paired or {}).get(rec.key)
            if declared == "keep":
                continue
            marker = KEY_PATTERN.search(rec.text)
            if declared:
                role = cast(ZoneRole, declared)
            elif marker:
                role = "body"
            elif lexicon.is_numberish(rec.text):
                role = "number"
            elif rank == 0 and not has_title and rec.box.y < 150:
                role = "title"
            elif (rec.font_pt or 0) >= 20 and rec.lines <= 2:
                role = "heading"
            elif rec.words <= 4 and rec.box.h <= 60:
                role = "label"
            else:
                role = "body"
            zones.append(_zone(rec, role, self.width, self.height))
        return zones

    # -- exclusions --------------------------------------------------------------------

    def _apply_exclusions(self, pattern: Pattern, slide, records: List[_Rec]) -> None:
        held = {z.key for z in pattern.zones if z.role in ("table", "chart")}
        # The chart part of a chart zone is replaced at render, never copied.
        chart_rids = {
            node.get(qn("r:id"))
            for rec in records
            if rec.key in held and rec.kind == "chart"
            for node in rec.shape._element.iter(qn("c:chart"))
        }
        for rec in records:
            if rec.key in held:
                continue
            if rec.kind in ("chart", "table") and rec.area >= _MAIN_CONTENT_AREA:
                pattern.included = False
                pattern.exclusion_reason = rec.kind  # type: ignore[assignment]
                return
        if (
            lexicon.is_library_title(pattern.label)
            or _icons(records) >= _LIBRARY_ITEMS
            or _resource_sheet(records, slide)
        ):
            pattern.included = False
            pattern.exclusion_reason = "library"
            return
        for rId, rel in slide.part.rels.items():
            if rel.is_external or rId in chart_rids:
                continue
            if not rel.reltype.endswith(COPYABLE_REL_SUFFIXES):
                pattern.included = False
                pattern.exclusion_reason = "unsupported_parts"
                pattern.warnings.append(
                    f"Slide holds a part that cannot be copied: {rel.reltype.rsplit('/', 1)[-1]}."
                )
                return


def _icons(records: List[_Rec]) -> int:
    """Small top-level items bearing no text that read as icons: pictures,
    groups or freeform glyphs, not preset dots, arrows or lines."""
    texty = {rec.path[0] for rec in records if rec.words}
    return sum(
        1
        for rec in records
        if len(rec.path) == 1
        and rec.path[0] not in texty
        and rec.area < _LIBRARY_ITEM_AREA
        and (
            rec.kind in ("picture", "group")
            or rec.shape._element.find(f".//{qn('a:custGeom')}") is not None
        )
    )


def _slot_role(rec: _Rec) -> Optional[ZoneRole]:
    """A placeholder's role, its Selection-pane name refining it. Only a
    content (`OBJECT`) placeholder may be declared an image."""
    role = _placeholder_role(rec)
    if role is None:
        return None
    declared = _declared_role(rec)
    if declared == "image" and rec.ph_type != "OBJECT":
        return role
    return cast(ZoneRole, declared) if declared else role


def _covers(records: List[_Rec], layout_slots: Dict[int, ZoneRole]) -> bool:
    """The slide holds every slot of its layout, whatever decoration it draws.
    A placeholder duplicated on the slide (two with one idx) stands in for a
    layout slot of the same role the slide lacks."""
    present: set[int] = set()
    spare: Counter[str] = Counter()
    for rec in records:
        role = _placeholder_role(rec) if rec.ph_idx is not None else None
        if role is None:
            continue
        if rec.ph_idx in present:
            spare[role] += 1
        present.add(cast(int, rec.ph_idx))
    for idx, role in layout_slots.items():
        if idx in present:
            continue
        if not spare[role]:
            return False
        spare[role] -= 1
    return True


def _draws(records: List[_Rec]) -> bool:
    """Anything on the slide besides its placeholders."""
    return any(
        (rec.kind in ("chart", "table", "picture", "group") and rec.ph_idx is None)
        or (rec.kind == "text" and rec.ph_idx is None and rec.text)
        for rec in records
    )


def _same_class(a: _Rec, b: _Rec) -> bool:
    def near(x: int, y: int, share: float) -> bool:
        return abs(x - y) <= max(8, share * max(x, y))

    return (
        near(a.box.w, b.box.w, _SAME_WIDTH)
        and near(a.box.h, b.box.h, _SAME_HEIGHT)
        and abs((a.font_pt or 0) - (b.font_pt or 0)) < 1
        and a.bold == b.bold
    )


def _apart(a: _Rec, b: _Rec) -> bool:
    return a.box.x >= b.box.x + b.box.w or b.box.x >= a.box.x + a.box.w


def _aligned(upper: List[_Rec], lower: List[_Rec], max_gap: int) -> bool:
    """``lower`` is a row of as many shapes under ``upper``, column by column."""
    if len(upper) != len(lower):
        return False
    gap = lower[0].box.y - max(r.box.y + r.box.h for r in upper)
    return -_SAME_BAND <= gap <= max_gap and all(
        abs((a.box.x + a.box.w / 2) - (b.box.x + b.box.w / 2)) <= _SAME_COLUMN
        for a, b in zip(upper, lower)
    )


def _rows_of_items(recs: List[_Rec]) -> List[List[List[_Rec]]]:
    """Candidate repeat groups among loose text shapes, each a list of items
    (an item lists its shapes top to bottom)."""
    rows: List[List[_Rec]] = []
    cells = [r for r in recs if r.box.w <= _ITEM_MAX_W and r.box.h <= _ITEM_MAX_H]
    for rec in sorted(cells, key=lambda r: (r.box.y, r.box.x)):
        for row in rows:
            if (
                abs(rec.box.y - row[0].box.y) <= _SAME_BAND
                and _same_class(rec, row[0])
                and all(_apart(rec, other) for other in row)
            ):
                row.append(rec)
                break
        else:
            rows.append([rec])
    rows = [sorted(row, key=lambda r: r.box.x) for row in rows if len(row) >= 2]
    rows.sort(key=lambda row: row[0].box.y)

    # Rows of other classes right under a row are the fields of its items.
    chains: List[List[List[_Rec]]] = []
    used: set[int] = set()
    for i, row in enumerate(rows):
        if i in used:
            continue
        chain = [row]
        used.add(i)
        for j in range(i + 1, len(rows)):
            if len(chain) == _MAX_FIELDS or _same_class(rows[j][0], row[0]):
                break  # the next item row of a grid starts here
            if j not in used and _aligned(chain[-1], rows[j], _FIELD_GAP):
                chain.append(rows[j])
                used.add(j)
        chains.append(chain)

    # Chains of the same classes stacked column by column are one grid.
    grids: List[List[List[List[_Rec]]]] = []
    for chain in chains:
        grid = next(
            (
                grid
                for grid in grids
                if len(grid[-1]) == len(chain)
                and all(_same_class(a[0], b[0]) for a, b in zip(grid[-1], chain))
                and _aligned(grid[-1][-1], chain[0], _ITEM_MAX_H)
            ),
            None,
        )
        if grid is None:
            grids.append([chain])
        else:
            grid.append(chain)

    found: List[List[List[_Rec]]] = []
    for grid in grids:
        width, fields = len(grid[0][0]), len(grid[0])
        has_text = any(row[0].kind == "text" for row in grid[0])
        if not has_text or (width < 3 and fields < 2):
            continue
        found.append(
            [[row[i] for row in chain] for chain in grid for i in range(width)]
        )
    return found


# Placeholder roles that may hold one field of a (heading, text) pair: Google
# Slides exports type a card's heading TITLE and its text SUBTITLE.
_PAIR_ROLES = ("title", "subtitle", "body", "heading", "label", "caption")


def _pair_gap(head: _Rec, text: _Rec, beside: bool) -> Optional[int]:
    """How far ``text`` sits right of ``head`` (``beside``), else above or
    under it, sharing half the smaller side; None when it is not its field."""
    a, b = head.box, text.box
    if beside:
        shared = min(a.y + a.h, b.y + b.h) - max(a.y, b.y)
        gap = b.x - (a.x + a.w)
        ok = shared * 2 >= min(a.h, b.h)
    else:
        shared = min(a.x + a.w, b.x + b.w) - max(a.x, b.x)
        gap = max(b.y - (a.y + a.h), a.y - (b.y + b.h))
        ok = shared * 2 >= min(a.w, b.w)
    return gap if ok and _PAIR_OVERLAP <= gap <= _PAIR_GAP else None


def _matched(
    heads: List[_Rec], texts: List[_Rec], beside: bool
) -> Optional[Tuple[int, List[List[_Rec]]]]:
    """Each head with its nearest free text, all in one arrangement, and the
    summed distance; None when a head has none."""
    free = list(texts)
    items: List[List[_Rec]] = []
    total = 0
    for head in heads:
        near = [
            (abs(gap), t)
            for t in free
            if (gap := _pair_gap(head, t, beside)) is not None
        ]
        if not near:
            return None
        distance, text = min(near, key=lambda pair: pair[0])
        free.remove(text)
        items.append([head, text])
        total += distance
    return total, items


def _pairs(recs: List[_Rec]) -> List[List[_Rec]]:
    """The most (heading, text) pairs: two size classes as many each, the
    bigger-font one's every shape with its own text beside or above/under it."""
    classes: List[List[_Rec]] = []
    for rec in recs:
        cls = next((c for c in classes if _same_class(rec, c[0])), None)
        if cls is None:
            classes.append([rec])
        else:
            cls.append(rec)
    best: List[List[_Rec]] = []
    for heads in classes:
        for texts in classes:
            n = len(heads)
            if n < 2 or len(texts) != n or n <= len(best):
                continue
            if (heads[0].font_pt or 0) <= (texts[0].font_pt or 0):
                continue
            found = [m for m in (_matched(heads, texts, b) for b in (True, False)) if m]
            if found:
                best = min(found, key=lambda m: m[0])[1]
    return _item_order(best)


def _lone_figure(recs: List[_Rec]) -> List[List[_Rec]]:
    """A big figure and its caption alone on the slide (a "big number" layout)."""
    if len(recs) != 2:
        return []
    figure, caption = sorted(recs, key=lambda r: -(r.font_pt or 0))
    if (
        (figure.font_pt or 0) < _FIGURE_MIN_PT
        or not lexicon.is_figure(figure.text)
        or _pair_gap(figure, caption, beside=False) is None
    ):
        return []
    return [[figure, caption]]


def _item_order(items: List[List[_Rec]]) -> List[List[_Rec]]:
    """Rows top to bottom, left to right; a staircase (one item per row, each
    further right) reads left to right."""

    def top(item: List[_Rec]) -> int:
        return min(rec.box.y for rec in item)

    def left(item: List[_Rec]) -> int:
        return min(rec.box.x for rec in item)

    rows: List[List[List[_Rec]]] = []
    for item in sorted(items, key=top):
        if rows and top(item) - top(rows[-1][0]) <= 2 * _SAME_BAND:
            rows[-1].append(item)
        else:
            rows.append([item])
    columns = {left(item) // _SAME_COLUMN for item in items}
    if len(rows) == len(items) and len(columns) == len(items):
        return sorted(items, key=left)
    return [item for row in rows for item in sorted(row, key=left)]


def _resource_sheet(records: List[_Rec], slide) -> bool:
    """A slide with no placeholder and no text but links: many drawn shapes
    (an editable-graphics sheet) or a linked picture (a promotion)."""
    if any(rec.ph_idx is not None for rec in records):
        return False
    if any(rec.words and not lexicon.is_url(rec.text) for rec in records):
        return False
    top_level = sum(1 for rec in records if len(rec.path) == 1)
    linked = any(rel.is_external for rel in slide.part.rels.values())
    return top_level >= _RESOURCE_SHAPES or (linked and top_level > 0)


def _mark_framed(pattern: Pattern, frames: List[Box]) -> None:
    """An image zone at least half covered by a frame shows that frame when
    left empty: the fit check then asks for a picture."""
    zones = pattern.zones + [z for g in pattern.groups for i in g.items for z in i]
    for zone in (z for z in zones if z.role == "image"):
        b = zone.box
        for f in frames:
            w = min(b.x + b.w, f.x + f.w) - max(b.x, f.x)
            h = min(b.y + b.h, f.y + f.h) - max(b.y, f.y)
            if w > 0 and h > 0 and w * h * 2 >= b.w * b.h:
                zone.framed = True
                break


def mark_takeaway(pattern: Pattern, said: Mapping[str, Sequence[str]]) -> None:
    """Flag the first text zone, in reading order, whose name or purpose
    (``said`` per zone key) names a key message; never by geometry, and never
    on a slide kind that carries no takeaway."""
    if pattern.kind in lexicon.TAKEAWAY_EXEMPT_KINDS:
        for zone in pattern.zones:
            zone.takeaway = False
        return
    if any(zone.takeaway for zone in pattern.zones):
        return
    headings = {z.key for z in chart_headings(pattern).values()}
    for zone in sorted(pattern.zones, key=lambda z: (z.box.y // 150, z.box.x)):
        if (
            zone.role in lexicon.TAKEAWAY_ROLES
            and zone.key not in headings
            and any(
                lexicon.name_matches(text, lexicon.TAKEAWAY_NAMES)
                for text in said.get(zone.key, ())
                if text
            )
        ):
            zone.takeaway = True
            return


def _quote_roles(pattern: Pattern) -> None:
    """A quote slide without a title says its quote in the top body and its
    author in the next: the plan's title and subtitle."""
    if pattern.kind != "quote" or any(
        z.role in ("title", "heading") for z in pattern.zones
    ):
        return
    bodies = sorted(
        (z for z in pattern.zones if z.role in ("body", "caption")),
        key=lambda z: z.box.y,
    )
    for zone, role in zip(bodies, ("title", "subtitle")):
        zone.role = role  # type: ignore[assignment]


# -- slide kind (heuristic; the vision pass refines it) ---------------------------

# `<p:sldLayout type=…>` values that state a kind; `obj`, `tx` and the custom
# ones say nothing the zones do not.
_LAYOUT_TYPE_KINDS: Dict[str, SlideKind] = {
    "title": "cover",
    "secHead": "section",
    "twoObj": "two_columns",
    "twoColTx": "two_columns",
    "twoTxTwoObj": "comparison",
    "picTx": "image_text",
}
# Kinds that never hold body content, whatever a name says ("3 images +
# message clé" is no quote).
_SPARSE_KINDS = ("cover", "section", "quote")
# A title this short names the slide ("Equipe", "Planche de processus"); a
# longer one is content and says nothing about the kind.
_NAMING_TITLE_WORDS = 4


def _has_body_content(pattern: Pattern) -> bool:
    """Content a cover or a divider never holds: a list, three or more text
    slots beside the title, or a body box sized for paragraphs."""
    if pattern.groups:
        return True
    text = [
        z for z in pattern.zones if z.role not in ("title", "subtitle", "image", "keep")
    ]
    if len(text) >= 3:
        return True
    return any(
        z.role in ("body", "heading")
        and (z.box.h >= 300 or len(z.sample.split()) >= 25)
        for z in text
    )


def _hinted_kind(
    pattern: Pattern, records: List[_Rec], name: str, layout_type: Optional[str]
) -> Optional[SlideKind]:
    """A kind the template states: its layout type, then its layout name, a
    naming title or label, then a declared shape name."""
    # The layout type describes the placeholders, not what was drawn over them.
    if layout_type in _LAYOUT_TYPE_KINDS and not pattern.groups:
        return _LAYOUT_TYPE_KINDS[layout_type]
    names = [name]
    title = next((z.sample for z in pattern.zones if z.role == "title"), "")
    if pattern.source == "slide":
        # The label is the title, or what a `label:` note says.
        for text in (title, pattern.label):
            if len(text.split()) <= _NAMING_TITLE_WORDS or lexicon.is_sample_text(text):
                names.append(text)
    names.extend(rec.name for rec in records if not lexicon.is_auto_name(rec.name))
    for candidate in names:
        kind = lexicon.kind_from_name(candidate)
        if kind is not None:
            return cast(SlideKind, kind)
    return None


def _group_kind(group: RepeatGroup, samples: List[str]) -> SlideKind:
    roles = set(group.item_zones)
    numbers = [z.sample for item in group.items for z in item if z.role == "number"]
    if numbers and all(lexicon.is_date_sample(n) for n in numbers):
        return "timeline"
    if "number" in roles:
        if not _sequential(numbers):
            return "figures"
        if not roles & {"body", "label"}:
            return "agenda"
        # Steps run along one line; numbered texts laid out in a grid are cards.
        firsts = [item[0].box for item in group.items if item]
        rows = len({b.y // 40 for b in firsts})
        columns = len({b.x // 40 for b in firsts})
        grid = 1 < rows < len(firsts) and 1 < columns < len(firsts)
        return "cards" if grid else "process"
    if lexicon.has_contact_words(samples):
        return "team"
    return "cards"


def _kind(
    pattern: Pattern,
    records: List[_Rec],
    name: str,
    layout_type: Optional[str],
    slide_number: int = 0,
) -> SlideKind:
    """The slide kind from the template's own hints, else from the slot roles.
    A slide holding real body content is never a cover, a section or a quote."""
    zones = [z for z in pattern.zones if z.role != "keep"]
    if any(z.role == "table" for z in zones):
        return "table"
    if any(z.role == "chart" for z in zones):
        return "chart"
    hinted = _hinted_kind(pattern, records, name, layout_type)
    if not zones and not pattern.groups:
        return "closing" if hinted == "closing" else "other"
    body = _has_body_content(pattern)
    titled = any(z.role in ("title", "heading") for z in zones)
    if hinted == "quote" and not titled and not pattern.groups and len(zones) <= 3:
        # A quote box is big; without a title it holds nothing else.
        return hinted
    if hinted is not None and not (hinted in _SPARSE_KINDS and body):
        return hinted
    if slide_number == 1 and not body and titled:
        return "cover"

    samples = [z.sample for z in zones] + [
        z.sample for g in pattern.groups for item in g.items for z in item
    ]
    if pattern.groups:
        return _group_kind(pattern.groups[0], samples)
    if len(zones) >= 3 and lexicon.has_contact_words(samples + [name]):
        return "closing"

    titles = [z for z in zones if z.role == "title"]
    subtitles = [z for z in zones if z.role == "subtitle"]
    headings = [z for z in zones if z.role == "heading"]
    bodies = [z for z in zones if z.role == "body"]
    images = [z for z in zones if z.role == "image"]
    labels = [z for z in zones if z.role in ("label", "caption", "number")]

    if titles and not bodies and not headings and not images and not body:
        if subtitles or len(labels) == 1 or slide_number == 1:
            return "cover"
        # A lone title over graphics on an example slide is decoration.
        if pattern.source == "slide" and len(zones) == 1:
            if any(rec.kind == "picture" for rec in records):
                return "other"
        return "section"
    if titles and len(bodies) == 1 and not headings and not images:
        under_title = bodies[0].box.y >= titles[0].box.y and bodies[0].box.h < 120
        if under_title and slide_number == 1 and not body:
            return "cover"
        return "bullets"
    if titles and not bodies and images and not headings:
        return "image_text"
    if len(bodies) >= 2 and not images:
        same_band = max(b.box.y for b in bodies) - min(b.box.y for b in bodies) < 100
        if same_band:
            return "comparison" if len(headings) >= 2 else "two_columns"
        if len(bodies) == 2 and len(headings) == 2:
            return "two_columns"  # two headed blocks stacked
    if images and (bodies or headings):
        return "image_text"
    if len(bodies) == 1 and not titles and not headings:
        return "quote"
    if images and not bodies and not headings:
        return "image_text"
    return "bullets"


def _sequential(numbers: List[str]) -> bool:
    values: List[int] = []
    for raw in numbers:
        digits = re.sub(r"[^0-9]", "", raw)
        if not digits:
            return False
        values.append(int(digits))
    if len(values) < 2:
        return False
    # In any order: a grid of steps lists them row by row.
    values.sort()
    return all(b - a == 1 for a, b in zip(values, values[1:])) or all(
        lexicon.is_year(n) for n in numbers
    )


def analyze_presentation(presentation) -> TemplateInventory:
    """Inventory an open presentation (CPU-bound; call off the event loop)."""
    return _Analyzer(presentation).run()


def analyze_template(pptx_source: Union[bytes, str]) -> AnalyzeResult:
    """Inventory ``.pptx`` bytes or a path, wrapping unreadable input as one error."""
    try:
        presentation = (
            open_presentation(pptx_source)
            if isinstance(pptx_source, bytes)
            else Presentation(pptx_source)
        )
        inventory = analyze_presentation(presentation)
    except Exception as exc:  # noqa: BLE001 - any unreadable .pptx is a template error
        return AnalyzeResult(
            errors=[
                AnalysisError(
                    code=CODE_INVALID_UPLOAD,
                    message=f"The uploaded file could not be read as a .pptx: {exc}",
                )
            ]
        )
    return AnalyzeResult(inventory=inventory, summary=inventory.summary())


__all__ = [
    "CODE_INVALID_UPLOAD",
    "COPYABLE_REL_SUFFIXES",
    "ANALYSIS_VERSION",
    "analyze_presentation",
    "analyze_template",
    "estimate_capacity",
]
