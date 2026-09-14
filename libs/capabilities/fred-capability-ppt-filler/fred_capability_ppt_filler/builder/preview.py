# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The patterns preview deck: one slide per inventory pattern, in inventory
order, so a PDF conversion yields one thumbnail page per pattern for the
review widget. Layout patterns are instantiated with their zone roles typed
into the placeholders; example slides are copied as they are."""

from __future__ import annotations

import io
from typing import Any, cast

from pptx import Presentation

from fred_capability_ppt_filler.builder.models import TemplateInventory
from fred_capability_ppt_filler.builder.render import (
    _all_layouts,
    _duplicate_slide,
    _remove_slides,
)


def build_patterns_deck(template_bytes: bytes, inventory: TemplateInventory) -> bytes:
    """Return ``.pptx`` bytes with one slide per pattern (CPU-bound; call off the loop)."""
    presentation = Presentation(io.BytesIO(template_bytes))
    layouts = _all_layouts(presentation)
    slides_by_id = {slide.slide_id: slide for slide in presentation.slides}
    original_ids = list(slides_by_id)
    for pattern in inventory.patterns:
        if pattern.source == "layout":
            if pattern.source_index >= len(layouts):
                continue
            slide = presentation.slides.add_slide(layouts[pattern.source_index])
            for shape in slide.placeholders:
                if not shape.has_text_frame:
                    continue
                shape = cast(Any, shape)
                zone = next(
                    (
                        z
                        for z in pattern.zones
                        if z.key == f"ph:{shape.placeholder_format.idx}"
                    ),
                    None,
                )
                shape.text_frame.text = zone.role.capitalize() if zone else ""
        else:
            source = slides_by_id.get(int(pattern.id.split(":", 1)[1]))
            if source is not None:
                _duplicate_slide(presentation, source)
    _remove_slides(presentation, original_ids)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
