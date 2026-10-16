# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""After-render check of duplicated example slides, and the per-slide report.

The renderer records into a :class:`DeckReport` what it truncated or could not
place; :func:`remove_leftovers` adds the example text it had to sweep away.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Collection, Dict, Iterable, List, Optional, cast

from pptx.oxml.ns import qn

from fred_capability_ppt_filler.fill import BuildNote, BuildNoteCode, BuiltSlide
from fred_capability_ppt_filler.traversal import walk_shapes

# Placeholders the template means to keep on every slide.
_FIELD_PLACEHOLDERS = frozenset({"FOOTER", "DATE", "SLIDE_NUMBER"})


@dataclass
class SlideReport:
    position: int
    # Example text found on the slide after filling, and removed.
    removed: List[str] = field(default_factory=list)
    # Roles of the zones whose text was cut at a word boundary.
    truncated: List[str] = field(default_factory=list)
    # Plan content no zone received ("subtitle", "2 item(s)", ...).
    lost: List[str] = field(default_factory=list)
    # A table set smaller as a whole to fit its rows (nothing cut).
    shrunk: List[str] = field(default_factory=list)

    @property
    def clean(self) -> bool:
        return not (self.removed or self.truncated or self.lost or self.shrunk)


@dataclass
class DeckReport:
    slides: Dict[int, SlideReport] = field(default_factory=dict)
    # Other findings for the agent, each a full line with its remedy.
    notes: List[str] = field(default_factory=list)
    # Per rendered slide, for the user: what it holds and what was changed.
    built: List[BuiltSlide] = field(default_factory=list)

    def slide(self, position: int) -> SlideReport:
        return self.slides.setdefault(position, SlideReport(position))

    def note(self, code: BuildNoteCode, count: Optional[int] = None, **kw: Any) -> None:
        """A note on the slide being rendered; a repeated code (and field) adds
        its count."""
        if not self.built:
            return
        notes = self.built[-1].notes
        field = kw.get("field")
        same = next((n for n in notes if n.code == code and n.field == field), None)
        if same is None:
            notes.append(BuildNote(code=code, count=count, **kw))
        elif count is not None:
            same.count = (same.count or 0) + count

    def issues(self) -> List[SlideReport]:
        """Slides with something to report, in deck order."""
        return [r for _, r in sorted(self.slides.items()) if not r.clean]


def _text(shape) -> str:
    return shape.text_frame.text.replace("\x0b", "\n").strip()


def _is_field(shape) -> bool:
    if shape.is_placeholder:
        kind = shape.placeholder_format.type
        if kind is not None and kind.name in _FIELD_PLACEHOLDERS:
            return True
    return shape.text_frame._txBody.find(f".//{qn('a:fld')}") is not None


def _text_shapes(slide) -> Iterable[tuple[str, Any]]:
    for placed in walk_shapes(slide.shapes):
        shape = cast(Any, placed.shape)
        if getattr(shape, "has_text_frame", False) and not _is_field(shape):
            yield "sp:" + "/".join(str(p) for p in placed.path), shape


def original_texts(slide) -> Dict[str, str]:
    """Shape-id path key → text of each text shape of an example slide."""
    return {key: text for key, shape in _text_shapes(slide) if (text := _text(shape))}


def clear_paragraphs(txBody) -> None:
    """Empty a text body, keeping its first paragraph's formatting."""
    paragraphs = txBody.findall(qn("a:p"))
    for extra in paragraphs[1:]:
        txBody.remove(extra)
    for child in list(paragraphs[0]):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            paragraphs[0].remove(child)


def remove_leftovers(
    slide,
    originals: Dict[str, str],
    handled: Collection[Any],
    shared: Collection[str],
) -> List[str]:
    """Remove the example text a duplicated slide still shows; returns it.

    ``handled`` holds the elements of the zones the renderer filled, cleared or
    was told to keep; ``shared`` the template's repeated texts (footers).
    """
    removed: List[str] = []
    for key, shape in list(_text_shapes(slide)):
        text = _text(shape)
        if (
            not text
            or originals.get(key) != text
            or text in shared
            or any(shape._element is element for element in handled)
        ):
            continue
        removed.append(text)
        if shape.is_placeholder:
            # An empty placeholder shows its prompt in PowerPoint: drop it.
            element = shape._element
            element.getparent().remove(element)
        else:
            clear_paragraphs(shape.text_frame._txBody)
    return removed


def settle_charts(slide, keep: bool = True) -> int:
    """Pair chart frames and chart parts one to one: a frame without its part goes
    (PowerPoint would repair it), a part no frame shows is unrelated (not saved).
    ``keep=False`` drops every chart: a slide shown as is never shows example data.
    Returns the number of frames removed."""
    rid = qn("r:id")
    rels = slide.part.rels
    frames = [
        node.getparent().getparent().getparent()
        for node in slide.shapes._spTree.iter(qn("c:chart"))
        if not keep or node.get(rid) not in rels
    ]
    for frame in frames:
        frame.getparent().remove(frame)
    shown = {node.get(rid) for node in slide.part._element.iter(qn("c:chart"))}
    for rId in [r for r in rels if rels[r].reltype.endswith("/chart")]:
        if rId not in shown:
            rels.pop(rId)
    return len(frames)


def chart_count(slide) -> int:
    """Native charts the slide shows, each linked to a chart part of its own."""
    rels = slide.part.rels
    parts = {
        id(rels[rId].target_part)
        for node in slide.shapes._spTree.iter(qn("c:chart"))
        if (rId := node.get(qn("r:id"))) in rels
        and rels[rId].reltype.endswith("/chart")
    }
    return len(parts)
