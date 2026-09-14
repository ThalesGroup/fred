# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The built-in theme used when an agent has no template.

Built in memory from python-pptx's default template (no binary in the repo):
a restrained palette and fonts written into the theme part, plus example
slides for the intents the default layouts lack (agenda, cards, steps,
statement, contact). It then goes through the same analysis as any uploaded
template, so "no template" is not a special code path anywhere else.
"""

from __future__ import annotations

import io
from functools import lru_cache
from typing import Any, Tuple, cast

from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from fred_capability_ppt_filler.builder.analysis import analyze_presentation
from fred_capability_ppt_filler.builder.models import TemplateInventory

# Default-template layout indices.
_L_TITLE, _L_CONTENT, _L_SECTION, _L_TWO, _L_COMPARISON, _L_TITLE_ONLY, _L_BLANK = (
    range(7)
)

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GREY = RGBColor(0x5F, 0x66, 0x72)
LIGHT = RGBColor(0xE4, 0xE9, 0xF0)

_THEME_COLORS = {
    "dk1": "1F2A44",
    "lt1": "FFFFFF",
    "dk2": "0E7C86",
    "lt2": "EEF1F5",
    "accent1": "0E7C86",
    "accent2": "1F2A44",
    "accent3": "5F6672",
    "accent4": "D97706",
    "accent5": "2563EB",
    "accent6": "9333EA",
}
# Arial maps to a metric-compatible font on every LibreOffice host, so the
# preview matches what PowerPoint shows.
_FONT = "Arial"


def _restyle_theme(presentation) -> None:
    """Rewrite the theme part's colour scheme and fonts in place."""
    master = presentation.slide_masters[0]
    theme_part = next(
        rel.target_part
        for rel in master.part.rels.values()
        if rel.reltype.endswith("/theme")
    )
    root = etree.fromstring(theme_part.blob)
    scheme = root.find(f".//{qn('a:clrScheme')}")
    if scheme is not None:
        for child in scheme:
            name = child.tag.rsplit("}", 1)[-1]
            if name in _THEME_COLORS:
                for spec in list(child):
                    child.remove(spec)
                srgb = child.makeelement(qn("a:srgbClr"), {"val": _THEME_COLORS[name]})
                child.append(srgb)
    for tag in ("a:majorFont", "a:minorFont"):
        node = root.find(f".//{qn(tag)}/{qn('a:latin')}")
        if node is not None:
            node.set("typeface", _FONT)
    theme_part._blob = etree.tostring(  # pyright: ignore[reportAttributeAccessIssue]
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def _textbox(
    shapes, x, y, w, h, text, size, *, bold=False, color=None, align=None, name=None
):
    box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for paragraph in frame.paragraphs:
        if align is not None:
            paragraph.alignment = align
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = _FONT
            if color is not None:
                run.font.color.rgb = color
    if name:
        box.name = name
    return box


def _titled(presentation, title: str):
    slide = presentation.slides.add_slide(presentation.slide_layouts[_L_TITLE_ONLY])
    cast(Any, slide.shapes.title).text = title
    return slide


def _card(shapes, x, y, w, h):
    card = shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.fill.background()
    card.adjustments[0] = 0.06
    return card


def _agenda(presentation) -> None:
    slide = _titled(presentation, "Agenda")
    for index in range(6):
        col, row = index % 2, index // 2
        x, y = 0.7 + col * 4.5, 1.7 + row * 1.6
        group = slide.shapes.add_group_shape()
        _textbox(
            group.shapes, x, y, 0.9, 0.7, f"0{index + 1}", 28, bold=True, color=TEAL
        )
        _textbox(
            group.shapes,
            x + 1.0,
            y + 0.1,
            3.3,
            0.6,
            "Titre du chapitre",
            16,
            bold=True,
            color=NAVY,
        )


def _cards(presentation) -> None:
    slide = _titled(presentation, "Chiffres clés")
    for index in range(4):
        x = 0.5 + index * 2.3
        group = slide.shapes.add_group_shape()
        _card(group.shapes, x, 1.7, 2.1, 3.4)
        bar = group.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x + 0.2), Inches(1.9), Inches(0.5), Inches(0.06)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = TEAL
        bar.line.fill.background()
        _textbox(
            group.shapes,
            x + 0.2,
            2.2,
            1.7,
            1.0,
            "60%",
            36,
            bold=True,
            color=TEAL,
            align=PP_ALIGN.CENTER,
        )
        _textbox(
            group.shapes,
            x + 0.2,
            3.4,
            1.7,
            0.35,
            "Libellé",
            12,
            color=GREY,
            align=PP_ALIGN.CENTER,
        )


def _steps(presentation) -> None:
    slide = _titled(presentation, "Démarche")
    for index in range(4):
        x = 0.5 + index * 2.3
        group = slide.shapes.add_group_shape()
        _textbox(
            group.shapes, x, 1.6, 1.0, 0.9, str(index + 1), 40, bold=True, color=TEAL
        )
        _textbox(
            group.shapes,
            x,
            2.6,
            2.1,
            0.5,
            f"Étape {index + 1}",
            16,
            bold=True,
            color=NAVY,
        )
        _textbox(group.shapes, x, 3.2, 2.1, 2.5, "Changez ce texte", 12, color=GREY)


def _statement(presentation) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[_L_BLANK])
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(2.4), Inches(10), Inches(2.7)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.name = "keep"
    _textbox(
        slide.shapes,
        0.8,
        2.8,
        8.4,
        1.9,
        "Cliquez ici pour saisir le message clé",
        28,
        bold=True,
        color=RGBColor(0xFF, 0xFF, 0xFF),
        align=PP_ALIGN.CENTER,
        name="body",
    )
    slide.notes_slide.notes_text_frame.text = "intent: statement\nlabel: Message clé"


def _contact(presentation) -> None:
    slide = _titled(presentation, "Contact")
    _textbox(slide.shapes, 0.8, 2.0, 5.0, 0.6, "Nom Prénom", 22, bold=True, color=NAVY)
    _textbox(slide.shapes, 0.8, 2.7, 5.0, 0.5, "Intitulé du poste", 16, color=GREY)
    _textbox(slide.shapes, 0.8, 3.5, 5.0, 0.4, "Numéro de téléphone", 14, color=NAVY)
    _textbox(slide.shapes, 0.8, 4.0, 5.0, 0.4, "email", 14, color=NAVY)
    slide.notes_slide.notes_text_frame.text = "intent: contact\nlabel: Contact"


def _closing(presentation) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[_L_BLANK])
    _textbox(
        slide.shapes,
        0.8,
        2.6,
        8.4,
        1.4,
        "Merci",
        44,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
        name="keep",
    )
    slide.notes_slide.notes_text_frame.text = "label: Merci"


@lru_cache(maxsize=1)
def default_theme() -> Tuple[bytes, TemplateInventory]:
    """The built-in theme's bytes and its inventory, built once per process."""
    presentation = Presentation()
    _restyle_theme(presentation)
    _agenda(presentation)
    _cards(presentation)
    _steps(presentation)
    _statement(presentation)
    _contact(presentation)
    _closing(presentation)
    inventory = analyze_presentation(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue(), inventory
