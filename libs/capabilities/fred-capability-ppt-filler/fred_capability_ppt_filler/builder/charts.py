# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Charts: the plan's data checked, measured against the zone that will show
it, and drawn as a native python-pptx chart styled by the template's theme.

The type mapping, chart-data construction, legend rule and series-length check
are adapted from ``fred-capability-pptx-generator``, by Shuigami.
"""

from __future__ import annotations

import copy
import re
from collections import Counter
from typing import Any, List, Optional, Sequence, Tuple, cast

from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.chart.datalabel import DataLabels
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_MARK,
)
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

from fred_capability_ppt_filler import text_metrics
from fred_capability_ppt_filler.builder import lexicon
from fred_capability_ppt_filler.builder.models import Zone
from fred_capability_ppt_filler.builder.plan import (
    MAX_SERIES,
    MAX_SLICES,
    STACKED_TYPES,
    XY_TYPES,
    PlanChart,
    PlanPoint,
    capacity_text,
)
from fred_capability_ppt_filler.traversal import PlacedShape, walk_shapes

XL_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "area": XL_CHART_TYPE.AREA,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "radar": XL_CHART_TYPE.RADAR_MARKERS,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "bubble": XL_CHART_TYPE.BUBBLE,
}
# The example chart's type in the plan's vocabulary (phase B reuses it).
_FROM_XL = {xl: name for name, xl in XL_TYPES.items()} | {
    XL_CHART_TYPE.LINE: "line",
    XL_CHART_TYPE.PIE_EXPLODED: "pie",
    XL_CHART_TYPE.DOUGHNUT_EXPLODED: "doughnut",
    XL_CHART_TYPE.COLUMN_STACKED_100: "stacked_column",
    XL_CHART_TYPE.BAR_STACKED_100: "stacked_bar",
    XL_CHART_TYPE.RADAR: "radar",
    XL_CHART_TYPE.RADAR_FILLED: "radar",
    XL_CHART_TYPE.XY_SCATTER_LINES: "scatter",
    XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS: "scatter",
    XL_CHART_TYPE.XY_SCATTER_SMOOTH: "scatter",
    XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS: "scatter",
    XL_CHART_TYPE.BUBBLE_THREE_D_EFFECT: "bubble",
}
_PERCENT_XL = {
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED_100,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED_100,
}
_HORIZONTAL = ("bar", "stacked_bar")
ROUND = ("pie", "doughnut")
_ACCENTS = tuple(f"accent{n}" for n in range(1, 7))
_COLOUR_TAGS = (qn("a:schemeClr"), qn("a:srgbClr"))
# Readability caps, and the share of the frame the plot gets beside its axes.
_MAX_CATEGORIES = {"line": 24, "area": 24}
_MAX_CATEGORIES_DEFAULT = 12
_RADAR_CATEGORIES = (3, 8)
_RADAR_SERIES = 3
_SCATTER_POINTS = (3, 50)
_BUBBLE_POINTS = 15
# Point names show beside their point up to this many points (owner decision).
_LABELLED_POINTS = 12
_POINT_LABEL_WORDS = 2
_BUBBLE_SCALE = 50
_BUBBLE_ALPHA = 75000
_SCATTER_MARKER = 9
# The longest "Bubble size: " caption prefix of the deck languages.
_SIZE_PREFIX = len("Taille des bulles : ")
_LABELLED_BARS = 8
# The catalogue quotes a label's room among this many categories.
QUOTED_CATEGORIES = 6
_PLOT_W = 0.85
_PLOT_H = 0.75
_CATEGORY_EM = 3.0
# Line heights a bar's category takes: its label on one line or two, and the gap.
_BAR_LINES = {1: 1.7, 2: 2.2}
# A chart zone's text size when no example chart shows one.
_CHART_PT = 12.0
_EMU_PER_PT = 12700


def type_of(chart: Any) -> Optional[str]:
    """The plan type an existing chart shows, or None for another kind of chart."""
    try:
        return _FROM_XL.get(chart.chart_type)
    except (KeyError, NotImplementedError, ValueError):
        return None


def chart_pt(zone: Zone) -> float:
    """One text size for axes, legend and labels: a chart zone's own (its example
    chart's), else two thirds of a text zone's."""
    if zone.role == "chart":
        return float(min(14, max(8, zone.font_pt or _CHART_PT)))
    return float(min(14, max(10, round((zone.font_pt or 18.0) * 2 / 3))))


def text_pt(chart: Any) -> Optional[float]:
    """The size a chart's text is drawn at: the most common one it sets."""
    sizes = Counter(
        node.get("sz")
        for node in chart._chartSpace.iter(qn("a:defRPr"), qn("a:rPr"))
        if node.get("sz")
    )
    return int(sizes.most_common(1)[0][0]) / 100 if sizes else None


def _frame_pt(zone: Zone, size: Tuple[int, int]) -> Tuple[float, float]:
    return (
        zone.box.w / 1000 * size[0] / _EMU_PER_PT,
        zone.box.h / 1000 * size[1] / _EMU_PER_PT,
    )


def max_categories(
    zone: Zone, kind: str, size: Tuple[int, int], label_lines: int = 2
) -> int:
    if kind in ROUND:
        return MAX_SLICES
    if kind == "radar":
        return _RADAR_CATEGORIES[1]
    width, height = _frame_pt(zone, size)
    pt = chart_pt(zone)
    if kind in _HORIZONTAL:
        lines = _BAR_LINES[label_lines]
        room = height * _PLOT_H / (pt * text_metrics.LINE_HEIGHT * lines)
    else:
        room = width * _PLOT_W / (pt * _CATEGORY_EM)
    cap = _MAX_CATEGORIES.get(kind, _MAX_CATEGORIES_DEFAULT)
    return max(2, min(cap, int(room)))


def label_room(
    zone: Zone, kind: str, size: Tuple[int, int], count: int
) -> Tuple[int, int]:
    """(chars per line, lines) of one category label among ``count``."""
    width, _ = _frame_pt(zone, size)
    plot = width * _PLOT_W
    if kind in ROUND:
        span, lines = plot / 3, 1
    elif kind == "radar":
        span, lines = plot / 4, 1
    elif kind in _HORIZONTAL:
        span, lines = plot * 0.4, 2
    else:
        span, lines = plot / max(1, count), 2
    return text_metrics.chars_per_line(span, chart_pt(zone), None, False), lines


def _name_room(zone: Zone, size: Tuple[int, int], count: int) -> int:
    width, _ = _frame_pt(zone, size)
    span = width * _PLOT_W / max(1, min(count, MAX_SERIES))
    return text_metrics.chars_per_line(span, chart_pt(zone), None, False)


def _title_room(zone: Zone, size: Tuple[int, int]) -> int:
    width, _ = _frame_pt(zone, size)
    return text_metrics.chars_per_line(width, chart_pt(zone) * 1.2, None, True)


def named(found: List[Tuple[str, str]], at: str) -> List[Tuple[str, str]]:
    """Problems of one chart of several, each naming it (``charts[1].title …``)."""
    if at == "chart":
        return found

    def renamed(text: str) -> str:
        return re.sub(
            r"^chart: ", f"{at}: ", re.sub(r"\bchart\.(?=\w)", f"{at}.", text)
        )

    return [
        (what if what.startswith(at) else f"{at}: {what}", fix)
        for what, fix in ((renamed(w), renamed(f)) for w, f in found)
    ]


def data_problems(chart: PlanChart) -> List[Tuple[str, str]]:
    """What is wrong with the chart's data whatever slide shows it: (what, fix)."""
    if chart.type in XY_TYPES:
        return _point_problems(chart)
    if any(s.points for s in chart.series):
        return [
            (
                f"a '{chart.type}' chart takes categories and values, not points",
                "give categories and series [{name, values}], or use type 'scatter' "
                "for points",
            )
        ]
    found: List[Tuple[str, str]] = []
    count = len(chart.categories)
    if not count or not chart.series:
        return [
            (
                "chart: no data; give categories and at least one series",
                "give `chart` {type, categories, series [{name, values}]}, or drop it",
            )
        ]
    if chart.type in ROUND and chart.asked_show:
        found.append(
            (
                f"chart.show '{chart.asked_show}' is not a label mode",
                "use 'share' (percentages), 'value' (the numbers) or 'both'",
            )
        )
    if chart.type in ROUND and len(chart.series) > 1:
        found.append(
            (
                f"a '{chart.type}' chart shows one series, {len(chart.series)} given",
                "keep one series, or use type 'column'",
            )
        )
    elif chart.type == "radar" and len(chart.series) > _RADAR_SERIES:
        found.append(
            (
                f"a 'radar' chart: {len(chart.series)} series; at most {_RADAR_SERIES} "
                "stay readable (the shapes overlap)",
                f"keep the {_RADAR_SERIES} that matter, or use type 'column'",
            )
        )
    elif len(chart.series) > MAX_SERIES:
        found.append(
            (
                f"chart: {len(chart.series)} series; at most {MAX_SERIES} stay readable",
                f"keep the {MAX_SERIES} that matter, or split them over two charts",
            )
        )
    if chart.type in ROUND and count > MAX_SLICES:
        found.append(
            (
                f"a '{chart.type}' chart: {count} slices; at most {MAX_SLICES} stay readable",
                f"keep {MAX_SLICES - 1} and group the smallest as 'Other', or use type 'bar'",
            )
        )
    low, high = _RADAR_CATEGORIES
    if chart.type == "radar" and not low <= count <= high:
        found.append(
            (
                f"radar: {count} criteria; a radar needs {low} to {high}",
                "use type 'column' or 'bar'"
                + (f", or keep the {high} that matter" if count > high else ""),
            )
        )
    for j, series in enumerate(chart.series):
        at = f"chart.series[{j}]"
        named = f" '{series.name}'" if series.name else ""
        if len(series.values) != count:
            found.append(
                (
                    f"{at}{named} has {len(series.values)} values for {count} categories",
                    "give one value per category, in the categories' order",
                )
            )
        for i, text in sorted(series.unread.items()):
            found.append(
                (
                    f"{at}.values[{i}] '{text}' is not a number",
                    "write a number, or drop that category",
                )
            )
        if len(chart.series) > 1 and not series.name:
            found.append((f"{at} has no name", "name each series: the legend shows it"))
    if chart.type in ROUND:
        values = [v for s in chart.series[:1] for v in s.values if v is not None]
        if any(v < 0 for v in values) or (values and not any(values)):
            found.append(
                (
                    f"a '{chart.type}' chart needs shares: positive values, not all 0",
                    "use type 'column' or 'bar' for these values",
                )
            )
    return found


def _point_problems(chart: PlanChart) -> List[Tuple[str, str]]:
    """A scatter's or bubble's data: numeric points, both axis titles, caps."""
    kind = chart.type
    points = chart.points()
    if not points:
        if chart.categories and any(s.values for s in chart.series):
            shown = ", ".join(f"'{c}'" for c in chart.categories[:3])
            return [
                (
                    f"a '{kind}' chart takes points; series[0] has values over "
                    f"categories {shown}…",
                    "use type 'line' or 'column', or give points [{x, y}]",
                )
            ]
        return [
            (
                f"chart: no points; a '{kind}' chart takes series [{{name, points}}]",
                "give each series its points [{x, y"
                + (", size" if kind == "bubble" else "")
                + "}], or drop the chart",
            )
        ]
    found: List[Tuple[str, str]] = []
    if len(chart.series) > MAX_SERIES:
        found.append(
            (
                f"chart: {len(chart.series)} series; at most {MAX_SERIES} stay readable",
                f"keep the {MAX_SERIES} that matter, or split them over two charts",
            )
        )
    low, high = _SCATTER_POINTS if kind == "scatter" else (1, _BUBBLE_POINTS)
    if len(points) < low:
        found.append(
            (
                f"a '{kind}' chart: {len(points)} points; it needs {low} to {high}",
                "give more points, or show these figures on a figures slide",
            )
        )
    elif len(points) > high:
        found.append(
            (
                f"a '{kind}' chart: {len(points)} points; at most {high} stay readable",
                f"keep the {high} that matter"
                + (", or use type 'scatter' (no size)" if kind == "bubble" else ""),
            )
        )
    for field, side, example in (
        ("x_title", "horizontal", "Price"),
        ("y_title", "vertical", "Units sold"),
    ):
        if not (getattr(chart, field) or "").strip():
            found.append(
                (
                    f"chart: no {field}; say what the {side} axis measures ('{example}')",
                    "set chart.x_title and chart.y_title to what each axis measures",
                )
            )
    for j, series in enumerate(chart.series):
        at = f"chart.series[{j}]"
        if series.values and not series.points:
            found.append((f"{at} has values, not points", "give its points [{x, y}]"))
        for i, point in enumerate(series.points):
            found += _point_problem(f"{at}.points[{i}]", point, kind)
        if len(chart.series) > 1 and not series.name:
            found.append((f"{at} has no name", "name each series: the legend shows it"))
    if kind == "scatter" and any(p.size is not None for p in points):
        found.append(
            (
                "points carry size: use type 'bubble'",
                "set chart.type to 'bubble', or drop the sizes",
            )
        )
    return found


def _point_problem(at: str, point: PlanPoint, kind: str) -> List[Tuple[str, str]]:
    found = [
        (f"{at}.{axis} '{text}' is not a number", "write a number, or drop that point")
        for axis, text in sorted(point.unread.items())
    ]
    missing = [
        a for a in ("x", "y") if getattr(point, a) is None and a not in point.unread
    ]
    if missing:
        found.append((f"{at} has no {' or '.join(missing)}", "give each point {x, y}"))
    if kind == "bubble" and "size" not in point.unread:
        if point.size is None:
            found.append((f"{at} has no size", "give each bubble a positive size"))
        elif point.size <= 0:
            found.append(
                (
                    f"{at}.size is {point.size:g}; a bubble needs a positive size",
                    "give a size above 0, or drop that point",
                )
            )
    words = len((point.label or "").split())
    if words > _POINT_LABEL_WORDS:
        found.append(
            (
                f"{at}.label '{point.label}' is {words} words; a point name holds "
                f"{_POINT_LABEL_WORDS}",
                f"shorten each point label to {_POINT_LABEL_WORDS} words (a name)",
            )
        )
    return found


def axis_title(title: Optional[str], unit: Optional[str]) -> str:
    """``Price (€)``: an axis title with its unit."""
    title, unit = (title or "").strip(), (unit or "").strip()
    return f"{title} ({unit})" if title and unit else title


def _xy_fit(
    zone: Zone, chart: PlanChart, size: Tuple[int, int]
) -> List[Tuple[str, str]]:
    """Each axis title holds one line of the plot's side."""
    width, height = _frame_pt(zone, size)
    pt = chart_pt(zone)
    found: List[Tuple[str, str]] = []
    for field, unit, span in (
        ("x_title", chart.x_unit, width * _PLOT_W),
        ("y_title", chart.unit, height * _PLOT_H),
    ):
        text = axis_title(getattr(chart, field), unit)
        room = text_metrics.chars_per_line(span, pt, None, False)
        if len(text) > room:
            found.append(
                (
                    f"chart.{field} is {len(text)} chars with its unit; it holds "
                    f"{capacity_text(room, 1)}",
                    f"shorten chart.{field} to {capacity_text(room, 1)}",
                )
            )
    return found


def fit_problems(
    zone: Zone, chart: PlanChart, size: Tuple[int, int]
) -> List[Tuple[str, str]]:
    """What does not fit this zone at a readable size: (what, fix)."""
    if chart.type in XY_TYPES:
        found = _xy_fit(zone, chart, size)
    else:
        found = _category_fit(zone, chart, size)
    if len(chart.series) > 1 or chart.type in ROUND:
        name_room = _name_room(zone, size, len(chart.series))
        for j, series in enumerate(chart.series):
            if len(series.name) > name_room:
                found.append(
                    (
                        f"chart.series[{j}].name is {len(series.name)} chars; the legend "
                        f"holds {capacity_text(name_room, 1)}",
                        f"shorten the series names to {capacity_text(name_room, 1)}",
                    )
                )
                break
    title_room = _title_room(zone, size)
    if chart.title and len(chart.title) > title_room:
        found.append(
            (
                f"chart.title is {len(chart.title)} chars; it holds "
                f"{capacity_text(title_room, 1)}",
                f"shorten chart.title to {capacity_text(title_room, 1)}, or leave it out "
                "(the slide title says it)",
            )
        )
    size_line = axis_title(chart.size_title, chart.size_unit)
    if (
        chart.type == "bubble"
        and size_line
        and len(size_line) + _SIZE_PREFIX > title_room
    ):
        found.append(
            (
                f"chart.size_title is {len(size_line)} chars with its unit; the caption "
                f"line holds {capacity_text(title_room - _SIZE_PREFIX, 1)}",
                f"shorten chart.size_title to {capacity_text(title_room - _SIZE_PREFIX, 1)}",
            )
        )
    return found


def _category_fit(
    zone: Zone, chart: PlanChart, size: Tuple[int, int]
) -> List[Tuple[str, str]]:
    found: List[Tuple[str, str]] = []
    count = len(chart.categories)
    # Bars whose labels all hold on one line each take less height.
    one_line = label_room(zone, chart.type, size, count)[0]
    short = chart.type in _HORIZONTAL and all(
        len(c) <= one_line for c in chart.categories
    )
    most = max_categories(zone, chart.type, size, 1 if short else 2)
    if count > most and chart.type not in ROUND + ("radar",):
        found.append(
            (
                f"chart: {count} categories; this chart holds {most}",
                f"keep {most} categories (group the rest, or a coarser period)",
            )
        )
    per_line, lines = label_room(zone, chart.type, size, min(count, most))
    room = capacity_text(per_line, lines)
    long = [(i, c) for i, c in enumerate(chart.categories) if len(c) > per_line * lines]
    if long:
        i, label = max(long, key=lambda pair: len(pair[1]))
        found.append(
            (
                f"chart.categories[{i}] is {len(label)} chars; each label holds {room}",
                f"shorten every category label to {room}",
            )
        )
    return found


# --- intents ---------------------------------------------------------------------------

# Categories in time order: a sort would scramble them, and one point of their
# single line or area cannot be coloured apart (LibreOffice ignores point markers).
_ORDERED = ("line", "area")
# Types whose order is their own (positions, or a polygon): never sorted.
_OWN_ORDER = ("radar",) + XY_TYPES
_UNSORTED = _ORDERED + _OWN_ORDER
# Types where one category cannot stand out: only a series can.
_NO_CATEGORY = _ORDERED + ("radar",)
# Opacity of muted text-colour points or series: below the accent on a light or
# a dark background alike; neighbours alternate so a legend still tells them apart.
_MUTED = (35000, 55000)


def _a(kind: str) -> str:
    return f"an {kind}" if kind == "area" else f"a {kind}"


def _names(values: Sequence[str]) -> List[str]:
    return [lexicon.normalize(v) for v in values]


def emphasis(chart: PlanChart) -> Optional[Tuple[str, int]]:
    """("series" | "category", index) of what ``highlight`` names, if it can show."""
    name = lexicon.normalize(chart.highlight or "")
    if not name:
        return None
    if len(chart.series) > 1 and chart.type not in ROUND:
        names = _names([s.name for s in chart.series])
        return ("series", names.index(name)) if name in names else None
    if chart.type in XY_TYPES:
        # One series: a point stands out by its label.
        labels = _names([p.label or "" for p in chart.series[0].points])
        return ("category", labels.index(name)) if name in labels else None
    categories = _names(chart.categories)
    if chart.type in _NO_CATEGORY or name not in categories:
        return None
    return "category", categories.index(name)


def _highlight_note(chart: PlanChart) -> Optional[str]:
    # Told after delivery, when the model can only reply: a fact for the user.
    asked = chart.asked_intents.get("highlight")
    if asked:
        return f"the chart has no emphasis: '{asked}' names more than one item."
    if not chart.highlight or emphasis(chart) is not None:
        return None
    name = chart.highlight
    series = [s.name for s in chart.series]
    if len(series) > 1 and chart.type not in ROUND:
        if lexicon.normalize(name) in _names(chart.categories):
            return (
                f"the chart has no emphasis: '{name}' is a category, and with "
                f"{len(series)} series only a series can stand out "
                f"({', '.join(series)})."
            )
        return (
            f"the chart has no emphasis: '{name}' is not one of its series "
            f"({', '.join(series)})."
        )
    if chart.type in XY_TYPES:
        labels = [p.label for p in chart.series[0].points if p.label] if series else []
        return (
            f"the chart has no emphasis: '{name}' is not one of its point labels "
            f"({', '.join(labels) or 'none given'})."
        )
    if chart.type in _NO_CATEGORY and lexicon.normalize(name) in _names(
        chart.categories
    ):
        return (
            f"the chart has no emphasis on '{name}': {_a(chart.type)} chart cannot "
            "emphasise one category (a column chart can)."
        )
    return (
        f"the chart has no emphasis: '{name}' is not one of its categories "
        f"({', '.join(chart.categories)})."
    )


def intent_notes(chart: PlanChart) -> List[Tuple[str, str]]:
    """(field, note) for each intent the chart is built without."""
    found: List[Tuple[str, str]] = []
    highlight = _highlight_note(chart)
    if highlight:
        found.append(("highlight", highlight))
    asked = chart.asked_intents.get("sort")
    if asked:
        found.append(
            (
                "sort",
                f"chart.sort '{asked}' is not an order; built in the order given. To "
                "sort, edit chart.sort to 'desc' (largest first) or 'asc'.",
            )
        )
    elif chart.sort in ("desc", "asc") and chart.type in _OWN_ORDER:
        found.append(
            (
                "sort",
                f"chart.sort '{chart.sort}' was not applied: a '{chart.type}' chart "
                "keeps its own order. Nothing to do; drop chart.sort, or use type "
                "'column' or 'bar' to sort.",
            )
        )
    elif chart.sort in ("desc", "asc") and chart.type in _ORDERED:
        found.append(
            (
                "sort",
                f"chart.sort '{chart.sort}' was not applied: {_a(chart.type)} chart keeps "
                "its categories in the order given (usually time). Nothing to do; drop "
                "chart.sort, or use type 'column' or 'bar' to sort.",
            )
        )
    asked = chart.asked_intents.get("stacked_as")
    if asked:
        found.append(
            (
                "stacked_as",
                f"chart.stacked_as '{asked}' is not a stacking; built with the values. "
                "Use 'percent' for stacked columns or bars that each fill 100 %, or "
                "drop it.",
            )
        )
    elif chart.stacked_as == "percent" and chart.type not in STACKED_TYPES:
        found.append(
            (
                "stacked_as",
                f"chart.stacked_as 'percent' was not applied: it needs type "
                f"'stacked_column' or 'stacked_bar'; a '{chart.type}' chart was built. "
                "Nothing to do; drop chart.stacked_as, or set chart.type to "
                "'stacked_column'.",
            )
        )
    return found


def _arranged(chart: PlanChart) -> PlanChart:
    """The chart in the order ``sort`` asks; the plan itself keeps the order written."""
    if chart.sort not in ("desc", "asc") or chart.type in _UNSORTED:
        return chart
    count = len(chart.categories)
    if chart.type in STACKED_TYPES:
        keys = [sum(s.values[i] or 0 for s in chart.series) for i in range(count)]
    else:
        keys = [v or 0 for v in chart.series[0].values]
    order = sorted(range(count), key=keys.__getitem__, reverse=chart.sort == "desc")
    if chart.type in _HORIZONTAL:
        # A bar chart draws its first category at the bottom: the first in order goes last.
        order.reverse()
    return chart.model_copy(
        update={
            "categories": [chart.categories[i] for i in order],
            "series": [
                s.model_copy(update={"values": [s.values[i] for i in order]})
                for s in chart.series
            ],
        }
    )


# --- drawing -------------------------------------------------------------------------


def _decimals(values: Sequence[Optional[float]]) -> int:
    return max(
        (len(f"{v:.2f}".rstrip("0").split(".")[1]) for v in values if v is not None),
        default=0,
    )


def number_format(chart: PlanChart, axis: str = "y") -> str:
    """``#,##0.0" M€"``: the data's own decimals (at most 2), then its unit; on a
    scatter or bubble, those of one axis."""
    if chart.type in XY_TYPES:
        values = [getattr(p, axis) for p in chart.points()]
        unit = {"x": chart.x_unit, "size": chart.size_unit}.get(axis, chart.unit)
    else:
        values = [v for s in chart.series for v in s.values]
        unit = chart.unit
    decimals = _decimals(values)
    base = "#,##0" + ("." + "0" * decimals if decimals else "")
    unit = (unit or "").replace('"', "").strip()
    if not unit:
        return base
    return f'{base}"{unit}"' if unit == "%" else f'{base}" {unit}"'


def _scheme(name: str) -> Any:
    colour = OxmlElement("a:schemeClr")
    colour.set("val", name)
    return colour


def _own_colour(sppr: Any) -> Optional[Any]:
    """The colour a series or point paints with: its fill, else its line."""
    if sppr is None:
        return None
    for path in (qn("a:solidFill"), f"{qn('a:ln')}/{qn('a:solidFill')}"):
        fill = sppr.find(path)
        colour = (
            next((c for c in fill if c.tag in _COLOUR_TAGS), None)
            if fill is not None
            else None
        )
        if colour is not None and colour.get("val") != "phClr":
            return colour
    return None


def example_palette(presentation: Any) -> List[Any]:
    """The series (or pie points) colours of the template's own charts, in order,
    then the theme accents they leave out; the accents alone without any."""
    found: List[Any] = []
    seen: set[str] = set()
    for slide in presentation.slides:
        for placed in walk_shapes(slide.shapes):
            shape = cast(Any, placed.shape)
            if not getattr(shape, "has_chart", False):
                continue
            for ser in shape.chart._chartSpace.iter(qn("c:ser")):
                owners = [ser] + ser.findall(qn("c:dPt"))
                for owner in owners:
                    colour = _own_colour(owner.find(qn("c:spPr")))
                    if colour is None:
                        continue
                    own = copy.deepcopy(colour)
                    for alpha in own.findall(qn("a:alpha")):
                        own.remove(alpha)
                    key = str(own.xml)
                    if key not in seen:
                        seen.add(key)
                        found.append(own)
    named = {c.get("val") for c in found if c.tag == qn("a:schemeClr")}
    return found + [_scheme(name) for name in _ACCENTS if name not in named]


def _default_palette() -> Sequence[Any]:
    return [_scheme(name) for name in _ACCENTS]


def _put(fill: Any, colour: Any, alpha: Optional[int]) -> None:
    """A copy of ``colour`` as the only colour of the ``a:solidFill`` ``fill``."""
    for child in list(fill):
        fill.remove(child)
    own = copy.deepcopy(colour)
    if alpha is not None:
        own.append(OxmlElement("a:alpha"))
        own[-1].set("val", str(alpha))
    fill.append(own)


def _fill(fmt: Any, colour: Any, alpha: Optional[int] = None) -> None:
    fmt.fill.solid()
    _put(fmt._element.find(f"{qn('c:spPr')}/{qn('a:solidFill')}"), colour, alpha)


def _line(
    fmt: Any, colour: Any, width: Optional[float], alpha: Optional[int] = None
) -> None:
    fmt.line.fill.solid()
    if width is not None:
        fmt.line.width = Pt(width)
    path = f"{qn('c:spPr')}/{qn('a:ln')}/{qn('a:solidFill')}"
    _put(fmt._element.find(path), colour, alpha)


def _tone(
    index: int, chosen: Optional[int], palette: Sequence[Any], alternate: bool = True
) -> Tuple[Any, Optional[int]]:
    """Colour and opacity of point or series ``index``: its own in the palette,
    or with one ``chosen``, the palette's first for it and muted text for the rest."""
    if chosen is None:
        return palette[index % len(palette)], None
    if index == chosen:
        return palette[0], None
    rank = index - (index > chosen) if alternate else 0
    return _scheme("tx1"), _MUTED[rank % 2]


def _text(font: Any, pt: float, bold: bool = False) -> None:
    font.size = Pt(pt)
    font.bold = bold
    font.color.theme_color = MSO_THEME_COLOR.TEXT_1
    font.color.brightness = 0.25


def _marker(owner: Any, colour: Any, alpha: Optional[int], size: int) -> None:
    owner.marker.style = XL_MARKER_STYLE.CIRCLE
    owner.marker.size = size
    _fill(owner.marker.format, colour, alpha)
    _line(owner.marker.format, colour, None, alpha)


def _colour(
    built: Any,
    chart: PlanChart,
    target: Optional[Tuple[str, int]],
    palette: Sequence[Any],
    painted: Optional[Sequence[int]] = None,
) -> None:
    """Palette colours per series or point. On the template's own chart (``painted``
    given), only the indexes listed and the muted ones of an emphasis are set: the
    emphasised one and the others keep the designer's colour."""
    plot = built.plots[0]
    chosen = target[1] if target else None

    def paints(i: int) -> bool:
        if painted is None:
            return True
        return i in painted if chosen is None else i != chosen

    if chart.type in ROUND:
        plot.vary_by_categories = True
        for i, point in enumerate(plot.series[0].points):
            if not paints(i):
                continue
            _fill(point.format, *_tone(i, chosen, palette))
            # A thin background-coloured edge keeps neighbouring slices apart.
            point.format.line.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
            point.format.line.width = Pt(1.5)
        return
    if target and target[0] == "category":
        # Each bar coloured on its own: LibreOffice lays a series' opacity over its points.
        for i, point in enumerate(plot.series[0].points):
            if not paints(i):
                continue
            if chart.type == "scatter":
                _marker(point, *_tone(i, chosen, palette, False), _SCATTER_MARKER)
            else:
                _fill(point.format, *_tone(i, chosen, palette, False))
        return
    for i, series in enumerate(plot.series):
        if not paints(i):
            continue
        colour, alpha = _tone(i, chosen, palette)
        if chart.type in ("line", "radar"):
            stressed = chosen is None or i == chosen
            width = 2.25 if chosen is None else 3.0 if stressed else 1.5
            _line(series.format, colour, width, alpha)
            if chart.type == "line":
                series.smooth = False
            if stressed:
                _marker(series, colour, alpha, 7)
            else:
                series.marker.style = XL_MARKER_STYLE.NONE
        elif chart.type == "scatter":
            _marker(series, colour, alpha, _SCATTER_MARKER)
        elif chart.type == "bubble":
            # Translucent, so a bubble behind another still shows.
            _fill(series.format, colour, alpha or _BUBBLE_ALPHA)
        else:
            _fill(series.format, colour, alpha)


def _labelled(chart: PlanChart) -> bool:
    """Values on the chart: a pie's slices, or a few single-series bars."""
    return chart.type in ROUND or (
        chart.type in ("column", "bar")
        and len(chart.series) == 1
        and len(chart.categories) <= _LABELLED_BARS
    )


def _labels(
    built: Any, chart: PlanChart, pt: float, fmt: str, target: Optional[Tuple[str, int]]
) -> None:
    plot = built.plots[0]
    point = target[1] if target and target[0] == "category" else None
    if chart.type in XY_TYPES:
        _point_labels(built, chart, pt)
        return
    if not _labelled(chart):
        return
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format_is_linked = False
    _text(labels.font, pt)
    if chart.type in ROUND:
        _round_labels(labels, pie_show(chart), fmt)
        if chart.type == "pie":
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
        elif point is None:
            # A doughnut's labels sit on the ring: the background colour reads on accents.
            labels.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
            labels.font.bold = True
        if point is not None:
            _stress_label(plot, point, light=chart.type == "doughnut")
        return
    labels.number_format = fmt
    labels.show_value = True
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    if point is not None:
        _stress_label(plot, point)


def _point_labels(built: Any, chart: PlanChart, pt: float) -> None:
    """Each point's name beside it, while the chart has few enough points."""
    points = chart.points()
    if len(points) > _LABELLED_POINTS or not any(p.label for p in points):
        return
    for drawn, planned in zip(built.plots[0].series, chart.series):
        for point, plan in zip(drawn.points, planned.points):
            if not plan.label:
                continue
            label = point.data_label
            label.position = XL_LABEL_POSITION.RIGHT
            label.text_frame.text = plan.label
            _text(label.text_frame.paragraphs[0].runs[0].font, pt)


def _stress_label(plot: Any, index: int, light: bool = False) -> None:
    """The emphasised point's label in bold, otherwise as the plot's shared labels
    (a series-level ``c:dLbls`` replaces them, so it repeats them)."""
    ser = plot._element.find(qn("c:ser"))
    shared = plot._element.find(qn("c:dLbls"))
    if shared is None:
        # A template's own chart often formats its labels on the series.
        shared = ser.find(qn("c:dLbls"))
    if shared is None:
        return
    children = [copy.deepcopy(c) for c in shared if c.tag != qn("c:dLbl")]
    own = ser.get_or_add_dLbls()
    for child in list(own):
        own.remove(child)
    label = own.makeelement(qn("c:dLbl"), {})
    label.append(label.makeelement(qn("c:idx"), {"val": str(index)}))
    skipped = (qn("c:showLeaderLines"), qn("c:leaderLines"), qn("c:extLst"))
    for child in children:
        own.append(child)
        if child.tag not in skipped:
            label.append(copy.deepcopy(child))
    own.insert(0, label)
    for run in label.iter(qn("a:defRPr")):
        run.set("b", "1")
        if light:
            # On the doughnut's accent ring, the label reads in the background colour.
            for colour in run.iter(qn("a:schemeClr")):
                colour.set("val", "bg1")
                for child in list(colour):
                    colour.remove(child)


def pie_show(chart: PlanChart) -> str:
    """What a pie's labels show unless asked: data with a unit (already %, M€)
    shows its own values, a share of their sum would misstate them."""
    return chart.show or ("value" if (chart.unit or "").strip() else "share")


def _round_labels(labels: Any, show: str, fmt: str) -> None:
    labels.show_percentage = show != "value"
    labels.show_value = show != "share"
    labels.number_format = "0%" if show == "share" else fmt
    if show == "both":
        # One label format applies to every field: linked, the value keeps the
        # series' format (its unit) and the percentage its own, on the next line.
        labels.number_format_is_linked = True
        _separator(labels._element, "\n")


def _separator(dlbls: Any, text: str) -> None:
    """``c:separator`` in its schema place: python-pptx has no accessor for it."""
    node = dlbls.find(qn("c:separator"))
    if node is None:
        node = dlbls.makeelement(qn("c:separator"), {})
        after = ("c:showLeaderLines", "c:leaderLines", "c:extLst")
        anchor = next((n for t in after for n in dlbls.iterchildren(qn(t))), None)
        if anchor is None:
            dlbls.append(node)
        else:
            anchor.addprevious(node)
    node.text = text


def _axes(built: Any, fmt: str) -> None:
    category, value = built.category_axis, built.value_axis
    value.has_major_gridlines = True
    gridlines = value.major_gridlines.format.line
    gridlines.color.theme_color = MSO_THEME_COLOR.TEXT_1
    gridlines.color.brightness = 0.85
    gridlines.width = Pt(0.75)
    value.format.line.fill.background()
    value.major_tick_mark = XL_TICK_MARK.NONE
    value.tick_labels.number_format = fmt
    value.tick_labels.number_format_is_linked = False
    category.major_tick_mark = XL_TICK_MARK.NONE
    category.format.line.color.theme_color = MSO_THEME_COLOR.TEXT_1
    category.format.line.color.brightness = 0.6


def _xy_axes(built: Any, chart: PlanChart, pt: float) -> None:
    """A scatter's or bubble's two value axes: each its number format and its
    title with the unit (a reused chart keeps its title format)."""
    for axis, name, title, unit in (
        (built.category_axis, "x", chart.x_title, chart.x_unit),
        (built.value_axis, "y", chart.y_title, chart.unit),
    ):
        axis.tick_labels.number_format = number_format(chart, name)
        axis.tick_labels.number_format_is_linked = False
        text = axis_title(title, unit)
        if not text:
            axis.has_title = False
            continue
        frame = axis.axis_title.text_frame if axis.has_title else None
        runs = frame.paragraphs[0].runs if frame is not None else []
        if runs:
            runs[0].text = text
            for extra in runs[1:]:
                extra._r.getparent().remove(extra._r)
            continue
        axis.axis_title.text_frame.text = text
        _text(axis.axis_title.text_frame.paragraphs[0].runs[0].font, pt)


def _hole(built: Any) -> None:
    for node in built._chartSpace.iter(qn("c:holeSize")):
        node.set("val", "55")


def _number(value: Optional[float]) -> float:
    return value if value is not None else 0.0


def _data(chart: PlanChart) -> Any:
    if chart.type in XY_TYPES:
        bubble = chart.type == "bubble"
        fmt = number_format(chart)
        xy = (
            BubbleChartData(number_format=fmt)
            if bubble
            else XyChartData(number_format=fmt)
        )
        for series in chart.series:
            added = xy.add_series(series.name)
            for p in series.points:
                values = [_number(p.x), _number(p.y)]
                added.add_data_point(*values, *([_number(p.size)] if bubble else []))
        return xy
    data = CategoryChartData(number_format=number_format(chart))
    data.categories = chart.categories
    for series in chart.series:
        data.add_series(series.name, [v if v is not None else 0 for v in series.values])
    return data


def _percent(chart: PlanChart) -> bool:
    return chart.type in STACKED_TYPES and chart.stacked_as == "percent"


def _title(built: Any, chart: PlanChart, pt: float) -> None:
    """The plan's caption, in the chart's own title format when it has one."""
    if not chart.title:
        # Untitled, a one-series chart would show its series name as a title.
        built.has_title = False
        return
    first, *rest = chart.title.split("\n")
    runs = built.chart_title.text_frame.paragraphs[0].runs if built.has_title else []
    frame = built.chart_title.text_frame
    if runs:
        runs[0].text = first
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
        for paragraph in frame.paragraphs[1:]:
            paragraph._p.getparent().remove(paragraph._p)
    else:
        built.has_title = True
        frame.text = first
        _text(frame.paragraphs[0].runs[0].font, pt * 1.2, True)
    # A bubble chart's size line: a plain line under the caption.
    for line in rest:
        run = frame.add_paragraph().add_run()
        run.text = line
        _text(run.font, pt)


def _legend(built: Any, chart: PlanChart, pt: float, keep: bool = False) -> None:
    """A legend for several series or a pie's slices (nothing else says which is
    which); with ``keep``, the chart's own legend stays as it is."""
    wanted = chart.type in ROUND or len(chart.series) > 1
    if keep and wanted and built.has_legend:
        return
    built.has_legend = wanted
    if wanted:
        built.legend.position = XL_LEGEND_POSITION.BOTTOM
        built.legend.include_in_layout = False
        _text(built.legend.font, pt)


def example_frame(placed: PlacedShape, chart: PlanChart) -> Any:
    """The template's own chart frame at ``placed`` when it shows the type asked
    and its parts came with the slide, else None (a fresh chart is drawn)."""
    shape = cast(Any, placed.shape)
    if not getattr(shape, "has_chart", False):
        return None
    # A chart part that did not copy left its frame without a link.
    link = shape._element.find(f".//{qn('c:chart')}")
    if link is None or link.get(qn("r:id")) not in shape.part.rels:
        return None
    return shape if type_of(shape.chart) == chart.type else None


def refill(
    frame: Any, zone: Zone, chart: PlanChart, palette: Sequence[Any] = ()
) -> None:
    """The plan's data in the template's own chart. Its colours, fonts and layout
    stay; only what depends on the data is set: legend, data labels, number
    format, the intents, and palette colours for series the example did not have."""
    built = frame.chart
    plot = built.plots[0]
    before = len(plot.series[0].points) if chart.type in ROUND else len(plot.series)
    chart = _arranged(chart)
    fmt = number_format(chart)
    pt = chart_pt(zone)
    built.replace_data(_data(chart))
    if chart.type in STACKED_TYPES:
        # The example may stack to 100 % while the plan does not, or the reverse.
        for grouping in plot._element.iter(qn("c:grouping")):
            grouping.set("val", "percentStacked" if _percent(chart) else "stacked")
    _title(built, chart, pt)
    _legend(built, chart, pt, keep=True)
    count = len(chart.categories) if chart.type in ROUND else len(chart.series)
    target = emphasis(chart)
    _colour(built, chart, target, palette or _default_palette(), range(before, count))
    _relabel(built, chart, pt, fmt, target)
    if chart.type in XY_TYPES:
        _xy_axes(built, chart, pt)
    elif chart.type not in ROUND:
        value = built.value_axis.tick_labels
        value.number_format = "0%" if _percent(chart) else fmt
        value.number_format_is_linked = False


def _relabel(
    built: Any, chart: PlanChart, pt: float, fmt: str, target: Optional[Tuple[str, int]]
) -> None:
    """Data labels of a reused chart: its own label format kept, their presence
    and number format set by the data, as on a fresh chart."""
    plot = built.plots[0]
    sets = [plot._element.find(qn("c:dLbls"))] + [
        ser.find(qn("c:dLbls")) for ser in plot._element.iter(qn("c:ser"))
    ]
    own = [s for s in sets if s is not None]
    if chart.type in XY_TYPES:
        for node in own:
            node.getparent().remove(node)
        _point_labels(built, chart, pt)
        return
    if not _labelled(chart) or not own:
        for node in own:
            node.getparent().remove(node)
        if _labelled(chart):
            _labels(built, chart, pt, fmt, target)
        return
    for node in own:
        labels = DataLabels(node)
        labels.number_format_is_linked = False
        if chart.type in ROUND:
            _round_labels(labels, pie_show(chart), fmt)
        else:
            labels.number_format = fmt
            labels.show_value = True
    if target and target[0] == "category":
        _stress_label(plot, target[1], light=chart.type == "doughnut")


def draw(
    slide: Any,
    host: PlacedShape,
    zone: Zone,
    chart: PlanChart,
    palette: Sequence[Any] = (),
) -> Any:
    """Replace ``host`` (a chart or content placeholder, or a duplicated chart
    frame) with a native chart at its box and drawing-order place."""
    for series in chart.series:
        if chart.type not in XY_TYPES and len(series.values) != len(chart.categories):
            raise ValueError(
                f"chart series '{series.name}' has {len(series.values)} value(s) for "
                f"{len(chart.categories)} categories"
            )
    chart = _arranged(chart)
    data = _data(chart)
    percent = _percent(chart)
    frame = slide.shapes.add_chart(
        _PERCENT_XL[chart.type] if percent else XL_TYPES[chart.type],
        Emu(host.left),
        Emu(host.top),
        Emu(host.width),
        Emu(host.height),
        data,
    )
    old = cast(Any, host.shape)._element
    # In a placeholder the chart keeps its link, as PowerPoint's "insert chart" does.
    placeholder = old.find(f".//{qn('p:nvPr')}/{qn('p:ph')}")
    if placeholder is not None:
        props = frame._element.find(f"{qn('p:nvGraphicFramePr')}/{qn('p:nvPr')}")
        props.insert(0, copy.deepcopy(placeholder))
    tree = old.getparent()
    if tree is not None:
        tree.replace(old, frame._element)

    built = frame.chart
    pt = chart_pt(zone)
    fmt = number_format(chart)
    _text(built.font, pt)
    built.font.name = "+mn-lt"  # the theme's body font, never a literal face
    _title(built, chart, pt)
    _legend(built, chart, pt)
    target = emphasis(chart)
    _colour(built, chart, target, palette or _default_palette())
    _labels(built, chart, pt, fmt, target)
    if chart.type in ROUND:
        _hole(built)
    else:
        _axes(built, "0%" if percent else fmt)
        if chart.type in XY_TYPES:
            _xy_axes(built, chart, pt)
        if chart.type == "bubble":
            built.plots[0].bubble_scale = _BUBBLE_SCALE
        if chart.type in ("column", "bar") + STACKED_TYPES:
            built.plots[0].gap_width = 60 if chart.type in STACKED_TYPES else 80
    return frame
