# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Single-open offline template checks, followed by scoped async folder resolution.

``analyze_template`` owns the presentation in a worker. ``resolve_image_folders``
only receives lightweight schema/errors and performs no deck parsing.
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional, Protocol, Union

from pptx import Presentation

from fred_capability_ppt_filler.parser import (
    KeyField,
    ParseResult,
    SlideSchema,
    TemplateError,
    parse_presentation_with_locations,
)
from fred_capability_ppt_filler.pptx_archive import open_presentation

# Resolution / geometry error codes (RFC "Image support"). The metadata codes live in
# ``parser.py``; these two need an outside lookup or the slide shapes, so they belong to
# this validation layer.
CODE_FOLDER_NOT_FOUND = "folder_not_found"
CODE_IMAGE_KEY_INVALID_LOCATION = "image_key_invalid_location"


class FolderResolver(Protocol):
    """Seam for resolving an author folder string to a DOCUMENT tag id.

    The real implementation (Story 04/05) wraps :class:`KfTagClient` with the request's
    scope (personal vs a specific team) already bound, so :meth:`resolve` takes only the
    folder string. Tests pass an in-memory fake. Resolution is async because the real
    lookup hits Knowledge Flow.
    """

    async def resolve(self, folder: str) -> Optional[str]:
        """Return the DOCUMENT tag id for ``folder`` (e.g. ``"images/flags"``) in the
        current space, or ``None`` if it does not exist."""
        ...


def _folder_not_found_message(key: str, slide_number: int, folder: str) -> str:
    """English fallback message for ``folder_not_found`` (frontend i18n's by code)."""
    return (
        f"{{{{{key}}}}} on slide {slide_number} points to the folder "
        f"'{folder}', which does not exist in this space."
    )


def _image_key_invalid_location_message(slide_number: int, keys: List[str]) -> str:
    """English fallback message for ``image_key_invalid_location`` (grouped per slide).

    Keys are grouped by slide to match how the UI groups errors; the heading tells the
    author how to fix it (move the key into a picture-capable shape).
    """
    placeholders = ", ".join(f"{{{{{key}}}}}" for key in keys)
    return (
        f"On slide {slide_number}, {placeholders} cannot be turned into a picture "
        "because the key sits in a table cell. Move it into a text box or a rectangle "
        "so a picture can replace it."
    )


def analyze_template(pptx_source: Union[bytes, str, Path]) -> ParseResult:
    """Open once for field and image checks; return only lightweight validation data."""
    if isinstance(pptx_source, bytes):
        presentation = open_presentation(pptx_source)
    else:
        presentation = Presentation(str(pptx_source))
    result, invalid_locations = parse_presentation_with_locations(presentation)
    for slide_schema in result.slides:
        keys = [
            field.key
            for field in slide_schema.keys
            if field.type == "image"
            and field.key in invalid_locations.get(slide_schema.slide, set())
        ]
        if keys:
            message = _image_key_invalid_location_message(slide_schema.slide, keys)
            result.errors.extend(
                TemplateError(
                    slide=slide_schema.slide,
                    key=key,
                    code=CODE_IMAGE_KEY_INVALID_LOCATION,
                    message=message,
                )
                for key in keys
            )
    return result


async def resolve_image_folders(
    parse_result: ParseResult,
    resolver: Optional[FolderResolver],
) -> ParseResult:
    """Resolve image folders after offline checks have released the parsed deck.

    Returns a new result; resolver=None keeps the stateless analysis offline.
    Each distinct folder is resolved once and existing errors are preserved.
    """
    # Resolve each DISTINCT non-empty folder once (several keys may share a folder).
    resolved: Dict[str, Optional[str]] = {}

    # Rebuild the schema so the resolved folder_tag_id is reflected in the returned
    # result. KeyField is a frozen-by-convention pydantic model; we mutate via
    # ``model_copy(update=...)`` to produce well-formed copies rather than poking fields.
    new_slides: List[SlideSchema] = []
    new_errors: List[TemplateError] = list(parse_result.errors)

    for slide_schema in parse_result.slides:
        slide_number = slide_schema.slide
        new_keys: List[KeyField] = []

        for key_field in slide_schema.keys:
            if key_field.type != "image":
                new_keys.append(key_field)
                continue

            updated = key_field

            # --- folder resolution (non-empty folders only) ---
            folder = key_field.folder
            if folder and resolver is not None:
                if folder not in resolved:
                    resolved[folder] = await resolver.resolve(folder)
                tag_id = resolved[folder]
                if tag_id is None:
                    new_errors.append(
                        TemplateError(
                            slide=slide_number,
                            key=key_field.key,
                            code=CODE_FOLDER_NOT_FOUND,
                            message=_folder_not_found_message(
                                key_field.key, slide_number, folder
                            ),
                        )
                    )
                else:
                    updated = key_field.model_copy(update={"folder_tag_id": tag_id})

            new_keys.append(updated)

        new_slides.append(SlideSchema(slide=slide_number, keys=new_keys))

    return ParseResult(schema=new_slides, errors=new_errors)
