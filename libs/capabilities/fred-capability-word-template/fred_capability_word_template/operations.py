# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Operations on a ``.docx`` and the operation log.

A batch is a list of operations whose addresses ALL refer to the document as
it was before the batch: every address is resolved to its XML element first,
the whole batch is validated (every problem listed, nothing applied), then
the operations apply in order. ``replay`` applies an operation log - a list
of batches - to a template, each batch against the result of the previous.
"""

from __future__ import annotations

import copy
import io
import re
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from typing import Annotated, Any, Literal, cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run
from pydantic import (
    AliasChoices,
    BaseModel,
    ConfigDict,
    Field,
    TypeAdapter,
    ValidationError,
)

from fred_capability_word_template.docmap import (
    DocMap,
    MapElement,
    block_children,
    build_map_from_document,
    is_toc,
    own_runs,
    own_text_units,
    word_numbered,
)
from fred_capability_word_template.inline import parse_inline_markdown
from fred_capability_word_template.render import (
    TODO_MARKER_RE,
    ListLook,
    PictureTarget,
    ensure_required_styles,
    markdown_features,
    render_blocks,
    renumber_figure_captions,
    split_by_ranges,
    template_default_font,
    usable_width,
)
from fred_capability_word_template.toc import refresh_table_of_contents

MAX_REPEAT = 50
_BULLET_LINE_RE = re.compile(r"^\s*[-*+]\s+", re.MULTILINE)
_NUMBERED_LINE_RE = re.compile(r"^\s*\d+[.)]\s+", re.MULTILINE)

# --- operation models ----------------------------------------------------------


class _Op(BaseModel):
    model_config = ConfigDict(populate_by_name=True, extra="forbid")


class SetText(_Op):
    """Replace a paragraph's or a cell's text, keeping its formatting."""

    op: Literal["set_text"] = "set_text"
    address: str
    text: str


class ReplaceText(_Op):
    """Replace every occurrence of ``find`` inside a paragraph, a cell, a
    header/footer or a text box by ``replace``, keeping the runs' formatting
    and leaving fields (page numbers, dates...) untouched."""

    op: Literal["replace_text"] = "replace_text"
    address: str
    find: str
    replace: str


class Write(_Op):
    """Replace a range of blocks (or a cell's content) with rendered Markdown."""

    op: Literal["write"] = "write"
    from_: str = Field(
        validation_alias=AliasChoices("from", "from_"), serialization_alias="from"
    )
    to: str | None = None
    markdown: str


class InsertAfter(_Op):
    """Insert rendered Markdown right after a paragraph or a table."""

    op: Literal["insert_after"] = "insert_after"
    address: str
    markdown: str


class Repeat(_Op):
    """Clone a range of blocks or table rows ``count`` more times, right after it."""

    op: Literal["repeat"] = "repeat"
    from_: str = Field(
        validation_alias=AliasChoices("from", "from_"), serialization_alias="from"
    )
    to: str | None = None
    count: int


class Remove(_Op):
    """Remove a range of blocks or table rows."""

    op: Literal["remove"] = "remove"
    from_: str = Field(
        validation_alias=AliasChoices("from", "from_"), serialization_alias="from"
    )
    to: str | None = None


class Move(_Op):
    """Move a range of body blocks before or after another block. With a
    heading as ``from`` and no ``to``, its whole section moves; ``after`` a
    heading means after its whole section."""

    op: Literal["move"] = "move"
    from_: str = Field(
        validation_alias=AliasChoices("from", "from_"), serialization_alias="from"
    )
    to: str | None = None
    before: str | None = None
    after: str | None = None


Operation = Annotated[
    SetText | ReplaceText | Write | InsertAfter | Repeat | Remove | Move,
    Field(discriminator="op"),
]
_OPERATION_ADAPTER: TypeAdapter[Operation] = TypeAdapter(Operation)


class BatchError(ValueError):
    """A batch was rejected; ``problems`` lists every reason, nothing was applied."""

    def __init__(self, problems: Sequence[str]) -> None:
        # Once each: one operation can hit the same problem twice (e.g. a
        # range whose "from" and "to" are the same unknown address).
        problems = list(dict.fromkeys(problems))
        self.problems = problems
        super().__init__(
            "The batch was rejected, nothing was applied:\n- " + "\n- ".join(problems)
        )


_BLOCK_LINE_RE = re.compile(r"^\s*(?:[-*+]|\d+[.)]|#{1,6})\s+")
_BLANK_LINE_RE = re.compile(r"\n[ \t]*\n")


_PLACE_RE = re.compile(r"(?:^|\.)[PC]\d+$")
_HEADING_OR_RULE_RE = re.compile(r"^\s*(?:#{1,6}\s|(?:-{3,}|\*{3,}|_{3,})\s*$)")
# A heading line wrapped in bold ("**## Context**") is meant as that heading.
_BOLD_HEADING_RE = re.compile(
    r"^([ \t]*)\*\*[ \t]*(#{1,6})[ \t]+(.+?)[ \t]*\*\*[ \t]*:?[ \t]*$", re.MULTILINE
)


def normalized(raw: Any) -> Any:
    """A model operation in the one form the engine and every guard read:
    canonical keys ("from", never "from_"), a ``remove`` of one ``address``
    as a ``remove`` from it, and a ``set_text`` of a paragraph or cell whose
    text holds several paragraphs or list items (or one bullet) as the
    ``write`` of that place (``set_text`` renders one paragraph), as is a
    single bullet, heading or horizontal rule line. Invalid operations as is,
    for the engine to refuse."""
    if not isinstance(raw, Mapping):
        return raw
    op = dict(raw)
    if op.get("op") == "remove" and "from" not in op and "from_" not in op:
        if "address" in op:
            op["from"] = op.pop("address")
    for key in ("text", "markdown"):
        if isinstance(op.get(key), str):
            op[key] = _BOLD_HEADING_RE.sub(r"\1\2 \3", op[key])
    text, address = op.get("text"), op.get("address")
    if (
        op.get("op") == "set_text"
        and isinstance(text, str)
        and isinstance(address, str)
        and _PLACE_RE.search(address.strip())
    ):
        lines = [line for line in text.splitlines() if line.strip()]
        blocks = _BLANK_LINE_RE.search(text) or any(
            _BLOCK_LINE_RE.match(line) for line in lines
        )
        one_block = len(lines) == 1 and (
            _BULLET_LINE_RE.match(lines[0]) or _HEADING_OR_RULE_RE.match(lines[0])
        )
        if (len(lines) >= 2 and blocks) or one_block:
            op = {"op": "write", "from": address, "markdown": text}
    try:
        return parse_operation(op).model_dump(by_alias=True, exclude_none=True)
    except ValidationError:
        return op


_HEADING_LINE_RE = re.compile(r"^(#{1,6})\s+(.*)$")
# A number typed before a heading: "4.", "3.1.4", "7", "IV.".
_TYPED_NUMBER_RE = re.compile(r"^(?:\d+(?:[.)]\d+)*[.)]?|[IVXLC]+[.)])\s+")
_MARKDOWN_RE = re.compile(r"\*\*|\n\s*\n|^\s*(?:[-*+]|\d+[.)]|#{1,6})\s", re.MULTILINE)


def _plain(text: str) -> str:
    text = _TYPED_NUMBER_RE.sub("", text.strip().lstrip("#").strip(" *_"))
    return " ".join(text.rstrip(" :.").split()).casefold()


def _numbered(paragraph, document: DocxDocument) -> bool:
    """Whether Word numbers this paragraph itself: a number typed in its
    text would show twice."""
    return word_numbered(paragraph, document.styles)


def fit_to_document(operations: Sequence[Any], docx_bytes: bytes) -> list[Any]:
    """A model batch adjusted to the document before it is stored (so a
    replay rebuilds the same document): a number typed before a heading Word
    numbers is dropped, a written text does not repeat the heading above it,
    a heading line written over a heading or right under one of its level
    titles that heading, and a ``replace_text`` of a whole paragraph with
    Markdown is a ``set_text``."""
    document = Document(io.BytesIO(docx_bytes))
    doc_map = build_map_from_document(document)
    order = {e.address: i for i, e in enumerate(doc_map.elements)}
    numbered_levels = {
        e.heading_level
        for e in doc_map.elements
        if e.heading_level is not None and _numbered(e.node, document)
    }

    def untyped_headings(markdown: str) -> str:
        lines = []
        for line in markdown.split("\n"):
            match = _HEADING_LINE_RE.match(line)
            if match and len(match.group(1)) in numbered_levels:
                title = match.group(2).strip().strip("*_ ").strip()
                line = f"{match.group(1)} {_TYPED_NUMBER_RE.sub('', title)}"
            lines.append(line)
        return "\n".join(lines)

    def heading_above(address: str) -> MapElement | None:
        index = order.get(address)
        if index is None:
            return None
        return next(
            (
                e
                for e in reversed(doc_map.elements[:index])
                if e.heading_level is not None and e.container == ""
            ),
            None,
        )

    def right_under(heading: MapElement, address: str) -> bool:
        """Nothing but empty paragraphs between ``heading`` and ``address``."""
        i, j = order.get(heading.address), order.get(address)
        if i is None or j is None or j <= i:
            return False
        return all(
            e.kind == "paragraph" and not e.text.strip()
            for e in doc_map.elements[i + 1 : j]
        )

    def title_of(line: str, level: int) -> str | None:
        match = _HEADING_LINE_RE.match(line.strip())
        if match is None or len(match.group(1)) != level:
            return None
        return match.group(2).strip() or None

    top_level = min(
        (e.heading_level for e in doc_map.elements if e.heading_level is not None),
        default=0,
    )
    targeted = {
        str(op.get(key))
        for op in operations
        if isinstance(op, Mapping)
        for key in ("address", "from", "to")
        if op.get(key)
    }
    titled = {
        str(op.get("address"))
        for op in operations
        if isinstance(op, Mapping) and op.get("op") == "set_text"
    }

    out: list[Any] = []
    for op in operations:
        if not isinstance(op, Mapping):
            out.append(op)
            continue
        op = dict(op)
        kind = op.get("op")
        if kind == "replace_text" and isinstance(op.get("replace"), str):
            target = doc_map.get(str(op.get("address", "")))
            whole = (
                target is not None
                and target.text.strip() == str(op.get("find", "")).strip()
            )
            if whole and _MARKDOWN_RE.search(op["replace"]):
                op = normalized(
                    {"op": "set_text", "address": op["address"], "text": op["replace"]}
                )
                kind = op.get("op")
        if kind == "set_text" and isinstance(op.get("text"), str):
            target = doc_map.get(str(op.get("address", "")))
            if target is not None and target.heading_level is not None:
                if _numbered(target.node, document):
                    op["text"] = _TYPED_NUMBER_RE.sub("", op["text"].lstrip())
        if kind in ("write", "insert_after") and isinstance(op.get("markdown"), str):
            markdown = untyped_headings(op["markdown"])
            start = str(op.get("from") or op.get("address"))
            heading = heading_above(start) if kind == "write" else None
            lines = markdown.split("\n")
            first = next((i for i, line in enumerate(lines) if line.strip()), None)
            filled = [line for line in lines if line.strip()]
            target = doc_map.get(start) if kind == "write" else None
            if (
                target is not None
                and target.heading_level is not None
                and op.get("to") in (None, start)
                and len(filled) == 1
                and (heading_line := _HEADING_LINE_RE.match(filled[0].strip()))
                and (title := heading_line.group(2).strip())
            ):
                # "## Title" sent over a heading, whatever its "#" count: that
                # heading's new title (a heading keeps its level).
                out.append({"op": "set_text", "address": start, "text": title})
                continue
            if (
                heading is not None
                and first is not None
                and len(filled) > 1
                and _plain(lines[first]) == _plain(heading.text)
            ):
                markdown = "\n".join(lines[first + 1 :]).lstrip("\n")
            elif (
                heading is not None
                and first is not None
                and len(filled) > 1
                and (heading.heading_level or 0) > top_level
                and heading.address not in targeted
                and right_under(heading, start)
                and (title := title_of(lines[first], heading.heading_level or 0))
            ):
                # The sub-section's title goes into the template heading above
                # it, not into a second heading of the same level.
                out.append(
                    {"op": "set_text", "address": heading.address, "text": title}
                )
                markdown = "\n".join(lines[first + 1 :]).lstrip("\n")
            elif (
                heading is not None
                and first is not None
                and len(filled) > 1
                and heading.address in titled
                and right_under(heading, start)
                and title_of(lines[first], heading.heading_level or 0)
            ):
                # The batch already titles that heading: no second title.
                markdown = "\n".join(lines[first + 1 :]).lstrip("\n")
            op["markdown"] = markdown
        out.append(op)
    return out


# A title typed with its section number as a plain paragraph: "3.2 Scope".
TYPED_TITLE_RE = re.compile(r"^\s*\d+(?:\.\d+)+\.?\s+[^\n]{1,80}?(?<![.!?;])\s*$")


@dataclass(frozen=True)
class SetTextFindings:
    renumbered: list[str]  # numbered headings only given a typed number
    typed_titles: list[str]  # plain paragraphs holding a typed numbered title
    unchanged: list[str]  # targets already holding the text sent


def inspect_set_texts(operations: Sequence[Any], docx_bytes: bytes) -> SetTextFindings:
    """What a model batch's set_text operations would really do: a typed
    number on a heading Word numbers is dropped (its place gives its
    number), a paragraph titled by hand ("3.2 Scope") is a sub-heading
    without its style, and a text already in place changes nothing."""
    found = SetTextFindings([], [], [])
    candidates = [
        op
        for op in operations
        if isinstance(op, Mapping)
        and op.get("op") == "set_text"
        and isinstance(op.get("text"), str)
    ]
    if not candidates:
        return found
    document = Document(io.BytesIO(docx_bytes))
    doc_map = build_map_from_document(document)
    for op in candidates:
        target = doc_map.get(str(op.get("address", "")))
        if target is None or target.kind != "paragraph":
            continue
        text = op["text"].strip()
        if text == target.text.strip():
            found.unchanged.append(target.address)
        elif target.heading_level is not None:
            if (
                _TYPED_NUMBER_RE.match(text)
                and _numbered(target.node, document)
                and _plain(text) == _plain(target.text)
            ):
                found.renumbered.append(target.address)
        elif (
            target.story == "body"
            and target.container == ""
            and TYPED_TITLE_RE.match(target.text)
            and "\n" not in text
        ):
            found.typed_titles.append(target.address)
    return found


def parse_operation(raw: Operation | Mapping[str, Any]) -> Operation:
    if isinstance(
        raw, SetText | ReplaceText | Write | InsertAfter | Repeat | Remove | Move
    ):
        return raw
    return _OPERATION_ADAPTER.validate_python(raw)


# --- address resolution --------------------------------------------------------

_SEGMENT_RE = re.compile(r"^(?:[HF]\d+[fe]?|[PTXRC]\d+)$")
_KIND_NAMES = {
    "P": "paragraphs",
    "T": "tables",
    "R": "rows",
    "C": "cells",
    "X": "text boxes",
}


def _describe_children(doc_map: DocMap, container: str) -> str:
    children = doc_map.children(container)
    groups: dict[str, list[str]] = {}
    for child in children:
        local = child.address.rsplit(".", 1)[-1]
        key = local[0]
        groups.setdefault(key, []).append(local)
    if not groups:
        return f"{container or 'the body'} has no addressable element"
    described = []
    for key, locals_ in groups.items():
        name = _KIND_NAMES.get(key, key)
        if key in "HF":
            described.append(f"{name} {', '.join(locals_)}")
        elif len(locals_) > 3:
            described.append(f"{name} {locals_[0]}–{locals_[-1]}")
        else:
            described.append(f"{name} {', '.join(locals_)}")
    return f"{container or 'the document'} has {'; '.join(described)}"


def _unknown(doc_map: DocMap, address: str) -> str:
    segments = address.split(".")
    if not all(_SEGMENT_RE.match(s) for s in segments):
        return (
            f"'{address}' is not an address (expected e.g. P12, T3, T3.R2.C4, "
            "T3.R2.C4.P1, H1.P2, F1.P1, X1.P1)"
        )
    prefix = ""
    for index in range(len(segments) - 1, 0, -1):
        candidate = ".".join(segments[:index])
        if candidate in doc_map.by_address:
            prefix = candidate
            break
    return f"'{address}' does not exist ({_describe_children(doc_map, prefix)})"


@dataclass
class _Plan:
    index: int
    op: Operation
    targets: list[MapElement] = field(default_factory=list)  # resolved range / target
    consumed: list[Any] = field(default_factory=list)  # nodes replaced or removed
    consumed_inside: list[Any] = field(
        default_factory=list
    )  # nodes whose content is replaced
    touched: list[Any] = field(default_factory=list)  # nodes the op addresses
    destination: tuple[str, Any] | None = None  # a move: ("before"|"after", node)


def _range(
    doc_map: DocMap, start: str, end: str | None, problems: list[str], label: str
) -> list[MapElement] | None:
    first = doc_map.get(start)
    last = doc_map.get(end) if end else first
    if first is None:
        problems.append(f"{label}: {_unknown(doc_map, start)}")
    if end and last is None:
        problems.append(f"{label}: {_unknown(doc_map, end)}")
    if first is None or last is None:
        return None
    family = "row" if first.kind == "row" else "block" if first.is_block else None
    if (
        family is None
        or (last.kind == "row") != (family == "row")
        or not (last.is_block or last.kind == "row")
    ):
        problems.append(
            f"{label}: a range goes from a paragraph or table to another paragraph or "
            f"table, or from a row to another row of the same table; '{first.address}' "
            f"is a {first.kind} and '{last.address}' is a {last.kind}"
        )
        return None
    if first.container != last.container:
        problems.append(
            f"{label}: '{first.address}' and '{last.address}' are not in the same place "
            f"({_where(first)} vs {_where(last)}); a range must stay in one container"
        )
        return None
    if first.position > last.position:
        problems.append(
            f"{label}: '{first.address}' comes after '{last.address}'; give the range "
            "in reading order"
        )
        return None
    siblings = [
        e
        for e in doc_map.children(first.container)
        if (e.kind == "row") == (family == "row") and (e.is_block or e.kind == "row")
    ]
    return [e for e in siblings if first.position <= e.position <= last.position]


def _where(element: MapElement) -> str:
    return f"in {element.container}" if element.container else "in the body"


def _no_tables_here(element: MapElement) -> str | None:
    """Why a written table cannot go where ``element`` is, if it cannot."""
    if element.kind == "cell" or ".C" in element.address:
        return "inside a table cell"
    if element.story != "body":
        return f"in a {element.story}"
    if element.in_textbox:
        return "inside a text box"
    return None


def _check_markdown(
    markdown: str,
    place: MapElement,
    images: Mapping[str, bytes],
    problems: list[str],
    label: str,
) -> None:
    features = markdown_features(markdown)
    where = _no_tables_here(place)
    if features.has_table and where:
        problems.append(
            f"{label}: Markdown tables cannot be written {where} ('{place.address}'); "
            "use paragraphs or lists there"
        )
    missing = [i for i in features.image_ids if i not in images]
    if missing:
        problems.append(
            f"{label}: image(s) {', '.join(repr(i) for i in missing)} were not provided"
        )


def _plan(
    doc_map: DocMap,
    index: int,
    op: Operation,
    images: Mapping[str, bytes],
    problems: list[str],
) -> _Plan:
    label = f"operation {index} ({op.op})"
    plan = _Plan(index=index, op=op)
    if isinstance(op, Move):
        return _plan_move(doc_map, plan, op, problems, label)
    if isinstance(op, SetText):
        target = doc_map.get(op.address)
        if target is None:
            problems.append(f"{label}: {_unknown(doc_map, op.address)}")
        elif target.kind not in ("paragraph", "cell"):
            hint = (
                f"; address one of its cells, e.g. {target.address}.R1.C1"
                if target.kind == "table"
                else ""
            )
            problems.append(
                f"{label}: set_text applies to a paragraph or a cell; '{target.address}' "
                f"is a {target.kind}{hint}"
            )
        else:
            plan.targets = [target]
            plan.touched = [target.node]
            if target.kind == "cell":
                plan.consumed_inside = [target.node]
            else:
                plan.consumed = [target.node]
        return plan
    if isinstance(op, ReplaceText):
        target = doc_map.get(op.address)
        if target is None:
            problems.append(f"{label}: {_unknown(doc_map, op.address)}")
        elif target.kind in ("table", "row"):
            problems.append(
                f"{label}: replace_text applies to a paragraph, a cell, a header, a "
                f"footer or a text box; '{target.address}' is a {target.kind}"
            )
        elif not op.find:
            problems.append(f"{label}: 'find' must not be empty")
        else:
            found = 0
            for paragraph in _paragraphs_of(target.node):
                matches, blocked = _find_in_paragraph(paragraph, op.find)
                found += len(matches)
                if blocked:
                    problems.append(
                        f"{label}: '{op.find}' is, at least in part, text Word "
                        f"computes in '{target.address}' (a field: page number, date, "
                        "file name...). Word updates it itself: leave it, and replace "
                        "only the text around it, if any"
                    )
                    break
            else:
                if not found:
                    text = " ⏎ ".join(
                        t for p in _paragraphs_of(target.node) if (t := _units_text(p))
                    )
                    problems.append(
                        f"{label}: '{op.find}' is not in '{target.address}' (its text: "
                        f'"{text[:300]}"); copy the exact text from the map'
                    )
            plan.targets = [target]
            plan.touched = [target.node]
        return plan
    if isinstance(op, InsertAfter):
        target = doc_map.get(op.address)
        if target is None:
            problems.append(f"{label}: {_unknown(doc_map, op.address)}")
        elif not target.is_block:
            hint = (
                " (to fill a cell, use write on the cell)"
                if target.kind == "cell"
                else ""
            )
            problems.append(
                f"{label}: insert_after needs a paragraph or a table; '{target.address}' is a "
                f"{target.kind}{hint}"
            )
        else:
            _check_markdown(op.markdown, target, images, problems, label)
            plan.targets = [target]
            plan.touched = [target.node]
        return plan
    assert isinstance(op, Write | Repeat | Remove)
    start = doc_map.get(op.from_)
    if isinstance(op, Write) and start is not None and start.kind == "cell":
        if op.to not in (None, op.from_):
            problems.append(
                f"{label}: writing into a cell takes the cell alone; use from='{op.from_}' "
                "with no 'to' (or to equal to from)"
            )
            return plan
        _check_markdown(op.markdown, start, images, problems, label)
        plan.targets = [start]
        plan.touched = [start.node]
        plan.consumed_inside = [start.node]
        return plan
    elements = _range(doc_map, op.from_, op.to, problems, label)
    if elements is None:
        return plan
    nodes = [e.node for e in elements]
    plan.targets = elements
    plan.touched = nodes
    if isinstance(op, Write):
        if elements[0].kind == "row":
            problems.append(
                f"{label}: write replaces paragraphs and tables, not table rows; write into "
                "the cells (e.g. 'T1.R2.C1') or use repeat/remove on rows"
            )
            return plan
        _check_markdown(op.markdown, elements[0], images, problems, label)
        plan.consumed = nodes
    elif isinstance(op, Repeat):
        if not 1 <= op.count <= MAX_REPEAT:
            problems.append(
                f"{label}: count must be between 1 and {MAX_REPEAT} (the number of EXTRA "
                f"copies); got {op.count}"
            )
    else:
        if elements[0].kind == "row":
            table = doc_map.get(elements[0].container)
            all_rows = table is not None and len(
                doc_map.children(table.address)
            ) == len(elements)
            if all_rows:
                problems.append(
                    f"{label}: this removes every row of {elements[0].container}; remove "
                    f"the table '{elements[0].container}' instead"
                )
        plan.consumed = nodes
    return plan


def _plan_move(
    doc_map: DocMap, plan: _Plan, op: Move, problems: list[str], label: str
) -> _Plan:
    """A move works on top-level body blocks after the cover page, never on
    the table of contents or a section break (Word's page setup)."""
    blocks = [
        e
        for e in doc_map.elements
        if e.container == ""
        and e.story == "body"
        and not e.in_textbox
        and e.kind in ("paragraph", "table")
    ]
    index = {e.address: i for i, e in enumerate(blocks)}

    def section(i: int) -> int:
        """The last block of the section headed by ``blocks[i]``."""
        level = blocks[i].heading_level
        if level is None:
            return i
        end = next(
            (
                j
                for j in range(i + 1, len(blocks))
                if (other := blocks[j].heading_level) is not None and other <= level
            ),
            len(blocks),
        )
        return end - 1

    if (op.before is None) == (op.after is None):
        problems.append(f"{label}: give exactly one of 'before' or 'after'")
        return plan
    first = index.get(op.from_)
    last = (
        section(first)
        if first is not None and op.to is None
        else index.get(op.to or "")
    )
    anchor = index.get(op.before or op.after or "")
    for address, found in (
        (op.from_, first),
        (op.to, last),
        (op.before or op.after, anchor),
    ):
        if address and found is None:
            problems.append(
                f"{label}: '{address}' is not a paragraph, table or heading of the "
                "body (a move works on whole blocks after the cover page)"
            )
    if first is None or last is None or anchor is None:
        return plan
    if last < first:
        first, last = last, first
    moved = blocks[first : last + 1]
    if first <= anchor <= last:
        problems.append(f"{label}: the destination is inside the moved range")
        return plan
    if any(
        is_toc(e) or e.node.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
        for e in moved
    ):
        problems.append(
            f"{label}: the range holds the table of contents or a section break; "
            "move only chapters, sub-sections or paragraphs"
        )
        return plan
    at = section(anchor) if op.after is not None else anchor
    if (op.after is not None and at == first - 1) or (
        op.before is not None and at == last + 1
    ):
        problems.append(
            f"{label}: this range is already there; the document would not change"
        )
        return plan
    target = blocks[at]
    plan.targets = moved
    plan.touched = [e.node for e in moved]
    plan.consumed = [e.node for e in moved]
    plan.destination = ("after" if op.after is not None else "before", target.node)
    return plan


def _inside(node, containers: Iterable[Any], *, strict: bool) -> bool:
    pool = {id(c) for c in containers}
    if not pool:
        return False
    if not strict and id(node) in pool:
        return True
    return any(id(a) in pool for a in node.iterancestors())


def _conflicts(plans: Sequence[_Plan], problems: list[str]) -> None:
    """Two operations conflict when one addresses content another replaces
    or removes (a node consumed, or inside a cell whose content is
    replaced, or the same cell filled twice)."""
    reported: set[tuple[int, int]] = set()
    for plan in plans:
        for other in plans:
            if other is plan:
                continue
            hit = any(
                _inside(node, plan.consumed, strict=False)
                or _inside(node, plan.consumed_inside, strict=True)
                or any(node is c for c in plan.consumed_inside)
                for node in other.touched
            )
            pair = (min(plan.index, other.index), max(plan.index, other.index))
            if hit and pair not in reported:
                reported.add(pair)
                problems.append(
                    f"operations {pair[0]} and {pair[1]} overlap: one replaces or removes "
                    "content the other addresses. Every address refers to the document "
                    "before the batch, so merge them into one operation or send them in "
                    "two batches"
                )


# --- template blocks kept (model batches) --------------------------------------
#
# The model changes how many copies of a template block there are, it never
# removes every copy to rebuild it in its own layout. The collaborator is free.

TableSignature = tuple[Any, ...]

_TABLE_ROWS_XPATH = "./w:tr | ./w:sdt/w:sdtContent/w:tr"
_ROW_CELLS_XPATH = "./w:tc | ./w:sdt/w:sdtContent/w:tc"
_KEEP_BLOCK = (
    "keep at least one copy of this block. To change how many there are, repeat "
    "it or remove the extra copies, then fill each copy; if a previous edit went "
    "wrong, undo it"
)


def _cell_key(cell) -> str:
    text = "".join(t.text or "" for t in cell.iter(qn("w:t")))
    return " ".join(text.split()).casefold()


def table_signature(table) -> TableSignature:
    """A table's structure, independent of what was filled in: with merged
    cells, its column count and merged-cell layout (consecutive rows of the
    same layout counted once, so added rows keep it); a plain grid, its
    column count and first cell's text (normalized): a label, where the
    rest of the first row may be a value being filled."""
    rows = table.xpath(_TABLE_ROWS_XPATH)
    layout: list[tuple[tuple[int, str | None], ...]] = []
    merged = False
    widest = 0
    for row in rows:
        pattern: list[tuple[int, str | None]] = []
        for cell in row.xpath(_ROW_CELLS_XPATH):
            span_el = cell.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}")
            vmerge_el = cell.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}")
            try:
                span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            except (TypeError, ValueError):
                span = 1
            vmerge = None
            if vmerge_el is not None:
                vmerge = (
                    "restart" if vmerge_el.get(qn("w:val")) == "restart" else "continue"
                )
            merged = merged or span > 1 or vmerge is not None
            pattern.append((span, vmerge))
        widest = max(widest, sum(span for span, _ in pattern))
        if not layout or layout[-1] != tuple(pattern):
            layout.append(tuple(pattern))
    grid = table.findall(f"{qn('w:tblGrid')}/{qn('w:gridCol')}")
    cols = len(grid) or widest
    if merged:
        return ("merged", cols, tuple(layout))
    first = (
        tuple(_cell_key(c) for c in rows[0].xpath(_ROW_CELLS_XPATH)[:1]) if rows else ()
    )
    return ("grid", cols, first)


def _body_tables(document: DocxDocument) -> list[Any]:
    return list(document.element.body.iter(qn("w:tbl")))


def template_heading_texts(docx_bytes: bytes) -> frozenset[str]:
    """The base document's headings, as compared text: the ones a model
    ``write`` must keep (a heading the model added may be rewritten)."""
    return frozenset(
        _plain(e.text)
        for e in build_map_from_document(Document(io.BytesIO(docx_bytes))).elements
        if e.heading_level is not None
    )


def chapters_with_template_sub_headings(
    docx_bytes: bytes, template_headings: frozenset[str]
) -> frozenset[str]:
    """The top-level body blocks of the chapters holding one of the
    template's sub-headings: their outline is set, unlike a chapter the
    template left without sub-headings."""
    top = [
        e
        for e in build_map_from_document(Document(io.BytesIO(docx_bytes))).elements
        if e.container == "" and e.story == "body" and e.kind in ("paragraph", "table")
    ]
    chapter_level = min(
        (e.heading_level for e in top if e.heading_level is not None), default=0
    )
    chapters: list[list[MapElement]] = []
    for element in top:
        if element.heading_level == chapter_level or not chapters:
            chapters.append([])
        chapters[-1].append(element)
    return frozenset(
        e.address
        for blocks in chapters
        if any(
            b.heading_level not in (None, chapter_level)
            and _plain(b.text) in template_headings
            for b in blocks
        )
        for e in blocks
    )


def template_sub_heading_levels(docx_bytes: bytes) -> frozenset[int]:
    """The heading levels the base uses below its chapters (its top level):
    the model may add sub-headings of these levels on its own."""
    levels = {
        e.heading_level
        for e in build_map_from_document(Document(io.BytesIO(docx_bytes))).elements
        if e.heading_level is not None
    }
    return frozenset(level for level in levels if level > min(levels, default=0))


def template_table_signatures(docx_bytes: bytes) -> frozenset[TableSignature]:
    """The signatures of the base document's body tables: the blocks the
    model's batches must keep at least one copy of. A one-cell table is a
    frame (an image's, a boxed note's), not a block to fill."""
    document = Document(io.BytesIO(docx_bytes))
    return frozenset(
        table_signature(t)
        for t in _body_tables(document)
        if sum(len(row.xpath(_ROW_CELLS_XPATH)) for row in t.xpath(_TABLE_ROWS_XPATH))
        > 1
    )


def _removed_tables(plan: _Plan) -> list[Any]:
    if not isinstance(plan.op, Remove | Write | SetText):
        return []
    # Removing a whole chapter (its heading included) is a content decision,
    # not a block being rebuilt.
    if any(t.heading_level is not None for t in plan.targets):
        return []
    found: list[Any] = []
    # A rewritten cell (consumed_inside) loses the tables nested in it too.
    for node in [*plan.consumed, *plan.consumed_inside]:
        found.extend(node.iter(qn("w:tbl")))
    return found


_FIELD_TAGS = (qn("w:fldChar"), qn("w:instrText"), qn("w:fldSimple"))
_TOC_KEPT = (
    "this is the table of contents: Word builds it from the headings and "
    "updates it itself (right-click > Update field); leave it as it is"
)
_FIELD_KEPT = (
    "this holds a field Word computes (page number, date, file name...); "
    "keep it and change only the text around it with replace_text"
)


def _toc_titles(doc_map: DocMap) -> set[int]:
    """The paragraphs titling a table of contents ("Contents"): a styled
    paragraph right above its first line (only empty paragraphs between)."""
    titles: set[int] = set()
    body = [e for e in doc_map.elements if e.kind == "paragraph" and e.container == ""]
    for i, e in enumerate(body):
        if not is_toc(e) or (i and is_toc(body[i - 1])):
            continue
        before = next((p for p in reversed(body[:i]) if p.text.strip()), None)
        styled = before is not None and (before.style or "normal").lower() not in (
            "normal",
            "default",
        )
        if styled and before is not None and before.heading_level is None:
            titles.add(id(before.node))
    return titles


def _check_computed_kept(
    doc_map: DocMap, plans: Sequence[_Plan], problems: list[str]
) -> None:
    """Reject a model operation that rewrites or removes what Word computes:
    the table of contents (its title included), or a paragraph holding a
    field."""
    by_node = {id(e.node): e for e in doc_map.elements if e.kind == "paragraph"}
    toc_titles = _toc_titles(doc_map)
    for plan in plans:
        if not isinstance(plan.op, SetText | Write | Remove):
            continue
        paragraphs = [
            p
            for node in [*plan.consumed, *plan.consumed_inside]
            for p in ([node] if node.tag == qn("w:p") else node.iter(qn("w:p")))
        ]
        label = f"operation {plan.index} ({plan.op.op})"
        elements = [by_node.get(id(p)) for p in paragraphs]
        if any(e is not None and is_toc(e) for e in elements) or (
            not isinstance(plan.op, Remove)
            and any(id(p) in toc_titles for p in paragraphs)
        ):
            problems.append(f"{label}: {_TOC_KEPT}")
        elif not isinstance(plan.op, Remove) and any(
            any(True for _ in p.iter(*_FIELD_TAGS)) for p in paragraphs
        ):
            problems.append(f"{label}: {_FIELD_KEPT}")


def _check_places(doc_map: DocMap, plans: Sequence[_Plan], problems: list[str]) -> None:
    """Reject a model operation that puts text where the template has no
    place for it: a header or footer Word never shows, a paragraph holding
    only an image (a logo), or another paragraph's text moved into one (each
    paragraph's style says what it holds: a date, a title...)."""
    rewritten: dict[str, tuple[int, str]] = {}
    for plan in plans:
        label = f"operation {plan.index} ({plan.op.op})"
        root = plan.targets[0].address.split(".")[0] if plan.targets else ""
        story = doc_map.get(root)
        if story is not None and "hidden" in story.flags:
            shown = re.sub(r"[fe]$", "", root)
            problems.append(
                f"{label}: {root} is a {story.kind} Word never shows in this "
                f"document; the one shown there is {shown}"
            )
            continue
        text = (
            plan.op.text
            if isinstance(plan.op, SetText)
            else plan.op.markdown
            if isinstance(plan.op, Write)
            else None
        )
        if text is None or len(plan.targets) != 1:
            continue
        target = plan.targets[0]
        if target.kind != "paragraph":
            continue
        if text.strip() and not target.text.strip() and "image" in target.flags:
            problems.append(
                f"{label}: {target.address} holds only an image (a logo); it is "
                "not a place for text"
            )
        if isinstance(plan.op, SetText):
            rewritten[target.address] = (plan.index, text.strip())
    for address, (index, text) in rewritten.items():
        moved_from = next(
            (
                other
                for other in rewritten
                if other != address
                and text
                and (element := doc_map.get(other)) is not None
                and element.text.strip() == text
            ),
            None,
        )
        if moved_from is not None:
            problems.append(
                f"operation {index} (set_text): this moves the text of "
                f"{moved_from} into {address}; each paragraph's style says what "
                'it holds (a "Date", a title...): keep every value in its own '
                "paragraph"
            )


def _check_headings_kept(
    doc_map: DocMap,
    plans: Sequence[_Plan],
    template_headings: frozenset[str] | None,
    problems: list[str],
) -> None:
    """Reject a model ``write`` over a chapter heading (renamed or not: the
    chapter would merge into the one before) or over another heading of the
    template (all headings when they are not given): it would turn a real
    heading into body text. Removing a chapter with its heading stays
    possible."""
    chapter_level = min(
        (e.heading_level for e in doc_map.elements if e.heading_level is not None),
        default=0,
    )
    for plan in plans:
        if not isinstance(plan.op, Write):
            continue
        headings = [
            t
            for t in plan.targets
            if t.heading_level is not None
            and (
                t.heading_level == chapter_level
                or template_headings is None
                or _plain(t.text) in template_headings
            )
        ]
        if headings:
            names = ", ".join(f'{h.address} "{h.text[:60]}"' for h in headings[:3])
            problems.append(
                f"operation {plan.index} (write): the range includes the heading "
                f"{names}; keep the template's headings: write only the content "
                "between them (start after the heading, stop before the next one), "
                "and do not repeat the heading in the Markdown. Rename a heading "
                "with set_text; remove a whole chapter with remove"
            )


def _check_template_blocks(
    document: DocxDocument,
    plans: Sequence[_Plan],
    protected: frozenset[TableSignature],
    problems: list[str],
) -> None:
    """Reject a batch that removes every copy of a template table, or that
    removes a copy of one while writing a Markdown table (a rebuild)."""
    consumed = [n for p in plans for n in p.consumed]
    consumed_inside = [n for p in plans for n in p.consumed_inside]
    remaining = [
        table_signature(t)
        for t in _body_tables(document)
        if not _inside(t, consumed, strict=False)
        and not _inside(t, consumed_inside, strict=True)
    ]
    removes_block = False
    for plan in plans:
        signatures = [
            s for t in _removed_tables(plan) if (s := table_signature(t)) in protected
        ]
        if not signatures:
            continue
        removes_block = True
        if any(s not in remaining for s in signatures):
            problems.append(
                f"operation {plan.index} ({plan.op.op}): this removes every copy of a "
                f"table block of the template; {_KEEP_BLOCK}"
            )
    if not removes_block:
        return
    for plan in plans:
        if isinstance(plan.op, Write | InsertAfter) and (
            markdown_features(plan.op.markdown).has_table
        ):
            problems.append(
                f"operation {plan.index} ({plan.op.op}): this batch removes a copy of "
                "a table block of the template and writes a Markdown table in its "
                "place; fill the template's own copies instead: repeat a copy to get "
                "more, remove the extra ones"
            )


# --- XML helpers ---------------------------------------------------------------

_KEEP_IN_PARAGRAPH = (qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd"))
_PICTURE_TAGS = (qn("w:drawing"), qn("w:pict"), qn("w:object"))
_W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_HEADING_PPR_TAGS = (qn("w:pStyle"), qn("w:numPr"), qn("w:outlineLvl"))


def _is_picture_run(child) -> bool:
    return (
        child.tag == qn("w:r")
        and any(True for _ in child.iter(*_PICTURE_TAGS))
        and not any(True for _ in child.iter(qn("w:t")))
    )


def _base_run_properties(paragraph) -> Any | None:
    runs = own_runs(paragraph)
    text_runs = [r for r in runs if any(True for _ in r.iter(qn("w:t")))]
    for run in text_runs or runs:
        properties = run.find(qn("w:rPr"))
        if properties is not None:
            return properties
    mark = paragraph.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
    return mark


def _clean_run_properties(properties) -> Any | None:
    if properties is None:
        return None
    clean = copy.deepcopy(properties)
    clean.tag = qn("w:rPr")
    for tag in (
        "w:highlight",
        "w:ins",
        "w:del",
        "w:moveFrom",
        "w:moveTo",
        "w:rPrChange",
    ):
        for child in clean.findall(qn(tag)):
            clean.remove(child)
    return clean


def _new_runs(text: str, base) -> list[Any]:
    spans = parse_inline_markdown(text) if text else []
    todo_ranges = [
        (m.start(), m.end())
        for m in TODO_MARKER_RE.finditer("".join(s.text for s in spans))
    ]
    out = []
    offset = 0
    for span in spans:
        for piece, highlighted in split_by_ranges(span.text, offset, todo_ranges):
            element = OxmlElement("w:r")
            if base is not None:
                element.append(copy.deepcopy(base))
            run = Run(cast(Any, element), cast(Any, None))
            run.text = piece  # "\n" becomes a line break, "\t" a tab
            if span.bold:
                run.bold = True
            if span.italic:
                run.italic = True
            if span.code:
                run.font.name = "Courier New"
            if highlighted:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            out.append(element)
        offset += len(span.text)
    return out


def set_paragraph_text(
    paragraph,
    text: str,
    *,
    in_body: bool = False,
    body_colour: str | None = None,
    placeholder_styles: frozenset[str] = frozenset(),
) -> None:
    """Replace a paragraph's text with ``text`` (inline Markdown, ``\\n`` as a
    line break): the paragraph properties, bookmarks and pictures stay, the
    first run's formatting is kept but for what marked a placeholder (its
    highlight, Word's placeholder style, and ``in_body`` the guidance colour,
    replaced by the template's ``body_colour``) and every other run goes."""
    base = _clean_run_properties(_base_run_properties(paragraph))
    if base is not None:
        for style in base.findall(qn("w:rStyle")):
            if style.get(qn("w:val")) in placeholder_styles:
                base.remove(style)
    if base is not None and in_body:
        for colour in base.findall(qn("w:color")):
            base.remove(colour)
        if body_colour is not None:
            colour = OxmlElement("w:color")
            colour.set(qn("w:val"), body_colour)
            base.append(colour)
    insert_at = None
    for position, child in enumerate(list(paragraph)):
        if child.tag in _KEEP_IN_PARAGRAPH or _is_picture_run(child):
            continue
        if insert_at is None:
            insert_at = position
        paragraph.remove(child)
    children = list(paragraph)
    if insert_at is None:
        insert_at = len(children)
    insert_at = min(insert_at, len(children))
    # Keep the text after the paragraph properties.
    if children and children[0].tag == qn("w:pPr"):
        insert_at = max(insert_at, 1)
    for offset, run in enumerate(_new_runs(text, base)):
        paragraph.insert(insert_at + offset, run)


def _first_paragraph(container) -> Any | None:
    return next((b for b in block_children(container) if b.tag == qn("w:p")), None)


def _clear_container(container, keep=None) -> None:
    """Remove the container's blocks (and wrappers), keeping its properties
    and, when given, the direct child holding ``keep``."""
    holder = None
    if keep is not None:
        holder = keep
        while holder.getparent() is not container:
            holder = holder.getparent()
    for child in list(container):
        if child.tag in (qn("w:tcPr"),) or child is holder:
            continue
        if child.tag in (qn("w:p"), qn("w:tbl"), qn("w:sdt"), qn("w:customXml")):
            container.remove(child)


def _empty_paragraph_like(reference) -> Any:
    paragraph = OxmlElement("w:p")
    properties = reference.find(qn("w:pPr")) if reference is not None else None
    if properties is not None:
        clean = copy.deepcopy(properties)
        for child in clean.findall(qn("w:sectPr")):
            clean.remove(child)
        paragraph.append(clean)
    return paragraph


def _adopt_container_look(elements: Sequence[Any], reference) -> None:
    """Rendered blocks written in a cell, header, footer or text box take
    the look of the paragraph they replace: plain paragraphs its paragraph
    properties, every run its run properties minus emphasis (bold, italic
    and underline come from the Markdown only)."""
    if reference is None:
        return
    ref_ppr = reference.find(qn("w:pPr"))
    ref_rpr = _clean_run_properties(_base_run_properties(reference))
    if ref_rpr is not None:
        for tag in ("w:b", "w:bCs", "w:i", "w:iCs", "w:u"):
            for child in ref_rpr.findall(qn(tag)):
                ref_rpr.remove(child)
    for element in elements:
        if element.tag != qn("w:p"):
            continue
        properties = element.find(qn("w:pPr"))
        styled = properties is not None and properties.find(qn("w:pStyle")) is not None
        if not styled and ref_ppr is not None:
            if properties is not None:
                element.remove(properties)
            clean = copy.deepcopy(ref_ppr)
            for tag in ("w:numPr", "w:sectPr"):
                for child in clean.findall(qn(tag)):
                    clean.remove(child)
            element.insert(0, clean)
        if ref_rpr is None:
            continue
        for run_element in element.iter(qn("w:r")):
            run = Run(cast(Any, run_element), cast(Any, None))
            bold, italic = run.bold, run.italic
            highlight = run.font.highlight_color
            code = run.font.name == "Courier New"
            old = run_element.find(qn("w:rPr"))
            if old is not None:
                run_element.remove(old)
            run_element.insert(0, copy.deepcopy(ref_rpr))
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if highlight is not None:
                run.font.highlight_color = highlight
            if code:
                run.font.name = "Courier New"


def _empty_keeping_section_break(paragraph) -> None:
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    properties = paragraph.find(qn("w:pPr"))
    if properties is not None:
        for child in list(properties):
            if child.tag in _HEADING_PPR_TAGS:
                properties.remove(child)


def _remove_nodes(nodes: Sequence[Any], story_root) -> None:
    """Remove blocks or rows; a paragraph carrying a section break is
    emptied instead, and bookmark ends of removed bookmarks go too (the
    table of contents is rebuilt after the batch)."""
    ids: set[str] = set()
    for node in nodes:
        for start in node.iter(qn("w:bookmarkStart")):
            if bookmark_id := start.get(qn("w:id")):
                ids.add(bookmark_id)
    for node in nodes:
        if (
            node.tag == qn("w:p")
            and node.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
        ):
            _empty_keeping_section_break(node)
            continue
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)
    for end in list(story_root.iter(qn("w:bookmarkEnd"))):
        if end.get(qn("w:id")) in ids:
            end_parent = end.getparent()
            if end_parent is not None:
                end_parent.remove(end)


class _Ids:
    """Fresh ids for drawings in cloned content (duplicate `wp:docPr` ids
    make Word report a corrupted file)."""

    def __init__(self, document: DocxDocument, roots: Iterable[Any]) -> None:
        values = [
            int(v)
            for root in roots
            for el in root.iter(qn("wp:docPr"))
            if (v := el.get("id")) and v.isdigit()
        ]
        self.next = max(values, default=0) + 1000

    def take(self) -> int:
        self.next += 1
        return self.next


def _clean_clone(clone, ids: _Ids) -> None:
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for element in list(clone.iter(qn(tag))):
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
    for element in clone.iter():
        for attribute in (f"{{{_W14}}}paraId", f"{{{_W14}}}textId"):
            if attribute in element.attrib:
                del element.attrib[attribute]
    for doc_pr in clone.iter(qn("wp:docPr")):
        doc_pr.set("id", str(ids.take()))
    for sdt_id in list(clone.iter(qn("w:id"))):
        parent = sdt_id.getparent()
        if parent is not None and parent.tag == qn("w:sdtPr"):
            parent.remove(sdt_id)


def _ensure_ends_with_paragraph(container) -> None:
    """A cell, header, footer or text box must end with a paragraph."""
    if container.tag not in (qn("w:tc"), qn("w:hdr"), qn("w:ftr"), qn("w:txbxContent")):
        return
    blocks = block_children(container)
    if not blocks or blocks[-1].tag != qn("w:p"):
        container.append(OxmlElement("w:p"))


def _sync_textbox_fallback(content) -> None:
    """Mirror an edited DrawingML text box into its VML fallback copy."""
    choice = next(
        (a for a in content.iterancestors() if a.tag == f"{{{_MC_NS}}}Choice"), None
    )
    if choice is None:
        return
    alternate = choice.getparent()
    if alternate is None:
        return
    fallback = alternate.find(f"{{{_MC_NS}}}Fallback")
    if fallback is None:
        return
    choice_boxes = list(choice.iter(qn("w:txbxContent")))
    fallback_boxes = list(fallback.iter(qn("w:txbxContent")))
    if len(choice_boxes) != len(fallback_boxes):
        return
    for source, target in zip(choice_boxes, fallback_boxes, strict=True):
        if source is content:
            for child in list(target):
                target.remove(child)
            for child in source:
                target.append(copy.deepcopy(child))


# --- replace_text ---------------------------------------------------------------

_OPAQUE = "\ufffc"  # a field character: never matched, never edited
_FIELD_CHAR = qn("w:fldChar")
_FLD_SIMPLE = qn("w:fldSimple")


@dataclass
class _Unit:
    node: Any
    kind: Literal["t", "tab", "br", "field"]
    start: int
    text: str
    shown: str  # the text as the map shows it (a field's result included)


def _paragraphs_of(node) -> list[Any]:
    return [node] if node.tag == qn("w:p") else list(node.iter(qn("w:p")))


def text_units(paragraph) -> list[_Unit]:
    """The paragraph's own text in reading order, as units: editable text
    (`w:t`), tabs and line breaks (matched as a space / a line break, removed
    when inside a match) and field text (codes and results, opaque)."""
    field_runs: set[int] = set()
    depth = 0
    # Held for the whole function: lxml reuses a proxy's id() once it is
    # garbage collected, so the ids below are only stable while it lives.
    runs = own_runs(paragraph)
    for run in runs:
        in_simple = any(a.tag == _FLD_SIMPLE for a in run.iterancestors())
        for child in run:
            if child.tag == _FIELD_CHAR:
                kind = child.get(qn("w:fldCharType"))
                if kind == "begin":
                    depth += 1
                elif kind == "end":
                    depth = max(0, depth - 1)
                field_runs.add(id(run))
        if depth or in_simple:
            field_runs.add(id(run))
    units: list[_Unit] = []
    offset = 0
    for node in own_text_units(paragraph):
        run = node.getparent()
        is_t = node.tag == qn("w:t")
        shown = (
            (node.text or "")
            if is_t
            else " "
            if node.tag == qn("w:tab")
            else ""
            if node.get(qn("w:type")) in ("page", "column")
            else "\n"
        )
        kind: Literal["t", "tab", "br", "field"]
        if (run is not None and id(run) in field_runs) or (
            not is_t and node.tag != qn("w:tab") and not shown
        ):
            text, kind = _OPAQUE * max(1, len(shown)), "field"
        elif is_t:
            text, kind = shown, "t"
        elif node.tag == qn("w:tab"):
            text, kind = shown, "tab"
        else:
            text, kind = shown, "br"
        units.append(_Unit(node, kind, offset, text, shown))
        offset += len(text)
    return units


def _units_text(paragraph) -> str:
    return "".join(u.shown for u in text_units(paragraph))


def _find_in_paragraph(paragraph, find: str) -> tuple[list[tuple[int, int]], bool]:
    """Non-overlapping matches of ``find`` in the paragraph's text, and
    whether a match would cross a field (then nothing may be replaced)."""
    units = text_units(paragraph)
    text = "".join(u.text for u in units)
    matches: list[tuple[int, int]] = []
    start = text.find(find)
    while start != -1:
        matches.append((start, start + len(find)))
        start = text.find(find, start + len(find))
    if not matches:
        # The map shows a field's result too: a `find` that runs through a
        # field is reported as such rather than "not found".
        return [], find in "".join(u.shown for u in units)
    return matches, False


def _replace_in_paragraph(paragraph, find: str, replace: str) -> None:
    matches, _blocked = _find_in_paragraph(paragraph, find)
    # Last match first: the offsets of earlier matches stay valid.
    for start, end in reversed(matches):
        units = [
            u
            for u in text_units(paragraph)
            if u.start < end and u.start + len(u.text) > start
        ]
        receiver = next((u for u in units if u.kind == "t"), None)
        if receiver is None:
            # The match is only tabs / line breaks: the text goes in a new run
            # right before the first of them.
            run = units[0].node.getparent()
            new_run = OxmlElement("w:r")
            properties = run.find(qn("w:rPr")) if run is not None else None
            if properties is not None:
                new_run.append(copy.deepcopy(properties))
            new_text = OxmlElement("w:t")
            new_text.text = replace
            new_text.set(qn("xml:space"), "preserve")
            new_run.append(new_text)
            if run is not None:
                run.addprevious(new_run)
        for unit in units:
            local_start = max(0, start - unit.start)
            local_end = min(len(unit.text), end - unit.start)
            if unit.kind in ("tab", "br"):
                parent = unit.node.getparent()
                if parent is not None:
                    parent.remove(unit.node)
                continue
            current = unit.node.text or ""
            inserted = replace if unit is receiver else ""
            unit.node.text = current[:local_start] + inserted + current[local_end:]
            unit.node.set(qn("xml:space"), "preserve")
            if unit is receiver and local_start == 0 and local_end == len(current):
                # The whole run was the placeholder: no longer highlighted.
                run = unit.node.getparent()
                properties = run.find(qn("w:rPr")) if run is not None else None
                if properties is not None:
                    for highlight in properties.findall(qn("w:highlight")):
                        properties.remove(highlight)


# --- template list look ----------------------------------------------------------


def numbering_root(document: DocxDocument) -> Any | None:
    try:
        return document.part.numbering_part.element
    except Exception:  # noqa: BLE001 - a document without numbering has no list look
        return None


def style_num_ids(document: DocxDocument) -> dict[str, str]:
    """Paragraph style id -> numbering id it applies (through ``basedOn``)."""
    direct: dict[str, str | None] = {}
    based: dict[str, str] = {}
    for style in document.styles.element.findall(qn("w:style")):
        style_id = style.get(qn("w:styleId")) or ""
        num = style.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        direct[style_id] = num.get(qn("w:val")) if num is not None else None
        parent = style.find(qn("w:basedOn"))
        if parent is not None:
            based[style_id] = parent.get(qn("w:val")) or ""
    out: dict[str, str] = {}
    for style_id in direct:
        seen: set[str] = set()
        current: str | None = style_id
        while current and current not in seen:
            seen.add(current)
            if direct.get(current):
                out[style_id] = cast(str, direct[current])
                break
            current = based.get(current)
    return out


def list_kind_of(numbering, num_id: str) -> str | None:
    """ "bullet", "number", or None for a numbering id that draws nothing."""
    if num_id == "0":
        return None
    num = next(
        (n for n in numbering.iter(qn("w:num")) if n.get(qn("w:numId")) == num_id),
        None,
    )
    abstract_ref = num.find(qn("w:abstractNumId")) if num is not None else None
    if abstract_ref is None:
        return None
    abstract = next(
        (
            a
            for a in numbering.iter(qn("w:abstractNum"))
            if a.get(qn("w:abstractNumId")) == abstract_ref.get(qn("w:val"))
        ),
        None,
    )
    if abstract is None:
        return None
    level = next(
        (lvl for lvl in abstract.findall(qn("w:lvl")) if lvl.get(qn("w:ilvl")) == "0"),
        None,
    )
    fmt = level.find(qn("w:numFmt")) if level is not None else None
    value = fmt.get(qn("w:val")) if fmt is not None else None
    if value == "bullet":
        return "bullet"
    if value in (None, "none"):
        return None
    return "number"


@dataclass(frozen=True)
class _FoundLook:
    style_id: str | None
    direct_num_id: str | None
    effective_num_id: str


def _find_list_looks(
    document: DocxDocument, candidates: Iterable[MapElement]
) -> dict[str, _FoundLook]:
    """The first list paragraph of each kind among ``candidates`` (in
    order): the template's own look for written bullets / numbered items."""
    numbering = numbering_root(document)
    if numbering is None:
        return {}
    style_nums = style_num_ids(document)
    found: dict[str, _FoundLook] = {}
    for element in candidates:
        if element.kind != "paragraph" or element.heading_level is not None:
            continue
        properties = element.node.find(qn("w:pPr"))
        style = properties.find(qn("w:pStyle")) if properties is not None else None
        style_id = style.get(qn("w:val")) if style is not None else None
        direct = (
            properties.find(f"{qn('w:numPr')}/{qn('w:numId')}")
            if properties is not None
            else None
        )
        direct_id = direct.get(qn("w:val")) if direct is not None else None
        effective = direct_id or (style_nums.get(style_id) if style_id else None)
        if not effective:
            continue
        kind = list_kind_of(numbering, effective)
        if kind is None or kind in found:
            continue
        found[kind] = _FoundLook(style_id, direct_id, effective)
        if len(found) == 2:
            break
    return found


def _restarted_numbering(document: DocxDocument, num_id: str) -> str | None:
    """A new numbering id sharing ``num_id``'s definition but starting over,
    so a written numbered list does not continue the template's count."""
    numbering = numbering_root(document)
    if numbering is None:
        return None
    num = next(
        (n for n in numbering.iter(qn("w:num")) if n.get(qn("w:numId")) == num_id),
        None,
    )
    abstract_ref = num.find(qn("w:abstractNumId")) if num is not None else None
    if abstract_ref is None:
        return None
    new_id = 1 + max(
        (int(n.get(qn("w:numId")) or 0) for n in numbering.iter(qn("w:num"))),
        default=0,
    )
    new_num = OxmlElement("w:num")
    new_num.set(qn("w:numId"), str(new_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), abstract_ref.get(qn("w:val")) or "0")
    new_num.append(reference)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    new_num.append(override)
    numbering.append(new_num)
    return str(new_id)


# --- applying a batch ----------------------------------------------------------


def _dominant_body_table(doc_map: DocMap) -> Any | None:
    """A copy of the first body table carrying the most frequent table style
    after the first heading (cover-page tables not counted), else None."""
    top = [e for e in doc_map.elements if e.container == "" and e.is_block]
    first_heading = next(
        (i for i, e in enumerate(top) if e.heading_level is not None), 0
    )
    counts: dict[str, int] = {}
    firsts: dict[str, Any] = {}
    for element in top[first_heading:]:
        if element.kind != "table":
            continue
        style = element.node.find(f"{qn('w:tblPr')}/{qn('w:tblStyle')}")
        value = style.get(qn("w:val")) if style is not None else None
        if value:
            counts[value] = counts.get(value, 0) + 1
            firsts.setdefault(value, element.node)
    if not counts:
        return None
    return copy.deepcopy(firsts[max(counts, key=counts.__getitem__)])


def _story_root(element: MapElement, document: DocxDocument):
    return document.element.body if element.story == "body" else element.part.element


def _body_text_colour(document: DocxDocument) -> str | None:
    """The explicit colour the template's own body text carries ("auto" or a
    dark colour over a coloured style, say), most of its text by length; None
    when the text takes its styles' colour."""
    counts: dict[str, int] = {}
    for run in document.element.body.iter(qn("w:r")):
        colour = run.find(f"{qn('w:rPr')}/{qn('w:color')}")
        text = "".join(t.text or "" for t in run.iter(qn("w:t")))
        if colour is None or not text.strip():
            continue
        value = colour.get(qn("w:val")) or ""
        if value == "auto" or _dark(value):
            counts[value] = counts.get(value, 0) + len(text)
    return max(counts, key=counts.__getitem__) if counts else None


def _dark(hex_colour: str) -> bool:
    try:
        r, g, b = (int(hex_colour[i : i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    return 0.2126 * r + 0.7152 * g + 0.0722 * b < 64


def _placeholder_styles(document: DocxDocument) -> frozenset[str]:
    """Word's "Placeholder Text" character style ids (localized ids differ)."""
    return frozenset(
        s.style_id for s in document.styles if s.name == "Placeholder Text"
    )


def _paint_runs(elements: Sequence[Any], colour: str) -> None:
    """Runs of written paragraphs without a colour of their own take the
    template's body text colour (a style may be grey under it)."""
    for element in elements:
        if element.tag != qn("w:p"):
            continue
        for run in element.iter(qn("w:r")):
            properties = run.get_or_add_rPr()
            if properties.find(qn("w:color")) is None:
                colour_el = OxmlElement("w:color")
                colour_el.set(qn("w:val"), colour)
                properties.append(colour_el)


class _Applier:
    def __init__(
        self, document: DocxDocument, doc_map: DocMap, images: Mapping[str, bytes]
    ) -> None:
        self.document = document
        self.map = doc_map
        self.images = images
        self.scratch_source: bytes | None = None
        self.dominant_table = _dominant_body_table(doc_map)
        self.width = usable_width(document)
        roots = [document.element.body] + [
            e.node for e in doc_map.elements if e.kind in ("header", "footer")
        ]
        self.ids = _Ids(document, roots)
        self.containers: list[Any] = []
        self._body_paragraphs = [
            e for e in doc_map.elements if e.kind == "paragraph" and e.story == "body"
        ]
        self.body_colour = _body_text_colour(document)
        self.placeholder_styles = _placeholder_styles(document)
        self.heading_styles = _heading_styles(doc_map)
        self.main_body_style = _main_body_style(doc_map)

    def _list_looks(
        self, place: Sequence[MapElement], markdown: str
    ) -> dict[str, ListLook]:
        """The template's own list look for the lists ``markdown`` writes at
        ``place``: the list paragraphs being replaced first, then the
        body's. A numbered list restarts at 1."""
        bullets = _BULLET_LINE_RE.search(markdown or "") is not None
        numbered = _NUMBERED_LINE_RE.search(markdown or "") is not None
        if not (bullets or numbered):
            return {}
        prefixes = tuple(f"{e.address}." for e in place)
        place_ids = {id(e) for e in place}
        inside = [
            e
            for e in self.map.elements
            if e.kind == "paragraph"
            and (id(e) in place_ids or e.address.startswith(prefixes))
        ]
        found = _find_list_looks(self.document, [*inside, *self._body_paragraphs])
        looks: dict[str, ListLook] = {}
        bullet = found.get("bullet") if bullets else None
        if bullet is not None:
            looks["bullet"] = ListLook(bullet.style_id, bullet.direct_num_id)
        number = found.get("number") if numbered else None
        if number is not None:
            restarted = _restarted_numbering(self.document, number.effective_num_id)
            looks["number"] = ListLook(
                number.style_id, restarted or number.direct_num_id
            )
        return looks

    def _render(
        self,
        markdown: str,
        place: MapElement,
        table_reference,
        body_style: str | None = None,
        targets: Sequence[MapElement] = (),
    ) -> list[Any]:
        if self.scratch_source is None:
            buffer = io.BytesIO()
            self.document.save(buffer)
            self.scratch_source = buffer.getvalue()
        stamp_font = _flows_in_body(place) and body_style is None
        rendered = render_blocks(
            self.scratch_source,
            markdown,
            pictures=PictureTarget(place.part, self.width, self.images),
            table_reference=table_reference
            if table_reference is not None
            else self.dominant_table,
            default_font=template_default_font(self.document) if stamp_font else None,
            list_looks=self._list_looks(targets or [place], markdown),
            heading_styles=self.heading_styles,
        )
        if body_style is not None:
            for element in rendered:
                if element.tag != qn("w:p"):
                    continue
                properties = element.get_or_add_pPr()
                if properties.find(qn("w:pStyle")) is None:
                    properties.get_or_add_pStyle().set(qn("w:val"), body_style)
        if self.body_colour is not None and _flows_in_body(place):
            _paint_runs(rendered, self.body_colour)
        # Rendered in a scratch copy: its pictures' ids may repeat each other.
        for element in rendered:
            for drawing in element.iter(qn("wp:docPr")):
                drawing.set("id", str(self.ids.take()))
        return rendered

    def _note_container(self, node) -> None:
        self.containers.append(node)

    def apply(self, plan: _Plan) -> None:
        op = plan.op
        if isinstance(op, SetText):
            target = plan.targets[0]
            if target.kind == "cell":
                first = _first_paragraph(target.node)
                _clear_container(target.node, keep=first)
                if first is None:
                    first = OxmlElement("w:p")
                    target.node.append(first)
                set_paragraph_text(
                    first, op.text, placeholder_styles=self.placeholder_styles
                )
            else:
                set_paragraph_text(
                    target.node,
                    op.text,
                    in_body=_flows_in_body(target),
                    body_colour=self.body_colour,
                    placeholder_styles=self.placeholder_styles,
                )
            self._note_container(target.node)
            return
        if isinstance(op, ReplaceText):
            target = plan.targets[0]
            for paragraph in _paragraphs_of(target.node):
                _replace_in_paragraph(paragraph, op.find, op.replace)
            self._note_container(target.node)
            return
        if isinstance(op, InsertAfter):
            anchor = plan.targets[0]
            rendered = self._render(
                op.markdown,
                anchor,
                None,
                (_body_style([anchor]) or self.main_body_style)
                if _flows_in_body(anchor)
                else None,
            )
            if not _flows_in_body(anchor):
                _adopt_container_look(
                    rendered, anchor.node if anchor.kind == "paragraph" else None
                )
            cursor = anchor.node
            for element in rendered:
                cursor.addnext(element)
                cursor = element
            self._note_container(anchor.node)
            return
        if isinstance(op, Write):
            first = plan.targets[0]
            if first.kind == "cell":
                reference = _first_paragraph(first.node)
                reference_copy = (
                    copy.deepcopy(reference) if reference is not None else None
                )
                rendered = self._render(op.markdown, first, None)
                _adopt_container_look(rendered, reference_copy)
                _clear_container(first.node)
                for element in rendered:
                    first.node.append(element)
                if not rendered:
                    first.node.append(_empty_paragraph_like(reference_copy))
                self._note_container(first.node)
                return
            nodes = [e.node for e in plan.targets]
            table = next(
                (copy.deepcopy(n) for n in nodes if n.tag == qn("w:tbl")), None
            )
            rendered = self._render(
                op.markdown,
                first,
                table,
                (_body_style(plan.targets) or self.main_body_style)
                if _flows_in_body(first)
                else None,
                plan.targets,
            )
            if not _flows_in_body(first):
                paragraph = next((n for n in nodes if n.tag == qn("w:p")), None)
                _adopt_container_look(
                    rendered,
                    copy.deepcopy(paragraph) if paragraph is not None else None,
                )
            for element in rendered:
                first.node.addprevious(element)
            parent = first.node.getparent()
            _remove_nodes(nodes, _story_root(first, self.document))
            if parent is not None:
                self._note_container(parent)
            return
        if isinstance(op, Move):
            assert plan.destination is not None
            where, anchor = plan.destination
            nodes = [e.node for e in plan.targets]
            for node in nodes:
                node.getparent().remove(node)
            if where == "before":
                for node in nodes:
                    anchor.addprevious(node)
            else:
                cursor = anchor
                for node in nodes:
                    cursor.addnext(node)
                    cursor = node
            return
        if isinstance(op, Repeat):
            nodes = [e.node for e in plan.targets]
            cursor = nodes[-1]
            for _ in range(op.count):
                for node in nodes:
                    clone = copy.deepcopy(node)
                    _clean_clone(clone, self.ids)
                    cursor.addnext(clone)
                    cursor = clone
            return
        assert isinstance(op, Remove)
        first = plan.targets[0]
        parent = first.node.getparent()
        _remove_nodes(
            [e.node for e in plan.targets],
            _story_root(first, self.document),
        )
        if parent is not None:
            self._note_container(parent)

    def finish(self) -> None:
        for node in self.containers:
            for candidate in [node, *node.iterancestors()]:
                if candidate.tag in (
                    qn("w:tc"),
                    qn("w:hdr"),
                    qn("w:ftr"),
                    qn("w:txbxContent"),
                ):
                    _ensure_ends_with_paragraph(candidate)
                    if candidate.tag == qn("w:txbxContent"):
                        _sync_textbox_fallback(candidate)
                    break
        renumber_figure_captions(self.document)


def _heading_styles(doc_map: DocMap) -> dict[int, str]:
    """The style the document's headings use at each level, its own one (a
    contract's "Title 2") over Word's built-in "Heading N" when it has both:
    a heading the model adds then looks, and is numbered, like its siblings."""
    counts: dict[int, Counter[str]] = {}
    for e in doc_map.elements:
        if e.heading_level is None or e.story != "body" or e.kind != "paragraph":
            continue
        style = e.node.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        if style is not None and (style_id := style.get(qn("w:val"))):
            builtin = re.fullmatch(r"heading \d", (e.style or "").lower()) is not None
            counts.setdefault(e.heading_level, Counter())[style_id] += (
                0 if builtin else 1000
            )
            counts[e.heading_level][style_id] += 1
    return {level: c.most_common(1)[0][0] for level, c in counts.items() if c}


def _body_style(elements: Sequence[MapElement]) -> str | None:
    """The paragraph style of the first plain text paragraph among
    ``elements`` (not a heading, not a list item): written paragraphs take
    it, so they look like the body text they replace - only the style, never
    the replaced text's direct formatting (an instruction's blue). A
    caption is not body text: written paragraphs never take its style."""
    for element in elements:
        if _is_body_text(element):
            return _style_id(element)
    return None


def _main_body_style(doc_map: DocMap) -> str | None:
    """The style most of the body text uses (a contract's "Body text"), None
    when most of it has none: written text takes it where the text it
    replaces has no style, rather than the document default's font."""
    counts = Counter(
        _style_id(e)
        for e in doc_map.elements
        if e.story == "body" and e.container == "" and _is_body_text(e)
    )
    return counts.most_common(1)[0][0] if counts else None


def _is_body_text(element: MapElement) -> bool:
    return (
        element.kind == "paragraph"
        and bool(element.text.strip())
        and element.heading_level is None
        and "list" not in element.flags
        and (element.style or "").lower() != "caption"
        and not is_toc(element)
    )


def _style_id(element: MapElement) -> str | None:
    style = element.node.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return style.get(qn("w:val")) if style is not None else None


def _flows_in_body(element: MapElement) -> bool:
    return (
        element.story == "body"
        and not element.in_textbox
        and element.kind != "cell"
        and ".C" not in element.address
    )


def _operation_number(problem: str) -> int:
    match = re.match(r"operations? (\d+)", problem)
    return int(match.group(1)) if match else 0


def apply_batch(
    docx_bytes: bytes,
    operations: Sequence[Operation | Mapping[str, Any]],
    images: Mapping[str, bytes] | None = None,
    *,
    protected_tables: frozenset[TableSignature] | None = None,
    template_headings: frozenset[str] | None = None,
    refresh_toc: bool = True,
) -> bytes:
    """Apply one batch to ``docx_bytes`` and return the new document.

    Every address refers to the document as given. The whole batch is
    validated first; on any problem :class:`BatchError` lists them all and
    nothing is applied. ``protected_tables`` (the model's batches: the
    base's ``template_table_signatures``) rejects a batch that leaves no
    copy of a template table it removes, or that writes over a heading of
    the template (``template_headings``; every heading when not given)."""
    images = images or {}
    problems: list[str] = []
    parsed: list[tuple[int, Operation]] = []
    for index, raw in enumerate(operations, start=1):
        try:
            parsed.append((index, parse_operation(raw)))
        except ValidationError as error:
            details = "; ".join(
                f"{'.'.join(str(p) for p in e['loc']) or 'operation'}: {e['msg']}"
                for e in error.errors()
            )
            problems.append(f"operation {index}: invalid ({details})")
    document = Document(io.BytesIO(docx_bytes))
    if any(isinstance(op, Write | InsertAfter) for _, op in parsed):
        ensure_required_styles(document)
    doc_map = build_map_from_document(document)
    plans = [_plan(doc_map, index, op, images, problems) for index, op in parsed]
    if len(parsed) > 1 and any(isinstance(op, Move) for _, op in parsed):
        problems.append(
            "a move goes in a batch of its own (every other address would shift)"
        )
    _conflicts(plans, problems)
    if protected_tables is not None:
        _check_headings_kept(doc_map, plans, template_headings, problems)
        _check_computed_kept(doc_map, plans, problems)
        _check_places(doc_map, plans, problems)
    if protected_tables:
        _check_template_blocks(document, plans, protected_tables, problems)
    if problems:
        raise BatchError(sorted(problems, key=_operation_number))
    applier = _Applier(document, doc_map, images)
    for plan in plans:
        applier.apply(plan)
    applier.finish()
    if refresh_toc:
        refresh_table_of_contents(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def replay(
    template_bytes: bytes,
    batches: Sequence[Sequence[Operation | Mapping[str, Any]]],
    images: Mapping[str, bytes] | None = None,
) -> bytes:
    """Apply an operation log to the template: each batch is
    resolved against the result of the previous ones."""
    result = template_bytes
    for index, batch in enumerate(batches, start=1):
        # The table of contents is rebuilt once, after the last batch.
        result = apply_batch(result, batch, images, refresh_toc=index == len(batches))
    return result


__all__ = [
    "MAX_REPEAT",
    "BatchError",
    "InsertAfter",
    "Move",
    "Operation",
    "Remove",
    "Repeat",
    "ReplaceText",
    "SetText",
    "TableSignature",
    "Write",
    "apply_batch",
    "fit_to_document",
    "normalized",
    "parse_operation",
    "replay",
    "set_paragraph_text",
    "table_signature",
    "inspect_set_texts",
    "chapters_with_template_sub_headings",
    "template_heading_texts",
    "template_sub_heading_levels",
    "template_table_signatures",
]
