# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Markdown -> Word blocks, adapted from writable_document's docx export:
headings (the template's heading styles), paragraphs with inline formatting,
lists on 3 levels, tables, code blocks and ``![caption](img:<id>)`` images;
``[À COMPLÉTER]`` markers highlighted, horizontal rules dropped. Rendered in a
scratch copy of the document, so every style id resolves in the target.
"""

from __future__ import annotations

import copy
import io
import itertools
import logging
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any, cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Length, Pt
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph

from fred_capability_word_template.inline import parse_inline_markdown

logger = logging.getLogger(__name__)

Font = tuple[str | None, Length | None]

# Every named style the renderer can reference. A template only defines the
# styles it used, and python-docx raises `KeyError` applying a missing one.
_REQUIRED_STYLE_NAMES = (
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
    "Heading 5",
    "Heading 6",
    "List Bullet",
    "List Bullet 2",
    "List Bullet 3",
    "List Number",
    "List Number 2",
    "List Number 3",
    "Table Grid",
    "Caption",
)

_default_styles_source: DocxDocument | None = None


def _get_default_styles_source() -> DocxDocument:
    global _default_styles_source
    if _default_styles_source is None:
        _default_styles_source = Document()
    return _default_styles_source


def ensure_required_styles(document: DocxDocument) -> None:
    """Copy in whichever required style ``document`` lacks, from
    python-docx's default template; the template's own styles are kept."""
    existing = {s.name for s in document.styles}
    source = _get_default_styles_source()
    for name in _REQUIRED_STYLE_NAMES:
        if name in existing:
            continue
        try:
            style_element = source.styles[name].element
        except KeyError:
            continue
        copied = copy.deepcopy(style_element)
        _copy_style_numbering(document, source, copied)
        document.styles.element.append(copied)


def _copy_style_numbering(target: DocxDocument, source: DocxDocument, style) -> None:
    """A copied list style points at a numbering definition by id, which may
    mean something else in the target: copy the definition under fresh ids."""
    num_id_el = style.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
    if num_id_el is None:
        return
    source_numbering = source.part.numbering_part.element
    source_num = next(
        (
            n
            for n in source_numbering.iter(qn("w:num"))
            if n.get(qn("w:numId")) == num_id_el.get(qn("w:val"))
        ),
        None,
    )
    abstract_ref = (
        source_num.find(qn("w:abstractNumId")) if source_num is not None else None
    )
    if abstract_ref is None:
        return
    source_abstract = next(
        (
            a
            for a in source_numbering.iter(qn("w:abstractNum"))
            if a.get(qn("w:abstractNumId")) == abstract_ref.get(qn("w:val"))
        ),
        None,
    )
    if source_abstract is None:
        return
    numbering = target.part.numbering_part.element
    new_abstract_id = 1 + max(
        (
            int(a.get(qn("w:abstractNumId")))
            for a in numbering.iter(qn("w:abstractNum"))
        ),
        default=-1,
    )
    new_num_id = 1 + max(
        (int(n.get(qn("w:numId"))) for n in numbering.iter(qn("w:num"))), default=0
    )
    abstract = copy.deepcopy(source_abstract)
    abstract.set(qn("w:abstractNumId"), str(new_abstract_id))
    for nsid in abstract.findall(qn("w:nsid")):
        abstract.remove(nsid)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(new_abstract_id))
    num.append(ref)
    numbering.append(num)
    num_id_el.set(qn("w:val"), str(new_num_id))


_HEADING_RE = re.compile(r"^(?P<hashes>#{1,6})\s+(?P<text>.*)$")
_BULLET_RE = re.compile(r"^(?P<indent>\s*)[-*+]\s+(?P<text>.*)$")
_ORDERED_RE = re.compile(r"^(?P<indent>\s*)\d+[.)]\s+(?P<text>.*)$")
_FENCE_RE = re.compile(r"^\s*```")
_RULE_RE = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")
IMAGE_REF_RE = re.compile(
    r"^\s*!\[(?P<caption>[^\]]*)\]\(img:(?P<image_id>[A-Za-z0-9_.-]+)\)\s*$"
)
_MISSING_IMAGE_TEXT = "[image indisponible]"
_TABLE_ROW_RE = re.compile(r"^\s*\|?.*\|.*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")
_LIST_INDENT_WIDTH = 2
_MAX_LIST_LEVEL = 3

# A gap marked by the model, highlighted yellow in the output.
TODO_MARKER_RE = re.compile(r"\[\s*(?:À|A)\s+COMPL[ÉE]TER[^\]]*\]", re.IGNORECASE)


@dataclass(frozen=True)
class MarkdownFeatures:
    has_table: bool
    heading_levels: frozenset[int]
    image_ids: tuple[str, ...]


def markdown_features(markdown: str) -> MarkdownFeatures:
    """What a Markdown text would render - checked before a batch applies."""
    lines = (markdown or "").splitlines()
    images = tuple(
        m.group("image_id") for line in lines if (m := IMAGE_REF_RE.match(line))
    )
    in_fence = False
    has_table = False
    levels: set[int] = set()
    for i, line in enumerate(lines):
        if _FENCE_RE.match(line):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        has_table = has_table or _is_table_start(lines, i)
        if heading := _HEADING_RE.match(line):
            levels.add(min(len(heading.group("hashes")), 6))
    return MarkdownFeatures(has_table, frozenset(levels), images)


def _list_level(indent: str) -> int:
    columns = indent.replace("\t", " " * _LIST_INDENT_WIDTH).count(" ")
    return min(columns // _LIST_INDENT_WIDTH + 1, _MAX_LIST_LEVEL)


def _list_style(base: str, indent: str) -> str:
    level = _list_level(indent)
    return base if level == 1 else f"{base} {level}"


@dataclass(frozen=True)
class ListLook:
    """How the template itself draws one kind of list ("bullet" or
    "number"): the paragraph style id of its list paragraphs and, when the
    numbering is set on the paragraph rather than by the style, the
    numbering id. Written list items copy it instead of "List Bullet" /
    "List Number"."""

    style_id: str | None
    num_id: str | None


def _apply_list_look(paragraph: Paragraph, look: ListLook, level: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    if look.style_id:
        properties.get_or_add_pStyle().set(qn("w:val"), look.style_id)
    if look.num_id is None and level == 1:
        return
    numbering = properties.get_or_add_numPr()
    numbering.get_or_add_ilvl().val = level - 1
    if look.num_id is not None:
        numbering.get_or_add_numId().val = int(look.num_id)


def split_by_ranges(
    text: str, offset: int, ranges: Sequence[tuple[int, int]]
) -> list[tuple[str, bool]]:
    """Cut ``text`` (at ``offset`` in the whole line) into pieces flagged as
    inside one of ``ranges`` or not; at least one piece."""
    cuts = {0, len(text)}
    for start, end in ranges:
        for point in (start - offset, end - offset):
            if 0 < point < len(text):
                cuts.add(point)
    bounds = sorted(cuts)
    pieces = [
        (
            text[a:b],
            any(start <= offset + a and offset + b <= end for start, end in ranges),
        )
        for a, b in itertools.pairwise(bounds)
    ]
    return pieces or [(text, False)]


def add_inline_runs(
    paragraph: Paragraph, text: str, *, default_font: Font | None = None
) -> None:
    """Render one line of inline Markdown as runs on ``paragraph``, the
    ``[À COMPLÉTER ...]`` markers highlighted yellow."""
    spans = parse_inline_markdown(text)
    todo_ranges = [
        (m.start(), m.end())
        for m in TODO_MARKER_RE.finditer("".join(s.text for s in spans))
    ]
    offset = 0
    for span in spans:
        for piece, highlighted in split_by_ranges(span.text, offset, todo_ranges):
            run = paragraph.add_run(piece)
            if span.bold:
                run.bold = True
            if span.italic:
                run.italic = True
            if highlighted:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if span.code:
                run.font.name = "Courier New"
            elif default_font is not None:
                name, size = default_font
                if name:
                    run.font.name = name
                if size:
                    run.font.size = size
        offset += len(span.text)


def _split_table_row(line: str) -> list[str]:
    placeholder = "\x00"
    cells = line.replace(r"\|", placeholder).split("|")
    if cells and cells[0].strip() == "":
        cells = cells[1:]
    if cells and cells[-1].strip() == "":
        cells = cells[:-1]
    return [c.replace(placeholder, "|").strip() for c in cells]


def _is_table_start(lines: list[str], i: int) -> bool:
    return (
        i + 1 < len(lines)
        and _TABLE_ROW_RE.match(lines[i]) is not None
        and "|" in lines[i]
        and _TABLE_SEP_RE.match(lines[i + 1]) is not None
    )


def _add_table(
    document: DocxDocument, rows: list[list[str]], *, default_font: Font | None
) -> None:
    n_cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(n_cols):
            text = row[c_idx] if c_idx < len(row) else ""
            paragraph = table.cell(r_idx, c_idx).paragraphs[0]
            add_inline_runs(paragraph, text, default_font=default_font)
            if r_idx == 0:
                for run in paragraph.runs:
                    run.bold = True


# --- written tables take a template table's look -----------------------------

_TCPR_BEFORE_SHD = tuple(
    qn(f"w:{tag}")
    for tag in ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders")
)
_COPIED_RUN_PROPERTIES = ("b", "bCs", "color", "sz")


def _cell_look(cell) -> tuple[Any, Any, Any, Any]:
    """(shading, paragraph style, justification, run properties) of a
    reference cell - run properties from its first run, else its paragraph mark."""
    paragraph = cell.find(qn("w:p"))
    run = next(paragraph.iter(qn("w:r")), None) if paragraph is not None else None
    run_properties = run.find(qn("w:rPr")) if run is not None else None
    paragraph_properties = (
        paragraph.find(qn("w:pPr")) if paragraph is not None else None
    )
    if run_properties is None and paragraph_properties is not None:
        run_properties = paragraph_properties.find(qn("w:rPr"))
    return (
        cell.find(f"{qn('w:tcPr')}/{qn('w:shd')}"),
        paragraph_properties.find(qn("w:pStyle"))
        if paragraph_properties is not None
        else None,
        paragraph_properties.find(qn("w:jc"))
        if paragraph_properties is not None
        else None,
        run_properties,
    )


def _copy_attributes(target, source) -> None:
    target.attrib.clear()
    target.attrib.update(source.attrib)


def _apply_cell_look(cell, look: tuple[Any, Any, Any, Any], *, header: bool) -> None:
    shading, style, justification, run_properties = look
    if shading is not None:
        cell_properties = cell.get_or_add_tcPr()
        for old in cell_properties.findall(qn("w:shd")):
            cell_properties.remove(old)
        before = [c for c in cell_properties if c.tag in _TCPR_BEFORE_SHD]
        if before:
            before[-1].addnext(copy.deepcopy(shading))
        else:
            cell_properties.insert(0, copy.deepcopy(shading))
    copied = [
        (name, source)
        for name in _COPIED_RUN_PROPERTIES
        if run_properties is not None
        and (source := run_properties.find(qn(f"w:{name}"))) is not None
    ]
    for paragraph in cell.iter(qn("w:p")):
        paragraph_properties = paragraph.get_or_add_pPr()
        if style is not None:
            _copy_attributes(paragraph_properties.get_or_add_pStyle(), style)
        if justification is not None:
            _copy_attributes(paragraph_properties.get_or_add_jc(), justification)
        if not copied:
            continue
        for run in paragraph.iter(qn("w:r")):
            properties = run.get_or_add_rPr()
            if header:
                properties._remove_b()  # CT_RPr has no public remover
            for name, source in copied:
                _copy_attributes(getattr(properties, f"get_or_add_{name}")(), source)


def apply_table_look(table, reference) -> None:
    """Give a rendered table the look of ``reference`` (a template table):
    its `w:tblPr` (fixed width and floating position dropped), the header
    cells' look from its first row and the body cells' from its second.
    The column grid is never copied: the column count may differ."""
    old = table.find(qn("w:tblPr"))
    table_properties = copy.deepcopy(reference.find(qn("w:tblPr")))
    if old is None or table_properties is None:
        return
    width = table_properties.find(qn("w:tblW"))
    if width is not None and width.get(qn("w:type")) not in ("pct", "auto"):
        table_properties.remove(width)
    for floating in table_properties.findall(qn("w:tblpPr")):
        table_properties.remove(floating)
    own_style = old.find(qn("w:tblStyle"))
    if table_properties.find(qn("w:tblStyle")) is None and own_style is not None:
        # An unstyled reference draws its borders per cell: keep "Table Grid".
        table_properties.insert(0, own_style)
    old.addprevious(table_properties)
    table.remove(old)
    reference_rows = reference.findall(qn("w:tr"))
    for index, row in enumerate(table.findall(qn("w:tr"))):
        source_index = min(index, 1)
        if source_index >= len(reference_rows):
            continue
        source_cells = reference_rows[source_index].findall(qn("w:tc"))
        if not source_cells:
            continue
        for column, cell in enumerate(row.findall(qn("w:tc"))):
            source = source_cells[min(column, len(source_cells) - 1)]
            _apply_cell_look(cell, _cell_look(source), header=index == 0)


# --- images ------------------------------------------------------------------


class PictureTarget:
    """Where rendered pictures must be registered: the story part the
    blocks finally land in, and the widest a picture may be there."""

    def __init__(
        self, part: Any, max_width: int | None, images: Mapping[str, bytes]
    ) -> None:
        self.part = part
        self.max_width = max_width
        self.images = images


def usable_width(document: DocxDocument) -> int | None:
    section = document.sections[-1]
    if section.page_width is None:
        return None
    return section.page_width - (section.left_margin or 0) - (section.right_margin or 0)


def _add_image_block(
    document: DocxDocument,
    image_id: str,
    caption: str,
    pictures: PictureTarget | None,
    default_font: Font | None,
) -> None:
    content = pictures.images.get(image_id) if pictures is not None else None
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline = None
    if content is not None and pictures is not None:
        try:
            inline = pictures.part.new_pic_inline(io.BytesIO(content))
        except Exception:  # noqa: BLE001 - any unreadable image renders a marker
            logger.warning("[WORD_TEMPLATE] image %s could not be embedded.", image_id)
    if inline is None or pictures is None:
        add_inline_runs(
            paragraph,
            f"{_MISSING_IMAGE_TEXT} {caption}".strip(),
            default_font=default_font,
        )
        return
    # Tag it so the "Rédaction" view and the map read it back as
    # ![caption](img:<id>) (see `docmap.placed_image`).
    inline.docPr.set("name", f"img:{image_id}")
    inline.docPr.set("descr", caption.strip())
    extent = inline.extent
    if pictures.max_width and extent.cx > pictures.max_width:
        extent.cy = int(extent.cy * pictures.max_width / extent.cx)
        extent.cx = int(pictures.max_width)
        graphic_extent = inline.find(".//" + qn("a:ext"))
        if graphic_extent is not None:
            graphic_extent.set("cx", str(extent.cx))
            graphic_extent.set("cy", str(extent.cy))
    paragraph.add_run()._r.add_drawing(inline)
    if caption.strip():
        _add_figure_caption(document, caption.strip(), default_font)


def _add_figure_caption(
    document: DocxDocument, caption: str, default_font: Font | None
) -> None:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_runs(paragraph, "Figure ", default_font=default_font)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " SEQ Figure \\* ARABIC ")
    number_run = OxmlElement("w:r")
    number_text = OxmlElement("w:t")
    number_text.text = "1"
    number_run.append(number_text)
    field.append(number_run)
    paragraph._p.append(field)
    add_inline_runs(paragraph, f" : {caption}", default_font=default_font)


def renumber_figure_captions(document: DocxDocument) -> None:
    """Rewrite every `SEQ Figure` field's cached number in document order."""
    number = 0
    for field in document.element.body.iter(qn("w:fldSimple")):
        if "SEQ Figure" not in (field.get(qn("w:instr")) or ""):
            continue
        number += 1
        texts = list(field.iter(qn("w:t")))
        if texts:
            texts[0].text = str(number)


# --- the rendering loop ------------------------------------------------------


def _add_code_block(document: DocxDocument, lines: list[str]) -> None:
    run = document.add_paragraph().add_run("\n".join(lines))
    run.font.name = "Courier New"
    run.font.size = Pt(9)


def template_default_font(document: DocxDocument) -> Font:
    normal = cast(ParagraphStyle, document.styles["Normal"])
    return normal.font.name, normal.font.size


def _render_markdown_into(
    document: DocxDocument,
    content_md: str,
    *,
    default_font: Font | None,
    pictures: PictureTarget | None,
    list_looks: Mapping[str, ListLook] | None = None,
    heading_styles: Mapping[int, str] | None = None,
) -> None:
    """Append ``content_md``'s rendered blocks to the end of ``document``'s body.
    A heading takes the template's own style for its level when given (a
    contract's "Title 2" style), else Word's built-in one."""
    lines = (content_md or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        image = IMAGE_REF_RE.match(line)
        if image:
            _add_image_block(
                document,
                image.group("image_id"),
                image.group("caption"),
                pictures,
                default_font,
            )
            i += 1
            continue
        if _FENCE_RE.match(line):
            block: list[str] = []
            i += 1
            while i < len(lines) and not _FENCE_RE.match(lines[i]):
                block.append(lines[i])
                i += 1
            i += 1
            _add_code_block(document, block)
            continue
        if not line.strip() or _RULE_RE.match(line):
            i += 1
            continue
        heading = _HEADING_RE.match(line)
        if heading:
            level = min(len(heading.group("hashes")), 6)
            style_id = (heading_styles or {}).get(level)
            if style_id:
                paragraph = document.add_paragraph()
                paragraph._p.get_or_add_pPr().get_or_add_pStyle().set(
                    qn("w:val"), style_id
                )
            else:
                paragraph = document.add_heading(level=level)
            add_inline_runs(paragraph, heading.group("text").strip())
            i += 1
            continue
        if _is_table_start(lines, i):
            header = _split_table_row(lines[i])
            i += 2
            body = []
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i]) and "|" in lines[i]:
                body.append(_split_table_row(lines[i]))
                i += 1
            _add_table(document, [header, *body], default_font=default_font)
            continue
        listed = _BULLET_RE.match(line) or _ORDERED_RE.match(line)
        if listed:
            bullet = _BULLET_RE.match(line) is not None
            look = (list_looks or {}).get("bullet" if bullet else "number")
            if look is None:
                base = "List Bullet" if bullet else "List Number"
                paragraph = document.add_paragraph(
                    style=_list_style(base, listed.group("indent"))
                )
            else:
                paragraph = document.add_paragraph()
                _apply_list_look(paragraph, look, _list_level(listed.group("indent")))
            add_inline_runs(
                paragraph, listed.group("text").strip(), default_font=default_font
            )
            i += 1
            continue
        # A plain paragraph: consecutive plain lines are one soft-wrapped paragraph.
        para_lines = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (
                not nxt.strip()
                or _RULE_RE.match(nxt)
                or _HEADING_RE.match(nxt)
                or _BULLET_RE.match(nxt)
                or _ORDERED_RE.match(nxt)
                or _FENCE_RE.match(nxt)
                or IMAGE_REF_RE.match(nxt)
                or _is_table_start(lines, i)
            ):
                break
            para_lines.append(nxt.strip())
            i += 1
        add_inline_runs(
            document.add_paragraph(), " ".join(para_lines), default_font=default_font
        )


def render_blocks(
    scratch_source: bytes,
    markdown: str,
    *,
    pictures: PictureTarget | None,
    table_reference: Any | None,
    default_font: Font | None,
    list_looks: Mapping[str, ListLook] | None = None,
    heading_styles: Mapping[int, str] | None = None,
) -> list[Any]:
    """Render ``markdown`` into a scratch document loaded from
    ``scratch_source`` (the target's own bytes, so style ids and numbering
    match) and return the rendered body elements. Each table takes the look
    of ``table_reference`` when given, else stays "Table Grid"; list items
    take ``list_looks`` (the template's own list look) when given."""
    scratch = Document(io.BytesIO(scratch_source))
    body = scratch.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    ensure_required_styles(scratch)
    _render_markdown_into(
        scratch,
        markdown,
        default_font=default_font,
        pictures=pictures,
        list_looks=list_looks,
        heading_styles=heading_styles,
    )
    elements = [el for el in body if el.tag != qn("w:sectPr")]
    if table_reference is not None:
        for element in elements:
            if element.tag == qn("w:tbl"):
                apply_table_look(element, table_reference)
    return elements
