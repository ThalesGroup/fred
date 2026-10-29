# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

from __future__ import annotations

import pytest
from conftest import add_vml_textbox, form_document, report_document, to_bytes
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import RGBColor
from fred_capability_word_template import build_map, parts, render_map, render_part
from fred_capability_word_template.docmap import (
    chapter_of,
    check_parts,
    color_name,
    unused_template_ranges,
)


def test_form_addresses_and_merged_full_width_row(form_bytes: bytes) -> None:
    doc_map = build_map(form_bytes)
    assert doc_map.get("T1.R1.C2") is not None
    assert doc_map["T1.R1.C2"].text == "Week 10"
    merged = doc_map.get("T2.R3.C1")
    assert merged is not None and merged.full_width and merged.col_span == 4
    # A merged cell appears once, at its first grid position.
    assert doc_map.get("T2.R3.C2") is None
    assert doc_map["T2.R4.C1.P2"].text == "second line"
    text = render_map(doc_map)
    assert '[C1 "Significant comments" - full width]' in text
    assert "[P1] ∅" in text
    assert "[T2.R4.C1.P1]" in text  # multi-paragraph cell is detailed


def test_headers_footers_and_textboxes() -> None:
    document = form_document()
    add_vml_textbox(document.add_paragraph("anchor"), "Boxed title")
    doc_map = build_map(to_bytes(document))
    assert doc_map["H1.P1"].text == "Header text"
    assert doc_map["F1.P1"].text == "Footer text"
    assert doc_map["X1.P1"].text == "Boxed title"
    # The text-box paragraph is not part of its anchor's own text.
    assert doc_map["P2"].text == "anchor"
    text = render_map(doc_map)
    assert "[H1] header (section 1)" in text
    assert '[X1.P1] "Boxed title"' in text


def test_nested_table_addresses() -> None:
    document = Document()
    outer = document.add_table(rows=1, cols=2)
    inner = outer.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "deep"
    doc_map = build_map(to_bytes(document))
    assert doc_map["T1.R1.C2.T1.R1.C1"].text == "deep"
    assert "[T1.R1.C2.T1] table" in render_map(doc_map)


def test_vertical_merge_spans_rows() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(2, 0)).text = "tall"
    doc_map = build_map(to_bytes(document))
    assert doc_map["T1.R1.C1"].row_span == 3
    assert doc_map.get("T1.R2.C1") is None
    assert "spans 3 rows" in render_map(doc_map)


def test_headings_and_colour(report_bytes: bytes) -> None:
    doc_map = build_map(report_bytes)
    text = render_map(doc_map)
    assert '[P4 heading 1] "Context"' in text
    assert '[P5] "Describe the context here." {blue}' in text
    assert '[P6] "Fixed sentence."' in text


def test_parts_split_at_top_headings(report_bytes: bytes) -> None:
    doc_map = build_map(report_bytes)
    found = parts(doc_map)
    assert [p.title for p in found if p.part_id != "hf"] == [
        "before the first heading",
        "Context",
        "Scope",
        "Annex",
    ]
    scope = render_part(doc_map, "2")
    assert "Scope instruction." in scope and "Annex" not in scope
    with pytest.raises(KeyError):
        render_part(doc_map, "9")


def test_parts_chunk_headingless_documents() -> None:
    document = Document()
    for index in range(95):
        document.add_paragraph(f"line {index}")
    found = [p for p in parts(build_map(to_bytes(document)), chunk_size=40)]
    assert [p.size for p in found] == [40, 40, 15]
    assert found[1].first == "P41"


def test_a_long_chapter_is_read_in_pieces_and_keeps_the_numbering() -> None:
    document = Document()
    document.add_heading("Missions", level=1)
    for index in range(100):
        document.add_paragraph(f"line {index}")
    document.add_heading("Places", level=1)
    found = parts(build_map(to_bytes(document)), chunk_size=40)
    assert [p.part_id for p in found] == ["1", "1b", "1c", "2"]
    assert [p.title for p in found] == [
        "Missions",
        "Missions (continued)",
        "Missions (continued)",
        "Places",
    ]
    assert [chapter_of(p.part_id) for p in found] == ["1", "1", "1", "2"]
    assert chapter_of("hf") == "hf" and chapter_of("12") == "12"


def test_form_part_title_and_hf_part(form_bytes: bytes) -> None:
    doc_map = build_map(form_bytes)
    found = parts(doc_map)
    assert found[0].title == "Weekly report"
    assert found[-1].part_id == "hf"
    assert "Footer text" in render_part(doc_map, "hf")


def test_empty_runs_collapse() -> None:
    document = Document()
    document.add_paragraph("a")
    for _ in range(4):
        document.add_paragraph()
    document.add_paragraph("b")
    assert "[P2–P5] ∅ ×4" in render_map(build_map(to_bytes(document)))


def test_color_names() -> None:
    assert color_name("0070C0") == "blue"
    assert color_name("FF0000") == "red"
    assert color_name("808080") == "grey"
    assert color_name("1F3864") == "dark blue"


def test_report_document_builds() -> None:
    assert render_map(build_map(to_bytes(report_document())))


def _unfinished(checks) -> list[str]:
    return [c.part_id for c in checks if not c.finished]


def test_the_check_follows_the_document_not_the_conversation() -> None:
    document = report_document()
    document.add_heading("Empty chapter", level=1)
    doc_map = build_map(to_bytes(document))
    guidance = frozenset({"Describe the context here."})
    # Chapter 1 keeps its guidance, the last chapter has nothing written;
    # the others have content (fixed text, a table).
    checks = check_parts(doc_map, guidance)
    assert _unfinished(checks) == ["1", "4"]
    (left,) = checks[1].leftovers
    assert (left.kind, left.text, left.where) == (
        "template",
        "Describe the context here.",
        'under "Context"',
    )
    assert checks[4].empty
    assert _unfinished(check_parts(doc_map, frozenset())) == ["4"]


def test_fields_markers_notes_and_toc_lines() -> None:
    document = Document()
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("Letter", level=1)
    document.add_paragraph("1 Letter 3", style="toc 1")
    document.add_paragraph("Written by [Name] for the team.")
    document.add_paragraph("Budget: [À COMPLÉTER : amount]")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Date"
    table.cell(0, 1).text = "00/00/0000"
    doc_map = build_map(to_bytes(document))
    (check, *_rest) = check_parts(
        doc_map,
        frozenset({"00/00/0000", "1 Letter 3"}),
        frozenset({"Written by [Name] for the team."}),
        notes={"Written by [Name] for the team.": '"Author"'},
        is_marker=lambda text: "COMPLÉTER" in text,
    )
    assert [(lo.kind, lo.where) for lo in check.leftovers] == [
        ("field", '"Author"'),
        ("marker", 'under "Letter"'),
        ("template", 'row "Date"'),
    ]
    assert check.places_left == 2 and not check.finished


def test_a_repeated_block_lists_each_copy() -> None:
    document = Document()
    for _ in range(2):
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Project code"
    checks = check_parts(build_map(to_bytes(document)), frozenset({"Project code"}))
    assert [lo.address for lo in checks[0].leftovers] == ["T1.R1.C1.P1", "T2.R1.C1.P1"]


def test_guidance_stays_coloured_in_a_guidance_heavy_template() -> None:
    # More blue guidance than black text must not make blue the body colour.
    document = Document()
    document.add_paragraph("Short fixed line.")
    for _ in range(3):
        run = document.add_paragraph().add_run("Describe the context here, at length.")
        run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    doc_map = build_map(to_bytes(document))
    assert [e.color for e in doc_map.elements if e.kind == "paragraph"] == [
        None,
        "blue",
        "blue",
        "blue",
    ]


def test_a_placeholder_in_a_body_text_box_is_checked() -> None:
    document = Document()
    document.add_heading("Cover", level=1)
    add_vml_textbox(document.add_paragraph(), "[Project name]")
    document.add_heading("Next", level=1)
    document.add_paragraph("Done.")
    (cover, _next) = check_parts(
        build_map(to_bytes(document)), frozenset({"[Project name]"})
    )
    assert [lo.text for lo in cover.leftovers] == ["[Project name]"]
    assert not cover.empty


def test_a_sub_heading_with_nothing_under_it_is_left() -> None:
    # The text went under the chapter heading; two sub-headings stay empty.
    document = Document()
    document.add_heading("General", level=1)
    document.add_heading("Scope", level=2)
    document.add_heading("Company", level=2)
    document.add_paragraph("The company runs three sites.")
    document.add_heading("IT department", level=2)
    document.add_paragraph("")
    (check,) = check_parts(build_map(to_bytes(document)), frozenset())
    assert [(lo.kind, lo.text) for lo in check.leftovers] == [
        ("empty", "Scope"),
        ("empty", "IT department"),
    ]
    assert not check.finished


def test_a_long_template_text_is_shown_whole_when_asked() -> None:
    document = Document()
    long_text = "Fixed company presentation. " * 30
    document.add_paragraph(long_text)
    doc_map = build_map(to_bytes(document))
    assert "(+" in render_part(doc_map, "1")
    whole = render_part(doc_map, "1", full=frozenset({long_text.strip()}))
    assert "(+" not in whole and long_text.strip() in whole


def test_a_document_without_headings_counts_its_parts_from_one(form_bytes) -> None:
    assert [p.part_id for p in parts(build_map(form_bytes)) if p.part_id != "hf"][
        0
    ] == "1"


def test_unused_template_ranges_keep_what_was_written() -> None:
    document = Document()
    document.add_heading("Chapter", level=1)
    document.add_paragraph("Written.")
    document.add_paragraph("Guidance")
    document.add_paragraph()
    document.add_table(rows=1, cols=1)
    document.add_paragraph("Caption")
    document.add_paragraph("[À COMPLÉTER : the owner]")
    document.add_heading("Sample", level=1)
    document.add_paragraph("Sample text")
    sample = frozenset({"Chapter", "Guidance", "Caption", "Sample", "Sample text"})
    found = unused_template_ranges(
        build_map(to_bytes(document)),
        sample,
        sample,
        is_marker=lambda text: text.startswith("[À COMPLÉTER"),
    )
    # The guidance, the empty frame and its caption go together (with the
    # empty paragraph between them); the marker stays; the unused chapter
    # goes whole. The written chapter keeps its heading.
    assert [(r.start, r.end) for r in found] == [("P3", "P5"), ("P7", "P8")]
    assert found[1].label == '"Sample" (heading and content)'
