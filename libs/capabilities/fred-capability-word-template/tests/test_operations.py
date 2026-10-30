# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

from __future__ import annotations

import io
from typing import cast

import pytest
from conftest import (
    PNG,
    add_bookmark,
    add_toc_entry,
    add_vml_textbox,
    load,
    report_document,
    to_bytes,
)
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import RGBColor
from docx.styles.style import ParagraphStyle
from fred_capability_word_template import (
    BatchError,
    Remove,
    Repeat,
    SetText,
    Write,
    apply_batch,
    build_map,
    parse_operation,
    replay,
)
from fred_capability_word_template.docmap import outline
from fred_capability_word_template.operations import (
    fit_to_document,
    inspect_set_texts,
    normalized,
    table_signature,
    template_heading_texts,
    template_table_signatures,
)


def _style(item) -> str | None:
    return item.style.name if item.style is not None else None


def _text(docx_bytes: bytes, address: str) -> str:
    element = build_map(docx_bytes).get(address)
    assert element is not None, address
    return element.text


# --- set_text ----------------------------------------------------------------


def test_set_text_keeps_first_run_format_and_drops_other_runs() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("Old ")
    first.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    first.font.highlight_color = WD_COLOR_INDEX.YELLOW
    paragraph.add_run("value").bold = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    out = load(
        apply_batch(to_bytes(document), [SetText(address="P1", text="New **bold**")])
    )
    result = out.paragraphs[0]
    assert result.text == "New bold"
    assert result.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert [r.bold for r in result.runs] == [None, True]
    # The placeholder highlight is gone once the field is filled, and in the
    # body the colour too (a template's guidance mark).
    assert all(r.font.highlight_color is None for r in result.runs)
    assert all(r.font.color.rgb is None for r in result.runs)


def test_set_text_in_a_cell_keeps_its_colour() -> None:
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Label").font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    out = load(
        apply_batch(to_bytes(document), [SetText(address="T1.R1.C1", text="Value")])
    )
    (run,) = out.tables[0].cell(0, 0).paragraphs[0].runs
    assert (run.text, run.font.color.rgb) == ("Value", RGBColor(0xFF, 0xFF, 0xFF))


@pytest.mark.parametrize(
    ("text", "written"),
    [
        ("First paragraph.\n\nSecond paragraph.", True),
        ("Points:\n- one\n- two", True),
        ("1. First\n2. Second", True),
        ("## Planning\nText.", True),
        ("44 Example Street\n75000 City", False),
        ("1. Introduction", False),
        ("- Steering the project.", True),
        ("One line.", False),
    ],
)
def test_set_text_with_several_paragraphs_is_written(text, written) -> None:
    op = {"op": "set_text", "address": "T1.R2.C1", "text": text}
    expected = {"op": "write", "from": "T1.R2.C1", "markdown": text}
    assert normalized(op) == (expected if written else op)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        # A table is not "one place": the set_text stays, for the engine to refuse.
        (
            {"op": "set_text", "address": "T2", "text": "A\n\nB"},
            {"op": "set_text", "address": "T2", "text": "A\n\nB"},
        ),
        ({"op": "remove", "address": "P83"}, {"op": "remove", "from": "P83"}),
        (
            {"op": "write", "from_": "P6", "markdown": "x"},
            {"op": "write", "from": "P6", "markdown": "x"},
        ),
        ({"op": "remove"}, {"op": "remove"}),
    ],
)
def test_operations_are_normalized_once_for_every_reader(raw, expected) -> None:
    assert normalized(raw) == expected


def test_several_paragraphs_become_paragraphs_in_the_body_look() -> None:
    document = Document()
    guidance = document.add_paragraph().add_run("Describe the context here.")
    guidance.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    document.add_paragraph("Next.")
    op = normalized(
        {"op": "set_text", "address": "P1", "text": "One.\n\nTwo:\n- a\n- b"}
    )
    out = load(apply_batch(to_bytes(document), [op]))
    texts = [p.text for p in out.paragraphs]
    assert texts == ["One.", "Two:", "a", "b", "Next."]
    assert all(r.font.color.rgb is None for p in out.paragraphs[:4] for r in p.runs)
    assert not any(p._p.findall(".//" + qn("w:br")) for p in out.paragraphs)


def test_set_text_line_break_and_todo_highlight() -> None:
    document = Document()
    document.add_paragraph("x")
    out = load(
        apply_batch(
            to_bytes(document),
            [{"op": "set_text", "address": "P1", "text": "a\nb [À COMPLÉTER : date]"}],
        )
    )
    paragraph = out.paragraphs[0]
    assert paragraph._p.find(f".//{qn('w:br')}") is not None
    highlighted = [r.text for r in paragraph.runs if r.font.highlight_color]
    assert highlighted == ["[À COMPLÉTER : date]"]


def test_set_text_on_cell_keeps_one_paragraph(form_bytes: bytes) -> None:
    result = apply_batch(form_bytes, [SetText(address="T2.R4.C1", text="only line")])
    doc_map = build_map(result)
    assert doc_map["T2.R4.C1"].text == "only line"
    assert doc_map.get("T2.R4.C1.P2") is None


def test_set_text_in_header_and_textbox(form_bytes: bytes) -> None:
    document = load(form_bytes)
    add_vml_textbox(document.add_paragraph("anchor"), "Box")
    result = apply_batch(
        to_bytes(document),
        [
            SetText(address="H1.P1", text="Project Alpha"),
            SetText(address="X1.P1", text="Cover title"),
        ],
    )
    assert _text(result, "H1.P1") == "Project Alpha"
    assert _text(result, "X1.P1") == "Cover title"


# --- write / insert_after ------------------------------------------------------


def test_write_replaces_range_with_rich_markdown(report_bytes: bytes) -> None:
    markdown = "Intro paragraph.\n\n## Sub part\n\n- first\n- second\n\n1. one\n"
    result = apply_batch(report_bytes, [Write(from_="P5", to="P6", markdown=markdown)])
    document = load(result)
    texts = [p.text for p in document.paragraphs]
    assert "Describe the context here." not in texts
    assert "Fixed sentence." not in texts
    index = texts.index("Intro paragraph.")
    styles = [_style(p) for p in document.paragraphs[index : index + 5]]
    assert styles == [
        "Normal",
        "Heading 2",
        "List Bullet",
        "List Bullet",
        "List Number",
    ]
    assert build_map(result)["P5"].text == "Intro paragraph."


def test_written_table_copies_the_replaced_table_look(report_bytes: bytes) -> None:
    markdown = "| Col | Val |\n|---|---|\n| a | 1 |\n"
    result = apply_batch(report_bytes, [Write(from_="P8", to="T1", markdown=markdown)])
    tables = load(result).tables
    assert len(tables) == 1
    assert _style(tables[0]) == "Light Grid Accent 1"
    assert tables[0].cell(1, 0).text == "a"


def test_written_table_falls_back_to_dominant_body_table(report_bytes: bytes) -> None:
    markdown = "| X | Y |\n|---|---|\n| 1 | 2 |\n"
    result = apply_batch(report_bytes, [Write(from_="P10", markdown=markdown)])
    tables = load(result).tables
    assert [_style(t) for t in tables] == [
        "Light Grid Accent 1",
        "Light Grid Accent 1",
    ]


def test_write_into_cell_with_list_keeps_cell_look(form_bytes: bytes) -> None:
    document = load(form_bytes)
    cell = document.tables[1].cell(3, 0)
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x0E, 0x28, 0x41)
    result = apply_batch(
        to_bytes(document),
        [Write(from_="T2.R4.C1", markdown="Status:\n\n- done\n- next")],
    )
    doc_map = build_map(result)
    assert doc_map["T2.R4.C1"].text == "Status:\ndone\nnext"
    out_cell = load(result).tables[1].cell(3, 0)
    assert [_style(p) for p in out_cell.paragraphs] == [
        "Normal",
        "List Bullet",
        "List Bullet",
    ]
    assert out_cell.paragraphs[1].runs[0].font.color.rgb == RGBColor(0x0E, 0x28, 0x41)


def test_insert_after_and_image(report_bytes: bytes) -> None:
    result = apply_batch(
        report_bytes,
        [
            {
                "op": "insert_after",
                "address": "P6",
                "markdown": "Added.\n\n![Plan](img:p1)",
            }
        ],
        images={"p1": PNG},
    )
    doc_map = build_map(result)
    assert doc_map["P7"].text == "Added."
    assert "image" in doc_map["P8"].flags
    assert doc_map["P9"].text.startswith("Figure 1 : Plan")


# --- repeat / remove -----------------------------------------------------------


def test_repeat_table_block_strips_bookmarks(form_bytes: bytes) -> None:
    document = load(form_bytes)
    add_bookmark(document.tables[1].cell(0, 0).paragraphs[0], "block", "7")
    result = apply_batch(to_bytes(document), [Repeat(from_="T2", count=2)])
    out = load(result)
    assert len(out.tables) == 4
    names = [b.get(qn("w:name")) for b in out.element.body.iter(qn("w:bookmarkStart"))]
    assert names == ["block"]
    assert build_map(result)["T4.R3.C1"].full_width


def test_repeat_range_and_rows() -> None:
    document = Document()
    document.add_paragraph("a")
    document.add_paragraph("b")
    table = document.add_table(rows=2, cols=1)
    table.cell(1, 0).text = "row"
    result = apply_batch(
        to_bytes(document),
        [Repeat(from_="P1", to="P2", count=1), Repeat(from_="T1.R2", count=3)],
    )
    out = load(result)
    assert [p.text for p in out.paragraphs] == ["a", "b", "a", "b"]
    assert len(out.tables[0].rows) == 5


def test_remove_range_and_its_toc_entries(report_bytes: bytes) -> None:
    result = apply_batch(report_bytes, [Remove(from_="P7", to="T1")])
    out = load(result)
    texts = [p.text for p in out.paragraphs]
    assert "Scope" not in texts and "2 Scope" not in texts
    assert "1 Context" in texts
    assert out.tables == []
    ends = [e.get(qn("w:id")) for e in out.element.body.iter(qn("w:bookmarkEnd"))]
    assert ends == ["1"]


def test_remove_keeps_section_break_paragraph() -> None:
    document = Document()
    document.add_paragraph("first section")
    document.add_section()
    document.add_paragraph("second section")
    breaker = document.element.body.findall(qn("w:p"))[1]
    assert breaker.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    result = apply_batch(to_bytes(document), [Remove(from_="P1", to="P3")])
    out = load(result)
    assert len(out.sections) == 2
    assert all(not p.text for p in out.paragraphs)


# --- batch semantics -----------------------------------------------------------


def test_invalid_batch_lists_every_problem_and_applies_nothing(
    form_bytes: bytes,
) -> None:
    with pytest.raises(BatchError) as caught:
        apply_batch(
            form_bytes,
            [
                SetText(address="T2.R4.C1", text="valid"),
                SetText(address="T9.R1.C1", text="x"),
                Repeat(from_="T2", count=0),
                Remove(from_="T2", to="T1"),
                Write(from_="T2.R4.C1", markdown="| a |\n|---|\n| 1 |"),
                {"op": "explode"},
                SetText(address="T2", text="x"),
                Remove(from_="P1", to="T2.R1.C1"),
            ],
        )
    problems = caught.value.problems
    # Sorted by operation; operation 1 is valid on its own but overlaps operation 5.
    assert len(problems) == 8
    assert problems[0].startswith("operations 1 and 5 overlap")
    assert "T9.R1.C1' does not exist" in problems[1] and "tables T1, T2" in problems[1]
    assert "between 1 and 50" in problems[2]
    assert "comes after" in problems[3]
    assert "cannot be written inside a table cell" in problems[4]
    assert problems[5].startswith("operation 6: invalid")
    assert "e.g. T2.R1.C1" in problems[6]
    assert "'T2.R1.C1' is a cell" in problems[7]


def test_unknown_row_hint(form_bytes: bytes) -> None:
    with pytest.raises(BatchError) as caught:
        apply_batch(form_bytes, [SetText(address="T2.R9.C1", text="x")])
    assert "T2 has rows R1–R4" in caught.value.problems[0]


def test_each_problem_is_listed_once(form_bytes: bytes) -> None:
    # "from" and "to" are the same unknown address: one line, not two.
    with pytest.raises(BatchError) as caught:
        apply_batch(
            form_bytes,
            [Remove(from_="T9", to="T9"), SetText(address="T9.R1.C1", text="x")],
        )
    problems = caught.value.problems
    assert len(problems) == len(set(problems)) == 2
    assert problems[0].startswith("operation 1 (remove): 'T9' does not exist")
    assert str(caught.value).count("operation 1 (remove)") == 1


def test_conflicting_operations_are_rejected(report_bytes: bytes) -> None:
    with pytest.raises(BatchError) as caught:
        apply_batch(
            report_bytes,
            [Remove(from_="P4", to="P6"), SetText(address="P5", text="x")],
        )
    assert "overlap" in caught.value.problems[0]


def test_missing_image_is_a_problem(report_bytes: bytes) -> None:
    with pytest.raises(BatchError) as caught:
        apply_batch(report_bytes, [Write(from_="P5", markdown="![x](img:nope)")])
    assert "'nope'" in caught.value.problems[0]


def test_addresses_refer_to_the_state_before_the_batch(form_bytes: bytes) -> None:
    result = apply_batch(
        form_bytes,
        [
            Repeat(from_="T2", count=1),
            SetText(address="T2.R2.C4", text="Milestone A"),
            Remove(from_="P1"),
            SetText(address="T1.R1.C2", text="Week 11"),
        ],
    )
    doc_map = build_map(result)
    # P1 (between T1 and T2) is gone, so the original block is now T2 and its clone T3.
    assert doc_map["T1.R1.C2"].text == "Week 11"
    assert doc_map["T2.R2.C4"].text == "Milestone A"
    assert doc_map["T3.R2.C4"].text == ""


def test_replay_three_batches(form_bytes: bytes) -> None:
    batches = [
        [Repeat(from_="T2", count=2)],
        # Addresses of batch 2 see the three blocks of batch 1.
        [SetText(address=f"T{n}.R2.C1", text=f"S{n - 1}") for n in (2, 3, 4)],
        [Write(from_="T4.R4.C1", markdown="- last block")],
    ]
    doc_map = build_map(replay(form_bytes, batches))
    assert [doc_map[f"T{n}.R2.C1"].text for n in (2, 3, 4)] == ["S1", "S2", "S3"]
    assert doc_map["T4.R4.C1"].text == "last block"
    assert doc_map["T3.R4.C1"].text == "old comment\nsecond line"


def test_operations_round_trip_as_json_with_from_alias() -> None:
    raw = {"op": "write", "from": "P1", "to": "P2", "markdown": "x"}
    operation = parse_operation(raw)
    assert isinstance(operation, Write) and operation.from_ == "P1"
    assert operation.model_dump(by_alias=True) == raw


def test_intra_word_underscores_are_kept_and_cell_writes_drop_emphasis(
    form_bytes: bytes,
) -> None:
    document = load(form_bytes)
    document.tables[1].cell(3, 0).paragraphs[0].runs[0].bold = True
    result = apply_batch(
        to_bytes(document),
        [
            SetText(address="T2.R2.C2", text="PRJ_19_002 and _italic_"),
            Write(from_="T2.R4.C1", markdown="- plain line"),
        ],
    )
    out = load(result)
    cell = out.tables[1].cell(1, 1).paragraphs[0]
    assert cell.text == "PRJ_19_002 and italic"
    assert [r.italic for r in cell.runs] == [None, True]
    assert out.tables[1].cell(3, 0).paragraphs[0].runs[0].bold is None


def test_written_paragraphs_take_the_replaced_body_style() -> None:
    document = Document()
    instruction = document.add_paragraph("Describe here.", style="Body Text")
    instruction.runs[0].font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    result = apply_batch(
        to_bytes(document), [Write(from_="P1", markdown="Written.\n\n- item")]
    )
    out = load(result)
    assert [_style(p) for p in out.paragraphs] == ["Body Text", "List Bullet"]
    assert out.paragraphs[0].runs[0].font.color.rgb is None


# --- template blocks kept (model batches) -------------------------------------


@pytest.fixture
def three_blocks(form_bytes: bytes) -> bytes:
    """The form with its merged 4-row project block three times (T2-T4)."""
    return apply_batch(form_bytes, [Repeat(from_="T2", count=2)])


def _tables(docx_bytes: bytes) -> int:
    return len(load(docx_bytes).tables)


def test_removing_every_copy_of_a_template_block_is_rejected(
    three_blocks: bytes,
) -> None:
    protected = template_table_signatures(three_blocks)
    with pytest.raises(BatchError) as error:
        apply_batch(
            three_blocks,
            [Remove(from_="T2", to="T4")],
            protected_tables=protected,
        )
    (problem,) = error.value.problems
    assert problem.startswith("operation 1 (remove): this removes every copy")
    assert "keep at least one copy of this block" in problem
    assert "undo it" in problem
    # Three removes in one batch, one per copy: the same.
    with pytest.raises(BatchError, match="every copy"):
        apply_batch(
            three_blocks,
            [Remove(from_="T2"), Remove(from_="T3"), Remove(from_="T4")],
            protected_tables=protected,
        )
    # A write over the copies is a removal too.
    with pytest.raises(BatchError, match=r"operation 1 \(write\): this removes"):
        apply_batch(
            three_blocks,
            [{"op": "write", "from": "T2", "to": "T4", "markdown": "text"}],
            protected_tables=protected,
        )


def test_removing_extra_copies_or_one_then_repeating_is_allowed(
    three_blocks: bytes,
) -> None:
    protected = template_table_signatures(three_blocks)
    two_removed = apply_batch(
        three_blocks, [Remove(from_="T3", to="T4")], protected_tables=protected
    )
    assert _tables(two_removed) == 2
    one_removed = apply_batch(
        three_blocks, [Remove(from_="T4")], protected_tables=protected
    )
    repeated = apply_batch(
        one_removed, [Repeat(from_="T3", count=2)], protected_tables=protected
    )
    assert _tables(repeated) == 5
    # Filled copies keep their structure: the last one can still not go.
    filled = apply_batch(
        two_removed,
        [SetText(address="T2.R1.C1", text="Perimeter")],
        protected_tables=protected,
    )
    with pytest.raises(BatchError, match="every copy"):
        apply_batch(filled, [Remove(from_="T2")], protected_tables=protected)


def test_rebuilding_a_block_as_a_markdown_table_is_rejected(
    three_blocks: bytes,
) -> None:
    protected = template_table_signatures(three_blocks)
    with pytest.raises(BatchError) as error:
        apply_batch(
            three_blocks,
            [
                Remove(from_="T3", to="T4"),
                {
                    "op": "insert_after",
                    "address": "T2",
                    "markdown": "| a | b |\n|---|---|\n| 1 | 2 |",
                },
            ],
            protected_tables=protected,
        )
    (problem,) = error.value.problems
    assert problem.startswith("operation 2 (insert_after): this batch removes a copy")
    # A Markdown table with no removal stays allowed.
    apply_batch(
        three_blocks,
        [
            {
                "op": "insert_after",
                "address": "T4",
                "markdown": "| a | b |\n|---|---|\n| 1 | 2 |",
            }
        ],
        protected_tables=protected,
    )


def test_a_table_the_model_created_is_not_protected(form_bytes: bytes) -> None:
    protected = template_table_signatures(form_bytes)
    with_table = apply_batch(
        form_bytes,
        [
            {
                "op": "insert_after",
                "address": "T2",
                "markdown": "| Scope | Code | Project | Milestone |\n|---|---|---|---|\n| a | b | c | d |",
            }
        ],
        protected_tables=protected,
    )
    assert _tables(with_table) == 3
    removed = apply_batch(with_table, [Remove(from_="T3")], protected_tables=protected)
    assert _tables(removed) == 2


def test_without_protection_and_for_chapters_removal_is_free(
    form_bytes: bytes, report_bytes: bytes
) -> None:
    # The collaborator's batches carry no protection.
    assert _tables(apply_batch(form_bytes, [Remove(from_="T2")])) == 1
    # A chapter (its heading included) goes with its only table.
    protected = template_table_signatures(report_bytes)
    removed = apply_batch(
        report_bytes, [Remove(from_="P7", to="T1")], protected_tables=protected
    )
    assert _tables(removed) == 0


def test_a_model_write_keeps_the_template_headings(report_bytes: bytes) -> None:
    protected = template_table_signatures(report_bytes)
    assert _text(report_bytes, "P4") == "Context"
    # Over the heading: refused, nothing applied.
    with pytest.raises(BatchError) as caught:
        apply_batch(
            report_bytes,
            [Write(from_="P4", to="P6", markdown="**1. Context**\n\nText.")],
            protected_tables=protected,
        )
    assert 'P4 "Context"' in caught.value.problems[0]
    assert "keep the template's headings" in caught.value.problems[0]
    # Between two headings: applied, the heading stays a real heading.
    written = apply_batch(
        report_bytes,
        [Write(from_="P5", to="P6", markdown="Text.")],
        protected_tables=protected,
    )
    heading = build_map(written).get("P4")
    assert heading is not None and heading.heading_level == 1
    # The collaborator's batches and a chapter removal stay free.
    apply_batch(report_bytes, [Write(from_="P4", to="P6", markdown="Text.")])
    apply_batch(report_bytes, [Remove(from_="P4", to="P6")], protected_tables=protected)


def test_signatures_ignore_filled_text_of_merged_blocks(form_bytes: bytes) -> None:
    document = load(form_bytes)
    before = table_signature(document.tables[1]._tbl)
    filled = apply_batch(form_bytes, [SetText(address="T2.R1.C1", text="Other")])
    after = table_signature(load(filled).tables[1]._tbl)
    assert before == after and before[0] == "merged"
    # A plain grid is told apart by its first-row texts.
    assert table_signature(document.tables[0]._tbl)[0] == "grid"


def test_removing_a_part_keeps_sentences_that_refer_to_it() -> None:
    # A cross-reference to the removed part is a sentence, not a TOC line.
    document = Document()
    add_toc_entry(document, "1 Context", "_Toc1")
    heading = document.add_heading("Context", level=1)
    add_bookmark(heading, "_Toc1", "1")
    add_bookmark(document.add_paragraph("Sample."), "_Ref1", "2")
    sentence = add_toc_entry(document, "as explained in the sample", "_Ref1")
    sentence.insert_paragraph_before("The requirements below apply,")
    document.add_heading("Scope", level=1)
    out = load(apply_batch(to_bytes(document), [Remove(from_="P2", to="P3")]))
    texts = [p.text for p in out.paragraphs]
    assert "1 Context" not in texts
    assert "as explained in the sample" in texts


@pytest.mark.parametrize(
    "op",
    [
        {"op": "set_text", "address": "T1.R1.C1", "text": "Flat"},
        {"op": "write", "from": "T1.R1.C1", "markdown": "Flat"},
    ],
)
def test_a_template_table_nested_in_a_cell_is_kept(op) -> None:
    document = Document()
    outer = document.add_table(rows=1, cols=1).cell(0, 0)
    inner = outer.add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "Label"
    docx = to_bytes(document)
    with pytest.raises(BatchError, match="every copy of a table block"):
        apply_batch(docx, [op], protected_tables=template_table_signatures(docx))


def test_set_text_drops_the_guidance_colour_inside_a_content_control() -> None:
    document = Document()
    guidance = document.add_paragraph()
    guidance.add_run("Describe the context.").font.color.rgb = RGBColor(
        0x00, 0x70, 0xC0
    )
    control = parse_xml(f"<w:sdt {nsdecls('w')}><w:sdtContent/></w:sdt>")
    guidance._p.addprevious(control)
    control[0].append(guidance._p)
    out = load(
        apply_batch(to_bytes(document), [SetText(address="P1", text="Written.")])
    )
    runs = [r for p in out.element.body.iter(qn("w:p")) for r in p.iter(qn("w:r"))]
    assert runs and all(r.find(f"{qn('w:rPr')}/{qn('w:color')}") is None for r in runs)


def test_an_extra_copy_of_a_filled_label_value_block_can_be_removed() -> None:
    document = Document()
    block = document.add_table(rows=2, cols=2)
    block.cell(0, 0).text = "Project"
    block.cell(1, 0).text = "Status"
    docx = to_bytes(document)
    protected = template_table_signatures(docx)
    docx = apply_batch(docx, [{"op": "repeat", "from": "T1", "count": 2}])
    docx = apply_batch(
        docx,
        [
            {"op": "set_text", "address": "T1.R1.C2", "text": "Alpha"},
            {"op": "set_text", "address": "T2.R1.C2", "text": "Beta"},
        ],
    )
    out = apply_batch(
        docx, [{"op": "remove", "from": "T3"}], protected_tables=protected
    )
    assert len(load(out).tables) == 2


def test_a_repeat_conflicts_with_an_edit_inside_its_range() -> None:
    # Every repeated block counts, not only the range's two ends.
    document = Document()
    for text in ("One.", "Two.", "Three."):
        document.add_paragraph(text)
    with pytest.raises(BatchError, match="overlap"):
        apply_batch(
            to_bytes(document),
            [
                {"op": "repeat", "from": "P1", "to": "P3", "count": 1},
                {"op": "set_text", "address": "P2", "text": "Changed."},
            ],
        )


def test_written_text_takes_the_templates_own_text_colour() -> None:
    # Grey styles, black text by an explicit "auto" colour: written text must
    # not fall back to the style's grey.
    document = Document()
    normal = document.styles["Normal"]
    assert isinstance(normal, ParagraphStyle)
    normal.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    fixed = document.add_paragraph().add_run("Fixed company text, in black.")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "auto")
    fixed._r.get_or_add_rPr().append(colour)
    guidance = document.add_paragraph().add_run("Describe it.")
    guidance.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    out = load(
        apply_batch(
            to_bytes(document),
            [
                {"op": "set_text", "address": "P2", "text": "Written."},
                {"op": "insert_after", "address": "P1", "markdown": "More."},
            ],
        )
    )
    for paragraph in out.paragraphs[1:]:
        values = [
            r._r.find(f"{qn('w:rPr')}/{qn('w:color')}").get(qn("w:val"))
            for r in paragraph.runs
        ]
        assert values and set(values) == {"auto"}, paragraph.text


def _toc_and_footer_document() -> bytes:
    document = Document()
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("1 Context 3", style="toc 1")
    printed = document.add_paragraph("Printed on ")
    begin = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>')
    instr = parse_xml(
        f"<w:r {nsdecls('w')}><w:instrText> PRINTDATE </w:instrText></w:r>"
    )
    end = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
    for run in (begin, instr, end):
        printed._p.append(run)
    document.add_heading("Context", level=1)
    return to_bytes(document)


@pytest.mark.parametrize(
    ("op", "message"),
    [
        ({"op": "remove", "from": "P1"}, "table of contents"),
        ({"op": "write", "from": "P1", "markdown": "1. Context"}, "table of contents"),
        (
            {"op": "set_text", "address": "P2", "text": "Printed on {DATE}"},
            "field Word computes",
        ),
    ],
)
def test_the_model_never_rewrites_what_word_computes(op, message) -> None:
    with pytest.raises(BatchError, match=message):
        apply_batch(_toc_and_footer_document(), [op], protected_tables=frozenset())


def test_the_collaborator_and_a_replay_are_not_limited() -> None:
    # Only the model's batches are guarded; a stored log always replays.
    apply_batch(
        _toc_and_footer_document(),
        [{"op": "set_text", "address": "P2", "text": "Printed on {DATE}"}],
    )


def _numbered_headings_document() -> bytes:
    """Word numbers the headings itself (numPr): typed numbers would double."""
    document = Document()
    for level, text in ((1, "Organisation"), (2, "Scope"), (3, "Skills")):
        heading = document.add_heading(text, level=level)
        heading._p.get_or_add_pPr().append(
            parse_xml(
                f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level - 1}"/>'
                '<w:numId w:val="1"/></w:numPr>'
            )
        )
        document.add_paragraph("Describe it here.")
    return to_bytes(document)


def test_text_written_over_an_unstyled_instruction_takes_the_main_body_style() -> None:
    document = Document()
    document.styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
    for text in ("One.", "Two.", "Three."):
        document.add_paragraph(text, style="Body")
    document.add_paragraph("Paste the presentation here.")
    out = Document(
        io.BytesIO(
            apply_batch(
                to_bytes(document),
                [
                    {
                        "op": "write",
                        "from": "P4",
                        "markdown": "Written.\n\n- **Point** : one.",
                    }
                ],
            )
        )
    )
    written = [p for p in out.paragraphs if p.text in ("Written.", "Point : one.")]
    assert len(written) == 2
    assert written[0].style is not None and written[0].style.name == "Body"
    # The style's font, not the document default's stamped over it.
    assert not any(r.font.name for p in written for r in p.runs)


def test_the_outline_shows_the_numbers_word_gives_the_headings() -> None:
    line = outline(build_map(_numbered_headings_document()), "1")
    assert line == (
        'OUTLINE of chapter 1 now: 1 "Organisation"; sub-headings: '
        '1.1 "Scope";   1.1.1 "Skills"'
    )


def test_typed_numbers_are_dropped_from_numbered_headings() -> None:
    docx = _numbered_headings_document()
    fitted = fit_to_document(
        [
            {"op": "set_text", "address": "P1", "text": "7 Organisation"},
            {
                "op": "insert_after",
                "address": "P6",
                "markdown": "### **4. Soft skills**\n\nText.",
            },
        ],
        docx,
    )
    assert fitted[0]["text"] == "Organisation"
    assert fitted[1]["markdown"].startswith("### Soft skills\n")


def test_written_text_does_not_repeat_the_heading_above() -> None:
    fitted = fit_to_document(
        [
            {
                "op": "write",
                "from": "P4",
                "markdown": "**1.1 Scope**\n\nThe scope is wide.",
            }
        ],
        _numbered_headings_document(),
    )
    assert fitted[0]["markdown"] == "The scope is wide."


def test_a_heading_line_wrapped_in_bold_is_a_heading() -> None:
    op = normalized(
        {"op": "write", "from": "P4", "markdown": "**## Context**\n\nText."}
    )
    assert op["markdown"] == "## Context\n\nText."


def test_a_sub_section_title_goes_into_the_template_sub_heading_above() -> None:
    docx = _numbered_headings_document()
    # P3 is the level-2 "Scope", P4 its body: the title renames P3.
    fitted = fit_to_document(
        [
            normalized(
                {"op": "write", "from": "P4", "markdown": "## 2.1 Perimeter\n\nText."}
            )
        ],
        docx,
    )
    assert fitted == [
        {"op": "set_text", "address": "P3", "text": "Perimeter"},
        {"op": "write", "from": "P4", "markdown": "Text."},
    ]
    # Under a chapter heading, or sent over a heading: no silent chapter rename,
    # but a heading line sent over a heading is that heading's title.
    chapter = {"op": "write", "from": "P2", "markdown": "# Other\n\nText."}
    assert fit_to_document([chapter], docx) == [chapter]
    over = normalized({"op": "set_text", "address": "P3", "text": "## New scope"})
    assert fit_to_document([over], docx) == [
        {"op": "set_text", "address": "P3", "text": "New scope"}
    ]
    # Whatever its "#" count: a chapter heading (P1, level 1) is renamed too.
    chapter = normalized({"op": "set_text", "address": "P1", "text": "## Planning"})
    assert fit_to_document([chapter], docx) == [
        {"op": "set_text", "address": "P1", "text": "Planning"}
    ]


def test_a_one_cell_frame_is_not_a_template_block() -> None:
    document = Document()
    document.add_heading("Budget", level=1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "image"
    document.add_paragraph("Caption")
    docx = to_bytes(document)
    assert template_table_signatures(docx) == frozenset()
    table = "| Item | Cost |\n|---|---|\n| Licences | 300 |"
    written = apply_batch(
        docx,
        [{"op": "write", "from": "T1", "to": "P2", "markdown": table}],
        protected_tables=template_table_signatures(docx),
    )
    cells = [e.text for e in build_map(written).elements if e.kind == "cell"]
    assert "Licences" in cells


def test_replacing_a_whole_paragraph_with_markdown_writes_it() -> None:
    fitted = fit_to_document(
        [
            {
                "op": "replace_text",
                "address": "P2",
                "find": "Describe it here.",
                "replace": "First.\n\n- one\n- two",
            }
        ],
        _numbered_headings_document(),
    )
    assert fitted[0] == {
        "op": "write",
        "from": "P2",
        "markdown": "First.\n\n- one\n- two",
    }


def test_a_horizontal_rule_is_dropped_and_a_heading_line_is_a_write():
    docx = to_bytes(report_document())
    after = build_map(
        apply_batch(
            docx, [{"op": "write", "from": "P5", "markdown": "One.\n\n---\n\nTwo."}]
        )
    )
    assert "---" not in [e.text for e in after.elements]
    for text in ("### Phase 1", "---"):
        assert (
            normalized({"op": "set_text", "address": "P5", "text": text})["op"]
            == "write"
        )


def test_only_the_template_headings_are_kept_from_a_model_write():
    template = to_bytes(report_document())
    # The model added a sub-heading under "Context" in an earlier edit.
    added = apply_batch(
        template, [{"op": "write", "from": "P5", "markdown": "## Added\n\nText."}]
    )
    doc_map = build_map(added)
    sub = next(e for e in doc_map.elements if e.text == "Added")
    shape = {
        "protected_tables": template_table_signatures(template),
        "template_headings": template_heading_texts(template),
    }
    # Its own heading may be rewritten; the template's may not.
    apply_batch(added, [{"op": "write", "from": sub.address, "markdown": "x"}], **shape)
    with pytest.raises(BatchError, match="keep the template's headings"):
        apply_batch(
            added, [{"op": "write", "from": "P4", "to": "P5", "markdown": "x"}], **shape
        )


def _model(docx: bytes, *ops: dict) -> bytes:
    return apply_batch(docx, list(ops), protected_tables=frozenset())


def test_a_footer_word_never_shows_is_not_a_place() -> None:
    document = Document()
    document.add_heading("Chapter", level=1)
    section = document.sections[0]
    section.footer.paragraphs[0].text = "Shown"
    section.first_page_footer.paragraphs[0].text = "Hidden"
    docx = to_bytes(document)
    footer = next(e for e in build_map(docx).elements if e.address == "F1f")
    assert "hidden" in footer.flags
    with pytest.raises(BatchError, match="the one shown there is F1"):
        _model(docx, {"op": "set_text", "address": "F1f.P1", "text": "Name"})
    _model(docx, {"op": "set_text", "address": "F1.P1", "text": "Name"})


def test_a_paragraph_holding_only_an_image_is_not_a_place_for_text() -> None:
    document = Document()
    document.add_paragraph().add_run().add_picture(io.BytesIO(PNG))
    docx = to_bytes(document)
    with pytest.raises(BatchError, match="holds only an image"):
        _model(docx, {"op": "set_text", "address": "P1", "text": "Name"})


def test_values_are_never_moved_between_paragraphs() -> None:
    document = Document()
    for text in ("Report", "28/10/2026", "Owner"):
        document.add_paragraph(text)
    docx = to_bytes(document)
    with pytest.raises(BatchError, match="moves the text of P2 into P3"):
        _model(
            docx,
            {"op": "set_text", "address": "P2", "text": "Jane Doe"},
            {"op": "set_text", "address": "P3", "text": "28/10/2026"},
        )


def test_the_title_of_the_table_of_contents_is_kept() -> None:
    document = Document()
    document.styles.add_style("Contents title", WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Cover")
    document.add_paragraph("Contents", style="Contents title")
    document.add_paragraph("1 Context", style="toc 1")
    document.add_heading("Context", level=1)
    docx = to_bytes(document)
    with pytest.raises(BatchError, match="table of contents"):
        _model(docx, {"op": "set_text", "address": "P2", "text": "28/10/2026"})
    _model(docx, {"op": "set_text", "address": "P1", "text": "New cover"})


def test_a_heading_titled_in_the_batch_gets_no_second_title() -> None:
    title = {"op": "set_text", "address": "P3", "text": "Architecture"}
    write = {"op": "write", "from": "P4", "markdown": "## Architecture of it\n\nText."}
    fitted = fit_to_document([title, write], _numbered_headings_document())
    assert fitted == [title, {"op": "write", "from": "P4", "markdown": "Text."}]


def _chapters(docx: bytes) -> list[str]:
    return [e.text for e in build_map(docx).elements if e.heading_level == 1]


def test_a_chapter_moves_with_its_whole_section() -> None:
    document = Document()
    for title in ("A", "B", "C"):
        document.add_heading(title, level=1)
        document.add_paragraph(f"{title} text")
    docx = to_bytes(document)
    # P5 is the heading "C": its section moves before "B" (P3), or "A"
    # moves after the whole section of "C".
    moved = apply_batch(docx, [{"op": "move", "from": "P5", "before": "P3"}])
    assert _chapters(moved) == ["A", "C", "B"]
    assert _text(moved, "P4") == "C text"
    after = apply_batch(docx, [{"op": "move", "from": "P1", "after": "P5"}])
    assert _chapters(after) == ["B", "C", "A"]
    with pytest.raises(BatchError, match="inside the moved range"):
        apply_batch(docx, [{"op": "move", "from": "P1", "to": "P4", "before": "P3"}])
    # "A" after the section just before it (none: it is first) / "B" after "A".
    with pytest.raises(BatchError, match="already there"):
        apply_batch(docx, [{"op": "move", "from": "P3", "after": "P1"}])
    with pytest.raises(BatchError, match="batch of its own"):
        apply_batch(
            docx,
            [
                {"op": "move", "from": "P5", "before": "P3"},
                {"op": "set_text", "address": "P2", "text": "x"},
            ],
        )


def test_a_typed_number_alone_is_a_renumbering() -> None:
    docx = _numbered_headings_document()
    retyped = {"op": "set_text", "address": "P1", "text": "3. Organisation"}
    renamed = {"op": "set_text", "address": "P1", "text": "3. Planning"}
    assert inspect_set_texts([retyped], docx).renumbered == ["P1"]
    assert inspect_set_texts([renamed], docx).renumbered == []


def test_set_text_findings_for_typed_titles_and_unchanged_text() -> None:
    document = Document()
    document.add_heading("Chapter", level=1)
    document.add_paragraph("3.2 Integrations")
    document.add_paragraph("Text.")
    docx = to_bytes(document)
    found = inspect_set_texts(
        [
            {"op": "set_text", "address": "P2", "text": "Integrations"},
            {"op": "set_text", "address": "P3", "text": "Text."},
        ],
        docx,
    )
    assert found.typed_titles == ["P2"] and found.unchanged == ["P3"]


def test_written_text_never_takes_a_caption_style() -> None:
    document = Document()
    document.add_heading("Chapter", level=1)
    document.add_paragraph()
    document.add_paragraph("Caption", style="Caption")
    docx = to_bytes(document)
    written = apply_batch(
        docx, [{"op": "write", "from": "P2", "to": "P3", "markdown": "Body.\n\nMore."}]
    )
    styles = {e.text: e.style for e in build_map(written).elements}
    assert (styles["Body."] or "").lower() != "caption"


def test_a_renamed_chapter_heading_is_never_written_over() -> None:
    template = to_bytes(report_document())
    shape = {
        "protected_tables": template_table_signatures(template),
        "template_headings": template_heading_texts(template),
    }
    renamed = apply_batch(
        template, [{"op": "set_text", "address": "P7", "text": "Planning"}], **shape
    )
    with pytest.raises(BatchError, match="Rename a heading with set_text"):
        apply_batch(
            renamed,
            [{"op": "write", "from": "P7", "to": "P8", "markdown": "Text."}],
            **shape,
        )


def test_a_written_heading_takes_the_template_s_own_style() -> None:
    document = Document()
    style = cast(
        ParagraphStyle,
        document.styles.add_style("Title2-Contract", WD_STYLE_TYPE.PARAGRAPH),
    )
    style.base_style = document.styles["Heading 2"]
    document.add_heading("Chapter", level=1)
    document.add_paragraph("Template sub-heading", style="Title2-Contract")
    document.add_paragraph("Text.")
    docx = to_bytes(document)
    template_sub = build_map(docx).get("P2")
    assert template_sub is not None and template_sub.heading_level == 2
    written = apply_batch(
        docx, [{"op": "write", "from": "P3", "markdown": "## Added\n\nMore."}]
    )
    added = next(e for e in build_map(written).elements if e.text == "Added")
    assert added.heading_level == 2 and added.style == "Title2-Contract"
