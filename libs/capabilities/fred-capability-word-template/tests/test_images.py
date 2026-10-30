# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Images in the document:
extraction from an attached Word file, the image routes, the
prompt fragment, placing an image with ``write`` and reading it back in the
"Rédaction" view - offline (fake store, synthetic documents)."""

from __future__ import annotations

import asyncio
import io
import random
import struct
import zlib
from collections.abc import Iterator
from typing import Any, cast

import pytest
from conftest import PNG, load, to_bytes
from docx import Document
from docx.oxml.ns import qn
from fastapi import FastAPI
from fastapi.testclient import TestClient
from fred_capability_word_template import documents as documents_module
from fred_capability_word_template import images as images_module
from fred_capability_word_template import router as router_module
from fred_capability_word_template import store as store_module
from fred_capability_word_template.capability import (
    WordTemplateCapability,
)
from fred_capability_word_template.config import (
    WordTemplateConfig,
)
from fred_capability_word_template.content import (
    ChapterBlock,
    TextSegment,
    build_content_parts,
)
from fred_capability_word_template.images import (
    batch_image_ids,
    extract_docx_images,
    images_fragment,
)
from fred_capability_word_template.middleware import (
    WordTemplateMiddleware,
)
from fred_capability_word_template.operations import BatchError, apply_batch
from fred_capability_word_template.router import build_router
from fred_capability_word_template.session_document import content_sha256
from fred_capability_word_template.store import BaseRef, ImageInfo, Writer
from fred_core import KeycloakUser, get_current_user
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
)
from fred_sdk.contracts.runtime import RuntimeServices
from langchain_core.messages import SystemMessage
from port_fakes import FakeAssets, FakeWordTemplateStore

_PREFIX = "/capabilities/word_template"
_AUTH = {"Authorization": "Bearer token-1"}
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _big_png(seed: int = 1, side: int = 48) -> bytes:
    """A valid PNG of random pixels, well above the 2 KB extraction floor."""
    rng = random.Random(seed)

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", side, side, 8, 2, 0, 0, 0)
    raw = b"".join(
        b"\x00" + bytes(rng.randrange(256) for _ in range(side * 3))
        for _ in range(side)
    )
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


DIAGRAM = _big_png()


def _docx_with_picture() -> bytes:
    """A heading, a diagram with its caption, a tiny icon (ignored)."""
    document = Document()
    document.add_heading("Architecture", level=1)
    document.add_paragraph("The target architecture is shown below.")
    document.add_picture(io.BytesIO(DIAGRAM))
    document.add_paragraph("Figure 1 : flux applicatifs", style="Caption")
    document.add_picture(io.BytesIO(PNG))  # far below 2 KB
    return to_bytes(document)


def _report() -> bytes:
    document = Document()
    document.add_heading("Context", level=1)
    document.add_paragraph("Describe the context here.")
    document.add_heading("Annex", level=1)
    document.add_paragraph("Annex text.")
    return to_bytes(document)


@pytest.fixture(autouse=True)
def _clean_caches() -> Iterator[None]:
    documents_module._BASES.clear()
    documents_module._STATES.clear()
    documents_module.SEEN.clear()
    yield


@pytest.fixture()
def fake() -> Iterator[FakeWordTemplateStore]:
    store = FakeWordTemplateStore()
    store_module.set_store_provider(lambda: store)
    try:
        yield store
    finally:
        store_module.clear_store_provider()


@pytest.fixture()
def http(fake) -> TestClient:
    app = FastAPI()
    app.include_router(build_router(), prefix=_PREFIX)
    app.dependency_overrides[get_current_user] = lambda: KeycloakUser(
        uid="u-1", username="u-1", roles=[]
    )
    return TestClient(app)


# --- extraction ---------------------------------------------------------------------


def test_extraction_keeps_pictures_with_the_text_next_to_them():
    found = extract_docx_images(_docx_with_picture())
    assert len(found) == 1  # the icon is too small
    assert found[0].content == DIAGRAM
    assert found[0].content_type == "image/png"
    assert found[0].context_text == "Figure 1 : flux applicatifs"
    assert found[0].position == 1


def test_extraction_falls_back_to_the_text_before_the_picture():
    document = Document()
    document.add_heading("Architecture", level=1)
    document.add_picture(io.BytesIO(DIAGRAM))
    (found,) = extract_docx_images(to_bytes(document))
    assert found.context_text == "Architecture"


def test_attached_docx_pictures_become_session_images_once(http, fake):
    docx = _docx_with_picture()
    for _ in range(2):
        response = http.post(
            f"{_PREFIX}/sessions/s-1/attachments",
            files={"file": ("Archi.docx", docx, _DOCX)},
            headers=_AUTH,
        )
        assert response.status_code == 201
    (info,) = asyncio.run(fake.list_images("s-1"))
    assert info.image_id == "img1"
    assert info.filename == "Archi.docx - image 1"
    assert info.context_text == "Figure 1 : flux applicatifs"
    assert info.source_attachment_id == "doc1"
    (attachment,) = asyncio.run(fake.list_attachments("s-1"))
    assert attachment.images_extracted is True


@pytest.mark.asyncio
async def test_attachments_kept_before_images_are_backfilled_on_first_read(fake):
    docx = _docx_with_picture()
    await fake.put_attachment("s-1", "Old.docx", content_sha256(docx), docx)
    images = await images_module.session_images("s-1")
    assert [i.image_id for i in images] == ["img1"]
    (attachment,) = await fake.list_attachments("s-1")
    assert attachment.images_extracted is True
    # Flagged: the file is not parsed again.
    fake.images.clear()
    assert await images_module.session_images("s-1") == []


# --- image routes -------------------------------------------------------------------


def _upload_image(http: TestClient, content: bytes, name: str = "schema.png"):
    return http.post(
        f"{_PREFIX}/sessions/s-1/images",
        files={"file": (name, content, "image/png")},
        headers=_AUTH,
    )


def test_image_upload_dedups_and_serves_the_bytes(http):
    response = _upload_image(http, DIAGRAM)
    assert response.status_code == 201
    assert response.json() == {
        "image_id": "img1",
        "filename": "schema.png",
        "content_type": "image/png",
        "size": len(DIAGRAM),
    }
    assert _upload_image(http, DIAGRAM, "again.png").json()["image_id"] == "img1"
    assert _upload_image(http, _big_png(2)).json()["image_id"] == "img2"
    got = http.get(f"{_PREFIX}/sessions/s-1/images/img1", headers=_AUTH)
    assert got.status_code == 200
    assert got.content == DIAGRAM and got.headers["content-type"] == "image/png"
    missing = http.get(f"{_PREFIX}/sessions/s-1/images/img9", headers=_AUTH)
    assert missing.status_code == 404


def test_image_upload_rejects_other_types_and_big_files(http, fake, monkeypatch):
    # Type by content, never by the declared name or MIME type.
    assert _upload_image(http, b"%PDF-1.7 not an image").status_code == 415
    monkeypatch.setattr(router_module, "MAX_IMAGE_BYTES", len(DIAGRAM) - 1)
    assert _upload_image(http, DIAGRAM).status_code == 413
    assert fake.images == {}


# --- placing an image -----------------------------------------------------------------


def test_batch_image_ids_reads_write_and_insert_after_only():
    batches = [
        [
            {"op": "write", "from": "P2", "markdown": "Text\n\n![Flux](img:img1)"},
            {"op": "insert_after", "address": "P3", "markdown": "![x](img:img2)"},
            {"op": "set_text", "address": "P1", "text": "![y](img:img3)"},
        ]
    ]
    assert batch_image_ids(batches) == {"img1", "img2"}


def _pictures(docx: bytes) -> list[tuple[str, str]]:
    names = []
    for node in load(docx).element.body.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    ):
        names.append((node.get("name"), node.get("descr")))
    return names


def test_write_places_the_image_and_the_view_reads_it_back():
    markdown = "Intro text.\n\n![Flux applicatifs](img:img1)\n\nAfter the figure."
    docx = apply_batch(
        _report(),
        [{"op": "write", "from": "P2", "markdown": markdown}],
        {"img1": DIAGRAM},
    )
    assert _pictures(docx) == [("img:img1", "Flux applicatifs")]
    captions = [
        p.text
        for p in load(docx).paragraphs
        if p.style is not None and p.style.name == "Caption"
    ]
    # python-docx's text skips the SEQ field holding the number.
    assert captions == ["Figure  : Flux applicatifs"]
    segment = _first_segment(docx)
    assert segment.markdown == markdown
    # Round trip: writing the view's Markdown back gives the same document.
    again = apply_batch(
        docx,
        [
            segment.edit.model_copy(update={"markdown": segment.markdown}).model_dump(
                by_alias=True, exclude_none=True
            )
        ],
        {"img1": DIAGRAM},
    )
    assert _first_segment(again).markdown == markdown


def _first_segment(docx: bytes) -> TextSegment:
    chapter = build_content_parts(docx)[0].blocks[0]
    assert isinstance(chapter, ChapterBlock)
    (segment,) = chapter.body
    assert isinstance(segment, TextSegment)
    return segment


def test_an_unknown_image_is_refused():
    with pytest.raises(BatchError) as error:
        apply_batch(
            _report(),
            [{"op": "write", "from": "P2", "markdown": "![x](img:img7)"}],
            {},
        )
    assert "img7" in error.value.problems[0]


def _ctx() -> CapabilityContext:
    return CapabilityContext(
        identity=CapabilityIdentity(
            user_id="u-1",
            session_id="s-1",
            team_id="team-1",
            agent_instance_id="agent-1",
        ),
        config=WordTemplateConfig(template_library=[]),
        turn_options=EmptyModel(),
        services=RuntimeServices(agent_assets=FakeAssets()),
    )


async def _attachment_base(fake: FakeWordTemplateStore) -> None:
    base = _report()
    await fake.put_attachment("s-1", "report.docx", content_sha256(base), base)
    await fake.mark_images_extracted("s-1", "doc1")
    await fake.select_base(
        "s-1", BaseRef("attachment", "doc1"), Writer("u-1"), expected_count=0
    )


@pytest.mark.asyncio
async def test_edit_document_places_a_session_image_and_exports_keep_it(fake, http):
    await _attachment_base(fake)
    await fake.put_image(
        "s-1",
        filename="schema.png",
        content_type="image/png",
        sha256=content_sha256(DIAGRAM),
        content=DIAGRAM,
    )
    tools: dict[str, Any] = {t.name: t for t in WordTemplateCapability().tools(_ctx())}
    await tools["read_template"].coroutine(part="1")
    content, artifact = await tools["edit_document"].coroutine(
        operations=[
            {"op": "write", "from": "P2", "markdown": "![Flux applicatifs](img:img1)"}
        ]
    )
    assert artifact.is_error is False, content
    # A replay from the base (no snapshot) reads the image from the store.
    fake.snapshots.clear()
    response = http.get(f"{_PREFIX}/sessions/s-1/document.docx", headers=_AUTH)
    assert response.status_code == 200
    assert _pictures(response.content) == [("img:img1", "Flux applicatifs")]
    view = http.get(f"{_PREFIX}/sessions/s-1/content", headers=_AUTH).json()
    segment = view["parts"][0]["blocks"][0]["body"][0]
    assert segment["markdown"] == "![Flux applicatifs](img:img1)"
    # The collaborator's edit route resolves images the same way.
    edit = http.post(
        f"{_PREFIX}/sessions/s-1/edits",
        json={
            "base_seq": view["seq"],
            "operations": [
                {**segment["edit"], "markdown": "Text\n\n" + segment["markdown"]}
            ],
        },
        headers=_AUTH,
    )
    assert edit.status_code == 200, edit.text


# --- prompt -------------------------------------------------------------------------------


def test_images_fragment_lists_ids_names_and_near_text():
    assert images_fragment([]) is None
    text = images_fragment(
        [
            ImageInfo("img1", "Archi.docx - image 1", "image/png", "s", 1, "Flux"),
            ImageInfo("img2", "photo.jpg", "image/jpeg", "t", 1),
        ]
    )
    assert text is not None
    assert '- img1 - Archi.docx - image 1 - near: "Flux"' in text
    assert "- img2 - photo.jpg" in text
    assert "![caption](img:<id>)" in text
    assert "Never say an image was placed unless" in text


class _Request:
    def __init__(self) -> None:
        self.model = None
        self.system_prompt = "AGENT PROMPT"
        self.overridden: SystemMessage | None = None

    def override(self, *, system_message: SystemMessage) -> _Request:
        self.overridden = system_message
        return self


@pytest.mark.asyncio
async def test_prompt_lists_images_backfilled_from_an_attachment(fake):
    await _attachment_base(fake)
    docx = _docx_with_picture()
    await fake.put_attachment("s-1", "Archi.docx", content_sha256(docx), docx)
    middleware = WordTemplateMiddleware(_ctx())

    async def handler(request):
        return request

    result = cast(
        _Request,
        await middleware.awrap_model_call(_Request(), handler),  # type: ignore[arg-type]
    )
    assert result.overridden is not None
    prompt = str(result.overridden.content)
    assert "IMAGES AVAILABLE FOR THE DOCUMENT" in prompt
    assert 'img1 - Archi.docx - image 1 - near: "Figure 1 : flux applicatifs"' in prompt


def test_rendered_picture_is_tagged_in_the_map():
    from fred_capability_word_template.docmap import build_map

    docx = apply_batch(
        _report(),
        [{"op": "write", "from": "P2", "markdown": "![A](img:img1)"}],
        {"img1": DIAGRAM},
    )
    flagged = [e for e in build_map(docx).elements if "image" in e.flags]
    assert flagged and "img:img1" in flagged[0].flags


def test_two_pictures_in_one_write_get_distinct_drawing_ids():
    document = Document()
    document.add_paragraph("Before.")
    markdown = "![First](img:img1)\n\n![Second](img:img1)"
    out = apply_batch(
        to_bytes(document),
        [{"op": "write", "from": "P1", "markdown": markdown}],
        {"img1": PNG},
    )
    ids = [e.get("id") for e in load(out).element.body.iter(qn("wp:docPr"))]
    assert len(ids) == 2 and len(set(ids)) == 2
