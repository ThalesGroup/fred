# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The "Rédaction" view: the structured view
of synthetic documents, and the docx -> Markdown converter's round trip
through ``write`` / ``set_text``."""

from __future__ import annotations

import pytest
from conftest import add_toc_entry, form_document, to_bytes
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from fred_capability_word_template.content import (
    BlockGroup,
    ChapterBlock,
    FieldsBlock,
    ParagraphBlock,
    RepeatableGroup,
    TableBlock,
    TextSegment,
    build_content_parts,
)
from fred_capability_word_template.inline import parse_inline_markdown
from fred_capability_word_template.operations import apply_batch


def _chapters_document():
    document = Document()
    document.add_paragraph("Cover")
    document.add_heading("Intro", level=1)
    document.add_paragraph("Intro text.")
    document.add_paragraph("point one", style="List Bullet")
    document.add_heading("Sub", level=2)
    paragraph = document.add_paragraph("Sub with ")
    paragraph.add_run("bold").bold = True
    document.add_heading("Next", level=1)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "wide"
    table.cell(1, 2).merge(table.cell(2, 2)).text = "tall"
    table.cell(1, 0).text = "a"
    document.add_heading("Empty", level=1)
    return document


def _add_page_field(paragraph) -> None:
    paragraph._p.append(
        parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>3</w:t></w:r>'
            "</w:fldSimple>"
        )
    )


def _blocks(parts, part_id: str):
    return next(p for p in parts if p.id == part_id).blocks


def test_chapters_and_sub_chapters_are_blocks_in_reading_order():
    parts = build_content_parts(to_bytes(_chapters_document()))
    assert [p.id for p in parts] == ["0", "1", "2", "3"]
    cover = _blocks(parts, "0")
    assert isinstance(cover[0], ParagraphBlock) and cover[0].text == "Cover"
    intro, sub = _blocks(parts, "1")
    assert isinstance(intro, ChapterBlock) and isinstance(sub, ChapterBlock)
    assert (intro.heading.text, intro.level, sub.level) == ("Intro", 1, 2)
    # The body stops at the sub-chapter's heading; the removal covers it.
    (segment,) = intro.body
    assert isinstance(segment, TextSegment)
    assert (segment.from_, segment.to) == ("P3", "P4")
    assert segment.markdown == "Intro text.\n\n- point one"
    assert segment.edit.op == "write" and segment.edit.to == "P4"
    assert (intro.from_, intro.to) == ("P2", "P6")
    assert intro.remove.model_dump(by_alias=True, exclude_none=True) == {
        "op": "remove",
        "from": "P2",
        "to": "P6",
    }
    assert sub.body[0].markdown == "Sub with **bold**"  # type: ignore[union-attr]
    (empty,) = _blocks(parts, "3")
    assert isinstance(empty, ChapterBlock)
    (placeholder,) = empty.body
    assert isinstance(placeholder, TextSegment) and placeholder.from_ is None
    assert placeholder.edit.op == "insert_after"
    assert placeholder.edit.address == empty.heading.address


def test_a_table_with_merged_cells_lists_each_cell_once():
    parts = build_content_parts(to_bytes(_chapters_document()))
    (chapter,) = _blocks(parts, "2")
    assert isinstance(chapter, ChapterBlock)
    (table,) = chapter.body
    assert isinstance(table, TableBlock)
    assert (table.address, table.n_rows, table.n_cols) == ("T1", 3, 3)
    first = table.rows[0].cells
    assert [(c.address, c.col_span, c.text) for c in first] == [
        ("T1.R1.C1", 2, "wide"),
        ("T1.R1.C3", 1, ""),
    ]
    tall = next(c for c in table.rows[1].cells if c.text == "tall")
    assert (tall.address, tall.row_span, tall.edit) == ("T1.R2.C3", 2, "set_text")
    # Merged-cell continuation rows are listed without the merged cell.
    assert [c.address for c in table.rows[2].cells] == ["T1.R3.C1", "T1.R3.C2"]
    # Rows inside a vertical merge cannot be duplicated or removed alone.
    assert table.rows[0].duplicate is not None
    assert table.rows[1].duplicate is None and table.rows[2].remove is None
    assert table.duplicate is not None and table.duplicate.op == "repeat"
    assert table.remove is not None and table.remove.from_ == "T1"


def test_a_heading_less_form_is_tables_and_paragraphs():
    parts = build_content_parts(to_bytes(form_document()))
    assert [p.id for p in parts] == ["1", "hf"]
    kinds = [b.kind for b in _blocks(parts, "1")]
    assert kinds == ["table", "paragraph", "table"]
    block = _blocks(parts, "1")[2]
    assert isinstance(block, TableBlock)
    multi = block.rows[3].cells[0]
    assert (multi.address, multi.col_span, multi.edit) == ("T2.R4.C1", 4, "write")
    assert multi.markdown == "old comment\n\nsecond line"
    header, footer = _blocks(parts, "hf")
    assert isinstance(header, FieldsBlock) and isinstance(footer, FieldsBlock)
    assert [(f.address, f.text, f.edit) for f in header.fields] == [
        ("H1.P1", "Header text", "set_text")
    ]


def test_a_footer_page_field_is_edited_with_replace_text_only():
    document = form_document()
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "Page "
    _add_page_field(footer)
    parts = build_content_parts(to_bytes(document))
    footer_block = _blocks(parts, "hf")[1]
    assert isinstance(footer_block, FieldsBlock)
    (place,) = footer_block.fields
    assert (place.address, place.text, place.edit) == (
        "F1.P1",
        "Page 3",
        "replace_text",
    )
    assert place.markdown is None
    assert [(p.text, p.field) for p in place.pieces or []] == [
        ("Page ", False),
        ("3", True),
    ]


def test_a_body_paragraph_with_a_field_breaks_the_chapter_text():
    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("before")
    _add_page_field(document.add_paragraph("Date: "))
    document.add_paragraph("after")
    (chapter,) = _blocks(build_content_parts(to_bytes(document)), "1")
    assert isinstance(chapter, ChapterBlock)
    assert [s.kind for s in chapter.body] == ["text", "paragraph", "text"]
    assert chapter.body[1].edit == "replace_text"  # type: ignore[union-attr]


def test_repeatable_groups_of_the_analysis():
    groups = [
        RepeatableGroup("T1", "T2", "one project"),
        RepeatableGroup("T2.R2", "T2.R4", "one line"),
    ]
    blocks = _blocks(build_content_parts(to_bytes(form_document()), groups=groups), "1")
    (group,) = blocks
    assert isinstance(group, BlockGroup)
    assert (group.from_, group.to, group.label) == ("T1", "T2", "one project")
    assert [b.kind for b in group.blocks] == ["table", "paragraph", "table"]
    assert group.duplicate.model_dump(by_alias=True, exclude_none=True) == {
        "op": "repeat",
        "from": "T1",
        "to": "T2",
        "count": 1,
    }
    table = group.blocks[2]
    assert isinstance(table, TableBlock)
    (rows,) = table.row_groups
    assert (rows.from_, rows.to, rows.duplicate.count) == ("T2.R2", "T2.R4", 1)


def test_a_single_table_group_labels_the_table():
    groups = [RepeatableGroup("T2", "T2", "one project")]
    blocks = _blocks(build_content_parts(to_bytes(form_document()), groups=groups), "1")
    assert isinstance(blocks[2], TableBlock) and blocks[2].repeat_label == "one project"


# --- Markdown round trip -------------------------------------------------------------

_ROUND_TRIP = [
    "First paragraph with **bold**, *italic* and `code`.\n\nSecond paragraph.",
    "- one\n- two\n  - nested\n\n1. first\n2. second",
    "Intro:\n\n1. a\n2. b\n\nAfter the list.",
    r"Price 2 \* 3, snake\_case, a \`tick\` and back\\slash.",
    "\\# not a heading\n\n1\\. not a list\n\n\\- not a bullet\n\n\\![not](img:x)",
    "[À COMPLÉTER : la date] then **bold [À COMPLÉTER : x]** text",
    "Tab\there and **bold** *it*",
]


def _context_segment(docx: bytes) -> TextSegment:
    chapter = next(
        b
        for part in build_content_parts(docx)
        for b in part.blocks
        if isinstance(b, ChapterBlock) and b.heading.text == "Context"
    )
    (segment,) = chapter.body
    assert isinstance(segment, TextSegment)
    return segment


@pytest.mark.parametrize("markdown", _ROUND_TRIP)
def test_markdown_write_markdown_is_stable(report_bytes, markdown):
    first = _context_segment(report_bytes)
    written = apply_batch(
        report_bytes,
        [
            {
                **first.edit.model_dump(by_alias=True, exclude_none=True),
                "markdown": markdown,
            }
        ],
    )
    again = _context_segment(written)
    assert again.markdown == markdown
    # And once more: the view's own Markdown is a fixed point.
    rewritten = apply_batch(
        written,
        [
            {
                **again.edit.model_dump(by_alias=True, exclude_none=True),
                "markdown": again.markdown,
            }
        ],
    )
    assert _context_segment(rewritten).markdown == markdown


@pytest.mark.parametrize(
    "markdown", ["**Week** 41", "line one\nline two", r"a \* b and c\_d", ""]
)
def test_set_text_round_trip_on_a_cell(form_bytes, markdown):
    written = apply_batch(
        form_bytes, [{"op": "set_text", "address": "T1.R1.C2", "text": markdown}]
    )
    blocks = _blocks(build_content_parts(written), "1")
    cell = blocks[0].rows[0].cells[1]  # type: ignore[union-attr]
    assert (cell.edit, cell.markdown) == ("set_text", markdown)


def test_escaped_markers_are_literal_in_the_inline_parser():
    spans = parse_inline_markdown(r"\*not italic\* and **bold \* star** \_x\_")
    assert [(s.text, s.bold, s.italic) for s in spans] == [
        ("*not italic* and ", False, False),
        ("bold * star", True, False),
        (" _x_", False, False),
    ]


def test_a_tabbed_paragraph_keeps_its_tab():
    document = Document()
    document.add_heading("H", level=1)
    paragraph = document.add_paragraph("a")
    paragraph._p.append(OxmlElement("w:r"))
    run = paragraph._p[-1]
    run.append(OxmlElement("w:tab"))
    text = OxmlElement("w:t")
    text.text = "b"
    run.append(text)
    (chapter,) = _blocks(build_content_parts(to_bytes(document)), "1")
    assert chapter.body[0].markdown == "a\tb"  # type: ignore[union-attr]
    assert qn("w:tab")


def test_editing_a_chapter_keeps_its_styles_and_links():
    document = Document()
    document.add_heading("Context", level=1)
    document.add_paragraph("Plain text.")
    document.add_paragraph("A quoted sentence.", style="Quote")
    linked = add_toc_entry(document, "see the annex", "_Ref1")
    document.add_heading("Next", level=1)
    (chapter,) = _blocks(build_content_parts(to_bytes(document)), "1")
    assert isinstance(chapter, ChapterBlock)
    texts = [s for s in chapter.body if isinstance(s, TextSegment)]
    # Each style its own segment; the linked paragraph is no free-text segment.
    assert [t.markdown for t in texts] == ["Plain text.", "A quoted sentence."]
    (place,) = [s for s in chapter.body if isinstance(s, ParagraphBlock)]
    assert place.edit == "replace_text" and place.text == linked.text


def test_a_long_chapter_is_one_part_of_the_view():
    # Read by the model in pieces "1", "1b", "1c"; shown whole.
    document = Document()
    document.add_heading("Missions", level=1)
    for index in range(100):
        document.add_paragraph(f"line {index}")
    document.add_heading("Places", level=1)
    parts = build_content_parts(to_bytes(document), chunk_size=40)
    assert [(p.id, p.title) for p in parts] == [("1", "Missions"), ("2", "Places")]
