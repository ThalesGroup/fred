# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Engine fixes of step 1b: ``replace_text`` (run-safe, fields untouched)
and written lists taking the template's own list look."""

from __future__ import annotations

import pytest
from conftest import load, to_bytes
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fred_capability_word_template import (
    BatchError,
    ReplaceText,
    apply_batch,
    build_map,
    parse_operation,
)


def _run(paragraph, text: str | None = None, *, field: str | None = None):
    run = OxmlElement("w:r")
    if field is not None:
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), field)
        run.append(char)
    elif text is not None:
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
    paragraph._p.append(run)
    return run


def _footer_with_page_field() -> bytes:
    """Footer "Report (Week 10, 2026)  <PAGE field showing 1>", the text
    split across runs like Word does."""
    document = Document()
    document.add_paragraph("Body")
    paragraph = document.sections[0].footer.paragraphs[0]
    _run(paragraph, "Report (Week ")
    _run(paragraph, "10,")
    _run(paragraph, " 2026)  ")
    _run(paragraph, field="begin")
    instr = OxmlElement("w:r")
    code = OxmlElement("w:instrText")
    code.text = " PAGE "
    instr.append(code)
    paragraph._p.append(instr)
    _run(paragraph, field="separate")
    _run(paragraph, "1")
    _run(paragraph, field="end")
    return to_bytes(document)


def _footer_xml(docx_bytes: bytes):
    return load(docx_bytes).sections[0].footer.paragraphs[0]._p


def test_replace_text_across_runs_keeps_the_page_field() -> None:
    result = apply_batch(
        _footer_with_page_field(),
        [
            {
                "op": "replace_text",
                "address": "F1",
                "find": "Week 10",
                "replace": "Week 41",
            }
        ],
    )
    footer = _footer_xml(result)
    texts = "".join(t.text or "" for t in footer.iter(qn("w:t")))
    assert texts == "Report (Week 41, 2026)  1"
    assert [c.get(qn("w:fldCharType")) for c in footer.iter(qn("w:fldChar"))] == [
        "begin",
        "separate",
        "end",
    ]
    assert [i.text for i in footer.iter(qn("w:instrText"))] == [" PAGE "]


def test_replace_text_never_matches_the_field_result() -> None:
    with pytest.raises(BatchError) as error:
        apply_batch(
            _footer_with_page_field(),
            [
                {
                    "op": "replace_text",
                    "address": "F1.P1",
                    "find": "2026)  1",
                    "replace": "x",
                }
            ],
        )
    assert "text Word computes" in str(error.value)


def test_replace_text_in_a_cell_every_occurrence_and_highlight_dropped(
    form_bytes: bytes,
) -> None:
    document = load(form_bytes)
    cell = document.tables[1].cell(3, 0)
    cell.paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    result = apply_batch(
        to_bytes(document),
        [ReplaceText(address="T2.R4.C1", find="old comment", replace="new comment")],
    )
    out = load(result).tables[1].cell(3, 0)
    assert out.paragraphs[0].text == "new comment"
    assert out.paragraphs[0].runs[0].font.highlight_color is None
    assert out.paragraphs[1].text == "second line"
    result = apply_batch(
        result,
        [{"op": "replace_text", "address": "T2.R4.C1", "find": "e", "replace": "E"}],
    )
    assert build_map(result)["T2.R4.C1"].text == "nEw commEnt\nsEcond linE"


def test_replace_text_problems(form_bytes: bytes) -> None:
    with pytest.raises(BatchError) as error:
        apply_batch(
            form_bytes,
            [
                {"op": "replace_text", "address": "T1", "find": "a", "replace": "b"},
                {
                    "op": "replace_text",
                    "address": "H1",
                    "find": "nothing",
                    "replace": "b",
                },
                {"op": "replace_text", "address": "P1", "find": "", "replace": "b"},
            ],
        )
    problems = error.value.problems
    assert "is a table" in problems[0]
    assert "'nothing' is not in 'H1'" in problems[1] and "Header text" in problems[1]
    assert "must not be empty" in problems[2]


def test_replace_text_conflicts_with_a_set_text_on_the_same_paragraph(
    form_bytes: bytes,
) -> None:
    with pytest.raises(BatchError, match="overlap"):
        apply_batch(
            form_bytes,
            [
                {"op": "set_text", "address": "F1.P1", "text": "x"},
                {
                    "op": "replace_text",
                    "address": "F1.P1",
                    "find": "Footer",
                    "replace": "y",
                },
            ],
        )


def test_replace_text_round_trips_as_json() -> None:
    op = parse_operation(
        {"op": "replace_text", "address": "F1", "find": "a", "replace": "b"}
    )
    assert op.model_dump(by_alias=True) == {
        "op": "replace_text",
        "address": "F1",
        "find": "a",
        "replace": "b",
    }


# --- written lists take the template's own list look ------------------------------


def _num_id_of_style(document, name: str) -> str:
    num = document.styles[name].element.find(
        f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}"
    )
    assert num is not None
    return num.get(qn("w:val"))


def _list_paragraph(document, text: str, num_id: str):
    paragraph = document.add_paragraph(text, style="List Paragraph")
    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_ilvl().val = 0
    numbering.get_or_add_numId().val = int(num_id)
    return paragraph


def _template_with_own_lists() -> tuple[bytes, str, str]:
    """Body lists drawn with "List Paragraph" + a numbering set on the
    paragraph, as Word does - never with the "List Bullet" style."""
    document = Document()
    bullet_id = _num_id_of_style(document, "List Bullet")
    number_id = _num_id_of_style(document, "List Number")
    document.add_heading("Chapter", level=1)
    document.add_paragraph("Instruction to replace.")
    _list_paragraph(document, "an existing item", bullet_id)
    _list_paragraph(document, "first step", number_id)
    return to_bytes(document), bullet_id, number_id


def _style_name(paragraph) -> str | None:
    return paragraph.style.name if paragraph.style is not None else None


def _numbering(paragraph) -> tuple[str | None, str | None]:
    numbering = paragraph._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
    if numbering is None:
        return None, None
    ilvl = numbering.find(qn("w:ilvl"))
    num = numbering.find(qn("w:numId"))
    return (
        ilvl.get(qn("w:val")) if ilvl is not None else None,
        num.get(qn("w:val")) if num is not None else None,
    )


def test_written_bullets_take_the_template_list_look() -> None:
    template, bullet_id, _number_id = _template_with_own_lists()
    result = apply_batch(
        template, [{"op": "write", "from": "P2", "markdown": "- one\n  - nested"}]
    )
    paragraphs = load(result).paragraphs
    written = [p for p in paragraphs if p.text in ("one", "nested")]
    assert [_style_name(p) for p in written] == ["List Paragraph", "List Paragraph"]
    assert [_numbering(p) for p in written] == [("0", bullet_id), ("1", bullet_id)]


def test_written_numbered_list_takes_the_look_and_restarts() -> None:
    template, _bullet_id, number_id = _template_with_own_lists()
    result = apply_batch(
        template, [{"op": "write", "from": "P2", "markdown": "1. alpha\n2. beta"}]
    )
    document = load(result)
    written = [p for p in document.paragraphs if p.text in ("alpha", "beta")]
    assert {_style_name(p) for p in written} == {"List Paragraph"}
    new_id = _numbering(written[0])[1]
    assert new_id is not None and new_id != number_id
    numbering = document.part.numbering_part.element
    num = next(n for n in numbering.iter(qn("w:num")) if n.get(qn("w:numId")) == new_id)
    original = next(
        n for n in numbering.iter(qn("w:num")) if n.get(qn("w:numId")) == number_id
    )
    abstract = qn("w:abstractNumId")
    assert num.find(abstract).get(qn("w:val")) == original.find(abstract).get(
        qn("w:val")
    )
    assert num.find(f"{qn('w:lvlOverride')}/{qn('w:startOverride')}") is not None


def test_without_template_lists_the_default_list_styles_stay() -> None:
    document = Document()
    document.add_paragraph("Instruction.")
    result = apply_batch(
        to_bytes(document), [{"op": "write", "from": "P1", "markdown": "- a\n1. b"}]
    )
    assert [_style_name(p) for p in load(result).paragraphs] == [
        "List Bullet",
        "List Number",
    ]
