# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The table of contents rebuilt from the headings after each batch."""

from __future__ import annotations

import io

from conftest import add_bookmark, to_bytes
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from fred_capability_word_template import apply_batch, build_map
from fred_capability_word_template.docmap import is_toc
from fred_capability_word_template.toc import with_update_fields

_RUN = f"<w:r {nsdecls('w')}>{{}}</w:r>"


def _field(kind: str) -> str:
    return _RUN.format(f'<w:fldChar w:fldCharType="{kind}"/>')


def _toc_document() -> bytes:
    """Chapters "A", "B", "C" (P4-P9 after the three TOC lines), with a
    Word-like TOC field: field runs, then one linked line per heading."""
    document = Document()
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    titles = ("A", "B", "C")
    for index, title in enumerate(titles):
        line = document.add_paragraph(style="toc 1")
        if index == 0:
            line._p.append(parse_xml(_field("begin")))
            line._p.append(
                parse_xml(_RUN.format('<w:instrText> TOC \\o "1-2" \\h </w:instrText>'))
            )
            line._p.append(parse_xml(_field("separate")))
        line._p.append(
            parse_xml(
                f'<w:hyperlink {nsdecls("w")} w:anchor="_Toc{index}">'
                f"<w:r><w:t>{title}</w:t></w:r><w:r><w:tab/></w:r>"
                '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                f"<w:r><w:instrText> PAGEREF _Toc{index} \\h </w:instrText></w:r>"
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                f"<w:r><w:t>{index + 2}</w:t></w:r>"
                '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink>'
            )
        )
        if index == len(titles) - 1:
            line._p.append(parse_xml(_field("end")))
    for index, title in enumerate(titles):
        heading = document.add_heading(title, level=1)
        add_bookmark(heading, f"_Toc{index}", str(index + 10))
        document.add_paragraph(f"{title} text")
    return to_bytes(document)


def _toc(docx: bytes) -> list[str]:
    return [e.text for e in build_map(docx).elements if is_toc(e)]


def test_the_table_of_contents_follows_moves_renames_and_new_headings():
    docx = _toc_document()
    assert _toc(docx) == ["A 2", "B 3", "C 4"]
    moved = apply_batch(docx, [{"op": "move", "from": "P8", "before": "P4"}])
    assert _toc(moved) == ["C 4", "A 2", "B 3"]
    renamed = apply_batch(moved, [{"op": "set_text", "address": "P4", "text": "Costs"}])
    # A new sub-heading gets a line, linked to a bookmark added on it, in
    # the nearest level's style; its page is the line before's until Word.
    added = apply_batch(
        renamed,
        [{"op": "write", "from": "P5", "markdown": "C text\n\n## Detail\n\nMore."}],
    )
    assert _toc(added) == ["Costs 4", "Detail 4", "A 2", "B 3"]
    document = Document(io.BytesIO(added))
    fields = [
        c.get(qn("w:fldCharType")) for c in document.element.body.iter(qn("w:fldChar"))
    ]
    assert fields.count("begin") == fields.count("end")
    instructions = "".join(
        t.text or "" for t in document.element.body.iter(qn("w:instrText"))
    )
    assert instructions.count("TOC") == 1 and instructions.count("PAGEREF") == 4


def test_a_removed_chapter_leaves_the_table_of_contents():
    docx = _toc_document()
    removed = apply_batch(docx, [{"op": "remove", "from": "P6", "to": "P7"}])
    assert _toc(removed) == ["A 2", "C 4"]


def test_the_download_asks_word_to_update_its_fields():
    updated = Document(io.BytesIO(with_update_fields(_toc_document())))
    setting = updated.settings.element.find(qn("w:updateFields"))
    assert setting is not None and setting.get(qn("w:val")) == "true"
