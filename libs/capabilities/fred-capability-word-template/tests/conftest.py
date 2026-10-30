# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Synthetic .docx fixtures built with python-docx (no real template data)."""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import RGBColor


def to_bytes(document: DocxDocument) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def load(docx_bytes: bytes) -> DocxDocument:
    return Document(io.BytesIO(docx_bytes))


def add_bookmark(paragraph, name: str, bookmark_id: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(1, start)
    paragraph._p.append(end)


def add_toc_entry(document: DocxDocument, text: str, anchor: str):
    paragraph = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return paragraph


def form_document() -> DocxDocument:
    """A heading-less form: a title table, then a 4-row project block whose
    rows 3 and 4 are merged full width (like a weekly report)."""
    document = Document()
    title = document.add_table(rows=1, cols=3)
    title.style = "Table Grid"
    title.cell(0, 0).text = "Weekly report"
    title.cell(0, 1).text = "Week 10"
    title.cell(0, 2).text = "Year 2026"
    document.add_paragraph()
    block = document.add_table(rows=4, cols=4)
    block.style = "Table Grid"
    for column, label in enumerate(["Scope", "Code", "Project", "Milestone"]):
        block.cell(0, column).text = label
    block.cell(1, 0).text = "S1"
    block.cell(2, 0).merge(block.cell(2, 3)).text = "Significant comments"
    merged = block.cell(3, 0).merge(block.cell(3, 3))
    merged.text = "old comment"
    merged.add_paragraph("second line")
    document.sections[0].header.paragraphs[0].text = "Header text"
    document.sections[0].footer.paragraphs[0].text = "Footer text"
    return document


def report_document() -> DocxDocument:
    """Chapters with headings, a blue instruction, a TOC and a styled table."""
    document = Document()
    document.add_paragraph("Cover title")
    add_toc_entry(document, "1 Context", "_Toc1")
    add_toc_entry(document, "2 Scope", "_Toc2")
    heading = document.add_heading("Context", level=1)
    add_bookmark(heading, "_Toc1", "1")
    instruction = document.add_paragraph()
    run = instruction.add_run("Describe the context here.")
    run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    document.add_paragraph("Fixed sentence.")
    heading2 = document.add_heading("Scope", level=1)
    add_bookmark(heading2, "_Toc2", "2")
    document.add_paragraph("Scope instruction.")
    table = document.add_table(rows=2, cols=2)
    table.style = "Light Grid Accent 1"
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.add_heading("Annex", level=1)
    document.add_paragraph("Annex text.")
    return document


def add_vml_textbox(paragraph, text: str) -> None:
    xml = (
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:pict><v:shape><v:textbox><w:txbxContent>"
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    paragraph._p.append(parse_xml(xml))


@pytest.fixture
def form_bytes() -> bytes:
    return to_bytes(form_document())


@pytest.fixture
def report_bytes() -> bytes:
    return to_bytes(report_document())


def _png() -> bytes:
    """A valid 2x2 grey PNG."""
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", 2, 2, 8, 0, 0, 0, 0)
    raw = b"\x00\x80\x80" * 2
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


PNG = _png()
