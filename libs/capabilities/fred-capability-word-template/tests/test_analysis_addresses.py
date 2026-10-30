# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The analysis' addresses are read as the model writes them: lists,
ranges, and range ends relative to their start."""

from __future__ import annotations

import pytest
from conftest import to_bytes
from docx import Document
from docx.shared import RGBColor
from fred_capability_word_template.analysis import TemplateAnalysis, address_ranges
from fred_capability_word_template.check import (
    fixed_texts,
    template_texts,
)
from fred_capability_word_template.content import repeatable_groups
from fred_capability_word_template.docmap import build_map


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("P5", [("P5", "P5")]),
        ("P5–P9", [("P5", "P9")]),
        ("P5 - P9", [("P5", "P9")]),
        ("P3, P5–P6", [("P3", "P3"), ("P5", "P6")]),
        ("H1f.P1; F1.P5", [("H1f.P1", "H1f.P1"), ("F1.P5", "F1.P5")]),
        ("T1.R3.C1.P3–P4", [("T1.R3.C1.P3", "T1.R3.C1.P4")]),
        ("F1f.T1.R1–R2", [("F1f.T1.R1", "F1f.T1.R2")]),
        ("T2–T4", [("T2", "T4")]),
        ("", []),
    ],
)
def test_address_ranges(text, expected):
    assert address_ranges(text) == expected


def _letter_map():
    """A sender cell of three lines, then four body paragraphs."""
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Sender"
    cell = table.cell(0, 1)
    cell.text = "First Last"
    cell.add_paragraph("Phone 00 00")
    cell.add_paragraph("Fax 00 00")
    for text in (
        "Dear Sir or Madam,",
        "Sample body text.",
        "More sample text.",
        "Closing line.",
    ):
        document.add_paragraph(text)
    return build_map(to_bytes(document))


def _analysis(*elements: tuple[str, str]) -> TemplateAnalysis:
    return TemplateAnalysis.model_validate(
        {
            "summary": "A letter.",
            "parts": [
                {
                    "title": "Letter",
                    "addresses": "T1–P4",
                    "elements": [{"address": a, "role": r} for a, r in elements],
                }
            ],
        }
    )


def test_lists_and_cell_ranges_cover_only_what_they_name():
    analysis = _analysis(
        ("T1.R1.C2.P2–P3", "fill"),
        ("P2, P3", "replace"),
        ("P1; T1.R1.C1.P1", "fixed"),
    )
    doc_map = _letter_map()
    assert template_texts(doc_map, analysis) == {
        "Phone 00 00",
        "Fax 00 00",
        "Sample body text.",
        "More sample text.",
    }
    assert fixed_texts(doc_map, analysis) == {"Dear Sir or Madam,", "Sender"}


def test_a_range_written_end_first_covers_the_same_elements():
    doc_map = _letter_map()
    assert (
        fixed_texts(doc_map, _analysis(("P3–P1", "fixed")))
        == fixed_texts(doc_map, _analysis(("P1–P3", "fixed")))
        == {"Dear Sir or Madam,", "Sample body text.", "More sample text."}
    )


def test_a_text_also_meant_to_be_replaced_is_not_protected():
    # Contradictory analysis: the sample text is both fixed and to replace.
    analysis = _analysis(("P1–P2", "fixed"), ("P2", "replace"))
    doc_map = _letter_map()
    assert fixed_texts(doc_map, analysis) == {"Dear Sir or Madam,"}
    assert "Sample body text." in template_texts(doc_map, analysis)


def test_an_unknown_address_skips_only_itself():
    analysis = _analysis(("P9, P2", "replace"))
    # The zero-filled phone and fax numbers count through their pattern.
    assert template_texts(_letter_map(), analysis) == {
        "Sample body text.",
        "Phone 00 00",
        "Fax 00 00",
    }


def test_repeatable_blocks_read_lists_and_relative_ends():
    analysis = TemplateAnalysis.model_validate(
        {
            "summary": "A form.",
            "parts": [],
            "repeatable_blocks": [
                {"addresses": "T2.R2–R4, T3", "one_per": "one project"},
            ],
        }
    )
    assert [(g.first, g.last, g.label) for g in repeatable_groups(analysis)] == [
        ("T2.R2", "T2.R4", "one project"),
        ("T3", "T3", "one project"),
    ]


def test_a_fixed_text_with_a_coloured_link_stays_protected():
    # Only a passage is coloured (a link): the analysis' "fixed" wins. A
    # paragraph coloured whole is guidance, whatever the analysis says.
    document = Document()
    fixed = document.add_paragraph("Our company runs three sites. More at ")
    fixed.add_run("example.org").font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    guidance = document.add_paragraph().add_run("Example: write it like this.")
    guidance.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    # The body colour is the document's dominant one: plain text sets it.
    document.add_paragraph("Plain body text that sets the usual colour, " * 3)
    doc_map = build_map(to_bytes(document))
    analysis = _analysis(("P1–P2", "fixed"))
    assert fixed_texts(doc_map, analysis) == {
        "Our company runs three sites. More at example.org"
    }


def test_a_role_over_a_whole_table_counts_only_its_fields():
    # The analysis marks the version table as placeholders: its headings
    # are labels; the blue sample row and the date pattern are the fields.
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Version"
    table.cell(0, 1).text = "Date"
    table.cell(1, 0).paragraphs[0].add_run("1.0").font.color.rgb = RGBColor(
        0x00, 0x70, 0xC0
    )
    table.cell(1, 1).text = "jj/mm/aaaa"
    document.add_paragraph("Plain body text that sets the usual colour, " * 3)
    doc_map = build_map(to_bytes(document))
    whole = template_texts(doc_map, _analysis(("T1", "fill")))
    assert {"1.0", "jj/mm/aaaa"} <= whole and not {"Version", "Date"} & whole
    # Named cells count whole, whatever their look.
    named = template_texts(doc_map, _analysis(("T1.R1.C1", "fill")))
    assert "Version" in named
