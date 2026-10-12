# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Synchronous python-pptx deck construction for the `pptx_generator` capability.

`build()` is the single entry point. It is CPU-bound and synchronous, so the
tool runs it via `asyncio.to_thread` under the `concurrency.py` gate, never
inline on the event loop. Rationale: README.md
"""

from __future__ import annotations

import copy
import functools
import io
import re
from collections.abc import Mapping, Sequence
from typing import Any, NamedTuple

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from fred_capability_pptx_generator import background_render, chart_render
from fred_capability_pptx_generator.inline_markdown import parse_inline_markdown
from fred_capability_pptx_generator.models import (
    IMAGE_FALLBACK_CHART_TYPES,
    NATIVE_CHART_TYPES,
    BackgroundSpec,
    BulletItem,
    ChartSpec,
    ColumnSpec,
    ShapeSpec,
    SlideSpec,
    TableSpec,
    TextStyle,
    ThemeSpec,
)
from fred_capability_pptx_generator.palettes import THEME_PRESETS, ThemePresetDef

# Engineering guardrails, not product limits: they bound how much CPU/memory
# one tool call can commit the pod to.
MAX_SLIDES = 50
MAX_TABLE_ROWS = 50  # includes the header row
MAX_TABLE_COLS = 20
MAX_SHAPES_PER_SLIDE = 20
MAX_COLUMNS = 4

# Row count -> font size step function: a short table reads big, a dense one
# still fits its box. Steps, not a formula, so tier boundaries are testable.
_TABLE_FONT_SIZE_TIERS: tuple[tuple[int, float], ...] = (
    (6, 16.0),
    (10, 13.0),
    (15, 11.0),
)
_TABLE_FONT_SIZE_FLOOR_PT = 9.0
# Minimum column width, as a fraction of an even split, so one long column
# cannot crush its neighbors to zero.
_TABLE_COLUMN_WIDTH_FLOOR_FRACTION = 0.1

# A shape overlapping body content by more than this fraction of the smaller
# area is dropped: observed failure was a badge shape covering bullet text.
# Near-full-slide shapes (backgrounds) never count as occupied.
_SHAPE_OVERLAP_AREA_FRACTION_THRESHOLD = 0.15
_BACKGROUND_SHAPE_AREA_FRACTION = 0.9

# "line" is absent: python-pptx draws it as a connector (see `_add_shape`).
# MSO_SHAPE.PENTAGON is the "home plate" arrow, hence REGULAR_PENTAGON.
_MSO_SHAPE_BY_TYPE: dict[str, MSO_SHAPE] = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "arrow_right": MSO_SHAPE.RIGHT_ARROW,
    "arrow_left": MSO_SHAPE.LEFT_ARROW,
    "arrow_up": MSO_SHAPE.UP_ARROW,
    "arrow_down": MSO_SHAPE.DOWN_ARROW,
    "star": MSO_SHAPE.STAR_5_POINT,
    "pentagon": MSO_SHAPE.REGULAR_PENTAGON,
    "hexagon": MSO_SHAPE.HEXAGON,
    "chevron": MSO_SHAPE.CHEVRON,
    "cloud": MSO_SHAPE.CLOUD,
}

_PP_ALIGN_BY_NAME: dict[str, PP_ALIGN] = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}

# Series color goes through a fill or a line color. Pie/doughnut (per-point)
# and scatter/bubble (markers) are in neither set and stay uncolored.
_SERIES_FILL_CHART_TYPES = frozenset({"bar", "column", "area"})
_SERIES_LINE_CHART_TYPES = frozenset({"line", "radar"})

_NATIVE_CHART_XL_TYPE: dict[str, XL_CHART_TYPE] = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "radar": XL_CHART_TYPE.RADAR,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "bubble": XL_CHART_TYPE.BUBBLE,
}

# Office layout names per kind, tried before the structural fallback since a
# name match is a stronger signal than placeholder structure.
_LAYOUT_NAMES_BY_KIND: dict[str, tuple[str, ...]] = {
    "title": ("Title Slide",),
    "section_header": ("Section Header",),
    "two_content": ("Two Content", "Comparison"),
    "columns": ("Title and Content",),
    "bullets": ("Title and Content",),
    "image": ("Title and Content", "Picture with Caption"),
    "table": ("Title and Content",),
    "chart": ("Title and Content",),
}
# Allowlist, not "everything but the title": furniture placeholders would
# swamp the count (observed: "Title Only" out-scored "Title and Content").
_GENERIC_CONTENT_PLACEHOLDER_TYPES = frozenset(
    {
        PP_PLACEHOLDER.BODY,
        PP_PLACEHOLDER.OBJECT,
        PP_PLACEHOLDER.VERTICAL_BODY,
        PP_PLACEHOLDER.VERTICAL_OBJECT,
    }
)
# Generic content placeholders each kind wants, for the structural fallback
# (a fixed layout index collapsed several kinds onto one layout).
_WANTED_GENERIC_CONTENT_BY_KIND: dict[str, int] = {
    "title": 0,
    "section_header": 1,
    "two_content": 2,
    "columns": 1,
    "bullets": 1,
    "image": 1,
    "table": 1,
    "chart": 1,
}
# A cover slide prefers a real SUBTITLE placeholder over a same-count
# generic one.
_WANTS_SUBTITLE_PLACEHOLDER = frozenset({"title"})

# A chapter-divider badge's default text ("01"): some templates type the
# badge as the title, so placeholder type alone can't identify it.
_SECTION_HEADER_BADGE_TEXT_RE = re.compile(r"^\d{2}$")

# Never a content slot: slide content written here overwrote the cover date
# and put a chart in the footer band. Footers get the deck title instead.
_FURNITURE_PLACEHOLDER_TYPES = frozenset(
    {
        PP_PLACEHOLDER.DATE,
        PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.SLIDE_NUMBER,
        PP_PLACEHOLDER.HEADER,
    }
)

# Rejected when duplicating a template slide: their relationships are far
# harder to re-wire than a picture's single blip reference.
_UNSUPPORTED_TEMPLATE_SHAPE_TYPES = frozenset(
    {
        MSO_SHAPE_TYPE.CHART,
        MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
        MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
        MSO_SHAPE_TYPE.MEDIA,
    }
)

# Fallback content box when no placeholder exists, as fractions of the real
# slide size: fixed Inches() overflowed differently-sized templates.
_DEFAULT_CONTENT_BOX_FRACTIONS = (0.05, 0.20, 0.90, 0.65)
_COLUMN_GAP = Inches(0.3)
# Template zones whose left edges differ by less than this share a column.
_COLUMN_ALIGN_TOLERANCE = Inches(0.25)

# Share of fully transparent pixels above which an image counts as an icon.
_ICON_TRANSPARENT_SHARE = 0.5
# Pillow formats `add_picture` can embed; anything else is transcoded to PNG.
_PPTX_EMBEDDABLE_FORMATS = frozenset({"BMP", "GIF", "JPEG", "PNG", "TIFF", "WMF"})

_Box = tuple[int, int, int, int]


def _default_content_box(presentation) -> _Box:
    left_frac, top_frac, width_frac, height_frac = _DEFAULT_CONTENT_BOX_FRACTIONS
    return (
        round(presentation.slide_width * left_frac),
        round(presentation.slide_height * top_frac),
        round(presentation.slide_width * width_frac),
        round(presentation.slide_height * height_frac),
    )


def build(
    title: str,
    slides: Sequence[SlideSpec],
    *,
    template_bytes: bytes | None,
    image_bytes_by_ref: Mapping[str, bytes],
    theme: ThemeSpec | None = None,
    warnings: list[str] | None = None,
) -> bytes:
    """Build the complete deck and return its `.pptx` bytes.

    `image_bytes_by_ref` must be prefetched: a worker thread cannot await a
    platform port. Structural problems raise `ValueError` aimed at the agent;
    misfit optional fields and overlapping shapes are dropped into `warnings`.
    """

    presentation = (
        Presentation(io.BytesIO(template_bytes)) if template_bytes else Presentation()
    )
    # The template's own slides are raw material for `template_slide_index`
    # only, captured before any slide is appended and stripped at the end.
    template_source_slides = list(presentation.slides)
    presentation.core_properties.title = title
    preset = _resolve_theme_preset(theme)
    accent_theme = _theme_with_preset_defaults(
        theme, preset, default_primary_color=preset.accent_color if preset else None
    )
    layouts = list(presentation.slide_layouts)
    layouts_by_name = {layout.name.strip().lower(): layout for layout in layouts}
    # The chapter badge's number is this divider's position among badge
    # dividers, never caller-supplied.
    next_badge_ordinal = 1
    for index, spec in enumerate(slides, start=1):
        used_badge_ordinal = _build_slide(
            presentation,
            spec,
            index=index,
            deck_title=title,
            image_bytes_by_ref=image_bytes_by_ref,
            theme=theme,
            preset=preset,
            accent_theme=accent_theme,
            layouts=layouts,
            layouts_by_name=layouts_by_name,
            template_source_slides=template_source_slides,
            badge_ordinal=next_badge_ordinal,
            warnings=warnings if warnings is not None else [],
        )
        if used_badge_ordinal:
            next_badge_ordinal += 1
    if template_source_slides:
        _strip_template_source_slides(presentation, template_source_slides)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _resolve_theme_preset(theme: ThemeSpec | None) -> ThemePresetDef | None:
    if theme is None or theme.preset is None:
        return None
    # Validated against THEME_PRESETS by models.ThemeSpec.
    return THEME_PRESETS[theme.preset]


def _preset_background_for_kind(preset: ThemePresetDef, kind: str) -> BackgroundSpec:
    """Title/section_header slides get the preset's strong background, every
    other kind a solid fill of its soft tint."""

    if kind in ("title", "section_header"):
        return preset.strong_background
    return BackgroundSpec(kind="solid", color=preset.soft_background_color)


def _theme_with_preset_defaults(
    theme: ThemeSpec | None,
    preset: ThemePresetDef | None,
    *,
    default_primary_color: str | None,
) -> ThemeSpec | None:
    """Layer `theme` over `preset`'s fonts and `default_primary_color`; every
    explicit `theme` field wins. Rationale: README.md"""

    if preset is None:
        return theme
    heading_font = (
        theme.heading_font
        if theme and theme.heading_font is not None
        else preset.heading_font
    )
    body_font = (
        theme.body_font if theme and theme.body_font is not None else preset.body_font
    )
    primary_color = (
        theme.primary_color
        if theme and theme.primary_color is not None
        else default_primary_color
    )
    return ThemeSpec(
        heading_font=heading_font,
        body_font=body_font,
        primary_color=primary_color,
        background=theme.background if theme else None,
    )


def _preset_strong_text_default(preset: ThemePresetDef | None, kind: str) -> str | None:
    if preset is None or kind not in ("title", "section_header"):
        return None
    return preset.strong_text_color


def _build_slide(
    presentation,
    spec: SlideSpec,
    *,
    index: int,
    image_bytes_by_ref: Mapping[str, bytes],
    theme: ThemeSpec | None,
    layouts: Sequence,
    deck_title: str = "",
    layouts_by_name: Mapping[str, object],
    template_source_slides: Sequence = (),
    preset: ThemePresetDef | None = None,
    accent_theme: ThemeSpec | None = None,
    badge_ordinal: int = 1,
    warnings: list[str],
) -> bool:
    """Build one slide. Returns whether it consumed `badge_ordinal`."""

    image_bytes = image_bytes_by_ref.get(spec.image_ref or "")
    if image_bytes is not None and _looks_like_icon(image_bytes):
        raise ValueError(
            f"slide {index}: image_ref '{spec.image_ref}' is a mostly "
            "transparent icon, not a photo — use it only as a column "
            "'icon_ref' and pick a real photo for 'image_ref'."
        )
    template_slide_index = spec.template_slide_index
    from_template = template_slide_index is not None
    if template_slide_index is not None:
        slide = _duplicate_slide(
            presentation,
            template_slide_index,
            template_source_slides,
            index=index,
        )
        if spec.kind != "columns":  # its icon frames are optional
            _require_image_for_picture_placeholders(slide, spec, index=index)
        unfilled_snapshot = _placeholder_run_snapshot(slide)
    else:
        layout = _resolve_layout(spec.kind, layouts, layouts_by_name)
        slide = presentation.slides.add_slide(layout)
        unfilled_snapshot = {}

    # Slide background > preset > theme.background. A template slide skips
    # both deck-wide defaults: painting over its designed background broke
    # the template's text contrast.
    background = spec.background
    if background is None and preset is not None and not from_template:
        background = _preset_background_for_kind(preset, spec.kind)
    if background is None and theme is not None and not from_template:
        background = theme.background
    _apply_background(
        presentation,
        slide,
        background,
        index=index,
        image_bytes_by_ref=image_bytes_by_ref,
        theme=accent_theme,
    )

    # A template slide also skips the deck-wide text color (one color cannot
    # suit both dark and light template slides); fonts still apply.
    text_theme = (
        (theme.model_copy(update={"primary_color": None}) if theme else None)
        if from_template
        else _theme_with_preset_defaults(
            theme,
            preset,
            default_primary_color=_preset_strong_text_default(preset, spec.kind),
        )
    )
    title_style = _resolve_text_style(
        spec.title_style, theme=text_theme, theme_font_attr="heading_font"
    )
    body_style = _resolve_text_style(
        spec.body_style, theme=text_theme, theme_font_attr="body_font"
    )
    section_header_badge = section_header_title_zone = None
    if spec.kind == "section_header" and from_template:
        section_header_badge, section_header_title_zone = (
            _section_header_badge_and_title_placeholders(slide)
        )

    if section_header_badge is not None:
        _fill_section_header_from_template(
            spec,
            section_header_badge,
            section_header_title_zone,
            ordinal=badge_ordinal,
            style=title_style,
            index=index,
        )
    else:
        _set_title(presentation, slide, spec.title, style=title_style)

    if spec.kind in ("title", "section_header"):
        if section_header_badge is None:
            _set_secondary_text(
                presentation,
                slide,
                spec.subtitle,
                style=body_style,
                template_slide_index=template_slide_index,
                index=index,
            )
        # else: already fully filled by _fill_section_header_from_template above.
    elif spec.kind == "bullets" or (
        spec.kind == "two_content" and not (spec.image_ref or spec.table or spec.chart)
    ):
        # A two_content slide with nothing for its right column is a bullets
        # slide — splitting the body would leave half of it empty.
        _fill_bullets_on_slide(
            presentation,
            slide,
            spec.bullets or [],
            subtitle=spec.subtitle,
            style=body_style,
        )
    elif spec.kind == "image":
        frame = _first_empty_picture_frame(slide)
        _place_image(
            slide,
            frame if frame is not None else _take_content_box(presentation, slide),
            spec.image_ref,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
        )
    elif spec.kind == "table":
        box = _take_content_box(presentation, slide)
        _place_table(slide, box, spec.table, index=index)
    elif spec.kind == "chart":
        box = _take_content_box(presentation, slide)
        _place_chart(slide, box, spec.chart, index=index)
    elif spec.kind == "two_content":
        _build_two_content(
            presentation,
            slide,
            spec,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
            body_style=body_style,
        )
    elif spec.kind == "columns":
        _build_columns(
            presentation,
            slide,
            spec,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
            body_style=body_style,
            warnings=warnings,
        )
    else:
        raise ValueError(f"slide {index}: unknown slide kind '{spec.kind}'")

    if spec.image_ref and spec.kind in _KINDS_WITHOUT_IMAGE:
        # A template picture frame would otherwise be dropped as unfilled.
        frame = _first_empty_picture_frame(slide)
        if frame is not None:
            _place_image(
                slide,
                frame,
                spec.image_ref,
                index=index,
                image_bytes_by_ref=image_bytes_by_ref,
            )
    if unfilled_snapshot:
        _remove_unfilled_placeholders(slide, unfilled_snapshot)
    if deck_title:
        _fill_footers(slide, deck_title)

    if spec.speaker_notes:
        slide.notes_slide.notes_text_frame.text = spec.speaker_notes

    if spec.shapes:
        if len(spec.shapes) > MAX_SHAPES_PER_SLIDE:
            raise ValueError(
                f"slide {index}: too many shapes ({len(spec.shapes)}); the "
                f"limit is {MAX_SHAPES_PER_SLIDE} per slide. Split them "
                "across more slides."
            )
        occupied_boxes = _occupied_content_boxes(presentation, slide)
        for shape_spec in spec.shapes:
            _add_shape(
                presentation,
                slide,
                shape_spec,
                index=index,
                occupied_boxes=occupied_boxes,
                warnings=warnings,
            )

    return section_header_badge is not None


def _resolve_layout(
    kind: str, layouts: Sequence, layouts_by_name: Mapping[str, object]
):
    if not layouts:
        raise ValueError(
            "the configured brand template has no slide layouts to build "
            "from — re-upload a valid PowerPoint template."
        )
    for name in _LAYOUT_NAMES_BY_KIND.get(kind, ()):
        layout = layouts_by_name.get(name.lower())
        if layout is not None:
            return layout
    return _best_structural_layout_match(kind, layouts)


def _generic_content_count(layout) -> int:
    return sum(
        1
        for placeholder in layout.placeholders
        if placeholder.placeholder_format.type in _GENERIC_CONTENT_PLACEHOLDER_TYPES
    )


def _has_placeholder_type(layout, placeholder_type) -> bool:
    return any(
        placeholder.placeholder_format.type == placeholder_type
        for placeholder in layout.placeholders
    )


def _best_structural_layout_match(kind: str, layouts: Sequence):
    """Pick the layout whose placeholders fit `kind` best when no name matches.
    A PICTURE (image) or SUBTITLE (title) placeholder ranks ahead of the
    content count, which alone can't tell same-sized layouts apart."""

    wanted = _WANTED_GENERIC_CONTENT_BY_KIND.get(kind, 1)
    best_layout = layouts[0]
    best_score: tuple[int, int, int] | None = None
    for layout in layouts:
        subtitle_penalty = int(
            kind in _WANTS_SUBTITLE_PLACEHOLDER
            and not _has_placeholder_type(layout, PP_PLACEHOLDER.SUBTITLE)
        )
        picture_penalty = int(
            kind == "image"
            and not _has_placeholder_type(layout, PP_PLACEHOLDER.PICTURE)
        )
        score = (
            subtitle_penalty,
            picture_penalty,
            abs(_generic_content_count(layout) - wanted),
        )
        if best_score is None or score < best_score:
            best_score = score
            best_layout = layout
    return best_layout


def _iter_shapes_recursive(shapes):
    """Yield every shape in `shapes`, descending into group shapes too."""

    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes_recursive(shape.shapes)


def _duplicate_slide(
    presentation,
    template_slide_index: int,
    template_source_slides: Sequence,
    *,
    index: int,
):
    """Append a deep copy of a template slide, keeping its design; the fill
    functions then overwrite its placeholders. python-pptx has no slide
    duplication API, so this walks the shape tree by hand."""

    if not template_source_slides:
        raise ValueError(
            f"slide {index}: 'template_slide_index' was set but no brand "
            "template is configured for this agent."
        )
    if template_slide_index >= len(template_source_slides):
        raise ValueError(
            f"slide {index}: template_slide_index {template_slide_index} "
            "does not exist — the uploaded template has "
            f"{len(template_source_slides)} slide(s) (valid indices: 0-"
            f"{len(template_source_slides) - 1})."
        )
    source_slide = template_source_slides[template_slide_index]

    # Recurse into groups: a nested chart/OLE/media shape would otherwise be
    # copied with dangling relationship references.
    for shape in _iter_shapes_recursive(source_slide.shapes):
        if shape.shape_type in _UNSUPPORTED_TEMPLATE_SHAPE_TYPES:
            raise ValueError(
                f"slide {index}: template_slide_index {template_slide_index} "
                f"contains a {shape.shape_type} shape, which isn't supported "
                "for duplication yet — remove it from that template slide or "
                "pick a different template_slide_index."
            )

    new_slide = presentation.slides.add_slide(source_slide.slide_layout)
    # add_slide() clones the layout's placeholders; drop them so the copied
    # shape tree isn't duplicated.
    for shape in list(new_slide.shapes):
        _remove_shape(shape)
    _copy_slide_background(source_slide, new_slide)
    for shape in source_slide.shapes:
        new_element = copy.deepcopy(shape._element)
        new_slide.shapes._spTree.append(new_element)
        for blip in new_element.findall(f".//{qn('a:blip')}"):
            _rewire_blip_relationship(source_slide, new_slide, blip)

    if source_slide.has_notes_slide:
        notes_text = source_slide.notes_slide.notes_text_frame.text
        if notes_text:
            new_slide.notes_slide.notes_text_frame.text = notes_text

    return new_slide


def _copy_slide_background(source_slide, new_slide) -> None:
    """Copy the source slide's explicit `<p:bg>`, if any: it is not a shape,
    so the shape-copy loop in `_duplicate_slide` never touches it."""

    source_bg = source_slide._element.cSld.find(qn("p:bg"))
    if source_bg is None:
        return
    new_csld = new_slide._element.cSld
    existing_bg = new_csld.find(qn("p:bg"))
    if existing_bg is not None:
        new_csld.remove(existing_bg)
    new_bg = copy.deepcopy(source_bg)
    new_csld.insert(0, new_bg)
    for blip in new_bg.findall(f".//{qn('a:blip')}"):
        _rewire_blip_relationship(source_slide, new_slide, blip)


def _rewire_blip_relationship(source_slide, new_slide, blip) -> None:
    """Re-point a deep-copied blip, whose rId belongs to the source slide, at
    a fresh relationship to the same image part."""

    old_r_id = blip.get(qn("r:embed"))
    if not old_r_id:
        return
    image_part = source_slide.part.related_part(old_r_id)
    new_r_id = new_slide.part.relate_to(image_part, RT.IMAGE)
    blip.set(qn("r:embed"), new_r_id)


def _strip_template_source_slides(presentation, source_slides: Sequence) -> None:
    """Remove the template's original slides from the output deck."""

    source_slide_ids = {slide.slide_id for slide in source_slides}
    sld_id_lst = presentation.slides._sldIdLst
    for sld_id_element in list(sld_id_lst):
        if int(sld_id_element.get("id")) in source_slide_ids:
            r_id = sld_id_element.get(qn("r:id"))
            sld_id_lst.remove(sld_id_element)
            presentation.part.rels.pop(r_id)
    # The template's PowerPoint sections only list its own (now removed) slides.
    for section_list in presentation.part._element.xpath(
        ".//*[local-name()='sectionLst']"
    ):
        extension = section_list.getparent()
        extension.getparent().remove(extension)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def _resolve_text_style(
    style: TextStyle | None, *, theme: ThemeSpec | None, theme_font_attr: str
) -> TextStyle | None:
    """Fill the unset font/color of a slide's text style from the theme."""

    theme_font = getattr(theme, theme_font_attr, None) if theme else None
    theme_color = theme.primary_color if theme else None
    font = style.font if style is not None and style.font is not None else theme_font
    color = (
        style.color if style is not None and style.color is not None else theme_color
    )
    if style is None and font is None and color is None:
        return None
    return TextStyle(
        font=font,
        size_pt=style.size_pt if style else None,
        color=color,
        bold=style.bold if style else None,
        italic=style.italic if style else None,
        align=style.align if style else None,
    )


def _apply_run_style(run, style: TextStyle | None) -> None:
    if style is None:
        return
    if style.font:
        run.font.name = style.font
    if style.size_pt is not None:
        run.font.size = Pt(style.size_pt)
    if style.color:
        run.font.color.rgb = _rgb(style.color)
    if style.bold is not None:
        run.font.bold = style.bold
    if style.italic is not None:
        run.font.italic = style.italic


def _enable_shrink_to_fit(text_frame) -> None:
    """Enable native shrink-on-overflow. python-pptx writes no `fontScale`;
    the renderer computes it (verified to work in the LibreOffice preview)."""

    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _resolve_align(style: TextStyle | None) -> PP_ALIGN | None:
    if style is None or style.align is None:
        return None
    return _PP_ALIGN_BY_NAME.get(style.align)


def _apply_text_frame_style(text_frame, style: TextStyle | None) -> None:
    """Apply `style` to every run of a single-block text frame."""

    if style is None:
        return
    align = _resolve_align(style)
    for paragraph in text_frame.paragraphs:
        if align is not None:
            paragraph.alignment = align
        for run in paragraph.runs:
            _apply_run_style(run, style)


def _set_title(
    presentation, slide, title: str | None, *, style: TextStyle | None = None
) -> None:
    if not title:
        return
    if slide.shapes.title is not None:
        slide.shapes.title.text_frame.text = title
        _enable_shrink_to_fit(slide.shapes.title.text_frame)
        _apply_text_frame_style(slide.shapes.title.text_frame, style)
        return
    left, _top, width, _height = _default_content_box(presentation)
    text_box = slide.shapes.add_textbox(left, Inches(0.3), width, Inches(1.0))
    text_box.text_frame.text = title
    _enable_shrink_to_fit(text_box.text_frame)
    _apply_text_frame_style(text_box.text_frame, style)


def _is_furniture(placeholder) -> bool:
    return placeholder.placeholder_format.type in _FURNITURE_PLACEHOLDER_TYPES


def _body_placeholder(slide):
    """First non-PICTURE content placeholder on the slide, in idx order, or
    None (text cannot be written into a picture frame)."""

    for placeholder in _content_placeholders(slide):
        if placeholder.placeholder_format.type != PP_PLACEHOLDER.PICTURE:
            return placeholder
    return None


def _largest_text_placeholder(slide):
    texts = [
        ph
        for ph in _content_placeholders(slide)
        if ph.placeholder_format.type != PP_PLACEHOLDER.PICTURE
    ]
    return max(texts, key=lambda ph: ph.width * ph.height, default=None)


def _first_empty_picture_frame(slide):
    return next(
        (ph for ph in _content_placeholders(slide) if _is_empty_picture_frame(ph)),
        None,
    )


def _content_placeholders(slide) -> list:
    """Every non-title, non-furniture placeholder on the slide, in idx order."""

    return [
        placeholder
        for placeholder in slide.placeholders
        if placeholder.placeholder_format.type
        not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        and not _is_furniture(placeholder)
    ]


def _fill_footers(slide, text: str) -> None:
    """Write the deck-wide footer into every FOOTER placeholder, keeping the
    template's first run formatting (a template footer often carries caps
    or a size override on the run itself)."""

    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.type != PP_PLACEHOLDER.FOOTER:
            continue
        text_frame = placeholder.text_frame
        first_run = next(placeholder._element.iter(qn("a:r")), None)
        run_props = first_run.find(qn("a:rPr")) if first_run is not None else None
        run_props = copy.deepcopy(run_props) if run_props is not None else None
        # The setter also drops fields and line breaks, not only runs.
        text_frame.text = text
        if run_props is not None:
            new_run = text_frame.paragraphs[0].runs[0]._r
            existing = new_run.find(qn("a:rPr"))
            if existing is not None:
                new_run.remove(existing)
            new_run.insert(0, run_props)


def _is_empty_picture_frame(placeholder) -> bool:
    # A designer-placed photo is a filled `p:pic`; only a `p:sp` is an empty frame.
    return (
        placeholder.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        and placeholder._element.tag == qn("p:sp")
    )


def _placeholder_run_snapshot(slide) -> dict[int, list]:
    """The run elements each content/title placeholder of a freshly
    duplicated template slide starts with — writing new text replaces them,
    which is how `_remove_unfilled_placeholders` tells untouched from filled."""

    return {
        placeholder.shape_id: list(placeholder._element.iter(qn("a:r")))
        for placeholder in slide.placeholders
        if not _is_furniture(placeholder)
    }


def _remove_unfilled_placeholders(slide, snapshot: Mapping[int, list]) -> None:
    """Drop the template's sample text ("Lorem ipsum", "Sous-titre"...) and
    empty picture frames that no SlideSpec field wrote into."""

    for placeholder in list(slide.placeholders):
        original_runs = snapshot.get(placeholder.shape_id)
        if original_runs is None:
            continue
        if _is_empty_picture_frame(placeholder):
            _remove_shape(placeholder)
            continue
        if not placeholder.has_text_frame:
            continue
        # Held references keep lxml proxies stable, so identity is reliable.
        current_run_ids = {id(run) for run in placeholder._element.iter(qn("a:r"))}
        untouched = all(id(run) in current_run_ids for run in original_runs)
        empty_decoration = _has_visible_outline(placeholder)
        if untouched and (
            original_runs
            or not (placeholder.text_frame.text.strip() or empty_decoration)
        ):
            _remove_shape(placeholder)


def _has_visible_outline(placeholder) -> bool:
    """An empty text frame drawn only as a border (e.g. a dotted box around a
    group of columns) is decoration, not sample text to clean up."""

    base = placeholder._base_placeholder
    for element in (placeholder._element, base._element if base else None):
        line = element.spPr.find(qn("a:ln")) if element is not None else None
        if line is not None and len(line) and line.find(qn("a:noFill")) is None:
            return True
    return False


def _require_image_for_picture_placeholders(
    slide, spec: SlideSpec, *, index: int
) -> None:
    if spec.image_ref:
        return
    if any(_is_empty_picture_frame(placeholder) for placeholder in slide.placeholders):
        raise ValueError(
            f"slide {index}: template_slide_index {spec.template_slide_index} "
            "is an image layout (it has picture frames) but this slide sets no "
            "'image_ref' — pick a text-only template slide instead, or set an "
            "'image_ref' from the document library."
        )


def _section_header_badge_and_title_placeholders(slide):
    """Find a chapter divider's badge/title zones by default content, not
    type. `(None, None)` unless the match is unambiguous."""

    placeholders = [
        placeholder
        for placeholder in slide.placeholders
        if not _is_furniture(placeholder)
    ]
    if len(placeholders) != 2:
        return None, None
    badge_candidates = [
        placeholder
        for placeholder in placeholders
        if _SECTION_HEADER_BADGE_TEXT_RE.match(placeholder.text_frame.text.strip())
    ]
    if len(badge_candidates) != 1:
        return None, None
    badge_placeholder = badge_candidates[0]
    title_placeholder = next(
        placeholder
        for placeholder in placeholders
        if placeholder is not badge_placeholder
    )
    if not title_placeholder.text_frame.text.strip():
        return None, None
    return badge_placeholder, title_placeholder


def _fill_section_header_from_template(
    spec: SlideSpec,
    badge_placeholder,
    title_placeholder,
    *,
    ordinal: int,
    style: TextStyle | None,
    index: int,
) -> None:
    """Fill a content-detected chapter-divider slide."""

    if not spec.title:
        raise ValueError(
            f"slide {index}: kind 'section_header' duplicating "
            f"template_slide_index {spec.template_slide_index} needs a "
            "'title' (the chapter name) to replace this template slide's "
            "own default title text."
        )
    badge_text = f"{ordinal:02d}"
    badge_placeholder.text_frame.text = badge_text
    title_placeholder.text_frame.text = spec.title
    _enable_shrink_to_fit(title_placeholder.text_frame)
    _apply_text_frame_style(title_placeholder.text_frame, style)

    # Fail loudly rather than ship a half-filled divider.
    badge_paragraphs = badge_placeholder.text_frame.paragraphs
    badge_written_cleanly = (
        len(badge_paragraphs) == 1
        and len(badge_paragraphs[0].runs) == 1
        and badge_paragraphs[0].runs[0].text == badge_text
    )
    if not badge_written_cleanly or title_placeholder.text_frame.text != spec.title:
        raise ValueError(
            f"slide {index}: could not cleanly fill the chapter-divider "
            f"template slide (template_slide_index={spec.template_slide_index}) "
            "— its badge/title zones did not take the new content as "
            "expected; re-check the template slide's structure."
        )


def _remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def _placeholder_box(placeholder) -> _Box:
    return (placeholder.left, placeholder.top, placeholder.width, placeholder.height)


def _take_content_box(presentation, slide) -> _Box:
    """Geometry of the body placeholder, removed so the caller can replace it
    with a table/chart/picture."""

    placeholder = _body_placeholder(slide)
    if placeholder is None:
        return _default_content_box(presentation)
    box = _placeholder_box(placeholder)
    _remove_shape(placeholder)
    return box


def _set_secondary_text(
    presentation,
    slide,
    text: str | None,
    *,
    style: TextStyle | None = None,
    template_slide_index: int | None = None,
    index: int = 0,
) -> None:
    if not text:
        return
    placeholder = _body_placeholder(slide)
    if placeholder is not None:
        text_frame = placeholder.text_frame
        _fill_bullets(text_frame, _lead_items(text_frame, text), style=style)
        return
    if template_slide_index is not None:
        # A free text box would land off-design, in the default font.
        raise ValueError(
            f"slide {index}: template_slide_index {template_slide_index} has "
            "no subtitle zone — leave 'subtitle' unset or pick another slide."
        )
    left, top, width, height = _default_content_box(presentation)
    text_box = slide.shapes.add_textbox(left, top, width, Emu(int(height / 3)))
    text_box.text_frame.text = text
    _enable_shrink_to_fit(text_box.text_frame)
    _apply_text_frame_style(text_box.text_frame, style)


_MAX_PARAGRAPH_LEVEL = 8
_KINDS_WITHOUT_IMAGE = ("title", "section_header", "bullets", "table", "chart")


def _sample_text_level(text_frame) -> int:
    """Lowest paragraph level of a template's sample text — where that zone's
    bullets actually start (a template may style level 0 as plain text and
    put its bulleted list at level 2). 0 when there is no sample text."""

    levels = [p.level for p in text_frame.paragraphs if p.text.strip()]
    return min(levels, default=0)


def _lead_items(text_frame, text: str) -> list[BulletItem]:
    """One paragraph per line of `text`, the i-th taking the zone's i-th
    sample level in reading order — a lead zone often pairs a bold lead
    line with a plain paragraph below it at another level."""

    levels: list[int] = []
    for paragraph in text_frame.paragraphs:
        if paragraph.text.strip() and paragraph.level not in levels:
            levels.append(paragraph.level)
    base = min(levels, default=0)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return [
        BulletItem(
            text=line,
            level=levels[min(position, len(levels) - 1)] - base if levels else 0,
        )
        for position, line in enumerate(lines)
    ]


def _sample_level_formats(text_frame) -> dict[int, tuple[Any, Any]]:
    """First sample paragraph's `a:pPr` and first-run `a:rPr` per level. A
    template often sets bullet/bold/color on the sample paragraphs rather
    than its layout, so `clear()` would otherwise lose the zone's look."""

    formats: dict[int, tuple[Any, Any]] = {}
    for paragraph in text_frame.paragraphs:
        if not paragraph.text.strip() or paragraph.level in formats:
            continue
        p_pr = paragraph._p.pPr
        first_run = paragraph._p.find(qn("a:r"))
        r_pr = first_run.find(qn("a:rPr")) if first_run is not None else None
        formats[paragraph.level] = (
            copy.deepcopy(p_pr) if p_pr is not None else None,
            _strip_hyperlinks(copy.deepcopy(r_pr)) if r_pr is not None else None,
        )
    return formats


def _strip_hyperlinks(r_pr):
    for tag in ("a:hlinkClick", "a:hlinkMouseOver"):
        for link in r_pr.findall(qn(tag)):
            r_pr.remove(link)
    return r_pr


def _fill_bullets(
    text_frame, bullets: Sequence[BulletItem], *, style: TextStyle | None = None
) -> None:
    """One paragraph per bullet, one run per inline-markdown span. A span's
    emphasis only turns bold/italic on, layering over `style`."""

    base_level = _sample_text_level(text_frame)
    level_formats = _sample_level_formats(text_frame)
    text_frame.clear()
    if not bullets:
        return
    _enable_shrink_to_fit(text_frame)
    align = _resolve_align(style)
    for position, bullet in enumerate(bullets):
        paragraph = (
            text_frame.paragraphs[0] if position == 0 else text_frame.add_paragraph()
        )
        level = min(base_level + bullet.level, _MAX_PARAGRAPH_LEVEL)
        p_pr, r_pr = level_formats.get(level, (None, None))
        if p_pr is not None:
            existing = paragraph._p.pPr
            if existing is not None:
                paragraph._p.remove(existing)
            paragraph._p.insert(0, copy.deepcopy(p_pr))
        paragraph.level = level
        if align is not None:
            paragraph.alignment = align
        for span in parse_inline_markdown(bullet.text):
            run = paragraph.add_run()
            run.text = span.text
            if r_pr is not None:
                run._r.insert(0, copy.deepcopy(r_pr))
            _apply_run_style(run, style)
            if span.bold:
                run.font.bold = True
            if span.italic:
                run.font.italic = True


def _fill_bullets_on_slide(
    presentation,
    slide,
    bullets: Sequence[BulletItem],
    *,
    subtitle: str | None = None,
    style: TextStyle | None = None,
) -> None:
    """Bullets go to the largest text zone; `subtitle` to the topmost other
    one (a template's lead paragraph above its list), when there is one."""

    placeholder = _largest_text_placeholder(slide)
    if placeholder is None:
        left, top, width, height = _default_content_box(presentation)
        text_box = slide.shapes.add_textbox(left, top, width, height)
        _fill_bullets(text_box.text_frame, bullets, style=style)
        return
    _fill_bullets(placeholder.text_frame, bullets, style=style)
    _fill_lead_zone(slide, placeholder, subtitle, style=style)


def _fill_lead_zone(
    slide, main_zone, subtitle: str | None, *, style: TextStyle | None
) -> None:
    """Write `subtitle` into the topmost text zone other than `main_zone`."""

    lead = min(
        (
            ph
            for ph in _content_placeholders(slide)
            if ph.placeholder_format.type != PP_PLACEHOLDER.PICTURE
            and ph.shape_id != main_zone.shape_id
        ),
        key=_top,
        default=None,
    )
    if subtitle and lead is not None:
        _fill_bullets(
            lead.text_frame, _lead_items(lead.text_frame, subtitle), style=style
        )


def _fit_inside_box(
    img_w: int, img_h: int, box_left: int, box_top: int, box_w: int, box_h: int
) -> _Box:
    """Fit the image inside the box, aspect preserved and centered. Duplicated
    from `ppt_filler` rather than imported cross-package (AUTHORING.md)."""

    if img_w <= 0 or img_h <= 0 or box_w <= 0 or box_h <= 0:
        return (box_left, box_top, box_w, box_h)

    aspect_ratio = img_w / img_h
    if box_w / box_h > aspect_ratio:
        height = box_h
        width = box_h * aspect_ratio
        top = box_top
        left = box_left + (box_w - width) / 2
    else:
        width = box_w
        height = box_w / aspect_ratio
        left = box_left
        top = box_top + (box_h - height) / 2
    return (round(left), round(top), round(width), round(height))


def _resolve_image_bytes(
    image_ref: str | None,
    *,
    index: int,
    image_bytes_by_ref: Mapping[str, bytes],
    field_label: str = "image_ref",
) -> bytes:
    if not image_ref:
        # Actionable on purpose: a bare schema complaint gets echoed to the
        # user verbatim by a model that doesn't retry with a fix.
        raise ValueError(
            f"slide {index}: '{field_label}' is required. Pass the "
            "document_uid of a document found via a document search/tree "
            "tool, or don't use an image here if no such tool is available "
            "in this conversation — use a different slide kind/background "
            "instead."
        )
    image_bytes = image_bytes_by_ref.get(image_ref)
    if image_bytes is None:
        raise ValueError(
            f"slide {index}: {field_label} '{image_ref}' could not be resolved "
            "to a document in the library. Pass the document_uid of a "
            "document you found via a document search/tree tool."
        )
    return image_bytes


@functools.lru_cache(maxsize=16)
def _looks_like_icon(image_bytes: bytes) -> bool:
    """A picto is mostly transparent pixels; a photo has none. Undecodable
    bytes are left for `_load_embeddable_image` to report."""

    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            if "A" not in image.getbands() and "transparency" not in image.info:
                return False
            histogram = image.convert("RGBA").getchannel("A").histogram()
    except (UnidentifiedImageError, OSError):
        return False
    return histogram[0] > _ICON_TRANSPARENT_SHARE * sum(histogram)


@functools.lru_cache(maxsize=16)
def _decode_and_maybe_transcode(image_bytes: bytes) -> tuple[bytes, int, int]:
    """Decode/transcode cached on the raw bytes, so an image reused across
    slides is decoded once. `maxsize` is small on purpose: entries can be
    several MB in a process-wide cache shared by every tenant."""

    with Image.open(io.BytesIO(image_bytes)) as image:
        img_w, img_h = image.size
        payload = image_bytes
        if image.format not in _PPTX_EMBEDDABLE_FORMATS:
            buffer = io.BytesIO()
            image.convert("RGBA").save(buffer, format="PNG")
            payload = buffer.getvalue()
    return payload, img_w, img_h


def _load_embeddable_image(
    image_bytes: bytes, *, image_ref: str, index: int
) -> tuple[bytes, int, int]:
    """Decode `image_bytes`, transcoding to PNG if python-pptx can't embed
    the source format directly. Returns `(payload, width, height)`."""

    try:
        return _decode_and_maybe_transcode(image_bytes)
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise ValueError(
            f"slide {index}: image_ref '{image_ref}' is not a usable image "
            f"[{type(exc).__name__}: {exc}]."
        ) from exc


def _place_image(
    slide,
    target,
    image_ref: str | None,
    *,
    index: int,
    image_bytes_by_ref: Mapping[str, bytes],
) -> None:
    """`target` is either an empty template picture frame, filled in place
    and cropped to its shape, or a plain `_Box` the image is fit inside."""

    image_bytes = _resolve_image_bytes(
        image_ref, index=index, image_bytes_by_ref=image_bytes_by_ref
    )
    payload, img_w, img_h = _load_embeddable_image(
        image_bytes, image_ref=image_ref or "", index=index
    )
    if not isinstance(target, tuple):
        _fill_picture_frame(target, payload)
        return

    left, top, width, height = target
    fit_left, fit_top, fit_width, fit_height = _fit_inside_box(
        img_w, img_h, left, top, width, height
    )
    slide.shapes.add_picture(
        io.BytesIO(payload), fit_left, fit_top, fit_width, fit_height
    )


def _fill_picture_frame(frame, payload: bytes) -> None:
    """Crop-to-fill the frame, keeping the slide's own frame formatting
    (position, preset geometry such as an ellipse, outline), which
    `insert_picture` alone drops in favour of the layout's."""

    sp_pr = copy.deepcopy(frame._element.spPr)
    picture = frame.insert_picture(io.BytesIO(payload))
    picture._element.replace(picture._element.spPr, sp_pr)


def _send_to_back(shape) -> None:
    """Move `shape` behind every other shape. python-pptx has no z-order API,
    and an appended picture would paint over the cloned placeholders."""

    sp_tree = shape._element.getparent()
    sp_tree.remove(shape._element)
    sp_tree.insert(2, shape._element)  # after nvGrpSpPr, grpSpPr


def _apply_background(
    presentation,
    slide,
    background: BackgroundSpec | None,
    *,
    index: int,
    image_bytes_by_ref: Mapping[str, bytes],
    theme: ThemeSpec | None = None,
) -> None:
    """Paint the caller-resolved `background`. python-pptx has no background
    image primitive, so "image" and "pattern" are a full-slide picture sent
    to the back. `theme` only supplies a pattern's fallback accent color.
    """

    if background is None:
        return
    if background.kind == "solid":
        if not background.color:
            raise ValueError(f"slide {index}: a 'solid' background needs 'color'")
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(background.color)
    elif background.kind == "gradient":
        colors = background.gradient_colors or []
        if len(colors) != 2:
            raise ValueError(
                f"slide {index}: a 'gradient' background needs exactly 2 "
                "'gradient_colors'."
            )
        slide.background.fill.gradient()
        stops = slide.background.fill.gradient_stops
        stops[0].color.rgb = _rgb(colors[0])
        stops[0].position = 0.0
        stops[1].color.rgb = _rgb(colors[1])
        stops[1].position = 1.0
        slide.background.fill.gradient_angle = 45.0  # simple diagonal
    elif background.kind == "image":
        image_bytes = _resolve_image_bytes(
            background.image_ref,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
            field_label="background.image_ref",
        )
        payload, _w, _h = _load_embeddable_image(
            image_bytes, image_ref=background.image_ref or "", index=index
        )
        picture = slide.shapes.add_picture(
            io.BytesIO(payload),
            0,
            0,
            presentation.slide_width,
            presentation.slide_height,
        )
        _send_to_back(picture)
    elif background.kind == "pattern":
        if not background.color:
            raise ValueError(
                f"slide {index}: a 'pattern' background needs 'color' (its base fill)."
            )
        accent_color = background.accent_color or (
            theme.primary_color if theme else None
        )
        if not accent_color:
            raise ValueError(
                f"slide {index}: a 'pattern' background needs 'accent_color' "
                "(or a deck-wide 'theme.primary_color' to fall back to)."
            )
        if background.pattern is None:
            # Defense-in-depth: BackgroundSpec already rejects this.
            raise ValueError(f"slide {index}: a 'pattern' background needs 'pattern'")
        render_fn = background_render.PATTERN_RENDERERS[background.pattern]
        aspect_ratio = presentation.slide_width / presentation.slide_height
        png_bytes = render_fn(background.color, accent_color, aspect_ratio=aspect_ratio)
        picture = slide.shapes.add_picture(
            io.BytesIO(png_bytes),
            0,
            0,
            presentation.slide_width,
            presentation.slide_height,
        )
        _send_to_back(picture)
    else:
        raise ValueError(f"slide {index}: unknown background kind '{background.kind}'")


def _fractional_box_to_emu(presentation, spec: ShapeSpec) -> _Box:
    """Convert a `ShapeSpec`'s fractional box to EMU on the real slide size."""

    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    left = round(spec.left * slide_width)
    top = round(spec.top * slide_height)
    width = round(spec.width * slide_width)
    height = round(spec.height * slide_height)
    return (left, top, width, height)


def _box_area(box: _Box) -> int:
    _left, _top, width, height = box
    return max(0, width) * max(0, height)


def _overlap_area(a: _Box, b: _Box) -> int:
    a_left, a_top, a_width, a_height = a
    b_left, b_top, b_width, b_height = b
    overlap_width = min(a_left + a_width, b_left + b_width) - max(a_left, b_left)
    overlap_height = min(a_top + a_height, b_top + b_height) - max(a_top, b_top)
    if overlap_width <= 0 or overlap_height <= 0:
        return 0
    return overlap_width * overlap_height


def _significantly_overlaps(a: _Box, b: _Box) -> bool:
    smaller_area = min(_box_area(a), _box_area(b))
    if smaller_area == 0:
        return False
    return (_overlap_area(a, b) / smaller_area) > _SHAPE_OVERLAP_AREA_FRACTION_THRESHOLD


def _occupied_content_boxes(presentation, slide) -> list[_Box]:
    """Boxes of the slide's body content, snapshotted before shapes are added
    (so shapes may layer on each other). The title (near full-width box) and
    full-bleed backgrounds are excluded to avoid false-positive overlaps."""

    slide_area = presentation.slide_width * presentation.slide_height
    title_shape = slide.shapes.title
    # python-pptx builds a new wrapper per access, so compare the shared
    # lxml `_element`, not the wrappers.
    title_element = title_shape._element if title_shape is not None else None
    boxes = []
    for shape in slide.shapes:
        if title_element is not None and shape._element is title_element:
            continue
        box = (shape.left, shape.top, shape.width, shape.height)
        if _box_area(box) >= _BACKGROUND_SHAPE_AREA_FRACTION * slide_area:
            continue
        boxes.append(box)
    return boxes


def _add_shape(
    presentation,
    slide,
    spec: ShapeSpec,
    *,
    index: int,
    occupied_boxes: Sequence[_Box],
    warnings: list[str],
) -> None:
    box = _fractional_box_to_emu(presentation, spec)
    # Dropped, not raised: drawing it would bury the text, failing would cost
    # the whole deck for one misplaced decoration.
    if any(_significantly_overlaps(box, occupied) for occupied in occupied_boxes):
        warnings.append(
            f"slide {index}: a '{spec.shape_type}' shape at "
            f"left={spec.left}, top={spec.top}, width={spec.width}, "
            f"height={spec.height} overlapped the slide's own body content "
            "(text, table, chart, or image) and was left out — place shapes "
            "in a margin or corner, or on a slide without body content."
        )
        return
    left, top, width, height = box

    if spec.shape_type == "line":
        # No MSO_SHAPE autoshape for a straight line — python-pptx models it
        # as a connector, which has no `.fill` and no text frame, so
        # fill_color/text/text_style are silently no-ops for this one type.
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, left, top, left + width, top + height
        )
    else:
        shape = slide.shapes.add_shape(
            _MSO_SHAPE_BY_TYPE[spec.shape_type], left, top, width, height
        )
        if spec.fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(spec.fill_color)
        if spec.text:
            shape.text_frame.text = spec.text
            _enable_shrink_to_fit(shape.text_frame)
            _apply_text_frame_style(shape.text_frame, spec.text_style)

    if spec.rotation is not None:
        shape.rotation = spec.rotation
    if spec.line_color:
        shape.line.color.rgb = _rgb(spec.line_color)
    if spec.line_width_pt is not None:
        shape.line.width = Pt(spec.line_width_pt)


def _table_font_size_pt(n_rows: int) -> float:
    """Font size for a (header + data) row count."""

    for max_rows, size in _TABLE_FONT_SIZE_TIERS:
        if n_rows <= max_rows:
            return size
    return _TABLE_FONT_SIZE_FLOOR_PT


def _table_column_widths(
    total_width: int, headers: Sequence[str], rows: Sequence[Sequence[str]]
) -> list[int]:
    """Split `total_width` proportionally to each column's longest text, above
    a per-column floor."""

    n_cols = len(headers)
    weights = []
    for col_index, header in enumerate(headers):
        max_len = len(str(header))
        for row in rows:
            if col_index < len(row):
                max_len = max(max_len, len(str(row[col_index])))
        weights.append(max(max_len, 1))

    floor = round(total_width * _TABLE_COLUMN_WIDTH_FLOOR_FRACTION / n_cols)
    remaining = total_width - floor * n_cols
    weight_sum = sum(weights)
    widths = [floor + round(remaining * weight / weight_sum) for weight in weights]
    # python-pptx won't reflow a frame whose column widths don't sum to its
    # width, so rounding drift goes into the last column.
    widths[-1] += total_width - sum(widths)
    return widths


def _apply_table_font_size(
    pptx_table, *, n_rows: int, n_cols: int, size_pt: float
) -> None:
    size = Pt(size_pt)
    for row_index in range(n_rows):
        for col_index in range(n_cols):
            # Every paragraph: a '\n' in the cell text starts a new one.
            text_frame = pptx_table.cell(row_index, col_index).text_frame
            for paragraph in text_frame.paragraphs:
                paragraph.font.size = size
                for run in paragraph.runs:
                    run.font.size = size


def _place_table(slide, box: _Box, table: TableSpec | None, *, index: int) -> None:
    if table is None:
        raise ValueError(f"slide {index}: kind 'table' requires a 'table' spec")
    if not table.headers:
        raise ValueError(f"slide {index}: table needs at least one header column")

    n_cols = len(table.headers)
    n_rows = 1 + len(table.rows)
    if n_rows > MAX_TABLE_ROWS or n_cols > MAX_TABLE_COLS:
        raise ValueError(
            f"slide {index}: table is too large ({n_rows} rows x {n_cols} cols); "
            f"the limit is {MAX_TABLE_ROWS} rows x {MAX_TABLE_COLS} cols. Split "
            "the data across more slides or summarize it."
        )

    left, top, width, height = box
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    pptx_table = graphic_frame.table
    for col_index, header in enumerate(table.headers):
        pptx_table.cell(0, col_index).text = str(header)
    for row_index, row in enumerate(table.rows, start=1):
        for col_index in range(n_cols):
            value = row[col_index] if col_index < len(row) else ""
            pptx_table.cell(row_index, col_index).text = str(value)

    for col_index, col_width in enumerate(
        _table_column_widths(width, table.headers, table.rows)
    ):
        pptx_table.columns[col_index].width = col_width
    _apply_table_font_size(
        pptx_table, n_rows=n_rows, n_cols=n_cols, size_pt=_table_font_size_pt(n_rows)
    )


def _add_native_chart(slide, chart: ChartSpec, box: _Box, *, index: int) -> None:
    xl_type = _NATIVE_CHART_XL_TYPE[chart.chart_type]
    if chart.chart_type == "bubble":
        chart_data = BubbleChartData()
        for series in chart.series:
            x_values = series.x_values or []
            sizes = series.sizes or []
            if len(x_values) != len(series.values) or len(sizes) != len(series.values):
                raise ValueError(
                    f"slide {index}: bubble series '{series.name}': x_values "
                    f"({len(x_values)}), values ({len(series.values)}) and sizes "
                    f"({len(sizes)}) must all have the same length"
                )
            series_data = chart_data.add_series(series.name)
            for x, y, size in zip(x_values, series.values, sizes):
                series_data.add_data_point(x, y, size)
    elif chart.chart_type == "scatter":
        chart_data = XyChartData()
        for series in chart.series:
            x_values = series.x_values or []
            if len(x_values) != len(series.values):
                raise ValueError(
                    f"slide {index}: scatter series '{series.name}': x_values "
                    f"({len(x_values)}) and values ({len(series.values)}) must "
                    "have the same length"
                )
            series_data = chart_data.add_series(series.name)
            for x, y in zip(x_values, series.values):
                series_data.add_data_point(x, y)
    else:
        chart_data = CategoryChartData()
        chart_data.categories = chart.categories
        for series in chart.series:
            if len(series.values) != len(chart.categories):
                raise ValueError(
                    f"slide {index}: chart series '{series.name}' has "
                    f"{len(series.values)} value(s) but there are "
                    f"{len(chart.categories)} categories — they must match"
                )
            chart_data.add_series(series.name, series.values)

    left, top, width, height = box
    graphic_frame = slide.shapes.add_chart(
        xl_type, left, top, width, height, chart_data
    )
    built_chart = graphic_frame.chart
    if chart.title:
        built_chart.has_title = True
        built_chart.chart_title.text_frame.text = chart.title
    else:
        built_chart.has_title = False
    # A pie/doughnut's slices are its categories: without a legend nothing
    # says which is which. Bottom, like PowerPoint's default chart style.
    built_chart.has_legend = (
        chart.chart_type in ("pie", "doughnut") or len(chart_data) > 1
    )
    if built_chart.has_legend:
        built_chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        built_chart.legend.include_in_layout = False
    _apply_series_colors(built_chart, chart)


def _apply_series_colors(built_chart, chart: ChartSpec) -> None:
    """Apply `ChartSeries.color` as a fill or line color; a no-op for chart
    types in neither set."""

    if chart.chart_type not in (_SERIES_FILL_CHART_TYPES | _SERIES_LINE_CHART_TYPES):
        return
    plot = built_chart.plots[0]
    for series, spec_series in zip(plot.series, chart.series):
        if not spec_series.color:
            continue
        color = _rgb(spec_series.color)
        if chart.chart_type in _SERIES_FILL_CHART_TYPES:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        else:
            series.format.line.color.rgb = color


def _render_fallback_chart(chart: ChartSpec, *, index: int) -> bytes:
    try:
        if chart.chart_type == "heatmap":
            if not chart.series:
                raise ValueError("heatmap chart needs at least one series")
            return chart_render.render_heatmap(
                [series.name for series in chart.series],
                chart.categories,
                [series.values for series in chart.series],
                title=chart.title,
            )
        if chart.chart_type in ("waterfall", "funnel"):
            values = chart.series[0].values if chart.series else []
            if chart.chart_type == "waterfall":
                return chart_render.render_waterfall(
                    chart.categories, values, title=chart.title
                )
            return chart_render.render_funnel(
                chart.categories, values, title=chart.title
            )
        if chart.chart_type == "gantt":
            if not chart.gantt_tasks:
                raise ValueError("gantt chart needs 'gantt_tasks'")
            return chart_render.render_gantt(chart.gantt_tasks, title=chart.title)
        raise ValueError(f"unsupported chart_type '{chart.chart_type}'")
    except ValueError as exc:
        raise ValueError(f"slide {index}: {exc}") from exc


def _place_chart(slide, box: _Box, chart: ChartSpec | None, *, index: int) -> None:
    if chart is None:
        raise ValueError(f"slide {index}: kind 'chart' requires a 'chart' spec")
    if chart.chart_type in NATIVE_CHART_TYPES:
        _add_native_chart(slide, chart, box, index=index)
    elif chart.chart_type in IMAGE_FALLBACK_CHART_TYPES:
        png_bytes = _render_fallback_chart(chart, index=index)
        left, top, width, height = box
        with Image.open(io.BytesIO(png_bytes)) as image:
            img_w, img_h = image.size
        fit_left, fit_top, fit_width, fit_height = _fit_inside_box(
            img_w, img_h, left, top, width, height
        )
        slide.shapes.add_picture(
            io.BytesIO(png_bytes), fit_left, fit_top, fit_width, fit_height
        )
    else:
        raise ValueError(f"slide {index}: unknown chart_type '{chart.chart_type}'")


def _two_content_boxes(presentation, slide) -> tuple[_Box, _Box]:
    placeholders = _content_placeholders(slide)
    if len(placeholders) >= 2:
        left_box = _placeholder_box(placeholders[0])
        right_box = _placeholder_box(placeholders[1])
        _remove_shape(placeholders[0])
        _remove_shape(placeholders[1])
        return left_box, right_box

    box = (
        _placeholder_box(placeholders[0])
        if placeholders
        else _default_content_box(presentation)
    )
    if placeholders:
        _remove_shape(placeholders[0])
    left, top, width, height = box
    half_width = Emu(int((width - _COLUMN_GAP) / 2))
    left_box = (left, top, half_width, height)
    right_box = (Emu(int(left + half_width + _COLUMN_GAP)), top, half_width, height)
    return left_box, right_box


def _build_two_content(
    presentation,
    slide,
    spec: SlideSpec,
    *,
    index: int,
    image_bytes_by_ref: Mapping[str, bytes],
    body_style: TextStyle | None = None,
) -> None:
    frame = _first_empty_picture_frame(slide) if spec.image_ref else None
    text_column = _largest_text_placeholder(slide) if frame is not None else None
    if frame is not None and text_column is not None:
        # Template text + picture layout: fill both zones in place so the
        # bullets keep the template's box size and formatting.
        _fill_bullets(text_column.text_frame, spec.bullets or [], style=body_style)
        _fill_lead_zone(slide, text_column, spec.subtitle, style=body_style)
        _place_image(
            slide,
            frame,
            spec.image_ref,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
        )
        return

    left_box, right_box = _two_content_boxes(presentation, slide)
    left, top, width, height = left_box
    left_text_box = slide.shapes.add_textbox(left, top, width, height)
    _fill_bullets(left_text_box.text_frame, spec.bullets or [], style=body_style)

    if spec.image_ref:
        _place_image(
            slide,
            right_box,
            spec.image_ref,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
        )
    elif spec.table:
        _place_table(slide, right_box, spec.table, index=index)
    elif spec.chart:
        _place_chart(slide, right_box, spec.chart, index=index)
    # else: right column left empty; bullets-only is still a valid call.


class _ColumnZone(NamedTuple):
    texts: list  # the column's text placeholders, top to bottom
    icon: Any  # its empty picture frame, if any


class _ColumnLayout(NamedTuple):
    zones: list[_ColumnZone]
    intro: Any  # wide text placeholder above the columns
    main_frame: Any  # picture frame that belongs to no column
    label: Any  # small tag above the column group, e.g. "Titre secondaire"


def _template_column_layout(presentation, slide) -> _ColumnLayout | None:
    """Read a template slide's columns from its placeholders' geometry: text
    zones sharing a left edge form a column, left to right. None when the
    slide has fewer than two such columns."""

    content = _content_placeholders(slide)
    half_width = presentation.slide_width / 2
    texts = [
        ph for ph in content if ph.placeholder_format.type != PP_PLACEHOLDER.PICTURE
    ]
    groups: list[list] = []
    for ph in sorted((ph for ph in texts if ph.width < half_width), key=_left):
        if groups and abs(_left(ph) - _left(groups[-1][0])) < _COLUMN_ALIGN_TOLERANCE:
            groups[-1].append(ph)
        else:
            groups.append([ph])
    if len(groups) < 2:
        return None

    # Columns repeat one structure; a stray zone sharing a column's edge
    # (a group label above it) is not part of it, so keep the bottom ones.
    depth = min(len(group) for group in groups)
    frames = [ph for ph in content if _is_empty_picture_frame(ph)]
    zones = []
    for group in groups:
        icon = next(
            (
                frame
                for frame in frames
                if abs(_left(frame) - _left(group[0])) < _COLUMN_ALIGN_TOLERANCE
            ),
            None,
        )
        if icon is not None:
            frames.remove(icon)
        zones.append(_ColumnZone(sorted(group, key=_top)[-depth:], icon))
    strays = [ph for group in groups for ph in sorted(group, key=_top)[:-depth]]

    intros = [
        ph for ph in texts if ph.width >= half_width and ph.text_frame.text.strip()
    ]
    return _ColumnLayout(
        zones,
        intro=min(intros, key=_top, default=None),
        main_frame=max(frames, key=lambda ph: ph.width * ph.height, default=None),
        label=min(strays, key=_top, default=None),
    )


def _left(shape) -> int:
    return shape.left or 0


def _top(shape) -> int:
    return shape.top or 0


def _has_several_levels(placeholder) -> bool:
    paragraphs = placeholder.text_frame.paragraphs
    return len({p.level for p in paragraphs if p.text.strip()}) > 1


def _column_items(column: ColumnSpec, *, nested: bool) -> list[BulletItem]:
    """Heading then bullets as one list, for a column with a single text
    zone. `nested` (the zone's sample text uses two levels) puts the bullets
    one level under the heading; otherwise the heading is set in bold."""

    if not column.heading:
        return list(column.bullets or [])
    if not nested:
        return [BulletItem(text=f"**{column.heading}**"), *(column.bullets or [])]
    return [BulletItem(text=column.heading)] + [
        BulletItem(text=b.text, level=min(b.level + 1, _MAX_PARAGRAPH_LEVEL))
        for b in column.bullets or []
    ]


def _build_columns(
    presentation,
    slide,
    spec: SlideSpec,
    *,
    index: int,
    image_bytes_by_ref: Mapping[str, bytes],
    body_style: TextStyle | None = None,
    warnings: list[str],
) -> None:
    columns = spec.columns or []
    if not columns:
        raise ValueError(f"slide {index}: kind 'columns' requires 'columns'")
    if len(columns) > MAX_COLUMNS:
        raise ValueError(
            f"slide {index}: too many columns ({len(columns)}); the limit is "
            f"{MAX_COLUMNS}. Split them across more slides."
        )

    layout = _template_column_layout(presentation, slide)
    if layout is None:
        _build_generic_columns(
            presentation,
            slide,
            spec,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
            body_style=body_style,
            warnings=warnings,
        )
        return

    if len(columns) > len(layout.zones):
        raise ValueError(
            f"slide {index}: template_slide_index {spec.template_slide_index} "
            f"has {len(layout.zones)} columns but {len(columns)} were given — "
            "pass at most that many, or pick another template slide."
        )
    where = f"slide {index}: template_slide_index {spec.template_slide_index}"
    if spec.image_ref and layout.main_frame is None:
        warnings.append(f"{where} has no main picture frame, 'image_ref' ignored.")
    if spec.group_label and layout.label is None:
        warnings.append(f"{where} has no label zone, 'group_label' ignored.")
    if any(c.icon_ref for c in columns) and not any(z.icon for z in layout.zones):
        warnings.append(f"{where} has no icon frames, 'icon_ref' ignored.")

    for column, zone in zip(columns, layout.zones):
        # One zone with one text level is a caption: a heading plus bullets
        # would overflow it and flatten into a single style.
        caption = len(zone.texts) == 1 and not _has_several_levels(zone.texts[0])
        if caption and len(_column_items(column, nested=False)) > 1:
            raise ValueError(
                f"slide {index}: template_slide_index {spec.template_slide_index} "
                "has one-sentence caption columns — give each column a single "
                "'bullets' entry and no 'heading', or pick a template slide "
                "whose columns have a heading zone."
            )

    for column, zone in zip(columns, layout.zones):
        if len(zone.texts) == 1:
            text_frame = zone.texts[0].text_frame
            nested = _has_several_levels(zone.texts[0])
            items = _column_items(column, nested=nested)
            _fill_bullets(text_frame, items, style=body_style)
        else:
            if column.heading:
                _fill_bullets(
                    zone.texts[0].text_frame,
                    [BulletItem(text=column.heading)],
                    style=body_style,
                )
            if column.bullets:
                _fill_bullets(
                    zone.texts[-1].text_frame, column.bullets, style=body_style
                )
        if column.icon_ref and zone.icon is not None:
            _place_image(
                slide,
                zone.icon,
                column.icon_ref,
                index=index,
                image_bytes_by_ref=image_bytes_by_ref,
            )

    if spec.group_label and layout.label is not None:
        label = layout.label.text_frame
        _fill_bullets(label, _lead_items(label, spec.group_label), style=body_style)
    if spec.subtitle and layout.intro is not None:
        intro = layout.intro.text_frame
        _fill_bullets(intro, _lead_items(intro, spec.subtitle), style=body_style)
    if spec.image_ref and layout.main_frame is not None:
        _place_image(
            slide,
            layout.main_frame,
            spec.image_ref,
            index=index,
            image_bytes_by_ref=image_bytes_by_ref,
        )


def _build_generic_columns(
    presentation,
    slide,
    spec: SlideSpec,
    *,
    index: int,
    image_bytes_by_ref: Mapping[str, bytes],
    body_style: TextStyle | None = None,
    warnings: list[str],
) -> None:
    """No column layout to reuse: split the content area into equal columns,
    with the subtitle as an intro line above them."""

    if spec.image_ref:
        warnings.append(
            f"slide {index}: a 'columns' slide without a template picture "
            "frame cannot show 'image_ref', ignored."
        )
    columns = spec.columns or []
    left, top, width, height = _take_content_box(presentation, slide)
    if spec.subtitle:
        intro_height = Emu(int(height / 5))
        intro = slide.shapes.add_textbox(left, top, width, intro_height)
        _fill_bullets(
            intro.text_frame, [BulletItem(text=spec.subtitle)], style=body_style
        )
        top, height = Emu(top + intro_height), Emu(height - intro_height)

    count = len(columns)
    column_width = Emu(int((width - _COLUMN_GAP * (count - 1)) / count))
    icon_size = Emu(min(Inches(0.8), int(height / 4)))
    for position, column in enumerate(columns):
        column_left = Emu(int(left + position * (column_width + _COLUMN_GAP)))
        text_top, text_height = top, height
        if column.icon_ref:
            _place_image(
                slide,
                (column_left, top, icon_size, icon_size),
                column.icon_ref,
                index=index,
                image_bytes_by_ref=image_bytes_by_ref,
            )
            text_top = Emu(top + icon_size + _COLUMN_GAP)
            text_height = Emu(height - icon_size - _COLUMN_GAP)
        items = _column_items(column, nested=False)
        text_box = slide.shapes.add_textbox(
            column_left, text_top, column_width, text_height
        )
        _fill_bullets(text_box.text_frame, items, style=body_style)
