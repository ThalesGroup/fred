# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Summarizes an uploaded brand template's own slides for the `generate_pptx`
tool description (RFC §13). Computed once in `capability.validate_config` at
config-save time and stored on `PptxGeneratorConfig.template_slide_catalog`
(mirroring `ppt_filler`'s placeholder-schema extraction) rather than fetched
per turn — `tools(ctx)` is synchronous and this capability ships no
`middleware()` escape hatch to await a fetch from (RFC §4.1)."""

from __future__ import annotations

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

_MAX_TITLE_CHARS = 60
_MAX_SAMPLE_CHARS = 30
_MAX_ZONES_PER_SLIDE = 6
_NON_CONTENT_TYPES = (
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.PICTURE,
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.SLIDE_NUMBER,
)
_BULLET_TAGS = ("buChar", "buAutoNum", "buBlip")
# Mirrors deck_builder.MAX_SLIDES — an uploaded "template" with an unbounded
# slide count would otherwise inject an unbounded block into the tool
# description on every turn.
_MAX_CATALOG_SLIDES = 50


def describe_template_slides(presentation) -> str | None:
    """One numbered line per slide already in `presentation` (a template
    opened by the caller), or None if it has no slides at all — the common
    case for a template used only as a masters/theme source (RFC §4.2.1)."""

    slides = list(presentation.slides)
    lines = [
        f"  {slide_index}: {_slide_summary(slide)}"
        for slide_index, slide in enumerate(slides[:_MAX_CATALOG_SLIDES])
    ]
    if not lines:
        return None
    if len(slides) > _MAX_CATALOG_SLIDES:
        lines.append(f"  ... and {len(slides) - _MAX_CATALOG_SLIDES} more slide(s)")
    return "\n".join(lines)


def _slide_summary(slide) -> str:
    title = _title_text(slide)
    kinds = _content_shape_kinds(slide)
    summary = f'"{title}"' if title else "(untitled)"
    if kinds:
        summary += f" [{', '.join(kinds)}]"
    return summary + "".join(f"\n      {line}" for line in _text_zone_lines(slide))


def _text_zone_lines(slide) -> list[str]:
    """One line per sample-text zone, listing each paragraph level's look, so
    the agent can tell a heading/lead level from a bulleted one. Omitted for
    the plain case (one zone, one level) where there is nothing to choose."""

    zones = sorted(
        (
            ph
            for ph in slide.placeholders
            if ph.placeholder_format.type not in _NON_CONTENT_TYPES
            and ph.has_text_frame
            and ph.text_frame.text.strip()
        ),
        key=lambda ph: (ph.top or 0, ph.left or 0),
    )
    described = [_zone_levels(slide, ph) for ph in zones[:_MAX_ZONES_PER_SLIDE]]
    if len(described) < 2 and all(len(levels) < 2 for levels in described):
        return []
    return [
        f"text zone {number}: " + "; ".join(levels)
        for number, levels in enumerate(described, start=1)
    ]


def _zone_levels(slide, placeholder) -> list[str]:
    """`level N = <look> ("sample…")` per distinct level, in reading order.
    N is relative to the zone's lowest level, matching `BulletItem.level`."""

    paragraphs = [p for p in placeholder.text_frame.paragraphs if p.text.strip()]
    base = min(p.level for p in paragraphs)
    seen: dict[int, str] = {}
    for paragraph in paragraphs:
        if paragraph.level in seen:
            continue
        chain = _level_style_chain(slide, placeholder, paragraph.level)
        looks = ["bullet" if _has_bullet(paragraph, chain) else "no bullet"]
        if _is_bold(paragraph, chain):
            looks.append("bold")
        sample = paragraph.text.strip().replace("\n", " ").replace("\x0b", " ")
        if len(sample) > _MAX_SAMPLE_CHARS:
            sample = sample[: _MAX_SAMPLE_CHARS - 1].rstrip() + "…"
        seen[paragraph.level] = (
            f'level {paragraph.level - base} = {", ".join(looks)} ("{sample}")'
        )
    return list(seen.values())


def _level_style_chain(slide, placeholder, level: int) -> list:
    """Inherited `a:lvlNpPr` elements for `level`, nearest first: slide,
    layout and master placeholder list styles, then the master body style."""

    tag = qn(f"a:lvl{level + 1}pPr")
    chain = []
    shape = placeholder
    while shape is not None:
        list_style = shape._element.find(f".//{qn('a:lstStyle')}")
        found = list_style.find(tag) if list_style is not None else None
        if found is not None:
            chain.append(found)
        shape = getattr(shape, "_base_placeholder", None)
    tx_styles = slide.slide_layout.slide_master._element.find(qn("p:txStyles"))
    body_style = tx_styles.find(qn("p:bodyStyle")) if tx_styles is not None else None
    found = body_style.find(tag) if body_style is not None else None
    if found is not None:
        chain.append(found)
    return chain


def _has_bullet(paragraph, chain: list) -> bool:
    own = paragraph._p.pPr
    for p_pr in ([own] if own is not None else []) + chain:
        for child in p_pr:
            name = child.tag.rsplit("}", 1)[-1]
            if name == "buNone":
                return False
            if name in _BULLET_TAGS:
                return True
    return False


def _is_bold(paragraph, chain: list) -> bool:
    first_run = paragraph._p.find(qn("a:r"))
    r_pr = first_run.find(qn("a:rPr")) if first_run is not None else None
    candidates = [r_pr] + [p_pr.find(qn("a:defRPr")) for p_pr in chain]
    for props in candidates:
        if props is not None and props.get("b") is not None:
            return props.get("b") in ("1", "true")
    return False


def _title_text(slide) -> str:
    title_shape = slide.shapes.title
    if title_shape is None or not title_shape.has_text_frame:
        return ""
    text = title_shape.text_frame.text.strip().replace("\n", " ")
    if len(text) > _MAX_TITLE_CHARS:
        text = text[: _MAX_TITLE_CHARS - 1].rstrip() + "…"
    return text


def _content_shape_kinds(slide) -> list[str]:
    """A short tag per notable non-title shape, so the agent can tell a
    bullet-heavy slide from a picture/table/chart slide without guessing
    from the title text alone."""

    title_shape = slide.shapes.title
    title_id = title_shape.shape_id if title_shape is not None else None
    kinds: list[str] = []
    for shape in slide.shapes:
        if title_id is not None and shape.shape_id == title_id:
            continue
        kind = _shape_kind_tag(shape)
        if kind and kind not in kinds:
            kinds.append(kind)
    return kinds


def _shape_kind_tag(shape) -> str | None:
    """Tag `shape` only when it's a PLACEHOLDER — the only kind of shape
    `deck_builder`'s content-fill functions ever remove/replace
    (`_body_placeholder`/`_take_content_box` iterate `slide.placeholders`
    only). A static, non-placeholder picture/table/chart/textbox sitting
    decoratively on the slide is left untouched either way, so tagging it
    here would wrongly promise the agent that its content overwrites it."""

    if not shape.is_placeholder:
        return None
    if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
        return "image"
    if shape.has_table:
        return "table"
    if shape.has_chart:
        return "chart"
    if shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    return None
