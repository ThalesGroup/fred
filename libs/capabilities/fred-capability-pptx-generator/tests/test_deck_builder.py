# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Each slide variant must produce a valid, openable `.pptx` with the right
shapes/content (RFC §4.3) — including table/chart embedding and the
native-vs-fallback chart routing decision."""

from __future__ import annotations

import copy
import io
from typing import Any, cast

import pydantic
import pytest
from fred_capability_pptx_generator import deck_builder
from fred_capability_pptx_generator.models import (
    BackgroundSpec,
    BulletItem,
    ChartSeries,
    ChartSpec,
    ColumnSpec,
    ShapeSpec,
    ShapeType,
    SlideSpec,
    TableSpec,
    TextStyle,
    ThemeSpec,
)
from fred_capability_pptx_generator.palettes import THEME_PRESETS
from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn as qn_
from pptx.presentation import Presentation as PresentationType
from pptx.util import Inches, Pt
from pydantic import ValidationError

_ALL_SHAPE_TYPES: list[ShapeType] = [
    "rectangle",
    "rounded_rectangle",
    "oval",
    "triangle",
    "diamond",
    "arrow_right",
    "arrow_left",
    "arrow_up",
    "arrow_down",
    "line",
    "star",
    "pentagon",
    "hexagon",
    "chevron",
    "cloud",
]


def _png_bytes(width: int = 40, height: int = 20) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (255, 0, 0)).save(buffer, format="PNG")
    return buffer.getvalue()


def _open(deck_bytes: bytes) -> PresentationType:
    return Presentation(io.BytesIO(deck_bytes))


def _build(slides, *, template_bytes=None, image_bytes_by_ref=None, warnings=None):
    return deck_builder.build(
        "Deck",
        slides,
        template_bytes=template_bytes,
        image_bytes_by_ref=image_bytes_by_ref or {},
        warnings=warnings,
    )


def test_title_slide_sets_title_and_subtitle():
    slides = [SlideSpec(kind="title", title="Q3 Results", subtitle="Board update")]
    prs = _open(_build(slides))

    assert len(prs.slides) == 1
    slide = prs.slides[0]
    assert slide.shapes.title.text_frame.text == "Q3 Results"
    other_texts = [
        ph.text_frame.text
        for ph in slide.placeholders
        if ph.placeholder_format.idx != 0
    ]
    assert "Board update" in other_texts


def test_section_header_slide_sets_title():
    slides = [SlideSpec(kind="section_header", title="Part Two", subtitle="Deep dive")]
    slide = _open(_build(slides)).slides[0]
    assert slide.shapes.title.text_frame.text == "Part Two"


def test_bullets_slide_has_nested_levels():
    bullets = [
        BulletItem(text="Top", level=0),
        BulletItem(text="Nested", level=1),
        BulletItem(text="Deeper", level=2),
    ]
    slides = [SlideSpec(kind="bullets", title="Agenda", bullets=bullets)]
    slide = _open(_build(slides)).slides[0]

    body = next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    paragraphs = body.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Top", "Nested", "Deeper"]
    assert [p.level for p in paragraphs] == [0, 1, 2]


def test_title_and_bullets_text_frames_shrink_to_fit():
    # Regression guard: text must shrink on overflow rather than spill out of
    # its box (verified empirically against the installed LibreOffice, which
    # also renders the PDF preview — it computes a real shrink from this
    # marker, not just PowerPoint proper).
    slides = [
        SlideSpec(
            kind="bullets",
            title="Agenda",
            bullets=[BulletItem(text="Item")],
        )
    ]
    slide = _open(_build(slides)).slides[0]

    title_shape = slide.shapes.title
    assert title_shape is not None
    title_tf = title_shape.text_frame
    assert title_tf.word_wrap is True
    assert title_tf.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    # python-pptx's stubs type a placeholder as a broad shape union that
    # doesn't statically expose `.text_frame` — every text-bearing
    # placeholder has one at runtime, hence the cast.
    body = cast(
        Any, next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    )
    assert body.text_frame.word_wrap is True
    assert body.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def test_shape_text_shrinks_to_fit():
    shape_spec = ShapeSpec(
        shape_type="rectangle", left=0.1, top=0.1, width=0.2, height=0.1, text="Step 1"
    )
    slides = [SlideSpec(kind="title", title="Flow", shapes=[shape_spec])]
    slide = _open(_build(slides)).slides[0]

    # Same stub-gap cast as above: `has_text_frame` isn't a pyright TypeGuard.
    shape = next(
        cast(Any, s)
        for s in slide.shapes
        if s.has_text_frame and cast(Any, s).text_frame.text == "Step 1"
    )
    assert shape.text_frame.word_wrap is True
    assert shape.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def test_table_slide_renders_headers_and_rows():
    table = TableSpec(headers=["Name", "Score"], rows=[["Alice", "9"], ["Bob", "7"]])
    slides = [SlideSpec(kind="table", title="Scores", table=table)]
    slide = _open(_build(slides)).slides[0]

    table_shapes = [shape for shape in slide.shapes if shape.has_table]
    assert len(table_shapes) == 1
    pptx_table = table_shapes[0].table
    assert pptx_table.cell(0, 0).text == "Name"
    assert pptx_table.cell(0, 1).text == "Score"
    assert pptx_table.cell(1, 0).text == "Alice"
    assert pptx_table.cell(2, 1).text == "7"


def test_table_missing_spec_raises():
    slides = [SlideSpec(kind="table", title="No table")]
    with pytest.raises(ValueError, match="requires a 'table' spec"):
        _build(slides)


def test_table_over_row_cap_raises():
    rows = [["x"] for _ in range(deck_builder.MAX_TABLE_ROWS)]  # + header = over cap
    table = TableSpec(headers=["Col"], rows=rows)
    slides = [SlideSpec(kind="table", title="Big", table=table)]
    with pytest.raises(ValueError, match="too large"):
        _build(slides)


def test_table_over_col_cap_raises():
    headers = [f"c{i}" for i in range(deck_builder.MAX_TABLE_COLS + 1)]
    table = TableSpec(headers=headers, rows=[])
    slides = [SlideSpec(kind="table", title="Wide", table=table)]
    with pytest.raises(ValueError, match="too large"):
        _build(slides)


def test_native_chart_slide_adds_editable_chart_with_title():
    chart = ChartSpec(
        chart_type="bar",
        title="Revenue",
        categories=["Q1", "Q2"],
        series=[ChartSeries(name="EMEA", values=[10, 20])],
    )
    slides = [SlideSpec(kind="chart", title="Revenue", chart=chart)]
    slide = _open(_build(slides)).slides[0]

    chart_shapes = [shape for shape in slide.shapes if shape.has_chart]
    assert len(chart_shapes) == 1
    assert chart_shapes[0].chart.chart_title.text_frame.text == "Revenue"


@pytest.mark.parametrize(
    ("chart_type", "series_count", "expect_legend"),
    [
        ("pie", 1, True),
        ("doughnut", 1, True),
        ("column", 1, False),
        ("column", 2, True),
    ],
)
def test_native_chart_legend_sits_below_when_shown(
    chart_type, series_count, expect_legend
):
    series = [
        ChartSeries(name=f"S{n}", values=[45.0, 35.0, 20.0])
        for n in range(series_count)
    ]
    chart = ChartSpec(
        chart_type=chart_type, categories=["Europe", "Intl", "France"], series=series
    )
    slide = _open(_build([SlideSpec(kind="chart", title="T", chart=chart)])).slides[0]

    (built,) = [cast(Any, s).chart for s in slide.shapes if s.has_chart]
    assert built.has_legend is expect_legend
    if expect_legend:
        assert built.legend.position == XL_LEGEND_POSITION.BOTTOM
        assert built.legend.include_in_layout is False


def test_native_scatter_chart_uses_xy_series():
    chart = ChartSpec(
        chart_type="scatter",
        series=[ChartSeries(name="S1", x_values=[1, 2, 3], values=[4, 5, 6])],
    )
    slides = [SlideSpec(kind="chart", title="Scatter", chart=chart)]
    slide = _open(_build(slides)).slides[0]
    assert any(shape.has_chart for shape in slide.shapes)


def test_native_category_chart_rejects_series_values_length_mismatch():
    chart = ChartSpec(
        chart_type="bar",
        categories=["Q1", "Q2", "Q3"],
        series=[ChartSeries(name="EMEA", values=[10, 20])],
    )
    slides = [SlideSpec(kind="chart", title="Bad", chart=chart)]
    with pytest.raises(ValueError, match="slide 1.*3 categories"):
        _build(slides)


def test_native_bubble_chart_length_mismatch_error_names_the_slide():
    chart = ChartSpec(
        chart_type="bubble",
        series=[ChartSeries(name="S1", x_values=[1, 2], values=[4], sizes=[10, 20])],
    )
    slides = [
        SlideSpec(kind="title", title="Intro"),
        SlideSpec(kind="chart", title="Bubble", chart=chart),
    ]
    with pytest.raises(ValueError, match=r"slide 2:.*bubble series"):
        _build(slides)


def test_native_bubble_chart_uses_bubble_series():
    chart = ChartSpec(
        chart_type="bubble",
        series=[
            ChartSeries(name="S1", x_values=[1, 2], values=[4, 5], sizes=[10, 20]),
        ],
    )
    slides = [SlideSpec(kind="chart", title="Bubble", chart=chart)]
    slide = _open(_build(slides)).slides[0]
    assert any(shape.has_chart for shape in slide.shapes)


@pytest.mark.parametrize(
    ("chart_type", "categories", "values"),
    [
        ("waterfall", ["Start", "Delta", "End"], [100, -20, 0]),
        ("funnel", ["Visits", "Signups", "Purchases"], [1000, 400, 120]),
    ],
)
def test_fallback_chart_slide_embeds_picture(chart_type, categories, values):
    chart = ChartSpec(
        chart_type=chart_type,
        categories=categories,
        series=[ChartSeries(name="Value", values=values)],
    )
    slides = [SlideSpec(kind="chart", title=chart_type, chart=chart)]
    slide = _open(_build(slides)).slides[0]

    pictures = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1
    # A fallback chart must never leave a native chart shape behind too.
    assert not any(shape.has_chart for shape in slide.shapes)


def test_fallback_chart_picture_is_aspect_fit_not_stretched():
    # The heatmap fallback's matplotlib figure (8x3in, aspect ~2.67) has a
    # different aspect ratio than the default content box (9x5in, aspect
    # 1.8) — if the picture were naively stretched to the box, both its
    # width AND height would exactly equal the box's; aspect-fit leaves at
    # least one dimension strictly smaller.
    chart = ChartSpec(
        chart_type="heatmap",
        categories=["Mon", "Tue"],
        series=[ChartSeries(name="Team A", values=[1, 2])],
    )
    slides = [SlideSpec(kind="chart", title="Heat", chart=chart)]
    slide = _open(_build(slides)).slides[0]

    pictures = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1
    picture = pictures[0]
    _, _, box_width, box_height = deck_builder._default_content_box(Presentation())
    assert picture.width < box_width or picture.height < box_height


def test_heatmap_fallback_uses_series_as_rows():
    chart = ChartSpec(
        chart_type="heatmap",
        categories=["Mon", "Tue"],
        series=[
            ChartSeries(name="Team A", values=[1, 2]),
            ChartSeries(name="Team B", values=[3, 4]),
        ],
    )
    slides = [SlideSpec(kind="chart", title="Heat", chart=chart)]
    slide = _open(_build(slides)).slides[0]
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)


def test_chart_spec_rejects_gantt_without_tasks():
    with pytest.raises(pydantic.ValidationError, match="gantt_tasks"):
        ChartSpec(chart_type="gantt", gantt_tasks=None)


def test_chart_spec_rejects_category_chart_without_categories():
    with pytest.raises(pydantic.ValidationError, match="categories"):
        ChartSpec(chart_type="bar", series=[ChartSeries(name="s", values=[1])])


def test_chart_spec_rejects_chart_without_series():
    with pytest.raises(pydantic.ValidationError, match="series"):
        ChartSpec(chart_type="bar", categories=["a"])


def test_chart_spec_scatter_does_not_require_categories():
    # scatter/bubble plot points, not categories — must not be rejected.
    ChartSpec(
        chart_type="scatter", series=[ChartSeries(name="s", x_values=[1], values=[2])]
    )


def test_gantt_fallback_requires_tasks():
    # `ChartSpec`'s own model_validator now rejects a gantt chart with no
    # gantt_tasks at the pydantic boundary (the earlier, clean-error gate) —
    # bypass it at both levels via model_construct to exercise deck_builder's
    # own defense-in-depth guard for the same case.
    chart = ChartSpec.model_construct(
        chart_type="gantt", categories=[], series=[], gantt_tasks=None
    )
    slides = [SlideSpec.model_construct(kind="chart", title="Plan", chart=chart)]
    with pytest.raises(ValueError, match="gantt_tasks"):
        _build(slides)


def test_chart_missing_spec_raises():
    slides = [SlideSpec(kind="chart", title="No chart")]
    with pytest.raises(ValueError, match="requires a 'chart' spec"):
        _build(slides)


def test_unknown_chart_type_raises():
    # Bypass pydantic's Literal validation (and the has-required-data
    # model_validator, by giving it data that satisfies it) to exercise the
    # builder's own guard for a chart_type string it doesn't recognize.
    chart = ChartSpec.model_construct(
        chart_type="pareto",
        categories=["a"],
        series=[ChartSeries(name="s", values=[1])],
    )
    slides = [SlideSpec.model_construct(kind="chart", title="Bad", chart=chart)]
    with pytest.raises(ValueError, match="unknown chart_type"):
        _build(slides)


def test_image_slide_embeds_picture():
    slides = [SlideSpec(kind="image", title="Logo", image_ref="doc-1")]
    slide = _open(_build(slides, image_bytes_by_ref={"doc-1": _png_bytes()})).slides[0]

    pictures = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1


def test_mostly_transparent_icon_is_rejected_as_image_ref():
    buffer = io.BytesIO()
    icon = Image.new("RGBA", (40, 40), (0, 0, 0, 0))
    icon.paste((0, 110, 184, 255), (15, 15, 25, 25))
    icon.save(buffer, format="PNG")
    slides = [SlideSpec(kind="image", title="T", image_ref="picto")]

    with pytest.raises(ValueError, match="transparent icon"):
        _build(slides, image_bytes_by_ref={"picto": buffer.getvalue()})


def test_unknown_image_ref_raises():
    slides = [SlideSpec(kind="image", title="Logo", image_ref="missing-doc")]
    with pytest.raises(
        ValueError, match="image_ref 'missing-doc' could not be resolved"
    ):
        _build(slides)


def test_image_missing_ref_raises():
    slides = [SlideSpec(kind="image", title="Logo")]
    with pytest.raises(ValueError, match="'image_ref' is required"):
        _build(slides)


def test_two_content_slide_has_bullets_and_table():
    table = TableSpec(headers=["A"], rows=[["1"]])
    slides = [
        SlideSpec(
            kind="two_content",
            title="Comparison",
            bullets=[BulletItem(text="Point one")],
            table=table,
        )
    ]
    slide = _open(_build(slides)).slides[0]

    text_shapes = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame and "Point one" in shape.text_frame.text
    ]
    table_shapes = [shape for shape in slide.shapes if shape.has_table]
    assert text_shapes
    assert len(table_shapes) == 1


def test_two_content_prefers_image_over_table_and_chart():
    slides = [
        SlideSpec(
            kind="two_content",
            title="Priority",
            bullets=[BulletItem(text="Left")],
            image_ref="doc-1",
            table=TableSpec(headers=["A"], rows=[]),
        )
    ]
    slide = _open(_build(slides, image_bytes_by_ref={"doc-1": _png_bytes()})).slides[0]

    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    assert not any(shape.has_table for shape in slide.shapes)


def test_speaker_notes_are_set():
    slides = [SlideSpec(kind="title", title="Cover", speaker_notes="Say hello first.")]
    slide = _open(_build(slides)).slides[0]
    assert slide.notes_slide.notes_text_frame.text == "Say hello first."


def test_multi_slide_deck_with_brand_template():
    template_buffer = io.BytesIO()
    Presentation().save(template_buffer)

    slides = [
        SlideSpec(kind="title", title="Cover"),
        SlideSpec(kind="bullets", title="Agenda", bullets=[BulletItem(text="Item")]),
    ]
    prs = _open(_build(slides, template_bytes=template_buffer.getvalue()))
    assert len(prs.slides) == 2


def test_unknown_slide_kind_raises():
    slide = SlideSpec.model_construct(kind="unsupported", title="Bad")
    with pytest.raises(ValueError, match="unknown slide kind"):
        _build([slide])


# ---------------------------------------------------------------------------
# Layout resolution — structural fallback for a custom-named template (RFC §14)
# ---------------------------------------------------------------------------


def _custom_named_template_bytes() -> bytes:
    """The default python-pptx template's layouts, RENAMED so none of them
    match `_LAYOUT_NAMES_BY_KIND` — same trap a real, non-English/custom
    brand template falls into. Their actual placeholder STRUCTURE (which
    `_best_structural_layout_match` reads) is untouched."""

    prs = Presentation()
    for layout in prs.slide_layouts:
        layout.name = f"Modèle personnalisé — {layout.name}"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_structural_fallback_gives_image_its_own_picture_placeholder_layout():
    # Before RFC §14: bullets/image/table/chart all fell back to the SAME
    # hardcoded index whenever names didn't match — an image slide would
    # never land on a layout with an actual picture placeholder in a
    # custom-named template. Checked on `slide_layout` (the resolved layout)
    # rather than the built slide's own placeholders: `_take_content_box`
    # removes and replaces the content placeholder regardless of kind, by
    # design, same as every other kind.
    template_bytes = _custom_named_template_bytes()
    slides = [SlideSpec(kind="image", title="Cover", image_ref="doc-1")]
    slide = _open(
        _build(
            slides,
            template_bytes=template_bytes,
            image_bytes_by_ref={"doc-1": _png_bytes()},
        )
    ).slides[0]
    assert any(
        ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        for ph in slide.slide_layout.placeholders
    )


def test_structural_fallback_does_not_collapse_title_and_bullets_onto_one_layout():
    template_bytes = _custom_named_template_bytes()
    slides = [
        SlideSpec(kind="title", title="Cover"),
        SlideSpec(kind="bullets", title="Agenda", bullets=[BulletItem(text="Item")]),
    ]
    prs = _open(_build(slides, template_bytes=template_bytes))
    title_layout_name = prs.slides[0].slide_layout.name
    bullets_layout_name = prs.slides[1].slide_layout.name
    assert title_layout_name != bullets_layout_name


def test_structural_fallback_prefers_layout_with_two_content_placeholders():
    template_bytes = _custom_named_template_bytes()
    slides = [
        SlideSpec(
            kind="two_content",
            title="Compare",
            bullets=[BulletItem(text="Left")],
            table=TableSpec(headers=["h"], rows=[]),
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]
    non_title_placeholders = [
        ph
        for ph in slide.slide_layout.placeholders
        if ph.placeholder_format.type
        not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    ]
    assert len(non_title_placeholders) >= 2


def test_resolve_layout_raises_a_clear_error_for_a_template_with_no_layouts():
    # A pathological/corrupted template `validate_config` only checked opens
    # as a `.pptx`, not that it has any layout to build from — this must
    # surface as the same agent-correctable ValueError as every other
    # structural problem, never an unhandled IndexError. Exercised directly
    # against `_resolve_layout` — constructing a real zero-layout `.pptx`
    # would mean reaching into python-pptx's private slide-master XML.
    with pytest.raises(ValueError, match="no slide layouts"):
        deck_builder._resolve_layout("title", [], {})


def test_default_content_box_scales_with_custom_slide_size():
    # Regression guard (RFC §14): a fixed Inches() fallback box, sized for
    # the default 10x7.5in presentation, could extend past the bounds of a
    # smaller or differently-proportioned brand template's actual slide —
    # including past a furniture placeholder (footer/page number) near the
    # bottom. The box must scale with the presentation it's built for.
    prs = Presentation()
    prs.slide_width = Inches(6)
    prs.slide_height = Inches(4)
    left, top, width, height = deck_builder._default_content_box(prs)
    assert left + width <= Inches(6)
    assert top + height <= Inches(4)


# ---------------------------------------------------------------------------
# Shapes (RFC §11.2)
# ---------------------------------------------------------------------------


def test_all_curated_shape_types_render_without_error():
    shapes = [
        ShapeSpec(shape_type=shape_type, left=0.1, top=0.1, width=0.2, height=0.2)
        for shape_type in _ALL_SHAPE_TYPES
    ]
    slides = [SlideSpec(kind="title", title="Shapes", shapes=shapes)]
    slide = _open(_build(slides)).slides[0]

    autoshapes = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    ]
    connectors = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE
    ]
    # 14 curated autoshapes + 1 connector ("line") == len(_ALL_SHAPE_TYPES).
    assert len(autoshapes) + len(connectors) == len(_ALL_SHAPE_TYPES)
    assert len(connectors) == 1


def test_shape_applies_fill_line_rotation_and_text_style():
    shape = ShapeSpec(
        shape_type="rectangle",
        left=0.1,
        top=0.1,
        width=0.3,
        height=0.2,
        fill_color="#1565C0",
        line_color="#000000",
        line_width_pt=2.0,
        rotation=15.0,
        text="Node",
        text_style=TextStyle(bold=True, color="#FFFFFF"),
    )
    slides = [SlideSpec(kind="title", title="Shape", shapes=[shape])]
    slide = _open(_build(slides)).slides[0]

    rect = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
    assert rect.fill.fore_color.rgb == RGBColor.from_string("1565C0")
    assert rect.line.color.rgb == RGBColor.from_string("000000")
    assert rect.rotation == 15.0
    assert rect.text_frame.text == "Node"
    run = rect.text_frame.paragraphs[0].runs[0]
    assert run.font.bold is True
    assert run.font.color.rgb == RGBColor.from_string("FFFFFF")


def test_line_shape_renders_as_connector_with_line_style():
    # left/top in the bottom margin, well below the title slide's title and
    # subtitle placeholders — this test is about connector/line-style
    # behavior, not overlap detection (covered separately).
    shape = ShapeSpec(
        shape_type="line",
        left=0.05,
        top=0.9,
        width=0.15,
        height=0.05,
        line_color="#FF0000",
    )
    slides = [SlideSpec(kind="title", title="Line", shapes=[shape])]
    slide = _open(_build(slides)).slides[0]

    connector = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE)
    assert connector.line.color.rgb == RGBColor.from_string("FF0000")


def test_line_shape_allows_zero_width_or_height():
    # A perfectly horizontal or vertical line has one dimension at 0 — this
    # must not be rejected the way a degenerate autoshape box would be.
    horizontal = ShapeSpec(shape_type="line", left=0.1, top=0.5, width=0.3, height=0.0)
    vertical = ShapeSpec(shape_type="line", left=0.5, top=0.1, width=0.0, height=0.3)
    slides = [SlideSpec(kind="title", title="Lines", shapes=[horizontal, vertical])]
    slide = _open(_build(slides)).slides[0]
    connectors = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
    assert len(connectors) == 2


def test_line_shape_applies_rotation():
    # Same non-overlapping placement rationale as the test above.
    shape = ShapeSpec(
        shape_type="line", left=0.05, top=0.9, width=0.15, height=0.05, rotation=30.0
    )
    slides = [SlideSpec(kind="title", title="Line", shapes=[shape])]
    slide = _open(_build(slides)).slides[0]

    connector = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE)
    assert connector.rotation == 30.0


def test_too_many_shapes_raises():
    shapes = [
        ShapeSpec(shape_type="rectangle", left=0.0, top=0.0, width=0.05, height=0.05)
        for _ in range(deck_builder.MAX_SHAPES_PER_SLIDE + 1)
    ]
    slides = [SlideSpec(kind="title", title="Too many", shapes=shapes)]
    with pytest.raises(ValueError, match="too many shapes"):
        _build(slides)


def test_shape_invalid_hex_color_raises():
    with pytest.raises(ValidationError, match="not a valid hex color"):
        ShapeSpec(
            shape_type="rectangle",
            left=0.0,
            top=0.0,
            width=0.1,
            height=0.1,
            fill_color="blue",
        )


# ---------------------------------------------------------------------------
# Shape/content overlap guard (RFC §12 amendment — labeled badge shapes
# dropped on top of running bullet text, both illegible)
# ---------------------------------------------------------------------------


def test_shape_overlapping_bullet_text_is_dropped_with_warning():
    # A small labeled diamond centered over the bullets content box: the
    # deck still builds, without the shape, and the agent is told why.
    badge = ShapeSpec(
        shape_type="diamond",
        left=0.35,
        top=0.4,
        width=0.15,
        height=0.15,
        fill_color="#D9432A",
        text="Sécurité",
    )
    slides = [
        SlideSpec(
            kind="bullets",
            title="Les risques majeurs",
            bullets=[BulletItem(text="Sécurité des données : fuites, non-conformité")],
            shapes=[badge],
        )
    ]
    warnings: list[str] = []
    prs = _open(_build(slides, warnings=warnings))

    texts = [
        cast(Any, shape).text_frame.text
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
    ]
    assert "Sécurité" not in texts
    assert len(warnings) == 1
    assert "'diamond' shape" in warnings[0]
    assert "overlapped the slide's own body content" in warnings[0]


def test_only_the_overlapping_shape_is_dropped():
    # A corner accent next to the overlapping badge must survive.
    badge = ShapeSpec(
        shape_type="diamond", left=0.35, top=0.4, width=0.15, height=0.15, text="Badge"
    )
    corner = ShapeSpec(
        shape_type="oval", left=0.8, top=0.02, width=0.18, height=0.15, text="Coin"
    )
    slides = [
        SlideSpec(
            kind="bullets",
            title="Les risques majeurs",
            bullets=[BulletItem(text="Sécurité des données")],
            shapes=[badge, corner],
        )
    ]
    warnings: list[str] = []
    prs = _open(_build(slides, warnings=warnings))

    texts = [
        cast(Any, shape).text_frame.text
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
    ]
    assert "Coin" in texts
    assert "Badge" not in texts
    assert len(warnings) == 1


def test_shape_overlapping_only_the_title_does_not_raise():
    # Deliberate exclusion, not an oversight: PowerPoint/python-pptx default
    # layouts give the title placeholder a box spanning nearly the full
    # slide width regardless of how short the title text actually is, so
    # checking against it would false-positive-reject the exact corner
    # accent near the title this capability's own prompt guidance
    # recommends. Only BODY content (bullets/table/chart/image) is guarded.
    shape = ShapeSpec(shape_type="oval", left=0.3, top=0.02, width=0.2, height=0.1)
    slides = [SlideSpec(kind="bullets", title="Agenda", shapes=[shape])]
    _build(slides)  # must not raise


def test_shape_in_margin_corner_does_not_raise():
    # A legitimate corner accent (RFC §12 prompt guidance) must still work —
    # this is a regression guard against the overlap check being too strict.
    corner_circle = ShapeSpec(
        shape_type="oval", left=0.8, top=0.02, width=0.18, height=0.15
    )
    slides = [
        SlideSpec(
            kind="bullets",
            title="Agenda",
            bullets=[BulletItem(text="Item")],
            shapes=[corner_circle],
        )
    ]
    _build(slides)  # must not raise


def test_shapes_may_overlap_each_other():
    # Deliberate: the guard only checks overlap against pre-existing SLIDE
    # CONTENT, snapshotted once — decorative shapes may still layer on top
    # of each other (e.g. a colored circle behind a smaller accent circle).
    background_circle = ShapeSpec(
        shape_type="oval", left=0.75, top=0.02, width=0.24, height=0.24
    )
    accent_circle = ShapeSpec(
        shape_type="oval", left=0.8, top=0.06, width=0.12, height=0.12
    )
    slides = [
        SlideSpec(
            kind="title", title="Cover", shapes=[background_circle, accent_circle]
        )
    ]
    _build(slides)  # must not raise


def test_pattern_background_does_not_count_as_occupied_content():
    # A full-bleed background picture/pattern must NOT make every shape
    # "overlap" — only actual content (title/text/table/chart/image) counts.
    corner_circle = ShapeSpec(
        shape_type="oval", left=0.8, top=0.02, width=0.18, height=0.15
    )
    slides = [
        SlideSpec(
            kind="title",
            title="Cover",
            background=BackgroundSpec(
                kind="pattern",
                pattern="dot_grid",
                color="#1A1A1A",
                accent_color="#4D4D4D",
            ),
            shapes=[corner_circle],
        )
    ]
    _build(slides)  # must not raise


# ---------------------------------------------------------------------------
# Backgrounds (RFC §11.2)
# ---------------------------------------------------------------------------


def test_solid_background_via_theme():
    slides = [SlideSpec(kind="title", title="Cover")]
    theme = ThemeSpec(background=BackgroundSpec(kind="solid", color="#112233"))
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]
    assert slide.background.fill.fore_color.rgb == RGBColor.from_string("112233")


def test_gradient_background_sets_both_stops():
    background = BackgroundSpec(kind="gradient", gradient_colors=["#FF0000", "#0000FF"])
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    slide = _open(_build(slides)).slides[0]

    stops = slide.background.fill.gradient_stops
    assert stops[0].color.rgb == RGBColor.from_string("FF0000")
    assert stops[1].color.rgb == RGBColor.from_string("0000FF")


def test_per_slide_background_overrides_theme_background():
    theme = ThemeSpec(background=BackgroundSpec(kind="solid", color="#000000"))
    slides = [
        SlideSpec(
            kind="title",
            title="Cover",
            background=BackgroundSpec(kind="solid", color="#FFFFFF"),
        )
    ]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]
    assert slide.background.fill.fore_color.rgb == RGBColor.from_string("FFFFFF")


def test_image_background_sent_behind_placeholders():
    background = BackgroundSpec(kind="image", image_ref="bg-1")
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    deck_bytes = _build(slides, image_bytes_by_ref={"bg-1": _png_bytes(100, 60)})
    slide = _open(deck_bytes).slides[0]

    shapes = list(slide.shapes)
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
    # The title placeholder must still be present and paint on top.
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER for shape in shapes)


def test_background_image_missing_ref_raises():
    background = BackgroundSpec(kind="image", image_ref="missing-bg")
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    with pytest.raises(ValueError, match="background.image_ref 'missing-bg'"):
        _build(slides)


def test_background_image_no_ref_raises_with_field_label():
    # Regression: this branch used to hardcode "kind 'image'" even for a
    # non-'image' slide with an 'image' background (field_label bug).
    background = BackgroundSpec(kind="image")
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    with pytest.raises(ValueError, match="'background.image_ref' is required"):
        _build(slides)


def test_background_gradient_requires_exactly_two_colors():
    with pytest.raises(ValidationError, match="exactly 2"):
        BackgroundSpec(kind="gradient", gradient_colors=["#FFFFFF"])


def test_background_gradient_requires_gradient_colors_when_omitted_entirely():
    # Regression guard: a plain `field_validator` on `gradient_colors` is
    # SKIPPED ENTIRELY when the field is left at its default (never passed),
    # so omitting it altogether previously slipped past pydantic and failed
    # much later, deep inside the synchronous build — not at the schema
    # boundary the other malformed-input errors are caught at.
    with pytest.raises(ValidationError, match="exactly 2"):
        BackgroundSpec(kind="gradient")


def test_theme_background_image_decode_is_cached_across_slides():
    # A deck-wide theme.background image is the same bytes on every slide —
    # it must be decoded/transcoded once, not once per slide.
    deck_builder._decode_and_maybe_transcode.cache_clear()
    theme = ThemeSpec(background=BackgroundSpec(kind="image", image_ref="bg-1"))
    slides = [
        SlideSpec(kind="title", title="One"),
        SlideSpec(kind="title", title="Two"),
        SlideSpec(kind="title", title="Three"),
    ]
    deck_builder.build(
        "Deck",
        slides,
        template_bytes=None,
        image_bytes_by_ref={"bg-1": _png_bytes(100, 60)},
        theme=theme,
    )

    info = deck_builder._decode_and_maybe_transcode.cache_info()
    assert info.misses == 1
    assert info.hits == 2


# ---------------------------------------------------------------------------
# Text styling (RFC §11.2)
# ---------------------------------------------------------------------------


def test_title_style_applies_font_size_color_bold_italic():
    style = TextStyle(
        font="Georgia",
        size_pt=32,
        color="#0000FF",
        bold=True,
        italic=True,
        align="center",
    )
    slides = [SlideSpec(kind="title", title="Styled", title_style=style)]
    slide = _open(_build(slides)).slides[0]

    run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Georgia"
    assert run.font.size == Pt(32)
    assert run.font.color.rgb == RGBColor.from_string("0000FF")
    assert run.font.bold is True
    assert run.font.italic is True


def test_body_style_falls_back_to_theme_when_slide_leaves_it_unset():
    theme = ThemeSpec(
        heading_font="Georgia", body_font="Verdana", primary_color="#123456"
    )
    slides = [
        SlideSpec(kind="bullets", title="Agenda", bullets=[BulletItem(text="Item")])
    ]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]

    body = next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    run = body.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Verdana"
    assert run.font.color.rgb == RGBColor.from_string("123456")


def test_slide_body_style_overrides_theme_font_but_keeps_theme_color():
    theme = ThemeSpec(body_font="Verdana", primary_color="#123456")
    slides = [
        SlideSpec(
            kind="bullets",
            title="Agenda",
            bullets=[BulletItem(text="Item")],
            body_style=TextStyle(font="Courier New"),
        )
    ]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]

    body = next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    run = body.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Courier New"
    assert run.font.color.rgb == RGBColor.from_string("123456")


# ---------------------------------------------------------------------------
# Bullets: inline **bold**/*italic* (RFC §11.3)
# ---------------------------------------------------------------------------


def test_slide_style_explicit_empty_font_is_not_overridden_by_theme():
    # A regression guard: `_resolve_text_style` must treat an explicitly-set
    # but falsy value (font="") as "set" and NOT fall back to the theme font
    # for it — only a field the slide's style genuinely leaves unset (None)
    # should inherit the theme's default.
    theme = ThemeSpec(heading_font="Georgia")
    slides = [SlideSpec(kind="title", title="Styled", title_style=TextStyle(font=""))]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]

    title_shape = slide.shapes.title
    assert title_shape is not None
    run = title_shape.text_frame.paragraphs[0].runs[0]
    assert run.font.name != "Georgia"


def test_bullet_inline_markdown_produces_multiple_styled_runs():
    bullets = [BulletItem(text="This is **bold** and *italic* text")]
    slides = [SlideSpec(kind="bullets", title="Rich", bullets=bullets)]
    slide = _open(_build(slides)).slides[0]

    body = next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    runs = body.text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == [
        "This is ",
        "bold",
        " and ",
        "italic",
        " text",
    ]
    assert runs[1].font.bold is True
    assert runs[3].font.italic is True
    assert runs[0].font.bold is not True
    assert runs[0].font.italic is not True


def test_bullet_without_markdown_yields_single_run():
    bullets = [BulletItem(text="Plain text")]
    slides = [SlideSpec(kind="bullets", title="Plain", bullets=bullets)]
    slide = _open(_build(slides)).slides[0]

    body = next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    runs = body.text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == ["Plain text"]


# ---------------------------------------------------------------------------
# Native chart series coloring (RFC §11.2)
# ---------------------------------------------------------------------------


def test_chart_series_color_applies_fill_for_bar_chart():
    chart = ChartSpec(
        chart_type="bar",
        categories=["Q1"],
        series=[ChartSeries(name="S1", values=[1], color="#00FF00")],
    )
    slides = [SlideSpec(kind="chart", title="Bar", chart=chart)]
    slide = _open(_build(slides)).slides[0]

    chart_shape = next(s for s in slide.shapes if s.has_chart)
    series = chart_shape.chart.plots[0].series[0]
    assert series.format.fill.fore_color.rgb == RGBColor.from_string("00FF00")


def test_chart_series_color_applies_line_for_line_chart():
    chart = ChartSpec(
        chart_type="line",
        categories=["Q1"],
        series=[ChartSeries(name="S1", values=[1], color="#00FF00")],
    )
    slides = [SlideSpec(kind="chart", title="Line", chart=chart)]
    slide = _open(_build(slides)).slides[0]

    chart_shape = next(s for s in slide.shapes if s.has_chart)
    series = chart_shape.chart.plots[0].series[0]
    assert series.format.line.color.rgb == RGBColor.from_string("00FF00")


def test_chart_series_color_ignored_for_pie_chart():
    chart = ChartSpec(
        chart_type="pie",
        categories=["Q1", "Q2"],
        series=[ChartSeries(name="S1", values=[1, 2], color="#00FF00")],
    )
    slides = [SlideSpec(kind="chart", title="Pie", chart=chart)]
    # Must not raise — pie colors per-point, not per-series (RFC §11.2, §11.5).
    slide = _open(_build(slides)).slides[0]
    assert any(s.has_chart for s in slide.shapes)


# ---------------------------------------------------------------------------
# Curated theme presets (RFC §12.2a)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("preset_name", sorted(THEME_PRESETS))
def test_every_preset_resolves_without_error_with_strong_soft_split(preset_name):
    theme = ThemeSpec(preset=preset_name)
    slides = [
        SlideSpec(kind="title", title="Cover"),
        SlideSpec(kind="bullets", title="Agenda", bullets=[BulletItem(text="Item")]),
    ]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    prs = _open(deck_bytes)
    preset = THEME_PRESETS[preset_name]

    # Strong slide: either a solid/gradient fill, or (for pattern presets) a
    # full-bleed picture sent behind everything else.
    strong_slide = prs.slides[0]
    if preset.strong_background.kind == "pattern":
        first_shape = strong_slide.shapes[0]
        assert first_shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    elif preset.strong_background.kind == "solid":
        assert preset.strong_background.color is not None
        assert strong_slide.background.fill.fore_color.rgb == RGBColor.from_string(
            preset.strong_background.color.lstrip("#")
        )
    title_shape = strong_slide.shapes.title
    assert title_shape is not None
    title_run = title_shape.text_frame.paragraphs[0].runs[0]
    assert title_run.font.color.rgb == RGBColor.from_string(
        preset.strong_text_color.lstrip("#")
    )

    # Soft slide: plain solid fill of the preset's pale tint; body text keeps
    # its layout default color (no preset color injected for non-strong
    # kinds, RFC §12.2a).
    soft_slide = prs.slides[1]
    assert soft_slide.background.fill.fore_color.rgb == RGBColor.from_string(
        preset.soft_background_color.lstrip("#")
    )
    body = cast(
        Any,
        next(ph for ph in soft_slide.placeholders if ph.placeholder_format.idx != 0),
    )
    body_run = body.text_frame.paragraphs[0].runs[0]
    assert body_run.font.color.type is None


def test_preset_explicit_slide_background_still_wins():
    theme = ThemeSpec(preset="corporate_navy")
    override = BackgroundSpec(kind="solid", color="#00FF00")
    slides = [SlideSpec(kind="title", title="Cover", background=override)]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]
    assert slide.background.fill.fore_color.rgb == RGBColor.from_string("00FF00")


def test_preset_explicit_title_style_color_still_wins():
    theme = ThemeSpec(preset="corporate_navy")
    slides = [
        SlideSpec(kind="title", title="Cover", title_style=TextStyle(color="#123456"))
    ]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]
    title_shape = slide.shapes.title
    assert title_shape is not None
    run = title_shape.text_frame.paragraphs[0].runs[0]
    assert run.font.color.rgb == RGBColor.from_string("123456")


def test_preset_fonts_and_primary_color_default_theme_fields():
    theme = ThemeSpec(preset="sustainability_green")
    slides = [SlideSpec(kind="bullets", title="Agenda", bullets=[BulletItem(text="x")])]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]
    title_shape = slide.shapes.title
    assert title_shape is not None
    title_run = title_shape.text_frame.paragraphs[0].runs[0]
    assert title_run.font.name == THEME_PRESETS["sustainability_green"].heading_font


def test_unknown_theme_preset_raises_clear_error():
    with pytest.raises(ValidationError, match="not a known theme preset"):
        ThemeSpec(preset="not-a-real-preset")


# ---------------------------------------------------------------------------
# Adaptive table layout (RFC §12.2b)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("n_data_rows", "expected_pt"),
    [(0, 16.0), (5, 16.0), (9, 13.0), (14, 11.0), (15, 9.0), (40, 9.0)],
)
def test_table_font_size_tiers_at_boundaries(n_data_rows, expected_pt):
    headers = ["Col"]
    rows = [["x"] for _ in range(n_data_rows)]
    table = TableSpec(headers=headers, rows=rows)
    slides = [SlideSpec(kind="table", title="T", table=table)]
    slide = _open(_build(slides)).slides[0]

    graphic_frame = cast(Any, next(shape for shape in slide.shapes if shape.has_table))
    header_paragraph = graphic_frame.table.cell(0, 0).text_frame.paragraphs[0]
    assert header_paragraph.font.size == Pt(expected_pt)


def test_table_font_size_applies_to_every_line_of_a_multiline_cell():
    table = TableSpec(headers=["Famille"], rows=[["Génération\n(NxT, DCL)"]])
    slides = [SlideSpec(kind="table", title="T", table=table)]
    slide = _open(_build(slides)).slides[0]

    graphic_frame = cast(Any, next(shape for shape in slide.shapes if shape.has_table))
    paragraphs = graphic_frame.table.cell(1, 0).text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Génération", "(NxT, DCL)"]
    assert {p.font.size for p in paragraphs} == {Pt(16.0)}
    assert {r.font.size for p in paragraphs for r in p.runs} == {Pt(16.0)}


def test_table_column_width_is_proportional_to_content_length():
    table = TableSpec(
        headers=["A", "A much much longer header that needs more room"],
        rows=[["1", "short"], ["2", "short"]],
    )
    slides = [SlideSpec(kind="table", title="T", table=table)]
    slide = _open(_build(slides)).slides[0]

    graphic_frame = cast(Any, next(shape for shape in slide.shapes if shape.has_table))
    columns = graphic_frame.table.columns
    assert columns[1].width > columns[0].width
    # Widths must still sum to the box width the table was placed into.
    _, _, box_width, _box_height = deck_builder._default_content_box(Presentation())
    assert columns[0].width + columns[1].width == box_width


def test_table_column_width_never_crushed_below_floor():
    # A pathologically long column must not starve the others down to zero.
    table = TableSpec(
        headers=["x" * 500, "y"],
        rows=[],
    )
    slides = [SlideSpec(kind="table", title="T", table=table)]
    slide = _open(_build(slides)).slides[0]

    graphic_frame = cast(Any, next(shape for shape in slide.shapes if shape.has_table))
    columns = graphic_frame.table.columns
    assert columns[1].width > 0


# ---------------------------------------------------------------------------
# Procedural pattern backgrounds (RFC §12.2c)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("pattern", ["soft_circles", "diagonal_bands", "dot_grid"])
def test_pattern_background_embeds_full_bleed_picture_sent_to_back(pattern):
    background = BackgroundSpec(
        kind="pattern", pattern=pattern, color="#101010", accent_color="#FF8800"
    )
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    deck_bytes = _build(slides)
    prs = _open(deck_bytes)
    slide = prs.slides[0]

    shapes = list(slide.shapes)
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
    assert shapes[0].width == prs.slide_width
    assert shapes[0].height == prs.slide_height
    # The title placeholder must still be present and paint on top.
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER for shape in shapes)


def test_pattern_background_falls_back_to_theme_primary_color_for_accent():
    background = BackgroundSpec(kind="pattern", pattern="dot_grid", color="#101010")
    theme = ThemeSpec(primary_color="#00FF00")
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    deck_bytes = deck_builder.build(
        "Deck", slides, template_bytes=None, image_bytes_by_ref={}, theme=theme
    )
    slide = _open(deck_bytes).slides[0]
    assert slide.shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE


def test_pattern_background_without_accent_color_anywhere_raises():
    background = BackgroundSpec(kind="pattern", pattern="dot_grid", color="#101010")
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    with pytest.raises(ValueError, match="needs 'accent_color'"):
        _build(slides)


def test_pattern_background_without_base_color_raises():
    background = BackgroundSpec.model_construct(
        kind="pattern", pattern="dot_grid", color=None, accent_color="#FF8800"
    )
    slides = [SlideSpec(kind="title", title="Cover", background=background)]
    with pytest.raises(ValueError, match="needs 'color'"):
        _build(slides)


def test_pattern_background_requires_pattern_name():
    with pytest.raises(ValidationError, match="needs a 'pattern' name"):
        BackgroundSpec(kind="pattern", color="#101010", accent_color="#FF8800")


# ---------------------------------------------------------------------------
# template_slide_index — duplicating a real template slide (RFC §13)
# ---------------------------------------------------------------------------


def _template_bytes(build_slides) -> bytes:
    """A one-off `.pptx` whose slides are added by `build_slides(prs)` —
    used as `template_bytes` for `template_slide_index` tests."""

    prs = Presentation()
    build_slides(prs)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _add_designed_slide(prs) -> None:
    """A "Title and Content" slide with the template author's own sample
    text, a decorative accent shape, and a solid background — the shape of
    slide a real brand template would ship for `template_slide_index` to
    duplicate."""

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    assert title_shape is not None
    title_shape.text_frame.text = "Original title"
    cast(Any, slide.placeholders[1]).text_frame.text = "Original body"
    accent = cast(
        Any,
        slide.shapes.add_shape(
            MSO_SHAPE.STAR_5_POINT, Inches(0.2), Inches(0.2), Inches(0.5), Inches(0.5)
        ),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor.from_string("112233")
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("445566")


def test_template_slide_index_fills_placeholders_and_keeps_decoration():
    template_bytes = _template_bytes(_add_designed_slide)
    slides = [
        SlideSpec(
            kind="bullets",
            title="New title",
            bullets=[BulletItem(text="New bullet")],
            template_slide_index=0,
        )
    ]
    prs = _open(_build(slides, template_bytes=template_bytes))

    assert len(prs.slides) == 1
    slide = prs.slides[0]
    title_shape = slide.shapes.title
    assert title_shape is not None
    assert title_shape.text_frame.text == "New title"
    body = cast(
        Any, next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    )
    assert body.text_frame.text == "New bullet"
    # The decorative accent shape (not addressed by any SlideSpec field)
    # survives the duplication untouched.
    accents = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(accents) == 1
    assert cast(Any, accents[0]).fill.fore_color.rgb == RGBColor.from_string("112233")
    # The template's own background is preserved — no theme/preset was set.
    assert slide.background.fill.fore_color.rgb == RGBColor.from_string("445566")


def test_template_slide_index_duplicates_picture_with_working_relationship():
    def _add_picture_slide(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        title_shape = slide.shapes.title
        assert title_shape is not None
        title_shape.text_frame.text = "Picture slide"
        image_buffer = io.BytesIO()
        Image.new("RGB", (12, 8), (5, 6, 7)).save(image_buffer, format="PNG")
        image_buffer.seek(0)
        slide.shapes.add_picture(
            image_buffer, Inches(0.5), Inches(0.5), Inches(2.0), Inches(1.5)
        )

    template_bytes = _template_bytes(_add_picture_slide)
    slides = [SlideSpec(kind="title", title="Reused", template_slide_index=0)]
    prs = _open(_build(slides, template_bytes=template_bytes))

    slide = prs.slides[0]
    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert cast(Any, pictures[0]).image.blob  # the embedded image part is readable


def test_template_source_slides_do_not_leak_into_output():
    template_bytes = _template_bytes(
        lambda prs: (_add_designed_slide(prs), _add_designed_slide(prs))
    )
    slides = [
        SlideSpec(kind="title", title="Cover"),
        SlideSpec(kind="bullets", title="From template", template_slide_index=1),
    ]
    prs = _open(_build(slides, template_bytes=template_bytes))

    # Only the two generated slides — the template's own two source slides
    # (which "Original title"/"Original body" came from) are stripped.
    assert len(prs.slides) == 2
    first_title = prs.slides[0].shapes.title
    second_title = prs.slides[1].shapes.title
    assert first_title is not None and first_title.text_frame.text == "Cover"
    assert second_title is not None and second_title.text_frame.text == "From template"


def test_template_sections_do_not_leak_into_output():
    def _add_slide_in_a_section(prs) -> None:
        _add_designed_slide(prs)
        slide_id = prs.slides[-1].slide_id
        prs.part._element.append(
            parse_xml(
                '<p:extLst xmlns:p="http://schemas.openxmlformats.org/'
                'presentationml/2006/main"><p:ext uri="{521415D9-36F7-43E2-'
                'AB2F-B90AF26B5E84}"><p14:sectionLst xmlns:p14="http://'
                'schemas.microsoft.com/office/powerpoint/2010/main">'
                '<p14:section name="Covers" id="{00000000-0000-0000-0000-'
                f'000000000001}}"><p14:sldIdLst><p14:sldId id="{slide_id}"/>'
                "</p14:sldIdLst></p14:section></p14:sectionLst></p:ext>"
                "</p:extLst>"
            )
        )

    template_bytes = _template_bytes(_add_slide_in_a_section)
    slides = [SlideSpec(kind="bullets", title="T", template_slide_index=0)]
    deck = _open(_build(slides, template_bytes=template_bytes))

    assert not deck.part._element.xpath(".//*[local-name()='sectionLst']")


def test_template_slide_index_out_of_range_raises():
    template_bytes = _template_bytes(_add_designed_slide)
    slides = [SlideSpec(kind="title", title="X", template_slide_index=5)]
    with pytest.raises(ValueError, match="does not exist"):
        _build(slides, template_bytes=template_bytes)


def test_template_slide_index_without_template_raises():
    slides = [SlideSpec(kind="title", title="X", template_slide_index=0)]
    with pytest.raises(ValueError, match="no brand template is configured"):
        _build(slides)


def test_template_slide_index_rejects_chart_shape():
    def _add_chart_slide(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        chart_data = ChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("S", (1, 2))
        slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, Inches(1), Inches(1), Inches(3), Inches(2), chart_data
        )

    template_bytes = _template_bytes(_add_chart_slide)
    slides = [SlideSpec(kind="title", title="X", template_slide_index=0)]
    with pytest.raises(ValueError, match="isn't supported for duplication"):
        _build(slides, template_bytes=template_bytes)


def test_template_slide_index_rejects_chart_nested_inside_a_group():
    # A chart's own `.shape_type` is GROUP once it's wrapped in a group
    # shape — the unsupported-shape check must still catch it, since
    # `_rewire_blip_relationship` only knows how to re-wire a picture's
    # blip, not a chart's embedded-workbook relationship.
    def _add_grouped_chart_slide(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        rectangle = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1)
        )
        group = slide.shapes.add_group_shape([rectangle])
        chart_data = ChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("S", (1, 2))
        group.shapes.add_chart(
            XL_CHART_TYPE.PIE, Inches(0), Inches(0), Inches(1), Inches(1), chart_data
        )

    template_bytes = _template_bytes(_add_grouped_chart_slide)
    slides = [SlideSpec(kind="title", title="X", template_slide_index=0)]
    with pytest.raises(ValueError, match="isn't supported for duplication"):
        _build(slides, template_bytes=template_bytes)


def test_template_slide_index_skips_preset_background_but_manual_theme_still_wins():
    template_bytes = _template_bytes(_add_designed_slide)
    preset_name = next(iter(THEME_PRESETS))
    theme = ThemeSpec(preset=preset_name)
    slides = [
        SlideSpec(kind="bullets", title="Reused", template_slide_index=0),
        # kind="bullets" (not "title") so the preset's "soft" treatment
        # applies — always a solid fill, unlike "strong" which some presets
        # render as a pattern picture instead of `slide.background.fill`.
        SlideSpec(kind="bullets", title="Generic"),
    ]
    deck_bytes = deck_builder.build(
        "Deck",
        slides,
        template_bytes=template_bytes,
        image_bytes_by_ref={},
        theme=theme,
    )
    prs = _open(deck_bytes)

    # The template-sourced slide keeps its OWN background, not the preset's.
    assert prs.slides[0].background.fill.fore_color.rgb == RGBColor.from_string(
        "445566"
    )
    # A generic (non-template) slide still gets the preset's soft treatment.
    preset = THEME_PRESETS[preset_name]
    assert prs.slides[1].background.fill.fore_color.rgb == RGBColor.from_string(
        preset.soft_background_color.lstrip("#")
    )


def test_template_slide_index_notes_copied_from_source():
    def _add_slide_with_notes(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        notes_text_frame = slide.notes_slide.notes_text_frame
        assert notes_text_frame is not None
        notes_text_frame.text = "Template author's notes"

    template_bytes = _template_bytes(_add_slide_with_notes)
    slides = [SlideSpec(kind="title", title="X", template_slide_index=0)]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]
    notes_text_frame = slide.notes_slide.notes_text_frame
    assert notes_text_frame is not None
    assert notes_text_frame.text == "Template author's notes"


# ---------------------------------------------------------------------------
# section_header + template_slide_index — a real customer template gives the
# numbered "chapter badge" shape type="title" and the real chapter-name
# shape a generic type="body", backwards from the rest of the template (RFC
# §13 amendment, 2026-09-22).
# ---------------------------------------------------------------------------


def _add_chapter_divider_slide(prs) -> None:
    """A section_header-style slide whose OOXML placeholder TYPES are
    inverted relative to their visual role: `type="title"` is the small
    numbered badge shape, `type="body"` (a generic content placeholder) is
    the real chapter title — the exact structure confirmed in a real
    customer template's XML."""

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    assert title_shape is not None
    title_shape.text_frame.text = "01"
    cast(
        Any, slide.placeholders[1]
    ).text_frame.text = "TITRE DU CHAPITRE\nSUR 1 OU 3 LIGNES\nMAXIMUM"


def _add_normal_section_header_slide(prs) -> None:
    """A NON-inverted section_header slide (type="title" really is the
    title) — used to prove the fix doesn't regress an ordinary template."""

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    assert title_shape is not None
    title_shape.text_frame.text = "Section title default"
    cast(Any, slide.placeholders[1]).text_frame.text = "Subtitle default"


def test_section_header_template_slide_fills_badge_and_title_by_content_not_type():
    template_bytes = _template_bytes(_add_chapter_divider_slide)
    slides = [
        SlideSpec(kind="section_header", title=name, template_slide_index=0)
        for name in ["Pourquoi l'IA générative ?", "Cadrage", "Déploiement"]
    ]
    prs = _open(_build(slides, template_bytes=template_bytes))

    assert len(prs.slides) == 3
    for position, (slide, spec) in enumerate(zip(prs.slides, slides), start=1):
        title_shape = slide.shapes.title
        assert title_shape is not None
        # The badge zone (OOXML type="title") gets ONLY the two-digit
        # number — never the chapter name.
        assert title_shape.text_frame.text == f"{position:02d}"
        body = cast(
            Any,
            next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0),
        )
        # The real title zone (OOXML type="body") gets the chapter name, and
        # none of the template's default text survives.
        assert body.text_frame.text == spec.title
        assert "TITRE DU CHAPITRE" not in body.text_frame.text


def test_section_header_template_slide_badge_is_single_run_no_stacking():
    template_bytes = _template_bytes(_add_chapter_divider_slide)
    slides = [
        SlideSpec(kind="section_header", title="Chapter One", template_slide_index=0)
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]
    title_shape = slide.shapes.title
    assert title_shape is not None
    paragraphs = title_shape.text_frame.paragraphs
    assert len(paragraphs) == 1
    assert len(paragraphs[0].runs) == 1
    assert paragraphs[0].runs[0].text == "01"


def test_section_header_template_slide_without_badge_keeps_type_based_fill():
    template_bytes = _template_bytes(_add_normal_section_header_slide)
    slides = [
        SlideSpec(
            kind="section_header",
            title="Real Title",
            subtitle="Real Subtitle",
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    title_shape = slide.shapes.title
    assert title_shape is not None
    assert title_shape.text_frame.text == "Real Title"
    body = cast(
        Any, next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    )
    assert body.text_frame.text == "Real Subtitle"


def test_section_header_template_slide_missing_title_raises():
    template_bytes = _template_bytes(_add_chapter_divider_slide)
    slides = [SlideSpec(kind="section_header", template_slide_index=0)]
    with pytest.raises(ValueError, match="needs a 'title'"):
        _build(slides, template_bytes=template_bytes)


def test_section_header_badge_numbering_only_counts_badge_slides():
    template_bytes = _template_bytes(
        lambda prs: (_add_chapter_divider_slide(prs), _add_designed_slide(prs))
    )
    slides = [
        SlideSpec(kind="title", title="Cover"),
        SlideSpec(kind="section_header", title="First", template_slide_index=0),
        SlideSpec(kind="bullets", title="Not a divider", template_slide_index=1),
        SlideSpec(kind="section_header", title="Second", template_slide_index=0),
    ]
    prs = _open(_build(slides, template_bytes=template_bytes))

    first_divider = prs.slides[1].shapes.title
    second_divider = prs.slides[3].shapes.title
    assert first_divider is not None and first_divider.text_frame.text == "01"
    assert second_divider is not None and second_divider.text_frame.text == "02"


# ---------------------------------------------------------------------------
# Template furniture (date/footer/slide number) and leftover sample text —
# reproduced from a real brand template whose dividers and covers carry
# furniture placeholders next to their content ones.
# ---------------------------------------------------------------------------


def _add_layout_furniture(slide, *, footer_text: str = "PRESENTATION TITLE") -> None:
    """Copy the layout's date/footer/slide-number placeholders onto `slide`
    (python-pptx's add_slide never does), the way a real template ships them."""

    for layout_placeholder in slide.slide_layout.placeholders:
        if layout_placeholder.placeholder_format.type in (
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.SLIDE_NUMBER,
        ):
            slide.shapes._spTree.append(copy.deepcopy(layout_placeholder._element))
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
            placeholder.text_frame.text = footer_text


def _add_divider_with_furniture(prs) -> None:
    _add_chapter_divider_slide(prs)
    _add_layout_furniture(prs.slides[-1])


def _add_cover_with_furniture(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_shape = slide.shapes.title
    assert title_shape is not None
    title_shape.text_frame.text = "Cover title default"
    cast(Any, slide.placeholders[1]).text_frame.text = "Subtitle default"
    _add_layout_furniture(slide)


def _placeholders_of_type(slide, placeholder_type) -> list:
    return [
        ph
        for ph in slide.placeholders
        if ph.placeholder_format.type == placeholder_type
    ]


def test_section_header_badge_numbering_ignores_furniture_placeholders():
    template_bytes = _template_bytes(_add_divider_with_furniture)
    slides = [
        SlideSpec(kind="section_header", title=name, template_slide_index=0)
        for name in ["One", "Two", "Three"]
    ]
    prs = _open(_build(slides, template_bytes=template_bytes))

    badges = [cast(Any, slide.shapes.title).text_frame.text for slide in prs.slides]
    assert badges == ["01", "02", "03"]


def test_template_footer_gets_deck_title_and_is_never_a_content_slot():
    template_bytes = _template_bytes(_add_cover_with_furniture)
    slides = [
        SlideSpec(
            kind="title",
            title="Cover",
            subtitle="Real subtitle",
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    (footer,) = _placeholders_of_type(slide, PP_PLACEHOLDER.FOOTER)
    assert footer.text_frame.text == "Deck"
    (subtitle,) = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 1]
    assert cast(Any, subtitle).text_frame.text == "Real subtitle"
    (date,) = _placeholders_of_type(slide, PP_PLACEHOLDER.DATE)
    assert date.text_frame.text != "Real subtitle"


def test_two_content_chart_is_never_placed_in_the_footer():
    def _add_body_slide_with_furniture(prs) -> None:
        _add_designed_slide(prs)
        _add_layout_furniture(prs.slides[-1])

    template_bytes = _template_bytes(_add_body_slide_with_furniture)
    chart = ChartSpec(
        chart_type="bar",
        categories=["A", "B"],
        series=[ChartSeries(name="S", values=[1, 2])],
    )
    slides = [
        SlideSpec(
            kind="two_content",
            title="T",
            bullets=[BulletItem(text="b")],
            chart=chart,
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    (footer,) = _placeholders_of_type(slide, PP_PLACEHOLDER.FOOTER)
    (chart_frame,) = [shape for shape in slide.shapes if shape.has_chart]
    assert chart_frame.height > Inches(2)
    assert chart_frame.top + chart_frame.height <= footer.top


def test_unfilled_template_sample_text_is_removed():
    template_bytes = _template_bytes(_add_cover_with_furniture)
    slides = [SlideSpec(kind="title", title="Cover only", template_slide_index=0)]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    texts = [cast(Any, ph).text_frame.text for ph in slide.placeholders]
    assert "Subtitle default" not in texts
    assert "Cover only" in texts


def test_template_image_layout_without_image_ref_raises():
    def _add_picture_placeholder_slide(prs) -> None:
        prs.slides.add_slide(prs.slide_layouts[8])  # "Picture with Caption"

    template_bytes = _template_bytes(_add_picture_placeholder_slide)
    slides = [
        SlideSpec(
            kind="bullets",
            title="T",
            bullets=[BulletItem(text="b")],
            template_slide_index=0,
        )
    ]
    with pytest.raises(ValueError, match="image layout"):
        _build(slides, template_bytes=template_bytes)


def test_two_content_without_right_content_fills_bullets_placeholder():
    template_bytes = _template_bytes(_add_designed_slide)
    slides = [
        SlideSpec(
            kind="two_content",
            title="T",
            bullets=[BulletItem(text="Only bullets")],
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    body = cast(
        Any, next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    )
    assert body.text_frame.text == "Only bullets"


def test_template_slide_ignores_deck_wide_background_and_text_color():
    template_bytes = _template_bytes(_add_designed_slide)
    theme = ThemeSpec(
        primary_color="#AA0000",
        background=BackgroundSpec(kind="solid", color="#00AA00"),
    )
    slides = [SlideSpec(kind="bullets", title="Reused", template_slide_index=0)]
    deck_bytes = deck_builder.build(
        "Deck",
        slides,
        template_bytes=template_bytes,
        image_bytes_by_ref={},
        theme=theme,
    )
    slide = _open(deck_bytes).slides[0]

    assert slide.background.fill.fore_color.rgb == RGBColor.from_string("445566")
    title_run = cast(Any, slide.shapes.title).text_frame.paragraphs[0].runs[0]
    assert title_run.font.color.type is None


def test_footer_fill_handles_empty_first_paragraph_and_fields():
    from pptx.oxml import parse_xml

    def _add_tricky_footer_slide(prs) -> None:
        _add_designed_slide(prs)
        slide = prs.slides[-1]
        _add_layout_furniture(slide)
        (footer,) = _placeholders_of_type(slide, PP_PLACEHOLDER.FOOTER)
        tx_body = footer._element.txBody
        for paragraph in list(tx_body.iterchildren(qn_("a:p"))):
            tx_body.remove(paragraph)
        ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        tx_body.append(parse_xml(f"<a:p {ns}/>"))
        tx_body.append(
            parse_xml(
                f'<a:p {ns}><a:r><a:rPr b="1"/><a:t>Page </a:t></a:r>'
                '<a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
                "<a:t>1</a:t></a:fld></a:p>"
            )
        )

    template_bytes = _template_bytes(_add_tricky_footer_slide)
    slides = [SlideSpec(kind="bullets", title="T", template_slide_index=0)]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    (footer,) = _placeholders_of_type(slide, PP_PLACEHOLDER.FOOTER)
    assert footer.text_frame.text == "Deck"
    assert footer.text_frame.paragraphs[0].runs[0].font.bold is True


def _add_round_picture_frame_slide(prs) -> None:
    """A "Picture with Caption" slide whose frame the author resized to a
    square and turned into a circle on the slide itself."""

    slide = prs.slides.add_slide(prs.slide_layouts[8])
    (frame,) = _placeholders_of_type(slide, PP_PLACEHOLDER.PICTURE)
    frame.left, frame.top = Inches(1), Inches(1)
    frame.width, frame.height = Inches(3), Inches(3)
    frame._element.spPr.get_or_add_prstGeom().set("prst", "ellipse")


def test_title_kind_places_image_ref_in_the_template_picture_frame():
    template_bytes = _template_bytes(_add_round_picture_frame_slide)
    slides = [
        SlideSpec(kind="title", title="T", image_ref="doc-1", template_slide_index=0)
    ]
    slide = _open(
        _build(
            slides,
            template_bytes=template_bytes,
            image_bytes_by_ref={"doc-1": _png_bytes()},
        )
    ).slides[0]

    assert [
        s
        for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and s._element.tag == qn_("p:pic")
    ]


def test_subtitle_on_template_slide_without_a_subtitle_zone_raises():
    def _add_title_only_slide(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # "Title Only"
        cast(Any, slide.shapes.title).text_frame.text = "Sample"

    template_bytes = _template_bytes(_add_title_only_slide)
    slides = [SlideSpec(kind="title", title="T", subtitle="S", template_slide_index=0)]
    with pytest.raises(ValueError, match="no subtitle zone"):
        _build(slides, template_bytes=template_bytes)


def test_multi_line_subtitle_follows_the_template_sample_levels():
    # Lead zone sample: a level-1 lead line above a level-0 paragraph.
    def _add_lead_zone_slide(prs) -> None:
        _add_designed_slide(prs)
        body = next(
            ph for ph in prs.slides[-1].placeholders if ph.placeholder_format.idx != 0
        )
        body.text_frame.text = "Sample lead"
        body.text_frame.paragraphs[0].level = 1
        body.text_frame.add_paragraph().text = "Sample paragraph"

    template_bytes = _template_bytes(_add_lead_zone_slide)
    slides = [
        SlideSpec(
            kind="title", title="T", subtitle="Lead\nDetail", template_slide_index=0
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    body = cast(
        Any, next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    )
    assert [(p.text, p.level) for p in body.text_frame.paragraphs] == [
        ("Lead", 1),
        ("Detail", 0),
    ]


@pytest.mark.parametrize("kind", ["image", "two_content"])
def test_image_fills_template_picture_frame_with_its_shape(kind):
    template_bytes = _template_bytes(_add_round_picture_frame_slide)
    slides = [
        SlideSpec(
            kind=kind,
            title="T",
            bullets=[BulletItem(text="b")],
            image_ref="doc-1",
            template_slide_index=0,
        )
    ]
    slide = _open(
        _build(
            slides,
            template_bytes=template_bytes,
            image_bytes_by_ref={"doc-1": _png_bytes(400, 200)},
        )
    ).slides[0]

    (picture,) = [ph for ph in slide.placeholders if hasattr(ph, "image")]
    picture = cast(Any, picture)
    assert (picture.left, picture.top) == (Inches(1), Inches(1))
    assert (picture.width, picture.height) == (Inches(3), Inches(3))
    # 2:1 image in a square frame: cropped to fill, not letterboxed.
    assert picture.crop_left == pytest.approx(0.25)
    assert picture.crop_right == pytest.approx(0.25)
    geometry = picture._element.spPr.find(qn_("a:prstGeom"))
    assert geometry is not None and geometry.get("prst") == "ellipse"


def _add_text_and_picture_slide(prs) -> None:
    """Small "chapeau" zone + large bullets zone next to a picture frame."""

    _add_round_picture_frame_slide(prs)
    slide = prs.slides[-1]
    (body,) = _placeholders_of_type(slide, PP_PLACEHOLDER.BODY)
    body.width, body.height = Inches(4), Inches(4)
    chapeau = copy.deepcopy(body._element)
    chapeau.nvSpPr.nvPr.ph.set("idx", "13")
    chapeau.nvSpPr.cNvPr.set("id", "99")
    body._element.addnext(chapeau)
    small = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == 13)
    small.width, small.height = Inches(4), Inches(0.5)
    small.text_frame.text = "Chapeau sample"


def test_two_content_fills_template_text_column_in_place():
    # Bullets must go into the large zone itself, not a new box sized like
    # the chapeau (which shrink-to-fit rendered unreadably small).
    template_bytes = _template_bytes(_add_text_and_picture_slide)
    slides = [
        SlideSpec(
            kind="two_content",
            title="T",
            bullets=[BulletItem(text="Point one"), BulletItem(text="Point two")],
            image_ref="doc-1",
            template_slide_index=0,
        )
    ]
    slide = _open(
        _build(
            slides,
            template_bytes=template_bytes,
            image_bytes_by_ref={"doc-1": _png_bytes()},
        )
    ).slides[0]

    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    (column,) = [
        cast(Any, ph)
        for ph in slide.placeholders
        if ph.has_text_frame and ph.placeholder_format.idx != 0
    ]
    assert column.text_frame.text == "Point one\nPoint two"
    assert column.height == Inches(4)


def test_two_content_subtitle_fills_the_template_chapeau():
    template_bytes = _template_bytes(_add_text_and_picture_slide)
    slides = [
        SlideSpec(
            kind="two_content",
            title="T",
            subtitle="Lead",
            bullets=[BulletItem(text="Point one")],
            image_ref="doc-1",
            template_slide_index=0,
        )
    ]
    slide = _open(
        _build(
            slides,
            template_bytes=template_bytes,
            image_bytes_by_ref={"doc-1": _png_bytes()},
        )
    ).slides[0]

    texts = {
        ph.placeholder_format.idx: cast(Any, ph).text_frame.text
        for ph in slide.placeholders
        if ph.has_text_frame and ph.placeholder_format.idx != 0
    }
    assert texts.pop(13) == "Lead"
    assert list(texts.values()) == ["Point one"]


def test_bullets_start_at_the_template_sample_text_level():
    # Template whose bulleted zone holds its sample text at level 2 (levels
    # 0/1 are plain paragraph / sub-heading styles): level 0 bullets land there.
    def _add_level_two_sample_slide(prs) -> None:
        _add_designed_slide(prs)
        body = next(
            ph for ph in prs.slides[-1].placeholders if ph.placeholder_format.idx != 0
        )
        body.text_frame.text = "Sample bullet"
        body.text_frame.paragraphs[0].level = 2

    template_bytes = _template_bytes(_add_level_two_sample_slide)
    slides = [
        SlideSpec(
            kind="bullets",
            title="T",
            bullets=[BulletItem(text="Top"), BulletItem(text="Child", level=1)],
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    body = cast(
        Any, next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    )
    assert [p.level for p in body.text_frame.paragraphs] == [2, 3]


def test_bullets_keep_the_template_sample_formatting_of_their_level():
    # Sample paragraphs carry their own bullet/bold overrides (not inherited
    # from the layout): a heading level must stay unbulleted and bold.
    def _add_styled_sample_slide(prs) -> None:
        _add_designed_slide(prs)
        body = next(
            ph for ph in prs.slides[-1].placeholders if ph.placeholder_format.idx != 0
        )
        heading = body.text_frame.paragraphs[0]
        heading.text = "Sample heading"
        heading._p.get_or_add_pPr().append(heading._p.makeelement(qn_("a:buNone")))
        heading.runs[0].font.bold = True
        item = body.text_frame.add_paragraph()
        item.text = "Sample item"
        item.level = 1

    template_bytes = _template_bytes(_add_styled_sample_slide)
    slides = [
        SlideSpec(
            kind="bullets",
            title="T",
            bullets=[BulletItem(text="Heading"), BulletItem(text="Item", level=1)],
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    body = cast(
        Any, next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    )
    heading, item = body.text_frame.paragraphs
    assert heading._p.pPr.find(qn_("a:buNone")) is not None
    assert heading.runs[0].font.bold is True
    assert item.level == 1
    assert item._p.pPr.find(qn_("a:buNone")) is None
    assert item.runs[0].font.bold is not True


def test_bullets_fill_largest_zone_and_subtitle_fills_the_lead_zone():
    template_bytes = _template_bytes(_add_comparison_slide)
    slides = [
        SlideSpec(
            kind="bullets",
            title="T",
            subtitle="Lead",
            bullets=[BulletItem(text="b1")],
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    # The two content zones differ in size only by rounding; either may win.
    texts = _column_texts(slide)
    assert texts.pop(1) == "Lead"
    assert list(texts.values()) == ["b1"]


def _add_comparison_slide(prs) -> None:
    """Two columns, each a heading zone above a text zone, with sample text."""

    slide = prs.slides.add_slide(prs.slide_layouts[4])  # "Comparison"
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx != 0:
            cast(Any, placeholder).text_frame.text = "Sample"


def _column_texts(slide) -> dict[int, str]:
    return {
        ph.placeholder_format.idx: cast(Any, ph).text_frame.text
        for ph in slide.placeholders
        if ph.placeholder_format.idx != 0 and ph.has_text_frame
    }


def test_columns_fill_template_column_zones_left_to_right():
    template_bytes = _template_bytes(_add_comparison_slide)
    slides = [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[
                ColumnSpec(heading="Left", bullets=[BulletItem(text="l1")]),
                ColumnSpec(heading="Right", bullets=[BulletItem(text="r1")]),
            ],
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert _column_texts(slide) == {1: "Left", 2: "l1", 3: "Right", 4: "r1"}


def _add_comparison_slide_with_group_label(prs) -> None:
    """Comparison slide plus a small label zone above the left column."""

    _add_comparison_slide(prs)
    slide = prs.slides[-1]
    header = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == 1)
    # Pin inherited geometry so the copy (with no layout twin) keeps it.
    header.left, header.top = header.left, header.top
    header.width, header.height = header.width, header.height
    label = copy.deepcopy(header._element)
    label.nvSpPr.nvPr.ph.set("idx", "30")
    label.nvSpPr.cNvPr.set("id", "99")
    header._element.addnext(label)
    zone = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == 30)
    zone.top, zone.height = header.top - Inches(0.6), Inches(0.4)


def _two_columns(**fields) -> list[SlideSpec]:
    return [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[
                ColumnSpec(heading="Left", bullets=[BulletItem(text="l1")]),
                ColumnSpec(heading="Right", bullets=[BulletItem(text="r1")]),
            ],
            template_slide_index=0,
            **fields,
        )
    ]


def test_columns_group_label_fills_the_template_label_zone():
    template_bytes = _template_bytes(_add_comparison_slide_with_group_label)
    slides = _two_columns(group_label="Key levers")
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    assert _column_texts(slide) == {
        1: "Left",
        2: "l1",
        3: "Right",
        4: "r1",
        30: "Key levers",
    }


def test_columns_group_label_without_a_label_zone_is_dropped_with_a_warning():
    template_bytes = _template_bytes(_add_comparison_slide)
    warnings: list[str] = []
    deck = _build(
        _two_columns(group_label="X"), template_bytes=template_bytes, warnings=warnings
    )

    assert "X" not in _column_texts(_open(deck).slides[0]).values()
    assert len(warnings) == 1 and "'group_label' ignored" in warnings[0]


def test_columns_image_ref_without_a_main_picture_frame_is_dropped_with_a_warning():
    template_bytes = _template_bytes(_add_comparison_slide)
    warnings: list[str] = []
    deck = _build(
        _two_columns(image_ref="doc-1"),
        template_bytes=template_bytes,
        image_bytes_by_ref={"doc-1": _png_bytes()},
        warnings=warnings,
    )

    slide = _open(deck).slides[0]
    assert _column_texts(slide) == {1: "Left", 2: "l1", 3: "Right", 4: "r1"}
    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(warnings) == 1
    assert warnings[0].startswith("slide 1:") and "'image_ref' ignored" in warnings[0]


def test_caption_columns_reject_a_heading_with_bullets():
    # "Two Content": one single-level text zone per column, i.e. captions.
    def _add_caption_columns_slide(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[3])
        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.idx != 0:
                cast(Any, placeholder).text_frame.text = "Caption sample"

    template_bytes = _template_bytes(_add_caption_columns_slide)
    with pytest.raises(ValueError, match="one-sentence caption columns"):
        _build(_two_columns(), template_bytes=template_bytes)

    captions = [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[
                ColumnSpec(bullets=[BulletItem(text="Left")]),
                ColumnSpec(bullets=[BulletItem(text="Right")]),
            ],
            template_slide_index=0,
        )
    ]
    slide = _open(_build(captions, template_bytes=template_bytes)).slides[0]
    assert sorted(_column_texts(slide).values()) == ["Left", "Right"]


def test_columns_drop_unused_template_columns():
    template_bytes = _template_bytes(_add_comparison_slide)
    slides = [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[ColumnSpec(heading="Only", bullets=[BulletItem(text="x")])],
            template_slide_index=0,
        )
    ]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    assert _column_texts(slide) == {1: "Only", 2: "x"}


def test_columns_reject_more_columns_than_the_template_has():
    template_bytes = _template_bytes(_add_comparison_slide)
    slides = [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[ColumnSpec(heading=str(n)) for n in range(3)],
            template_slide_index=0,
        )
    ]
    with pytest.raises(ValueError, match="has 2 columns but 3 were given"):
        _build(slides, template_bytes=template_bytes)


def test_columns_drop_icon_ref_without_icon_frames():
    template_bytes = _template_bytes(_add_comparison_slide)
    slides = [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[ColumnSpec(heading="A", icon_ref="doc-1"), ColumnSpec()],
            template_slide_index=0,
        )
    ]
    warnings: list[str] = []
    deck = _build(
        slides,
        template_bytes=template_bytes,
        image_bytes_by_ref={"doc-1": _png_bytes()},
        warnings=warnings,
    )

    slide = _open(deck).slides[0]
    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(warnings) == 1 and "'icon_ref' ignored" in warnings[0]


def test_columns_without_template_split_the_content_area():
    slides = [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[
                ColumnSpec(heading=f"H{n}", bullets=[BulletItem(text=f"b{n}")])
                for n in range(3)
            ],
        )
    ]
    slide = _open(_build(slides)).slides[0]

    boxes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert [cast(Any, b).text_frame.text for b in boxes] == [
        "H0\nb0",
        "H1\nb1",
        "H2\nb2",
    ]
    lefts = [b.left for b in boxes]
    assert lefts == sorted(lefts) and len(set(lefts)) == 3


def test_columns_without_template_drop_image_ref_with_a_warning():
    slides = [
        SlideSpec(
            kind="columns",
            title="T",
            columns=[ColumnSpec(heading="A"), ColumnSpec(heading="B")],
            image_ref="doc-1",
        )
    ]
    warnings: list[str] = []
    deck = _build(slides, image_bytes_by_ref={"doc-1": _png_bytes()}, warnings=warnings)

    slide = _open(deck).slides[0]
    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(warnings) == 1 and "'image_ref', ignored" in warnings[0]


def test_columns_requires_columns():
    with pytest.raises(ValueError, match="requires 'columns'"):
        _build([SlideSpec(kind="columns", title="T")])


def test_empty_outlined_template_frame_is_kept():
    # An empty text frame drawn only as a border groups the columns visually;
    # the sample-text cleanup must not remove it.
    def _add_outlined_frame_slide(prs) -> None:
        _add_designed_slide(prs)
        body = next(
            ph for ph in prs.slides[-1].placeholders if ph.placeholder_format.idx != 0
        )
        body.text_frame.text = ""
        body.line.color.rgb = RGBColor.from_string("0000FF")

    template_bytes = _template_bytes(_add_outlined_frame_slide)
    slides = [SlideSpec(kind="title", title="T", template_slide_index=0)]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    assert any(ph.placeholder_format.idx != 0 for ph in slide.placeholders)


def test_designer_placed_picture_is_kept_and_needs_no_image_ref():
    def _add_filled_picture_slide(prs) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[8])  # "Picture with Caption"
        picture_frame = next(
            ph
            for ph in slide.placeholders
            if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        )
        cast(Any, picture_frame).insert_picture(io.BytesIO(_png_bytes()))

    template_bytes = _template_bytes(_add_filled_picture_slide)
    slides = [SlideSpec(kind="title", title="T", template_slide_index=0)]
    slide = _open(_build(slides, template_bytes=template_bytes)).slides[0]

    assert [
        s
        for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
        and s.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ]
