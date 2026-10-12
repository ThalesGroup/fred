# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""`describe_template_slides` must produce an accurate, stable catalog for the
`generate_pptx` tool description, or the agent guesses `template_slide_index`."""

from __future__ import annotations

import io
from typing import Any, cast

from fred_capability_pptx_generator.template_inventory import describe_template_slides
from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches


def _set_title(slide, text: str) -> None:
    title_shape = slide.shapes.title
    assert title_shape is not None
    title_shape.text_frame.text = text


def test_no_slides_returns_none():
    assert describe_template_slides(Presentation()) is None


def test_untitled_slide_is_labeled():
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # Blank: no title placeholder
    assert describe_template_slides(prs) == "  0: (untitled)"


def test_titled_slide_with_body_text_is_tagged_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    _set_title(slide, "Agenda")
    cast(Any, slide.placeholders[1]).text_frame.text = "Item one"
    assert describe_template_slides(prs) == '  0: "Agenda" [text]'


def test_long_title_is_truncated():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _set_title(slide, "X" * 100)
    catalog = describe_template_slides(prs)
    assert catalog is not None
    assert catalog.endswith('…"')
    assert len(catalog) < 100


def test_picture_placeholder_is_tagged_image():
    # Layout 8 ("Picture with Caption") has a genuine PICTURE-type
    # placeholder — `deck_builder._body_placeholder` can find and replace
    # this one, unlike a plain inserted picture (see the test below).
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])
    _set_title(slide, "Cover")
    picture_placeholder = next(
        ph
        for ph in slide.placeholders
        if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    )
    buffer = io.BytesIO()
    Image.new("RGB", (10, 10), (1, 2, 3)).save(buffer, format="PNG")
    buffer.seek(0)
    cast(Any, picture_placeholder).insert_picture(buffer)
    assert describe_template_slides(prs) == '  0: "Cover" [image]'


def test_non_placeholder_picture_is_not_tagged():
    # A plain inserted picture is not a placeholder, so deck_builder never
    # replaces it with `image_ref`; tagging it "image" would mislead the agent.
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _set_title(slide, "Cover")
    buffer = io.BytesIO()
    Image.new("RGB", (10, 10), (1, 2, 3)).save(buffer, format="PNG")
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Inches(1), Inches(1), Inches(2), Inches(2))
    assert describe_template_slides(prs) == '  0: "Cover"'


def test_non_placeholder_table_is_not_tagged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(3), Inches(2))
    assert describe_template_slides(prs) == "  0: (untitled)"


def test_non_placeholder_chart_is_not_tagged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = ChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("S", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(1), Inches(1), Inches(3), Inches(2), chart_data
    )
    assert describe_template_slides(prs) == "  0: (untitled)"


def test_multiple_slides_are_numbered_in_order():
    prs = Presentation()
    for title in ("First", "Second", "Third"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        _set_title(slide, title)
    catalog = describe_template_slides(prs)
    assert catalog == '  0: "First"\n  1: "Second"\n  2: "Third"'


def test_catalog_is_capped_for_a_very_large_template():
    prs = Presentation()
    for i in range(55):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        _set_title(slide, f"Slide {i}")
    catalog = describe_template_slides(prs)
    assert catalog is not None
    lines = catalog.split("\n")
    assert len(lines) == 51  # 50 listed slides + one "N more" summary line
    assert lines[-1] == "  ... and 5 more slide(s)"


def test_multi_level_zone_lists_each_level_look():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    _set_title(slide, "Agenda")
    text_frame = cast(Any, slide.placeholders[1]).text_frame
    text_frame.text = "Lead paragraph"
    text_frame.paragraphs[0].runs[0].font.bold = True
    item = text_frame.add_paragraph()
    item.text = "An item"
    item.level = 1

    catalog = describe_template_slides(prs)

    assert catalog == (
        '  0: "Agenda" [text]\n'
        '      text zone 1: level 0 = bullet, bold ("Lead paragraph"); '
        'level 1 = bullet ("An item")'
    )


def test_several_text_zones_are_listed_top_to_bottom():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[4])  # "Comparison"
    _set_title(slide, "Compare")
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx != 0:
            cast(
                Any, placeholder
            ).text_frame.text = f"Zone {placeholder.placeholder_format.idx}"

    catalog = describe_template_slides(prs)

    assert catalog is not None
    zone_lines = catalog.splitlines()[1:]
    assert len(zone_lines) == 4
    assert zone_lines[0].startswith("      text zone 1: level 0 =")
