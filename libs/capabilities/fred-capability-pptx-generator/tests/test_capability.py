# Copyright Thales 2026
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""`PptxGeneratorCapability` tests: registry boot invariant, `validate_config`
save-time state machine, and the `generate_pptx` tool's happy/error paths
against a stubbed `RuntimeServices` (RFC §4, `docs/swift/capabilities/AUTHORING.md`
testing expectations)."""

from __future__ import annotations

import io
from importlib.metadata import EntryPoint
from typing import Any, cast

import pytest
from fred_capability_pptx_generator import capability as capability_mod
from fred_capability_pptx_generator import deck_builder
from fred_capability_pptx_generator.capability import (
    BRAND_TEMPLATE_KEY,
    BRAND_TEMPLATE_SLOT,
    MAX_TEMPLATE_UPLOAD_BYTES,
    PptxGeneratorCapability,
    PptxGeneratorConfig,
)
from fred_capability_pptx_generator.models import (
    BackgroundSpec,
    ColumnSpec,
    SlideSpec,
    TableSpec,
    ThemeSpec,
)
from fred_capability_pptx_generator.parts import PptxGeneratedPart
from fred_runtime.capabilities.registry import (
    FRED_CAPABILITIES_ENTRY_POINT_GROUP,
    CapabilityRegistry,
)
from fred_sdk.contracts.capability import (
    CapabilityContext,
    CapabilityIdentity,
    EmptyModel,
    SaveContext,
    UploadedFile,
)
from fred_sdk.contracts.context import LinkKind, PublishedArtifact
from fred_sdk.contracts.runtime import (
    AgentAssetPort,
    DocumentContentPort,
    DocumentRawContent,
    RuntimeServices,
    WorkspaceFsPort,
)
from pptx import Presentation

_ENTRY_POINT_VALUE = "fred_capability_pptx_generator.capability:PptxGeneratorCapability"


def _blank_pptx_bytes() -> bytes:
    buffer = io.BytesIO()
    Presentation().save(buffer)
    return buffer.getvalue()


def _identity(session_id: str | None = "s-1") -> CapabilityIdentity:
    return CapabilityIdentity(user_id="u-1", session_id=session_id, team_id=None)


def _ctx(
    cap: PptxGeneratorCapability,
    *,
    services: RuntimeServices,
    config: dict | None = None,
) -> CapabilityContext[PptxGeneratorConfig, EmptyModel]:
    return CapabilityContext(
        identity=_identity(),
        config=cap.ConfigModel.model_validate(config or {}),
        turn_options=EmptyModel(),
        services=services,
    )


class _FakeWorkspaceFs(WorkspaceFsPort):
    """Records every `write()` call; nothing else is exercised by this tool."""

    def __init__(self) -> None:
        self.writes: list[tuple[str, bytes, str | None, str | None]] = []

    def bind(self, binding) -> None:
        pass

    async def write(
        self, path: str, content: bytes, *, content_type=None, title=None
    ) -> PublishedArtifact:
        self.writes.append((path, content, content_type, title))
        file_name = title or path.rsplit("/", 1)[-1]
        return PublishedArtifact(
            key=path,
            file_name=file_name,
            size=len(content),
            href=f"https://fs.example/{path}",
            document_uid=None,
            mime=content_type,
            title=title,
        )


class _FakeAgentAssets(AgentAssetPort):
    def __init__(self, initial: dict[str, bytes] | None = None) -> None:
        self._store: dict[str, bytes] = dict(initial or {})
        self.stored: list[tuple[str, bytes]] = []

    async def store(self, key, content, *, content_type=None, filename=None) -> str:
        self._store[key] = content
        self.stored.append((key, content))
        return key

    async def fetch(self, key: str) -> bytes:
        if key not in self._store:
            raise KeyError(key)
        return self._store[key]

    async def delete(self, key: str) -> None:
        self._store.pop(key, None)


class _FakeDocumentContent(DocumentContentPort):
    def __init__(self, content_by_uid: dict[str, bytes]) -> None:
        self._content_by_uid = content_by_uid

    async def fetch_raw(self, document_uid: str) -> DocumentRawContent:
        if document_uid not in self._content_by_uid:
            raise RuntimeError(f"no such document: {document_uid}")
        return DocumentRawContent(content=self._content_by_uid[document_uid])


def _png_bytes() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (10, 10), (0, 0, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def no_pdf(monkeypatch):
    """PDF conversion unavailable (e.g. no LibreOffice in the test env) —
    the degraded, download-only path."""

    async def _none(_bytes: bytes, *, suffix: str):
        return None

    monkeypatch.setattr(capability_mod, "convert_office_bytes_to_pdf", _none)


def _enable_preview(monkeypatch, *, pdf_bytes: bytes = b"%PDF-1.5 fake") -> None:
    async def _pdf(_bytes: bytes, *, suffix: str):
        return pdf_bytes

    monkeypatch.setattr(capability_mod, "convert_office_bytes_to_pdf", _pdf)


# ---------------------------------------------------------------------------
# Registration + boot invariant
# ---------------------------------------------------------------------------


def test_capability_registers_via_entry_point_and_boots() -> None:
    registry = CapabilityRegistry()
    entry = EntryPoint(
        name="pptx_generator",
        value=_ENTRY_POINT_VALUE,
        group=FRED_CAPABILITIES_ENTRY_POINT_GROUP,
    )
    registered = registry.discover(entry_points=[entry])

    assert registered == ["pptx_generator"]
    assert isinstance(registry.capability("pptx_generator"), PptxGeneratorCapability)
    registry.validate(env={})


def test_manifest_uses_default_execution_models() -> None:
    # This capability implements only tools() (RFC §4.1) — the default
    # ("react", "graph") is correct as-is, unlike ppt_filler/html_artifact.
    assert PptxGeneratorCapability.manifest.execution_models == ("react", "graph")


def test_manifest_declares_asset_slot_and_chat_surface() -> None:
    manifest = PptxGeneratorCapability.manifest
    assert [slot.key for slot in manifest.assets] == [BRAND_TEMPLATE_SLOT]
    assert manifest.assets[0].min_count == 0
    assert manifest.assets[0].max_count == 1
    assert [part.__name__ for part in manifest.chat_parts] == ["PptxGeneratedPart"]
    assert [panel.widget for panel in manifest.side_panels] == ["pptx_preview_pane"]
    assert manifest.icon == "co_present"
    # No owned table, no router — v1 has neither (RFC §3).
    assert manifest.router is None
    assert manifest.tables == []


def test_tools_returns_single_generate_pptx_tool() -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices())
    tools = cap.tools(ctx)
    assert [tool.name for tool in tools] == ["generate_pptx"]


# ---------------------------------------------------------------------------
# validate_config — save-time state machine
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_validate_config_without_upload_passes_through() -> None:
    cap = PptxGeneratorCapability()
    save_ctx = SaveContext(identity=_identity(), services=RuntimeServices())
    stored = await cap.validate_config(PptxGeneratorConfig(), {}, save_ctx)
    assert stored.template_key is None


@pytest.mark.asyncio
async def test_validate_config_with_upload_stores_template_and_persists_key() -> None:
    assets = _FakeAgentAssets()
    save_ctx = SaveContext(
        identity=_identity(), services=RuntimeServices(agent_assets=assets)
    )
    cap = PptxGeneratorCapability()
    uploads = {
        BRAND_TEMPLATE_SLOT: [
            UploadedFile(filename="brand.pptx", content=_blank_pptx_bytes())
        ]
    }
    stored = await cap.validate_config(PptxGeneratorConfig(), uploads, save_ctx)

    assert stored.template_key == BRAND_TEMPLATE_KEY
    assert assets.stored[0][0] == BRAND_TEMPLATE_KEY


@pytest.mark.asyncio
async def test_validate_config_rejects_unreadable_upload() -> None:
    save_ctx = SaveContext(
        identity=_identity(), services=RuntimeServices(agent_assets=_FakeAgentAssets())
    )
    cap = PptxGeneratorCapability()
    uploads = {
        BRAND_TEMPLATE_SLOT: [
            UploadedFile(filename="brand.pptx", content=b"not a pptx")
        ]
    }
    with pytest.raises(ValueError, match="could not be read"):
        await cap.validate_config(PptxGeneratorConfig(), uploads, save_ctx)


@pytest.mark.asyncio
async def test_validate_config_rejects_oversized_upload() -> None:
    save_ctx = SaveContext(
        identity=_identity(), services=RuntimeServices(agent_assets=_FakeAgentAssets())
    )
    cap = PptxGeneratorCapability()
    huge = b"x" * (MAX_TEMPLATE_UPLOAD_BYTES + 1)
    uploads = {BRAND_TEMPLATE_SLOT: [UploadedFile(filename="brand.pptx", content=huge)]}
    with pytest.raises(ValueError, match="exceeds"):
        await cap.validate_config(PptxGeneratorConfig(), uploads, save_ctx)


@pytest.mark.asyncio
async def test_validate_config_fails_loud_without_agent_assets_port() -> None:
    save_ctx = SaveContext(identity=_identity(), services=RuntimeServices())
    cap = PptxGeneratorCapability()
    uploads = {
        BRAND_TEMPLATE_SLOT: [
            UploadedFile(filename="brand.pptx", content=_blank_pptx_bytes())
        ]
    }
    with pytest.raises(RuntimeError, match="agent_assets"):
        await cap.validate_config(PptxGeneratorConfig(), uploads, save_ctx)


def _pptx_bytes_with_titled_slide(title: str) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.title
    assert title_shape is not None
    title_shape.text_frame.text = title
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_validate_config_with_slideless_template_leaves_catalog_unset() -> None:
    # A blank Presentation() has masters/layouts but zero slides — the
    # common "theme-only" brand template (RFC §4.2.1) — so there's nothing
    # to list (RFC §13).
    save_ctx = SaveContext(
        identity=_identity(), services=RuntimeServices(agent_assets=_FakeAgentAssets())
    )
    cap = PptxGeneratorCapability()
    uploads = {
        BRAND_TEMPLATE_SLOT: [
            UploadedFile(filename="brand.pptx", content=_blank_pptx_bytes())
        ]
    }
    stored = await cap.validate_config(PptxGeneratorConfig(), uploads, save_ctx)
    assert stored.template_slide_catalog is None


@pytest.mark.asyncio
async def test_validate_config_computes_template_slide_catalog() -> None:
    save_ctx = SaveContext(
        identity=_identity(), services=RuntimeServices(agent_assets=_FakeAgentAssets())
    )
    cap = PptxGeneratorCapability()
    uploads = {
        BRAND_TEMPLATE_SLOT: [
            UploadedFile(
                filename="brand.pptx",
                content=_pptx_bytes_with_titled_slide("Agenda"),
            )
        ]
    }
    stored = await cap.validate_config(PptxGeneratorConfig(), uploads, save_ctx)
    assert stored.template_slide_catalog == '  0: "Agenda"'


# ---------------------------------------------------------------------------
# tools() — description reflects the configured template's slide catalog
# ---------------------------------------------------------------------------


def test_tools_description_omits_template_section_when_no_catalog() -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices())
    description = cap.tools(ctx)[0].description
    assert "TEMPLATE SLIDES" not in description


def test_tools_description_includes_template_catalog_when_configured() -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(),
        config={
            "template_key": BRAND_TEMPLATE_KEY,
            "template_slide_catalog": '  0: "Agenda"\n  1: "Thank you"',
        },
    )
    description = cap.tools(ctx)[0].description
    assert "TEMPLATE SLIDES" in description
    assert "template_slide_index" in description
    assert '  0: "Agenda"' in description
    assert '  1: "Thank you"' in description


# ---------------------------------------------------------------------------
# generate_pptx — happy paths
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_pptx_happy_path_with_preview(monkeypatch) -> None:
    _enable_preview(monkeypatch)
    workspace = _FakeWorkspaceFs()
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices(workspace_fs=workspace))
    tool = cap.tools(ctx)[0]

    content, artifact = await tool.coroutine(
        title="My Deck", slides=[SlideSpec(kind="title", title="Cover")]
    )

    assert artifact.is_error is False
    assert len(artifact.ui_parts) == 1
    part = artifact.ui_parts[0]
    assert isinstance(part, PptxGeneratedPart)
    assert part.type == "pptx_generated"
    assert "preview opens beside the chat" in content
    assert workspace.writes[0][0] == "outputs/s-1/My Deck.pptx"


@pytest.mark.asyncio
async def test_generate_pptx_degrades_to_download_when_pdf_conversion_fails(
    no_pdf,
) -> None:
    workspace = _FakeWorkspaceFs()
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices(workspace_fs=workspace))
    tool = cap.tools(ctx)[0]

    content, artifact = await tool.coroutine(
        title="My Deck", slides=[SlideSpec(kind="title", title="Cover")]
    )

    assert artifact.is_error is False
    assert len(artifact.ui_parts) == 1
    assert artifact.ui_parts[0].kind == LinkKind.download
    assert "could not be generated" in content


@pytest.mark.asyncio
async def test_generate_pptx_uses_configured_brand_template(no_pdf) -> None:
    template_bytes = _blank_pptx_bytes()
    assets = _FakeAgentAssets({BRAND_TEMPLATE_KEY: template_bytes})
    workspace = _FakeWorkspaceFs()
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(workspace_fs=workspace, agent_assets=assets),
        config={"template_key": BRAND_TEMPLATE_KEY},
    )
    tool = cap.tools(ctx)[0]

    _content, artifact = await tool.coroutine(
        title="Branded", slides=[SlideSpec(kind="title", title="Cover")]
    )

    assert artifact.is_error is False
    assert workspace.writes


@pytest.mark.asyncio
async def test_generate_pptx_with_image_reference(no_pdf) -> None:
    workspace = _FakeWorkspaceFs()
    document_content = _FakeDocumentContent({"doc-1": _png_bytes()})
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(
            workspace_fs=workspace, document_content=document_content
        ),
    )
    tool = cap.tools(ctx)[0]

    _content, artifact = await tool.coroutine(
        title="Deck", slides=[SlideSpec(kind="image", title="Logo", image_ref="doc-1")]
    )

    assert artifact.is_error is False


@pytest.mark.asyncio
async def test_generate_pptx_fetches_column_icons(no_pdf) -> None:
    document_content = _FakeDocumentContent({"icon-1": _png_bytes()})
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(
            workspace_fs=_FakeWorkspaceFs(), document_content=document_content
        ),
    )
    tool = cap.tools(ctx)[0]
    columns = [ColumnSpec(heading="A", icon_ref="icon-1"), ColumnSpec(heading="B")]

    _content, artifact = await cast(Any, tool).coroutine(
        title="Deck", slides=[SlideSpec(kind="columns", title="T", columns=columns)]
    )

    assert artifact.is_error is False


# ---------------------------------------------------------------------------
# generate_pptx — error paths
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_pptx_rejects_empty_slides(no_pdf) -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices(workspace_fs=_FakeWorkspaceFs()))
    tool = cap.tools(ctx)[0]

    content, artifact = await tool.coroutine(title="Empty", slides=[])

    assert artifact.is_error is True
    assert "NOT generated" in content


@pytest.mark.asyncio
async def test_generate_pptx_rejects_over_max_slides(no_pdf) -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices(workspace_fs=_FakeWorkspaceFs()))
    tool = cap.tools(ctx)[0]
    slides = [
        SlideSpec(kind="title", title=f"Slide {i}")
        for i in range(deck_builder.MAX_SLIDES + 1)
    ]

    content, artifact = await tool.coroutine(title="Too Big", slides=slides)

    assert artifact.is_error is True
    assert "Too many slides" in content


@pytest.mark.asyncio
async def test_generate_pptx_reports_unknown_image_ref(no_pdf) -> None:
    document_content = _FakeDocumentContent({})
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(
            workspace_fs=_FakeWorkspaceFs(), document_content=document_content
        ),
    )
    tool = cap.tools(ctx)[0]

    content, artifact = await tool.coroutine(
        title="Deck",
        slides=[SlideSpec(kind="image", title="Logo", image_ref="missing-doc")],
    )

    assert artifact.is_error is True
    assert "image_ref" in content


@pytest.mark.asyncio
async def test_generate_pptx_fails_loud_without_document_content_port(no_pdf) -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices(workspace_fs=_FakeWorkspaceFs()))
    tool = cap.tools(ctx)[0]

    with pytest.raises(RuntimeError, match="document_content"):
        await tool.coroutine(
            title="Deck",
            slides=[SlideSpec(kind="image", title="Logo", image_ref="doc-1")],
        )


@pytest.mark.asyncio
async def test_generate_pptx_reports_oversized_table(no_pdf) -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices(workspace_fs=_FakeWorkspaceFs()))
    tool = cap.tools(ctx)[0]
    rows = [["x"] for _ in range(deck_builder.MAX_TABLE_ROWS)]
    table = TableSpec(headers=["Col"], rows=rows)

    content, artifact = await tool.coroutine(
        title="Deck", slides=[SlideSpec(kind="table", title="Big", table=table)]
    )

    assert artifact.is_error is True
    assert "too large" in content
    assert not artifact.ui_parts


@pytest.mark.asyncio
async def test_generate_pptx_fails_loud_without_workspace_fs_port(no_pdf) -> None:
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices())
    tool = cap.tools(ctx)[0]

    with pytest.raises(RuntimeError, match="workspace_fs"):
        await tool.coroutine(
            title="Deck", slides=[SlideSpec(kind="title", title="Cover")]
        )


@pytest.mark.asyncio
async def test_generate_pptx_folds_slide_background_image_ref_into_prefetch(
    no_pdf,
) -> None:
    workspace = _FakeWorkspaceFs()
    document_content = _FakeDocumentContent({"bg-1": _png_bytes()})
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(
            workspace_fs=workspace, document_content=document_content
        ),
    )
    tool = cap.tools(ctx)[0]

    _content, artifact = await tool.coroutine(
        title="Deck",
        slides=[
            SlideSpec(
                kind="title",
                title="Cover",
                background=BackgroundSpec(kind="image", image_ref="bg-1"),
            )
        ],
    )

    assert artifact.is_error is False


@pytest.mark.asyncio
async def test_generate_pptx_folds_theme_background_image_ref_into_prefetch(
    no_pdf,
) -> None:
    workspace = _FakeWorkspaceFs()
    document_content = _FakeDocumentContent({"bg-1": _png_bytes()})
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(
            workspace_fs=workspace, document_content=document_content
        ),
    )
    tool = cap.tools(ctx)[0]

    _content, artifact = await tool.coroutine(
        title="Deck",
        slides=[SlideSpec(kind="title", title="Cover")],
        theme=ThemeSpec(background=BackgroundSpec(kind="image", image_ref="bg-1")),
    )

    assert artifact.is_error is False


@pytest.mark.asyncio
async def test_generate_pptx_reports_unresolved_background_image_ref(no_pdf) -> None:
    document_content = _FakeDocumentContent({})
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(
            workspace_fs=_FakeWorkspaceFs(), document_content=document_content
        ),
    )
    tool = cap.tools(ctx)[0]

    content, artifact = await tool.coroutine(
        title="Deck",
        slides=[SlideSpec(kind="title", title="Cover")],
        theme=ThemeSpec(
            background=BackgroundSpec(kind="image", image_ref="missing-bg")
        ),
    )

    assert artifact.is_error is True
    assert "background.image_ref" in content


@pytest.mark.asyncio
async def test_busy_pod_rejects_generation_without_writing(no_pdf, monkeypatch) -> None:
    monkeypatch.setattr(capability_mod, "acquire_heavy_job_slot", lambda: False)
    workspace = _FakeWorkspaceFs()
    cap = PptxGeneratorCapability()
    ctx = _ctx(cap, services=RuntimeServices(workspace_fs=workspace))
    tool = cap.tools(ctx)[0]

    content, artifact = await tool.coroutine(
        title="Deck", slides=[SlideSpec(kind="title", title="Cover")]
    )

    assert artifact.is_error is True
    assert "busy" in content.lower()
    assert workspace.writes == []


@pytest.mark.asyncio
async def test_busy_pod_rejects_before_fetching_images(no_pdf, monkeypatch) -> None:
    # Regression guard: the heavy-job slot is now reserved BEFORE the image
    # prefetch (and the brand-template fetch), not just before the build —
    # an already-overloaded pod must reject fast without spending I/O.
    class _CountingDocumentContent(_FakeDocumentContent):
        def __init__(self, content_by_uid: dict[str, bytes]) -> None:
            super().__init__(content_by_uid)
            self.fetch_calls = 0

        async def fetch_raw(self, document_uid: str):
            self.fetch_calls += 1
            return await super().fetch_raw(document_uid)

    monkeypatch.setattr(capability_mod, "acquire_heavy_job_slot", lambda: False)
    document_content = _CountingDocumentContent({"doc-1": _png_bytes()})
    workspace = _FakeWorkspaceFs()
    cap = PptxGeneratorCapability()
    ctx = _ctx(
        cap,
        services=RuntimeServices(
            workspace_fs=workspace, document_content=document_content
        ),
    )
    tool = cap.tools(ctx)[0]

    content, artifact = await cast(Any, tool).coroutine(
        title="Deck",
        slides=[SlideSpec(kind="image", title="Cover", image_ref="doc-1")],
    )

    assert artifact.is_error is True
    assert "busy" in content.lower()
    assert document_content.fetch_calls == 0
    assert workspace.writes == []
